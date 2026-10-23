"""Drive the existing frozen selftest across a real TerminateProcess boundary."""
import argparse
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time

p = argparse.ArgumentParser()
p.add_argument('--repo', required=True)
p.add_argument('--exe', required=True)
p.add_argument('--out', required=True)
p.add_argument('--paragraphs', type=int, default=4000)
a = p.parse_args()
repo, out = Path(a.repo).resolve(), Path(a.out).resolve()
out.mkdir(parents=True, exist_ok=False)
sys.path[:0] = [str(repo), str(repo / 'scripts')]
from qt_desktop_probe import StubEndpoint
from qt_crash_recovery_probe import _durable_state
from folith.qt.selftest import make_docx

fixture = out / 'recovery.docx'
make_docx(fixture, [f'第 {n} 段，术语 MT。' for n in range(a.paragraphs)])
stub = StubEndpoint(out, delay=0.08)
env = dict(os.environ, QT_QPA_PLATFORM='windows',
           FOLITH_QT_NONINTERACTIVE='1', FOLITH_EXIT_FORCE_PROMPT='0',
           FOLITH_QT_SETTINGS_DIR=str(out / 'settings'),
           FOLITH_QT_SELFTEST_PROVIDER='自定义中转站',
           FOLITH_QT_SELFTEST_MODEL='stub-model',
           FOLITH_QT_SELFTEST_API_KEY='stub-key',
           FOLITH_QT_SELFTEST_TIMEOUT='2400')
report = {'fixture_kind': 'generated deterministic Chinese paragraphs',
          'requested_paragraphs': a.paragraphs, 'started_at': time.time(),
          'exe': str(Path(a.exe).resolve()), 'qt_platform_requested': 'windows'}
child = None
try:
    env['FOLITH_QT_SELFTEST_BASE_URL'] = stub.start()
    command = [a.exe, '--selftest', '--selftest-source', str(fixture),
               '--data-dir', str(out / 'data'), '--log-dir', str(out / 'logs'),
               '--result', str(out / 'flow.json')]
    handle = (out / 'first-process.log').open('w', encoding='utf-8')
    child = subprocess.Popen(command, env=env, stdout=handle, stderr=subprocess.STDOUT)
    deadline = time.monotonic() + 600
    while time.monotonic() < deadline:
        if child.poll() is not None:
            raise AssertionError(f'first frozen process exited early: {child.returncode}')
        candidates = list((out / 'data').glob('*/state.json'))
        if len(candidates) == 1:
            job = candidates[0].parent.name
            try:
                state, replay = _durable_state(out / 'data', job)
                runtime = json.loads((candidates[0].parent / 'runtime_state.json').read_text(encoding='utf-8'))
                if len(state.get('pairs') or []) >= 24 and runtime.get('status') == 'running':
                    break
            except (OSError, ValueError, AssertionError):
                pass
        time.sleep(0.5)
    else:
        raise AssertionError('no active durable checkpoint')
    child.kill()
    child.wait(timeout=30)
    handle.close()
    shutil.copytree(out / 'data', out / 'after-kill')
    expected, replay = _durable_state(out / 'data', job)
    old_pairs = expected['pairs']
    report['killed'] = {'exit_code': child.returncode, 'pairs': len(old_pairs),
                        'segments': len(expected['paras']), 'replay': replay}
    events = out / 'data' / job / 'events.jsonl'
    event_offset = events.stat().st_size
    with (out / 'second-process.log').open('w', encoding='utf-8') as handle:
        child = subprocess.Popen(command, env=env, stdout=handle, stderr=subprocess.STDOUT)
        code = child.wait(timeout=3600)
    flow = json.loads((out / 'flow.json').read_text(encoding='utf-8'))
    assert code == 0 and flow['ok'], flow.get('error') or flow.get('failures')
    final, replay = _durable_state(out / 'data', job)
    assert len(final['pairs']) == len(final['paras']) == a.paragraphs
    # The ordinary flow deliberately human-edits row 1 and reviews all rows.
    # Its target text must stay identical to the committed target, while those
    # audit fields are expected to change through the product's real actions.
    for before, after in zip(old_pairs, final['pairs']):
        for key in ('source', 'target', 'initial_target'):
            assert before.get(key) == after.get(key), key
        if before.get('segment_uid'):
            assert before['segment_uid'] == after.get('segment_uid')
    with events.open('rb') as handle:
        handle.seek(event_offset)
        offsets = [e['offset'] for line in handle if line.strip()
                   for e in [json.loads(line)] if e.get('phase') == 'generation_started']
    assert offsets and min(offsets) >= len(old_pairs), offsets[:5]
    from docx import Document
    exported = next(Path(f) for f in flow['saved_files'] if f.endswith('_translation.docx'))
    text = '\n'.join(p.text for p in Document(exported).paragraphs)
    position = 0
    for pair in final['pairs']:
        target = str(pair.get('target') or '')
        assert target
        found = text.find(target, position)
        assert found >= 0, 'export target missing or reordered'
        position = found + len(target)
    report.update(ok=True, recovered_pairs=len(final['pairs']),
                  committed_targets_preserved=True, committed_batches_not_regenerated=True,
                  export_targets_in_source_order=True, frozen_identity=flow['identity'],
                  reviewed=flow['review'], delivery=flow['delivery'],
                  heartbeat=flow.get('heartbeat'), replay_after_resume=replay)
except BaseException as exc:
    import traceback
    report.update(ok=False, error=f'{type(exc).__name__}: {exc}', traceback=traceback.format_exc())
finally:
    if child is not None and child.poll() is None:
        child.kill()
        child.wait(timeout=30)
    stub.stop()
    report['duration_s'] = round(time.time() - report['started_at'], 2)
    (out / 'recovery-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False, indent=2))
sys.exit(0 if report['ok'] else 1)
