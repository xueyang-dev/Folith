"""Folith 核心逻辑层（与 Streamlit UI 解耦，便于测试）。

职责：大模型路由、文档清洗、术语抽取、双语翻译、报告生成、任务进度持久化。
"""
import copy
import hashlib
import io
import json
import csv
import inspect
from functools import wraps
from contextlib import contextmanager
from concurrent.futures import ThreadPoolExecutor
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
import tempfile
import unicodedata
import time
import traceback
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from html import escape as html_escape
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence, Set, Tuple
from urllib.parse import urlparse

import fitz  # PyMuPDF
import httpx
import pandas as pd
from defusedxml import ElementTree as ET
from docx import Document
from google import genai
from openai import OpenAI, DefaultHttpxClient

from folith import models as _models
from folith import state_migration as _state_migration
from folith import context as _context
from folith import knowledge as _knowledge
from folith import repair as _repair
from folith import translation_evidence as _translation_evidence
from folith import checkpoint as _checkpoint
from folith import snapshots as _snapshots
from folith import entity_registry as _entity_registry
from folith import dynamic_context as _dynamic_context
from folith import model_roles as _model_roles
from folith import pdf_ingestion as _pdf_ingestion
from folith import source_cleanup as _source_cleanup
from folith import segmentation as _segmentation
from folith import source_quality as _source_quality
from folith import document_ir as _document_ir
from folith import translation_protocol as _translation_protocol
from folith import persistence_profiler as _persistence_profiler
from folith import critical_path as _critical_path
from folith.storage import job_repository as _job_repo
from folith.storage import state_journal as _state_journal
from folith import translation_target as _translation_target
from folith import finalization as _finalization
from folith import rendered_qa as _rendered_qa
from folith import project as _project
from folith import translation_memory as _tm_scope
from folith import translation_cache as _translation_cache
from folith import knowledge_cache as _knowledge_cache_module
from folith import usage as _usage
from folith import llm_transport as _llm_transport
from folith import request_control as _request_control
from folith import workflow_errors as _workflow_errors
from folith import job_lock as _job_lock
from folith import store_lock as _store_lock
from folith import review_scheduler as _review_scheduler
from folith import review_gate as _review_gate
from folith import ocr_backends as _ocr_backends
from folith import ocr_language as _ocr_language
from folith import ocr_models as _ocr_models
from folith import ocr_online as _ocr_online
from folith import text_layer_quality as _text_layer_quality
from folith.ocr_cache import (
    OCRPageCache as _OCRPageCache,
    OCRPageIdentity as _OCRPageIdentity,
    OCR_PREPROCESSING_VERSION as _OCR_PREPROCESSING_VERSION,
    read_legacy_checkpoint_pages as _read_legacy_ocr_pages,
)
from folith.performance import PipelineProfiler
from folith.textual import has_textual_content, normalize_language, text_direction
from folith.trace import trace, TraceLevel, RunContext

# ================= 常量 =================
# 任务进度与过程文件的本地存储目录（已加入 .gitignore）
OUTPUT_DIR = Path(os.environ.get("FOLITH_OUTPUT_DIR") or os.environ.get("FOLIOTHREAD_OUTPUT_DIR", "outputs"))

DELIVERY_CONFIG_DEFAULTS = {
    "enable_annotate": False,
    "enable_report": False,
    "deliver_plain_docx": True,
    "deliver_bilingual_docx": True,
    "deliver_pdf": False,
    "deliver_original_pdf": False,
    "deliver_epub": False,
    "deliver_terms_xlsx": True,
    "deliver_tbx": False,
    "deliver_tmx": False,
    "deliver_jsonl": False,
    "deliver_evidence": True,
    "deliver_cases": False,
    "deliver_academic_workspace": False,
    "deliver_review_report": False,
}


def default_delivery_config():
    return dict(DELIVERY_CONFIG_DEFAULTS)


def normalize_delivery_config(config, *, enable_report=None, enable_annotate=None):
    normalized = default_delivery_config()
    if isinstance(config, dict):
        normalized.update({key: bool(config[key]) for key in normalized if key in config})
    if enable_report is not None:
        normalized["enable_report"] = bool(enable_report)
    if enable_annotate is not None:
        normalized["enable_annotate"] = bool(enable_annotate)
    return normalized

# 提供商注册表：kind=openai_compat 走 OpenAI SDK 兼容接口（官方与中转站通用）。
# base_url 为 None 表示使用官方 SDK 默认路由；custom_base_url 表示地址由用户在
# UI 中填写（通用中转站）。OpenCode Go 只列出官方标记为
# /chat/completions 的模型；/messages 与 /responses 模型不走此路由。
PROVIDERS = {
    "OpenCode Go": {
        "kind": "openai_compat",
        "base_url": "https://opencode.ai/zen/go/v1",
        "proxy_bypass": True,
        "default_model": "glm-5.2",
        "models": [
            "glm-5.2", "glm-5.1",
            "deepseek-v4.1-flash", "deepseek-v4-pro", "deepseek-v4-flash",
            "kimi-k3", "kimi-k2.7-code", "kimi-k2.6",
            "mimo-v2.5-pro", "mimo-v2.5", "hy3", "grok-4.5",
        ],
    },
    "DeepSeek": {
        "kind": "openai_compat",
        "base_url": "https://api.deepseek.com",
        "default_model": "deepseek-flash",
        "models": ["deepseek-flash", "deepseek-v4-flash", "deepseek-v4-pro"],
    },
    "OpenAI": {
        "kind": "openai",
        # Keep the stable general models and expose current reasoning families
        # so the conditional reasoning-effort control has a real target.
        "models": [
            "gpt-4o-mini", "gpt-4o", "gpt-4.1-mini", "gpt-4.1",
            "gpt-5.2", "gpt-5.1", "gpt-5", "gpt-5-mini", "o3", "o4-mini",
        ],
        "capabilities": {
            "supports_json_schema": True,
            "supports_json_object": True,
            "supports_response_format": True,
        },
    },
    "Gemini": {
        "kind": "gemini",
        "default_model": "gemini-3.6-flash",
        "models": ["gemini-3.6-flash", "gemini-3.5-flash",
                   "gemini-3.5-flash-lite", "gemini-2.5-flash", "gemini-2.5-pro"],
    },
    "OpenRouter": {
        "kind": "openai_compat",
        "base_url": "https://openrouter.ai/api/v1",
        "models": [],
        "model_hint": "如 anthropic/claude-sonnet-4、openai/gpt-5、google/gemini-3-flash",
    },
    "SiliconFlow": {
        "kind": "openai_compat",
        "base_url": "https://api.siliconflow.cn/v1",
        "models": [],
        "model_hint": "如 Qwen/Qwen3-235B-A22B、deepseek-ai/DeepSeek-V3.2",
    },
    "Moonshot (Kimi)": {
        "kind": "openai_compat",
        "base_url": "https://api.moonshot.cn/v1",
        "models": [],
        "model_hint": "如 kimi-k2.5、moonshot-v1-8k",
    },
    "Zhipu (GLM)": {
        "kind": "openai_compat",
        "base_url": "https://open.bigmodel.cn/api/paas/v4",
        "models": [],
        "model_hint": "如 glm-4.5、glm-5",
    },
    "Qwen (DashScope)": {
        "kind": "openai_compat",
        "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1",
        "models": [],
        "model_hint": "如 qwen-max、qwen3-235b-a22b",
    },
    "自定义中转站": {
        "kind": "openai_compat",
        "base_url": None,
        "custom_base_url": True,
        "models": [],
        "model_hint": "填写中转站提供的模型名（OpenAI /chat/completions 兼容）",
    },
}

# 兼容旧引用：仅保留 模型名 -> 列表 的视图
MODELS = {name: cfg["models"] for name, cfg in PROVIDERS.items()}

# 会话线程级中转地址：自定义中转站的 base_url 由 UI 写入，call_llm 自动优先使用。
# 每个 Streamlit 会话线程独立，多设备同时使用不会互相串扰。
_LLM_CTX = threading.local()

# Runtime status is deliberately file-backed: the UI can poll it while the
# worker is inside a blocking provider request, and a browser refresh does not
# erase the last known activity.
_RUNTIME_CTX = threading.local()
_RUNTIME_LOCK = threading.RLock()
_RUNTIME_WORKERS = {}
_RUNTIME_WORKERS_LOCK = threading.RLock()



_REASONING_EFFORT_VALUES = ("none", "minimal", "low", "medium", "high", "xhigh", "max")



def _runtime_job_id():
    return getattr(_RUNTIME_CTX, "job_id", None)



# ================= 权威状态的跨进程事务锁 =================
# `state.json` / `project.json` 的原子替换只保证"文件不会被写坏"，不保证
# "读 → 改 → 写"整体是一个事务：两个进程可以各自读到同一版本，然后互相覆盖
# （丢失更新）。这里提供唯一的、按存储对象划分的跨进程事务锁；锁文件命名、
# 语义、锁序与超时行为集中在 `folith/store_lock.py` 里说明。
#
# 锁序只有一种：**state 锁 → project 锁**。worker 不参与该锁（它的独占性由
# `.worker.lock` 提供），因此不会出现反向的 project → state 嵌套。



def _runtime_cancel_requested(job_id):
    _workflow_errors.check(job_id)
    return bool((load_runtime_state(job_id) or {}).get("cancel_requested"))


def _runtime_pause_checkpoint(job_id):
    """Pause cooperatively between operations; keep in-memory results and lease."""
    if not job_id or not load_runtime_state(job_id).get("pause_requested"):
        return
    previous = load_runtime_state(job_id)
    update_runtime_state(job_id, status="paused", phase="paused",
                         event="任务已暂停，可继续处理", event_name="job_paused",
                         event_visibility="user", event_category="lifecycle")
    while load_runtime_state(job_id).get("pause_requested"):
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        time.sleep(0.25)
    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")
    update_runtime_state(job_id, status="running", phase=previous.get("phase") or "running",
                         phase_label=previous.get("phase_label") or "正在执行",
                         event="任务已继续", event_name="job_continued",
                         event_visibility="user", event_category="lifecycle")


def request_job_pause(job_id):
    if not is_job_worker_alive(job_id):
        return False
    update_runtime_state(job_id, pause_requested=True, status="pausing",
                         event="已请求暂停，当前操作结束后暂停", event_name="pause_requested",
                         event_visibility="user", event_category="lifecycle")
    return True


def continue_job(job_id):
    if not is_job_worker_alive(job_id):
        return False
    update_runtime_state(job_id, pause_requested=False, status="running", phase="running",
                         event="正在继续任务", event_name="continue_requested",
                         event_visibility="user", event_category="lifecycle")
    return True


# ================= 基础工具函数 =================



def _llm_cancel_check():
    """Cancellation probe bound to the job currently running on this thread."""
    job_id = _runtime_job_id()
    if not job_id:
        return None
    return lambda: _runtime_cancel_requested(job_id)



# Per-thread provider accounting. One translation batch runs entirely on one
# worker thread, so a before/after delta gives that batch's real API latency,
# request count and failure count without any cross-thread bookkeeping.
_LLM_LATENCY_CTX = threading.local()



# ================= PDF 确定性段落提取 =================
# 经验（来自 localize-anything 与全书实测）：分段/清洗必须确定性，不能交给 LLM。
# 旧流程把 ~2500 字符的任意文本块交给模型"清洗"，导致两类系统性缺陷：
#   1. 块边界落在句中 -> 句子被拦腰截断（"…pecking at" / "crumbs and bones…"）；
#   2. 模型自由裁量 -> 对白、引语被随意拆分或合并，分段结果不可复现。
# 新流程直接读 PDF 版面：块(block)->行(line)->首行缩进判定段落，连字符修复、
# 跨页段落合并、页眉页脚/页码剔除全部确定完成。

# 句末终结符（用于判断段落是否未完结、需要与下一段合并）
_SENTENCE_TERMINAL = set('.!?"”’…:;)')

# 纯装饰符号行（章节分隔花饰等），无字母/数字，不是正文
_ORNAMENT_RE = re.compile(r"^[\s*•·▪◦‣❦❧—–\-]{1,12}$")

# 常见缩写（句点不计入句界）
_ABBREV_RE = re.compile(
    r"\b(?:Lt|Col|Gen|Maj|Capt|Sgt|Brig|Mr|Mrs|Ms|Dr|St|No|Vol|pp|"
    r"e\.g|i\.e|vs|etc|a\.m|p\.m|U\.S|A\.F|B\.C|A\.D)\.", re.IGNORECASE)



def extract_pdf_paragraphs(file_bytes):
    """Compatibility API backed by the layout-aware PDF ingestion module."""
    return _pdf_ingestion.extract_pdf_paragraphs(file_bytes)


def _find_tesseract():
    """Find Tesseract even when the app was launched from Finder/Explorer.

    GUI launchers often receive a smaller PATH than an interactive shell.  The
    explicit environment variable is useful for portable installs and takes
    precedence over the usual PATH/Homebrew/Windows locations.
    """
    configured = os.environ.get("FOLIOTHREAD_TESSERACT_CMD", "").strip()
    candidates = [configured, shutil.which("tesseract")]
    if os.name == "nt":
        candidates.extend([
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        ])
    else:
        candidates.extend([
            "/opt/homebrew/bin/tesseract",
            "/usr/local/bin/tesseract",
            "/usr/bin/tesseract",
        ])
    for candidate in candidates:
        if not candidate:
            continue
        path = Path(candidate)
        if path.is_file() and os.access(path, os.X_OK):
            return str(path)
    return ""


# LibreOffice stays optional: it is only used for the render QA.  GUI launchers
# receive a small PATH, so the standard install locations are checked explicitly
# (roadmap §3.4/§14: "补 Applications / Windows 安装位置发现并验证桌面 PATH").
_SOFFICE_CANDIDATES = (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/libreoffice",
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/usr/bin/soffice",
    "/usr/bin/libreoffice",
    "/usr/local/bin/soffice",
    "/snap/bin/libreoffice",
)


def _find_soffice():
    """Find LibreOffice even when the app was launched from Finder/Explorer.

    Missing LibreOffice must stay an accurate NOT_RUN; this function never
    invents an engine, it only reports what is actually installed.
    """
    configured = os.environ.get("FOLIOTHREAD_SOFFICE_CMD", "").strip()
    candidates = [configured, shutil.which("soffice"), shutil.which("libreoffice")]
    candidates.extend(_SOFFICE_CANDIDATES)
    for candidate in candidates:
        if not candidate:
            continue
        path = Path(candidate)
        if path.is_file() and os.access(path, os.X_OK):
            return str(path)
    return ""


def _tesseract_languages(command):
    """Return installed OCR languages, ignoring Tesseract control models."""
    try:
        proc = subprocess.run(
            [command, "--list-langs"], capture_output=True, text=True,
            encoding="utf-8", errors="replace", timeout=15,
        )
    except Exception:  # noqa: BLE001 - diagnostics must not block OCR
        return []
    if proc.returncode != 0:
        return []
    languages = []
    for line in (proc.stdout or "").splitlines():
        value = line.strip()
        if value and re.fullmatch(r"[A-Za-z0-9_]+", value):
            languages.append(value)
    return languages



def _runtime_int_option(explicit, saved, env_name, default, *, minimum=1, maximum=32):
    value = explicit
    if value is None:
        value = saved
    if value is None:
        value = os.environ.get(env_name, "")
    try:
        value = int(value)
    except (TypeError, ValueError):
        value = int(default)
    return max(minimum, min(value, maximum))


def _runtime_float_option(explicit, saved, env_name, default, *, minimum=0.0, maximum=300.0):
    value = explicit
    if value is None:
        value = saved
    if value is None:
        value = os.environ.get(env_name, "")
    try:
        value = float(value)
    except (TypeError, ValueError):
        value = float(default)
    return max(minimum, min(value, maximum))



def installed_ocr_languages():
    """Installed Tesseract packs, for the visible language-pack validation."""
    command = _find_tesseract()
    return _tesseract_languages(command) if command else []



# Stable internal entry points used by existing tests/benchmarks.
from folith.ocr_tesseract import (
    _ocr_join_lines, _order_line_words, _iter_tesseract_tsv_rows,
    _parse_tesseract_tsv, _ocr_page_with_tesseract,
)
from folith.engine.runtime_telemetry import (  # noqa: F401  (re-export: core keeps its public surface)
    _RuntimeTelemetryCoalescer,
    _runtime_telemetry_coalescer,
    _runtime_telemetry_request_failed,
    flush_runtime_telemetry,
    reset_runtime_persistence_stats,
    reset_runtime_telemetry,
    runtime_persistence_stats,
    runtime_telemetry_stats,
)
from folith.engine.ocr_config import (  # noqa: F401  (re-export: core keeps its public surface)
    DEFAULT_OCR_MODE,
    OCR_ADAPTIVE_HIGH_DPI,
    OCR_ADAPTIVE_LOW_DPI,
    OCR_ADAPTIVE_MIN_CONFIDENCE,
    OCR_ADAPTIVE_MIN_TEXT_DENSITY,
    OCR_CHECKPOINT_VERSION,
    OCR_CONFIDENCE_THRESHOLD_DEFAULT,
    OCR_DPI_BASELINE,
    OCR_MODES,
    OCR_ONLINE_CONFIG_FILE,
    OCR_ONLINE_CREDENTIAL_LABELS,
    OCR_ONLINE_CREDENTIAL_SOURCES,
    OCR_ONLINE_DEFAULT_MAX_TOKENS,
    OCR_ONLINE_MAX_MAX_TOKENS,
    OCR_ONLINE_MIN_MAX_TOKENS,
    OCR_ONLINE_REASONING_MODES,
    OCR_PROFILE_CACHE_DIRNAME,
    OCR_QUEUE_SIZE_DEFAULT,
    OCR_WORKERS_DEFAULT,
    _OCR_BACKEND_VERSION_CACHE,
    _OCR_MODE_ALIASES,
    _OCR_PAGE_NUMBER_RE,
    _effective_ocr_config,
    _effective_ocr_source,
    _legacy_select_tesseract_language,
    _ocr_cache_dir,
    _ocr_checkpoint_path,
    _ocr_env_flag,
    _ocr_pdf_text,
    _ocr_queue_size_default,
    _ocr_runtime_metadata,
    _ocr_text_to_paragraphs,
    _ocr_workers_default,
    _resolve_ocr_language_config,
    _resolve_online_ocr_base_url,
    _select_tesseract_language,
    normalize_ocr_mode,
    ocr_backend_report,
    ocr_engine_status,
    ocr_language_options,
    ocr_language_status,
    ocr_mode_options,
    ocr_online_config_path,
    ocr_online_status,
    ocr_paddle_model_status,
    ocr_profile_cache_dir,
    ocr_quality_dpi,
    ocr_quality_options,
    ocr_source_label,
    ocr_source_options,
    paddle_model_options,
    resolve_ocr_online_config,
    save_ocr_online_config,
)
from folith.engine.document_export import (  # noqa: F401  (re-export: core keeps its public surface)
    _add_formatted_runs,
    _add_markdown_table,
    _apply_doc_fonts,
    _apply_paragraph_direction,
    _apply_run_fonts,
    _colored_cell,
    _docx_part_texts,
    _docx_unsupported_parts,
    _expand_span_end,
    _find_span,
    _is_latin_word_char,
    _markdown_table_row,
    dict_to_excel,
    findings_report_md,
    glossary_to_excel,
    markdown_to_word,
    pairs_to_word,
    paragraphs_to_word,
    translations_to_pdf,
    translations_to_word,
)
from folith.engine.llm_provider import (  # noqa: F401  (re-export: core keeps its public surface)
    _fetch_gemini_models,
    _llm_request_controller,
    _provider_error_payload,
    _record_llm_thread_stats,
    _trace_http_request,
    _trace_http_response,
    clean_xml_chars,
    is_rate_limited,
    llm_thread_stats,
    llm_thread_stats_delta,
    normalize_openai_base_url,
    parse_json_array,
    parse_translation_array,
    provider_error_message,
    provider_error_status,
    reasoning_effort_options,
    reset_llm_clients,
    resolve_openai_base_url,
    set_llm_base_url,
    set_llm_reasoning_effort,
)
from folith.engine.glossary_checks import (  # noqa: F401  (re-export: core keeps its public surface)
    _count_sentences,
    _deterministic_finding,
    _globalize_batch_findings,
    _local_name,
    _termbase_df_to_entries,
    check_glossary_compliance,
    extract_auto_terms,
    extract_preserved_tokens,
    find_residuals,
    glossary_block,
    glossary_to_terms,
    import_tmx,
    is_incomplete_translation,
    normalize_glossary,
    parse_termbase,
    parse_termbase_csv,
    parse_termbase_tbx,
)
from folith.engine.tm_store import (  # noqa: F401  (re-export: core keeps its public surface)
    _TranslationMemory,
    _ensure_tm_dir,
    _migrate_legacy_system_tm,
    _tm_eligible,
    _tm_hit,
    _write_tm_bytes,
    copy_system_tm_to_project,
    knowledge_cache_path,
    state_target_lang,
    tm_index,
    tm_normalize,
    tm_path,
    tm_project_id,
    translation_cache_path,
)
from folith.engine.locks import (  # noqa: F401  (re-export: core keeps its public surface)
    _exclusive_job_execution,
    _serialized_job_mutation,
    _serialized_project_mutation,
    job_state_lock_path,
    job_state_mutation_lock,
    mutate_job_state,
    project_lock_path,
    project_mutation_lock,
)
from folith.engine.projects import (  # noqa: F401  (re-export: core keeps its public surface)
    _create_project_record,
    _project_backup_dir,
    _project_memory_injection,
    _project_now_iso,
    archive_project,
    assign_job_to_project,
    assign_jobs_to_project,
    copy_default_tm_to_project,
    create_project,
    default_project_view,
    delete_project,
    ensure_default_project,
    ensure_system_project,
    export_project_memory,
    find_project_by_name,
    import_project_memory,
    is_system_project,
    is_system_project_id,
    list_active_project_options,
    list_active_projects,
    list_project_conflicts,
    list_project_jobs,
    list_projects,
    list_unassigned_jobs,
    load_project,
    project_for_job,
    project_injection,
    project_memory_view,
    project_name_for_id,
    project_sections,
    project_summary,
    promote_job_to_project,
    rename_project,
    require_project,
    resolve_all_project_conflicts,
    resolve_project_conflict,
    resolve_project_ref,
    resolved_project_id,
    restore_project,
    save_project,
    system_project_id,
    system_project_view,
    task_project_id_for_new_job,
    update_project,
    validate_project_name,
)
from folith.engine.source_quality import (  # noqa: F401  (re-export: core keeps its public surface)
    _apply_source_quality_admissions,
    _block_bbox,
    _body_text_bbox,
    _build_extraction_report,
    _confirmed_source_quality_indexes,
    _filter_confirmed_source_quality,
    _grounded_block_index,
    _segment_source_quality_evidence,
    _source_quality_dispositions,
    _source_quality_finding,
    _source_quality_gate_enabled,
    _source_segmentation_mode,
    _text_layer_decision,
    audit_source_quality,
    cleanup_source_paragraphs,
    confirm_source_quality_segment,
    extract_document_paragraphs,
    extraction_failure_detail,
    friendly_extraction_reason,
    segment_source_paragraphs,
)
from folith.engine.glossary_ops import (  # noqa: F401  (re-export: core keeps its public surface)
    _apply_glossary_staleness,
    _glossary_mutation_note,
    _refresh_entity_consistency,
    accept_entity_variants,
    add_glossary_entry,
    apply_entity_repair,
    bypass_freeze,
    confirm_translation_style_rule,
    delete_glossary_entry,
    entity_consistency_queue,
    freeze_glossary,
    load_context_artifacts,
    preview_entity_repair,
    review_knowledge_candidate,
    save_document_profile,
    save_glossary_draft,
    set_entity_translation,
    set_glossary_entry_status,
    unfreeze_glossary,
    update_glossary_entry,
    write_profile_artifacts,
)
from folith.engine.segments import (  # noqa: F401  (re-export: core keeps its public surface)
    restore_translation_edit,
    _auto_exclude_non_translatable_segments,
    _drop_segment_at,
    _excluded_segment_record,
    _excluded_segments_md,
    _exclusion_restore_index,
    _non_translatable_exclusions,
    _recheck_delivery_invariants_for_segments,
    _reconcile_translation_failures,
    _recount_reviewed_segments,
    _segment_source_list,
    _shift_paragraph_ranges,
    active_pair_indexes,
    assign_missing_segment_uids,
    can_undo_translation_segment_structure,
    clear_translation_drafts,
    excluded_segment_records,
    excluded_segments_manifest,
    excluded_segments_summary,
    is_excluded,
    load_translation_drafts,
    mutate_translation_segments,
    next_segment_uid,
    save_translation_drafts,
    segment_uid,
    translation_drafts_path,
    translation_terms_for_pair,
    translation_visible_indexes,
    undo_translation_segment_structure,
)
from folith.engine.delivery_state import (  # noqa: F401  (re-export: core keeps its public surface)
    _academic_workspace_archive,
    _delivery_timing,
    _finalization_artifacts,
    _invalidate_final_delivery_state,
    _invalidate_translation_reviews,
    _is_pdf_source,
    _mark_translation_truth_changed,
    _reconcile_final_delivery_snapshot,
    _record_delivery_validation_findings,
    _reset_final_qa,
    academic_status_label,
    approve_delivery,
    create_delivery_snapshot,
    current_translation_hash,
    delivery_snapshot_archive,
    delivery_snapshot_assets,
    delivery_snapshot_status,
    delivery_status_label,
    dependency_impact_view,
    list_delivery_snapshots,
    original_layout_pdf_applicable,
    progress_label,
    recovery_summary,
    task_status_label,
    translation_truth_view,
    validate_delivery_translation_state,
    validate_translation_pairs,
    validate_translation_target,
)



def _ocr_backend_version(backend, command=""):
    """The real engine version, part of the OCR cache identity.

    A silent Tesseract upgrade changes glyph output, so a cached page produced
    by another version must not be reused.  When the version cannot be probed we
    record ``unknown`` rather than pretending it is stable.
    """
    key = (str(backend or ""), str(command or ""))
    if key in _OCR_BACKEND_VERSION_CACHE:
        return _OCR_BACKEND_VERSION_CACHE[key]
    version = ""
    if str(backend or "") == "tesseract" and command:
        try:
            proc = subprocess.run(
                [command, "--version"], capture_output=True, text=True,
                encoding="utf-8", errors="replace", timeout=15)
            lines = (proc.stdout or proc.stderr or "").splitlines()
            if lines:
                match = re.search(r"(\d+\.\d+\.\d+)", lines[0])
                version = match.group(1) if match else lines[0].strip()[:40]
        except Exception:  # noqa: BLE001 - version probing must not block OCR
            version = ""
    _OCR_BACKEND_VERSION_CACHE[key] = version or "unknown"
    return _OCR_BACKEND_VERSION_CACHE[key]


# ================= OCR 来源：线上 OCR 配置 =================



def load_ocr_online_config():
    """Credential-bearing OCR connection profile; never enters a task snapshot."""
    try:
        path = ocr_online_config_path()
        if not path.exists():
            return None
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            return None
        return {
            "provider": str(data.get("provider") or ""),
            "model": str(data.get("model") or ""),
            "api_key": str(data.get("api_key") or ""),
            "base_url": normalize_openai_base_url(data.get("base_url")) or "",
            "max_tokens": int(data.get("max_tokens") or _ocr_online.DEFAULT_MAX_TOKENS),
            "temperature": float(data.get("temperature") or 0.0),
            "reasoning_mode": str(data.get("reasoning_mode") or "provider_default"),
            "timeout_s": float(data.get("timeout_s") or _ocr_online.DEFAULT_TIMEOUT_S),
            "max_retries": int(data.get("max_retries") or _ocr_online.DEFAULT_MAX_RETRIES),
        }
    except Exception:  # noqa: BLE001 - a damaged profile is "not configured"
        return None



def resolve_paddle_model(model_id):
    """Resolve a task-level Paddle model selection, failing on an unknown id."""
    return _ocr_models.try_resolve(model_id)



def _ocr_page_quality_ok(record, *, min_confidence=None, min_density=None):
    """Quality signal that decides whether a page needs a higher-DPI retry.

    Deliberately conservative: it only says "not good enough" on evidence the
    rest of the pipeline already treats as a source-quality problem —
    replacement characters, very low alphanumeric density, or low OCR
    confidence.  It never claims a page is perfect, only that re-running it at a
    higher DPI has a plausible chance of helping.
    """
    threshold = (OCR_ADAPTIVE_MIN_CONFIDENCE if min_confidence is None
                 else float(min_confidence))
    density_floor = (OCR_ADAPTIVE_MIN_TEXT_DENSITY if min_density is None
                     else float(min_density))
    blocks = (record or {}).get("blocks") or []
    texts = []
    confidences = []
    for block in blocks:
        text = re.sub(r"\s+", " ", str(block.get("text") or "").strip())
        if not text:
            continue
        texts.append(text)
        try:
            confidence = float(block.get("confidence"))
        except (TypeError, ValueError):
            confidence = None
        if confidence is not None:
            confidences.append(confidence)
    if not texts:
        return False
    joined = " ".join(texts)
    if "\ufffd" in joined or "\u25a1" in joined:
        return False
    alnum = sum(1 for char in joined if char.isalnum())
    if alnum / max(1, len(joined)) < density_floor:
        return False
    if confidences and (sum(confidences) / len(confidences)) < threshold:
        return False
    return True


def _ocr_pdf_text_with_warnings(
    file_bytes, max_pages=None, on_progress=None, *, workers=None, queue_size=None,
    checkpoint_dir=None, cancel_check=None, profiler=None, return_details=False,
    dpi=None, adaptive_dpi=None, backend=None, profile_cache_dir=None,
    ocr_language=None, ocr_psm=None, paddle_model=None, online_config=None,
):
    """OCR a scanned PDF with the local backend and return ``(text, warnings)``.

    ``max_pages=3`` is used by the quick-profile preview.  ``None`` means all
    pages and is used by the actual translation pipeline.  The function never
    sends page images or OCR text anywhere; it only invokes the local engine.

    ``ocr_language`` is the **user-selected** OCR language (a Tesseract code or
    registry key, optionally a ``ara+eng`` combination).  It is never inferred
    from the installed packs or the document.  When it names a pack that is not
    installed the function fails visibly and returns no text, because OCR run
    with the wrong pack produces plausible-looking text that poisons every later
    stage.  ``None`` is only for legacy tasks and is handled with an explicit
    warning.

    Durable caching is page-granular (``ocr_cache/manifest.json`` plus one file
    per page).  Because a page is keyed by *page index* rather than by the
    requested index set, the three representative pages produced by the
    quick-profile pass are reused by the full run instead of being OCR'd twice.

    ``profile_cache_dir`` adds a second, **read-only** cache root.  It exists
    because the quick profile runs before a job exists, so its pages cannot live
    in the job directory; it writes to the shared content-addressed profile
    cache instead, and the formal run reads that root in addition to its own
    job-scoped cache.  A page is only reused when its full identity matches, so
    a different DPI / backend / version / document still recomputes.

    ``adaptive_dpi`` is an opt-in experiment: a low-DPI first pass, with a
    higher-DPI rerun only for pages whose quality signal fails.
    """
    warnings = []
    command = _find_tesseract()
    requested_engine = _ocr_backends.normalize_backend(backend)
    try:
        effective_psm = int(ocr_psm) if ocr_psm is not None else _ocr_language.DEFAULT_PSM
    except (TypeError, ValueError):
        effective_psm = _ocr_language.DEFAULT_PSM
    version = _ocr_backend_version("tesseract", command) if command else ""
    try:
        resolved_backend, backend_warnings = _ocr_backends.resolve_backend(
            requested_engine, tesseract_command=command, tesseract_version=version,
            page_ocr=_ocr_page_with_tesseract, language=ocr_language or "eng",
            psm=effective_psm, paddle_model=paddle_model,
            online_config=online_config)
    except _ocr_backends.OCRUnavailable as exc:
        result = ("", [str(exc)], {"ocr_error": exc.code,
                  "requested_engine": requested_engine, "actual_engine": None,
                  "ocr_source": requested_engine})
        return result if return_details else result[:2]
    warnings.extend(backend_warnings)
    installed = []
    if resolved_backend.name == "tesseract":
        if not command:
            result = ("", ["未找到本机 Tesseract OCR。请安装 Tesseract，或设置 FOLIOTHREAD_TESSERACT_CMD。"],
                      {"ocr_error": "runtime_missing", "requested_engine": requested_engine})
            return result if return_details else result[:2]
        installed = _tesseract_languages(command)
        resolved_language, language_warnings, language_error = _resolve_ocr_language_config(ocr_language, installed)
        if language_error is not None:
            details = {**{k: v for k, v in language_error.items() if k not in {"error", "message"}},
                "ocr_error": language_error["error"], "installed_languages": list(installed or []),
                "installed_language_available": False, "ocr_backend": "tesseract", "requested_engine": requested_engine}
            result = ("", [language_error["message"]], details)
            return result if return_details else result[:2]
        warnings.extend(language_warnings)
    else:
        try:
            resolved_language = _ocr_language.resolve_ocr_language(ocr_language or "eng")
        except ValueError as exc:
            result = ("", [str(exc)], {"ocr_error": "unsupported", "requested_engine": requested_engine})
            return result if return_details else result[:2]
    language = resolved_language.engine_code
    text_direction_value = resolved_language.direction
    workers = max(1, min(int(workers or _ocr_workers_default()), resolved_backend.max_workers))
    queue_size = max(1, min(int(queue_size or _ocr_queue_size_default(workers)), 32))
    runtime_metadata = resolved_backend.runtime_info(language).to_dict()
    runtime_metadata.update(requested_engine=requested_engine, actual_engine=resolved_backend.name,
        ocr_source=resolved_backend.name,
        workers=workers, fallbacks=backend_warnings, confidence_kind=resolved_backend.capabilities()["confidence_kind"])
    configured_device = str(os.environ.get("FOLIOTHREAD_OCR_DEVICE", "auto")).lower()
    if resolved_backend.name == "tesseract" and configured_device not in {"", "auto", "cpu"}:
        warnings.append(f"OCR 配置请求 device={configured_device}，但当前 Tesseract runtime 实际使用 CPU。")
    backend_name = resolved_backend.name
    backend_model = runtime_metadata["model"]
    backend_version = runtime_metadata["backend_version"]
    if runtime_metadata.get("model_version") == "models_not_installed":
        try:
            if cancel_check and cancel_check():
                raise RuntimeError("任务已请求取消")
            resolved_backend.load(on_progress, cancel_check=cancel_check)
            runtime_metadata.update(resolved_backend.runtime_info(language).to_dict())
        except _ocr_backends.OCRUnavailable as exc:
            result = ("", [str(exc)], {**runtime_metadata, "ocr_error": exc.code})
            resolved_backend.close()
            return result if return_details else result[:2]
    model_version = runtime_metadata.get("model_version") or ""
    cache_backend_version = backend_version
    if model_version:
        cache_backend_version += "|models=" + model_version
    if runtime_metadata.get("inference_config"):
        cache_backend_version += "|config=" + runtime_metadata["inference_config"]
    # Backend-specific, secret-free configuration that changes the bytes of a
    # page: the Paddle model set, or the online target/model/prompt/parameters.
    # It is empty for the historical Tesseract default, so pages cached before
    # this contract existed stay reusable.
    cache_backend_options = _ocr_backends.backend_options(backend_name, runtime_metadata)
    runtime_metadata["backend_options"] = cache_backend_options
    if resolved_backend.capabilities().get("geometry") is False:
        warnings.append("该 OCR 来源只返回文字、不返回行框；原页对照只能定位页面，无法按坐标高亮。")
    base_dpi = max(72, int(dpi or OCR_DPI_BASELINE))
    language_metadata = {
        **resolved_language.to_dict(),
        "ocr_language_label": resolved_language.label,
        "ocr_language_engine_code": runtime_metadata.get("engine_language", language),
        "ocr_text_direction": text_direction_value,
        "ocr_direction": text_direction_value,
        "ocr_dpi": base_dpi,
        "ocr_psm": effective_psm,
        "ocr_backend": backend_name,
        "ocr_source": backend_name,
        "ocr_config_version": _ocr_language.OCR_CONFIG_VERSION,
        "installed_languages": list(installed or []),
        "installed_language_available": True,
        "explicit_ocr_language": bool(str(ocr_language or "").strip()),
    }
    runtime_metadata = {**runtime_metadata, **language_metadata}
    adaptive = (_ocr_env_flag("FOLIOTHREAD_OCR_ADAPTIVE_DPI", False)
                if adaptive_dpi is None else bool(adaptive_dpi))
    first_dpi = OCR_ADAPTIVE_LOW_DPI if adaptive else base_dpi
    retry_dpi = max(base_dpi, OCR_ADAPTIVE_HIGH_DPI) if adaptive else base_dpi
    source_hash = hashlib.sha256(file_bytes or b"").hexdigest()
    cache = _OCRPageCache(_ocr_cache_dir(checkpoint_dir))
    # Read-only secondary cache holding the quick-profile representative pages.
    # It lives outside the job directory because profiling happens before a job
    # exists.  A page found here is byte-identical to what this run would
    # produce (the identity match proves it), so it is reused rather than
    # recomputed.  Nothing is ever written to it from here.
    _profile_dir = Path(profile_cache_dir) if profile_cache_dir else None
    _profile_cache_root = _ocr_cache_dir(_profile_dir) if _profile_dir else None
    if _profile_cache_root is not None and _profile_cache_root == _ocr_cache_dir(checkpoint_dir):
        _profile_cache_root = None
    profile_cache = (_OCRPageCache(_profile_cache_root)
                     if _profile_cache_root is not None else None)
    legacy_path = _ocr_checkpoint_path(checkpoint_dir)

    def _identity(index, dpi_value):
        return _OCRPageIdentity(
            source_sha256=source_hash, page_index=int(index),
            backend=backend_name, model=backend_model,
            backend_version=cache_backend_version, language=language,
            dpi=int(dpi_value), preprocessing_version=_OCR_PREPROCESSING_VERSION,
            psm=effective_psm, direction=text_direction_value,
            ocr_config_version=_ocr_language.OCR_CONFIG_VERSION,
            backend_options=cache_backend_options)

    def _load_cached(loader_dpi):
        records, _manifest = cache.load(
            source_sha256=source_hash, indices=indices, language=language,
            backend=backend_name, model=backend_model,
            backend_version=cache_backend_version, dpi=int(loader_dpi),
            psm=effective_psm, direction=text_direction_value,
            ocr_config_version=_ocr_language.OCR_CONFIG_VERSION,
            backend_options=cache_backend_options)
        if profile_cache is not None:
            missing = [index for index in indices if str(index) not in records]
            if missing:
                extra, _ = profile_cache.load(
                    source_sha256=source_hash, indices=missing, language=language,
                    backend=backend_name, model=backend_model,
                    backend_version=cache_backend_version, dpi=int(loader_dpi),
                    psm=effective_psm, direction=text_direction_value,
                    ocr_config_version=_ocr_language.OCR_CONFIG_VERSION,
                    backend_options=cache_backend_options)
                for key, record in extra.items():
                    records.setdefault(key, record)
        return records

    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            count = doc.page_count
            if count == 0:
                result = ("", ["PDF 不包含任何页面，无法执行 OCR。"], {})
                return result if return_details else result[:2]
            if max_pages is None:
                indices = list(range(count))
            else:
                indices = sorted({0, count // 2, count - 1})[:max_pages]
            # Lookup order is best-evidence-first: a page already computed at
            # the retry DPI beats the first-pass DPI, which beats a legacy
            # checkpoint (adopted only when its identity matches exactly).
            page_records = {}
            dpi_by_page = {}
            if adaptive:
                for loader_dpi in (retry_dpi, first_dpi):
                    for key, record in _load_cached(loader_dpi).items():
                        if key not in page_records:
                            page_records[key] = record
                            dpi_by_page[key] = loader_dpi
            else:
                for key, record in _load_cached(first_dpi).items():
                    page_records[key] = record
                    dpi_by_page[key] = first_dpi
            legacy_pages = _read_legacy_ocr_pages(
                None if model_version else legacy_path, source_sha256=source_hash, indices=indices,
                language=language, backend=backend_name, model=backend_model,
                dpi=base_dpi, psm=effective_psm, direction=text_direction_value,
                ocr_config_version=_ocr_language.OCR_CONFIG_VERSION)
            for key, record in legacy_pages.items():
                if key not in page_records:
                    page_records[key] = record
                    dpi_by_page[key] = base_dpi
            cached_pages = len(page_records)
            failed_pages = []
            render_stage = profiler.start_stage(
                "rasterize", concurrency=1, page_count=len(indices),
                metadata={"dpi": first_dpi, "baseline_dpi": base_dpi,
                          "adaptive_dpi": adaptive, "queue_size": queue_size,
                          "cached_pages": cached_pages},
            ) if profiler is not None else None
            ocr_stage = profiler.start_stage(
                "ocr", concurrency=workers, page_count=len(indices),
                metadata={**runtime_metadata, "queue_size": queue_size,
                          "cached_pages": cached_pages, "adaptive_dpi": adaptive,
                          "first_dpi": first_dpi, "retry_dpi": retry_dpi},
            ) if profiler is not None else None
            page_lock = threading.RLock()
            queue_lock = threading.RLock()
            stop_event = threading.Event()
            worker_errors = []
            completed_pages = cached_pages
            failure_count = 0
            retried_pages = 0
            dpi_used_by_page = {}
            page_timings = {}
            page_requests = {}
            adaptive_retried = []
            # Constant-size run record: written once when the run starts so a
            # crash leaves evidence of what was in flight, then once at the end.
            cache.write_manifest(
                source_sha256=source_hash, indices=indices, language=language,
                backend=backend_name, model=backend_model,
                backend_version=cache_backend_version, dpi=first_dpi, status="running",
                completed_pages=cached_pages,
                psm=effective_psm, direction=text_direction_value,
                ocr_config_version=_ocr_language.OCR_CONFIG_VERSION,
                backend_options=cache_backend_options,
                extra={"adaptive_dpi": adaptive, "baseline_dpi": base_dpi,
                       "first_dpi": first_dpi, "retry_dpi": retry_dpi,
                       **language_metadata})

            def check_cancel():
                if cancel_check and cancel_check():
                    raise RuntimeError("任务已请求取消")

            def progress_source_label():
                """Name the source that is *actually* reading the pages."""
                if backend_name == "online":
                    return f"线上 OCR API · {runtime_metadata.get('model') or ''}".strip()
                model_label = str(runtime_metadata.get("model_label") or "")
                if backend_name == "paddleocr" and runtime_metadata.get(
                        "model_id") not in {None, "", "official"} and model_label:
                    return f"{resolved_backend.display_name} · {model_label}"
                return resolved_backend.display_name

            def emit_page_progress():
                if not on_progress:
                    return
                with queue_lock:
                    current = completed_pages
                    failed = failure_count
                    retried = retried_pages
                detail = ""
                if failed:
                    detail += f"，失败 {failed} 页"
                if retried:
                    detail += f"，重试 {retried} 次"
                try:
                    on_progress(
                        f"【阶段一】扫描 PDF OCR（{current}/{len(indices)} 页；"
                        f"{progress_source_label()} · {resolved_language.label}{detail}）…"
                    )
                except Exception:  # noqa: BLE001 - UI progress is optional
                    pass

            def run_pass(pass_indices, pass_dpi):
                """Bounded worker fan-out at one DPI; returns failed (page, error)."""
                nonlocal completed_pages
                if not pass_indices:
                    return []
                check_cancel()
                resolved_backend.load(on_progress, cancel_check=cancel_check)
                runtime_metadata.update(resolved_backend.runtime_info(language).to_dict())
                if ocr_stage is not None:
                    ocr_stage.record["metadata"].update(runtime_metadata)
                work_queue = queue.Queue(maxsize=queue_size)
                local_failed = []
                pass_completed = 0

                def ocr_worker():
                    nonlocal completed_pages, pass_completed, failure_count, retried_pages
                    while True:
                        task = work_queue.get()
                        try:
                            if task is None:
                                return
                            idx, png_bytes, enqueued_at, page_width, page_height = task
                            check_cancel()
                            started = time.perf_counter()
                            evidence = resolved_backend.recognize_page(
                                png_bytes, page_index=idx, page_number=idx + 1,
                                language=language, dpi=pass_dpi)
                            error = evidence.error
                            page_timings[str(idx)] = evidence.timings
                            if getattr(evidence, "request", None):
                                with page_lock:
                                    page_requests[str(idx)] = dict(evidence.request)
                                    if int(evidence.request.get("attempt_count") or 1) > 1:
                                        retried_pages += int(evidence.request["attempt_count"]) - 1
                            finished = time.perf_counter()
                            if error:
                                with page_lock:
                                    local_failed.append((idx + 1, error))
                                    failure_count += 1
                            else:
                                record = evidence.to_record(page_width, page_height)
                                with page_lock:
                                    page_records[str(idx)] = record
                                    dpi_by_page[str(idx)] = pass_dpi
                                    dpi_used_by_page[str(idx)] = pass_dpi
                                    # Page-granular durability: this write holds
                                    # only the page that just finished.
                                    stored = cache.store_page(_identity(idx, pass_dpi), record)
                                    if cache.directory is not None and stored is None:
                                        message = "OCR 页面缓存写入失败；检查可用磁盘空间，恢复时该页需重新识别。"
                                        if message not in warnings:
                                            warnings.append(message)
                            with queue_lock:
                                pass_completed += 1
                                completed_pages += 1
                            if ocr_stage is not None:
                                ocr_stage.item(
                                    f"page-{idx + 1}", started,
                                    finished_monotonic=finished,
                                    queue_wait_s=max(0.0, started - enqueued_at),
                                    status="completed" if not error else "failed",
                                    metadata={"kind": "page", "page_number": idx + 1,
                                              "dpi": pass_dpi},
                                )
                            emit_page_progress()
                        except BaseException as exc:  # noqa: BLE001 - propagate cancel
                            worker_errors.append(exc)
                            stop_event.set()
                        finally:
                            work_queue.task_done()

                threads = [threading.Thread(target=ocr_worker,
                                            name=f"ocr-worker-{i}", daemon=True)
                           for i in range(workers)]
                try:
                    for thread in threads:
                        thread.start()
                    for idx in pass_indices:
                        check_cancel()
                        if stop_event.is_set():
                            break
                        started = time.perf_counter()
                        pix = doc[idx].get_pixmap(dpi=pass_dpi, alpha=False)
                        png_bytes = pix.tobytes("png")
                        rendered_at = time.perf_counter()
                        queued_at = rendered_at
                        while True:
                            try:
                                page_width = float(doc[idx].rect.width)
                                page_height = float(doc[idx].rect.height)
                                work_queue.put(
                                    (idx, png_bytes, queued_at, page_width, page_height),
                                    timeout=0.2,
                                )
                                queued_at = time.perf_counter()
                                break
                            except queue.Full:
                                check_cancel()
                                if stop_event.is_set():
                                    break
                        if render_stage is not None:
                            render_stage.item(
                                f"page-{idx + 1}", started,
                                finished_monotonic=queued_at,
                                queue_wait_s=max(0.0, queued_at - rendered_at),
                                metadata={"kind": "page", "page_number": idx + 1,
                                          "dpi": pass_dpi},
                            )
                        if stop_event.is_set():
                            break
                    while not stop_event.is_set() and pass_completed < len(pass_indices):
                        check_cancel()
                        time.sleep(0.1)
                    if worker_errors:
                        raise worker_errors[0]
                finally:
                    if worker_errors:
                        stop_event.set()
                    if stop_event.is_set():
                        resolved_backend.close()
                    for _ in threads:
                        while True:
                            try:
                                work_queue.put(None, timeout=0.2)
                                break
                            except queue.Full:
                                if not any(thread.is_alive() for thread in threads):
                                    break
                    for thread in threads:
                        thread.join(timeout=10)
                return list(local_failed)

            cancelled = False
            try:
                emit_page_progress()
                failed_pages.extend(run_pass(
                    [index for index in indices if str(index) not in page_records],
                    first_dpi))
                if adaptive and not stop_event.is_set():
                    retry_indices = [
                        index for index in indices
                        if str(index) in page_records
                        and dpi_by_page.get(str(index), first_dpi) < retry_dpi
                        and not _ocr_page_quality_ok(page_records[str(index)])
                    ]
                    if retry_indices:
                        adaptive_retried = [index + 1 for index in retry_indices]
                        warnings.append(
                            f"自适应 OCR 分辨率：{len(retry_indices)} 页首轮质量信号不达标，"
                            f"已用 {retry_dpi} DPI 重跑（baseline {base_dpi} DPI）。")
                        failed_pages.extend(run_pass(retry_indices, retry_dpi))
            except BaseException:
                cancelled = True
                stop_event.set()
                raise
            finally:
                if render_stage is not None:
                    render_stage.finish(status="cancelled" if cancelled else "completed")
                if ocr_stage is not None:
                    ocr_stage.finish(status="cancelled" if cancelled else "completed")
                cache.write_manifest(
                    source_sha256=source_hash, indices=indices, language=language,
                    backend=backend_name, model=backend_model,
                    backend_version=cache_backend_version, dpi=first_dpi,
                    status="cancelled" if cancelled else "completed",
                    completed_pages=sum(1 for index in indices
                                        if str(index) in page_records),
                    failed_pages=len(failed_pages),
                    psm=effective_psm, direction=text_direction_value,
                    ocr_config_version=_ocr_language.OCR_CONFIG_VERSION,
                    backend_options=cache_backend_options,
                    extra={"adaptive_dpi": adaptive, "baseline_dpi": base_dpi,
                           "first_dpi": first_dpi, "retry_dpi": retry_dpi,
                           "retried_pages": adaptive_retried,
                           **language_metadata,
                           "dpi_used_by_page": {
                               key: dpi_used_by_page[key]
                               for key in sorted(dpi_used_by_page)}})

            if failed_pages:
                sample = "；".join(f"第 {page} 页：{detail}" for page, detail in failed_pages[:3])
                more = f"（另有 {len(failed_pages) - 3} 页失败）" if len(failed_pages) > 3 else ""
                warnings.append(f"OCR 部分页面失败：{sample}{more}")
            postprocess_started = time.perf_counter()
            paragraph_confidences = []
            chunks = []
            repeated = {}
            all_blocks = []
            for idx in indices:
                record = page_records.get(str(idx)) or {}
                height = float(record.get("page_height") or 0)
                for block in record.get("blocks") or []:
                    text = re.sub(r"\s+", " ", str(block.get("text") or "").strip())
                    if not text:
                        continue
                    normalized = re.sub(r"\d+", "#", text).casefold()
                    repeated[normalized] = repeated.get(normalized, 0) + 1
                    all_blocks.append((idx, height, block, text, normalized))
            repeat_threshold = max(2, int((len(indices) * 0.2) + 0.999)) if len(indices) > 1 else 99
            by_page = {}
            geometry_pages = 0
            coordinate_less_pages = 0
            for idx, height, block, text, normalized in all_blocks:
                mapped = block.get("bbox_pdf")
                # A block whose geometry is explicitly unknown must never be
                # treated as "at the top of the page": that would silently drop
                # repeated body lines as if they were running heads.
                known = bool(block.get("geometry_known", True)) and bool(mapped)
                top = float(mapped[1] if mapped else block.get("top") or 0)
                bottom = float(mapped[3] if mapped else block.get("bottom") or 0)
                is_edge = known and (top < height * 0.12 or bottom > height * 0.88)
                if known:
                    geometry_pages += 1
                if _OCR_PAGE_NUMBER_RE.fullmatch(text) and is_edge:
                    continue
                if repeated.get(normalized, 0) >= repeat_threshold and is_edge:
                    continue
                by_page.setdefault(idx, []).append((text, block.get("confidence")))
            for record in page_records.values():
                if record.get("geometry_known") is False:
                    coordinate_less_pages += 1
            for idx in indices:
                page_chunks = by_page.get(idx) or []
                if not page_chunks:
                    continue
                chunks.append("\n\n".join(text for text, _ in page_chunks))
                paragraph_confidences.extend(conf for _, conf in page_chunks)
            valid_confidences = [float(value) for value in paragraph_confidences
                                 if value is not None]
            completed_indices = [index for index in indices
                                 if str(index) in page_records]
            usage_totals = {"reported": False, "prompt_tokens": None,
                            "completion_tokens": None, "reasoning_tokens": None,
                            "total_tokens": None, "amount": None}
            usage_unknown_pages = []
            for key in sorted(page_requests, key=lambda value: int(value)):
                record = page_requests[key].get("usage") or {}
                if not record.get("reported"):
                    usage_unknown_pages.append(int(key) + 1)
                    continue
                usage_totals["reported"] = True
                for field in ("prompt_tokens", "completion_tokens",
                              "reasoning_tokens", "total_tokens", "amount"):
                    current = usage_totals[field]
                    reported = record.get(field)
                    if reported is None:
                        continue
                    usage_totals[field] = (current or 0) + reported
            details = {
                **runtime_metadata,
                "page_count": len(indices),
                "completed_pages": len(completed_indices),
                "page_order": [index + 1 for index in completed_indices],
                "failed_pages": len(failed_pages),
                # Never let a partly recognised book be reported as complete.
                "all_pages_completed": len(completed_indices) == len(indices),
                "checkpoint_file": str(cache.manifest_path) if cache.manifest_path else "",
                "ocr_cache_dir": str(cache.directory) if cache.directory else "",
                "ocr_cache_stats": cache.stats(),
                "ocr_cache_bytes_written": cache.bytes_written(),
                "ocr_dpi_baseline": base_dpi,
                "ocr_dpi_first_pass": first_dpi,
                "ocr_dpi_retry": retry_dpi,
                "ocr_adaptive_dpi": adaptive,
                "ocr_adaptive_retried_pages": adaptive_retried,
                "ocr_dpi_used_by_page": {key: dpi_by_page[key]
                                         for key in sorted(dpi_by_page)},
                "ocr_cached_pages": cached_pages,
                "ocr_geometry_blocks": geometry_pages,
                "ocr_pages_without_geometry": coordinate_less_pages,
                "ocr_usage": usage_totals,
                "ocr_usage_unknown_pages": usage_unknown_pages,
                "ocr_retried_requests": int(retried_pages),
                "ocr_page_requests": {key: page_requests[key]
                                      for key in sorted(page_requests, key=lambda v: int(v))},
                "paragraph_confidences": paragraph_confidences,
                "mean_confidence": (
                    sum(valid_confidences) / len(valid_confidences)
                    if valid_confidences else None
                ),
                "confidence_threshold": OCR_CONFIDENCE_THRESHOLD_DEFAULT if backend_name == "tesseract" else None,
                "postprocessing_s": time.perf_counter() - postprocess_started,
                "model_load_s": getattr(resolved_backend, "load_seconds", 0.0),
                "page_timings": page_timings,
                "evidence_version": 1,
                "raw_evidence_directory": str(cache.pages_dir) if cache.pages_dir else "",
                "cache_identity": _identity(0, first_dpi).to_dict(),
                "page_failures": [{"page_number": n, "error": e} for n, e in failed_pages],
            }
            text = "\n\n".join(chunks)
            if not text:
                warnings.append("OCR 未识别到有效文本。")
            result = (text, warnings, details)
            return result if return_details else result[:2]
    except _ocr_backends.OCRUnavailable as exc:
        result = ("", [str(exc)], {**runtime_metadata, "ocr_error": exc.code})
        return result if return_details else result[:2]
    except Exception as exc:  # noqa: BLE001 - caller can show a recoverable warning
        if "任务已请求取消" in str(exc):
            raise
        # 同样是"给用户看的第一句"：把已知失败模式（加密/损坏/不完整）翻译成
        # 可执行的中文，技术原文仍写进运行时技术日志。
        result = ("", [f"OCR 处理 PDF 失败："
                       f"{friendly_extraction_reason(str(exc)[:240])}"], {})
        return result if return_details else result[:2]

    finally:
        resolved_backend.close()



EXTRACTION_REPORT_VERSION = 1
_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_M_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"



def _pdf_extraction_facts(file_bytes):
    """Page/image facts for the import report (never raises)."""
    facts = {"pages": None, "images": 0, "text_pages": 0, "empty_pages": 0}
    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as document:
            facts["pages"] = document.page_count
            for page in document:
                if str(page.get_text() or "").strip():
                    facts["text_pages"] += 1
                else:
                    facts["empty_pages"] += 1
                try:
                    facts["images"] += len(page.get_images(full=True))
                except Exception:  # noqa: BLE001
                    continue
    except Exception:  # noqa: BLE001
        return facts
    return facts


_GENERIC_EXTRACTION_WARNING = "OCR 未生成可用段落"


#: 已知的提取失败模式 -> 用户能照着做的中文说明。技术原文仍然保留在警告里
#: （写进运行时技术日志与诊断包），但**呈现给用户的第一句**必须是他能行动的内容：
#: 之前损坏的 DOCX 会把 zipfile 的英文异常原样抛到界面上（"There is no item named
#: '[Content_Types].xml' …"），加密 PDF 则显示 "document closed or encrypted"。
_EXTRACTION_FAILURE_HINTS = (
    ("encrypt", "这个 PDF 已加密，无法读取；请先解除密码保护再导入。"),
    ("password", "这个 PDF 已加密，无法读取；请先解除密码保护再导入。"),
    ("content_types", "无法读取这个 DOCX：文件可能损坏，或它不是真正的 Word 文档。"),
    ("not a zip file", "无法读取这个 DOCX：文件可能损坏，或它不是真正的 Word 文档。"),
    ("failed to open stream", "无法读取这个 PDF：文件可能损坏，或它不是真正的 PDF。"),
    ("no objects found", "无法读取这个 PDF：文件可能损坏，或它不是真正的 PDF。"),
    ("cannot open", "无法读取这个文件：它可能已损坏，或不是支持的文档格式。"),
    ("eof marker", "无法读取这个 PDF：文件不完整（可能没有下载或复制完整）。"),
    ("syntax error", "无法解析这个 PDF：文件可能已损坏。"),
)

#: 文档里确实没有文本（不是错误，而是没有可翻译的内容）。
_NO_TEXT_REASON = ("文档里没有可翻译的文本：可能是空文档，或内容都在图片/表格里；"
                   "扫描件请在设置里启用 OCR 后重新导入。")



def extract_document_paragraphs_with_report(filename, file_bytes, *,
                                            ocr_max_pages=3, on_progress=None,
                                            profiler=None, checkpoint_dir=None,
                                            ocr_workers=None, ocr_queue_size=None,
                                            cancel_check=None,
                                            profile_cache_dir=None,
                                            ocr_language=None, ocr_dpi=None,
                                            ocr_psm=None, ocr_mode=None, ocr_backend=None,
                                            ocr_paddle_model=None, ocr_online_config=None,
                                            source_language=None,
                                            understanding_provider=None, spatial_engine=None):
    """Extract paragraphs together with an import-scope report.

    Returns ``(paragraphs, warnings, report)``.  The report is what makes the
    import honest: it states how much text was read, which parts of the
    document were *not* included in translation, and that the export is a
    regenerated translation document rather than a format-faithful round trip.

    ``ocr_language`` / ``ocr_dpi`` / ``ocr_psm`` are the user's OCR
    configuration; they are forwarded to the OCR stage and recorded in the
    extraction report.

    ``ocr_mode`` (``auto`` / ``force`` / ``off``) and ``source_language``
    implement the ingestion contract: a PDF text layer is used only when the
    quality gate accepts it, ``force`` always OCRs, and ``off`` never does.  The
    declared source language feeds the gate; it is never used to re-detect or
    override the user's language.
    """
    resolved_ocr_mode = normalize_ocr_mode(ocr_mode)
    warnings = []
    paragraphs = []
    unsupported = []
    facts = {}
    ocr_used = False
    text_layer_report = {}
    ocr_trigger = ""
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            classify_stage = profiler.start_stage(
                "pdf_classify", concurrency=1, metadata={"backend": "PyMuPDF"}
            ) if profiler is not None else None
            facts = _pdf_extraction_facts(file_bytes)
            if classify_stage is not None:
                classify_stage.record["item_count"] = int(facts.get("pages") or 0)
                classify_stage.record["page_count"] = int(facts.get("pages") or 0)
                classify_stage.finish(
                    status="completed", text_pages=int(facts.get("text_pages") or 0),
                    empty_pages=int(facts.get("empty_pages") or 0),
                )
            # Text existence is not text quality: judge the layer before
            # trusting it, and let the user's ocr_mode override the verdict.
            text_layer_report = _text_layer_quality.assess_pdf_text_layer(
                file_bytes, source_language=source_language, max_pages=3)
            # Sampling is deliberately shallow (first/middle/last).  If those
            # pages are empty but the document does have text pages elsewhere,
            # the layer is partial, not absent — do not silently force OCR on a
            # real digital document because its first/last page happens to be a
            # cover.
            if (text_layer_report.get("status") == "absent"
                    and int(facts.get("text_pages") or 0) > 0):
                text_layer_report = {
                    **text_layer_report, "status": "uncertain",
                    "reasons": [*(text_layer_report.get("reasons") or []),
                                "text_layer_partial_sampling"],
                }
            use_text_layer, ocr_trigger, mode_warning = _text_layer_decision(
                text_layer_report, resolved_ocr_mode)
            if mode_warning:
                warnings.append(mode_warning)
            layout_stage = profiler.start_stage(
                "layout_recovery", concurrency=1,
                metadata={"backend": "PyMuPDF", "page_count": facts.get("pages"),
                          "ocr_mode": resolved_ocr_mode,
                          "text_layer_quality": text_layer_report.get("status")},
            ) if profiler is not None else None
            paragraphs = []
            if use_text_layer:
                paragraphs = [clean_xml_chars(p)
                              for p in extract_pdf_paragraphs(file_bytes)]
                if not paragraphs and resolved_ocr_mode != "off":
                    # The gate accepted the layer but extraction produced
                    # nothing usable; fall through to OCR instead of shipping an
                    # empty document.
                    use_text_layer = False
                    ocr_trigger = "text_layer_empty_after_extraction"
            if layout_stage is not None:
                layout_stage.record["item_count"] = len(paragraphs)
                layout_stage.record["page_count"] = int(facts.get("pages") or 0)
                layout_stage.finish(status="completed")
            ocr_details = {}
            visual_indices = []
            if use_text_layer and resolved_ocr_mode != "off":
                with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
                    visual_indices = [i for i,p in enumerate(pdf)
                                      if not p.get_text().strip() and p.get_images()]
            if (not use_text_layer or visual_indices) and (understanding_provider is not None
                                      or ocr_backend is None
                                      or _ocr_backends.normalize_backend(ocr_backend) == "online"):
                from folith.document_understanding import VLMDocumentProvider, understand_visual_pdf, understand_mixed_pdf
                from folith.spatial_index import SpatialTextIndex
                if understanding_provider is None:
                    config = ocr_online_config
                    if config is None or isinstance(config, dict):
                        config, error = resolve_ocr_online_config(config)
                        if error:
                            raise ValueError(error["message"])
                    understanding_provider = VLMDocumentProvider(config, cancel_check=cancel_check)
                if spatial_engine is None:
                    command = _find_tesseract()
                    def spatial_engine(page):
                        index = SpatialTextIndex(page.page_id)
                        if not command:
                            index.error = "Local spatial engine unavailable"
                            return index
                        # Single cheap recognition pass. No quality retries,
                        # OCR repair, or dependency on installed language packs.
                        blocks, error = _ocr_page_with_tesseract(
                            command, page.image, ocr_language or "eng", psm=ocr_psm)
                        if error:
                            index.error = str(error)
                            return index
                        return SpatialTextIndex.from_ocr(page, blocks or [],
                                                        image_width=page.image_width,
                                                        image_height=page.image_height)
                understand = understand_mixed_pdf if visual_indices else understand_visual_pdf
                args = (filename, file_bytes, visual_indices, understanding_provider) if visual_indices else (filename, file_bytes, understanding_provider)
                paragraphs, document = understand(
                    *args, spatial_engine=spatial_engine, max_pages=ocr_max_pages, dpi=150,
                    on_progress=on_progress, cancel_check=cancel_check, checkpoint_dir=checkpoint_dir)
                report = _build_extraction_report(filename, paragraphs, facts=facts, warnings=warnings)
                report.update({"semantic_source":"mixed" if visual_indices else "vlm", "document_ir":document,
                               "ocr_used":False, "ocr_mode":resolved_ocr_mode,
                               "ocr_trigger":ocr_trigger, "pdf_text_layer_quality":text_layer_report.get("status"),
                               # The routing decision must be explainable in the
                               # same shape as the native path: this branch is
                               # exactly where a rejected layer ends up.
                               "pdf_text_layer_quality_score":text_layer_report.get("score"),
                               "pdf_text_layer_rejection_reasons":list(text_layer_report.get("reasons") or []),
                               "pdf_text_layer_evidence":text_layer_report.get("evidence") or {},
                               "page_diagnostics":document["provenance"]["page_diagnostics"]})
                if (ocr_max_pages is not None and len(document["provenance"]["page_diagnostics"])
                        < (len(visual_indices) if visual_indices else len(document["pages"]))):
                    report["status"] = "partial"
                return paragraphs, warnings, report
            if not use_text_layer:
                scope = "全文" if ocr_max_pages is None else "前/中/后代表页"
                if ocr_trigger == "ocr_mode_force":
                    warnings.append(f"已按「强制 OCR」忽略文本层，正在 OCR {scope}…")
                elif ocr_trigger == "text_layer_quality_rejected":
                    warnings.append(
                        f"PDF 文本层质量检测未通过（{'、'.join(text_layer_report.get('reasons') or []) or 'quality'}），"
                        f"正在改用 OCR {scope}…")
                else:
                    warnings.append(f"PDF 无可用文本层，正在自动 OCR {scope}…")
                ocr_fn = _ocr_pdf_text_with_warnings
                ocr_kwargs = {
                    "max_pages": ocr_max_pages,
                    "on_progress": on_progress,
                }
                supports_kwargs = False
                try:
                    parameters = inspect.signature(ocr_fn).parameters
                    supports_kwargs = any(
                        parameter.kind == inspect.Parameter.VAR_KEYWORD
                        for parameter in parameters.values()
                    )
                    optional_ocr_kwargs = {
                        "workers": ocr_workers,
                        "queue_size": ocr_queue_size,
                        "checkpoint_dir": checkpoint_dir,
                        "cancel_check": cancel_check,
                        "profiler": profiler,
                        "return_details": True,
                        "profile_cache_dir": profile_cache_dir,
                        "ocr_language": ocr_language,
                        "dpi": ocr_dpi,
                        "ocr_psm": ocr_psm,
                        "backend": ocr_backend,
                        "paddle_model": ocr_paddle_model,
                        "online_config": ocr_online_config,
                    }
                    for key, value in optional_ocr_kwargs.items():
                        if supports_kwargs or key in parameters:
                            ocr_kwargs[key] = value
                except (TypeError, ValueError):
                    optional_ocr_kwargs = {}
                ocr_result = ocr_fn(file_bytes, **ocr_kwargs)
                if isinstance(ocr_result, tuple) and len(ocr_result) >= 3:
                    ocr_text, ocr_warnings, ocr_details = ocr_result[:3]
                else:
                    ocr_text, ocr_warnings = ocr_result
                warnings.extend(ocr_warnings)
                paragraphs = _ocr_text_to_paragraphs(
                    ocr_text,
                    direction=str((ocr_details or {}).get("ocr_text_direction")
                                  or "ltr"))
                ocr_used = True
                if not paragraphs:
                    # A blocking OCR error (missing language pack, no engine,
                    # unreadable pages) already carries the precise reason; a
                    # generic follow-up would become the surfaced message and
                    # hide it from the user.
                    if not (ocr_details or {}).get("ocr_error"):
                        warnings.append(
                            "OCR 未生成可用段落；请安装语言包或手动选择风格")
            images = int(facts.get("images") or 0)
            if images:
                unsupported.append({
                    "kind": "image", "label": "图片", "count": images,
                    "detail": f"{images} 张图片里的文字不在文本层内，不参与翻译",
                    "sample": ""})
            empty_pages = int(facts.get("empty_pages") or 0)
            if empty_pages and not ocr_used:
                unsupported.append({
                    "kind": "page", "label": "无文本层页面",
                    "count": empty_pages,
                    "detail": f"{empty_pages} 页没有文本层（可能是扫描图），"
                              "其内容未进入翻译",
                    "sample": ""})
                warnings.append(
                    f"{empty_pages} 页没有文本层；如需翻译请重新导入并允许 OCR。")
        elif name.endswith(".docx"):
            doc_word = Document(io.BytesIO(file_bytes))
            for p in doc_word.paragraphs:
                for sub_p in re.split(r"\n+", clean_xml_chars(p.text)):
                    t = sub_p.strip()
                    if len(t) > 1 and not _ORNAMENT_RE.match(t):
                        paragraphs.append(t)
            unsupported = _docx_unsupported_parts(doc_word)
        else:
            warnings.append("不支持的文档格式，无法自动画像")
    except Exception as exc:  # noqa: BLE001
        if "任务已请求取消" in str(exc):
            raise
        warnings.append(f"预提取失败：{friendly_extraction_reason(exc)}")
        facts["extraction_error"] = friendly_extraction_reason(exc)
        if hasattr(exc, "diagnostics"):
            facts["understanding_failure"] = exc.diagnostics
    report = _build_extraction_report(
        filename, paragraphs, ocr_used=ocr_used, unsupported=unsupported,
        facts=facts, warnings=warnings)
    if facts.get("extraction_error"):
        report["status"] = "failed"
        report["error"] = facts["extraction_error"]
        if facts.get("understanding_failure"):
            report["page_diagnostics"] = [facts["understanding_failure"]]
    # Why did this document (not) go through OCR?  Recorded as durable evidence
    # so the UI, the audit trail and the benchmark can all answer it.
    report["pdf_text_layer_present"] = bool(
        (text_layer_report.get("evidence") or {}).get("pages_with_text"))
    report["pdf_text_layer_quality"] = text_layer_report.get("status")
    report["pdf_text_layer_quality_score"] = text_layer_report.get("score")
    report["pdf_text_layer_rejection_reasons"] = list(
        text_layer_report.get("reasons") or [])
    report["pdf_text_layer_evidence"] = text_layer_report.get("evidence") or {}
    report["ocr_mode"] = resolved_ocr_mode
    report["ocr_trigger"] = ocr_trigger or (
        "ocr_used" if ocr_used else "text_layer_accepted")
    report["ocr_used"] = bool(ocr_used)
    if ocr_used:
        report["ocr"] = dict(ocr_details or {})
        report["ocr"]["paragraph_count"] = len(paragraphs)
        report["ocr"]["paragraph_confidences"] = list(
            report["ocr"].get("paragraph_confidences") or []
        )
    if name.endswith(".pdf") and not ocr_used and paragraphs and facts.get("pages"):
        from folith.document_understanding import native_pdf_document
        report["document_ir"] = native_pdf_document(filename, file_bytes, paragraphs)
        report["semantic_source"] = "native_pdf"
    return paragraphs, warnings, report



def _apply_source_segmentation(state, cleaned_paragraphs, *, mode=None):
    """Persist paragraph structure and replace only the translation units."""
    source_paragraphs = [str(item or "").strip() for item in cleaned_paragraphs]
    segments, metadata = segment_source_paragraphs(
        source_paragraphs, state=state, mode=mode)
    previous_segmentation = state.get("segmentation") or {}
    if isinstance(previous_segmentation, dict) and previous_segmentation.get("locked"):
        metadata["locked"] = True
    state["source_paragraphs"] = list(source_paragraphs)
    state["segmentation"] = metadata
    from folith.document_ir import segment_spans
    document = (state.get("extraction_report") or {}).get("document_ir")
    if document:
        spans = segment_spans(document, source_paragraphs, segments, metadata)
        state["segment_source_spans"] = spans
        from folith.document_ir import segment_marks
        state["segment_source_marks"] = segment_marks(document, spans, segments)
    else:
        state.pop("segment_source_spans", None)
        state.pop("segment_source_marks", None)
    return segments, metadata



# Extraction provenances that produce trusted semantic blocks rather than rough
# OCR.  `mixed` means every page was read by the VLM (native PDF pages keep their
# text layer); neither carries OCR-fragment risk.
_VLM_SEMANTIC_SOURCES = frozenset({"vlm", "mixed"})



def _audit_source_quality_for_state(state, *, apply_admissions=True):
    paragraphs = (state or {}).get("paras") or []
    if not _source_quality_gate_enabled(state):
        semantic_source = ((state or {}).get("extraction_report") or {}).get(
            "semantic_source")
        return {
            "version": _source_quality.VERSION,
            "status": "disabled",
            "reason": ("vlm_semantic_source"
                       if semantic_source in _VLM_SEMANTIC_SOURCES
                       else "not_pdf_source"),
            "checked_count": len(paragraphs),
            "flagged_count": 0,
            "flags": [],
            "clusters": [],
        }
    report = _filter_confirmed_source_quality(
        audit_source_quality(paragraphs), state, paragraphs)

    # Also audit the pre-cleanup PDF rows. A cleanup join can turn many short,
    # obviously suspect OCR fragments into one long string which no longer
    # matches the short-fragment rules. Preserve those findings through the
    # cleanup output map and paragraph-to-segment ranges.
    raw_paragraphs = (state or {}).get("source_paras") or []
    cleanup = (state or {}).get("source_cleanup") or {}
    segmentation = (state or {}).get("segmentation") or {}
    output_inputs = cleanup.get("output_input_indices")
    paragraph_ranges = segmentation.get("paragraph_ranges") or []
    if not raw_paragraphs or not paragraph_ranges:
        return (_apply_source_quality_admissions(state, report)
                if apply_admissions else report)
    if not isinstance(output_inputs, list):
        # Safe for old jobs only when the stored raw and cleaned paragraph
        # arrays still have a direct one-to-one relationship.
        if len(raw_paragraphs) == len((state or {}).get("source_paragraphs") or []):
            output_inputs = [[index] for index in range(len(raw_paragraphs))]
        else:
            return (_apply_source_quality_admissions(state, report)
                    if apply_admissions else report)

    segments_by_paragraph = {}
    for row in paragraph_ranges:
        if not isinstance(row, dict):
            continue
        paragraph_index = row.get("paragraph_index")
        start, end = row.get("start"), row.get("end")
        if (isinstance(paragraph_index, int) and isinstance(start, int)
                and isinstance(end, int) and 0 <= start <= end <= len(paragraphs)):
            segments_by_paragraph[paragraph_index] = list(range(start, end))

    existing = {
        int(flag.get("segment_index")): flag
        for flag in report.get("flags") or []
        if isinstance(flag, dict) and isinstance(flag.get("segment_index"), int)
    }
    raw_report = audit_source_quality(raw_paragraphs)
    lineage_flags = []
    for raw_flag in raw_report.get("flags") or []:
        raw_index = raw_flag.get("segment_index")
        if not isinstance(raw_index, int):
            continue
        output_indices = [
            output_index for output_index, source_indices in enumerate(output_inputs)
            if isinstance(source_indices, list) and raw_index in source_indices
        ]
        affected = [
            segment_index
            for output_index in output_indices
            for segment_index in segments_by_paragraph.get(output_index, [])
        ]
        if not affected:
            continue
        raw_text = str(raw_flag.get("text") or "").strip()
        cleaned_paragraphs = (state or {}).get("source_paragraphs") or []
        excused = False
        for output_index in output_indices:
            cleaned_p = ""
            if output_index < len(cleaned_paragraphs):
                cleaned_p = str(cleaned_paragraphs[output_index] or "")
            if not cleaned_p:
                cleaned_p = " ".join(str(paragraphs[si] or "") for si in segments_by_paragraph.get(output_index, []))

            if _source_quality.is_integrated_numeric_fragment(raw_text, cleaned_p):
                excused = True
                break

            previous_is_mapped = raw_index > 0 and (raw_index - 1 in output_inputs[output_index])
            if previous_is_mapped and _source_quality.is_integrated_line_continuation(
                    raw_text, raw_paragraphs[raw_index - 1], cleaned_p):
                excused = True
                break

            raw_reasons_set = set(raw_flag.get("reasons") or [])
            if (raw_reasons_set <= {"含未配对的闭合括号", "小写开头的疑似残片或 OCR 杂质", "与相邻可疑片段形成连续异常", "小写且未闭合的短片段"}
                    and not _source_quality._has_unmatched_bracket(cleaned_p)
                    and not _source_quality._reasons(cleaned_p)
                    and raw_text and raw_text.casefold() in cleaned_p.casefold()):
                excused = True
                break

            if (len(cleaned_p) >= 50
                    and re.search(r"[.!?。！？]$", cleaned_p)
                    and not _source_quality._has_unmatched_bracket(cleaned_p)
                    and not _source_quality._reasons(cleaned_p)
                    and raw_text and raw_text.casefold() in cleaned_p.casefold()
                    and (raw_index == 0 or (raw_index - 1) not in output_inputs[output_index]
                         or not re.search(r"[.!?。！？]$", raw_paragraphs[raw_index - 1]))
                    and not (raw_reasons_set & {
                        "孤立数字行，疑似页码或批注", "含未配对的闭合括号",
                        "短片段含异常空格连字符", "孤立字母数字杂质",
                    })):
                excused = True
                break


        if excused:
            continue

        exact = [index for index in affected
                 if raw_text and raw_text.casefold() in
                 str(paragraphs[index] or "").casefold()]
        target_segments = exact if exact else affected
        for segment_index in target_segments:
            if _source_quality.is_integrated_numeric_fragment(
                    raw_text, str(paragraphs[segment_index] or "")):
                continue
            previous_is_mapped = raw_index > 0 and any(
                raw_index - 1 in output_inputs[output_index]
                and segment_index in segments_by_paragraph.get(output_index, [])
                for output_index in output_indices)
            if previous_is_mapped and _source_quality.is_integrated_line_continuation(
                    raw_text, raw_paragraphs[raw_index - 1],
                    str(paragraphs[segment_index] or "")):
                continue
            current = existing.get(segment_index)
            if current is None:
                current = {
                    "segment_index": segment_index,
                    "text": str(paragraphs[segment_index] or ""),
                    "reasons": [],
                    "severity": "blocking",
                    "adjacent_flagged_segments": 1,
                }
                existing[segment_index] = current
                report.setdefault("flags", []).append(current)
            reasons = current.setdefault("reasons", [])
            for reason in raw_flag.get("reasons") or []:
                note = f"清洗前片段「{raw_text[:80]}」：{reason}"
                if note not in reasons:
                    reasons.append(note)
            source_indices = current.setdefault("source_input_indices", [])
            if raw_index not in source_indices:
                source_indices.append(raw_index)
            lineage_flags.append({
                "segment_index": segment_index,
                "source_input_index": raw_index,
                "text": raw_text,
                "reasons": list(raw_flag.get("reasons") or []),
            })
    report["flags"].sort(key=lambda item: int(item.get("segment_index", -1)))
    report["flagged_count"] = len(report["flags"])
    if report["flags"]:
        report["status"] = "blocked"
    report["source_lineage_flags"] = lineage_flags
    report = _filter_confirmed_source_quality(report, state, paragraphs)
    if apply_admissions:
        report = _apply_source_quality_admissions(state, report)
    return report



def call_llm(provider, api_key, model, system_prompt, user_prompt,
             temperature=0.1, base_url=None, response_format=None,
             reasoning_effort=None):
    """底层大模型统一路由（超时 150 秒，模型可配置）。

    支持官方接口与 OpenAI /chat/completions 兼容中转站：
    base_url 显式传入 > 提供商默认 base_url > 会话线程级自定义中转地址。

    传输层 client 按 provider / base_url / proxy 模式 / 超时配置复用
    （见 :mod:`folith.llm_transport`），不再每个请求重建连接池；
    API Key 只以不可逆指纹参与复用键，不进入日志、cache key 明文或性能产物。
    运行时遥测走可合并通道，见 :func:`_runtime_telemetry_request_started`。
    """
    cfg = PROVIDERS.get(provider)
    if not cfg:
        return ""
    runtime_job_id = _runtime_job_id()
    if runtime_job_id:
        _runtime_pause_checkpoint(runtime_job_id)
        if _runtime_cancel_requested(runtime_job_id):
            raise RuntimeError("任务已请求取消")
        _runtime_telemetry_request_started(runtime_job_id)
    _llm_started = time.perf_counter()
    _request_run = trace.get_active_run(runtime_job_id)
    if _request_run:
        _request_run.begin_model_request(provider, model, role=_usage.CURRENT_ROLE.get(),
                                        batch_id=getattr(_LLM_CTX, "batch_id", None),
                                        prompt_cost=getattr(_LLM_CTX, "prompt_cost", None),
                                        system_chars=len(system_prompt or ""),
                                        user_chars=len(user_prompt or ""),
                                        segment_count=(getattr(_LLM_CTX, "prompt_cost", None) or {}).get(
                                            "segment_count", getattr(_LLM_CTX, "input_segments", None)))
    try:
        result = _call_llm_transport(
            cfg, provider, api_key, model, system_prompt, user_prompt,
            temperature=temperature, base_url=base_url,
            response_format=response_format, reasoning_effort=reasoning_effort,
            runtime_job_id=runtime_job_id, started_at=_llm_started)
    except BaseException as exc:
        try:
            _workflow_errors.raise_if_fatal(exc, runtime_job_id)
        except _workflow_errors.WorkflowStopped as stopped:
            exc = stopped
        _llm_dur = time.perf_counter() - _llm_started
        _record_llm_thread_stats(_llm_dur, failed=True)
        run = trace.get_active_run(runtime_job_id)
        if run:
            run.record_model_call(
                provider=provider, model=model, duration_s=_llm_dur,
                input_chars=len(system_prompt or "") + len(user_prompt or ""),
                status="error", error_type=type(exc).__name__,
                metadata={"usage_available": False},
            )
            run.record_error(
                stage="model_request", operation=model,
                error_type=type(exc).__name__, message=str(exc),
            )
        if runtime_job_id:
            # 失败必须立即持久化：既让 UI 不再停在"等待模型"，也保证崩溃恢复
            # 看到的是真实结果而不是一个永远不会结束的在途请求。
            _runtime_telemetry_request_failed(runtime_job_id)
        raise exc
    finally:
        if _request_run:
            _request_run.end_model_request()
    _record_llm_thread_stats(time.perf_counter() - _llm_started, failed=False)
    if runtime_job_id:
        _runtime_telemetry_request_finished(runtime_job_id)
        if _runtime_cancel_requested(runtime_job_id):
            raise RuntimeError("任务已请求取消")
    return result



def _call_llm_transport(cfg, provider, api_key, model, system_prompt, user_prompt,
                        *, temperature=0.1, base_url=None, response_format=None,
                        reasoning_effort=None, runtime_job_id=None, started_at=None):
    """Perform one provider request on a pooled transport client."""
    if cfg["kind"] == "gemini":
        def _build_gemini_client():
            try:
                return genai.Client(
                    api_key=api_key,
                    http_options=genai.types.HttpOptions(
                        timeout=_llm_transport.GEMINI_TIMEOUT_MS))
            except (AttributeError, TypeError):
                return genai.Client(api_key=api_key)

        client = _llm_transport.get_registry().acquire(
            _llm_transport.build_gemini_identity(provider=provider, api_key=api_key),
            _build_gemini_client)
        config_kwargs = {"temperature": temperature, "system_instruction": system_prompt}
        if response_format and cfg.get("capabilities", {}).get("supports_response_format"):
            config_kwargs["response_mime_type"] = "application/json"
        try:
            res = client.models.generate_content(
                model=model,
                contents=user_prompt,
                config=genai.types.GenerateContentConfig(**config_kwargs),
            )
            out_text = (res.text or "").strip()
            usage = getattr(res, "usage_metadata", None)
            run = trace.get_active_run(runtime_job_id)
            if run and started_at is not None:
                run.record_model_call(
                    provider=provider,
                    model=model,
                    duration_s=time.perf_counter() - started_at,
                    request_id=getattr(res, "request_id", None),
                    input_chars=len(system_prompt or "") + len(user_prompt or ""),
                    output_chars=len(out_text),
                    prompt_tokens=getattr(usage, "prompt_token_count", 0) if usage else 0,
                    completion_tokens=getattr(usage, "candidates_token_count", 0) if usage else 0,
                    cached_tokens=getattr(usage, "cached_content_token_count", 0) if usage else 0,
                    total_tokens=getattr(usage, "total_token_count", 0) if usage else 0,
                    ttft_s=None,
                    metadata={"usage_available": usage is not None},
                )
        finally:
            if not _llm_transport.client_pool_enabled():
                # Escape hatch: without a pool nobody owns the client, so this
                # request must release it instead of leaking one per call.
                _llm_transport.close_client(client)
        return out_text

    # OpenAI 官方与所有 OpenAI 兼容接口共用 SDK 路由
    kwargs = {"api_key": api_key, "timeout": _llm_transport.OPENAI_TIMEOUT,
              "max_retries": _llm_transport.OPENAI_MAX_RETRIES}
    resolved_base = resolve_openai_base_url(provider, base_url)
    resolved_reasoning = (reasoning_effort if reasoning_effort is not None
                          else getattr(_LLM_CTX, "reasoning_effort", None))
    if resolved_reasoning and resolved_reasoning not in reasoning_effort_options(
            provider, model):
        # The thread setting belongs to the translator role.  A separate
        # reviewer or an unknown relay model must not inherit it accidentally.
        resolved_reasoning = None
    if resolved_base:
        kwargs["base_url"] = resolved_base
    proxy_bypass = bool(cfg.get("proxy_bypass"))

    def _build_openai_client():
        build_kwargs = dict(kwargs)
        # Keep SDK defaults on the ordinary route and the existing bypass
        # transport on the proxy-free route; hooks never inspect bodies/headers.
        client_type = httpx.Client if proxy_bypass else DefaultHttpxClient
        client_kwargs = {"timeout": _llm_transport.OPENAI_TIMEOUT,
                         "event_hooks": {"request": [_trace_http_request],
                                         "response": [_trace_http_response]}}
        if proxy_bypass:
            client_kwargs["trust_env"] = False
        bypass_client = client_type(**client_kwargs)
        build_kwargs["http_client"] = bypass_client
        try:
            return OpenAI(**build_kwargs)
        except BaseException:
            if bypass_client is not None:
                bypass_client.close()
            raise

    client = _llm_transport.get_registry().acquire(
        _llm_transport.build_openai_identity(
            provider=provider, base_url=resolved_base, proxy_bypass=proxy_bypass,
            api_key=api_key),
        _build_openai_client)
    request = {
        "model": model,
        "messages": [{"role": "system", "content": system_prompt},
                     {"role": "user", "content": user_prompt}],
    }
    # Reasoning models may reject sampling controls such as temperature;
    # when the user explicitly selects reasoning, omit it for official
    # recognized reasoning families and let the model's policy take over.
    if resolved_reasoning:
        request["reasoning_effort"] = resolved_reasoning
        if not str(model or "").lower().startswith(
                ("o1", "o3", "o4-mini", "gpt-5", "gpt-6",
                 "deepseek-r1", "qwq")):
            request["temperature"] = temperature
    else:
        request["temperature"] = temperature
    if provider == "OpenAI" and str(model or "").lower().startswith(
            ("o1", "o3", "o4", "gpt-5", "gpt-6")) and "-chat" not in str(model).lower():
        if resolved_reasoning != "none":
            request.pop("temperature", None)
    if response_format and cfg.get("capabilities", {}).get(
            "supports_response_format"):
        request["response_format"] = response_format
    if provider == "OpenCode Go":
        session_id = runtime_job_id or getattr(_LLM_CTX, "opencode_session", None)
        if not session_id:
            session_id = str(uuid.uuid4())
            _LLM_CTX.opencode_session = session_id
        request["extra_headers"] = {
            "User-Agent": "Folith/0.5.0",
            "x-opencode-session": str(session_id),
        }
    try:
        res = client.chat.completions.create(**request)
        out_text = (res.choices[0].message.content or "").strip()
        usage = getattr(res, "usage", None)
        pt_details = getattr(usage, "prompt_tokens_details", None) if usage else None
        run = trace.get_active_run(runtime_job_id)
        if run and started_at is not None:
            run.record_model_call(
                provider=provider,
                model=model,
                duration_s=time.perf_counter() - started_at,
                request_id=getattr(res, "id", None),
                input_chars=len(system_prompt or "") + len(user_prompt or ""),
                output_chars=len(out_text),
                prompt_tokens=getattr(usage, "prompt_tokens", 0) if usage else 0,
                completion_tokens=getattr(usage, "completion_tokens", 0) if usage else 0,
                cached_tokens=getattr(pt_details, "cached_tokens", 0) if pt_details else 0,
                total_tokens=getattr(usage, "total_tokens", 0) if usage else 0,
                ttft_s=None,
                metadata={"usage_available": usage is not None},
            )
    finally:
        if not _llm_transport.client_pool_enabled():
            # Escape hatch: without a pool nobody owns the client, so this
            # request must release it instead of leaking one per call.
            _llm_transport.close_client(client)
    return out_text



def test_provider(provider, api_key, model, base_url=None):
    """连通性测试：发送一个最小请求，返回 (ok, message)。"""
    t0 = time.time()
    try:
        out = call_llm(provider, api_key, model,
                       "你是连接测试助手。", "请只回复两个字：OK",
                       temperature=0.0, base_url=base_url)
    except Exception as exc:
        return False, provider_error_message(exc)
    elapsed = time.time() - t0
    if not (out or "").strip():
        return False, "返回内容为空（请检查 API Key / 模型名 / 余额）"
    return True, f"响应「{(out or '').strip()[:24]}」· 耗时 {elapsed:.1f}s"



def fetch_provider_models(provider, api_key, base_url=None):
    """获取服务商的文本生成模型目录，返回 (ok, models, message)。"""
    cfg = PROVIDERS.get(provider) or {}
    if cfg.get("kind") == "gemini":
        return _fetch_gemini_models(api_key)
    if cfg.get("kind") not in ("openai", "openai_compat"):
        return False, [], "当前服务商不支持 OpenAI 兼容模型目录"
    if not (api_key or "").strip() and provider not in ("OpenCode Go", "OpenRouter"):
        return False, [], "请先填写 API Key"
    try:
        resolved_base = resolve_openai_base_url(provider, base_url)
    except ValueError as exc:
        return False, [], str(exc)
    if not resolved_base and provider == "OpenAI":
        resolved_base = "https://api.openai.com/v1"
    if not resolved_base:
        return False, [], "请先填写 API 地址"
    try:
        headers = {"Accept": "application/json"}
        if (api_key or "").strip():
            headers["Authorization"] = f"Bearer {api_key.strip()}"
        transport = {"trust_env": False} if cfg.get("proxy_bypass") else {}
        directory_url = f"{resolved_base}/models"
        if provider == "SiliconFlow":
            directory_url += "?sub_type=chat"
        response = httpx.get(
            directory_url, headers=headers,
            timeout=20.0, **transport,
        )
        response.raise_for_status()
        payload = response.json()
    except httpx.HTTPStatusError as exc:
        return False, [], provider_error_message(exc, "获取模型失败")
    except httpx.HTTPError as exc:
        return False, [], provider_error_message(exc, "获取模型失败")
    except ValueError:
        return False, [], "获取模型失败：服务商返回了无法解析的响应。"

    rows = payload.get("data") if isinstance(payload, dict) else payload
    if not rows and isinstance(payload, dict):
        rows = payload.get("models")
    if not isinstance(rows, list):
        return False, [], "模型目录为空或返回格式不受支持"
    models = []
    for row in rows or []:
        if isinstance(row, dict):
            architecture = row.get("architecture") or {}
            inputs = row.get("input_modalities", architecture.get("input_modalities"))
            outputs = row.get("output_modalities", architecture.get("output_modalities"))
            if (inputs is not None and "text" not in inputs) or (outputs is not None and "text" not in outputs):
                continue
        model_id = row if isinstance(row, str) else (
            row.get("id") or row.get("model") or row.get("name")
            if isinstance(row, dict) else None)
        if model_id and str(model_id).strip():
            models.append(str(model_id).strip())
    if provider == "OpenAI":
        models = [model for model in models if model.startswith(("gpt-", "chatgpt-", "o1", "o3", "o4", "ft:gpt-", "ft:o"))
                  and not any(part in model for part in (
                      "audio", "realtime", "transcribe", "tts", "image", "codex", "-pro", "deep-research"))]
    if provider == "OpenCode Go":
        # Go's public catalog has IDs but no endpoint metadata. These documented
        # families use chat/completions; Responses/Messages need separate adapters.
        models = [model for model in models if model.startswith((
            "deepseek-", "glm-", "kimi-", "longcat-", "mimo-", "hy", "space-bunny-"))]
    models = sorted(set(models), key=str.casefold)
    if not models:
        return False, [], "模型目录为空或返回格式不受支持"
    return True, models, f"已获取 {len(models)} 个可用模型"



# ================= 确定性检查（对齐 localize-anything 的机械检查）=================
PRESERVE_RE = re.compile(
    r'(?P<placeholder>%[sd]|%1\$[sd]|\{[A-Za-z_][A-Za-z0-9_]*\}|\{\{[A-Za-z_][A-Za-z0-9_]*\}\})'
    r'|(?P<url>https?://\S+|www\.\S+)'
    r'|(?P<email>[\w.+-]+@[\w-]+(?:\.[\w-]+)+)'
    r'|(?P<doi>10\.\d{4,9}/[^\s]+)'
    r'|(?P<citation>\[\d+(?:[-,]\s*\d+)*\])',
    re.IGNORECASE,
)

PRESERVE_SEVERITY = {
    "placeholder": "blocking",   # 占位符损坏 = 结构破坏，绝不可自动放行
    "url": "actionable",
    "email": "actionable",
    "doi": "actionable",
    "citation": "actionable",
}



def check_translation_batch(sources, targets, glossary, target_lang,
                            section_profile=None):
    """确定性检查一批译文：空译、保留项丢失、源语残留、锁定术语合规（scope 感知）。"""
    findings = []
    # Phase 8 §9：术语表的 normalize + 已编译视图在批次内只派生一次。以前
    # `check_glossary_compliance` 每个 segment 都自己 `normalize_glossary(glossary)`，
    # 一批 4 段就是 4 次 O(G) 归一化 + 4 轮全表术语扫描。这里显式预派生（同内容同
    # revision 会命中已编译视图缓存），结果与逐段自行派生完全一致。
    normalized_glossary = normalize_glossary(glossary)
    from folith.terminology import compiled_glossary as _compiled_glossary
    glossary_matcher = _compiled_glossary(normalized_glossary)
    for i, (src, tgt) in enumerate(zip(sources, targets)):
        raw_src, raw_tgt = src, tgt
        src = "" if src is None else str(src)
        tgt = "" if tgt is None else str(tgt)
        if not tgt.strip():
            findings.append(_deterministic_finding(
                i, "blocking", "completeness", "译文为空",
                "原文存在，但当前段落没有任何译文内容，无法确认信息是否被完整传达。",
                "补译本段后，检查是否覆盖原文的全部句子、限制条件和专有名词。",
                source_span=src, target_span=""))
            continue
        target_report = _translation_target.validate_translation_target(
            raw_src, raw_tgt, segment_index=i)
        for issue in target_report["issues"]:
            findings.append(_deterministic_finding(
                i, "blocking", "format_integrity",
                issue["message"],
                "目标文本仍带有模型 transport 或解释包装，不能作为普通正文交付。",
                "重新解析或重新翻译本段，只保留最终目标文本后再检查。",
                source_span=src, target_span=tgt,
                reason=issue["message"], kind="translation_target_invariant",
                invariant_code=issue["code"]))
        # 完整性检查：拦截截断译文。
        # 实测根因：审校/修复环节的整段替换把长段译文换成了一句修正。
        if is_incomplete_translation(src, tgt):
            reason = (f"疑似漏译/截断：原文 {len(src)} 字符/{_count_sentences(src)} 句，"
                      f"译文仅 {len(tgt.strip())} 字符/{_count_sentences(tgt)} 句")
            findings.append(_deterministic_finding(
                i, "blocking", "completeness", "译文疑似遗漏或被截断",
                "原文长度和句子数量与当前译文不匹配，译文可能没有覆盖完整内容。",
                "对照原文逐句补齐缺失内容，并确认修复后的译文通过完整性复验。",
                source_span=src, target_span=tgt, reason=reason))
        source_words = re.findall(r"[A-Za-z]+", src or "")
        whole_source_preserved = any(
            str(entry.get("behavior") or "") == "preserve"
            and str(entry.get("status") or "") == "locked"
            and re.sub(r"\s+", " ", str(entry.get("source") or "")).strip().casefold()
            == re.sub(r"\s+", " ", src or "").strip().casefold()
            for entry in glossary or [])
        if target_lang != "English" and len(source_words) >= 2 \
                and re.sub(r"\s+", " ", src or "").strip().casefold() \
                == re.sub(r"\s+", " ", tgt or "").strip().casefold() \
                and not whole_source_preserved:
            findings.append(_deterministic_finding(
                i, "actionable", "translation_completion", "译文仍保留整段源文",
                "当前译文与英文原文完全相同，除非这是明确的保留项，否则可能尚未完成翻译。",
                "确认该段是否属于项目规定的保留内容；若不是，请重新翻译并保留专有名词的必要形式。",
                source_span=src, target_span=tgt,
                reason="译文与英文源段完全相同，疑似整段未翻译"))
        for token, kind in extract_preserved_tokens(src).items():
            if token not in tgt:
                severity = PRESERVE_SEVERITY.get(kind, "actionable")
                findings.append(_deterministic_finding(
                    i, severity, "format_integrity",
                    f"译文遗漏必须保留的{kind}「{token}」",
                    f"原文包含必须保留的 {kind}「{token}」，但当前译文中找不到该内容。",
                    "补回该保留项后重新检查占位符、链接或引用格式，确认其余译文未被破坏。",
                    source_span=token, target_span="",
                    reason=f"保留项 {kind}「{token}」在译文中丢失", kind=kind))
        for residual, sev in find_residuals(src, tgt, target_lang):
            findings.append(_deterministic_finding(
                i, sev, "source_language_residue", f"译文残留源语片段「{residual}」",
                "当前译文仍包含原文语言片段，可能意味着该部分未完成翻译，或未经确认地保留了源语。",
                "确认该片段是否为有意保留的专名或术语；若不是，请翻译后检查术语和上下文一致性。",
                source_span=residual if residual in src else None,
                target_span=residual, reason=f"疑似残留源语片段「{residual}」",
                kind="source_residue", detected_text=residual))
        findings.extend(check_glossary_compliance(
            src, tgt, glossary, segment_id=i, section_profile=section_profile,
            entries=normalized_glossary, matcher=glossary_matcher))
        for f in findings:
            if "segment_index" not in f:
                f["segment_index"] = i
            if "segment_id" not in f or f.get("segment_id") is None:
                f["segment_id"] = i
    return findings



# ================= 语义批次（对齐 localize-anything 的上下文批次）=================
BATCH_SIZE = 4
MAX_BATCH_CHARS = 1600
TRANSLATION_MAX_BATCH_CHARS = 2400
DEFAULT_TRANSLATION_CONCURRENCY = 1

# ================= 权威快照节奏（translation hot path）=================
# `state.json` 始终是唯一的权威完整快照；这里只决定自动翻译过程中**多久**把它
# 落一次盘。每批一次全量重写会让累计写入量随文档长度平方增长（5000 段实测
# 655× 写放大），而崩溃时最多只会重算"尚未落盘的批次"。
#
# 上界由两个触发器**先到者**决定：
#   * 每 N 个批次（every_batches）；
#   * 距上次落盘超过 T 秒（max_interval_s）。
# 因此最坏崩溃窗口 = min(N 个批次, T 秒) 内完成的工作，是一个**硬上限**，
# 而不是"可能丢掉几十分钟"。
#
# 以下事件**永远**立即落盘，不受节奏影响：取消、致命失败、阶段结束、进入/离开
# 审校、人工编辑、术语冻结/变更、交付与定稿、显式保存、worker 交接/退出。
STATE_SNAPSHOT_EVERY_BATCHES = 8
STATE_SNAPSHOT_MAX_INTERVAL_S = 3.0
#: Replay budget, not a batch count.  With the journal carrying every committed
#: batch, a crash costs *replay* rather than *re-translation*, so the cadence no
#: longer has to bound model work -- it only has to bound how much journal a
#: recovery reads.  That is a byte budget, and using one is what decouples the
#: snapshot count from the segment count: the old batch-count cadence produced
#: O(segments) full snapshots, each of size O(segments).
STATE_SNAPSHOT_JOURNAL_BYTES = 8 * 1024 * 1024
#: The part that actually flattens the curve.  A *constant* replay budget gives
#: a *constant* number of snapshots per byte of journal -- and the journal is
#: O(document) -- so the snapshot count stays O(document) and the amplification
#: stays O(document), just with a smaller constant.  Measured: 8 MiB fixed gave
#: 1 / 3 / 10 / 48 snapshots at 1k / 5k / 10k / 20k segments, i.e. still linear.
#:
#: Making the budget a multiple of the snapshot's own size fixes the shape: the
#: snapshot grows with the document, so the budget does too, and the number of
#: snapshots grows only logarithmically.  Replay memory stays O(document) --
#: which is what the state already costs to hold -- so this bounds the thing the
#: budget existed to bound, without reintroducing the original O(n^2) writes.
STATE_SNAPSHOT_JOURNAL_FRACTION = 4.0
#: The second half of the fix, and the half that actually binds.
#:
#: Scaling the byte budget alone changed nothing measurable: measured with the
#: fraction budget but the 3s interval still live, the snapshot counts were
#: 1 / 3 / 11 / 64 at 1k / 5k / 10k / 20k -- *worse* than the constant budget's
#: 1 / 3 / 10 / 48, and still linear.  The reason is that the wall-clock trigger
#: was the binding one all along: 64 snapshots over a 209s run is one snapshot
#: every 3.3s, and 11 over 32.5s is one every 3.0s.  Run time is O(document), so
#: a fixed-interval trigger makes the snapshot count O(document) no matter what
#: the byte budget says.
#:
#: The interval trigger is a leftover from the pre-journal design, where the
#: only way to bound lost work was to bound *time* -- a crash could only be
#: repaired by re-translating, so the snapshot had to be fresh in seconds.  With
#: the journal every committed batch is already fsync'd, so a crash loses zero
#: committed work and the snapshot's only remaining job is to compact the
#: journal.  Compaction is a function of the journal's *size*, not of elapsed
#: time, so with the journal on the time trigger is disabled by default and the
#: cadence becomes purely byte-driven -- which is what makes it O(1).
#:
#: ``0.0`` means "no wall-clock trigger".  An explicit
#: ``FOLIOTHREAD_STATE_SNAPSHOT_INTERVAL_S`` still wins, so the old behaviour
#: stays reachable for benchmarks and for the pre-journal regression tests.
STATE_SNAPSHOT_JOURNAL_MAX_INTERVAL_S = 0.0


def state_journal_enabled():
    """Whether the translation hot path journals its commits.

    Off restores the bounded batch-count cadence verbatim, which is the
    pre-journal behaviour and is always correct -- just more expensive.
    """
    raw = os.environ.get("FOLITH_STATE_JOURNAL")
    if raw is None:
        raw = os.environ.get("FOLIOTHREAD_STATE_JOURNAL")
    if raw is None:
        return True
    return str(raw).strip().lower() not in {"0", "off", "false", "no"}


def state_snapshot_cadence(total_batches=None, journal=None):
    """Resolve the automatic-snapshot cadence (env overridable for benchmarks).

    Two triggers exist and whichever fires first wins:

    ``every_batches``
        The bounded batch-count cadence.  ``0`` means "no batch-count trigger",
        which is the default once the journal is on.  It is additionally capped
        at ~10% of the document, so a crash on a *short* job cannot throw away a
        large fraction of it.
    ``journal_bytes``
        Snapshot once this many bytes of un-applied journal have accumulated.
        ``0`` means "no byte trigger", which is the default when the journal is
        off.

    Keeping both means the old behaviour is exactly recoverable (journal off, or
    an explicit ``FOLIOTHREAD_STATE_SNAPSHOT_EVERY``), which is what makes the
    change auditable rather than a leap of faith.
    """
    if journal is None:
        journal = state_journal_enabled()
    every = STATE_SNAPSHOT_EVERY_BATCHES
    interval = STATE_SNAPSHOT_MAX_INTERVAL_S
    journal_bytes = STATE_SNAPSHOT_JOURNAL_BYTES
    raw = os.environ.get("FOLIOTHREAD_STATE_SNAPSHOT_EVERY")
    if raw:
        try:
            every = max(1, int(raw))
        except (TypeError, ValueError):
            pass
    else:
        # No explicit cadence: the journal, when present, replaces the
        # batch-count trigger entirely.
        every = 0 if journal else STATE_SNAPSHOT_EVERY_BATCHES
    raw = os.environ.get("FOLIOTHREAD_STATE_SNAPSHOT_INTERVAL_S")
    if raw:
        try:
            interval = max(0.0, float(raw))
        except (TypeError, ValueError):
            pass
    elif journal:
        # The journal makes the time trigger redundant *and* harmful: it is the
        # one trigger whose firing rate is proportional to run time, i.e. to
        # document length.  See STATE_SNAPSHOT_JOURNAL_MAX_INTERVAL_S.
        interval = STATE_SNAPSHOT_JOURNAL_MAX_INTERVAL_S
    fraction = STATE_SNAPSHOT_JOURNAL_FRACTION
    raw = os.environ.get("FOLITH_STATE_SNAPSHOT_JOURNAL_BYTES")
    if raw is None:
        raw = os.environ.get("FOLIOTHREAD_STATE_SNAPSHOT_JOURNAL_BYTES")
    if raw:
        try:
            journal_bytes = max(0, int(raw))
        except (TypeError, ValueError):
            pass
    raw = os.environ.get("FOLITH_STATE_SNAPSHOT_JOURNAL_FRACTION")
    if raw is None:
        raw = os.environ.get("FOLIOTHREAD_STATE_SNAPSHOT_JOURNAL_FRACTION")
    if raw:
        try:
            fraction = max(0.0, float(raw))
        except (TypeError, ValueError):
            pass
    if not journal:
        journal_bytes = 0
        fraction = 0.0
    every = max(0, int(every))
    if every and total_batches:
        every = max(1, min(every, -(-int(total_batches) // 10)))
    return {
        "every_batches": every,
        "max_interval_s": float(interval),
        "journal_bytes": int(journal_bytes),
        "journal_fraction": float(fraction),
    }


def state_snapshot_crash_window(batch_size=None, total_batches=None):
    """Worst-case segments a crash can force back through the pipeline.

    Two regimes, and the number is honest about which one is active:

    * with the journal, a crash forces **zero** re-translation -- recovery
      replays committed batches from ``state_journal.jsonl`` without calling a
      model, so the window is bounded in journal bytes, not segments;
    * without it, the bounded cadence still applies: at most ``every_batches``
      batches of committed-but-unsnapshotted translation.

    Reported in the benchmark output and asserted by the persistence tests.
    """
    cadence = state_snapshot_cadence(total_batches=total_batches)
    if not cadence["every_batches"]:
        return 0
    return max(1, int(batch_size or BATCH_SIZE)) * cadence["every_batches"]


# ---------------------------------------------------------------------------
# State journal: the per-batch delta of the translation hot path
# ---------------------------------------------------------------------------
# The record shape is fixed and named.  It is deliberately *not* a diff: the
# three collections below are the only ones the translation commit touches, and
# each is described in the way it actually changes.
#
#   set     -- whole-value changes (replaced, or mutated in place inside the
#              batch).  Compared against a pre-batch snapshot of just these keys.
#   extend  -- append-only lists.  The pre-append length travels with the items
#              so a double application is detected instead of duplicating.
#   patch   -- in-place updates to items earlier records already carried, keyed
#              by the item's own identity.  Produced by the mutator itself.
_JOURNAL_SET_KEYS = (
    "translation_truth", "dependency_impact", "academic_state", "final_qa",
    "p3_done", "p3_md", "p3_sections", "delivery_status",
    "delivery_approved_by_human", "delivery_approval", "stage",
    "review_stats", "knowledge_candidates", "translation_continuity",
    "entity_registry", "knowledge_feedback_policy",
    "knowledge_feedback_failures", "has_blocking",
)
_JOURNAL_EXTEND_KEYS = (
    "findings", "review_evidence", "repair_overlays", "human_actions",
    "knowledge_events", "context_packet_log", "glossary_injection_log",
)


def new_journal_patches():
    """Collector for the typed in-place patches one batch applies."""
    return {}


def journal_baseline(state):
    """Capture the named hot-path fields before one batch commits.

    A deep copy is required for the keys that are mutated in place (for example
    ``review_stats`` and ``knowledge_feedback_policy``); a shallow one would
    alias the live object and report "unchanged" forever.
    """
    return {
        "set": {key: copy.deepcopy(state[key])
                for key in _JOURNAL_SET_KEYS if key in state},
        "len": {key: len(state.get(key) or []) for key in _JOURNAL_EXTEND_KEYS},
    }


def build_batch_journal_record(state, baseline, patches, *, batch, offset,
                               segment_ids, pairs):
    """Build one ``translation_batch_committed`` record, or ``None``.

    ``None`` means the delta cannot be described faithfully -- an unaddressable
    item, or a list a normalizer shrank mid-batch.  The caller then takes an
    immediate snapshot instead.  That fallback is the whole safety argument for
    this optimisation: an undescribable delta costs performance, never
    durability.
    """
    if not baseline or patches.get("incomplete"):
        return None
    record = {
        "schema": _state_journal.JOURNAL_SCHEMA_VERSION,
        "mutation": _state_journal.MUTATION_TRANSLATION_BATCH,
        "mutation_id": uuid.uuid4().hex,
        "ts": _finalization.now_iso(),
        "batch": int(batch),
        "offset": int(offset),
        "segment_ids": list(segment_ids),
        "pairs": list(pairs),
    }
    set_values = {}
    for key in _JOURNAL_SET_KEYS:
        if key not in state:
            continue
        current = state[key]
        if key in baseline["set"] and current == baseline["set"][key]:
            continue
        set_values[key] = current
    if set_values:
        record["set"] = set_values
    extend_values = {}
    for key in _JOURNAL_EXTEND_KEYS:
        items = state.get(key)
        if not isinstance(items, list):
            return None
        start = baseline["len"][key]
        if len(items) < start:
            return None
        if len(items) > start:
            extend_values[key] = {"base": start, "items": items[start:]}
    if extend_values:
        record["extend"] = extend_values
    patch_values = {key: value for key, value in patches.items()
                    if value and key != "incomplete"}
    if patch_values:
        record["patch"] = patch_values
    return record


def make_batches(paragraphs, batch_size=BATCH_SIZE, max_chars=MAX_BATCH_CHARS,
                 semantic_units=None):
    """把段落聚成批次；提供 semantic_units 时不跨单元边界。"""
    if semantic_units:
        batches = []
        ranges = []
        for unit in semantic_units:
            if not isinstance(unit, dict):
                continue
            try:
                start = int(unit.get("start_segment"))
                end = int(unit.get("end_segment"))
            except (TypeError, ValueError):
                continue
            if 0 <= start <= end < len(paragraphs):
                ranges.append((start, end))
        for start, end in sorted(ranges):
            batches.extend(make_batches(paragraphs[start:end + 1], batch_size, max_chars))
        covered = {index for start, end in ranges for index in range(start, end + 1)}
        range_size = sum(end - start + 1 for start, end in ranges)
        if batches and len(covered) == len(paragraphs) and range_size == len(paragraphs):
            return batches
    batches, cur, n = [], [], 0
    for p in paragraphs:
        if cur and (len(cur) >= batch_size or n + len(p) > max_chars):
            batches.append(cur)
            cur, n = [], 0
        cur.append(p)
        n += len(p)
    if cur:
        batches.append(cur)
    return batches


def _translation_evidence_index(
    paras, pairs, batch_pairs, glossary, document_profile,
    document_synopsis, section_digests, findings_all, blind=False,
    candidate_targets=None, blocked_source_indexes=None,
):
    # The index copies every committed pair (`list(pairs)`) plus the whole
    # glossary and findings list, so its cost is a function of document length,
    # not batch length.  Recording items-scanned is what makes that visible.
    with _critical_path.phase("review_input_evidence_index") as _index_sink:
        _index_sink.traversed(
            len(pairs) + len(batch_pairs) + len(glossary) + len(findings_all))
        blocked = {int(index) for index in (blocked_source_indexes or set())
                   if isinstance(index, int) and not isinstance(index, bool)}
        safe_paras = ["" if index in blocked else source
                      for index, source in enumerate(paras or [])]
        safe_pairs = []
        for index, pair in enumerate(list(pairs or []) + list(batch_pairs or [])):
            if index not in blocked or not isinstance(pair, dict):
                safe_pairs.append(pair)
                continue
            safe_pair = dict(pair)
            for key in ("source", "target", "initial_target", "accepted_target",
                        "suggested_target"):
                safe_pair[key] = ""
            safe_pairs.append(safe_pair)
        safe_findings = [
            finding for finding in findings_all or []
            if not (isinstance(finding, dict)
                    and finding.get("type") == "source_quality_gate"
                    and isinstance(finding.get("segment_index"), int)
                    and not isinstance(finding.get("segment_index"), bool)
                    and finding.get("segment_index") in blocked)
        ]
        return _translation_evidence.TranslationEvidenceIndex(
            safe_paras, safe_pairs, glossary, document_profile,
            document_synopsis, section_digests, safe_findings,
            blind=blind, candidate_targets=candidate_targets)


def _review_finding_record(finding, review_event_id):
    """Project a Translation Core finding onto the existing persisted surface."""
    segment_id = finding.get("segment_id")
    severity = finding.get("severity")
    return {
        **finding,
        "segment_id": segment_id, "segment_index": segment_id,
        "severity": severity, "type": "review",
        "category": finding.get("category") or "semantic_accuracy",
        "summary": str(finding.get("summary") or finding.get("reason") or
                       "审校发现问题"),
        "source_span": finding.get("source_span"),
        "target_span": finding.get("target_span"),
        "explanation": finding.get("explanation"),
        "recommendation": finding.get("recommendation"),
        "confidence": finding.get("confidence"),
        "detector": finding.get("detector") or "Semantic QA",
        "diagnostic_version": finding.get("diagnostic_version"),
        "reason": str(finding.get("reason") or finding.get("summary") or
                      "审校发现问题"),
        "evidence_refs": list(finding.get("evidence_refs") or []),
        "review_event_id": review_event_id,
    }


def _runtime_review_context(
    state, offset, batch_len, glossary_text, style_rules, target_lang,
):
    # Phase 7 (§5) verified this call site rather than rewrote it.  It must keep
    # ``select_target_context``: ``limit=2`` really does truncate here, so the
    # level-priority rule is load-bearing, and ``recent_target_context`` (the
    # generation-side fast path) would silently change which previous targets
    # the reviewer sees.  The contract is therefore still *committed / reviewed*
    # context: ``state["pairs"]`` filtered by the caller below to
    # ``human_accepted | reviewed | tm_approved`` -- never provisional output.
    #
    # Phase 5 already removed the prefix materialisation from the helper
    # (``list(pairs)[:offset]`` plus a dict per historical segment); what
    # remains is an O(offset) *dereference* walk that cannot early-exit, because
    # only the top provenance slot is allowed to stop the scan.  That is a
    # semantic property, not slack: see docs/translation-review-context-scaling.md.
    with _critical_path.phase("review_input_previous_context") as _ctx_sink:
        _ctx_sink.traversed(max(0, int(offset or 0)))
        paras = state.get("paras") or []
        pairs = state.get("pairs") or []
        profile = state.get("document_profile") or {}
        digests = state.get("section_digests") or []
        previous_target = _context.select_target_context(pairs, offset, limit=2)
        dependency_ids = _runtime_review_dependency_ids(
            state, offset, batch_len, previous_target)
        target_dependency_ids = set(range(
            max(0, offset), max(0, offset) + max(0, batch_len)))
        target_dependency_ids.update(
            int(item["segment_index"])
            for item in previous_target or []
            if isinstance(item, dict)
            and isinstance(item.get("segment_index"), int)
            and not isinstance(item.get("segment_index"), bool)
        )
        return {
            "document_profile": profile,
            "document_synopsis": state.get("document_synopsis") or {},
            "section_profile": _batch_section_profile(profile, offset, batch_len) or {},
            "section_digest": _context.digest_for_segment(digests, offset) or {},
            "previous_source_context": list(paras[max(0, offset - 2):offset]),
            "previous_target_context": [
                dict(item) for item in previous_target
                if item.get("level") in {"human_accepted", "reviewed", "tm_approved"}
            ],
            "next_source_context": list(
                paras[offset + batch_len:offset + batch_len + 2]),
            "target_language": target_lang,
            "style_constraints": style_rules or "",
            "advisory_terminology_context": glossary_text or "",
            "_dependency_segment_ids": dependency_ids,
            "_dependency_target_segment_ids": sorted(target_dependency_ids),
        }


def _runtime_review_dependency_ids(state, offset, batch_len, previous_target=None):
    """Return segment IDs whose source/target can appear in review context."""
    paragraphs = state.get("paras") or []
    ids = set(range(max(0, offset), max(0, offset) + max(0, batch_len)))
    ids.update(range(max(0, offset - 2), max(0, offset)))
    ids.update(range(
        max(0, offset) + max(0, batch_len),
        min(len(paragraphs), max(0, offset) + max(0, batch_len) + 2),
    ))
    ids.update(
        int(item["segment_index"])
        for item in previous_target or []
        if isinstance(item, dict)
        and isinstance(item.get("segment_index"), int)
        and not isinstance(item.get("segment_index"), bool)
    )
    return sorted(ids)


def _batch_section_profile(document_profile, offset, batch_len):
    """按全局段区间匹配 section profile（用于相关术语的 section:<id> scope）。"""
    if not document_profile:
        return None
    for sec in document_profile.get("sections") or []:
        if not isinstance(sec, dict):
            continue
        try:
            start = int(sec.get("start_segment"))
            end = int(sec.get("end_segment"))
        except (TypeError, ValueError):
            continue
        if start <= offset and offset + batch_len - 1 <= end:
            return sec
    return None


# ================= 翻译记忆（对齐 localize-anything 的 TM：仅收录审校通过段落）=================
DEFAULT_TARGET_LANG = "简体中文"



# ---------------- 翻译记忆的作用域键 ----------------
# 一条翻译记忆的身份是「目标语言 + 原文」，**不是**「原文」。
# 同一段原文在 Français 与简体中文任务里的正确译文不同；只用原文当键，一种
# 语言的译文就会串进另一种语言的任务，并以"已审校"的名义静默通过交付检查。
# 因此作用域键是结构性的：不给出目标语言就**无法**命中任何条目。
# 定义放在 `folith.translation_memory`，核心层与展示层共用同一份格式。
TM_SCOPE_SEP = _tm_scope.TM_SCOPE_SEP
tm_scope_key = _tm_scope.tm_scope_key
tm_unscope_key = _tm_scope.tm_unscope_key
tm_record_language = _tm_scope.tm_record_language
tm_record = _tm_scope.tm_record
tm_put = _tm_scope.tm_put
tm_discard = _tm_scope.tm_discard
tm_legacy_keys = _tm_scope.tm_legacy_keys


# 归一化只做**可证明等价**的处理：同一段文字在不同来源下的排版差异。
# 刻意不做语义近似（编辑距离、词序、同义替换）：翻译记忆是错误放大器，
# 一次错配会复制到整篇文档，所以匹配必须是"同一句话"而不是"相似的话"。
_TM_CHAR_MAP = str.maketrans({
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2212": "-", "\u00ad": "",
    "\u2026": "...", "\u00a0": " ", "\u3000": " ",
})
_TM_SPACE_RE = re.compile(r"\s+")



def tm_lookup(tm, source, index=None, *, target_lang=None):
    """查找翻译记忆条目，返回 (记录, 命中方式)。

    **必须显式给出目标语言**：语言未知（旧条目没有语言标注）或目标语言不同
    -> 不命中。这是发布阻断级的正确性约束，不是可选过滤。

    先精确匹配，再退到归一化匹配。归一化命中是"同一段文字的排版差异"，
    不涉及语义猜测，因此可以安全复用译文；命中方式会被记录，便于审计。
    """
    language = normalize_language(target_lang)
    if not language:
        return None, ""
    cleaned = str(source or "").replace("\n", " ")
    for key in (tm_scope_key(target_lang, cleaned), cleaned):
        if not key:
            continue
        if _tm_hit((tm or {}).get(key), language, key):
            return (tm or {}).get(key), "exact"
    normalized = tm_normalize(cleaned)
    if not normalized:
        return None, ""
    scoped_index = index if index is not None else tm_index(tm, target_lang)
    raw_key = scoped_index.get(normalized)
    if raw_key is None or raw_key == cleaned:
        return None, ""
    record = (tm or {}).get(raw_key)
    if _tm_hit(record, language, raw_key):
        return record, "normalized"
    return None, ""



def load_tm(project_id=None):
    """加载指定项目的翻译记忆并自清洗：非法条目直接丢弃。

    翻译记忆是错误放大器（一次错译会复制到全书），因此加载即消毒，
    防止旧版本或异常写入留下的污染条目继续命中。

    语言无法证明的旧条目**保留在文件里**（不丢用户数据），但不会被
    `tm_lookup` 命中：证明不了目标语言的记忆不能自动复用。可用
    `tm_legacy_keys` 把它们单独列出来。

    `project_id=None` 表示默认项目（即历史的全局记忆）。
    """
    p = tm_path(project_id)
    if p.is_file():
        try:
            raw = json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            return _TranslationMemory()
        return _TranslationMemory({k: v for k, v in raw.items()
                if isinstance(v, dict) and v.get("reviewed")
                and _tm_eligible(tm_unscope_key(k)[1], v.get("target"))})
    return _TranslationMemory()


#: Parent directories whose TM file we have already created in this process.
#: ``mkdir(parents=True, exist_ok=True)`` still costs a syscall when the
#: directory exists, and the promotion path used to pay it once per batch.
_TM_DIR_READY = set()
_TM_DIR_LOCK = threading.Lock()



def reset_tm_dir_cache():
    """Forget the directory hint (tests that delete/repoint ``OUTPUT_DIR``)."""
    with _TM_DIR_LOCK:
        _TM_DIR_READY.clear()


def save_tm(tm, project_id=None):
    """持久化项目翻译记忆。

    `tm.json` 是**派生的物化缓存**，不是权威真值：`checkpoint.reconcile_translation_memory`
    会在每次 `translate_stage` 开始时用 `state["pairs"]`（已审校段落）加上尚未确认的
    `tm_promotion_pending` 事件把它重建出来。因此这个函数的写路径可以独立优化，
    而不会改变权威边界。

    写路径与 `state.json` 共用同一个紧凑序列化器与"编码一次再写 bytes"的做法；
    目录创建也只在进程内做一次。
    """
    p = tm_path(project_id)
    with _persistence_profiler.phase("tm_prepare"):
        _ensure_tm_dir(p)
    with _job_lock.exclusive(p.with_suffix(".write.lock")):
        if isinstance(tm, _TranslationMemory):
            # The lock covers re-read, delta merge and replacement across tasks
            # and processes. Unchanged stale entries must not resurrect deletes.
            merged = dict(load_tm(project_id))
            for key in tm.baseline.keys() - tm.keys():
                merged.pop(key, None)
            for key, value in tm.items():
                if key not in tm.baseline or value != tm.baseline[key]:
                    merged[key] = value
        else:
            # Explicit replacement mappings (project import/initialisation).
            merged = dict(tm)
        tmp = p.with_suffix(".json.tmp")
        try:
            with _persistence_profiler.phase("tm_json_serialization") as sink:
                encoded = json.dumps(merged, **_job_repo.COMPACT_JSON).encode("utf-8")
                sink.serialized(len(encoded))
            with _persistence_profiler.phase("tm_file_write") as sink:
                _write_tm_bytes(tmp, encoded)
                sink.written(len(encoded))
            with _persistence_profiler.phase("tm_atomic_replace"):
                tmp.replace(p)
        finally:
            tmp.unlink(missing_ok=True)
        if isinstance(tm, _TranslationMemory):
            tm.baseline = {key: dict(value) for key, value in tm.items()}



# ================= Project：跨任务复用的已确认项目记忆 =================
# Project 层本身是领域模块（folith/project.py，显式接收 root）；
# 这里只是把它接到本地 OUTPUT_DIR，并负责"任务属于哪个项目"的读写。
#
# 身份模型（v0.5 定死）：
#   - 每个项目都有不可变 UUID；display name 只是标签，不是主键，也不进路由；
#   - 系统工作区「未分类」是**真实项目**（真实 UUID + is_system=true），承载
#     所有没有 project_id 的任务；它不能重命名 / 归档 / 删除；
#   - 字符串 "default" 只是历史别名，任何入口都先归一到系统项目 UUID。

SYSTEM_PROJECT_ID = _project.SYSTEM_PROJECT_ID
SYSTEM_PROJECT_NAME = _project.SYSTEM_PROJECT_NAME
# 向后兼容别名：值现在是系统项目的真实 UUID，不再是字符串 "default"。
DEFAULT_PROJECT_ID = _project.DEFAULT_PROJECT_ID
DEFAULT_PROJECT_NAME = _project.DEFAULT_PROJECT_NAME
# 任务归属「未分类」时写进 state["project_id"] 的值就是系统项目 UUID。
UNCLASSIFIED_PROJECT_ID = SYSTEM_PROJECT_ID



# ---- CRUD：创建 / 读取 / 更新 / 归档 / 删除 ----

PROJECT_DESCRIPTION_LIMIT = 600
PROJECT_NAME_LIMIT = 80



# 任务正在运行的状态：此时它已经读取过一个项目的记忆，改归属会让同一个任务
# 在两套记忆之间漂移，因此这些状态一律拒绝移动。
# 集合只有一个来源：`RUNTIME_ACTIVE_STATUSES`（历史上这里曾有一份逐字重复的副本）。
def job_is_active(job_id, state=None):
    """任务是否正在运行（不可改归属）。"""
    try:
        view = build_job_runtime_view(job_id, state)
        status = view.get("runtime_status") or view.get("status")
    except Exception:  # 运行状态不可读时不阻止用户操作，由后续流程兜底
        return False
    return str(status or "") in RUNTIME_ACTIVE_STATUSES



# ================= 翻译 / 修复 / 审校（对齐 localize-anything 的三通道）=================
def _translator_system(glossary_text, style_rules, target_lang):
    return (f"你是一位专业翻译人员，请将用户提供的段落翻译成{target_lang}。\n"
            "先识别原文体裁并沿用其语气；执行学术翻译专家常见工作时，也不得把小说、"
            "回忆录、对话或商业文本改写成论文语体。只翻译可翻译的正文，译文须与原文"
            "一一对应并保持顺序。\n"
            "人名、地名、机构名、品牌名和作品名不应一律保留源语：遵循已提供的术语表、"
            "实体提示和全书既定译法；没有既定译法时，选择目标语言中自然、连贯的表达，"
            "并在后文保持一致。只有明确标为“保留原文”的词条才按其规则原样保留。URL、"
            "邮箱、DOI、引用标注（如 [12]）和其他机器标识符保持准确，不要改动其字符。\n"
            f"{glossary_text}\n"
            f"{style_rules}\n"
            "翻译前请在内部检查：指代与照应关系、术语和专名、理论/非字面用法、"
            "临时造词、长句逻辑关系与修辞功能；优先复用已提供的术语、实体和连续性选择，"
            "人工锁定实体/术语优先于审校、TM 和生成式建议；生成式实体提示不得覆盖锁定术语。"
            "避免明显的逐词直译和词典首义机械替换。\n"
            "不要输出分析过程，只输出最终译文。\n"
            "请严格输出合法的 JSON 字符串数组，不要包含任何解释文字。")


def _invoke_llm(call_fn, provider, api_key, model, system_prompt, user_prompt,
                temperature, response_format=None):
    """Call old and new provider/test doubles through one compatibility gate."""
    kwargs = {"temperature": temperature}
    if response_format is not None:
        kwargs["response_format"] = response_format
    try:
        return call_fn(provider, api_key, model, system_prompt, user_prompt, **kwargs)
    except TypeError:
        try:
            return call_fn(provider, api_key, model, system_prompt, user_prompt,
                           temperature=temperature)
        except TypeError:
            return call_fn(provider, api_key, model, system_prompt, user_prompt)


def _native_translation_response_format(provider, model, expected):
    capabilities = _model_roles.provider_capabilities(PROVIDERS, provider, model)
    if capabilities.get("supports_json_schema") and capabilities.get(
            "supports_response_format"):
        return _translation_protocol.json_schema_for_translations(expected)
    return None


def translate_batch(segments, ctx_prev, ctx_next, glossary_text, style_rules, target_lang,
                    provider, api_key, model, context_packet=None, call_llm_fn=None,
                    prompt_stats=None, result_cache=None):
    """翻译一个语义批次，返回与 segments 等长的译文列表；失败抛出 RuntimeError。

    ``prompt_stats`` (optional list) collects one compact cost record per real
    provider request: counts and a key fingerprint only, never prompt text.

    ``result_cache`` (optional ``folith.translation_cache.TranslationResultCache``)
    reuses an identical, already-paid computation.  The identity is the rendered
    request itself, so any change that would change what the model sees — source,
    glossary, context, style, model, generation settings — is a miss by
    construction.  A cache hit skips the provider call only; ordering, TM,
    review and persistence are untouched.
    """
    numbered = "\n".join(f"{i + 1}. {s}" for i, s in enumerate(segments))
    render_stats = {}
    if context_packet is not None:
        # Keep the system prefix invariant across batches.  Glossary/style are
        # carried once in the ordered context packet below.
        sys_prompt = _translator_system("", "", target_lang)
        user_prompt, render_stats = _context.render_context_packet_with_stats(
            context_packet)
    else:
        sys_prompt = _translator_system(glossary_text, style_rules, target_lang)
        context = ""
        if ctx_prev:
            context += "【前文上下文】：\n" + "\n".join(f"- {s}" for s in ctx_prev) + "\n\n"
        if ctx_next:
            context += "【后文上下文】：\n" + "\n".join(f"- {s}" for s in ctx_next) + "\n\n"
        user_prompt = f"{context}待翻译段落（按序号返回等长译文数组）：\n{numbered}"
        render_stats = {
            "category_chars": {
                "glossary": len(glossary_text or ""),
                "style": len(style_rules or ""),
                "neighbor_context": len(context),
            },
            "source_chars": len(numbered),
        }
    call_fn = call_llm_fn or call_llm
    response_format = _native_translation_response_format(
        provider, model, len(segments))
    cache_key = None
    if result_cache is not None:
        cache_key = _translation_cache.prompt_fingerprint(
            provider=provider, model=model, system_prompt=sys_prompt,
            user_prompt=user_prompt, temperature=0.3,
            response_format=response_format)
        cached = result_cache.get(cache_key)
        if cached is not None and len(cached) == len(segments):
            if prompt_stats is not None:
                prompt_stats.append(_translation_prompt_cost(
                    segments=segments, system_prompt=sys_prompt,
                    user_prompt=user_prompt, render_stats=render_stats,
                    attempt=0, output_chars=sum(len(item) for item in cached),
                    cache="hit"))
            return list(cached)
    # Rate limiting is handled by the shared provider controller: it paces
    # request starts and pushes one cooldown every worker observes, instead of
    # each worker sleeping a fixed 15 s and then stampeding the provider.
    controller = _llm_request_controller(provider, model)
    cancel_check = _llm_cancel_check()
    last_err = None
    for _attempt in range(3):
        if _attempt and _runtime_job_id():
            update_runtime_state(
                _runtime_job_id(),
                event=f"翻译请求未完成，正在重试（第 {_attempt + 1}/3 次尝试）",
                event_name="translation_retry", event_visibility="user",
                event_category="external_wait", phase="retry",
                phase_label=f"正在重试翻译请求（第 {_attempt + 1}/3 次尝试）",
                operation_id="translation_retry",
                operation_label=f"正在重试翻译请求（第 {_attempt + 1}/3 次尝试）",
            )
        try:
            with trace.span("model_wait", operation="provider_gate", job_id=_runtime_job_id()):
                controller.acquire(cancel_check)
            previous_cost = getattr(_LLM_CTX, "prompt_cost", None)
            _LLM_CTX.prompt_cost = _translation_prompt_cost(
                segments, sys_prompt, user_prompt, render_stats, _attempt + 1, 0)
            try:
                res = _invoke_llm(
                    call_fn, provider, api_key, model, sys_prompt, user_prompt,
                    temperature=0.3, response_format=response_format)
            finally:
                _LLM_CTX.prompt_cost = previous_cost
            controller.note_success()
            targets = _translation_protocol.parse_translation_response(
                res, len(segments))
            if result_cache is not None and cache_key is not None:
                # Only a fully parsed, count-matching batch is cacheable.  A
                # partial or malformed reply never reaches this line.
                result_cache.put(cache_key, targets)
            if prompt_stats is not None:
                prompt_stats.append(_translation_prompt_cost(
                    segments=segments, system_prompt=sys_prompt,
                    user_prompt=user_prompt, render_stats=render_stats,
                    attempt=_attempt + 1, output_chars=len(str(res or "")),
                    cache="miss"))
            return targets
        except Exception as e:
            _workflow_errors.raise_if_fatal(e, _runtime_job_id())
            if "任务已请求取消" in str(e):
                raise
            last_err = e
            backoff = controller.note_rate_limited() if is_rate_limited(e) else 0.0
            if _attempt < 2:
                trace.record_retry("translation", "translate_batch", _attempt + 2,
                                   backoff, type(e).__name__, job_id=_runtime_job_id())

    if prompt_stats is not None:
        prompt_stats.append(_translation_prompt_cost(
            segments=segments, system_prompt=sys_prompt, user_prompt=user_prompt,
            render_stats=render_stats, attempt=3, output_chars=0, failed=True))
    raise RuntimeError(
        f"批次翻译失败（{len(segments)} 段）："
        f"{last_err or '模型返回格式异常或数量不匹配'}"
    )


DYNAMIC_CONTEXT_RECORD_LIMIT = 600


def _record_dynamic_context_batch(state, batch_index, offset, segment_count, packet):
    """Append one compact experimental record for a dynamic_v2 batch.

    Counts, candidate ids and provenance only: never prompt text, never secrets.
    Bounded so an experimental run cannot bloat persisted job state.
    """
    diagnostics = (packet or {}).get("dynamic_diagnostics") or {}
    if not diagnostics:
        return None
    try:
        rendered, render_stats = _context.render_context_packet_with_stats(packet)
        prompt_chars = len(rendered)
        category_chars = dict(render_stats.get("category_chars") or {})
    except Exception:
        prompt_chars = 0
        category_chars = {}
    try:
        source_chars = sum(len(str(item or ""))
                           for item in (packet or {}).get("current_batch") or [])
    except Exception:
        source_chars = 0
    record = {
        "batch": int(batch_index),
        "offset": int(offset),
        "segment_count": int(segment_count),
        "experiment_id": diagnostics.get("experiment_id") or "",
        "context_mode": "dynamic_v2",
        "context_budget_tokens": int(diagnostics.get("context_budget_tokens") or 0),
        "offline": bool(diagnostics.get("offline")),
        "ablations": list(diagnostics.get("ablations") or []),
        "candidate_counts": dict(diagnostics.get("candidate_counts") or {}),
        "selected_ids": list(diagnostics.get("selected_ids") or []),
        "selected_by_source": {
            key: list(value) for key, value
            in (diagnostics.get("selected_by_source") or {}).items()},
        "selected_tokens_by_source": dict(
            diagnostics.get("selected_tokens_by_source") or {}),
        "total_selected_tokens": int(diagnostics.get("selected_tokens_total") or 0),
        "rejected_count": int(diagnostics.get("rejected_count") or 0),
        "mandatory_over_budget": bool(diagnostics.get("mandatory_over_budget")),
        "dedupe_removed_count": int(diagnostics.get("dedupe_removed_count") or 0),
        "fallback": bool(diagnostics.get("fallback")),
        "fallback_reason": str(diagnostics.get("fallback_reason") or ""),
        "distant_stats": dict(diagnostics.get("distant_stats") or {}),
        "distant_retrieved": int(diagnostics.get("distant_retrieved") or 0),
        "distant_selected": int(diagnostics.get("distant_selected") or 0),
        "selected_scores": dict(diagnostics.get("selected_scores") or {}),
        "profile": dict(diagnostics.get("profile") or {}),
        "source_chars": int(source_chars),
        "context_chars": int(render_stats.get("context_prefix_chars") or 0)
        if prompt_chars else 0,
        "prompt_chars": int(prompt_chars),
        "prompt_tokens_estimate": _usage_estimate(prompt_chars),
        "source_tokens_estimate": _usage_estimate(source_chars),
        "batch_source_tokens": _usage_estimate(source_chars),
        "wall_s": 0.0,
        "category_chars": category_chars,
    }
    records = state.setdefault("dynamic_context_records", [])
    if not isinstance(records, list):
        records = []
        state["dynamic_context_records"] = records
    records.append(record)
    if len(records) > DYNAMIC_CONTEXT_RECORD_LIMIT:
        del records[:len(records) - DYNAMIC_CONTEXT_RECORD_LIMIT]
    return record


def _usage_estimate(chars):
    """ceil(chars / 4) using the shared usage-ledger estimator."""
    try:
        return int(_usage.estimate_tokens(chars))
    except Exception:
        return 0


def _translation_prompt_cost(segments, system_prompt, user_prompt, render_stats,
                             attempt, output_chars, failed=False, cache="miss"):
    """One compact, text-free prompt cost record for a translation request.

    Counts only.  The prompt text itself is never stored or logged here, so this
    stays usable under the existing privacy/debug conventions.
    """
    category_chars = dict((render_stats or {}).get("category_chars") or {})
    system_chars = len(str(system_prompt or ""))
    user_chars = len(str(user_prompt or ""))
    categories = {
        "instructions": system_chars,
        "source": int((render_stats or {}).get("source_chars") or 0),
    }
    for name, chars in category_chars.items():
        if name == "source":
            continue
        categories[name] = categories.get(name, 0) + int(chars or 0)
    accounted = sum(categories.values()) - system_chars
    if user_chars > accounted:
        categories["other"] = categories.get("other", 0) + (user_chars - accounted)
    return {
        "segment_count": len(segments),
        "source_chars": sum(len(str(item or "")) for item in segments),
        "prompt_chars": system_chars + user_chars,
        "output_chars": int(output_chars or 0),
        "category_chars": categories,
        "attempt": int(attempt),
        "failed": bool(failed),
        "cache": cache,
    }


def _record_translation_prompt_cost(state, records):
    """Aggregate per-request prompt cost into the job state (counts only).

    Callers hold the stage's commit lock, so this stays a plain accumulation.
    """
    records = [item for item in (records or []) if isinstance(item, dict)]
    if not records:
        return
    sink = state.setdefault("translation_prompt_cost", {
        "requests": 0, "failed_requests": 0, "segments": 0,
        "prompt_chars": 0, "output_chars": 0, "source_chars": 0,
        "cache_hits": 0, "cache_misses": 0,
        "category_chars": {}, "requests_by_segment_count": {},
    })
    for item in records:
        sink["requests"] = int(sink.get("requests") or 0) + 1
        if item.get("failed"):
            sink["failed_requests"] = int(sink.get("failed_requests") or 0) + 1
        if str(item.get("cache") or "") == "hit":
            sink["cache_hits"] = int(sink.get("cache_hits") or 0) + 1
        else:
            sink["cache_misses"] = int(sink.get("cache_misses") or 0) + 1
        for key in ("segments", "prompt_chars", "output_chars", "source_chars"):
            source_key = "segment_count" if key == "segments" else key
            sink[key] = int(sink.get(key) or 0) + int(item.get(source_key) or 0)
        for name, chars in (item.get("category_chars") or {}).items():
            bucket = sink.setdefault("category_chars", {})
            bucket[name] = int(bucket.get(name) or 0) + int(chars or 0)
        counts = sink.setdefault("requests_by_segment_count", {})
        size = str(int(item.get("segment_count") or 0))
        counts[size] = int(counts.get(size) or 0) + 1


def repair_batch(sources, targets, findings, glossary_text, style_rules, target_lang,
                 provider, api_key, model, call_llm_fn=None):
    """根据确定性检查发现的问题自动修复一批译文；返回与 sources 等长的译文列表。"""
    numbered = "\n".join(
        f"{i + 1}. 原文：{s}\n   译文：{t}" for i, (s, t) in enumerate(zip(sources, targets)))
    issues = "\n".join(f"- 段落 {f['segment_index'] + 1}: [{f['severity']}] {f['reason']}"
                       for f in findings)
    sys_prompt = _translator_system(glossary_text, style_rules, target_lang)
    user_prompt = ("以下译文未通过检查，请仅修正有问题的段落，其余段落保持原样，"
                   f"返回与段落数相同的 JSON 字符串数组：\n\n{numbered}\n\n问题清单：\n{issues}")
    call_fn = call_llm_fn or call_llm
    response_format = _native_translation_response_format(
        provider, model, len(sources))
    controller = _llm_request_controller(provider, model)
    cancel_check = _llm_cancel_check()
    for _attempt in range(3):
        try:
            with trace.span("model_wait", operation="provider_gate", job_id=_runtime_job_id()):
                controller.acquire(cancel_check)
            res = _invoke_llm(
                call_fn, provider, api_key, model, sys_prompt, user_prompt,
                temperature=0.2, response_format=response_format)
            controller.note_success()
            arr = _translation_protocol.parse_translation_response(
                res, len(sources))
            return [clean_xml_chars(item).strip() for item in arr]
        except Exception as e:
            if "任务已请求取消" in str(e):
                raise
            if is_rate_limited(e):
                controller.note_rate_limited()
    raise RuntimeError("自动修复失败：模型返回格式异常")


def review_translation_batch(sources, targets, glossary_text, style_rules, target_lang,
                             provider, api_key, model):
    """独立审校一个批次（与翻译分离的 prompt/上下文），返回 (findings, failed)。"""
    numbered = "\n".join(
        f"{i + 1}. 原文：{s}\n   译文：{t}" for i, (s, t) in enumerate(zip(sources, targets)))
    sys_prompt = (f"你是一位独立的翻译审校专家，负责审查机器译文。请检查：语义准确性、术语一致性、"
                  f"漏译/增译、目标语言自然度与风格。只报告真实存在的问题，"
                  f"不要为低风险或主观偏好制造 finding。\n"
                  f"severity 只允许以下三种：blocking（结构/占位符/语义严重错误）、"
                  f"actionable（应修正的问题）、informational（建议）。\n"
                  "如果整批译文没有问题，请严格返回空数组 []，不要输出任何 informational 备注。\n"
                  f"{glossary_text}\n"
                  f"{style_rules}\n"
                  '请严格输出 JSON 数组，每项格式：{"segment_index": 0, "severity": "actionable", '
                  '"reason": "问题说明", "suggested_target": "可选：修正后的译文"}')
    user_prompt = f"待审校段落（目标语言：{target_lang}）：\n{numbered}"
    controller = _llm_request_controller(provider, model)
    cancel_check = _llm_cancel_check()
    for _attempt in range(3):
        try:
            with trace.span("model_wait", operation="provider_gate", job_id=_runtime_job_id()):
                controller.acquire(cancel_check)
            res = call_llm(provider, api_key, model, sys_prompt, user_prompt, temperature=0.2)
            controller.note_success()
            arr = parse_json_array(res)
            if arr is None:
                return [], True
            return arr, False
        except Exception as e:
            if "任务已请求取消" in str(e):
                raise
            if is_rate_limited(e):
                controller.note_rate_limited()
            else:
                break
    return [], True


# ================= 自动标注（三色学习重点）=================
# 红=生僻词/难词；黄=专业名词（特殊译法）；青绿=翻译难点句（特别译法）。
ANNOT_BATCH_SIZE = 10
ANNOT_MAX_PER_SEG = {"rare": 3, "domain": 3, "hard": 2}

# 常用英语词表（en_50k 字幕语料前 14000 词），用于把 LLM 滥标的"生僻词"挡回去
_DATA_DIR = Path(_models.__file__).resolve().parent / "resources"
_COMMON_WORDS = None


def _common_words():
    global _COMMON_WORDS
    if _COMMON_WORDS is None:
        try:
            _COMMON_WORDS = set(
                _DATA_DIR.joinpath("en_common.txt").read_text(encoding="utf-8").splitlines())
        except OSError:
            _COMMON_WORDS = set()
    return _COMMON_WORDS


_INFLECTION_SUFFIXES = (("ily", "y"), ("ness", ""), ("ment", ""), ("tion", ""),
                        ("sion", ""), ("ing", ""), ("ed", ""), ("er", ""),
                        ("est", ""), ("es", ""), ("ly", ""), ("s", ""))


def _base_form(word):
    """词形还原（一次）：grinning->grin、speedily->speedy、boxes->box。"""
    w = word
    for suffix, repl in _INFLECTION_SUFFIXES:
        if w.endswith(suffix) and len(w) > len(suffix) + 2:
            w = w[:-len(suffix)] + repl
            break
    if len(w) > 3 and w[-1] == w[-2] and w[-1] not in "eio":
        w = w[:-1]  # grinning -> grinn -> grin
    return w


def _is_common_word(token):
    """token 或其词形还原是否在常用词表内。"""
    common = _common_words()
    if not common:
        return False
    w = token.lower().strip('"\'(),;:!?')
    if not w:
        return True
    if w in common:
        return True
    base = _base_form(w)
    return base != w and base in common


_KINSHIP_TITLE_RE = re.compile(
    r"^(?:grandma|grandpa|grandmother|grandfather|aunt|uncle|mr|mrs|ms|dr|sir|lady|lord|"
    r"mother|father|mom|dad|brother|sister|captain|colonel|major|general|rabbi|"
    r"professor|doctor)\b", re.IGNORECASE)


def _rare_ok(span_text, token_freq):
    """生僻词门槛：单 token、非常用词（含词形还原）、全书出现次数少。"""
    span = (span_text or "").strip()
    if not span or " " in span or "\u00a0" in span:
        return False  # 只接受单词（可带连字符），短语/名句一律不要
    w = span.lower().strip('"\'(),;:!?')
    if len(w) < 4:
        return False
    if _is_common_word(w):
        return False
    if token_freq.get(w, 0) >= 8:
        return False  # 全书反复出现的词不算生僻
    return True


_COMMON_STOP_TERMS = {
    # 常见国家名
    "germany", "france", "china", "japan", "united states", "usa", "uk",
    "united kingdom", "russia", "italy", "spain", "canada", "australia",
    "india", "brazil", "egypt", "israel", "syria", "jordan", "lebanon",
    "iran", "iraq", "saudi arabia", "turkey", "greece", "poland", "ukraine",
    "sweden", "norway", "finland", "denmark", "netherlands", "belgium",
    "switzerland", "austria", "portugal", "mexico", "argentina", "korea",
    "vietnam", "thailand", "singapore",
    # 洲与大区
    "europe", "asia", "america", "africa", "middle east", "latin america",
    # 常见城市与通识地理实体
    "beijing", "shanghai", "berlin", "paris", "london", "tokyo",
    "new york", "washington", "jerusalem", "tel aviv", "cairo", "damascus",
    "amman", "beirut", "baghdad", "tehran", "rome", "athens", "moscow",
    "california", "texas", "sinai", "negev",
    # 通识宗教/神学通用词
    "god", "lord",
}

_COMMON_FIRST_NAMES = {
    # 常用英美/希伯来/国际人名
    "sarah", "yaakov", "david", "john", "michael", "mary", "james", "robert",
    "william", "joseph", "charles", "thomas", "daniel", "matthew", "anthony",
    "mark", "paul", "andrew", "joshua", "george", "edward", "brian", "kevin",
    "ronald", "timothy", "jason", "jeffrey", "ryan", "jacob", "gary", "nicholas",
    "eric", "jonathan", "stephen", "larry", "justin", "scott", "brandon",
    "benjamin", "samuel", "gregory", "alexander", "patrick", "frank", "jack",
    "adam", "peter", "henry", "ethan", "jeremy", "walter", "christian", "noah",
    "arthur", "jesse", "dylan", "bryan", "joe", "alan", "philip", "louis",
    "nadav", "moran", "ralik", "anat", "haggai", "ora", "avsha", "esti", "maya",
    "antoine", "evan", "gordon", "mageni", "ze’ev", "ze'ev", "raz", "oron",
    "yarden", "ilan", "ramon", "udi", "ben-amitai", "chemi", "ami", "yossik",
    "patchi", "uri", "yotam", "dan", "shuki", "danni", "yossi", "gabi", "moti",
}

_CHINESE_KINSHIP_TITLES = (
    "爷爷", "奶奶", "姥爷", "姥姥", "叔叔", "阿姨", "伯父", "伯母", "舅舅", "舅妈",
    "父亲", "母亲", "爸爸", "妈妈", "哥哥", "弟弟", "姐姐", "妹妹",
    "先生", "女士", "夫人", "小姐", "神父", "拉比", "牧师", "教授", "医生"
)

_NON_PERSON_SUFFIXES = (
    "节", "日", "市", "省", "州", "国", "湖", "山", "海", "洋", "岛", "湾",
    "条约", "协定", "行动", "计划", "公司", "战", "战争", "军", "队", "法案",
    "书", "篇", "卷", "号", "机", "弹", "炮", "服", "路", "桥", "门", "网", "图"
)

_GENERIC_CLASSIFIERS = {
    "kibbutz", "mount", "mt", "lake", "river", "sea", "fort", "port",
    "cape", "saint", "st", "valley", "bay", "camp", "gulf"
}


def _is_person_name(src_text, tgt_text=""):
    """判断是否为人物姓名或称谓（严禁作为专业术语标注）。"""
    s = (src_text or "").strip()
    t = (tgt_text or "").strip()
    if not s and not t:
        return False
    # 1. 英语称谓前缀（如 Grandpa Yaakov, Grandma Sarah, Mr. Smith）
    if _KINSHIP_TITLE_RE.search(s):
        return True
    # 2. 中文称谓后缀/包含（如 雅科夫爷爷、萨拉奶奶）
    if any(k in t for k in _CHINESE_KINSHIP_TITLES):
        return True
    # 3. 中文音译人名（由 2-4 个音译段用·或•连接，如“莎拉·马格尼”、“泽埃夫·拉兹”、“伊兰·拉蒙”、“乌迪·本-阿米泰”）
    if t and "《" not in t and "》" not in t:
        if not any(t.endswith(sfx) for sfx in _NON_PERSON_SUFFIXES):
            if re.fullmatch(r"[\u4e00-\u9fa5]{1,4}(?:[·•][\u4e00-\u9fa5\-]{1,5}){1,3}", t):
                return True
    # 4. 英文人名库匹配（1-3个词，包含人名库词汇且非军事/装备术语）
    toks = re.findall(r"[A-Za-z]+(?:['’\-][A-Za-z]+)*", s.lower())
    if toks and len(toks) <= 3:
        if any(tok in _COMMON_FIRST_NAMES for tok in toks):
            if not any(tok in {"war", "operation", "air", "missile", "jet", "force", "army", "corps", "wing", "squadron"} for tok in toks):
                return True
    return False


def _is_split_proper_noun(src_span, src_text):
    """检测标注片段是否为一个复合专名/地名短语的一部分（如 Kibbutz Geva 中仅标了 Kibbutz）。"""
    if not src_span or not src_text:
        return False
    s_start, s_end = src_span
    span_str = src_text[s_start:s_end].strip()
    if span_str.lower() in _GENERIC_CLASSIFIERS:
        after = src_text[s_end:]
        m_after = re.match(r"^\s+([A-Z][a-z]+)", after)
        before = src_text[:s_start]
        m_before = re.search(r"([A-Z][a-z]+)\s+$", before)
        if m_after or m_before:
            return True
    return False


_PRONOUN_OR_CONVERSATIONAL = {
    "you", "i", "we", "my", "your", "our", "ours", "them", "us", "me",
    "don't", "doesn't", "didn't", "won't", "can't", "please", "shall"
}


def _domain_ok(span_text, tgt_text=""):
    """专业名词门槛：称谓+人名、通识地名、全常用词短语、过长标题/句子不算专业名词。"""
    span = (span_text or "").strip()
    if not span:
        return False
    if span.lower() in _COMMON_STOP_TERMS:
        return False
    if _is_person_name(span, tgt_text):
        return False
    tokens = re.findall(r"[A-Za-z]+(?:['’\-][A-Za-z]+)*", span)
    if not tokens:
        return False
    if len(tokens) > 5 or len(span) > 40:
        return False  # 术语不应是长句或整行标题
    if any(t.lower().replace("’", "'") in _PRONOUN_OR_CONVERSATIONAL for t in tokens):
        return False  # 包含人称代词或祈使口语的句子/标题严禁标为术语
    if all(_is_common_word(t) for t in tokens):
        return False  # Translation from Hebrew 之类全是常用词
    return True


def _hard_ok(span_text, full_source_text):
    """难点句门槛：
    1. 必须是具体有特殊句式/翻译技巧的短语或分句，严禁将整段长文粗暴全量标记；
    2. 片段长度不能过长（超过 120 字符的整段/长句染色失去重点意义）。
    """
    span = (span_text or "").strip()
    if not span:
        return False
    full = (full_source_text or "").strip()
    if len(full) > 35 and len(span) >= len(full) * 0.8:
        return False
    if len(span) > 120:
        return False
    return True


def _normalize_annotations(annotations):
    """标注字典键归一化：JSON 落盘后键变字符串，统一回 int；越界/非列表值丢弃。

    这是本项目反复踩过的坑（标注渲染、过滤、续跑三处各自处理过一次），
    统一入口避免再次各修各的。
    """
    out = {}
    for k, v in (annotations or {}).items():
        try:
            gi = int(k)
        except (TypeError, ValueError):
            continue
        if isinstance(v, list):
            out[gi] = v
    return out


def _clean_annotations(annotations, pairs):
    """过滤 + 去重 + 数量上限；术语表覆盖与模型识别统一经过门槛控制。"""
    token_freq = {}
    for pr in pairs:
        for tok in re.findall(r"[A-Za-z]+", pr["source"].lower()):
            token_freq[tok] = token_freq.get(tok, 0) + 1
    cleaned = {}
    term_counts = {}
    for gi, items in _normalize_annotations(annotations).items():
        if not (0 <= gi < len(pairs)):
            continue  # 段已被删除等场景：越界键丢弃
        src_text = pairs[gi]["source"]
        seen, kept, counts = set(), [], {"rare": 0, "domain": 0, "hard": 0}
        for it in items:
            atype = it["type"]
            s_span = it.get("src_span")
            t_span = it.get("tgt_span")
            span_text = src_text[s_span[0]:s_span[1]] if s_span else ""
            tgt_text = pairs[gi]["target"][t_span[0]:t_span[1]] if t_span and 0 <= t_span[0] < t_span[1] <= len(pairs[gi]["target"]) else ""
            overlay = atype == "domain" and str(it.get("note", "")).startswith("术语：")
            if atype == "rare" and not _rare_ok(span_text, token_freq):
                continue
            if atype == "domain":
                term_from_note = ""
                if overlay:
                    term_from_note = str(it.get("note", "")).split("->")[0].replace("术语：", "").strip().lower()
                if span_text.strip().lower() in _COMMON_STOP_TERMS or term_from_note in _COMMON_STOP_TERMS:
                    continue
                if not _domain_ok(span_text, tgt_text):
                    continue
                if overlay and term_from_note and not _domain_ok(term_from_note, tgt_text):
                    continue
                if _is_split_proper_noun(s_span, src_text):
                    continue
                tokens = re.findall(r"[A-Za-z]+(?:['’\-][A-Za-z]+)*", span_text)
                if len(tokens) > 5 or len(span_text) > 40:
                    continue
                term_key = (term_from_note or span_text).strip().lower()
                if term_counts.get(term_key, 0) >= 2:
                    continue
            if atype == "hard" and not _hard_ok(span_text, src_text):
                continue
            key = (atype, tuple(s_span) if s_span else None,
                   tuple(it["tgt_span"]) if it.get("tgt_span") else None)
            if key in seen:
                continue
            seen.add(key)
            if counts[atype] >= ANNOT_MAX_PER_SEG[atype]:
                continue
            counts[atype] += 1
            if atype == "domain":
                term_key = span_text.strip().lower()
                term_counts[term_key] = term_counts.get(term_key, 0) + 1
            kept.append(it)
        if kept:
            cleaned[gi] = kept
    return cleaned


def _annotator_system(target_lang):
    return (f"你是一位翻译教学专家与专业的双语对照审阅专家。请从下列{target_lang}双语对照中，极为克制地标注三类学习与审阅重点（宁缺毋滥，每段若无真正亮点则不标）：\n"
            "1. rare：极度低频的生僻词（英语母语者也罕见的单词，最多一个带连字符复合词，如 chicory、cacophony）。"
            "严禁标注任何常用词或短语。\n"
            "2. domain：高价值核心专业术语（译文采用了行业规范/约定俗成的专名译法）。"
            "仅限 1-4 个单词的行业核心术语短语（如 Urban Air Mobility、propulsion systems、anti-aircraft guns）；"
            "【特别警告】：严禁标注任何人物姓名或亲属称谓（如 Grandpa Yaakov、Grandma Sarah、Ze'ev Raz 等人物名字一律严禁标为术语）；"
            "严禁标注常见国家/城市/行政区地名（如 Egypt、Germany、Paris、California）；"
            "严禁将专名短语/复合实体截断只标一半（例如当原文为具体地名、机构或专名如 Kibbutz Geva 时，严禁仅挑出 Kibbutz 等通用词素导致实体被切断）；"
            "严禁标注普通短语、整句标题或常用词。\n"
            "3. hard：采用了精妙翻译技巧的具体短语或从句（如后置定语分词前置、主动被动转换、倒装语序重组）。"
            "【特别警告】：hard 必须是具体的短语或分句（不超过 15 个英文单词），严禁标注整个长句，严禁标注整段文字！"
            "如果是词汇层面的普通选词，严禁标为 hard！普通直译句一律不标。\n"
            "规则：\n"
            "- 严格追求高质量与少而精：绝大多数普通段落应该返回空，每个段落最多 1 个 rare、1 个 domain、1 个 hard；\n"
            "- src 必须是原文中完整词或短语片段（严禁截断单词，西文字词必须包含完整单词边界），tgt 必须是对应译文中的完整字符串片段；\n"
            "- note 用一句话简短说明标注理由或所用译法；\n"
            '严格输出 JSON 数组，每项格式：{"seg": 1, "type": "rare", "src": "...", '
            '"tgt": "...", "note": "..."}。seg 为段落序号（从 1 开始）。不要输出任何解释文字。')


class AnnotationBatchResult(list):
    """An empty successful response differs from an unavailable annotation."""
    def __init__(self, items=(), *, status="success", error=""):
        super().__init__(items)
        self.status = status
        self.error = error


def annotate_batch(pairs_slice, target_lang, provider, api_key, model):
    """标注一个批次，返回 [{seg, type, src, tgt, note}]（seg 为 0 基）；解析失败返回 []。"""
    numbered = "\n\n".join(
        f"--- 段落 {i + 1} ---\n原文：{p['source']}\n译文：{p['target']}"
        for i, p in enumerate(pairs_slice))
    sys_prompt = _annotator_system(target_lang)
    user_prompt = f"待标注双语段落：\n{numbered}"
    controller = _llm_request_controller(provider, model)
    cancel_check = _llm_cancel_check()
    for _attempt in range(3):
        try:
            with trace.span("model_wait", operation="provider_gate", job_id=_runtime_job_id()):
                controller.acquire(cancel_check)
            res = call_llm(provider, api_key, model, sys_prompt, user_prompt, temperature=0.2)
            controller.note_success()
            arr = parse_json_array(res)
            if arr is None:
                return AnnotationBatchResult(status="parse_failed", error="模型未返回标注数组")
            out = []
            for item in arr:
                if not isinstance(item, dict):
                    continue
                seg = item.get("seg")
                if not isinstance(seg, int) or not (1 <= seg <= len(pairs_slice)):
                    continue
                atype = item.get("type")
                if atype not in ANNOTATION_COLORS:
                    continue
                src = str(item.get("src") or "").strip()
                tgt = str(item.get("tgt") or "").strip()
                note = str(item.get("note") or "").strip()
                if not src:
                    continue
                out.append({"seg": seg - 1, "type": atype, "src": src,
                            "tgt": tgt, "note": note})
            if arr and not out:
                return AnnotationBatchResult(status="parse_failed", error="标注字段或段落索引无效")
            return AnnotationBatchResult(out, status="success" if out else "empty")
        except Exception as e:
            _workflow_errors.raise_if_fatal(e, _runtime_job_id())
            if "请求取消" in str(e):
                raise
            if is_rate_limited(e):
                # Shared provider cooldown instead of a local fixed sleep, so
                # annotation retries cannot stampede alongside translation.
                backoff = controller.note_rate_limited()
                if _attempt < 2:
                    trace.record_retry("annotation", "annotate_batch", _attempt + 2,
                                       backoff, type(e).__name__, job_id=_runtime_job_id())
                continue
            return AnnotationBatchResult(status="request_failed", error=str(e)[:500])
    return AnnotationBatchResult(status="request_failed", error="请求限流，重试次数已用尽")


def _compose_spans(spans, text_len):
    """按边界切分 + 优先级（rare>domain>hard）覆盖，返回互不重叠的 (start,end,type) 列表。

    难点句常覆盖整段，词级标注嵌在其中：切到所有起点/终点后，每段取覆盖它的最高优先级。
    """
    priority = {"rare": 0, "domain": 1, "hard": 2}
    clipped = []
    for s, e, t in spans:
        s = max(0, min(int(s), text_len))
        e = max(s, min(int(e), text_len))
        if s < e:
            clipped.append((s, e, t))
    if not clipped:
        return []
    bounds = sorted({0, text_len} | {x for s, e, _ in clipped for x in (s, e)})
    composed = []
    for a, b in zip(bounds, bounds[1:]):
        if a >= b:
            continue
        mid = (a + b) / 2
        covering = [(priority[t], t) for s, e, t in clipped if s <= mid < e]
        if covering:
            composed.append((a, b, min(covering)[1]))
    return composed


def _batch_timing(job_id, stage, batch_id, phase, **metadata):
    run = trace.get_active_run(job_id)
    if run:
        run.record_point("batch_timing", stage=stage, batch_id=batch_id, phase=phase, **metadata)


def _annotation_results(tasks, job_id, provider, api_key, model, target_lang, workers):
    """Bounded requests; shared state is only changed by the ordered consumer."""
    inherited_base = getattr(_LLM_CTX, "base_url", None)
    inherited_reasoning = getattr(_LLM_CTX, "reasoning_effort", None)

    def request(task):
        bi, offset, pairs, indexes = task
        _RUNTIME_CTX.job_id = job_id
        _LLM_CTX.base_url = inherited_base
        _LLM_CTX.reasoning_effort = inherited_reasoning
        _LLM_CTX.batch_id = bi + 1
        _LLM_CTX.input_segments = len(indexes)
        token = _usage.CURRENT_ROLE.set("annotation")
        _batch_timing(job_id, "annotation", bi + 1, "worker_started", segment_count=len(pairs))
        try:
            if job_id:
                _runtime_pause_checkpoint(job_id)
                if _runtime_cancel_requested(job_id):
                    raise RuntimeError("任务已请求取消")
            with trace.span("annotation", operation="batch", job_id=job_id,
                            metadata={"batch_id": bi + 1, "offset": offset, "segment_count": len(pairs)}):
                items = annotate_batch([pairs[i] for i in indexes], target_lang, provider, api_key, model) if indexes else None
            _batch_timing(job_id, "annotation", bi + 1, "result_ready")
            return task, items
        except BaseException as exc:
            _batch_timing(job_id, "annotation", bi + 1, "worker_failed", error_type=type(exc).__name__)
            raise
        finally:
            _usage.CURRENT_ROLE.reset(token)
            _RUNTIME_CTX.__dict__.clear()
            _LLM_CTX.__dict__.clear()

    pending = {}
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="ft-annot") as executor:
        try:
            for i in range(min(workers, len(tasks))):
                _batch_timing(job_id, "annotation", tasks[i][0] + 1, "queued")
                pending[i] = executor.submit(request, tasks[i])
            for i in range(len(tasks)):
                if job_id:
                    _runtime_pause_checkpoint(job_id)
                    if _runtime_cancel_requested(job_id):
                        raise RuntimeError("任务已请求取消")
                yield pending.pop(i).result()
                # Refill only after the consumer has durably committed this result.
                next_i = i + workers
                if next_i < len(tasks):
                    _batch_timing(job_id, "annotation", tasks[next_i][0] + 1, "queued")
                    pending[next_i] = executor.submit(request, tasks[next_i])
        finally:
            for future in pending.values():
                future.cancel()


def annotate_stage(state, job_id, glossary, provider, api_key, model, target_lang,
                   on_caption=None):
    """三色自动标注：LLM 识别 + 术语表确定性覆盖（专业名词=黄色必标）。"""
    pairs = state["pairs"]
    quality_report = _audit_source_quality_for_state(state)
    state["source_quality_gate"] = quality_report
    blocked_source_indexes = {
        item.get("segment_index")
        for item in quality_report.get("flags") or []
        if isinstance(item, dict) and isinstance(item.get("segment_index"), int)
    }
    annotations = _normalize_annotations(state.get("annotations"))
    for index in blocked_source_indexes:
        annotations.pop(index, None)
    batch_results = state.setdefault("annotation_batches", {})
    failed_offsets = {int(key) for key, result in batch_results.items()
                      if result.get("status") in {"request_failed", "parse_failed", "stale"}}
    if state.get("annotations_done") and not failed_offsets:
        if blocked_source_indexes:
            state["annotations"] = annotations
            save_job_state(job_id, state)
        return state
    batches = make_batches([p["source"] for p in pairs], batch_size=ANNOT_BATCH_SIZE,
                           max_chars=2600)
    # 断点按“已标注段数”记录（而非批次号），避免批大小调整后续跑错位
    done_offset = state.get("annotations_done_offset") or 0
    start_bi = next((bi for bi, b in enumerate(batches)
                     if sum(len(x) for x in batches[:bi]) >= done_offset), len(batches))
    if done_offset:
        if on_caption:
            on_caption(f"↩️ 从第 {done_offset + 1} 段（批次 {start_bi + 1}/{len(batches)}）继续标注...")
    if job_id:
        update_runtime_state(
            job_id, heartbeat=True, stage="annotation", stage_id="annotation",
            operation="正在标注", operation_id="annotation", operation_label="正在标注",
            stage_progress={"completed": done_offset, "total": len(pairs), "unit": "段"}
        )

    tasks = []
    offset = 0
    for bi, batch in enumerate(batches):
        batch_offset = offset
        offset += len(batch)
        if bi < start_bi and batch_offset not in failed_offsets:
            continue
        slice_pairs = pairs[batch_offset:offset]
        indexes = [i for i in range(len(slice_pairs)) if batch_offset + i not in blocked_source_indexes]
        tasks.append((bi, batch_offset, slice_pairs, indexes))
    workers = min(2, max(1, int((state.get("pipeline_config") or {}).get("translation_concurrency", 2))))
    run = trace.get_active_run(job_id)
    if run:
        run.record_point("annotation_plan", concurrency=workers, batch_count=len(batches),
                         pending_batch_count=len(tasks), segment_count=len(pairs), batch_size=ANNOT_BATCH_SIZE)
    results = _annotation_results(tasks, job_id, provider, api_key, model, target_lang, workers)
    # 1) LLM 标注：请求有界并发，按原批次顺序保存。
    try:
        for (bi, offset, slice_pairs, eligible_local_indexes), items in results:
            if job_id:
                _runtime_pause_checkpoint(job_id)
                if _runtime_cancel_requested(job_id):
                    raise RuntimeError("任务已请求取消")
            if on_caption and bi % 10 == 0:
                on_caption(f"🎨 自动标注第 {offset + 1}-{offset + len(slice_pairs)} 段"
                           f"（共 {len(pairs)} 段，批次 {bi + 1}/{len(batches)}）...")
            _batch_timing(job_id, "annotation", bi + 1, "commit_started")
            if items is None:
                state["annotations"] = annotations
                state["annotations_done_offset"] = offset + len(slice_pairs)
                save_job_state(job_id, state)
                _batch_timing(job_id, "annotation", bi + 1, "commit_finished")
                continue
            batch_results[str(offset)] = {
                "status": getattr(items, "status", "success" if items else "empty"),
                "error": getattr(items, "error", ""),
                "segment_count": len(slice_pairs),
            }
            if batch_results[str(offset)]["status"] in {"success", "empty"}:
                for gi in range(offset, offset + len(slice_pairs)):
                    annotations.pop(gi, None)
            for item in items:
                local_index = item.get("seg")
                if (not isinstance(local_index, int)
                        or not 0 <= local_index < len(eligible_local_indexes)):
                    continue
                gi = offset + eligible_local_indexes[local_index]
                src, tgt, atype, note = item["src"], item["tgt"], item["type"], item["note"]
                src_span = _find_span(pairs[gi]["source"], src)
                tgt_span = _find_span(pairs[gi]["target"], tgt) if tgt else None
                # 原文片段必须能精确匹配定位，绝不允许整段粗暴兜底；译文未定位时不强制整段染色
                if src_span is None:
                    continue
                annotations.setdefault(gi, []).append(
                    {"type": atype, "src_span": list(src_span),
                     "tgt_span": list(tgt_span) if tgt_span else None, "note": note})
            # 每批落盘：断点粒度 = 一个批次
            state["annotations"] = annotations
            state["annotations_done_offset"] = max(done_offset, offset + len(slice_pairs))
            save_job_state(job_id, state)
            _batch_timing(job_id, "annotation", bi + 1, "commit_finished")
            if job_id:
                update_runtime_state(
                    job_id, heartbeat=True,
                    stage_progress={"completed": min(offset + len(slice_pairs), len(pairs)),
                                    "total": len(pairs), "unit": "段"}
                )
    finally:
        results.close()

    # 2) 术语表确定性覆盖：专业名词（特殊译法）-> 黄色（受停用词、长度和全文频次控制）
    term_counts = {}
    for gi, pr in enumerate(pairs):
        if gi in blocked_source_indexes:
            continue
        for entry in glossary:
            term = (entry.get("source") or "").strip()
            tgt_term = (entry.get("target") or "").strip()
            term_lower = term.lower()
            if len(term) < 2 or term not in pr["source"]:
                continue
            if not _domain_ok(term, tgt_term):
                continue
            tokens = re.findall(r"[A-Za-z]+(?:['’\-][A-Za-z]+)*", term)
            if len(tokens) > 5 or len(term) > 40:
                continue
            if term_counts.get(term_lower, 0) >= 2:
                continue
            span = _find_span(pr["source"], term)
            tgt_span = _find_span(pr["target"], tgt_term) if tgt_term else None
            if not span or not tgt_span:
                continue
            if _is_split_proper_noun(span, pr["source"]):
                continue
            term_counts[term_lower] = term_counts.get(term_lower, 0) + 1
            annotations.setdefault(gi, []).append(
                {"type": "domain", "src_span": list(span),
                 "tgt_span": list(tgt_span),
                 "note": f"术语：{term} -> {tgt_term}"})

    # 3) 确定性过滤（常用词/称谓/全常用词短语）+ 去重 + 数量上限
    cleaned = _clean_annotations(annotations, pairs)
    state["annotations"] = cleaned
    failed_batches = sum(result.get("status") in {"request_failed", "parse_failed", "stale"}
                         for result in batch_results.values())
    state["annotations_done"] = failed_batches == 0
    state["annotations_status"] = "partial" if failed_batches else "success"
    state["annotations_failed_batches"] = failed_batches
    save_job_state(job_id, state)
    run = trace.get_active_run(job_id)
    if run:
        run.record_point("milestone", phase="annotation_complete")
    if job_id:
        update_runtime_state(
            job_id, heartbeat=True,
            stage_progress={"completed": len(pairs), "total": len(pairs), "unit": "段"}
        )
    if on_caption:
        total = sum(len(v) for v in cleaned.values())
        on_caption(f"{'⚠️ 自动标注部分完成' if failed_batches else '✅ 自动标注完成'}："
                   f"{total} 处（失败批次 {failed_batches}/{len(batches)}；可继续处理补跑失败批次）")
    return state


DEFAULT_REVIEW_OVERLAP = False
DEFAULT_MAX_REVIEW_IN_FLIGHT = 2
DEFAULT_REVIEW_OVERLAP_STRATEGY = "generated"


def review_overlap_enabled(review_overlap=None, pipeline_config=None) -> bool:
    """Return whether review overlap is enabled.

    Priority:
    1. review_overlap parameter (if explicitly given)
    2. pipeline_config["review_overlap"] / pipeline_config["review_overlap_enabled"]
    3. env FOLITH_REVIEW_OVERLAP / FOLIOTHREAD_REVIEW_OVERLAP
    4. False (default)
    """
    if review_overlap is not None:
        return bool(review_overlap)
    if isinstance(pipeline_config, dict):
        if "review_overlap" in pipeline_config:
            return bool(pipeline_config["review_overlap"])
        if "review_overlap_enabled" in pipeline_config:
            return bool(pipeline_config["review_overlap_enabled"])
    env_val = os.getenv("FOLITH_REVIEW_OVERLAP") or os.getenv("FOLIOTHREAD_REVIEW_OVERLAP")
    if env_val is not None:
        return env_val.strip().lower() in ("1", "true", "yes", "on")
    return DEFAULT_REVIEW_OVERLAP


def review_overlap_max_in_flight(max_in_flight=None, pipeline_config=None) -> int:
    """Return maximum in-flight reviews for review overlap."""
    if max_in_flight is not None:
        try:
            return max(1, min(8, int(max_in_flight)))
        except (TypeError, ValueError):
            pass
    if isinstance(pipeline_config, dict) and pipeline_config.get("max_review_in_flight") is not None:
        try:
            return max(1, min(8, int(pipeline_config["max_review_in_flight"])))
        except (TypeError, ValueError):
            pass
    env_val = os.getenv("FOLITH_MAX_REVIEW_IN_FLIGHT") or os.getenv("FOLIOTHREAD_MAX_REVIEW_IN_FLIGHT")
    if env_val is not None:
        try:
            return max(1, min(8, int(env_val)))
        except (TypeError, ValueError):
            pass
    return DEFAULT_MAX_REVIEW_IN_FLIGHT


def review_overlap_strategy_mode(strategy=None, pipeline_config=None) -> str:
    """Return review overlap strategy mode ('generated' or 'stable')."""
    val = strategy
    if not val and isinstance(pipeline_config, dict):
        val = pipeline_config.get("review_overlap_strategy")
    if not val:
        val = os.getenv("FOLITH_REVIEW_OVERLAP_STRATEGY") or os.getenv("FOLIOTHREAD_REVIEW_OVERLAP_STRATEGY")
    val = str(val or "").strip().lower()
    return val if val in ("generated", "stable") else DEFAULT_REVIEW_OVERLAP_STRATEGY


def _execute_batch_review(
    snapshot: _review_scheduler.BatchReviewSnapshot,
    reviewer_call: Callable,
    reviewer_config: Dict[str, Any],
    job_id: str,
    enable_review: bool,
) -> _review_scheduler.ReviewResult:
    """Worker task that executes semantic review & suggestion verification on an immutable snapshot.

    Operates strictly on snapshot data; performs zero shared state mutations.
    """
    _llm_before = llm_thread_stats()
    started_at = time.perf_counter()
    if not enable_review:
        return _review_scheduler.ReviewResult(
            batch_index=snapshot.batch_index,
            offset=snapshot.offset,
            review_succeeded=False,
            review_failed=False,
            rfindings=[],
            review_trace=None,
            review_event_id=None,
            formal_records=[],
            promoted_review_events=[],
            repair_overlays=[],
            review_evidence=[],
            accepted_suggestions={},
            started_at_s=started_at,
            finished_at_s=time.perf_counter(),
            total_wall_s=0.0,
        )
    if not snapshot.formal_segment_ids:
        return _review_scheduler.ReviewResult(
            batch_index=snapshot.batch_index,
            offset=snapshot.offset,
            review_succeeded=True,
            review_failed=False,
            rfindings=[],
            review_trace=None,
            review_event_id=None,
            formal_records=[],
            promoted_review_events=[],
            repair_overlays=[],
            review_evidence=[],
            accepted_suggestions={},
            started_at_s=started_at,
            finished_at_s=time.perf_counter(),
            total_wall_s=0.0,
        )
    if snapshot.review_gate_skip:
        # Deterministic gating decided this batch does not need independent
        # review.  Report it as skipped with explicit provenance: no review
        # happened, so nothing is marked reviewed, approved or TM-eligible.
        skip_event_id = (
            f"translation-review-{job_id}-{snapshot.batch_index}-skipped-low-risk")
        skip_trace = {
            "blind": False,
            "segment_ids": list(snapshot.formal_segment_ids),
            "rounds": [],
            "requests": [],
            "evidence_ids": [],
            "decision": _review_gate.SKIPPED_LOW_RISK,
            "review_identity": {},
            "review_mode": _review_gate.REVIEW_MODE_RISK_SELECTIVE,
            "skip_reason": snapshot.review_gate_reason,
            "risk_signals": dict(snapshot.review_gate_signals or {}),
            "completion_receipt": {
                "status": _review_gate.SKIPPED_LOW_RISK,
                "reviewed_segment_ids": [],
                "reviewed_unit_count": 0,
                "finding_count": 0,
                "evidence_ids": [],
                "review_identity": {},
            },
        }
        return _review_scheduler.ReviewResult(
            batch_index=snapshot.batch_index,
            offset=snapshot.offset,
            review_succeeded=False,
            review_failed=False,
            rfindings=[],
            review_trace=skip_trace,
            review_event_id=skip_event_id,
            formal_records=[],
            promoted_review_events=[],
            repair_overlays=[],
            review_evidence=[],
            accepted_suggestions={},
            review_skipped=True,
            review_skip_reason=snapshot.review_gate_reason,
            review_skip_trace=skip_trace,
            started_at_s=started_at,
            finished_at_s=time.perf_counter(),
            total_wall_s=0.0,
        )

    rfindings = []
    review_failed = False
    review_succeeded = False
    review_trace = None
    review_event_id = None
    formal_records = []
    promoted_review_events = []
    repair_overlays = []
    review_evidence = []
    accepted_suggestions = {}

    rfindings, failed, review_trace = (
        _translation_evidence.review_translation_batch_with_evidence(
            [snapshot.source_texts[segment_id - snapshot.offset]
             for segment_id in snapshot.formal_segment_ids],
            [snapshot.generated_targets[segment_id - snapshot.offset]
             for segment_id in snapshot.formal_segment_ids],
            snapshot.glossary_text,
            snapshot.style_rules,
            snapshot.target_lang,
            reviewer_config["provider"],
            reviewer_config["api_key"],
            reviewer_config["model"],
            snapshot.evidence_index,
            call_llm=reviewer_call,
            segment_ids=snapshot.formal_segment_ids,
            translation_core_packet=snapshot.formal_packet,
        )
    )

    review_event_id = (
        f"translation-review-{job_id}-{snapshot.batch_index}-formal-"
        f"{snapshot.base_evidence_count}"
    )
    if review_trace:
        review_trace["review_event_id"] = review_event_id

    if failed:
        review_failed = True
    else:
        review_succeeded = True
        global_to_local = {
            snapshot.offset + index: index for index in range(len(snapshot.batch_pairs))
        }
        for rf in rfindings:
            sev = rf.get("severity")
            if sev not in ("blocking", "actionable", "informational"):
                continue
            segment_id = rf.get("segment_id")
            idx = global_to_local.get(segment_id)
            if idx is None:
                continue
            record = _review_finding_record(rf, review_event_id)
            if (
                sev == "actionable"
                and rf.get("suggested_target")
                and idx not in snapshot.blocked_local_indexes
                and not snapshot.batch_pairs[idx].get("from_tm")
            ):
                suggested = clean_xml_chars(rf["suggested_target"]).replace("\n", " ").strip()
                if suggested:
                    old_target = snapshot.generated_targets[idx]
                    overlay = _repair.create_overlay(
                        [old_target],
                        [suggested],
                        [rf],
                        "review_suggested",
                        sources=[snapshot.source_texts[idx]],
                        finding_segment_ids=[segment_id],
                    )
                    recheck = _globalize_batch_findings(
                        check_translation_batch(
                            [snapshot.source_texts[idx]],
                            [suggested],
                            snapshot.glossary,
                            snapshot.target_lang,
                            section_profile=snapshot.section_profile,
                        ),
                        segment_id,
                    )
                    blind_trace = None
                    blind_findings = []
                    blind_failed = False
                    if not any(f["severity"] in ("blocking", "actionable") for f in recheck):
                        blind_index = _translation_evidence_index(
                            snapshot.paras,
                            snapshot.committed_pairs,
                            snapshot.batch_pairs,
                            snapshot.glossary,
                            snapshot.document_profile,
                            snapshot.document_synopsis,
                            snapshot.section_digests,
                            snapshot.findings_all,
                            blind=True,
                            candidate_targets={segment_id: suggested},
                        )
                        suggested_packet = _translation_evidence.build_runtime_review_packet(
                            snapshot.state_ref,
                            [snapshot.batch_pairs[idx]],
                            [segment_id],
                            snapshot.glossary,
                            deterministic_checks=recheck,
                            review_context=snapshot.review_context,
                            candidate_targets={segment_id: suggested},
                            blind=True,
                        )
                        blind_findings, blind_failed, blind_trace = (
                            _translation_evidence.review_translation_batch_with_evidence(
                                [snapshot.source_texts[idx]],
                                [suggested],
                                snapshot.glossary_text,
                                snapshot.style_rules,
                                snapshot.target_lang,
                                reviewer_config["provider"],
                                reviewer_config["api_key"],
                                reviewer_config["model"],
                                blind_index,
                                call_llm=reviewer_call,
                                blind=True,
                                segment_ids=[segment_id],
                                review_identity={
                                    "input_hash": overlay["input_hash"],
                                    "candidate_hash": overlay["candidate_hash"],
                                },
                                translation_core_packet=suggested_packet,
                            )
                        )
                    overlay = _repair.evaluate_overlay(
                        overlay,
                        recheck,
                        blind_findings,
                        blind_failed,
                        review_identity=(blind_trace or {}).get("review_identity"),
                    )
                    repair_overlays.append({
                        "batch": snapshot.batch_index,
                        "offset": snapshot.offset,
                        "batch_local_ordinal": idx,
                        "segment_id": segment_id,
                        **overlay,
                        "blind_trace": blind_trace,
                    })
                    if overlay["status"] == "accepted":
                        accepted_suggestions[idx] = suggested
                        if blind_trace:
                            promoted_review_events.append((
                                blind_trace,
                                f"translation-review-{job_id}-{snapshot.batch_index}-suggested-"
                                f"{segment_id}-{len(promoted_review_events)}",
                                segment_id,
                            ))
                    else:
                        record["suggested_target"] = suggested
                        formal_records.append(record)
                        if blind_trace:
                            review_evidence.append({
                                "batch": snapshot.batch_index,
                                "phase": "suggested_shadow_review",
                                **blind_trace,
                            })
                    continue
            formal_records.append(record)

    finished_at = time.perf_counter()
    _llm_delta = llm_thread_stats_delta(_llm_before)
    return _review_scheduler.ReviewResult(
        batch_index=snapshot.batch_index,
        offset=snapshot.offset,
        review_succeeded=review_succeeded,
        review_failed=review_failed,
        rfindings=rfindings,
        review_trace=review_trace,
        review_event_id=review_event_id,
        formal_records=formal_records,
        promoted_review_events=promoted_review_events,
        repair_overlays=repair_overlays,
        review_evidence=review_evidence,
        accepted_suggestions=accepted_suggestions,
        api_latency_s=float(_llm_delta.get("api_latency_s", 0.0) or 0.0),
        llm_requests=int(_llm_delta.get("requests", 0) or 0),
        llm_failures=int(_llm_delta.get("failures", 0) or 0),
        started_at_s=started_at,
        finished_at_s=finished_at,
        total_wall_s=finished_at - started_at,
    )


@_exclusive_job_execution(1)
def translate_stage(state, job_id, glossary, provider, api_key, model, target_lang,
                    style_rules, enable_review, use_tm=True, document_profile=None,
                    on_status=None, on_caption=None, translator_config=None,
                    reviewer_config=None, batch_size=None, max_batch_chars=None,
                    auxiliary_config=None, knowledge_feedback_interval=None,
                    translation_concurrency=None, profiler=None,
                    review_overlap=None, max_review_in_flight=None,
                    review_overlap_strategy=None, use_translation_cache=True,
                    review_mode=None, use_knowledge_cache=True,
                    context_mode=None, context_budget_tokens=None,
                    context_ablations=None, context_offline=None):
    """阶段二：语义批次翻译 + 确定性检查/修复 + 独立审校 + 翻译记忆。

    对齐 localize-anything 经验：
    - 语义批次：≤4 段一组，携带前后文，保留每批落盘的断点粒度；
    - 概念化术语表：锁定术语强制首选译名/禁止译名，保留项强制原样；
    - 确定性检查：占位符/URL/引用等保留项、残留原文、锁定术语合规，问题自动修复一轮；
    - 独立审校：actionable 建议经确定性复验后应用，blocking 记录给用户确认；
    - 翻译记忆：仅审校通过的段落入库，精确命中直接复用。

    ``context_mode`` 选择上下文装配路径。**生产默认是 ``"fixed"``**：

    - ``"fixed"``（``PRODUCTION_CONTEXT_MODE``，生产/稳定/默认）：现有
      ``compile_context_packet`` 装配，渲染结果与历史版本逐字节一致；
    - ``"dynamic_v2"``（``EXPERIMENTAL_CONTEXT_MODE``，实验/研究/显式选择）：
      ``folith.dynamic_context`` 的多源检索装配，只改变注入的上下文，不改变
      批次、提示词骨架、审校、TM 或状态语义。**未推广，不再扩展**；
      见 ``docs/dynamic-context-v2-adr.md``。

    缺失、未知或非法取值一律回落到 ``"fixed"``——实验路径只能由显式的
    ``"dynamic_v2"`` 进入，不可能因为配置漂移而变成隐式路径。
    """
    state["translation_core_review_required"] = bool(enable_review)
    # ---- context selector resolution (production default: fixed) -------------
    # Resolved once per stage.  This is the ONLY place the mode is decided, so
    # "is the experiment on?" has exactly one answer per run.
    effective_context_mode = _dynamic_context.resolve_context_mode(context_mode, state)
    try:
        effective_context_budget = int(
            context_budget_tokens if context_budget_tokens is not None
            else state.get("context_budget_tokens")
            or _dynamic_context.DEFAULT_CONTEXT_BUDGET_TOKENS)
    except (TypeError, ValueError):
        effective_context_budget = _dynamic_context.DEFAULT_CONTEXT_BUDGET_TOKENS
    effective_context_budget = max(1, effective_context_budget)
    if context_ablations is None:
        context_ablations = state.get("context_ablations") or ()
    effective_context_ablations = tuple(
        name for name in (context_ablations or ())
        if name in _dynamic_context.ABLATION_NAMES)
    if context_offline is None:
        context_offline = bool(state.get("context_offline"))
    effective_context_offline = bool(context_offline)
    # Persist the experiment configuration only when the experiment is active:
    # a production `fixed` job's state must stay byte-identical to before, and
    # `pipeline_config` (the serialized task configuration) never carries these
    # keys at all.
    if effective_context_mode != _dynamic_context.PRODUCTION_CONTEXT_MODE:
        state["context_mode"] = effective_context_mode
        state["context_budget_tokens"] = effective_context_budget
        state["context_ablations"] = list(effective_context_ablations)
        state["context_offline"] = effective_context_offline
    # Review policy: ``always`` (the production default, unchanged) or
    # ``risk_selective`` (deterministic pre-review gating).  Unrecognised values
    # fall back to ``always``, so configuration can never silently weaken review.
    effective_review_mode = _review_gate.resolve_review_mode(
        review_mode, state.get("pipeline_config"), state)
    state["review_mode"] = effective_review_mode
    # 严格术语治理门禁：存在待审核候选术语且未冻结（且未显式跳过）时，
    # 任何入口都禁止开始翻译。导入的锁定术语视为已固定，不构成阻塞。
    pending = [e for e in (glossary or [])
               if (e.get("status") or "").lower() == "candidate"]
    if state.get("quality_mode") and pending \
            and not state.get("glossary_frozen") \
            and not state.get("quality_bypass"):
        raise RuntimeError(
            "严格术语治理：仍有候选术语未审核（术语表尚未冻结），"
            "禁止开始翻译（请在术语审核面板冻结后继续）")
    translator_config = _model_roles.normalize_role_config(
        translator_config, fallback_provider=provider, fallback_model=model,
        fallback_api_key=api_key)
    reviewer_config = _model_roles.normalize_role_config(
        reviewer_config, fallback_provider=provider, fallback_model=model,
        fallback_api_key=api_key)
    auxiliary_config = _model_roles.normalize_role_config(
        auxiliary_config, fallback_provider=provider, fallback_model=model,
        fallback_api_key=api_key)
    usage_ledger = state.setdefault("llm_usage", _usage.empty_usage())
    translator_call = _model_roles.make_role_call(
        _usage.tracked_call(call_llm, usage_ledger, role="translation"),
        translator_config, job_id=job_id)
    reviewer_call = _model_roles.make_role_call(
        _usage.tracked_call(call_llm, usage_ledger, role="review"),
        reviewer_config, job_id=job_id)
    auxiliary_call = _model_roles.make_role_call(
        _usage.tracked_call(call_llm, usage_ledger, role="auxiliary"),
        auxiliary_config, job_id=job_id)
    _tm_project = tm_project_id(state)
    # 目标语言是 TM 作用域的一部分：没有它就不建立任何命中路径（fail closed）。
    _tm_lang = str(target_lang or "").strip() or str(state.get("target_lang") or "").strip()
    tm = load_tm(_tm_project) if use_tm else {}
    tm_norm_index = tm_index(tm, _tm_lang) if use_tm else {}
    if use_tm:
        def _tm_scope(source, target):
            """把目标语言编进记忆键；语言无法证明时不恢复任何记忆。"""
            key = tm_scope_key(_tm_lang, source)
            if not key:
                return None
            return key, tm_record(target, _tm_lang)

        recovered, pending_events = _checkpoint.reconcile_translation_memory(
            tm, state, job_dir(job_id), scope=_tm_scope)
        if recovered:
            save_tm(tm, _tm_project)
            state["tm_recovered_count"] = state.get("tm_recovered_count", 0) + pending_events
            save_job_state(job_id, state)
    # Exact execution cache.  Separate from TM on purpose: TM is reviewed
    # linguistic memory, this is "the same request was already paid for".
    # Losing it costs API calls, never correctness, so it fails open and is
    # written on the same cadence as the state snapshot.
    _result_cache = (_translation_cache.load(translation_cache_path(_tm_project))
                     if use_translation_cache
                     else _translation_cache.TranslationResultCache())
    # Knowledge-feedback observation reuse.  Same project scoping as TM and the
    # exact translation cache: the remote extraction is a pure function of the
    # rendered request, so a repeated batch must not be paid for twice.  The
    # local merge always re-runs, so this can never skip state reconstruction.
    _knowledge_cache = (_knowledge_cache_module.load(
        knowledge_cache_path(_tm_project)) if use_knowledge_cache
        else _knowledge_cache_module.KnowledgeObservationCache())
    paras = state["paras"]
    pairs = state["pairs"]
    # Source cleanup repairs text, but a repair pass cannot prove that every
    # OCR item is a real paragraph.  Quarantine short/structurally suspicious
    # items before they can become model translation requests.
    source_quality_gate = _audit_source_quality_for_state(state)
    state["source_quality_gate"] = source_quality_gate
    blocked_source_indexes = {
        int(item["segment_index"]): item
        for item in source_quality_gate.get("flags") or []
        if isinstance(item, dict) and isinstance(item.get("segment_index"), int)
    }
    truncated_indexes = []

    # Bounded snapshot cadence for automatic translation generation.  A full
    # authoritative snapshot per batch makes cumulative writes grow with the
    # square of the document length; deferring to a bounded window keeps
    # `state.json` authoritative while cutting that growth by the cadence
    # factor.  The crash window is a hard cap (see state_snapshot_crash_window).
    # `_cadence` is filled in once the batch plan is known (the cap is a
    # fraction of the document).
    _cadence = {"every_batches": 1, "max_interval_s": STATE_SNAPSHOT_MAX_INTERVAL_S,
                "journal_bytes": 0, "journal_fraction": 0.0}
    _pending_snapshot_batches = 0
    _last_snapshot_at = time.perf_counter()
    _snapshot_stats = {"taken": 0, "deferred": 0}
    # The journal carries every committed batch, so a crash is repaired by
    # replay instead of re-translation.  Sequence numbers continue from the
    # snapshot's cursor, which keeps them monotonic across the rotations that
    # follow each snapshot.
    _journal_enabled = state_journal_enabled()
    _journal_seq = _state_journal.cursor(state)
    _journal_retire_pending = False
    _journal_bytes_written = _state_journal.size(job_id) if _journal_enabled else 0
    _journal_stats = {"records": 0, "bytes": 0, "unusable": 0, "budget_bytes": 0}
    # Replay budget: a floor for small documents, otherwise a multiple of the
    # snapshot's own size.  Recomputed at every flush (see `_refresh_budget`),
    # never per batch -- a `stat` per batch would be a syscall on the hot path
    # for a number that only changes when a snapshot is written.
    _journal_budget_bytes = _cadence["journal_bytes"]

    def _refresh_budget():
        """Resize the replay budget against the snapshot that just landed."""
        nonlocal _journal_budget_bytes
        floor = _cadence["journal_bytes"]
        fraction = _cadence.get("journal_fraction") or 0.0
        if not floor:
            _journal_budget_bytes = 0
            return
        if not fraction:
            _journal_budget_bytes = floor
            return
        try:
            snapshot_bytes = job_state_path(job_id).stat().st_size
        except OSError:
            snapshot_bytes = 0
        _journal_budget_bytes = max(floor, int(snapshot_bytes * fraction))
        _journal_stats["budget_bytes"] = _journal_budget_bytes
    # TM promotions ride the same window as the state snapshot.  `tm.json` is a
    # derived cache (see `save_tm`), so its write cadence is a performance knob
    # rather than a truth knob -- but it must stay *consistent* with the
    # snapshot, which is why the two flush together instead of on independent
    # schedules.  State is written first: that keeps the snapshot ahead of the
    # cache, which is the direction recovery can repair from.
    _pending_tm_entries = []
    _pending_tm_window = {"batch": None, "offset": 0}

    def _flush_result_cache():
        """Persist cache entries added since the last flush.

        Only writes when there is something new; the file is a derived side
        cache, so a failed write is a missed optimisation, never an error.
        """
        if not use_translation_cache or not _result_cache.dirty():
            return
        started = time.perf_counter()
        started_cpu = time.thread_time()
        _translation_cache.save(_result_cache, translation_cache_path(_tm_project))
        finished = time.perf_counter()
        _critical_path.profiler().record(
            "commit_translation_cache_flush", finished - started,
            cpu_s=time.thread_time() - started_cpu)
        _critical_path.count("translation_cache_flush_count")

    def _flush_knowledge_cache():
        """Persist extraction results added since the last flush.

        A derived side cache: a failed write is a missed optimisation, never an
        error, and the merge path does not depend on it.
        """
        if not use_knowledge_cache or not _knowledge_cache.dirty():
            return
        started = time.perf_counter()
        started_cpu = time.thread_time()
        _knowledge_cache_module.save(_knowledge_cache,
                                     knowledge_cache_path(_tm_project))
        finished = time.perf_counter()
        _critical_path.profiler().record(
            "commit_knowledge_cache_flush", finished - started,
            cpu_s=time.thread_time() - started_cpu)
        _critical_path.count("knowledge_cache_flush_count")

    def _flush_tm_promotions():
        """Persist accumulated promotions, paired with their evidence events.

        The `tm_promotion_pending` / `tm_promotion_done` pair is emitted around
        the actual write, so a crash can still distinguish "promoted" from
        "merely intended" -- the same contract as one-write-per-batch, just
        with a coarser granularity.
        """
        if not _pending_tm_entries:
            return
        entries = list(_pending_tm_entries)
        _pending_tm_entries.clear()
        batch_index = _pending_tm_window["batch"]
        offset = _pending_tm_window["offset"]
        safe_append_event({
            "batch": batch_index, "offset": offset,
            "phase": "tm_promotion_pending", "entries": entries,
        })
        started = time.perf_counter()
        started_cpu = time.thread_time()
        save_tm(tm, _tm_project)
        finished = time.perf_counter()
        _record_batch_stage(
            "tm_commit", batch_index if batch_index is not None else 0,
            (started, finished), tm_entry_count=len(entries))
        _critical_path.profiler().record(
            "commit_tm_promotion_flush", finished - started,
            cpu_s=time.thread_time() - started_cpu)
        _critical_path.count("tm_promotion_flush_count")
        safe_append_event({
            "batch": batch_index, "offset": offset,
            "phase": "tm_promotion_done", "entries": entries,
        })

    def _flush_state_snapshot(reason):
        """Write the authoritative snapshot now and reset the cadence window.

        Order is fixed and load-bearing.  The cursor is stamped onto the state
        *before* the snapshot is written, so the snapshot that reaches the disk
        always claims exactly the journal records it contains -- never fewer
        (which would replay them twice) and never more (which would drop them).
        The journal is retired only after the atomic replace has succeeded.
        """
        nonlocal _pending_snapshot_batches, _last_snapshot_at
        nonlocal _journal_bytes_written, _journal_retire_pending
        stopped = _workflow_errors.failure(job_id)
        if stopped is not None and stopped.reason == "disk_full":
            raise stopped
        if _journal_seq:
            _state_journal.set_cursor(state, _journal_seq)
        # `save_job_state` performs the compaction step (it retires the journal
        # once the snapshot provably contains every record).  Compaction order
        # is fixed: new snapshot -> temp -> atomic replace -> retire journal.
        save_job_state(job_id, state)
        if _journal_retire_pending:
            _state_journal.rotate(job_id)
            _journal_retire_pending = False
            _journal_bytes_written = 0
        if _journal_enabled:
            _journal_bytes_written = _state_journal.size(job_id)
            _refresh_budget()
        _flush_tm_promotions()
        _flush_result_cache()
        _flush_knowledge_cache()
        _pending_snapshot_batches = 0
        _last_snapshot_at = time.perf_counter()
        _snapshot_stats["taken"] += 1
        _critical_path.count("state_snapshot_taken")
        safe_append_event({
            "batch": None, "offset": len(pairs), "phase": "state_snapshot",
            "reason": reason, "pairs_count": len(pairs),
        })

    def _commit_translation_batch(batch_pairs, offset, bi, baseline, patches,
                                  segment_ids):
        nonlocal truncated_indexes, _pending_snapshot_batches, _cadence
        nonlocal _journal_seq, _journal_bytes_written, _journal_enabled
        nonlocal _journal_retire_pending
        pairs.extend(batch_pairs)
        changed_indexes = list(range(offset, offset + len(batch_pairs)))
        if changed_indexes:
            # Measured separately from the save itself: recording the truth
            # mutation is real per-batch work that the opaque ``state_commit``
            # stage used to hide.
            with _persistence_profiler.phase("truth_marking"):
                _mark_translation_truth_changed(
                    job_id, state, changed_indexes,
                    "翻译流水线写入 CURRENT_TRANSLATION；记录该批次影响范围",
                    actor="pipeline", action="translation_batch",
                    stale_translation_reviews=False, patches=patches)
        truncated_indexes = []
        # The record is built here, after every state mutation for this batch and
        # before the cadence decision, so it describes the committed batch
        # exactly and is on disk before the window that may swallow it.
        record = None
        if _journal_enabled:
            record = build_batch_journal_record(
                state, baseline, patches, batch=bi, offset=offset,
                segment_ids=segment_ids, pairs=batch_pairs)
        if record is not None:
            _journal_seq += 1
            record["seq"] = _journal_seq
            try:
                _, _written = _state_journal.append_mutation(job_id, record)
            except OSError as exc:
                _workflow_errors.raise_if_fatal(exc, job_id)
                _written = None
            if _written is None:
                record = None
                _journal_seq -= 1
            else:
                _journal_bytes_written += _written
                _journal_stats["records"] += 1
                _journal_stats["bytes"] += _written
        # Only a *missing* record while the journal is on is a problem.  With the
        # journal off there is nothing to describe, and the bounded cadence is
        # the durability mechanism -- exactly as before this round.
        fallback = record is None and _journal_enabled
        if fallback:
            _journal_stats["unusable"] += 1
        _pending_snapshot_batches += 1
        elapsed = time.perf_counter() - _last_snapshot_at
        due = False
        if _cadence["every_batches"] \
                and _pending_snapshot_batches >= _cadence["every_batches"]:
            due = True
        elif _cadence["max_interval_s"] \
                and elapsed >= _cadence["max_interval_s"]:
            # 0.0 disables the trigger outright; without this guard a 0.0 would
            # mean "snapshot on every batch", the opposite of the intent.
            due = True
        elif _journal_budget_bytes \
                and _journal_bytes_written >= _journal_budget_bytes:
            due = True
        if fallback:
            # Preserve the old durable journal until the complete replacement
            # snapshot succeeds, including when another cadence trigger fired.
            due = True
            _journal_enabled = False
            _journal_retire_pending = True
            _cadence = state_snapshot_cadence(
                total_batches=len(batches), journal=False)
            _refresh_budget()
            safe_append_event({
                "batch": bi, "offset": offset,
                "phase": "state_journal_disabled",
                "reason": "unusable_batch_delta",
            })
        if due:
            _flush_state_snapshot("cadence")
        else:
            _snapshot_stats["deferred"] += 1

    def _save_translation_failure():
        nonlocal truncated_indexes
        if truncated_indexes:
            _mark_translation_truth_changed(
                job_id, state, truncated_indexes,
                "断点恢复截断未完成批次；CURRENT_TRANSLATION 已改变",
                actor="pipeline", action="translation_checkpoint_truncate")
            truncated_indexes = []
        # A truncated/failed batch is a mandatory immediate snapshot: it must
        # not wait for the cadence window.
        _flush_state_snapshot("translation_failure")

    # 批次大小直接影响调用次数与耗时（本仓库 82 段文档实测：4/2400 → 32 批，
    # 6/4800 → 17 批）。默认沿用模块常量以保持既有行为，只有显式配置时才改变。
    effective_batch_size = int(batch_size or BATCH_SIZE)
    effective_max_chars = int(max_batch_chars or TRANSLATION_MAX_BATCH_CHARS)
    state["batch_plan"] = {
        "batch_size": effective_batch_size,
        "max_batch_chars": effective_max_chars,
    }
    batches = make_batches(
        paras, batch_size=effective_batch_size, max_chars=effective_max_chars,
        semantic_units=state.get("semantic_units")
        or state.get("section_digests") or None)
    state["batch_plan"]["batch_count"] = len(batches)
    # The snapshot cadence is bounded by a fraction of the document, so it can
    # only be resolved once the batch plan exists.
    _cadence = state_snapshot_cadence(
        total_batches=len(batches), journal=_journal_enabled)
    _refresh_budget()
    state["batch_plan"]["snapshot_cadence"] = {
        "every_batches": _cadence["every_batches"],
        "max_interval_s": _cadence["max_interval_s"],
        "journal_enabled": bool(_journal_enabled),
        "journal_bytes": _cadence["journal_bytes"],
        "journal_fraction": _cadence["journal_fraction"],
        "journal_budget_bytes": _journal_budget_bytes,
        "crash_window_segments": state_snapshot_crash_window(
            batch_size=effective_batch_size, total_batches=len(batches)),
    }
    # Knowledge extraction is a continuity aid, not a second translation pass.
    # Its per-batch cost is bounded by a deterministic local trigger rather than
    # by the review mode: a batch is observed when it carries a name-like form
    # the project does not already hold, plus a periodic checkpoint that bounds
    # staleness for term hints.  Measured on the committed workload this keeps
    # 100% of proper-name continuity knowledge and 88.5% of prompt-injected
    # items at 42.6% of the calls.  An explicit ``knowledge_feedback_interval``
    # of 1 (or strict ``quality_mode``) restores the mandatory every-batch path.
    if knowledge_feedback_interval is None:
        feedback_interval = _knowledge.KNOWLEDGE_FEEDBACK_INTERVAL
    else:
        try:
            feedback_interval = max(1, int(knowledge_feedback_interval))
        except (TypeError, ValueError):
            feedback_interval = 1
    force_knowledge_feedback = bool(
        state.get("quality_mode") or feedback_interval <= 1)
    state["knowledge_feedback_policy"] = {
        "interval": feedback_interval,
        "force_every_batch": force_knowledge_feedback,
        "batch_count": len(batches),
        "observed_batches": [],
        "skipped_batches": [],
        "skip_reasons": {},
        "signal_batches": [],
        "cache_hits": 0,
        "cache_misses": 0,
    }
    registry = _entity_registry.EntityRegistry(state.get("entity_registry") or [])

    # 断点：从第一个未完成批次继续；若中间批次不完整则截断重译
    cum_end, start_batch = 0, 0
    for bi, b in enumerate(batches):
        prev_end = cum_end
        cum_end += len(b)
        if cum_end > len(pairs):
            if len(pairs) > prev_end:
                truncated_indexes = list(range(prev_end, len(pairs)))
                del pairs[prev_end:]
            start_batch = bi
            break
    else:
        start_batch = len(batches)

    stats = state.setdefault("review_stats", {
        "reviewed_segments": 0, "batches_reviewed": 0,
        "batches_skipped_low_risk": 0, "skipped_segments": 0,
        "blocking": 0, "actionable": 0, "informational": 0, "review_failed": 0,
    })
    findings_all = state.setdefault("findings", [])
    existing_source_gate_indexes = {
        int(item.get("segment_index"))
        for item in findings_all
        if isinstance(item, dict)
        and item.get("type") == "source_quality_gate"
        and isinstance(item.get("segment_index"), int)
    }
    for segment_index, flag in blocked_source_indexes.items():
        if segment_index in existing_source_gate_indexes:
            continue
        findings_all.append(_source_quality_finding(flag))
    from folith.terminology import (
        detect_glossary_conflicts as _detect_conflicts,
        glossary_block as _glossary_block,
        select_glossary_for_segments as _select_glossary,
    )
    frozen_hash = (state.get("glossary_frozen") or {}).get("glossary_hash")
    sections = (document_profile or {}).get("sections") or []
    section_digests = state.get("section_digests") or []
    document_synopsis = state.get("document_synopsis") or {}

    offsets = []
    _cur_offset = 0
    for b in batches:
        offsets.append(_cur_offset)
        _cur_offset += len(b)

    effective_concurrency = _runtime_int_option(
        translation_concurrency,
        (state.get("pipeline_config") or {}).get("translation_concurrency"),
        "FOLIOTHREAD_TRANSLATION_CONCURRENCY",
        DEFAULT_TRANSLATION_CONCURRENCY,
        minimum=1,
        maximum=16,
    )
    effective_concurrency = max(1, min(int(effective_concurrency), max(1, len(batches) - start_batch)))
    state["batch_plan"]["concurrency"] = effective_concurrency
    run = trace.get_active_run(job_id)
    if run:
        run.record_point("translation_plan", concurrency=effective_concurrency,
                         batch_count=len(batches), resume_start_batch=start_batch + 1,
                         segment_count=len(paras))

    checkpoint_lock = threading.RLock()
    commit_lock = threading.RLock()
    active_batches = set()
    runtime_batches_saved = [start_batch]

    def _translation_runtime(operation_id, label, *, saved=False):
        if not (on_status or on_caption):
            return
        with commit_lock:
            if saved:
                runtime_batches_saved[0] += 1
            update_runtime_state(
                job_id, status="running", stage="translation", stage_id="translation",
                operation=operation_id, operation_id=operation_id, operation_label=label,
                phase=operation_id, phase_label=label, progress=True,
                stage_progress={"completed": runtime_batches_saved[0], "total": len(batches),
                                "unit": "批次", "in_flight": [i + 1 for i in sorted(active_batches)]},
                event=label, event_name=operation_id,
                event_visibility="user", event_category="progress")

    # Translation-phase instrumentation. Same vocabulary as the PDF stages: one
    # stage per real sub-phase, wall time and item counts only. Recording a
    # timing never changes whether a batch succeeds, fails or is resumed.
    _translation_stages = {}
    _telemetry_start = runtime_persistence_stats()
    if profiler is not None:
        for _stage_name, _stage_meta in (
            ("translation_generation", {"phase": "generation"}),
            ("translation_deterministic_qa", {"phase": "deterministic_qa"}),
            ("translation_repair", {"phase": "repair"}),
            ("translation_semantic_review", {"phase": "semantic_review"}),
            ("knowledge_feedback", {"phase": "knowledge_feedback"}),
            ("state_commit", {"phase": "state_commit"}),
            ("tm_commit", {"phase": "tm_commit"}),
            ("runtime_telemetry", {"phase": "runtime_telemetry"}),
        ):
            _translation_stages[_stage_name] = profiler.start_stage(
                _stage_name, concurrency=effective_concurrency,
                batch_count=len(batches), metadata=_stage_meta)

    def _record_batch_stage(stage_name, bi, timing, **extra_metadata):
        stage = _translation_stages.get(stage_name)
        if stage is None or not timing:
            return
        started, finished = timing
        stage.item(
            f"batch-{bi + 1}", started, finished_monotonic=finished,
            queue_wait_s=float(extra_metadata.pop("queue_wait_s", 0.0) or 0.0),
            api_latency_s=float(extra_metadata.pop("api_latency_s", 0.0) or 0.0),
            retry_count=int(extra_metadata.pop("retry_count", 0) or 0),
            metadata={"kind": "batch", "batch_index": bi + 1, **extra_metadata},
        )


    def safe_append_event(event):
        with checkpoint_lock:
            try:
                _checkpoint.append_event(job_dir(job_id), event)
            except OSError as exc:
                _workflow_errors.raise_if_fatal(exc, job_id)
                raise

    # ---- experimental context selector runtime (dynamic_v2 only) ------------
    # Nothing in this block exists for the production `fixed` path: no index, no
    # lock, no candidate generation, no dynamic budget.  `_dynamic_runtime` stays
    # None unless the run explicitly opted into the experiment, which makes the
    # isolation observable rather than merely intended.
    _dynamic_runtime = (
        {"index": None, "index_ready": False}
        if effective_context_mode == _dynamic_context.EXPERIMENTAL_CONTEXT_MODE
        else None
    )
    _dynamic_index_lock = threading.Lock() if _dynamic_runtime is not None else None

    def _dynamic_distant_index():
        if _dynamic_runtime is None:  # production path never reaches this
            return None
        if _dynamic_runtime["index_ready"]:
            return _dynamic_runtime["index"]
        with _dynamic_index_lock:
            if not _dynamic_runtime["index_ready"]:
                try:
                    index = _dynamic_context.DistantIndex.build(
                        paras,
                        document_profile=document_profile,
                        glossary_entries=glossary,
                        entity_records=registry.to_list(),
                    )
                except Exception:
                    # A broken index must not break translation: dynamic_v2 then
                    # simply contributes no distant candidates.
                    index = None
                _dynamic_runtime["index"] = index
                _dynamic_runtime["index_ready"] = True
        return _dynamic_runtime["index"]

    inherited_base_url = getattr(_LLM_CTX, "base_url", None)
    inherited_reasoning = getattr(_LLM_CTX, "reasoning_effort", None)
    inherited_job_id = getattr(_RUNTIME_CTX, "job_id", job_id)

    def _process_batch_payload(bi, run_review=True, effective_pairs=None):
        _LLM_CTX.batch_id = bi + 1
        if inherited_base_url is not None:
            _LLM_CTX.base_url = inherited_base_url
        if inherited_reasoning is not None:
            _LLM_CTX.reasoning_effort = inherited_reasoning
        if inherited_job_id is not None:
            _RUNTIME_CTX.job_id = inherited_job_id

        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        batch = batches[bi]
        offset = offsets[bi]
        safe_append_event({
            "batch": bi, "offset": offset, "phase": "generation_started",
            "segment_count": len(batch),
        })
        # Phase timers. `_mark(name)` closes the interval that started at the
        # previous mark, so every stage below reports its own wall time.
        timings = {}
        _phase_mark = [time.perf_counter()]
        _llm_before = llm_thread_stats()

        def _mark(name):
            now = time.perf_counter()
            timings[name] = (_phase_mark[0], now)
            _phase_mark[0] = now

        with commit_lock:
            active_batches.add(bi)
        if on_caption:
            active_str = " ".join(f"#{i + 1}" for i in sorted(active_batches))
            on_caption(f"🌍 正在翻译第 {offset + 1}-{offset + len(batch)} 段（并发批次 {active_str}，共 {len(paras)} 段）...")
        _translation_runtime("context_build", f"正在构建第 {bi + 1} 批翻译上下文")

        try:
            ctx_prev = [
                para for index, para in enumerate(paras[max(0, offset - 2):offset],
                                                 start=max(0, offset - 2))
                if index not in blocked_source_indexes
            ]
            ctx_next = [
                para for index, para in enumerate(
                    paras[min(len(paras), offset + len(batch)):
                          min(len(paras), offset + len(batch) + 2)],
                    start=min(len(paras), offset + len(batch)))
                if index not in blocked_source_indexes
            ]
            with commit_lock:
                pairs_to_query = pairs if effective_pairs is None else effective_pairs
                # 这里只需要"最近 2 段"。原式传 limit=max(2, offset) 再取 [-2:] —— 等价，
                # 但 `select_target_context` 会 list(pairs)[:offset] 并为每个历史段落建字典，
                # 每批 O(offset)、整篇 O(n²)（20k 段实测占 setup 65%）。尾部快路径与它逐字
                # 等价，由 tests/context_target_selection_equivalence_test.py 逐例差分守住。
                previous_target = _context.recent_target_context(
                    pairs_to_query, offset, 2,
                    exclude_indexes=blocked_source_indexes)
            section_digest = _context.digest_for_segment(section_digests, offset)

            # 1) 翻译记忆精确命中直接复用
            batch_pairs = [None] * len(batch)
            to_translate = []  # (index, clean_source)
            blocked_local_indexes = set()
            for i, para in enumerate(batch):
                clean_src = para.replace('\n', ' ')
                source_index = offset + i
                source_flag = blocked_source_indexes.get(source_index)
                if source_flag is not None:
                    blocked_local_indexes.add(i)
                    batch_pairs[i] = {
                        "source": clean_src, "target": "",
                        "initial_target": "",
                        "target_provenance": "source_review_required",
                        "reviewed": False,
                        "review_status": "source_review_required",
                        "from_tm": False,
                        "source_quality": dict(source_flag),
                    }
                    continue
                _tm_t0 = time.perf_counter()
                hit, tm_match = tm_lookup(tm, clean_src, tm_norm_index,
                                          target_lang=_tm_lang)
                _tm_dur_ms = (time.perf_counter() - _tm_t0) * 1000.0
                trace.record_tm_lookup(hit_type=tm_match if hit else "miss",
                                       duration_ms=_tm_dur_ms, job_id=job_id)
                if hit:
                    batch_pairs[i] = {"source": clean_src, "target": hit["target"],
                                      "initial_target": hit["target"],
                                      "accepted_target": hit["target"],
                                      "target_provenance": "tm_approved",
                                      "reviewed": True, "review_status": "tm_approved",
                                      "from_tm": True, "tm_match": tm_match}
                    with commit_lock:
                        state["tm_used_count"] = state.get("tm_used_count", 0) + 1
                elif not has_textual_content(clean_src):
                    batch_pairs[i] = {"source": clean_src, "target": clean_src,
                                      "initial_target": clean_src,
                                      "accepted_target": clean_src,
                                      "target_provenance": "reviewed",
                                      "reviewed": True, "review_status": "reviewed_clean",
                                      "from_tm": False}
                else:
                    to_translate.append((i, clean_src))

            texts = [t for _, t in to_translate]
            section_profile = _batch_section_profile(document_profile, offset, len(batch))
            with commit_lock:
                provisional_hints = _knowledge.provisional_hints(
                    state.get("knowledge_candidates") or [],
                    authoritative_entries=glossary, sources=texts)
                entity_hints = registry.hints_for(
                    texts, glossary_entries=glossary, limit=12)
            _term_t0 = time.perf_counter()
            selected, injected_ids = _select_glossary(
                texts, glossary + provisional_hints,
                document_profile, section_profile)
            _term_dur_ms = (time.perf_counter() - _term_t0) * 1000.0
            trace.record_terminology_lookup(matches=len(selected), duration_ms=_term_dur_ms, job_id=job_id)
            glossary_text = _glossary_block(selected)
            # Experimental-selector bookkeeping; stays inert for `fixed`.
            _dynamic_batch_record = None
            _dynamic_batch_started = 0.0
            with trace.span("translation", operation="context_build", job_id=job_id, items_in=len(texts)):
                if effective_context_mode == _dynamic_context.EXPERIMENTAL_CONTEXT_MODE:
                    # EXPERIMENTAL / RESEARCH / OPT-IN ONLY (not promoted, not
                    # extended).  Same packet shape, same renderer, same
                    # batch/protocol downstream; only *which* context is injected
                    # differs.  The production branch below is byte-identical to
                    # the accepted baseline.
                    context_packet = _dynamic_context.build_dynamic_packet(
                        current_batch=texts,
                        offset=offset,
                        total_segments=len(paras),
                        previous_source=ctx_prev,
                        next_source=ctx_next,
                        previous_target=previous_target,
                        glossary_entries=selected,
                        entity_hints=entity_hints,
                        knowledge_hints=provisional_hints,
                        tm=tm,
                        tm_lang=_tm_lang,
                        document_profile=document_profile,
                        document_synopsis=document_synopsis,
                        section_digest=section_digest,
                        section_profile=section_profile,
                        style_rules=style_rules,
                        entity_records=registry.to_list(),
                        source_segments=paras,
                        pairs=pairs_to_query,
                        blocked_segment_ids=sorted(blocked_source_indexes),
                        local_segment_ids=sorted(
                            list(range(max(0, offset - 2), offset))
                            + list(range(offset + len(batch),
                                         min(len(paras), offset + len(batch) + 2)))),
                        distant_index=_dynamic_distant_index(),
                        token_budget=effective_context_budget,
                        ablations=effective_context_ablations,
                        offline=effective_context_offline,
                        experiment_id=str(state.get("experiment_id") or ""),
                    )
                    _dynamic_batch_record = _record_dynamic_context_batch(
                        state, bi, offset, len(texts), context_packet)
                    _dynamic_batch_started = time.perf_counter()
                else:
                    # PRODUCTION PATH (context_mode = "fixed", the default).
                    # Unchanged from the accepted baseline; the experimental
                    # selector is never consulted here.
                    context_packet = _context.compile_context_packet(
                        document_profile=document_profile,
                        document_synopsis=document_synopsis,
                        section_digest=section_digest,
                        glossary_text=glossary_text,
                        previous_source=ctx_prev,
                        previous_target=previous_target,
                        next_source=ctx_next,
                        current_batch=texts,
                        style_rules=style_rules,
                        entity_hints=entity_hints,
                        context_budget_chars=max(5200, min(12000, effective_max_chars * 3)),
                    )
            review_context = _runtime_review_context(
                state, offset, len(batch), glossary_text, style_rules, target_lang)

            # 2) 未命中段落批次翻译
            _mark("_setup")
            if to_translate:
                _translation_runtime("translation_generation", f"正在翻译第 {offset + 1}–{offset + len(batch)} 段")
                # Per-request prompt cost, aggregated by the caller after the
                # batch commits.  Counts only: the prompt text never leaves the
                # request assembly.
                batch_prompt_cost = []
                try:
                    targets = translate_batch(texts, ctx_prev, ctx_next, glossary_text, style_rules,
                                              target_lang, translator_config["provider"],
                                              translator_config["api_key"],
                                              translator_config["model"],
                                              context_packet=context_packet,
                                              call_llm_fn=translator_call,
                                              prompt_stats=batch_prompt_cost,
                                              result_cache=_result_cache)
                except RuntimeError as batch_error:
                    _workflow_errors.raise_if_fatal(batch_error, job_id)
                    if "任务已请求取消" in str(batch_error):
                        raise
                    failure = {
                        "batch": bi, "offset": offset, "segment_count": len(texts),
                        "reason": str(batch_error)[:500], "recovered": False,
                    }
                    with commit_lock:
                        state.setdefault("translation_failures", []).append(failure)
                    safe_append_event({
                        "batch": bi, "offset": offset,
                        "phase": "translation_protocol_failed",
                        "reason": str(batch_error)[:240],
                    })
                    if len(texts) == 1:
                        with commit_lock:
                            _save_translation_failure()
                        raise
                    if on_caption:
                        on_caption(f"⚠️ 批次 {bi + 1} 翻译返回格式异常，降级为逐段翻译...")
                    targets = []
                    try:
                        for t in texts:
                            if _runtime_cancel_requested(job_id):
                                raise RuntimeError("任务已请求取消")
                            single_packet = dict(context_packet, current_batch=[t])
                            targets.append(translate_batch(
                                [t], ctx_prev, ctx_next, glossary_text, style_rules, target_lang,
                                translator_config["provider"], translator_config["api_key"],
                                translator_config["model"], context_packet=single_packet,
                                call_llm_fn=translator_call,
                                prompt_stats=batch_prompt_cost,
                                result_cache=_result_cache)[0])
                    except Exception as single_error:
                        failure["reason"] = str(single_error)[:500]
                        with commit_lock:
                            _save_translation_failure()
                        raise
                    failure["recovered"] = True
                with commit_lock:
                    _record_translation_prompt_cost(state, batch_prompt_cost)
                for (i, src), tgt in zip(to_translate, targets):
                    cleaned_tgt = clean_xml_chars(tgt).replace('\n', ' ')
                    batch_pairs[i] = {"source": src, "target": cleaned_tgt,
                                      "initial_target": cleaned_tgt,
                                      "target_provenance": "generated",
                                      "reviewed": False, "review_status": "not_reviewed",
                                      "from_tm": False}

            batch_sources = [p["source"] for p in batch_pairs]
            batch_targets = [p["target"] for p in batch_pairs]
            if _dynamic_batch_record is not None:
                # Stamp the real generation wall time on the experimental record
                # now that the batch's provider work is finished.
                _dynamic_batch_record["wall_s"] = round(
                    time.perf_counter() - _dynamic_batch_started, 3)
            for p in batch_pairs:
                p["glossary_entry_ids"] = list(injected_ids)
                p["glossary_hash_used"] = frozen_hash
            safe_append_event({
                "batch": bi, "offset": offset, "phase": "generation_done",
            })
            _mark("translation_generation")
            _translation_runtime("translation_qa", f"正在检查第 {bi + 1} 批译文")
            findings = [
                finding for finding in check_translation_batch(
                    batch_sources, batch_targets, glossary, target_lang,
                    section_profile=section_profile)
                if finding.get("segment_index") not in blocked_local_indexes
            ]
            _mark("translation_deterministic_qa")

            # 3) 确定性问题自动修复（一轮）
            fixable = [
                f for f in findings
                if f["severity"] in ("blocking", "actionable")
                and f.get("segment_index") not in blocked_local_indexes
            ]
            reviewable_local_indexes = [
                index for index in range(len(batch_pairs))
                if index not in blocked_local_indexes
            ]
            reviewable_segment_ids = [offset + index
                                      for index in reviewable_local_indexes]
            repair_overlays_to_record = []
            review_evidence_to_record = []
            if fixable and len(fixable) <= 8:
                if on_caption:
                    on_caption(f"🔧 批次 {bi + 1} 发现 {len(fixable)} 个确定性问题，正在自动修复...")
                try:
                    repair_local_indexes = [
                        index for index in range(len(batch_pairs))
                        if index not in blocked_local_indexes
                    ]
                    repair_local_position = {
                        local_index: position
                        for position, local_index in enumerate(repair_local_indexes)
                    }
                    repair_findings = [
                        {**finding,
                         "segment_index": repair_local_position[finding["segment_index"]]}
                        for finding in fixable
                    ]
                    repaired = repair_batch(
                        [batch_sources[index] for index in repair_local_indexes],
                        [batch_targets[index] for index in repair_local_indexes],
                        repair_findings, glossary_text,
                                            style_rules, target_lang,
                                            translator_config["provider"],
                                            translator_config["api_key"],
                                            translator_config["model"],
                                            call_llm_fn=translator_call)
                    if repaired and len(repaired) == len(repair_local_indexes):
                        formal_targets = list(batch_targets)
                        shadow_targets = list(formal_targets)
                        for position, j in enumerate(repair_local_indexes):
                            p = batch_pairs[j]
                            if (not p["from_tm"] and repaired[position]
                                    and repaired[position].strip()):
                                candidate = clean_xml_chars(
                                    repaired[position]).replace('\n', ' ')
                                if is_incomplete_translation(batch_sources[j], candidate) \
                                        and not is_incomplete_translation(batch_sources[j], formal_targets[j]):
                                    continue
                                shadow_targets[j] = candidate
                        overlay = _repair.create_overlay(
                            formal_targets, shadow_targets, fixable, "deterministic",
                            sources=batch_sources,
                            finding_segment_ids=[offset + f["segment_index"] for f in fixable])
                        shadow_findings = check_translation_batch(
                            batch_sources, shadow_targets, glossary, target_lang,
                            section_profile=section_profile)
                        shadow_findings = _globalize_batch_findings(shadow_findings, offset)
                        shadow_findings = [
                            finding for finding in shadow_findings
                            if (isinstance(finding.get("segment_index"), int)
                                and finding["segment_index"] - offset
                                not in blocked_local_indexes)
                        ]
                        blind_findings, blind_failed, blind_trace = [], False, None
                        if enable_review and not any(
                                f["severity"] in ("blocking", "actionable")
                                for f in shadow_findings):
                            with _critical_path.phase("repair_shadow_state_copy") as _copy_sink:
                                with commit_lock:
                                    current_pairs = list(pairs)
                                    current_findings_all = list(findings_all)
                                _copy_sink.traversed(
                                    len(current_pairs) + len(current_findings_all))
                            shadow_index = _translation_evidence_index(
                                paras, current_pairs, batch_pairs, glossary, document_profile,
                                document_synopsis, section_digests, current_findings_all,
                                blind=True,
                                candidate_targets={offset + j: shadow_targets[j]
                                                   for j in reviewable_local_indexes},
                                blocked_source_indexes=blocked_source_indexes)
                            reviewable_pairs = [batch_pairs[j]
                                                for j in reviewable_local_indexes]
                            reviewable_targets = [shadow_targets[j]
                                                  for j in reviewable_local_indexes]
                            shadow_packet = _translation_evidence.build_runtime_review_packet(
                                state, reviewable_pairs, reviewable_segment_ids,
                                glossary,
                                deterministic_checks=shadow_findings,
                                review_context=review_context,
                                candidate_targets={offset + j: shadow_targets[j]
                                                   for j in reviewable_local_indexes},
                                blind=True,
                            )
                            blind_findings, blind_failed, blind_trace = \
                                _translation_evidence.review_translation_batch_with_evidence(
                                    [batch_sources[j] for j in reviewable_local_indexes],
                                    reviewable_targets, glossary_text, style_rules,
                                    target_lang, reviewer_config["provider"],
                                    reviewer_config["api_key"], reviewer_config["model"],
                                    shadow_index, call_llm=reviewer_call, blind=True,
                                    segment_ids=reviewable_segment_ids,
                                    review_identity={
                                        "input_hash": overlay["input_hash"],
                                        "candidate_hash": overlay["candidate_hash"],
                                    }, translation_core_packet=shadow_packet)
                            if blind_trace:
                                review_evidence_to_record.append({
                                    "batch": bi, "phase": "shadow_repair", **blind_trace})
                        overlay = _repair.evaluate_overlay(
                            overlay, shadow_findings, blind_findings, blind_failed,
                            review_identity=(blind_trace or {}).get("review_identity"))
                        repair_overlays_to_record.append({
                            "batch": bi, "offset": offset, **overlay,
                            "blind_trace": blind_trace,
                        })
                        promoted = _repair.promoted_targets(overlay)
                        for j, p in enumerate(batch_pairs):
                            if not p["from_tm"]:
                                p["target"] = promoted[j]
                    batch_targets = [p["target"] for p in batch_pairs]
                    findings = [
                        finding for finding in check_translation_batch(
                            batch_sources, batch_targets, glossary, target_lang,
                            section_profile=section_profile)
                        if finding.get("segment_index") not in blocked_local_indexes
                    ]
                except Exception as exc:
                    repair_overlays_to_record.append({
                        "batch": bi, "offset": offset, "source": "deterministic",
                        "status": "rejected", "rejection": "repair_error",
                        "error": str(exc)[:240],
                    })
            safe_append_event({
                "batch": bi, "offset": offset, "phase": "deterministic_qa_done",
            })
            _mark("translation_repair")

            # 4) 独立审校
            review_succeeded = False
            rfindings = []
            review_failed = False
            review_trace = None
            review_event_id = None
            formal_records = []
            promoted_review_events = []
            formal_segment_ids = reviewable_segment_ids
            evidence_index = None
            formal_packet = None
            # Deterministic pre-review gate.  In ``always`` mode this is a
            # no-op; in ``risk_selective`` mode a batch with no deterministic
            # risk signal is recorded as skipped instead of reviewed.  The
            # decision is made here, before any review input is built, so a
            # skipped batch pays neither the packet nor the request.
            review_skipped = False
            review_skip_reason = ""
            review_signals = None
            if enable_review and effective_review_mode == _review_gate.REVIEW_MODE_RISK_SELECTIVE \
                    and formal_segment_ids:
                review_signals = _review_gate.batch_signals_from_context(
                    findings=findings,
                    fixable=fixable,
                    repair_attempted=bool(repair_overlays_to_record),
                    batch_pairs=batch_pairs,
                    blocked_local_indexes=blocked_local_indexes,
                    protocol_failure=any(
                        isinstance(item, dict) and item.get("batch") == bi
                        for item in (state.get("translation_failures") or [])),
                    findings_all=findings_all,
                    segment_ids=formal_segment_ids,
                )
                run_review, review_skip_reason = \
                    _review_gate.should_run_independent_review(review_signals)
                review_skipped = not run_review
            if run_review and enable_review and not review_skipped:
                _translation_runtime("translation_review", f"正在审校第 {bi + 1} 批译文")
                if formal_segment_ids:
                    with _critical_path.phase("review_input_state_copy") as _copy_sink:
                        with commit_lock:
                            current_pairs = list(pairs)
                            current_findings_all = list(findings_all)
                        _copy_sink.traversed(len(current_pairs) + len(current_findings_all))
                    _critical_path.count("reviewer_calls_formal")
                    evidence_index = _translation_evidence_index(
                        paras, current_pairs, batch_pairs, glossary, document_profile,
                        document_synopsis, section_digests, current_findings_all,
                        blocked_source_indexes=blocked_source_indexes)
                    reviewable_pairs = [batch_pairs[index]
                                        for index in reviewable_local_indexes]
                    reviewable_sources = [batch_sources[index]
                                          for index in reviewable_local_indexes]
                    reviewable_targets = [batch_targets[index]
                                          for index in reviewable_local_indexes]
                    review_findings = [
                        finding for finding in findings
                        if finding.get("segment_index") not in blocked_local_indexes
                    ]
                    formal_packet = _translation_evidence.build_runtime_review_packet(
                        state, reviewable_pairs, formal_segment_ids, glossary,
                        deterministic_checks=_globalize_batch_findings(
                            review_findings, offset),
                        review_context=review_context,
                    )
                    rfindings, failed, review_trace = \
                        _translation_evidence.review_translation_batch_with_evidence(
                            reviewable_sources, reviewable_targets, glossary_text,
                            style_rules, target_lang, reviewer_config["provider"],
                            reviewer_config["api_key"], reviewer_config["model"],
                            evidence_index, call_llm=reviewer_call,
                            segment_ids=formal_segment_ids,
                            translation_core_packet=formal_packet)
                    review_event_id = (
                        f"translation-review-{job_id}-{bi}-formal-"
                        f"{len(state.get('review_evidence') or []) + len(review_evidence_to_record)}")
                    review_trace["review_event_id"] = review_event_id
                    if failed:
                        review_failed = True
                        for pair in batch_pairs:
                            if (not pair.get("from_tm")
                                    and pair.get("target_provenance")
                                    != "source_review_required"):
                                pair["review_status"] = "review_failed"
                    else:
                        review_succeeded = True
                    global_to_local = {
                        offset + index: index for index in range(len(batch_pairs))
                    }
                    for rf in rfindings:
                        sev = rf.get("severity")
                        if sev not in ("blocking", "actionable", "informational"):
                            continue
                        segment_id = rf.get("segment_id")
                        idx = global_to_local.get(segment_id)
                        if idx is None:
                            continue
                        record = _review_finding_record(rf, review_event_id)
                        if sev == "actionable" and rf.get("suggested_target") \
                                and idx not in blocked_local_indexes \
                                and not batch_pairs[idx]["from_tm"]:
                            with _critical_path.phase("reviewer_suggestion_overlay"):
                                suggested = clean_xml_chars(rf["suggested_target"]).replace('\n', ' ').strip()
                                if suggested:
                                    old_target = batch_pairs[idx]["target"]
                                    overlay = _repair.create_overlay(
                                        [old_target], [suggested], [rf], "review_suggested",
                                        sources=[batch_sources[idx]],
                                        finding_segment_ids=[segment_id])
                                    recheck = _globalize_batch_findings(check_translation_batch(
                                        [batch_sources[idx]], [suggested], glossary, target_lang,
                                        section_profile=section_profile), segment_id)
                                    blind_trace = None
                                    blind_findings = []
                                    blind_failed = False
                                    if not any(f["severity"] in ("blocking", "actionable")
                                               for f in recheck):
                                        with _critical_path.phase("review_input_state_copy") as _copy_sink:
                                            with commit_lock:
                                                cur_pairs = list(pairs)
                                                cur_f_all = list(findings_all)
                                            _copy_sink.traversed(len(cur_pairs) + len(cur_f_all))
                                        blind_index = _translation_evidence_index(
                                            paras, cur_pairs, batch_pairs, glossary, document_profile,
                                            document_synopsis, section_digests, cur_f_all,
                                            blind=True, candidate_targets={segment_id: suggested},
                                            blocked_source_indexes=blocked_source_indexes)
                                        suggested_packet = \
                                            _translation_evidence.build_runtime_review_packet(
                                                state, [batch_pairs[idx]], [segment_id], glossary,
                                                deterministic_checks=recheck,
                                                review_context=review_context,
                                                candidate_targets={segment_id: suggested},
                                                blind=True,
                                            )
                                        with _critical_path.phase("reviewer_blind_validation"):
                                            blind_findings, blind_failed, blind_trace = \
                                                _translation_evidence.review_translation_batch_with_evidence(
                                                    [batch_sources[idx]], [suggested], glossary_text,
                                                    style_rules, target_lang,
                                                    reviewer_config["provider"],
                                                    reviewer_config["api_key"],
                                                    reviewer_config["model"], blind_index,
                                                    call_llm=reviewer_call, blind=True,
                                                    segment_ids=[segment_id], review_identity={
                                                        "input_hash": overlay["input_hash"],
                                                        "candidate_hash": overlay["candidate_hash"],
                                                    }, translation_core_packet=suggested_packet)
                                        _critical_path.count("reviewer_calls_blind")
                                    overlay = _repair.evaluate_overlay(
                                        overlay, recheck, blind_findings, blind_failed,
                                        review_identity=(blind_trace or {}).get("review_identity"))
                                    repair_overlays_to_record.append({
                                        "batch": bi, "offset": offset,
                                        "batch_local_ordinal": idx,
                                        "segment_id": segment_id, **overlay,
                                        "blind_trace": blind_trace,
                                    })
                                    if overlay["status"] == "accepted":
                                        batch_pairs[idx]["target"] = suggested
                                        if blind_trace:
                                            promoted_review_events.append((
                                                blind_trace,
                                                f"translation-review-{job_id}-{bi}-suggested-"
                                                f"{segment_id}-{len(promoted_review_events)}",
                                                segment_id,
                                            ))
                                    else:
                                        record["suggested_target"] = suggested
                                        formal_records.append(record)
                                        if blind_trace:
                                            review_evidence_to_record.append({
                                                "batch": bi, "phase": "suggested_shadow_review",
                                                **blind_trace,
                                            })
                                    continue
                        formal_records.append(record)

                else:
                    review_succeeded = True
                _mark("translation_semantic_review")
            elif enable_review and review_skipped and formal_segment_ids:
                # Explicit skip provenance.  The batch is *not* reviewed, *not*
                # approved and *not* TM-eligible; the audit trail records the
                # deterministic signal set that justified the decision.
                review_event_id = (
                    f"translation-review-{job_id}-{bi}-skipped-low-risk")
                review_trace = {
                    "blind": False,
                    "segment_ids": list(formal_segment_ids),
                    "rounds": [],
                    "requests": [],
                    "evidence_ids": [],
                    "decision": _review_gate.SKIPPED_LOW_RISK,
                    "review_identity": {},
                    "review_mode": effective_review_mode,
                    "skip_reason": review_skip_reason,
                    "risk_signals": dict(review_signals or {}),
                    "completion_receipt": {
                        "status": _review_gate.SKIPPED_LOW_RISK,
                        "reviewed_segment_ids": [],
                        "reviewed_unit_count": 0,
                        "finding_count": 0,
                        "evidence_ids": [],
                        "review_identity": {},
                    },
                }
                _mark("translation_semantic_review_skipped")
            elif not run_review and enable_review and formal_segment_ids and not review_skipped:
                with commit_lock:
                    current_pairs = list(pairs) if effective_pairs is None else list(effective_pairs)
                    current_findings_all = list(findings_all)
                evidence_index = _translation_evidence_index(
                    paras, current_pairs, batch_pairs, glossary, document_profile,
                    document_synopsis, section_digests, current_findings_all,
                    blocked_source_indexes=blocked_source_indexes)
                reviewable_pairs = [batch_pairs[index]
                                    for index in reviewable_local_indexes]
                formal_packet = _translation_evidence.build_runtime_review_packet(
                    state, reviewable_pairs, formal_segment_ids, glossary,
                    deterministic_checks=_globalize_batch_findings(findings, offset),
                    review_context=review_context,
                )

            return {
                "bi": bi,
                "offset": offset,
                "batch": batch,
                "timings": timings,
                "llm_stats": llm_thread_stats_delta(_llm_before),
                "batch_pairs": batch_pairs,
                "batch_sources": batch_sources,
                "injected_ids": injected_ids,
                "entity_hints": entity_hints,
                "context_packet": context_packet,
                "section_digest": section_digest,
                "section_profile": section_profile,
                "blocked_local_indexes": blocked_local_indexes,
                "review_succeeded": review_succeeded,
                "review_skipped": review_skipped,
                "review_skip_reason": review_skip_reason,
                "review_signals": review_signals,
                "rfindings": rfindings,
                "review_failed": review_failed,
                "review_trace": review_trace,
                "review_event_id": review_event_id,
                "formal_records": formal_records,
                "promoted_review_events": promoted_review_events,
                "formal_segment_ids": formal_segment_ids,
                "repair_overlays": repair_overlays_to_record,
                "review_evidence": review_evidence_to_record,
                "findings": findings,
                "review_context": review_context,
                "glossary_text": glossary_text,
                "evidence_index": evidence_index,
                "formal_packet": formal_packet,
            }
        finally:
            with commit_lock:
                active_batches.discard(bi)

    def _commit_batch(bundle):
        bi = bundle["bi"]
        _LLM_CTX.batch_id = bi + 1
        offset = bundle["offset"]
        # Captured before anything in this batch touches state: the journal
        # describes the delta against exactly this baseline, and no save happens
        # between here and the record being built, so the baseline cannot drift.
        _journal_patches = new_journal_patches()
        _journal_before = journal_baseline(state) if _journal_enabled else None
        batch = bundle["batch"]
        batch_pairs = bundle["batch_pairs"]
        batch_sources = bundle["batch_sources"]
        injected_ids = bundle["injected_ids"]
        context_packet = bundle["context_packet"]
        section_digest = bundle["section_digest"]
        section_profile = bundle["section_profile"]
        blocked_local_indexes = bundle["blocked_local_indexes"]
        review_succeeded = bundle["review_succeeded"]
        review_skipped = bool(bundle.get("review_skipped"))
        review_skip_reason = str(bundle.get("review_skip_reason") or "")
        rfindings = bundle["rfindings"]
        review_failed = bundle["review_failed"]
        review_trace = bundle["review_trace"]
        review_event_id = bundle["review_event_id"]
        formal_records = bundle["formal_records"]
        promoted_review_events = bundle["promoted_review_events"]
        formal_segment_ids = bundle["formal_segment_ids"]

        # Record this batch's translation-phase timings. Provider latency and
        # request/failure counts come from the worker thread's own accounting,
        # so retries are visible per batch rather than inferred.
        _llm_delta = bundle.get("llm_stats") or {}
        _segment_count = len(bundle.get("batch") or [])
        with _critical_path.phase("commit_timing_record"):
            for _stage_name, _stage_timing in (bundle.get("timings") or {}).items():
                if str(_stage_name).startswith("_"):
                    continue
                _record_batch_stage(
                    _stage_name, bi, _stage_timing,
                    api_latency_s=(_llm_delta.get("api_latency_s") or 0.0)
                    if _stage_name == "translation_generation" else 0.0,
                    retry_count=int(_llm_delta.get("failures") or 0)
                    if _stage_name == "translation_generation" else 0,
                    segment_count=_segment_count,
                    llm_requests=int(_llm_delta.get("requests") or 0),
                    llm_failures=int(_llm_delta.get("failures") or 0),
                )
        _batch_timings = bundle.get("timings") or {}
        _gen_timing = _batch_timings.get("translation_generation")
        _batch_dur_ms = 0.0
        if _gen_timing:
            _batch_dur_ms = max(0.0, (_gen_timing[1] - _gen_timing[0]) * 1000.0)
        trace.record_batch(
            batch_id=bi + 1,
            segment_count=_segment_count,
            source_chars=sum(len(p.get("source", "")) for p in (bundle.get("batch_pairs") or [])),
            duration_ms=_batch_dur_ms,
            tm_hits=sum(1 for p in (bundle.get("batch_pairs") or []) if p.get("from_tm")),
            model_called=bool(bundle.get("to_translate")),
            job_id=job_id,
        )

        with _critical_path.phase("commit_log_append") as _log_sink:
            state.setdefault("context_packet_log", []).append({
                "batch": bi,
                "offset": offset,
                **_context.context_metadata(context_packet),
            })
            state.setdefault("glossary_injection_log", []).append({
                "batch": bi,
                "offset": offset,
                "entry_ids": injected_ids,
                "glossary_version": (state.get("glossary_frozen") or {}).get("version"),
                "glossary_hash": (state.get("glossary_frozen") or {}).get("glossary_hash"),
                "entity_hint_count": len(bundle.get("entity_hints") or []),
                "document_synopsis_summary": (document_synopsis or {}).get("summary", ""),
                "section_digest_summary": (section_digest or {}).get("summary", ""),
            })
            if bundle.get("repair_overlays"):
                state.setdefault("repair_overlays", []).extend(bundle["repair_overlays"])
            if bundle.get("review_evidence"):
                state.setdefault("review_evidence", []).extend(bundle["review_evidence"])
            _log_sink.traversed(len(state.get("context_packet_log") or []))

        if enable_review and review_skipped:
            # Skipped by deterministic gating: record explicit provenance and
            # count it separately from both "reviewed" and "not reviewed".
            stats["batches_skipped_low_risk"] = \
                stats.get("batches_skipped_low_risk", 0) + 1
            stats["skipped_segments"] = (
                stats.get("skipped_segments", 0) + len(formal_segment_ids))
            if review_trace and review_event_id:
                _translation_evidence.register_runtime_review_event(
                    state, {"batch": bi, **review_trace}, [],
                    review_event_id, formal_segment_ids,
                    phase="risk_selective_skip", patches=_journal_patches)
            safe_append_event({
                "batch": bi, "offset": offset,
                "phase": "semantic_review_skipped_low_risk",
                "reason": review_skip_reason,
            })
        elif enable_review:
            stats["batches_reviewed"] += 1
            if review_failed:
                stats["review_failed"] += 1
            else:
                for rec in formal_records:
                    findings_all.append(rec)
            if review_trace and review_event_id:
                _translation_evidence.register_runtime_review_event(
                    state, {"batch": bi, **review_trace}, formal_records,
                    review_event_id, formal_segment_ids,
                    patches=_journal_patches)
            for promoted_trace, promoted_event_id, promoted_segment_id in \
                    promoted_review_events:
                _translation_evidence.register_runtime_review_event(
                    state, {"batch": bi, **promoted_trace}, [],
                    promoted_event_id, [promoted_segment_id],
                    phase="suggested_shadow_review", patches=_journal_patches)
            safe_append_event({
                "batch": bi, "offset": offset, "phase": "semantic_review_done",
            })
        else:
            safe_append_event({
                "batch": bi, "offset": offset, "phase": "semantic_review_skipped",
            })

        # 审校可能修改过译文：对最终译文整体复验一次确定性检查
        with _critical_path.phase("commit_recheck_deterministic") as _chk_sink:
            _chk_sink.traversed(len(batch_sources) + len(glossary))
            findings = [
                finding for finding in check_translation_batch(
                    batch_sources, [p["target"] for p in batch_pairs], glossary,
                    target_lang, section_profile=section_profile)
                if finding.get("segment_index") not in blocked_local_indexes
            ]

        # 批内冲突检测（跨段同术语多译法）——在 TM 入库前执行
        with _critical_path.phase("commit_conflict_detection") as _cf_sink:
            safe_batch_pairs = []
            for index, pair in enumerate(batch_pairs):
                if index not in blocked_local_indexes:
                    safe_batch_pairs.append(pair)
                else:
                    safe_batch_pairs.append({**pair, "source": "", "target": ""})
            _cf_sink.traversed(len(glossary) * len(safe_batch_pairs))
            for cf in _detect_conflicts(safe_batch_pairs, glossary, sections=sections):
                segment_id = offset + cf["segment_id"]
                cf["segment_index"] = segment_id
                cf["segment_id"] = segment_id
                findings_all.append(cf)

        # 记录仍未解决的确定性问题
        with _critical_path.phase("commit_findings_merge") as _merge_sink:
            for f in findings:
                if f["severity"] in ("blocking", "actionable", "informational"):
                    segment_id = offset + f["segment_index"]
                    record = dict(f)
                    record.update({"segment_id": segment_id,
                                   "segment_index": segment_id,
                                   "severity": f["severity"], "type": "check",
                                   "detector": f.get("detector") or "Deterministic QA"})
                    findings_all.append(record)
            _merge_sink.traversed(len(findings_all))

        # 5) 批后知识反馈：只进入 candidate queue，不改变 frozen glossary.
        _knowledge_started = time.perf_counter()
        _knowledge_started_cpu = time.thread_time()
        accepted_for_knowledge = []
        review_bad_segments = {
            finding.get("segment_id")
            for finding in (rfindings if enable_review and review_succeeded else [])
            if finding.get("severity") in ("blocking", "actionable")
        }
        for j, pair in enumerate(batch_pairs):
            if j in blocked_local_indexes:
                continue
            local_findings = [
                finding for finding in findings
                if finding.get("segment_index") == j
                and finding.get("severity") in ("blocking", "actionable")
            ]
            if (not pair.get("from_tm") and not local_findings
                    and offset + j not in review_bad_segments):
                accepted_for_knowledge.append(j)
        # Everything the project already knows about a name-like form: the
        # frozen glossary, the entity registry (including aliases) and the
        # provisional candidate queue.  Recomputed per batch because the queue
        # grows; the sets are small and the check is a casefolded membership
        # test, so this never reaches the provider.
        _knowledge_known_forms = {
            str(entry.get("source") or "").casefold()
            for entry in glossary if isinstance(entry, dict)
        }
        _knowledge_known_forms.update(
            str(item.get("source") or item.get("source_expression") or "").casefold()
            for item in (state.get("knowledge_candidates") or [])
            if isinstance(item, dict))
        for _record in registry.records:
            _knowledge_known_forms.add(
                str(_record.get("source_form") or "").casefold())
            _knowledge_known_forms.update(
                str(alias or "").casefold()
                for alias in _record.get("aliases") or [])
        _knowledge_known_forms.discard("")
        knowledge_due, knowledge_reason = _knowledge.should_run_knowledge_feedback(
            batch_index=bi, batch_count=len(batches),
            interval=feedback_interval,
            sources=[batch_sources[j] for j in accepted_for_knowledge],
            known_forms=_knowledge_known_forms,
            accepted_count=len(accepted_for_knowledge),
            force_every_batch=force_knowledge_feedback,
        )
        knowledge_cache_hit = False
        if knowledge_due:
            _policy = state["knowledge_feedback_policy"]
            _policy["observed_batches"].append(bi)
            if knowledge_reason == "new_named_entity":
                _policy["signal_batches"].append(bi)
            knowledge_segment_ids = [offset + j for j in accepted_for_knowledge]
            knowledge_source_context = [
                "" if index in blocked_source_indexes else source
                for index, source in enumerate(paras)
            ]
            knowledge_pair_context = []
            for index, pair in enumerate(pairs):
                if index in blocked_source_indexes and isinstance(pair, dict):
                    knowledge_pair_context.append({
                        **pair, "source": "", "target": "",
                        "initial_target": "", "accepted_target": "",
                    })
                else:
                    knowledge_pair_context.append(pair)
            # The only provider call on the commit path; timed apart so a slow
            # auxiliary model is never mistaken for commit CPU.  A cache hit
            # skips the remote call and still runs the full merge below.
            _knowledge_hits_before = _knowledge_cache.hits
            with _critical_path.phase("commit_knowledge_llm") as _kb_sink:
                knowledge_candidates, knowledge_events, knowledge_warning = \
                    _knowledge.observe_batch(
                        [batch_sources[j] for j in accepted_for_knowledge],
                        [batch_pairs[j]["target"] for j in accepted_for_knowledge],
                        knowledge_source_context, knowledge_pair_context,
                        glossary, offset, auxiliary_config["provider"],
                        auxiliary_config["api_key"], auxiliary_config["model"],
                        existing_candidates=state.get("knowledge_candidates") or [],
                        call_llm=auxiliary_call, segment_ids=knowledge_segment_ids,
                        observation_provenance=(
                            "reviewed" if review_succeeded else "generated_continuity"
                        ),
                        cache=_knowledge_cache)
                knowledge_cache_hit = (
                    _knowledge_cache.hits > _knowledge_hits_before)
                if knowledge_cache_hit:
                    _critical_path.count("knowledge_feedback_cache_hit")
                _kb_sink.traversed(len(knowledge_segment_ids))
                _kb_sink.received(sum(len(batch_sources[j])
                                      for j in accepted_for_knowledge))
                _critical_path.count("knowledge_observe_batch_calls")
        else:
            state["knowledge_feedback_policy"]["skipped_batches"].append(bi)
            reasons = state["knowledge_feedback_policy"].setdefault(
                "skip_reasons", {})
            reasons[knowledge_reason] = reasons.get(knowledge_reason, 0) + 1
            knowledge_candidates, knowledge_events, knowledge_warning = \
                state.get("knowledge_candidates") or [], [], None
        for event in knowledge_events:
            event["batch"] = bi
            state.setdefault("knowledge_events", []).append(event)
            if event.get("type") == "entity_observation":
                registry.observe(
                    event.get("source"), event.get("observed_target"),
                    entity_type=_entity_registry.entity_type_from_kind(
                        event.get("kind")),
                    segment_id=event.get("segment_id"),
                    provenance=event.get("provenance") or "generated_observation",
                    confidence=0.7 if event.get("provenance") == "reviewed" else 0.35,
                )
            if event.get("type") == "target_conflict":
                segment_id = event.get("segment_id", offset)
                findings_all.append({
                    "segment_id": segment_id, "segment_index": segment_id,
                    "severity": "actionable",
                    "type": "knowledge_conflict",
                    "category": "terminology_consistency",
                    "summary": "观察到的译法与锁定术语不一致",
                    "source_span": event.get("source"),
                    "target_span": event.get("observed_target"),
                    "explanation": "翻译流中观察到的译法与项目锁定术语的首选译名不同，可能造成术语漂移。",
                    "recommendation": "核对当前语境是否构成合理例外；若不是，请统一为锁定术语的首选译名。",
                    "confidence": None, "detector": "Knowledge QA",
                    "diagnostic_version": 1,
                    "reason": event.get("reason") or "翻译流观察译法与锁定术语不一致",
                    "source": event.get("source"),
                    "observed_target": event.get("observed_target"),
                    "preferred_target": event.get("preferred_target"),
                })
        state["knowledge_candidates"] = knowledge_candidates
        state["translation_continuity"] = list(knowledge_candidates)
        state["entity_registry"] = registry.to_list()
        if knowledge_warning:
            state["knowledge_feedback_failures"] = \
                state.get("knowledge_feedback_failures", 0) + 1
            state.setdefault("knowledge_events", []).append({
                "type": "extract_failed", "batch": bi, "offset": offset,
                "reason": knowledge_warning,
            })
        # Provenance distinguishes three outcomes: a live model call, a replay
        # of an identical already-paid request, and a batch the policy skipped.
        # None of them is ever recorded as one of the others.
        if not knowledge_due:
            safe_append_event({
                "batch": bi, "offset": offset, "phase": "knowledge_feedback_skipped",
                "reason": knowledge_reason,
            })
        elif knowledge_cache_hit:
            _policy = state["knowledge_feedback_policy"]
            _policy["cache_hits"] = _policy.get("cache_hits", 0) + 1
            safe_append_event({
                "batch": bi, "offset": offset,
                "phase": "knowledge_feedback_cache_hit",
                "candidate_count": len(knowledge_candidates),
                "reason": knowledge_reason,
            })
        else:
            _policy = state["knowledge_feedback_policy"]
            _policy["cache_misses"] = _policy.get("cache_misses", 0) + 1
            safe_append_event({
                "batch": bi, "offset": offset, "phase": "knowledge_feedback_done",
                "candidate_count": len(knowledge_candidates),
                "reason": knowledge_reason,
            })
        _knowledge_finished = time.perf_counter()
        _record_batch_stage("knowledge_feedback", bi,
                            (_knowledge_started, _knowledge_finished),
                            candidate_count=len(knowledge_candidates),
                            segment_count=len(batch_pairs))
        # Same interval, recorded on the critical-path timeline so it can be
        # compared against the other commit buckets without re-reading stages.
        _critical_path.profiler().record(
            "commit_knowledge_feedback", _knowledge_finished - _knowledge_started,
            cpu_s=time.thread_time() - _knowledge_started_cpu)

        # 6) 审校通过的段落 -> 翻译记忆
        if enable_review:
            with _critical_path.phase("commit_tm_eligibility") as _tm_sink:
                _tm_sink.traversed(len(batch_pairs) * max(1, len(findings_all)))
                for j, p in enumerate(batch_pairs):
                    if j in blocked_local_indexes:
                        continue
                    segment_id = offset + j
                    seg_findings = [
                        f for f in findings_all
                        if f.get("segment_id", f.get("segment_index")) == segment_id
                    ]
                    if not p.get("from_tm"):
                        if review_skipped:
                            # Not independently reviewed: explicit, and never
                            # promoted to TM or presented as reviewed/approved.
                            p["review_status"] = _review_gate.SKIPPED_LOW_RISK
                            p["reviewed"] = False
                        elif not review_succeeded:
                            p["review_status"] = "review_failed"
                            p["reviewed"] = False
                        elif seg_findings:
                            p["review_status"] = "reviewed_with_findings"
                            p["reviewed"] = False
                        else:
                            p["review_status"] = "reviewed_clean"
                    if review_succeeded and not review_skipped and not seg_findings \
                            and not p["from_tm"] \
                            and _tm_eligible(p["source"], p["target"]):
                        p["reviewed"] = True
                        p["accepted_target"] = p["target"]
                        p["target_provenance"] = "tm_approved" if p.get("from_tm") else "reviewed"
                        if use_tm:
                            _key = tm_put(tm, p["source"], p["target"], _tm_lang)
                            if _key:
                                _norm = tm_normalize(p["source"])
                                if _norm and _norm not in tm_norm_index:
                                    tm_norm_index[_norm] = _key
                        stats["reviewed_segments"] += 1

        _translation_runtime("state_commit", f"正在写入第 {bi + 1} 批译文")
        _snapshot_taken_before = _snapshot_stats["taken"]
        _state_commit_started = time.perf_counter()
        _state_commit_started_cpu = time.thread_time()
        _commit_translation_batch(
            batch_pairs, offset, bi, _journal_before, _journal_patches,
            [_finalization.segment_id(job_id, offset + index, pair)
             for index, pair in enumerate(batch_pairs)])
        _state_commit_finished = time.perf_counter()
        _translation_runtime("state_commit_done", f"已保存第 {bi + 1} 批译文", saved=True)
        _critical_path.profiler().record(
            "commit_state_commit", _state_commit_finished - _state_commit_started,
            cpu_s=time.thread_time() - _state_commit_started_cpu)
        _record_batch_stage("state_commit", bi,
                            (_state_commit_started, _state_commit_finished),
                            segment_count=len(batch_pairs),
                            snapshot_taken=(
                                _snapshot_stats["taken"] > _snapshot_taken_before))
        safe_append_event({
            "batch": bi, "offset": offset, "phase": "state_commit_done",
            "pairs_count": len(pairs),
        })
        tm_entries = _checkpoint.batch_entries([
            pair for pair in batch_pairs
            if not pair.get("from_tm")
            and pair.get("review_status") == "reviewed_clean"
        ]) if enable_review and use_tm and review_succeeded else []
        if tm_entries:
            # Accumulate onto the snapshot window instead of rewriting the whole
            # project TM once per batch.  Recovery does not depend on this
            # write: `checkpoint.reconcile_translation_memory` rebuilds the
            # cache from `state["pairs"]` plus any unacknowledged
            # `tm_promotion_pending` event.
            _pending_tm_entries.extend(tm_entries)
            _pending_tm_window["batch"] = bi
            _pending_tm_window["offset"] = offset


    def _prepare_batch(bi, run_review=True, effective_pairs=None, *, already_queued=False):
        if not already_queued:
            _batch_timing(job_id, "translation", bi + 1, "queued")
        _batch_timing(job_id, "translation", bi + 1, "worker_started", review_in_worker=run_review)
        try:
            bundle = _process_batch_payload(bi, run_review, effective_pairs)
            _batch_timing(job_id, "translation", bi + 1, "result_ready", review_in_worker=run_review)
            return bundle
        except BaseException as exc:
            _batch_timing(job_id, "translation", bi + 1, "worker_failed", error_type=type(exc).__name__)
            raise

    def _submit_batch(executor, bi):
        _batch_timing(job_id, "translation", bi + 1, "queued")
        return executor.submit(_prepare_batch, bi, already_queued=True)

    def _commit_batch_with_timing(bundle):
        bi = bundle["bi"]
        _batch_timing(job_id, "translation", bi + 1, "commit_started")
        status = "failed"
        try:
            _commit_batch(bundle)
            status = "completed"
        finally:
            _batch_timing(job_id, "translation", bi + 1, "commit_finished", status=status)

    overlap_active = (
        review_overlap_enabled(review_overlap, state.get("pipeline_config"))
        and bool(enable_review)
        and effective_concurrency <= 1
    )
    state["batch_plan"]["review_overlap_enabled"] = bool(overlap_active)

    try:
        if overlap_active and (len(batches) - start_batch) > 1:
            in_flight_cap = review_overlap_max_in_flight(
                max_review_in_flight, state.get("pipeline_config")
            )
            strategy_mode = review_overlap_strategy_mode(
                review_overlap_strategy, state.get("pipeline_config")
            )
            state["batch_plan"]["review_overlap_in_flight_cap"] = in_flight_cap
            state["batch_plan"]["review_overlap_strategy"] = strategy_mode
            scheduler = _review_scheduler.ReviewScheduler(
                max_in_flight=in_flight_cap, reviewer_workers=in_flight_cap
            )
            staged_batches = {}
            staged_pairs = {}
            gen_timings = {}
            next_commit = start_batch

            def _drain_one_commit(target_bi):
                bundle = staged_batches.pop(target_bi)
                idle_wait_start = time.perf_counter()
                rev_res = scheduler.wait_for_batch(target_bi)
                idle_wait = time.perf_counter() - idle_wait_start
                if idle_wait > 0.0001:
                    scheduler.record_translator_idle(idle_wait)

                overlap_s = 0.0
                for g_bi, (g_start, g_end) in gen_timings.items():
                    if g_bi > target_bi:
                        overlap_interval = max(
                            0.0,
                            min(rev_res.finished_at_s, g_end) - max(rev_res.started_at_s, g_start),
                        )
                        overlap_s += overlap_interval
                overlap_s = min(rev_res.total_wall_s, overlap_s)
                exposed_s = max(0.0, rev_res.total_wall_s - overlap_s)
                scheduler.record_hidden_latency(overlap_s)
                scheduler.record_exposed_latency(exposed_s)

                bundle["review_succeeded"] = rev_res.review_succeeded
                bundle["review_skipped"] = bool(rev_res.review_skipped)
                bundle["review_skip_reason"] = str(rev_res.review_skip_reason or "")
                bundle["rfindings"] = rev_res.rfindings
                bundle["review_failed"] = rev_res.review_failed
                bundle["review_trace"] = rev_res.review_trace
                bundle["review_event_id"] = rev_res.review_event_id
                bundle["formal_records"] = rev_res.formal_records
                bundle["promoted_review_events"] = rev_res.promoted_review_events
                bundle["repair_overlays"].extend(rev_res.repair_overlays)
                bundle["review_evidence"].extend(rev_res.review_evidence)

                if rev_res.review_failed:
                    for pair in bundle["batch_pairs"]:
                        if (not pair.get("from_tm")
                                and pair.get("target_provenance")
                                != "source_review_required"):
                            pair["review_status"] = "review_failed"

                if rev_res.accepted_suggestions:
                    for idx, suggested in rev_res.accepted_suggestions.items():
                        bundle["batch_pairs"][idx]["target"] = suggested

                bundle["timings"]["translation_semantic_review"] = (
                    rev_res.started_at_s, rev_res.finished_at_s
                )
                bundle["llm_stats"]["api_latency_s"] = (
                    bundle["llm_stats"].get("api_latency_s", 0.0) + rev_res.api_latency_s
                )
                bundle["llm_stats"]["requests"] = (
                    bundle["llm_stats"].get("requests", 0) + rev_res.llm_requests
                )
                bundle["llm_stats"]["failures"] = (
                    bundle["llm_stats"].get("failures", 0) + rev_res.llm_failures
                )

                _commit_batch_with_timing(bundle)
                staged_pairs.pop(target_bi, None)

            try:
                for bi in range(start_batch, len(batches)):
                    _runtime_pause_checkpoint(job_id)
                    if _runtime_cancel_requested(job_id):
                        scheduler.cancel()
                        raise RuntimeError("任务已请求取消")

                    if strategy_mode == "stable":
                        while next_commit < bi:
                            _drain_one_commit(next_commit)
                            next_commit += 1

                    effective_pairs = list(pairs)
                    for prev_bi in sorted(staged_pairs.keys()):
                        if prev_bi < bi:
                            effective_pairs.extend(staged_pairs[prev_bi])

                    gen_start = time.perf_counter()
                    bundle = _prepare_batch(
                        bi, run_review=False, effective_pairs=effective_pairs
                    )
                    gen_end = time.perf_counter()
                    gen_timings[bi] = (gen_start, gen_end)

                    staged_batches[bi] = bundle
                    staged_pairs[bi] = [dict(p) for p in bundle["batch_pairs"]]

                    # Drain ready prior batches (< bi) or when backpressure requires it
                    while next_commit < bi and (
                        scheduler.in_flight_count() >= in_flight_cap
                        or scheduler.has_result(next_commit)
                    ):
                        _drain_one_commit(next_commit)
                        next_commit += 1

                    snapshot = _review_scheduler.BatchReviewSnapshot.create(
                        batch_index=bi,
                        offset=bundle["offset"],
                        batch_pairs=bundle["batch_pairs"],
                        glossary=glossary,
                        glossary_text=bundle.get("glossary_text") or "",
                        glossary_ids=bundle["injected_ids"],
                        glossary_hash=frozen_hash or "",
                        style_rules=style_rules,
                        target_lang=target_lang,
                        section_profile=bundle["section_profile"],
                        deterministic_findings=bundle["findings"],
                        blocked_local_indexes=bundle["blocked_local_indexes"],
                        review_context=bundle["review_context"],
                        formal_segment_ids=bundle["formal_segment_ids"],
                        evidence_index=bundle["evidence_index"],
                        formal_packet=bundle["formal_packet"],
                        translator_model=str(translator_config.get("model") or ""),
                        paras=paras,
                        committed_pairs=list(pairs),
                        findings_all=list(findings_all),
                        document_profile=document_profile or {},
                        document_synopsis=document_synopsis or {},
                        section_digests=section_digests or [],
                        state_ref=state,
                        base_evidence_count=(
                            len(state.get("review_evidence") or [])
                            + len(bundle.get("review_evidence") or [])
                        ),
                        review_gate_skip=bool(bundle.get("review_skipped")),
                        review_gate_reason=str(bundle.get("review_skip_reason") or ""),
                        review_gate_signals=bundle.get("review_signals"),
                    )
                    scheduler.submit(
                        snapshot,
                        lambda s: _execute_batch_review(
                            s, reviewer_call, reviewer_config, job_id, enable_review
                        ),
                    )

                while next_commit < len(batches):
                    if _runtime_cancel_requested(job_id):
                        scheduler.cancel()
                        raise RuntimeError("任务已请求取消")
                    _drain_one_commit(next_commit)
                    next_commit += 1
            finally:
                scheduler.shutdown(wait=False)
                overlap_stats = scheduler.stats_snapshot()
                metrics = _review_scheduler.compute_overlap_metrics(overlap_stats)
                state["review_overlap"] = metrics
                if profiler is not None and "translation_semantic_review" in _translation_stages:
                    _translation_stages["translation_semantic_review"].record["metadata"][
                        "review_overlap"
                    ] = metrics
        elif effective_concurrency <= 1 or (len(batches) - start_batch) <= 1:
            for bi in range(start_batch, len(batches)):
                _runtime_pause_checkpoint(job_id)
                if _runtime_cancel_requested(job_id):
                    raise RuntimeError("任务已请求取消")
                bundle = _prepare_batch(bi)
                _commit_batch_with_timing(bundle)
        else:
            window = min(len(batches) - start_batch, effective_concurrency * 2)
            futures = {}
            with ThreadPoolExecutor(max_workers=effective_concurrency, thread_name_prefix="ft-trans") as executor:
                for bi in range(start_batch, start_batch + window):
                    futures[bi] = _submit_batch(executor, bi)

                for bi in range(start_batch, len(batches)):
                    _runtime_pause_checkpoint(job_id)
                    if _runtime_cancel_requested(job_id):
                        for f in futures.values():
                            f.cancel()
                        raise RuntimeError("任务已请求取消")
                    next_bi = bi + window
                    if next_bi < len(batches):
                        futures[next_bi] = _submit_batch(executor, next_bi)

                    bundle = futures[bi].result()
                    _commit_batch_with_timing(bundle)
                    del futures[bi]
    finally:
        # Mandatory immediate snapshot on leaving the loop: normal completion,
        # user cancellation and fatal failure all converge here, so a committed
        # batch can never be left sitting inside the deferred cadence window.
        # The TM check matters on its own: the last batch's promotions are only
        # accumulated *after* the cadence decision, so a final cadence flush can
        # leave them pending with `_pending_snapshot_batches` already back to 0.
        if _pending_snapshot_batches or _pending_tm_entries \
                or (_journal_enabled
                    and _journal_seq > _state_journal.cursor(state)):
            _flush_state_snapshot("loop_exit")
        if sys.exc_info()[1] is not None and profiler is not None:
            error = sys.exc_info()[1]
            status = "cancelled" if "请求取消" in str(error) else "failed"
            try:
                for stage in _translation_stages.values():
                    stage.finish(status=status, persist=False)
                profiler.persist()
            except Exception:
                pass  # Keep the original workflow/storage failure authoritative.

    # 全局冲突检测（跨批次），与批内结果去重
    batch_conflict_keys = {(f.get("type"), f.get("entry_id"),
                            f.get("segment_index"), True)
                           for f in findings_all if f.get("conflict")}
    safe_pairs = []
    for index, pair in enumerate(pairs):
        if index not in blocked_source_indexes:
            safe_pairs.append(pair)
        else:
            safe_pairs.append({**pair, "source": "", "target": ""})
    for cf in _detect_conflicts(safe_pairs, glossary, sections=sections):
        cf["segment_index"] = cf["segment_id"]
        key = (cf.get("type"), cf.get("entry_id"), cf.get("segment_index"), True)
        if key not in batch_conflict_keys:
            findings_all.append(cf)
    state["entity_registry"] = registry.to_list()
    existing_entity_conflicts = {
        (item.get("type"), item.get("segment_index"), item.get("source"))
        for item in findings_all
    }
    for conflict in registry.consistency_findings(pairs):
        key = (conflict.get("type"), conflict.get("segment_index"),
               conflict.get("source"))
        if key not in existing_entity_conflicts:
            findings_all.append(conflict)
            existing_entity_conflicts.add(key)

    state["state_journal"] = {
        "enabled": bool(_journal_enabled),
        "records": _journal_stats["records"],
        "bytes": _journal_stats["bytes"],
        "unusable_batches": _journal_stats["unusable"],
        "budget_bytes": _journal_budget_bytes,
        "cursor": _state_journal.cursor(state),
        "snapshots": _snapshot_stats["taken"],
        "deferred_batches": _snapshot_stats["deferred"],
    }
    stats["blocking"] = sum(1 for f in findings_all if f["severity"] == "blocking")
    stats["actionable"] = sum(1 for f in findings_all if f["severity"] == "actionable")
    stats["informational"] = sum(1 for f in findings_all if f["severity"] == "informational")
    state["has_blocking"] = stats["blocking"] > 0

    # Runtime-telemetry cost for this run: how much wall time the (now
    # coalesced) runtime persistence actually consumed.  Recorded as a stage so
    # it can be compared against generation/review without touching the state
    # machine.
    if profiler is not None and _translation_stages:
        with _critical_path.phase("commit_telemetry"):
            _telemetry_end = runtime_persistence_stats()
            _critical_path.count("telemetry_state_writes",
                                 _telemetry_end["state_writes"]
                                 - _telemetry_start["state_writes"])
            _critical_path.count("telemetry_event_appends",
                                 _telemetry_end["event_appends"]
                                 - _telemetry_start["event_appends"])
            _critical_path.count("telemetry_fsyncs",
                                 _telemetry_end["fsyncs"]
                                 - _telemetry_start["fsyncs"])
        _telemetry_stage = _translation_stages.get("runtime_telemetry")
        if _telemetry_stage is not None:
            _telemetry_stage.record["metadata"].update({
                "state_writes": int(_telemetry_end["state_writes"]
                                    - _telemetry_start["state_writes"]),
                "event_appends": int(_telemetry_end["event_appends"]
                                     - _telemetry_start["event_appends"]),
                "fsyncs": int(_telemetry_end["fsyncs"]
                              - _telemetry_start["fsyncs"]),
                "state_write_seconds": round(
                    _telemetry_end["state_write_seconds"]
                    - _telemetry_start["state_write_seconds"], 6),
                "event_append_seconds": round(
                    _telemetry_end["event_append_seconds"]
                    - _telemetry_start["event_append_seconds"], 6),
                "batch_count": len(batches),
            })
            _telemetry_stage.finish(status="completed", persist=False)
        for _stage_name, _stage in _translation_stages.items():
            if _stage_name != "runtime_telemetry":
                _stage.finish(status="completed", persist=False)
        # One artifact write for the whole translation phase instead of one per
        # stage: the stages share a single performance.json, so closing them
        # together must not re-serialize every prior stage eight times.
        profiler.persist()
    if use_translation_cache:
        with commit_lock:
            _flush_result_cache()
            state["translation_cache"] = _result_cache.stats()
    if use_knowledge_cache:
        with commit_lock:
            _flush_knowledge_cache()
            state["knowledge_cache"] = _knowledge_cache.stats()
    return state



# ================= 文档/表格生成 =================
EN_FONT = "Times New Roman"
CN_FONT = "宋体"

# 自动标注三色：生僻词=红、专业名词（特殊译法）=黄、翻译难点句=青绿
ANNOTATION_COLORS = {"rare": "C00000", "domain": "BF8F00", "hard": "008080"}
ANNOTATION_LABELS = {"rare": "生僻词/难词", "domain": "专业名词（特殊译法）",
                     "hard": "翻译难点句（特别译法）"}



def _is_markdown_table_separator(line):
    cells = _markdown_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)



# ================= 任务持久化（真正的断点续传）=================
def _ensure_output_dir():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def is_onboarded():
    """首次使用引导标记：成功配置并测试通过 AI 引擎后置位。"""
    return (OUTPUT_DIR / ".onboarded").exists()


def mark_onboarded():
    _ensure_output_dir()
    (OUTPUT_DIR / ".onboarded").touch()


def provider_config_path():
    return OUTPUT_DIR / "provider_config.json"


def save_provider_config(provider, model, api_key, base_url=None, reviewer=None,
                         reasoning_effort=None):
    """把 AI 引擎配置落盘（本地单机工具，0600 权限），重启后自动加载。"""
    _ensure_output_dir()
    normalized_base = normalize_openai_base_url(base_url) if base_url else ""
    normalized_reasoning = (str(reasoning_effort or "").strip().lower()
                            if str(reasoning_effort or "").strip().lower()
                            in _REASONING_EFFORT_VALUES else "")
    cfg = {
        "provider": provider or "",
        "model": model or "",
        "api_key": api_key or "",
        "base_url": normalized_base,
        "reasoning_effort": normalized_reasoning,
    }
    if isinstance(reviewer, dict) and reviewer.get("provider") and reviewer.get("model"):
        cfg["reviewer"] = {
            "provider": reviewer.get("provider", ""),
            "model": reviewer.get("model", ""),
            "api_key": reviewer.get("api_key", ""),
            "base_url": normalize_openai_base_url(reviewer.get("base_url"))
            if reviewer.get("base_url") else "",
        }
    path = provider_config_path()
    temp_path = path.with_suffix(path.suffix + ".tmp")
    temp_path.write_text(json.dumps(cfg, ensure_ascii=False, indent=2),
                         encoding="utf-8")
    try:
        temp_path.chmod(0o600)
    except OSError:
        pass
    temp_path.replace(path)
    return path


def load_provider_config():
    """读取已保存的 AI 引擎配置；文件缺失或损坏时返回 None。"""
    try:
        path = provider_config_path()
        if not path.exists():
            return None
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict) or not data.get("provider"):
            return None
        result = {
            "provider": str(data.get("provider", "")),
            "model": str(data.get("model", "") or ""),
            "api_key": str(data.get("api_key", "") or ""),
            "base_url": normalize_openai_base_url(data.get("base_url"))
            if data.get("base_url") else "",
            "reasoning_effort": str(data.get("reasoning_effort", "") or "")
            .strip().lower(),
        }
        if result["reasoning_effort"] not in _REASONING_EFFORT_VALUES:
            result["reasoning_effort"] = ""
        reviewer = data.get("reviewer")
        if isinstance(reviewer, dict) and reviewer.get("provider"):
            result["reviewer"] = {
                "provider": str(reviewer.get("provider", "")),
                "model": str(reviewer.get("model", "") or ""),
                "api_key": str(reviewer.get("api_key", "") or ""),
                "base_url": normalize_openai_base_url(reviewer.get("base_url"))
                if reviewer.get("base_url") else "",
            }
        return result
    except Exception:  # noqa: BLE001 - 配置文件损坏时按未保存处理
        return None


def job_dir(job_id):
    return OUTPUT_DIR / job_id


def job_state_path(job_id):
    return job_dir(job_id) / "state.json"


RUNTIME_SCHEMA_VERSION = "2"
RUNTIME_STATE_FILE = "runtime_state.json"
RUNTIME_EVENTS_FILE = "runtime_events.jsonl"
RUNTIME_TECHNICAL_LOG = "runtime_technical.log"
RUNTIME_HEARTBEAT_SECONDS = 5
RUNTIME_LEASE_SECONDS = 60
RUNTIME_STALL_SECONDS = 45
RUNTIME_ACTIVE_STATUSES = {
    "resume_requested", "queued", "starting", "running",
    "waiting_external", "cancelling", "pausing", "paused",
}

_RUNTIME_EVENT_META = {
    "resume_requested": ("user", "lifecycle"),
    "pipeline_resumed": ("user", "progress"),
    "llm_request_started": ("user", "external_wait"),
    "llm_response_received": ("user", "progress"),
    "job_completed": ("user", "lifecycle"),
    "job_failed": ("user", "error"),
    "job_cancelled": ("user", "lifecycle"),
    "cancel_requested": ("user", "lifecycle"),
    "interrupted": ("user", "error"),
    "stalled": ("user", "error"),
    "retry_requested": ("user", "lifecycle"),
    "job_queued": ("technical", "orchestration"),
    "worker_started": ("technical", "orchestration"),
    "worker_released": ("technical", "orchestration"),
    "job_checkpointed": ("technical", "checkpoint"),
    "checkpoint_saved": ("technical", "checkpoint"),
    "state_flushed": ("technical", "checkpoint"),
    "lease_acquired": ("technical", "orchestration"),
    "lease_renewed": ("technical", "orchestration"),
    "evidence": ("user", "progress"),
    "literature_evidence": ("user", "progress"),
    "research_model": ("user", "progress"),
    "literature_claims": ("user", "progress"),
    "argument_plan": ("user", "progress"),
    "selected_cases": ("user", "progress"),
    "outline": ("user", "progress"),
    "sections": ("user", "progress"),
    "validation": ("user", "progress"),
    "review": ("user", "progress"),
    "quality_repair": ("user", "progress"),
    "academic_quality": ("user", "progress"),
}


def runtime_state_path(job_id):
    return job_dir(job_id) / RUNTIME_STATE_FILE


def runtime_events_path(job_id):
    return job_dir(job_id) / RUNTIME_EVENTS_FILE


def runtime_technical_log_path(job_id):
    return job_dir(job_id) / RUNTIME_TECHNICAL_LOG


def _utc_now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _runtime_defaults():
    return {
        "schema_version": RUNTIME_SCHEMA_VERSION,
        "status": "idle",
        "pipeline": "document_pipeline",
        "stage": "",
        "stage_id": "",
        "stage_index": None,
        "stage_total": None,
        "operation": "",
        "operation_id": "",
        "operation_label": "",
        "stage_progress": {},
        "section_id": None,
        "phase": "",
        "phase_label": "",
        "started_at": None,
        "operation_started_at": None,
        "last_heartbeat_at": None,
        "last_progress_at": None,
        "completed_units": 0,
        "total_units": 0,
        "overall_progress": None,
        "last_event": "",
        "events": [],
        "last_technical_event": "",
        "cancel_requested": False,
        "pause_requested": False,
        "attempt": 0,
        "resume_request_id": None,
        "error": None,
        "worker": {
            "owner_pid": None,
            "worker_id": None,
            "lease_expires_at": None,
        },
    }


def load_runtime_state(job_id):
    """Read the durable worker status without changing the workflow state."""
    if not job_id:
        return _runtime_defaults()
    try:
        raw = json.loads(runtime_state_path(job_id).read_text(encoding="utf-8"))
    except (OSError, TypeError, ValueError):
        return _runtime_defaults()
    if not isinstance(raw, dict):
        return _runtime_defaults()
    result = _runtime_defaults()
    result.update(raw)
    result["schema_version"] = RUNTIME_SCHEMA_VERSION
    if result.get("status") == "cancel_requested":
        result["status"] = "cancelling"
    worker = result.get("worker")
    result["worker"] = {**_runtime_defaults()["worker"], **worker} \
        if isinstance(worker, dict) else dict(_runtime_defaults()["worker"])
    result["events"] = [item for item in result.get("events") or []
                         if isinstance(item, dict)
                         and item.get("visibility") == "user"][-24:]
    return result


def _runtime_event_code(message):
    value = re.sub(r"[^a-z0-9]+", "_", str(message or "").lower()).strip("_")
    return value[:80] or "runtime_event"


def _runtime_event_meta(event_name, visibility=None, category=None):
    default_visibility, default_category = _RUNTIME_EVENT_META.get(
        str(event_name or ""), ("technical", "debug"))
    return visibility or default_visibility, category or default_category


def _append_runtime_event(job_id, event, *, stage=None, operation=None, metadata=None,
                          visibility=None, category=None):
    event_name = event.get("event") if isinstance(event, dict) else _runtime_event_code(event)
    visibility, category = _runtime_event_meta(event_name, visibility, category)
    record = {
        "timestamp": _utc_now_iso(),
        "job_id": job_id,
        "pipeline": event.get("pipeline") if isinstance(event, dict) else None,
        "stage": stage or "",
        "operation": operation or "",
        "event": event_name,
        "message": event.get("message") if isinstance(event, dict) else str(event),
        "visibility": visibility,
        "category": category,
    }
    if record["pipeline"] is None:
        record["pipeline"] = "document_pipeline"
    if metadata:
        record["metadata"] = metadata
    path = runtime_events_path(job_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    started = time.perf_counter()
    try:
        with path.open("a", encoding="utf-8") as stream:
            stream.write(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n")
            stream.flush()
            os.fsync(stream.fileno())
    except OSError as exc:
        _workflow_errors.raise_if_fatal(exc, job_id)
        raise
    with _RUNTIME_PERSIST_LOCK:
        _RUNTIME_PERSIST_STATS["event_appends"] += 1
        _RUNTIME_PERSIST_STATS["fsyncs"] += 1
        _RUNTIME_PERSIST_STATS["event_append_seconds"] += (
            time.perf_counter() - started)
    return record


def _runtime_event_lines_reverse(path):
    """Read complete JSONL lines from the tail without loading the whole log."""
    with path.open("rb") as handle:
        position = handle.seek(0, os.SEEK_END)
        pending = b""
        while position:
            size = min(position, 8192)
            position -= size
            handle.seek(position)
            lines = (handle.read(size) + pending).split(b"\n")
            pending = lines[0]
            for line in reversed(lines[1:]):
                if line:
                    yield line.decode("utf-8")
        if pending:
            yield pending.decode("utf-8")


def read_runtime_events(job_id, limit=None, *, visibility=None, category=None):
    path = runtime_events_path(job_id)
    if not path.is_file():
        return []
    events = []
    try:
        newest_first = bool(limit and limit > 0)
        lines = (_runtime_event_lines_reverse(path) if newest_first else
                 path.read_text(encoding="utf-8").splitlines())
        for line in lines:
            try:
                value = json.loads(line)
            except (TypeError, ValueError):
                continue
            if isinstance(value, dict):
                event_visibility, event_category = _runtime_event_meta(
                    value.get("event"), value.get("visibility"), value.get("category"))
                normalized = {**value, "visibility": event_visibility,
                              "category": event_category}
                if visibility and event_visibility != visibility:
                    continue
                if category and event_category != category:
                    continue
                events.append(normalized)
                if newest_first and len(events) >= limit:
                    break
    except OSError:
        return []
    if newest_first:
        return list(reversed(events))
    return events[-limit:] if limit else events


# ================= 运行时遥测合并（高并发下的写放大治理）=================
# 背景：`call_llm` 过去对**每个请求**做两次 `update_runtime_state`，每次都是
# 「读全量 runtime_state.json → 合并 → 整文件重写」，并在 append runtime event
# 时 flush + fsync。并发 4 时，一次批次循环就能产生几十次全量重写和 fsync，
# 而这些写入里没有一个是业务真值。
#
# 现在把「在途请求状态」这类高频、非真值遥测改为**前缘立即 + 尾缘合并**：
#   * 前缘：窗口内第一次更新立即落盘（UI 马上能看到"等待模型响应"）；
#   * 尾缘：窗口内后续更新合并成一次写入（默认 350 ms，落在 250–500 ms 区间）；
#   * 在途请求数归零时立即落盘（UI 立刻回到"处理中/完成"）；
#   * 失败立即落盘（不能把失败藏在一个还没到期的窗口里）。
#
# 业务真值不受影响：取消、失败、worker lease、阶段边界和业务 checkpoint 仍然
# 直接调用 `update_runtime_state`，走原有同步持久化路径。
RUNTIME_TELEMETRY_FLUSH_SECONDS = 0.35

_RUNTIME_PERSIST_LOCK = threading.RLock()
_RUNTIME_PERSIST_STATS = {
    "state_writes": 0,
    "event_appends": 0,
    "fsyncs": 0,
    "state_write_seconds": 0.0,
    "event_append_seconds": 0.0,
}



_TELEMETRY_LOCK = threading.RLock()
_TELEMETRY_COALESCERS = {}
_TELEMETRY_COALESCER_STATS = {"created": 0}



def _runtime_telemetry_request_started(job_id):
    coalescer = _runtime_telemetry_coalescer(job_id)
    if coalescer is not None:
        coalescer.request_started()


def _runtime_telemetry_request_finished(job_id):
    coalescer = _runtime_telemetry_coalescer(job_id)
    if coalescer is not None:
        coalescer.request_finished()



def update_runtime_state(job_id, *, event=None, progress=False, heartbeat=False,
                         event_name=None, event_metadata=None,
                         event_visibility=None, event_category=None, **changes):
    """Atomically merge a worker status update into runtime_state.json."""
    if not job_id:
        return _runtime_defaults()
    now = _utc_now_iso()
    with _RUNTIME_LOCK, _job_lock.exclusive(job_dir(job_id) / ".runtime-write.lock"):
        current = load_runtime_state(job_id)
        # ``None`` is normally omitted so optional status fields do not erase
        # an existing value accidentally.  ``error`` is different: callers
        # pass ``error=None`` when a retry or a successful checkpoint clears a
        # stale failure from an earlier attempt.  Keep that explicit clear
        # while retaining the omission behavior for other optional fields.
        cancel_active = current.get("cancel_requested") and changes.get("cancel_requested") is not False
        if cancel_active and changes.get("status") in {"running", "waiting_external"}:
            changes = {**changes, "status": "cancelling", "phase": "cancelling",
                       "phase_label": "正在取消"}
        if current.get("pause_requested") and changes.get("pause_requested") is not False and changes.get("status") in {"running", "waiting_external"}:
            changes = {**changes, "status": current.get("status") if current.get("status") == "paused" else "pausing"}
        if changes.get("operation_label") and changes["operation_label"] != current.get("operation_label"):
            changes["operation_started_at"] = now
        if changes.get("phase") and changes["phase"] != current.get("phase"):
            changes["phase_started_at"] = now
        if changes.get("status") in {"completed", "cancelled", "failed"}:
            changes["finished_at"] = now
        current.update({key: value for key, value in changes.items()
                        if value is not None or key == "error"})
        if heartbeat:
            current["last_heartbeat_at"] = now
            worker = dict(current.get("worker") or {})
            if worker.get("worker_id"):
                worker["lease_expires_at"] = (
                    datetime.now(timezone.utc) + timedelta(seconds=RUNTIME_LEASE_SECONDS)
                ).isoformat(timespec="seconds")
                current["worker"] = worker
        if progress:
            current["last_progress_at"] = now
        if event:
            message = str(event).strip()
            if message:
                event_name = event_name or _runtime_event_code(message)
                visibility, category = _runtime_event_meta(
                    event_name, event_visibility, event_category)
                current["last_activity_at"] = now
                inline = {
                    "stage": current.get("stage_id") or current.get("stage"),
                    "operation": current.get("operation_id") or current.get("operation"),
                    "at": now, "timestamp": now, "event": event_name,
                    "message": message, "visibility": visibility,
                    "category": category,
                }
                if event_metadata:
                    inline["metadata"] = event_metadata
                duplicate = visibility == "user" and bool(current.get("events")) \
                    and current["events"][-1].get("event") == event_name \
                    and current["events"][-1].get("message") == message \
                    and current["events"][-1].get("metadata") == event_metadata
                if visibility == "user":
                    current["last_event"] = message
                    if not duplicate:
                        current["events"] = (current.get("events") or []) + [inline]
                        current["events"] = current["events"][-24:]
                else:
                    current["last_technical_event"] = message
                if not duplicate:
                    _append_runtime_event(
                        job_id,
                        {"event": event_name, "message": message,
                         "pipeline": current.get("pipeline")},
                        stage=current.get("stage_id") or current.get("stage"),
                        operation=current.get("operation_id") or current.get("operation"),
                        metadata=event_metadata, visibility=visibility, category=category)
        directory = job_dir(job_id)
        directory.mkdir(parents=True, exist_ok=True)
        tmp = directory / f"{RUNTIME_STATE_FILE}.{uuid.uuid4().hex}.tmp"
        started = time.perf_counter()
        try:
            tmp.write_text(json.dumps(current, ensure_ascii=False, indent=2),
                           encoding="utf-8")
            tmp.replace(runtime_state_path(job_id))
        except OSError as exc:
            _workflow_errors.raise_if_fatal(exc, job_id)
            raise
        finally:
            tmp.unlink(missing_ok=True)
        with _RUNTIME_PERSIST_LOCK:
            _RUNTIME_PERSIST_STATS["state_writes"] += 1
            _RUNTIME_PERSIST_STATS["state_write_seconds"] += (
                time.perf_counter() - started)
        return current


def _runtime_pid_alive(pid):
    if not isinstance(pid, int) or pid <= 0:
        return False
    try:
        os.kill(pid, 0)
    except OSError:
        return False
    return True


def _runtime_parse_time(value):
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
    except (TypeError, ValueError):
        return None


def _runtime_worker_registered(job_id, worker_id=None):
    with _RUNTIME_WORKERS_LOCK:
        worker = _RUNTIME_WORKERS.get(job_id)
        return bool(worker and worker.is_alive() and (
            worker_id is None or getattr(worker, "worker_id", None) == worker_id))


def _annotation_blocks_completion(state):
    """标注是否仍然阻止"业务完成"。

    只有**显式要求**标注（`enable_annotate is True`）且尚未完成时才阻止。
    缺失该字段不能当作"要求标注"：产品自身的默认是关闭
    （`DELIVERY_CONFIG_DEFAULTS["enable_annotate"] = False`），而旧任务里往往
    根本没有这个键。把它当作 True 会让一个已经完成的任务被报成
    `idle_incomplete`，于是概览一边说"可以准备交付"、一边渲染出「继续处理」。
    """
    if not isinstance(state, dict):
        return False
    return state.get("enable_annotate") is True and (
        not state.get("annotations_done") or workflow_outcomes(state)["annotation"]["status"] == "partial")


def _runtime_business_complete(state):
    return bool(state and state.get("p1_done") and state.get("p2_done") and (
        state.get("p3_done") or not state.get("report_enabled", True)) and not
        _annotation_blocks_completion(state))


def _lost_worker_message(runtime, state, *, stalled=False):
    """中断/停滞时给用户一个有信息量的原因。

    以前这里只有一句笼统的"上次运行已中断"，用户无法判断是模型卡住了、还是应用
    被关掉了。现在带上：中断时正在做什么、已经提交了多少段、能否从断点继续。
    """
    label = str(runtime.get("phase_label") or runtime.get("operation_label")
                or runtime.get("phase") or "").strip()
    pairs = len((state or {}).get("pairs") or [])
    paras = len((state or {}).get("paras") or [])
    if stalled:
        parts = ["超过预期时间没有收到运行信号"]
    else:
        parts = ["上次运行的进程已退出（应用被关闭或中断），翻译线程随之终止"]
    if label:
        parts.append(f"中断时正在：{label}")
    if paras:
        parts.append(f"已完成 {pairs}/{paras} 段，可从断点继续")
    return "；".join(parts)


def _runtime_mark_lost(job_id, status, message):
    worker = load_runtime_state(job_id).get("worker") or {}
    return update_runtime_state(
        job_id, status=status, phase=status, phase_label=message,
        worker=worker if _job_execution_alive(job_id) else {
            "owner_pid": None, "worker_id": None, "lease_expires_at": None},
        event=message, event_name=status)


def _job_execution_alive(job_id):
    """A stale lease does not prove that its owner has stopped executing."""
    if is_job_worker_alive(job_id):
        return True
    owner = (load_runtime_state(job_id).get("worker") or {}).get("owner_pid")
    return owner != os.getpid() and _runtime_pid_alive(owner)


def get_job_runtime_status(job_id, state=None):
    """Return the durable runtime status and reconcile dead local workers."""
    state = state if isinstance(state, dict) else load_job_state(job_id)
    runtime = load_runtime_state(job_id)
    failure = _workflow_errors.failure(job_id)
    if failure is not None and failure.reason == "disk_full":
        return {**runtime, "status": "failed", "phase": "failed", "phase_label": "磁盘空间不足",
                "error": {"type": "WorkflowStopped", "message": "磁盘空间不足，任务已停止发起新请求；请释放空间后重试。"}}
    status = runtime.get("status") or "idle"
    if status == "completed" and state and _annotation_blocks_completion(state):
        return {**runtime, "status": "idle_incomplete", "phase": "idle_incomplete",
                "phase_label": "正文已完成，标注仍有失败批次"}
    if status in RUNTIME_ACTIVE_STATUSES:
        worker = runtime.get("worker") or {}
        pid = worker.get("owner_pid")
        worker_id = worker.get("worker_id")
        lease = _runtime_parse_time(worker.get("lease_expires_at"))
        heartbeat = _runtime_parse_time(runtime.get("last_heartbeat_at"))
        last_progress = _runtime_parse_time(runtime.get("last_progress_at"))
        now = datetime.now(timezone.utc)
        registered = _runtime_worker_registered(job_id, worker_id)
        transition_age = (now - last_progress).total_seconds() if last_progress else None
        if status == "resume_requested" and not pid \
                and transition_age is not None and transition_age <= RUNTIME_STALL_SECONDS:
            return runtime
        if not _runtime_pid_alive(pid):
            return _runtime_mark_lost(job_id, "interrupted",
                                   _lost_worker_message(runtime, state))
        if lease and lease < now or heartbeat and (
                now - heartbeat).total_seconds() > RUNTIME_STALL_SECONDS:
            if registered or _job_execution_alive(job_id):
                return _runtime_mark_lost(job_id, "stalled",
                                       _lost_worker_message(
                                           runtime, state, stalled=True))
            return _runtime_mark_lost(job_id, "interrupted",
                                   _lost_worker_message(runtime, state))
        if not registered and pid == os.getpid():
            # Streamlit can re-execute the app module while a worker thread
            # from the same process is still handling a long model request.
            # A module reload recreates the in-memory registry, but the durable
            # lease and heartbeat still prove that this process owns the run.
            # Do not turn that live request into a false "interrupted" state.
            fresh_heartbeat = heartbeat and (
                now - heartbeat).total_seconds() <= RUNTIME_STALL_SECONDS
            lease_valid = not lease or lease >= now
            if fresh_heartbeat and lease_valid:
                return runtime
            if status in {"queued", "starting"} \
                    and transition_age is not None and transition_age <= RUNTIME_STALL_SECONDS:
                return runtime
            return _runtime_mark_lost(job_id, "interrupted",
                                      _lost_worker_message(runtime, state))
        return runtime
    if status == "idle":
        inferred = "completed" if _runtime_business_complete(state) else "idle_incomplete"
        if inferred != status:
            # 只影响本次读取的返回值，**不落盘**：读取不应有写副作用。
            # 此前的实现会在这里 update_runtime_state，于是"看一眼任务"就会给
            # 从未运行过的任务写出 runtime_state.json。真实状态由 worker 在任务
            # 运行时写入；这里的推断每次都能重新得出同样结果，不需要持久化。
            return {**runtime, "status": inferred, "phase": inferred,
                    "phase_label": "已完成" if inferred == "completed" else "尚未完成"}
    return runtime


#: Canonical runtime-action policy: one status -> the actions a user may take.
#: Every product surface (Qt workbench, Streamlit runtime strip, CLI) renders
#: from this table instead of re-deriving legality from status literals.
#: ``pause``/``continue`` are runtime capabilities (``request_job_pause`` /
#: ``continue_job``), not presentation choices, which is why they live here.
RUNTIME_ACTIONS = {
    "resume_requested": ("details",),
    "queued": ("cancel", "details"),
    "starting": ("cancel", "details"),
    "running": ("pause", "cancel", "details"),
    "waiting_external": ("pause", "cancel", "details"),
    "stalled": ("retry", "cancel", "details"),
    "interrupted": ("resume", "retry", "details"),
    "failed": ("retry", "resume", "details"),
    "pausing": ("cancel", "details"),
    "paused": ("continue", "resume", "cancel", "details"),
    "cancelling": ("details",),
    "cancelled": ("resume", "details"),
    "idle_incomplete": ("resume",),
    "waiting_manual": ("view",),
    "completed": ("view",),
}

#: The single action a surface should emphasise for a status (or ``None``).
PRIMARY_RUNTIME_ACTIONS = {
    "idle_incomplete": "resume", "interrupted": "resume",
    "cancelled": "resume", "failed": "retry", "stalled": "retry",
}


def runtime_actions_for_status(status):
    """Return the canonical user actions for one runtime status.

    Read-only and side-effect free: it is the policy shared by every frontend,
    so it must not depend on the caller or on presentation concerns.
    """
    return list(RUNTIME_ACTIONS.get(str(status or ""), ()))


def primary_runtime_action_for_status(status):
    """Return the emphasised action for a status, or ``None``."""
    return PRIMARY_RUNTIME_ACTIONS.get(str(status or ""))


def build_job_runtime_view(job_id, state=None):
    """Canonical user-facing runtime view shared by every product surface."""
    state = state if isinstance(state, dict) else load_job_state(job_id)
    runtime = get_job_runtime_status(job_id, state)
    status = runtime.get("status") or "idle_incomplete"
    if status == "idle" and not _runtime_business_complete(state):
        status = "idle_incomplete"
    if status == "idle_incomplete" and state and \
            state.get("stage") == "TERMS_PREPARED" and state.get("quality_mode") \
            and state.get("glossary") is not None and not state.get("glossary_frozen") \
            and not state.get("quality_bypass"):
        status = "waiting_manual"
    labels = {
        "resume_requested": "正在恢复任务", "queued": "正在恢复任务",
        "starting": "正在恢复任务", "running": "正在运行",
        "waiting_external": "正在运行",
        "stalled": "暂无运行信号", "interrupted": "上次运行已中断",
        "pausing": "正在暂停", "paused": "已暂停", "failed": "当前步骤失败", "cancelling": "正在取消", "cancelled": "任务已取消",
        "completed": "已完成", "idle_incomplete": "未完成",
        "waiting_manual": "待术语确认",
    }
    events = read_runtime_events(job_id, 5, visibility="user") \
        or runtime.get("events") or []
    completed, total = int(runtime.get("completed_units") or 0), \
        int(runtime.get("total_units") or 0)
    academic_present = bool(state and ((state.get("academic_state") or {}).get("artifacts")
                                       or (state.get("academic_state") or {}).get(
                                           "current_stage") not in {None, "", "not_started"}))
    if not runtime.get("segment_operation") and state and state.get("p2_done") and (
            state.get("report_enabled", True) or academic_present):
        completed, total = _academic_runtime_progress(job_id, state)
    if status == "completed" and total:
        completed = total
    progress = (completed, total) if total else None
    operation = runtime.get("operation_label") or ""
    is_report = bool(state and state.get("p2_done") and (
        state.get("report_enabled", True) or academic_present))
    surface_label = "实践报告" if is_report else "任务处理"
    if is_report and (not operation or operation in {"准备工作流", "准备任务", "继续处理"}):
        operation = _academic_resume_context(state, runtime)["operation_label"]
    if status in {"resume_requested", "queued", "starting"}:
        detail = "正在读取最近检查点…"
        if not operation or operation in {"准备工作流", "准备任务"}:
            operation = "正在从最近检查点继续学术写作" if is_report \
                else "正在从最近进度继续处理"
    elif status == "waiting_external":
        detail = "正在等待模型响应"
    elif status == "running":
        detail = runtime.get("phase_label") or "已恢复 · 正在执行"
    elif status == "interrupted":
        detail = "当前进度已安全保存，可以继续处理。"
    elif status == "stalled":
        detail = "暂时没有新的运行信号，可以重试当前步骤。"
    elif status == "idle_incomplete":
        detail = "当前进度已保存，可以继续处理。"
        if state and workflow_outcomes(state)["annotation"]["status"] == "partial":
            detail = "正文已完成；标注部分失败，可在审校工作台补跑。"
    elif status == "failed":
        error = runtime.get("error") or {}
        detail = error.get("message") if isinstance(error, dict) else str(error)
    else:
        detail = runtime.get("phase_label") or ""
    primary_action = primary_runtime_action_for_status(status)
    return {
        "status": status,
        "workflow_outcomes": workflow_outcomes(state or {}),
        "status_label": labels.get(status, status),
        "runtime_status": status,
        "headline_status": labels.get(status, status),
        "badge": status,
        "surface_label": surface_label,
        "headline": operation or surface_label,
        "detail": detail,
        "current_operation": operation,
        "operation_id": runtime.get("operation_id") or runtime.get("operation") or "",
        "progress": progress,
        "progress_completed": completed,
        "progress_total": total,
        "stage_progress": runtime.get("stage_progress") or {},
        "available_actions": runtime_actions_for_status(status),
        "primary_action": primary_action,
        "show_no_worker_warning": status in {
            "idle_incomplete", "interrupted", "cancelled"},
        "last_activity": events[-1].get("message") if events else "",
        "last_activity_at": events[-1].get("timestamp") if events else None,
        "user_events": events,
        "events": events,
        "runtime": runtime,
    }


def _append_runtime_technical_log(job_id, text):
    """向技术日志追加一行。

    进程被直接杀掉时不会有机会写任何东西，所以关键节点必须**主动**留痕：
    worker 启动/释放都记一条，这样"有启动、没有释放"本身就说明了中断方式。
    """
    try:
        path = runtime_technical_log_path(job_id)
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("a", encoding="utf-8") as stream:
            stream.write(f"[{_utc_now_iso()}] {text}\n")
            stream.flush()
    except OSError:
        pass  # 日志失败不得影响任务本身


def _write_runtime_technical_log(job_id, exc):
    path = runtime_technical_log_path(job_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as stream:
        stream.write(f"[{_utc_now_iso()}] {type(exc).__name__}: {exc}\n")
        stream.write(traceback.format_exc())
        stream.write("\n")
        stream.flush()


def _runtime_stage_info(label):
    """Extract the human-facing stage contract from existing status labels."""
    text = str(label or "").strip()

    def nested_progress(value):
        """Parse durable item progress without inventing an ETA."""
        batch_match = re.search(r"批次\s*(\d+)\s*/\s*(\d+)", value)
        page_match = re.search(r"(\d+)\s*/\s*(\d+)\s*页", value)
        sentence_match = re.search(r"(\d+)\s*/\s*(\d+)\s*句", value)
        match = batch_match or page_match or sentence_match
        if not match:
            return {}
        unit = "批次" if batch_match else ("页" if page_match else "句")
        completed, total = int(match.group(1)), int(match.group(2))
        progress = {
            "unit": unit,
            "completed": max(0, completed),
            "total": max(0, total),
            "in_flight": [],
        }
        active_match = re.search(r"处理中\s+((?:#\d+\s*)+)", value)
        if active_match:
            progress["in_flight"] = [
                int(item[1:]) for item in active_match.group(1).split()
                if item.startswith("#") and item[1:].isdigit()
            ]
        return progress

    def operation_without_progress(value):
        operation = re.sub(r"^.*?】\s*", "", value).strip(" .…") or value
        operation = re.sub(
            r"\s*[（(](?=[^）)]*(?:批次|页|句)[^）)]*[）)])[^）)]*[）)]?", "",
            operation,
        )
        # The expression above deliberately handles the normal full-width
        # parenthesis form; keep a small fallback for labels using ASCII
        # parentheses so an individual provider cannot pollute the headline.
        operation = re.sub(r"\s*\([^()]*?(?:批次|页|句)[^()]*\)", "", operation)
        return operation.strip(" .…") or value

    stage_progress = nested_progress(text)
    operation_label = operation_without_progress(text)
    if "扫描 PDF OCR" in text or "OCR" in text and "页" in text:
        return {
            "stage": "document_processing",
            "stage_id": "ocr",
            "operation": "ocr",
            "operation_id": "ocr",
            "operation_label": operation_label,
            "stage_progress": stage_progress,
        }
    if "LLM 原文纠错" in text or "原文纠错与断行" in text:
        return {
            "stage": "document_processing",
            "stage_id": "llm_cleanup",
            "operation": "llm_cleanup",
            "operation_id": "llm_cleanup",
            "operation_label": operation_label,
            "stage_progress": stage_progress,
        }
    if "按句构建" in text or "翻译单元" in text:
        return {
            "stage": "document_processing",
            "stage_id": "segment_build",
            "operation": "segment_build",
            "operation_id": "segment_build",
            "operation_label": operation_label,
            "stage_progress": stage_progress,
        }
    match = re.search(r"学术写作\s+(\d+)\s*/\s*(\d+)", text)
    if match:
        index, total = int(match.group(1)), int(match.group(2))
        section_match = re.search(r"第\s*([\w.-]+)\s*节", text)
        stage_id = "academic_writing"
        for needle, candidate in (
                ("证据库", "evidence"), ("研究问题", "research_model"),
                ("论点", "argument_plan"), ("案例", "selected_cases"),
                ("提纲", "outline"), ("撰写正文", "sections"),
                ("确定性", "validation"), ("审稿", "review"),
                ("修订", "quality_repair"), ("质量", "academic_quality")):
            if needle in text:
                stage_id = candidate
                break
        return {
            "stage": "academic_writing",
            "stage_id": stage_id,
            "stage_index": index,
            "stage_total": total,
            "operation": stage_id,
            "operation_id": "section_rewrite" if section_match else stage_id,
            "operation_label": operation_label,
            "section_id": section_match.group(1) if section_match else None,
            "stage_progress": stage_progress,
        }
    if "文档" in text or "排版" in text:
        stage = "document_processing"
    elif "翻译" in text:
        stage = "translation"
    elif "术语" in text:
        stage = "terminology"
    elif "标注" in text:
        stage = "annotation"
    elif "报告" in text or "学术" in text:
        stage = "academic_writing"
    else:
        stage = "pipeline"
    return {
        "stage": stage,
        "stage_id": stage,
        "operation": stage,
        "operation_id": stage,
        "operation_label": operation_label,
        "stage_progress": stage_progress,
    }


def _runtime_overall_progress(stage_info, state):
    """Return a conservative progress estimate; only completion may be 1.0."""
    if not isinstance(state, dict):
        return 0.0
    if state.get("p1_done") and state.get("p2_done") \
            and (state.get("p3_done") or not state.get("report_enabled", True)) \
            and not _annotation_blocks_completion(state):
        return 1.0
    return None


def _academic_runtime_progress(job_id, state):
    """Count only durable academic artifact checkpoints; no estimated percent."""
    if not isinstance(state, dict):
        return 0, 0
    academic = state.get("academic_state") or {}
    artifacts = academic.get("artifacts") or {}
    names = ("evidence", "research_model", "argument_plan", "selected_cases",
             "outline", "sections", "validation", "review",
             "literature_support_review", "academic_quality", "report")
    completed = sum(1 for name in names if (artifacts.get(name) or {}).get("file"))
    return completed, len(names)


_ACADEMIC_RESUME_META = {
    "evidence": (1, "evidence", "继续构建学术证据"),
    "literature_evidence": (2, "literature_evidence", "继续整理文献证据"),
    "research_model": (3, "research_model", "继续建立研究模型"),
    "literature_claims": (4, "literature_claims", "继续整理文献主张"),
    "argument_plan": (5, "argument_plan", "继续规划论点"),
    "case_analysis": (6, "outline", "继续规划案例分析"),
    "outline": (6, "outline", "继续生成学术提纲"),
    "writing": (7, "sections", "继续撰写报告章节"),
    "validation": (8, "validation", "继续验证报告"),
    "review": (9, "review", "继续执行学术复核"),
    "repair": (10, "quality_repair", "继续修订受影响章节"),
    "academic_quality": (11, "academic_quality", "继续评估学术质量"),
}


def _academic_resume_context(state, runtime=None):
    academic = (state or {}).get("academic_state") or {}
    current_stage = str(academic.get("current_stage") or "")
    index, operation, label = _ACADEMIC_RESUME_META.get(
        current_stage, (1, "academic_writing", "继续学术写作"))
    runtime = runtime or {}
    section_id = runtime.get("section_id")
    if current_stage == "repair" and section_id:
        operation = "section_rewrite"
        label = f"继续重新生成第 {section_id} 节"
    return {
        "pipeline": "academic_writing",
        "stage": "academic_writing",
        "stage_id": operation,
        "stage_index": index,
        "stage_total": 11,
        "operation": operation,
        "operation_id": operation,
        "operation_label": label,
        "stage_progress": {},
        "section_id": section_id or "",
    }


def _runtime_heartbeat_loop(job_id, stop_event):
    while not stop_event.wait(RUNTIME_HEARTBEAT_SECONDS):
        current = load_runtime_state(job_id)
        if current.get("status") not in RUNTIME_ACTIVE_STATUSES:
            return
        try:
            _workflow_errors.check(job_id)
            update_runtime_state(job_id, heartbeat=True)
        except _workflow_errors.WorkflowStopped:
            return


def _runtime_status_callback(job_id, state, label):
    _runtime_pause_checkpoint(job_id)
    stage_info = _runtime_stage_info(label)
    current = load_runtime_state(job_id)
    operation = stage_info.get("operation") or current.get("operation")
    operation_started_at = current.get("operation_started_at")
    if operation != current.get("operation") or not operation_started_at:
        operation_started_at = _utc_now_iso()
    saved_state = load_job_state(job_id)
    state = saved_state or (state if isinstance(state, dict) else {})
    changes = dict(stage_info)
    completed_units, total_units = (_academic_runtime_progress(job_id, state)
                                    if stage_info.get("stage") == "academic_writing" else (0, 0))
    changes.update({
        "operation": operation,
        "operation_id": stage_info.get("operation_id") or operation,
        "operation_label": stage_info.get("operation_label"),
        "stage_progress": stage_info.get("stage_progress") or {},
        "section_id": stage_info.get("section_id") or "",
        "operation_started_at": operation_started_at,
        "overall_progress": _runtime_overall_progress(stage_info, state),
        "status": "running",
        "phase": "running",
        "phase_label": "正在处理",
        "completed_units": completed_units,
        "total_units": total_units,
    })
    update_runtime_state(job_id, progress=True, event=label,
                         event_name=stage_info.get("stage_id") or operation,
                         event_visibility="user", event_category="progress",
                         **changes)
    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")


def _runtime_caption_callback(job_id, text):
    message = str(text or "").strip()
    if not message:
        return
    _runtime_pause_checkpoint(job_id)
    stage_info = _runtime_stage_info(message)
    changes = {"operation_label": message, "phase_label": message}
    if stage_info["stage"] != "pipeline":
        changes.update({key: stage_info[key] for key in (
            "stage", "stage_id", "operation", "operation_id", "operation_label")})
    update_runtime_state(job_id, progress=True, event=message, status="running",
                         event_visibility="user", event_category="progress", **changes)
    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")


def _run_job_worker(job_id, filename, file_bytes, pipeline_kwargs, base_url=None,
                    segment_operation=None):
    stop_event = threading.Event()
    heartbeat = threading.Thread(target=_runtime_heartbeat_loop,
                                 args=(job_id, stop_event), daemon=True)
    _RUNTIME_CTX.job_id = job_id
    set_llm_base_url(base_url)
    set_llm_reasoning_effort((pipeline_kwargs or {}).get("reasoning_effort"))
    heartbeat.start()
    # 主动留痕：进程被直接杀掉时没有机会写日志，所以"有启动、没有释放"本身就是
    # 最有价值的证据（本次排查正是靠它区分"被关闭"与"抛异常"）。
    _append_runtime_technical_log(
        job_id,
        f"worker started pid={os.getpid()} thread={threading.current_thread().name}"
        f" attempt={load_runtime_state(job_id).get('attempt')}"
        f" resume={bool(load_runtime_state(job_id).get('resume_request_id'))}")
    try:
        kwargs = dict(pipeline_kwargs or {})
        kwargs.pop("on_status", None)
        kwargs.pop("on_caption", None)
        state = load_job_state(job_id) or new_job_state(filename)
        queued = load_runtime_state(job_id)
        update_runtime_state(
            job_id, status="starting", phase="starting",
            phase_label="正在读取最近检查点", heartbeat=True, progress=True,
            event="后台 worker 已接管任务", event_name="worker_started",
            event_visibility="technical", event_category="orchestration",
            pipeline=("segment_operation" if segment_operation else
                      "academic_writing" if kwargs.get("enable_report")
                      else "document_pipeline"))
        update_runtime_state(
            job_id, status="running", phase="running",
            phase_label="已恢复 · 正在执行" if queued.get("resume_request_id")
            else "正在执行", heartbeat=True)
        if segment_operation:
            kind = segment_operation["kind"]
            indexes = segment_operation["indexes"]
            kwargs.pop("reasoning_effort", None)
            if kind == "review":
                result = review_translation_segments(job_id, indexes, **kwargs)
                failed = result[1].get("failed_segment_ids") or []
                if failed:
                    raise RuntimeError(f"{len(failed)} 个段落审校未完成，请在审校队列中重试。")
            else:
                result = retranslate_segments(
                    job_id, indexes, **kwargs,
                    on_caption=lambda text: _runtime_caption_callback(job_id, text))
                if len(result[1]) != len(indexes):
                    raise RuntimeError("部分段落未完成重译，请查看段落状态后重试。")
                readiness = _translation_evidence.translation_review_readiness(result[0])
                if set(indexes) & set(readiness.get("failed_segment_ids") or []):
                    raise RuntimeError("重译已保存，但部分段落审校未完成，请在审校队列中重试。")
            if not _runtime_cancel_requested(job_id):
                update_runtime_state(
                    job_id, segment_operation_result="completed",
                    completed_units=len(indexes), progress=True,
                    event="所选段落审校已完成" if kind == "review" else "所选段落重译已完成",
                    event_name="segment_operation_completed", event_visibility="user",
                    event_category="progress")
        else:
            result = run_job_pipeline(
                job_id, filename, file_bytes, **kwargs,
                on_status=lambda label: _runtime_status_callback(job_id, state, label),
                on_caption=lambda text: _runtime_caption_callback(job_id, text),
            )
        result_state = load_job_state(job_id) or state
        if _runtime_cancel_requested(job_id):
            if not segment_operation:
                trace.finish_run(job_id=job_id, status="cancelled")
            update_runtime_state(job_id, status="cancelled", phase="cancelled",
                                 phase_label="已取消", event="任务已取消",
                                 event_name="job_cancelled", progress=True,
                                 event_visibility="user", event_category="lifecycle",
                                 worker={"owner_pid": None, "worker_id": None,
                                         "lease_expires_at": None})
        elif _runtime_business_complete(result_state):
            if not segment_operation:
                trace.finish_run(job_id=job_id, status="completed")
            update_runtime_state(job_id, status="completed", phase="completed",
                                 phase_label="已完成", event="任务已完成",
                                 event_name="job_completed", progress=True,
                                 event_visibility="user", event_category="lifecycle",
                                 heartbeat=True, overall_progress=1.0,
                                 worker={"owner_pid": None, "worker_id": None,
                                         "lease_expires_at": None})
        else:
            if not segment_operation:
                trace.finish_run(job_id=job_id, status="checkpointed")
            update_runtime_state(job_id, status="idle_incomplete", phase="idle_incomplete",
                                 phase_label="等待继续", event="阶段已保存，等待继续",
                                 event_name="job_checkpointed", progress=True,
                                 event_visibility="technical", event_category="checkpoint",
                                 heartbeat=True,
                                 worker={"owner_pid": None, "worker_id": None,
                                         "lease_expires_at": None})
        return result
    # 捕获 BaseException 而不是 Exception：worker 只捕 Exception 时，若线程抛出
    # SystemExit 一类的 BaseException，`finally` 不会执行 → 心跳线程继续续租 →
    # lease 永不过期 → 应用还活着时这个任务会**真的硬卡死**（状态一直停在活跃态，
    # start_job_worker / resume_job 都会拒绝）。这里必须失败闭合。
    except BaseException as exc:  # noqa: BLE001 - worker must publish failure to UI
        cancelled = bool(load_runtime_state(job_id).get("cancel_requested")) or "请求取消" in str(exc)
        if not segment_operation:
            trace.finish_run(job_id=job_id, status="cancelled" if cancelled else "failed", failure_reason=str(exc))
        if not cancelled:
            try:
                _write_runtime_technical_log(job_id, exc)
            except OSError:
                pass
        current = load_runtime_state(job_id)
        provider_status = provider_error_status(exc)
        user_error_message = (
            provider_error_message(exc, "任务执行失败")
            if provider_status["status"] != "unknown"
            else str(exc)[:500] or "任务执行失败")
        error = None if cancelled else {
            "type": type(exc).__name__,
            "message": user_error_message,
            "stage": current.get("stage_id") or current.get("stage"),
            "operation": current.get("operation_id") or current.get("operation"),
            "timestamp": _utc_now_iso(),
            "technical_log": RUNTIME_TECHNICAL_LOG,
        }
        try:
            update_runtime_state(
                job_id, status="cancelled" if cancelled else "failed",
                phase="cancelled" if cancelled else "failed",
                phase_label="已取消" if cancelled else "步骤失败",
                error=error,
                event="任务已取消" if cancelled else f"步骤失败：{user_error_message[:180]}",
                event_name="job_cancelled" if cancelled else "job_failed",
                event_visibility="user",
                event_category="lifecycle" if cancelled else "error",
                progress=True, heartbeat=True,
                worker={"owner_pid": None, "worker_id": None,
                        "lease_expires_at": None})
        except _workflow_errors.WorkflowStopped as stopped:
            if stopped.reason != "disk_full":
                raise
        return None
    finally:
        _append_runtime_technical_log(
            job_id, f"worker released pid={os.getpid()}"
                    f" thread={threading.current_thread().name}")
        stop_event.set()
        heartbeat.join(timeout=RUNTIME_HEARTBEAT_SECONDS + 1)
        set_llm_base_url(None)
        # Terminal boundary: flush anything still inside the telemetry window,
        # then drop the coalescer so a resumed run starts from a clean slate.
        # The pooled transport client for this worker thread is released here
        # too, so a long-lived process does not accumulate idle sockets.
        try:
            flush_runtime_telemetry(job_id)
            reset_runtime_telemetry(job_id)
            _llm_transport.close_thread_clients()
            try:
                update_runtime_state(
                    job_id, event="后台运行已释放", event_name="worker_released",
                    event_visibility="technical", event_category="orchestration")
            except _workflow_errors.WorkflowStopped as stopped:
                if stopped.reason != "disk_full":
                    raise
        finally:
            _RUNTIME_CTX.__dict__.clear()
            with _RUNTIME_WORKERS_LOCK:
                if _RUNTIME_WORKERS.get(job_id) is threading.current_thread():
                    _RUNTIME_WORKERS.pop(job_id, None)
                execution_lock = getattr(threading.current_thread(), "execution_lock", None)
                if execution_lock is not None:
                    execution_lock.close()


def start_job_worker(job_id, filename, file_bytes, pipeline_kwargs, base_url=None,
                     resume_request_id=None, *, segment_operation=None):
    """Start one resumable pipeline worker; repeated UI reruns are idempotent."""
    with _RUNTIME_WORKERS_LOCK:
        if _job_execution_alive(job_id):
            return False
        execution_lock = _job_lock.acquire(job_dir(job_id) / ".worker.lock", blocking=False)
        if execution_lock is None:
            return False
        try:
            _workflow_errors.reset(job_id)
            if file_bytes is not None:
                save_source(job_id, file_bytes)
            runtime = load_runtime_state(job_id)
            status = runtime.get("status") or "idle"
            matching_resume = bool(resume_request_id) and (
                status not in RUNTIME_ACTIVE_STATUSES or (
                    status == "resume_requested" and runtime.get("resume_request_id") == resume_request_id))
            if status in RUNTIME_ACTIVE_STATUSES and not matching_resume:
                return False
            if not matching_resume:
                runtime = get_job_runtime_status(job_id)
                if runtime.get("status") in RUNTIME_ACTIVE_STATUSES:
                    return False
            state = load_job_state(job_id) or new_job_state(filename)
            worker_id = uuid.uuid4().hex
            attempt = int(runtime.get("attempt") or 0) + 1
            started_at = _utc_now_iso()
            lease_expires_at = (datetime.now(timezone.utc) +
                                timedelta(seconds=RUNTIME_LEASE_SECONDS)).isoformat(
                                    timespec="seconds")
            completed_units, total_units = (0, 0)
            if state.get("p2_done") and pipeline_kwargs.get("enable_report"):
                completed_units, total_units = _academic_runtime_progress(job_id, state)
            context = _academic_resume_context(state, runtime) \
                if matching_resume and pipeline_kwargs.get("enable_report") else {
                    "pipeline": "document_pipeline", "stage": "pipeline",
                    "stage_id": "pipeline", "operation": "pipeline",
                    "operation_id": "pipeline", "operation_label": "准备任务",
                    "stage_progress": {},
                    "section_id": "", "stage_index": None, "stage_total": None,
                }
            if segment_operation:
                kind = segment_operation["kind"]
                context.update(
                    pipeline="segment_operation", stage="translation",
                    stage_id="segment_review" if kind == "review" else "segment_retranslation",
                    operation=kind, operation_id=kind,
                    operation_label="所选段落审校" if kind == "review" else "所选段落重译")
                completed_units, total_units = 0, len(segment_operation["indexes"])
            if matching_resume and status != "resume_requested":
                runtime = update_runtime_state(
                    job_id, status="resume_requested", phase="resume_requested",
                    phase_label="正在恢复任务", resume_request_id=resume_request_id,
                    started_at=started_at, operation_started_at=started_at,
                    last_progress_at=started_at, completed_units=completed_units,
                    total_units=total_units, overall_progress=None,
                    worker={"owner_pid": os.getpid(), "worker_id": worker_id,
                            "lease_expires_at": lease_expires_at},
                    event="已从断点恢复任务", event_name="resume_requested",
                    event_visibility="user", event_category="lifecycle",
                    event_metadata={"resume_request_id": resume_request_id}, **context)
            update_runtime_state(
                job_id, status="queued", phase="starting", phase_label="准备中",
                **context, started_at=runtime.get("started_at") if matching_resume
                else started_at,
                operation_started_at=started_at, last_progress_at=started_at,
                last_heartbeat_at=started_at, cancel_requested=False, pause_requested=False, error=None,
                attempt=attempt, resume_request_id=resume_request_id or "",
                segment_operation=segment_operation or {}, segment_operation_result="pending",
                completed_units=completed_units, total_units=total_units,
                overall_progress=None,
                worker={"owner_pid": os.getpid(), "worker_id": worker_id,
                        "lease_expires_at": lease_expires_at},
                event="已排入后台 worker", event_name="job_queued",
                event_visibility="technical", event_category="orchestration")
            worker = threading.Thread(
                target=_run_job_worker,
                args=(job_id, filename, file_bytes, dict(pipeline_kwargs or {}), base_url,
                      segment_operation),
                name=f"folith-{job_id}", daemon=True)
            worker.worker_id = worker_id
            worker.execution_lock = execution_lock
            _RUNTIME_WORKERS[job_id] = worker
            worker.start()
            execution_lock = None  # owned by the worker until its final cleanup
            return True
        finally:
            if execution_lock is not None:
                execution_lock.close()


def start_segment_worker(job_id, kind, indexes, *, base_url=None, **kwargs):
    """Run a selected review or retranslation on the existing job worker."""
    if kind not in {"review", "retranslate"}:
        raise ValueError(f"未知段落操作：{kind}")
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    indexes = sorted(set(int(index) for index in indexes))
    if not indexes or any(index < 0 or index >= len(state.get("pairs") or [])
                          for index in indexes):
        raise ValueError("请选择仍然存在的翻译段落。")
    if kind == "review":
        kwargs["base_url"] = base_url or ""
    return start_job_worker(
        job_id, state.get("filename") or "", None, kwargs, base_url=base_url,
        segment_operation={"kind": kind, "indexes": indexes})


def resume_job(job_id, filename, pipeline_kwargs, base_url=None,
               resume_request_id=None):
    """Idempotently request resume and publish the transition before worker start."""
    with _RUNTIME_WORKERS_LOCK:
        runtime = get_job_runtime_status(job_id)
        if runtime.get("status") in {"paused", "pausing"} and is_job_worker_alive(job_id):
            return continue_job(job_id)
        if _job_execution_alive(job_id):
            return False
        if runtime.get("status") in RUNTIME_ACTIVE_STATUSES:
            return False
        state = load_job_state(job_id) or new_job_state(filename)
        request_id = resume_request_id or uuid.uuid4().hex
        if runtime.get("resume_request_id") == request_id:
            return False
        pipeline_kwargs = dict(pipeline_kwargs or {})
        academic = state.get("academic_state") or {}
        if state.get("p2_done") and (academic.get("artifacts") or academic.get(
                "current_stage") not in {None, "", "not_started"}):
            pipeline_kwargs["enable_report"] = True
        return start_job_worker(
            job_id, filename, None, pipeline_kwargs, base_url=base_url,
            resume_request_id=request_id)


def is_job_worker_alive(job_id):
    if not job_id:
        return False
    with _RUNTIME_WORKERS_LOCK:
        worker = _RUNTIME_WORKERS.get(job_id)
        if worker and worker.is_alive():
            return True
        # If in-memory registry lost the reference (e.g. Streamlit reloaded the module),
        # look for any active worker thread belonging to this job in the current process.
        target_name = f"folith-{job_id}"
        for thread in threading.enumerate():
            if thread.name == target_name and thread.is_alive():
                _RUNTIME_WORKERS[job_id] = thread
                return True
        return False


def request_job_cancel(job_id, *, force=False):
    """Request cancellation between provider calls; an active HTTP call finishes first.

    If no active worker is running, or if force=True, marks the runtime state as cancelled immediately.
    """
    if not job_id:
        return False
    alive = is_job_worker_alive(job_id)
    if alive or _job_execution_alive(job_id):
        update_runtime_state(job_id, status="cancelling", cancel_requested=True,
                             phase="cancelling", phase_label="正在取消",
                             event="已请求取消任务", event_name="cancel_requested",
                             progress=True)
        return True

    # If no worker is alive or force cancel was explicitly requested:
    # Immediately transition the job state out of active running status.
    update_runtime_state(
        job_id, status="cancelled", phase="cancelled", phase_label="已取消",
        cancel_requested=False, error=None,
        event="任务已取消", event_name="job_cancelled",
        event_visibility="user", event_category="lifecycle",
        progress=True, heartbeat=True,
        worker={"owner_pid": None, "worker_id": None, "lease_expires_at": None})
    return True


def retry_job_step(job_id):
    """Invalidate only the failed academic operation and its downstream work."""
    if _job_execution_alive(job_id):
        return False
    runtime = get_job_runtime_status(job_id)
    if runtime.get("status") not in {"failed", "stalled", "interrupted", "cancelled"}:
        return False
    operation = runtime.get("operation_id") or runtime.get("operation") or ""
    scopes = {
        "validation": "validation", "review": "review",
        "literature_support_review": "literature_review",
        "academic_quality": "quality", "quality_repair": "quality",
        "section_rewrite": "section", "sections": "writer", "repair": "writer",
        "outline": "planning",
        "argument_plan": "planning", "selected_cases": "planning",
        "research_model": "planning", "evidence": "all",
    }
    state = load_job_state(job_id)
    if state and (runtime.get("stage") == "academic_writing" or
                  operation in scopes):
        error_message = str((runtime.get("error") or {}).get("message") or "")
        scope = "case_analysis" if operation == "section_rewrite" and \
            "missing case target subsection" in error_message else \
            scopes.get(operation, "writer")
        invalidate_academic_report(job_id, scope, runtime.get("section_id"))
    update_runtime_state(
        job_id, status="idle_incomplete", phase="retry_ready", phase_label="等待重试",
        cancel_requested=False, error=None, event="已准备重试当前步骤",
        event_name="retry_requested", worker={"owner_pid": None, "worker_id": None,
                                               "lease_expires_at": None})
    return True



def compliance_profile_view(job_id, state=None):
    """Evaluate source-backed compliance and explicit project constraints."""
    state = state if state is not None else load_job_state(job_id) or {}
    from folith import compliance
    artifacts = {
        name: load_academic_artifact(job_id, name)
        for name in ("evidence", "report", "validation", "outline",
                     "selected_cases", "literature_sources",
                     "final_docx_validation")
    }
    profile_id = str(state.get("compliance_profile_id") or
                     compliance.DEFAULT_PROFILE_ID)
    profile = compliance.compliance_profile(profile_id)
    result = compliance.evaluate_compliance(
        state, artifacts, profile, state.get("p3_md") or "")
    language = compliance.evaluate_language_constraints(
        state, state.get("p3_md") or "")
    result["language_constraints"] = language
    language_constraints = language.get("constraints") or []
    result["counts"]["pass"] += sum(
        item.get("status") == "pass" for item in language_constraints)
    result["counts"]["fail"] += len(language.get("failures") or [])
    result["counts"]["manual_review"] += sum(
        item.get("status") == "manual_review" for item in language_constraints)
    project = result.setdefault("project_constraints", {})
    project.setdefault("failures", []).extend(
        f"language:{item.get('kind')}:{item.get('value')}"
        for item in language.get("failures") or [])
    if language.get("failures"):
        project["status"] = "fail"
        result["status"] = "fail"
    if language.get("status") == "manual_review" and \
            result.get("status") != "fail":
        result["status"] = "manual_review"
    return result



@_serialized_job_mutation
def generate_report_qa(job_id, state=None, *, save_file=False):
    """Build the concise QA report and bind it to current artifact hashes."""
    from folith import academic_evidence, academic_writer
    state = state or load_job_state(job_id) or {}
    report_record = load_academic_artifact(job_id, "report") or {}
    render_record = load_academic_artifact(job_id, "libreoffice_render") or {}
    final_docx = load_academic_artifact(job_id, "final_docx_validation") or {}
    compliance = compliance_profile_view(job_id, state)
    case_review = _finalization.case_review_gate(
        state, load_academic_artifact(job_id, "selected_cases"))
    final_qa = _finalization.normalize_final_qa(state.get("final_qa"))
    final_qa["structural_qa"] = "PASS" if final_docx.get("status") in {
        "pass", "pass_with_warnings"} else "FAIL" if final_docx.get(
        "status") == "fail" else "NOT_RUN"
    translation_hash = current_translation_hash(state)
    report_hash = str(report_record.get("content_hash") or "")
    docx_hash = str(render_record.get("source_docx_hash") or
                    final_docx.get("source_docx_hash") or "")
    markdown = _rendered_qa.render_qa_markdown(
        translation_hash=translation_hash, report_hash=report_hash,
        docx_hash=docx_hash, render_record=render_record,
        pdf_qa=render_record.get("analysis") or {}, compliance=compliance,
        case_review=case_review, final_qa=final_qa,
        placeholders=next((x.get("actual") or [] for x in compliance.get(
            "rules") or [] if x.get("rule_id") == "author_placeholders"), []))
    value = {
        "schema_version": _rendered_qa.VERSION,
        "generated_at": _finalization.now_iso(),
        "translation_truth_hash": translation_hash,
        "report_content_hash": report_hash,
        "source_docx_hash": docx_hash,
        "rendered_pdf_hash": render_record.get("rendered_pdf_hash"),
        "final_qa": final_qa,
        "compliance_status": compliance.get("status"),
        "case_review_status": case_review.get("status"),
        "content_hash": academic_evidence.stable_hash({
            "translation": translation_hash, "report": report_hash,
            "docx": docx_hash, "pdf": render_record.get("rendered_pdf_hash"),
            "compliance": compliance.get("status"),
            "case_review": case_review.get("status"), "final_qa": final_qa,
        }),
        "markdown": markdown,
    }
    if save_file:
        academic_writer._save_artifact(
            state, job_dir(job_id), "report_qa", value, str(value["content_hash"]),
            _rendered_qa.VERSION,
            input_artifact_ids=["report", "final_docx_validation",
                                "libreoffice_render"],
            input_segment_ids=[])
        (job_dir(job_id) / "report-qa.md").write_text(markdown, encoding="utf-8")
        save_job_state(job_id, state)
    return value


@_serialized_job_mutation
def save_compliance_record(job_id, state=None):
    """Persist the current compliance and language-constraint artifacts."""
    from folith import academic_evidence, academic_writer, compliance
    state = state or load_job_state(job_id) or {}
    result = compliance_profile_view(job_id, state)
    academic = state.get("academic_state") or {}
    records = academic.get("artifacts") or {}
    compliance_dependency = academic_evidence.stable_hash({
        "profile_id": result.get("profile_id"),
        "report": (records.get("report") or {}).get("content_hash"),
        "evidence": (records.get("evidence") or {}).get("content_hash"),
        "selected_cases": (records.get("selected_cases") or {}).get("content_hash"),
        "literature_sources": (records.get("literature_sources") or {}).get(
            "content_hash"),
        "outline": (records.get("outline") or {}).get("content_hash"),
    })
    academic_writer._save_artifact(
        state, job_dir(job_id), "compliance", result,
        compliance_dependency, compliance.compliance_profile(
            str(state.get("compliance_profile_id") or compliance.DEFAULT_PROFILE_ID)
        ).get("schema_version"),
        input_artifact_ids=["report", "evidence", "selected_cases",
                            "literature_sources"])
    language = result.get("language_constraints") or {}
    settings = state.get("research_settings") or {}
    language_dependency = academic_evidence.stable_hash({
        "report": (records.get("report") or {}).get("content_hash"),
        "language_constraints": {
            key: settings.get(key) for key in (
                "forbidden_report_phrases", "allowed_theory_labels",
                "required_terminology", "protected_names",
                "protected_work_titles")
        },
    })
    academic_writer._save_artifact(
        state, job_dir(job_id), "language_constraints", language,
        language_dependency, compliance.VERSION,
        input_artifact_ids=["report"])
    state["compliance_record"] = result
    state["language_constraint_record"] = language
    save_job_state(job_id, state)
    return result


def _case_artifact_case(job_id, case_id):
    selected = load_academic_artifact(job_id, "selected_cases") or {}
    case = next((item for item in selected.get("cases") or []
                 if str(item.get("case_id")) == str(case_id)), None)
    return selected, case


def _mark_case_downstream_stale(job_id, state, case_id, reason, *, actor="user",
                                action="case_changed", root_stale=True):
    from folith import academic_writer
    root_id = f"case:{case_id}"
    before_root = dict((state.get("academic_state") or {}).get(
        "artifacts", {}).get(root_id) or {})
    propagated = academic_writer.propagate_artifact_staleness(
        state, input_artifact_ids=[f"case:{case_id}"])
    if not root_stale and before_root:
        academic = state.setdefault("academic_state", {})
        academic.setdefault("artifacts", {})[root_id] = before_root
        academic_writer._write_status_mirror(academic, root_id, before_root)
    enriched = dict(state)
    enriched["_finalization_artifacts"] = _finalization_artifacts(job_id)
    impact = _finalization.build_dependency_impact(
        enriched, job_id, [], reason, changed_case_ids=[str(case_id)])
    state["dependency_impact"] = impact
    stale_names = [item.get("id") for item in impact.get("affected") or []
                   if item.get("kind") == "artifact"]
    academic = state.setdefault("academic_state", {})
    # The old fallback is retained only for pre-Stage-2 records. Canonical
    # records are invalidated by their case direct edge and reverse traversal.
    if not propagated and stale_names:
        academic_writer._invalidate_names(state, stale_names, reason)
    if propagated and any(name in propagated for name in {"sections", "report"}):
        state["p3_done"] = False
        academic["status"] = "stale"
    _reset_final_qa(state, reason)
    _invalidate_final_delivery_state(state)
    state.setdefault("human_actions", []).append({
        "finding_id": f"case:{case_id}", "action": action, "note": reason,
        "timestamp": _finalization.now_iso(), "actor": actor,
    })
    return state


@_serialized_job_mutation
def review_academic_case(job_id, case_id, status, note="", actor="user"):
    """Record one author decision without changing provenance or content."""
    state = load_job_state(job_id)
    if state is None:
        return None, False, "任务不存在"
    selected, case = _case_artifact_case(job_id, case_id)
    if case is None:
        return state, False, "找不到案例"
    from folith import case_provenance
    normalized = case_provenance.with_provenance(case)
    review_status = str(status or "").strip().lower()
    if review_status not in case_provenance.REVIEW_STATUSES:
        return state, False, "案例审校状态无效"
    state.setdefault("case_reviews", {})[str(case_id)] = {
        "review_status": review_status,
        "case_origin": normalized.get("case_origin"),
        "text_role": dict(normalized.get("text_role") or {}),
        "review_reason": str(note or "")[:700],
        "note": str(note or "")[:700],
        "reviewed_at": _finalization.now_iso(),
        "updated_at": _finalization.now_iso(),
        "actor": actor,
        "translation_truth_version": int(
            (state.get("translation_truth") or {}).get("version") or 0),
        "content_stale": False,
        "stale_reason": None,
        "stale_at": None,
    }
    state.setdefault("human_actions", []).append({
        "finding_id": f"case:{case_id}",
        "action": "case_approved" if review_status == "approved"
        else "case_excluded",
        "note": "批准案例纳入学术分析" if review_status == "approved"
        else f"排除案例：{str(note or '人工排除')[:180]}",
        "timestamp": _finalization.now_iso(),
        "actor": actor,
    })
    _mark_case_downstream_stale(
        job_id, state, str(case_id),
        "作者更新案例审核状态，需重组案例相关写作下游",
        actor=actor, action="case_review_changed", root_stale=False)
    save_job_state(job_id, state)
    return state, True, "已保存案例审校状态"


@_serialized_job_mutation
def replace_rejected_case(job_id, case_id, *, actor="user"):
    """Replace one author-rejected case from the already validated pool."""
    from folith import academic_writer, academic_quality, case_provenance
    state = load_job_state(job_id)
    if state is None:
        return None, False, ["任务不存在"]
    selected, old_case = _case_artifact_case(job_id, case_id)
    if selected is None or old_case is None:
        return state, False, ["找不到案例"]
    review = (state.get("case_reviews") or {}).get(str(case_id)) or {}
    if str(review.get("review_status")) != "rejected":
        return state, False, ["只能替换已被作者排除的案例"]
    selected_ids = {str(x.get("case_id")) for x in selected.get("cases") or []}
    reviews = state.get("case_reviews") or {}
    overrides = state.get("case_review_overrides") or {}
    candidates = []
    if case_provenance.is_synthetic(old_case):
        synthetic = load_academic_artifact(job_id, "synthetic_validation") or {}
        for candidate in synthetic.get("items") or []:
            cid = str(candidate.get("case_id") or "")
            validation = candidate.get("validation") or {}
            review_record = reviews.get(cid) if isinstance(reviews, dict) else None
            override = overrides.get(cid) if isinstance(overrides, dict) else None
            if not cid or cid in selected_ids or candidate is old_case:
                continue
            if review_record and review_record.get("review_status") == "rejected":
                continue
            if override and override.get("baseline_status") == "rejected":
                continue
            if not validation.get("academic_case_eligible"):
                continue
            difficulty = candidate.get("difficulty") or {}
            evidence = candidate.get("synthetic_evidence") or {}
            old_group = str(old_case.get("difficulty_group") or "")
            target_match = bool(old_group and str(difficulty.get("group") or "") == old_group)
            candidates.append((not target_match,
                               difficulty.get("academic_value") != "high",
                               difficulty.get("confidence") != "high",
                               evidence.get("material_difference") != "pass",
                               candidate.get("segment_index", 0), candidate))
    else:
        evidence_artifact = load_academic_artifact(job_id, "evidence") or {}
        argument_plan = load_academic_artifact(job_id, "argument_plan") or {}
        candidate = academic_quality.select_replacement_case(
            str(case_id), list(old_case.get("supports_claims") or []),
            selected, argument_plan, evidence_artifact)
        if candidate:
            candidates.append((False, False, False, False,
                               candidate.get("segment_index", 0), candidate))
    candidates.sort(key=lambda item: item[:-1])
    if not candidates:
        return state, False, ["没有可用的已验证替换候选；请保持该案例排除并调整案例数量"]
    candidate = case_provenance.with_provenance(dict(candidates[0][-1]))
    candidate.update({
        "supports_claims": sorted(set(old_case.get("supports_claims") or [])),
        "research_questions": sorted(set(old_case.get("research_questions") or [])),
        "argument_role": old_case.get("argument_role", "supporting"),
        "difficulty_group": old_case.get("difficulty_group") or
        candidate.get("difficulty_group"),
        "difficulty_subsection": old_case.get("difficulty_subsection"),
        "strategy_subsection": old_case.get("strategy_subsection"),
        "target_subsection": old_case.get("target_subsection"),
        "review_status": "unreviewed",
        "replacement_of": str(case_id),
        "selection_rationale": f"replacement of rejected {case_id}: existing validated pool",
    })
    selected["cases"] = [
        candidate if str(x.get("case_id")) == str(case_id)
        else x for x in selected.get("cases") or []
    ]
    record = academic_writer.artifact_record(state, "selected_cases")
    academic_writer._save_artifact(
        state, job_dir(job_id), "selected_cases", selected,
        str(record.get("dependency_hash") or ""),
        str(record.get("version") or "case-review-v1"))
    _mark_case_downstream_stale(
        job_id, state, str(candidate.get("case_id")),
        "作者排除案例后从已验证候选池替换，需重组其写作下游",
        actor=actor, action="case_replaced", root_stale=False)
    state.setdefault("case_reviews", {})[str(candidate.get("case_id"))] = {
        "review_status": "unreviewed",
        "case_origin": candidate.get("case_origin"),
        "text_role": dict(candidate.get("text_role") or {}),
        "review_reason": "替换被排除案例；需作者重新终审",
        "reviewed_at": _finalization.now_iso(),
        "updated_at": _finalization.now_iso(),
        "actor": actor,
        "translation_truth_version": int(
            (state.get("translation_truth") or {}).get("version") or 0),
        "content_stale": False,
    }
    state.setdefault("human_actions", []).append({
        "finding_id": f"case:{candidate.get('case_id')}",
        "action": "case_replaced",
        "note": f"从已验证候选池替换 {case_id}",
        "timestamp": _finalization.now_iso(),
        "actor": actor,
    })
    save_job_state(job_id, state)
    return state, True, [str(candidate.get("case_id"))]


@_serialized_job_mutation
def update_synthetic_baseline(job_id, case_id, text, *, status="modified", note="", actor="user"):
    """Modify or reject a synthetic baseline; never mutate translation truth."""
    state = load_job_state(job_id)
    if state is None:
        return None, False, "任务不存在"
    _selected, case = _case_artifact_case(job_id, case_id)
    if case is None:
        return state, False, "找不到案例"
    from folith import case_provenance
    if not case_provenance.is_synthetic(case):
        return state, False, "只有合成对照案例可以修改或拒绝模拟初译"
    if status not in {"modified", "rejected", "approved"}:
        return state, False, "模拟初译状态无效"
    record = dict((state.setdefault("case_review_overrides", {}).get(str(case_id)) or {}))
    if status == "modified":
        text = str(text or "").strip()
        if not text:
            return state, False, "模拟初译不能为空"
        record["synthetic_baseline_text"] = text
    record.update({
        "baseline_status": status,
        "note": str(note or "")[:700],
        "updated_at": _finalization.now_iso(),
        "actor": actor,
    })
    state["case_review_overrides"][str(case_id)] = record
    reason = ("修改模拟初译，需重跑该案例下游"
              if status == "modified" else "拒绝模拟初译，需替换或重新确认该案例")
    _finalization.mark_case_reviews_stale(state, [str(case_id)], reason)
    _mark_case_downstream_stale(
        job_id, state, str(case_id), reason, actor=actor,
        action="synthetic_baseline_modified" if status == "modified"
        else "synthetic_baseline_rejected")
    save_job_state(job_id, state)
    return state, True, "已保存模拟初译决定"


@_serialized_job_mutation
def record_final_qa(job_id, field, status, note="", actor="user"):
    """Persist one of the four independent final-QA facts."""
    if field not in _finalization.QA_FIELDS:
        raise ValueError(f"未知最终 QA 项：{field}")
    state = load_job_state(job_id)
    if state is None:
        return None
    allowed = {"PASS", "FAIL", "NOT_RUN"} if field in {
        "structural_qa", "libreoffice_render"} else {"CONFIRMED", "NOT_CONFIRMED"}
    if status not in allowed:
        raise ValueError(f"{field} 状态无效：{status}")
    qa = _finalization.normalize_final_qa(state.get("final_qa"))
    qa[field] = status
    qa["translation_truth_version"] = int(
        (state.get("translation_truth") or {}).get("version") or 0)
    notes = dict(qa.get("notes") or {})
    if note:
        notes[field] = str(note)[:700]
    qa["notes"] = notes
    qa["updated_at"] = _finalization.now_iso()
    state["final_qa"] = qa
    save_job_state(job_id, state)
    return state


@_serialized_job_mutation
def run_libreoffice_render_qa(job_id, state=None):
    """Render through LibreOffice, then run separate deterministic PDF QA."""
    state = state or load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    docx = report_docx_bytes(job_id, state)
    document_kind = "report"
    if docx is None:
        report_record = (state.get("academic_state") or {}).get(
            "artifacts", {}).get("report") or {}
        if state.get("report_enabled") and state.get("p3_done") and \
                report_record.get("status") in {"stale", "missing", "failed"}:
            raise RuntimeError("当前实践报告 artifact 已 stale，请先按影响范围重建报告")
        try:
            docx = build_delivery_assets(job_id, state).get("translation.docx")
            document_kind = "translation"
        except Exception as exc:
            raise RuntimeError(f"当前 DOCX 不可生成：{str(exc)[:180]}") from exc
    if not docx:
        raise RuntimeError("当前没有可渲染的 DOCX")
    previous_render = load_academic_artifact(job_id, "libreoffice_render") or {}
    current_qa = _finalization.normalize_final_qa(state.get("final_qa"))
    if (previous_render.get("rendered_pdf_hash") or
            previous_render.get("qa_status") == "PASS" or
            current_qa.get("author_visual_review") == "CONFIRMED" or
            current_qa.get("word_final_review") == "CONFIRMED"):
        _reset_final_qa(state, "LibreOffice render rerun; author and Word reviews reset")
        if current_qa.get("structural_qa") in {"PASS", "FAIL"}:
            state["final_qa"]["structural_qa"] = current_qa["structural_qa"]
    soffice = _find_soffice()
    if not soffice:
        qa = _finalization.normalize_final_qa(state.get("final_qa"))
        qa.update({"libreoffice_render": "NOT_RUN", "rendered_at": None,
                   "page_count": None, "source_docx_hash": _rendered_qa.sha256(docx),
                   "rendered_pdf_hash": None,
                   "updated_at": _finalization.now_iso(),
                   "translation_truth_version": int(
                       (state.get("translation_truth") or {}).get("version") or 0)})
        qa.setdefault("notes", {})["libreoffice_render"] = \
            "LibreOffice not installed; render NOT_RUN"
        state["final_qa"] = qa
        from folith import academic_writer
        academic_writer._save_artifact(
            state, job_dir(job_id), "libreoffice_render", {
                "schema_version": _rendered_qa.VERSION,
                "status": "not_run", "qa_status": "NOT_RUN",
                "render_engine": "libreoffice", "render_engine_version": None,
                "source_docx_hash": _rendered_qa.sha256(docx),
                "rendered_pdf_hash": None, "rendered_at": None, "page_count": None,
                "stale_reason": "LibreOffice unavailable; render was not run",
                "analysis": {"warnings": [], "manual_reviews": [{
                    "type": "libreoffice_unavailable", "severity": "manual_review"}]},
            }, "no-engine", _rendered_qa.VERSION,
            input_artifact_ids=["final_docx_validation"], status="missing")
        generate_report_qa(job_id, state, save_file=True)
        save_job_state(job_id, state)
        return state, qa
    from folith import academic_evidence, academic_writer
    version_result = subprocess.run([soffice, "--version"], capture_output=True,
                                    text=True, timeout=10, check=False)
    engine_version = (version_result.stdout or version_result.stderr or "").strip()
    with tempfile.TemporaryDirectory(prefix=f"folith-lo-{job_id}-") as tmp:
        tmp_path = Path(tmp)
        source_path = tmp_path / "current.docx"
        source_path.write_bytes(docx)
        profile = tmp_path / "profile"
        profile.mkdir()
        command = [soffice, "--headless", "-env:UserInstallation=file://"
                   + str(profile), "--convert-to", "pdf", "--outdir", tmp,
                   str(source_path)]
        result = subprocess.run(command, capture_output=True, text=True,
                                timeout=120, check=False)
        pdf_path = tmp_path / "current.pdf"
        if result.returncode != 0 or not pdf_path.is_file():
            detail = (result.stderr or result.stdout or "未知 LibreOffice 错误").strip()
            qa = _finalization.normalize_final_qa(state.get("final_qa"))
            qa.update({"libreoffice_render": "FAIL",
                       "rendered_at": _finalization.now_iso(),
                       "translation_truth_version": int(
                           (state.get("translation_truth") or {}).get("version") or 0),
                       "updated_at": _finalization.now_iso()})
            qa.setdefault("notes", {})["libreoffice_render"] = detail[:700]
            state["final_qa"] = qa
            render_value = {
                "status": "fail", "qa_status": "FAIL", "document_kind": document_kind,
                "detail": detail[:700],
                "render_engine": "libreoffice", "render_engine_version": engine_version,
                "source_docx_hash": _rendered_qa.sha256(docx),
                "rendered_pdf_hash": None,
                "rendered_at": _finalization.now_iso(), "page_count": None,
                "stale_reason": {"code": "render_failed",
                                 "source_type": "artifact",
                                 "source_id": "final_docx_validation"},
            }
            academic_writer._save_artifact(
                state, job_dir(job_id), "libreoffice_render", render_value,
                academic_evidence.stable_hash({
                    "final_docx_validation": (state.get("academic_state") or {}).get(
                        "artifacts", {}).get("final_docx_validation", {}).get(
                            "content_hash"),
                    "version": _finalization.VERSION,
                }), _finalization.VERSION,
                input_artifact_ids=["final_docx_validation"], status="failed",
                stale_reason={"code": "render_failed", "source_type": "artifact",
                              "source_id": "final_docx_validation"})
            generate_report_qa(job_id, state, save_file=True)
            save_job_state(job_id, state)
            raise RuntimeError(f"LibreOffice 渲染失败：{detail[:180]}")
        pdf_bytes = pdf_path.read_bytes()
    output = job_dir(job_id) / "libreoffice-render.pdf"
    output.write_bytes(pdf_bytes)
    analysis = _rendered_qa.analyze_pdf(pdf_bytes)
    page_count = int(analysis.get("page_count") or 0)
    page_metrics = analysis.get("pages") or []
    render_status = "PASS" if page_count and not analysis.get(
        "definite_failures") else "FAIL"
    qa = _finalization.normalize_final_qa(state.get("final_qa"))
    qa.update({
        "libreoffice_render": render_status,
        "rendered_at": _finalization.now_iso(),
        "page_count": page_count,
        "source_docx_hash": _rendered_qa.sha256(docx),
        "rendered_pdf_hash": _rendered_qa.sha256(pdf_bytes),
        "translation_truth_version": int(
            (state.get("translation_truth") or {}).get("version") or 0),
        "updated_at": _finalization.now_iso(),
    })
    qa.setdefault("notes", {})["document_kind"] = document_kind
    qa["page_metrics"] = page_metrics
    state["final_qa"] = qa
    render_value = _rendered_qa.build_render_record(
        document_kind=document_kind, source_docx=docx, rendered_pdf=pdf_bytes,
        engine="libreoffice", engine_version=engine_version, analysis=analysis)
    render_value["qa_status"] = render_status
    render_value["status"] = "pass" if render_status == "PASS" else "fail"
    academic_writer._save_artifact(
        state, job_dir(job_id), "libreoffice_render", render_value,
        academic_evidence.stable_hash({
            "final_docx_validation": (state.get("academic_state") or {}).get(
                "artifacts", {}).get("final_docx_validation", {}).get("content_hash"),
            "version": _finalization.VERSION,
        }), _finalization.VERSION,
        input_artifact_ids=["final_docx_validation"],
        status="valid" if qa["libreoffice_render"] == "PASS" else "failed")
    generate_report_qa(job_id, state, save_file=True)
    save_job_state(job_id, state)
    return state, qa



def new_job_state(filename):
    state = {
        "filename": filename,
        "p1_done": False,
        "p2_done": False,
        "p3_done": False,
        "report_enabled": False,
        "paras": [],
        "pairs": [],
        "auto_terms": {},
        "findings": [],
        "review_stats": {
            "reviewed_segments": 0, "batches_reviewed": 0,
            "batches_skipped_low_risk": 0, "skipped_segments": 0,
            "blocking": 0, "actionable": 0, "informational": 0, "review_failed": 0,
        },
        "tm_used_count": 0,
        "has_blocking": False,
        "p3_md": "",
        "p3_sections": [],
        "theory": "",
        "warnings": [],
        "annotations": {},
        "annotations_done": False,
        "annotations_done_offset": 0,
    }
    # 术语治理 / 交付门禁新增字段（默认值集中在 state_migration，保持单一来源）
    state.update(_state_migration._default_new_fields())
    return state


def load_job_state(job_id):
    """加载任务状态。使用文件签名缓存（mtime/size）消除重复反序列化，并返回安全深拷贝。"""
    return _job_repo.load_job_state_cached(job_id)


def save_job_state(job_id, state):
    """原子写入（先写临时文件再替换），避免中断写坏 state.json，并同步派生 summary.json。

    写完快照后是同一次 compaction：只要这份快照已经包含 journal 里的每一条记录，
    journal 就没有剩余职责。判断交给 ``rotate_if_applied``——它按 ``journal_cursor``
    复核，因此「把落后于 journal 的内存状态写盘」这种合法调用永远不会丢掉快照里
    没有的记录。
    """
    # External saves share the worker lifetime lock: a worker cannot start
    # between the ownership check and the atomic snapshot replacement.
    execution_lock = None
    if _runtime_job_id() != job_id:
        runtime = load_runtime_state(job_id)
        if runtime.get("segment_operation") and runtime.get("status") in RUNTIME_ACTIVE_STATUSES:
            raise RuntimeError("所选段落正在后台处理，请等待完成后保存；草稿可以保留。")
        execution_lock = _job_lock.acquire(job_dir(job_id) / ".worker.lock", blocking=False)
        if execution_lock is None:
            raise RuntimeError("任务正在后台处理，请等待完成后保存；草稿可以保留。")
    try:
        _reconcile_translation_failures(state)
        try:
            _job_repo.save_job_state_atomic(job_id, state)
        except OSError as exc:
            _workflow_errors.raise_if_fatal(exc, job_id)
            raise
        try:
            _state_journal.rotate_if_applied(job_id, state)
        except OSError:
            pass
    finally:
        if execution_lock is not None:
            execution_lock.close()



# ================= 未保存译文草稿的持久化 =================
# 草稿**不是**文档状态的一部分：它不进 state.json、不改 reviewed/review_status、
# 不进入交付资产。它只回答一个问题——"用户敲进去但还没点保存的内容"，在会话被
# 重建（刷新、关标签页、进程重启）之后还在不在。此前它只活在 st.session_state 里，
# 刷新即丢；会话会重置这件事我们改不了，但"草稿只存在会话里"是可以改的。
TRANSLATION_DRAFTS_VERSION = 1
TRANSLATION_DRAFTS_FILE = "translation_drafts.json"



# ================= 段落身份 / 排除 / 结构快照 =================
# `pairs` 与 `paras` 是贯穿流水线的位置契约：审校发现、检查点、交付资产全部按
# index 引用同一条记录。拆分、合并、插入、删除会改变**每一个后续 index 的含义**，
# 而界面需要"同一段在结构操作之后仍然是同一段"（编辑框、选中态、Agent 候选都不
# 能跟着旧索引漂移）。`segment_uid` 只为界面身份服务，数据契约保持按索引不变。
SEGMENT_UID_FIELD = "segment_uid"
_UID_SEQ_FIELD = "segment_uid_seq"

SEGMENT_STRUCTURE_OPERATION_LABELS = {
    "source_edit": "修改原文",
    "split": "拆分段落",
    "merge": "合并段落",
    "insert": "插入空段",
    "delete": "删除空段",
    "exclude": "排除段落",
    "include": "恢复段落",
}



@_serialized_job_mutation
def save_translation_edit(job_id, index, target, actor="user"):
    """Save one human translation edit through the state/business layer.

    A manual edit is a new working translation: it preserves the paragraph
    identity, clears segment review, removes TM reuse provenance, and invalidates
    mutable final approval while leaving historical frozen snapshots untouched.
    """
    from folith import delivery as _delivery

    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    pairs = state.get("pairs") or []
    if not (0 <= index < len(pairs)):
        raise IndexError(f"段落索引超出范围：{index}")
    pair = pairs[index]
    if is_excluded(pair):
        raise ValueError("该段落已被排除，不会进入翻译与交付；请先恢复本段再编辑译文。")
    assign_missing_segment_uids(state, job_id=job_id)
    new_target = str(target or "").strip()
    if not pair.get("human_edited"):
        pair["_translation_edit_restore"] = {
            "target": pair.get("target") or "",
            "reviewed": bool(pair.get("reviewed")),
            "review_status": pair.get("review_status", "not_reviewed"),
            "target_provenance": pair.get("target_provenance", "generated"),
            "from_tm": bool(pair.get("from_tm")),
        }
    pair["target"] = new_target
    pair["human_edited"] = True
    pair["reviewed"] = False
    pair["review_status"] = "not_reviewed"
    pair["target_provenance"] = "human_edit"
    pair["from_tm"] = False
    _recount_reviewed_segments(state)

    _mark_translation_truth_changed(
        job_id, state, [index], "人工修改 CURRENT_TRANSLATION；相关案例与学术下游需要重建",
        actor=actor, action="translation_edit")
    _recheck_delivery_invariants_for_segments(state, [index], actor=actor)
    save_job_state(job_id, state)
    return state



def list_jobs():
    """列出本地任务。**纯读取**：目录不存在时返回空列表，不创建它。"""
    jobs = []
    if not OUTPUT_DIR.is_dir():
        return jobs
    for d in sorted(OUTPUT_DIR.iterdir()):
        if not d.is_dir():
            continue
        s = _job_repo.load_job_state_cached(d.name)
        if s is not None:
            jobs.append({"job_id": d.name, "state": s})
    return jobs


def list_job_summaries():
    """列出本地任务的轻量级摘要（派生读取模型）。"""
    return _job_repo.list_job_summaries()


def read_runtime_view(job_id):
    """读取任务的轻量级运行时视图，避免每次轮询读取并解析完整 state.json。"""
    return _job_repo.read_runtime_view(job_id)


def delete_job(job_id, *, allow_active=False):
    """永久删除一个翻译任务（含其输出目录）。

    正在运行的任务默认拒绝删除：worker 还在往这个目录写状态，删掉会留下一个
    半写状态并让运行时报错。返回是否真的删除。

    The removal happens under the task's state transaction lock, so a concurrent
    ``save_job_state`` from another surface can neither be lost nor resurrect the
    deleted task directory (the atomic replace would otherwise re-create it).
    """
    if not str(job_id or "").strip():
        return False
    d = job_dir(job_id)
    if not d.exists():
        return False
    if not allow_active and job_is_active(job_id):
        raise ValueError("任务正在运行，无法删除；请先取消或等它结束。")
    with job_state_mutation_lock(job_id):
        if not d.exists():
            return False
        shutil.rmtree(d)
        _job_repo.invalidate_cache(job_id)
    return True


def file_job_id(file_bytes):
    """**文档身份**：文件内容哈希。同一份文档重传得到同一个值。

    注意：这**不是任务身份**。同一份文档在不同项目 / 不同目标语言下是彼此
    独立的本地化任务，内容哈希无法区分它们——用内容哈希当任务 ID 会让
    "同一个文件、另一个项目、另一种目标语言"静默打开旧任务。任务身份见
    `task_job_id` / `resolve_task_id`。
    """
    return hashlib.sha256(file_bytes).hexdigest()[:16]


# 任务身份 = 文档身份 + 本地化上下文（项目 + 目标语言 + 源语言）。
# 文档内容相同不等于本地化任务相同：项目决定注入哪一套项目记忆与术语，
# 目标语言决定译文本身。把它们丢掉，就等于让内容哈希替用户决定
# "这两个语义不同的活是同一个活"。
_TASK_ID_FIELDS = ("project_id", "target_lang", "source_lang")


def task_context(project_id=None, target_lang=None, source_lang=None):
    """本地化上下文的规范化三元组（用于任务身份与"是否同一个任务"的比较）。"""
    return {
        "project_id": resolved_project_id({"project_id": project_id}),
        "target_lang": normalize_language(target_lang),
        "source_lang": normalize_language(source_lang),
    }


def task_context_of(state):
    """一个已存在任务记录的本地化上下文（缺字段按"未指定"处理）。"""
    state = state or {}
    return task_context(state.get("project_id"), state.get("target_lang"),
                        state.get("source_lang"))


def task_job_id(file_bytes, *, project_id=None, target_lang=None,
                source_lang=None):
    """**任务身份**：文档身份 + 本地化上下文，确定性推导，不读磁盘。"""
    context = task_context(project_id, target_lang, source_lang)
    digest = "\x1f".join([file_job_id(file_bytes)]
                         + [context[field] for field in _TASK_ID_FIELDS])
    return hashlib.sha256(digest.encode("utf-8")).hexdigest()[:16]


def resolve_task_id(file_bytes, *, project_id=None, target_lang=None,
                    source_lang=None):
    """这份文档在**这个本地化上下文**下应当打开的任务 ID。

    三条规则，顺序即优先级：

    1. 按上下文推导出的任务已存在 -> 续做它（"同一个文件、同一个项目、
       同一种目标语言"就是同一个任务，重传即续传）；
    2. 否则，若历史版本留下的、以**内容哈希**命名的旧任务存在，且它记录的
       项目 / 目标语言与请求的上下文**可证明一致** -> 沿用它（旧任务不因
       这次修复而失联）；
    3. 否则 -> 用推导出的新 ID 建一个独立任务。

    规则 2 里的"可证明一致"是关键：旧任务没有目标语言记录时**不会**被认领，
    宁可让用户重建一个任务，也不把另一种语言的旧任务当成这个活。
    """
    context = task_context(project_id, target_lang, source_lang)
    candidate = task_job_id(file_bytes, project_id=project_id,
                            target_lang=target_lang, source_lang=source_lang)
    if load_job_state(candidate) is not None:
        return candidate
    legacy_id = file_job_id(file_bytes)
    if legacy_id != candidate:
        legacy_state = load_job_state(legacy_id)
        if legacy_state is not None and task_context_of(legacy_state) == context:
            return legacy_id
    return candidate


def resolve_job_id(job_ref):
    """根据任务 ID 或短哈希/短编码解析出确切已存在的 job_id。

    - 精确匹配：磁盘存在该任务且能成功加载状态 -> 直接返回；
    - 短哈希前缀匹配：长度 >= 4，大小写无关匹配本地任务目录；
      若有匹配且状态有效，按最近修改时间返回最匹配的合法任务 ID；
    - 未找到返回 None。
    """
    token = str(job_ref or "").strip()
    if not token:
        return None
    if load_job_state(token) is not None:
        return token
    if len(token) >= 4 and OUTPUT_DIR.is_dir():
        token_lower = token.casefold()
        matches = []
        for d in sorted(OUTPUT_DIR.iterdir(),
                        key=lambda p: p.stat().st_mtime if p.is_dir() else 0,
                        reverse=True):
            if d.is_dir() and d.name.casefold().startswith(token_lower):
                if load_job_state(d.name) is not None:
                    matches.append(d.name)
        if matches:
            return matches[0]
    return None


def save_source(job_id, file_bytes):
    """留存源文件，刷新页面后即使不重新上传也能继续（如重做阶段一）。

    **原子写入**：`write_bytes` 会先截断再写，而读取方（另一个 worker、启动前的
    `load_source`、OCR 阶段）可能正好落在截断与写入之间，读到 0 字节的"空源文件"
    ——随后 `fitz.open(stream=b"")` 只会报 "Cannot open empty stream"，把一次并发
    读写变成一次无法解释的提取失败（实测出现过）。临时文件 + `replace` 让读者要么
    看到旧内容、要么看到完整新内容。
    """
    d = job_dir(job_id)
    d.mkdir(parents=True, exist_ok=True)
    target = d / "source.bin"
    payload = bytes(file_bytes or b"")
    temp = d / f".source.bin.{uuid.uuid4().hex}.tmp"
    temp.write_bytes(payload)
    temp.replace(target)


def load_source(job_id):
    p = job_dir(job_id) / "source.bin"
    return p.read_bytes() if p.is_file() else None


@_serialized_job_mutation
def save_report_template(job_id, filename, template_bytes):
    """Parse and persist the uploaded DOCX template before academic stages run."""
    from folith import report_template

    raw = _bytes(template_bytes)
    contract = report_template.parse_docx_template(filename, raw)
    directory = job_dir(job_id)
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "report-template.docx").write_bytes(raw)
    (directory / "template-contract.json").write_text(
        json.dumps(contract, ensure_ascii=False, indent=2), encoding="utf-8")
    state = load_job_state(job_id)
    if state is not None:
        identity = contract.get("template_identity") or {}
        state["report_template"] = {
            "filename": identity.get("filename"),
            "template_id": identity.get("template_id"),
            "template_hash": identity.get("sha256"),
            "schema_version": contract.get("schema_version"),
            "status": "parsed",
        }
        state["report_template_contract"] = contract
        settings = dict(state.get("research_settings") or {})
        settings["report_template_contract"] = contract
        state["research_settings"] = settings
        save_job_state(job_id, state)
    return contract


@_serialized_job_mutation
def clear_report_template(job_id):
    """Remove the saved report template and its contract from a task."""
    directory = job_dir(job_id)
    for name in ("report-template.docx", "template-contract.json"):
        path = directory / name
        if path.is_file():
            path.unlink()
    state = load_job_state(job_id)
    if state is not None:
        state["report_template"] = None
        state["report_template_contract"] = None
        settings = dict(state.get("research_settings") or {})
        settings.pop("report_template_contract", None)
        settings.pop("template_contract", None)
        state["research_settings"] = settings
        save_job_state(job_id, state)
    return state


def load_report_template(job_id):
    """Load the persisted DOCX bytes and canonical contract, if configured."""
    from folith import report_template

    directory = job_dir(job_id)
    docx_path = directory / "report-template.docx"
    contract_path = directory / "template-contract.json"
    state = load_job_state(job_id)
    contract = None
    if contract_path.is_file():
        try:
            value = json.loads(contract_path.read_text(encoding="utf-8"))
            contract = value if isinstance(value, dict) else None
        except (OSError, ValueError):
            contract = None
    if contract is None and state:
        contract = state.get("report_template_contract") or \
            (state.get("research_settings") or {}).get("report_template_contract")
    if not contract or not docx_path.is_file():
        return None
    raw = docx_path.read_bytes()
    return {
        "bytes": raw,
        "contract": contract,
        "metadata": state.get("report_template") if state else None,
        "summary": report_template.contract_summary(contract),
    }


def _bytes(value):
    return value.getvalue() if hasattr(value, "getvalue") else bytes(value)


@_serialized_job_mutation
def report_docx_bytes(job_id, state=None, frozen_assets=None):
    """Render the structured report with its template, or use the legacy fallback."""
    if frozen_assets and frozen_assets.get("stage3_report.docx"):
        return frozen_assets["stage3_report.docx"]
    state = state or load_job_state(job_id) or {}
    report = state.get("p3_md")
    if not report:
        return None
    template = load_report_template(job_id)
    if template:
        from folith import academic_evidence, academic_writer, compliance
        from folith import final_docx, report_template
        artifact = load_academic_artifact(job_id, "report")
        report_record = academic_writer.artifact_record(state, "report")
        if isinstance((state.get("academic_state") or {}).get("artifacts", {}).get(
                "report"), dict) and report_record.get("status") != "valid":
            return None
        if not artifact:
            raise report_template.TemplateParseError(
                "模板化报告缺少结构化 report artifact，请重新生成报告。")
        if artifact.get("report_status") not in {
                "generated", "review_required", "literature_required"} or artifact.get(
                "template_compliance") not in {"pass", "pass_with_warnings"}:
            return None
        rendered = _bytes(report_template.render_report_docx(
            artifact, template["bytes"], template["contract"]))
        source_docx_hash = _rendered_qa.sha256(rendered)
        previous_docx = academic_writer.artifact_record(
            state, "final_docx_validation")
        previous_qa = _finalization.normalize_final_qa(state.get("final_qa"))
        previous_docx_hash = str(
            previous_docx.get("source_docx_hash") or
            previous_qa.get("source_docx_hash") or "")
        if (previous_docx_hash and previous_docx_hash != source_docx_hash) or (
                isinstance(previous_docx, dict) and previous_docx and (
                    previous_qa.get("author_visual_review") == "CONFIRMED" or
                    previous_qa.get("word_final_review") == "CONFIRMED")):
            academic_writer.propagate_artifact_staleness(
                state, input_artifact_ids=["final_docx_validation"])
            _reset_final_qa(state, "DOCX bytes changed; render and human reviews reset")
        final_validation = final_docx.validate_final_docx(rendered, artifact)
        final_validation.update({
            "source_docx_hash": source_docx_hash,
            "report_content_hash": (state.get("academic_state") or {}).get(
                "artifacts", {}).get("report", {}).get("content_hash") or
                artifact.get("content_hash"),
            "layout_facts": compliance.inspect_docx_layout(rendered),
        })
        final_validation["content_hash"] = academic_evidence.stable_hash({
            key: value for key, value in final_validation.items()
            if key != "content_hash"
        })
        final_validation_dep = academic_evidence.stable_hash({
            "report": (state.get("academic_state") or {}).get("artifacts", {}).get(
                "report", {}).get("content_hash") or artifact.get("content_hash"),
            "template": template.get("contract", {}).get("template_identity") or
            template.get("contract", {}).get("template_hash"),
            "version": final_docx.SCHEMA_VERSION,
        })
        academic_writer._save_artifact(
            state, job_dir(job_id), "final_docx_validation", final_validation,
            final_validation_dep, final_docx.SCHEMA_VERSION,
            input_artifact_ids=["report"],
            status="valid" if final_validation.get("status") != "fail" else "failed")
        qa = _finalization.normalize_final_qa(state.get("final_qa"))
        qa.update({
            "structural_qa": "FAIL" if final_validation.get("status") == "fail"
            else "PASS",
            "source_docx_hash": source_docx_hash,
            "updated_at": _finalization.now_iso(),
        })
        state["final_qa"] = qa
        save_compliance_record(job_id, state)
        generate_report_qa(job_id, state, save_file=True)
        save_job_state(job_id, state)
        if final_validation.get("status") == "fail":
            return None
        return rendered
    if state.get("report_status") in {
            "incomplete", "failed_template_validation", "review_required"}:
        return None
    from folith import academic_writer
    if isinstance((state.get("academic_state") or {}).get("artifacts", {}).get(
            "report"), dict) and academic_writer.artifact_record(
                state, "report").get("status") != "valid":
        return None
    return _bytes(markdown_to_word(report, state.get("theory") or ""))



def _delivery_asset_bundle(job_id, state, target_lang, provider, model):
    """Build the configured delivery set used by both preview and snapshots."""
    from folith import assets as _assets
    from folith import report_evidence as _report_evidence

    validation = validate_delivery_translation_state(state)
    _record_delivery_validation_findings(state, validation)
    if validation["blocking"]:
        save_job_state(job_id, state)
        reasons = "；".join(
            issue.get("message", "目标文本未通过交付检查")
            for issue in validation.get("issues") or []
        )
        raise RuntimeError(
            "交付被 Translation Target Invariant 阻止"
            + (f"：{reasons}" if reasons else "")
        )

    source = load_source(job_id)
    snapshot_state = dict(state)
    if source is not None:
        snapshot_state["_source_bin"] = source
    filename = state.get("filename", "")
    configured = state.get("delivery_config")
    if isinstance(configured, dict) and configured:
        config = normalize_delivery_config(
            configured, enable_report=state.get("report_enabled", False),
            enable_annotate=state.get("enable_annotate", False))
        bundle = {}
        pairs = state.get("pairs") or []
        glossary = state.get("glossary") or []
        # Requested artifacts that cannot exist for this source.  Initialised
        # outside the translation branch because the delivery manifest is built
        # for a job without completed translations as well.
        skipped_assets: Dict[str, str] = {}
        if state.get("source_cleanup_done"):
            bundle["source_cleanup.json"] = (
                json.dumps(state.get("source_cleanup") or {},
                           ensure_ascii=False, indent=2) + "\n"
            ).encode("utf-8")
        if state.get("p2_done") and pairs:
            if config["deliver_plain_docx"]:
                bundle["translation.docx"] = _bytes(translations_to_word(
                    pairs, target_lang=state.get("target_lang") or target_lang))
            if config["deliver_bilingual_docx"]:
                bundle["bilingual.docx"] = _bytes(pairs_to_word(
                    pairs, source_lang=state.get("source_lang"),
                    target_lang=state.get("target_lang") or target_lang))
            if config["deliver_pdf"]:
                bundle["translation.pdf"] = translations_to_pdf(
                    pairs, target_lang=state.get("target_lang") or target_lang)
            if config["deliver_original_pdf"]:
                # The original-layout PDF keeps the source document's page
                # geometry and carries the translated text in place.  It only
                # exists for a PDF source; a failure on a real PDF is surfaced,
                # never silently replaced by the reflow PDF.
                if not _is_pdf_source(filename, source):
                    skipped_assets["original_layout.pdf"] = (
                        "original-layout rendering requires a PDF source")
                else:
                    from folith import pdfrender as _pdfrender
                    cache_dir = job_dir(job_id) / "original_pdf_cache"
                    try:
                        rendered = _pdfrender.render_original_layout(
                            job_id, state, source, page_cache_dir=cache_dir)
                    except Exception as exc:
                        raise RuntimeError(
                            "原排版 PDF 生成失败："
                            f"{type(exc).__name__}: {exc}") from exc
                    bundle["original_layout.pdf"] = rendered.data
                    bundle["original_layout_report.json"] = (
                        json.dumps({k: v for k, v in rendered.report.items()
                                    if not str(k).startswith("_")},
                                   ensure_ascii=False, indent=2) + "\n"
                    ).encode("utf-8")
            if config["deliver_epub"]:
                from folith.epub import build_translation_epub
                epub_bytes, epub_report = build_translation_epub(
                    state, source, job_id=job_id)
                bundle["translation.epub"] = epub_bytes
                bundle["epub_export_report.json"] = (
                    json.dumps(epub_report, ensure_ascii=False, indent=2) + "\n"
                ).encode("utf-8")
            if config["enable_annotate"]:
                bundle["annotated_bilingual.docx"] = _bytes(pairs_to_word(
                    pairs, annotations=state.get("annotations"),
                    colors=ANNOTATION_COLORS, source_lang=state.get("source_lang"),
                    target_lang=state.get("target_lang") or target_lang))
            if config["deliver_terms_xlsx"]:
                bundle["terms.xlsx"] = _bytes(glossary_to_excel(
                    glossary, state.get("auto_terms")))
            if config["deliver_tbx"]:
                bundle["terms.tbx"] = _assets.build_tbx(glossary, *_assets.asset_languages(state, target_lang))
            if config["deliver_tmx"]:
                bundle["memory.tmx"] = _assets.build_tmx(state, *_assets.asset_languages(state, target_lang), job_id=job_id)
            if config["deliver_jsonl"]:
                bundle["bilingual.jsonl"] = _assets.build_jsonl(
                    state, job_id=job_id).encode("utf-8")
            if config["deliver_evidence"]:
                bundle["segment_evidence.jsonl"] = \
                    _report_evidence.export_segment_evidence_jsonl(
                        state, job_id).encode("utf-8")
            if config["deliver_review_report"]:
                bundle["review_report.md"] = findings_report_md(state).encode("utf-8")
        # 被排除的内容不进译文文档，但必须在交付包里留下痕迹：否则"我排除了
        # 12 段"会变成一次静默删除，用户在交付物上再也看不到它。
        excluded_manifest = excluded_segments_manifest(state)
        if excluded_manifest["count"]:
            bundle["excluded_segments.md"] = _excluded_segments_md(
                excluded_manifest).encode("utf-8")
            bundle["excluded_segments.json"] = (
                json.dumps(excluded_manifest, ensure_ascii=False, indent=2)
                + "\n").encode("utf-8")
        if config["deliver_cases"]:
            selected_cases = load_academic_artifact(job_id, "selected_cases")
            if selected_cases:
                bundle["selected_cases.json"] = (
                    json.dumps(selected_cases, ensure_ascii=False, indent=2) + "\n"
                ).encode("utf-8")
        if config["deliver_academic_workspace"]:
            workspace = _academic_workspace_archive(job_id)
            if workspace:
                bundle["academic_workspace.zip"] = workspace
        if config["enable_report"] and state.get("p3_md"):
            report_docx = report_docx_bytes(job_id, state)
            if report_docx is not None:
                bundle["report.docx"] = report_docx
            bundle["report.md"] = state["p3_md"].encode("utf-8")
        generated_assets = sorted([*bundle, "delivery_manifest.json"])
        manifest = _assets.build_delivery_manifest(
            snapshot_state, job_id, target_lang, provider, model,
            generated_assets=generated_assets, source_filename=filename,
            translator_config=state.get("translator_config"),
            reviewer_config=state.get("reviewer_config"),
            skipped_assets=skipped_assets)
        bundle["delivery_manifest.json"] = (
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        return bundle

    # Existing jobs without a saved selection retain their historical bundle.
    bundle = {}
    if state.get("p1_done") and state.get("paras"):
        # ``paras`` is the sentence-sized CAT contract; the stage-1 source
        # artifact should retain the recovered paragraph layout when available.
        cleaned_source = state.get("source_paragraphs") or state.get("paras")
        bundle["stage1_cleaned.docx"] = _bytes(paragraphs_to_word(
            cleaned_source, source_lang=state.get("source_lang")))
    if state.get("source_cleanup_done"):
        bundle["source_cleanup.json"] = (
            json.dumps(state.get("source_cleanup") or {},
                       ensure_ascii=False, indent=2) + "\n"
        ).encode("utf-8")
    if state.get("auto_terms"):
        bundle["auto_terms.xlsx"] = _bytes(dict_to_excel(state["auto_terms"]))
    if state.get("p2_done") and state.get("pairs"):
        bundle["stage2_bilingual.docx"] = _bytes(pairs_to_word(
            state["pairs"], annotations=state.get("annotations"), colors=ANNOTATION_COLORS,
            source_lang=state.get("source_lang"), target_lang=state.get("target_lang") or target_lang))
    if state.get("p3_md"):
        report_docx = report_docx_bytes(job_id, state)
        if report_docx is not None:
            bundle["stage3_report.docx"] = report_docx
    bundle.update(_assets.export_all(
        snapshot_state, job_id, target_lang, provider, model,
        source_filename=filename, source_bin=source))
    # 旧任务没有保存过 delivery_config，但排除段仍然必须在交付包中
    # 留下可追溯清单；否则历史任务会把排除误表现成静默删除。
    excluded_manifest = excluded_segments_manifest(state)
    if excluded_manifest["count"]:
        bundle["excluded_segments.md"] = _excluded_segments_md(
            excluded_manifest).encode("utf-8")
        bundle["excluded_segments.json"] = (
            json.dumps(excluded_manifest, ensure_ascii=False, indent=2)
            + "\n").encode("utf-8")
    if state.get("p2_done"):
        bundle["segment_evidence.jsonl"] = _report_evidence.export_segment_evidence_jsonl(
            state, job_id).encode("utf-8")
        if state.get("findings"):
            bundle["review_report.md"] = findings_report_md(state).encode("utf-8")
    manifest = _assets.build_delivery_manifest(
        snapshot_state, job_id, target_lang, provider, model,
        generated_assets=sorted(bundle), source_filename=filename,
        translator_config=state.get("translator_config"),
        reviewer_config=state.get("reviewer_config"))
    bundle["delivery_manifest.json"] = (
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    return bundle



def build_delivery_assets(job_id, state=None, target_lang="", provider="", model=""):
    state = state or load_job_state(job_id) or {}
    with _delivery_timing(job_id, state, "delivery_assets"):
        return _delivery_asset_bundle(
            job_id, state, target_lang or state.get("target_lang") or "",
            provider or state.get("provider") or "",
            model or state.get("model") or "")


def _reset_source_dependent_state(state):
    """Invalidate artifacts derived from source paragraphs after cleanup."""
    state.update({
        "profile_done": False,
        "document_profile": None,
        "profile_warnings": [],
        "semantic_units": [],
        "section_digests": [],
        "document_synopsis": None,
        "understanding_done": False,
        "understanding_warnings": [],
        "context_packet_log": [],
        "llm_usage": _usage.empty_usage(),
        "knowledge_feedback_policy": {},
        "targeted_final_review": {
            "status": "not_run", "segment_ids": [],
            "reviewed_segment_ids": [], "failed_segment_ids": [],
        },
        "knowledge_candidates": [],
        "translation_continuity": [],
        "knowledge_events": [],
        "confirmed_style_rules": [],
        "knowledge_feedback_failures": 0,
        "entity_registry": [],
        "translation_failures": [],
        "review_evidence": [],
        "repair_overlays": [],
        "auto_term_entries": [],
        "auto_terms": {},
        "glossary": [],
        "glossary_draft": [],
        "glossary_frozen": None,
        "glossary_versions": [],
        "glossary_injection_log": [],
        "pairs": [],
        # Structural snapshots belong to the previous source, even if the
        # re-extraction happens to produce the same paragraph count.
        "segment_structure_history": [],
        "segment_structure_undo_log": [],
        "excluded_segments": [],
        "findings": [],
        "has_blocking": False,
        "review_stats": {
            "reviewed_segments": 0, "batches_reviewed": 0,
            "batches_skipped_low_risk": 0, "skipped_segments": 0,
            "blocking": 0, "actionable": 0,
            "informational": 0, "review_failed": 0,
        },
        "p2_done": False,
        "p3_done": False,
        "p3_md": "",
        "p3_sections": [],
        "annotations": {},
        "annotations_done": False,
        "annotations_done_offset": 0,
        "annotation_batches": {},
        "annotations_failed_batches": 0,
        "batch_plan": {},
        "delivery_status": "draft",
        "delivery_approved_by_human": False,
        "delivery_approval": None,
        "delivery_validation": {},
        "delivery_snapshots": [],
        "latest_delivery_snapshot_version": None,
        "exported_assets": [],
        "translation_truth": {
            "authority": "CURRENT_TRANSLATION",
            "version": 0,
            "last_changed_at": None,
            "last_change": None,
        },
        "source_quality_gate": {},
        "source_quality_auto_exclusions": [],
    })


def rescan_pdf_source(job_id, *, filename=None, provider, api_key, model,
                      base_url=None, on_status=None, max_batch_chars=None,
                      max_batch_items=None, parallelism=None, ocr_workers=None,
                      ocr_queue_size=None, confidence_threshold=None,
                      max_retries=None, request_interval_seconds=None,
                      segmentation_mode=None, ocr_language=None, ocr_dpi=None,
                      ocr_psm=None, ocr_quality=None, ocr_mode=None, ocr_backend=None,
                      ocr_paddle_model=None, ocr_online=None,
                      _allow_active=False):
    """Re-extract and LLM-clean a saved PDF without running later stages.

    This is the explicit "重新从 PDF 扫一遍" operation.  It deliberately
    invalidates every artifact derived from the previous source text, including
    partial translation pairs, so the next resume cannot pair old translations
    with newly merged or corrected source paragraphs.

    ``ocr_language`` may change the OCR language for the rescan.  Because the
    OCR checkpoint identity and the cleanup fingerprint both include the
    language, the old pages are never reused, and every source-derived artifact
    is reset below.
    """
    if not str(job_id or "").strip():
        raise ValueError("缺少任务 ID")
    if not _allow_active and (is_job_worker_alive(job_id) or job_is_active(job_id)):
        raise RuntimeError("任务正在运行，无法重新扫描原文")
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    resolved_filename = str(filename or state.get("filename") or "")
    if not resolved_filename.lower().endswith(".pdf"):
        raise ValueError("原文重新扫描只支持 PDF 任务")
    source = load_source(job_id)
    if source is None:
        raise ValueError("任务没有已保存的 PDF 源文件，请重新上传后再试")

    previous_pairs = len(state.get("pairs") or [])
    previous_paragraphs = len(state.get("paras") or [])
    old_base_url = getattr(_LLM_CTX, "base_url", None)
    effective_cleanup_base_url = base_url if base_url is not None else old_base_url
    saved_config = state.get("pipeline_config") or {}
    effective_ocr_language, effective_ocr_dpi, effective_ocr_psm, \
        effective_ocr_quality, resolved_ocr = _effective_ocr_config(
            ocr_language, ocr_dpi, ocr_psm, saved_config, ocr_quality)
    previous_ocr_language = str(saved_config.get("ocr_language") or "")
    ocr_language_changed = (
        str(effective_ocr_language or "") != previous_ocr_language)
    effective_ocr_backend, effective_paddle_model, effective_online_descriptor, \
        effective_online_config = _effective_ocr_source(
            ocr_backend, saved_config, ocr_paddle_model=ocr_paddle_model,
            ocr_online=ocr_online)
    previous_ocr_backend = _ocr_backends.normalize_backend(
        saved_config.get("ocr_backend", "tesseract"))
    previous_paddle_model = str(saved_config.get("ocr_paddle_model") or "official")
    # A different online target/model/prompt/parameter set is a different OCR
    # configuration even though the source name ("online") is unchanged.
    previous_online_descriptor = saved_config.get("ocr_online") \
        if isinstance(saved_config.get("ocr_online"), dict) else {}
    previous_online_key = str(previous_online_descriptor.get("cache_options") or "")
    effective_online_key = str((effective_online_descriptor or {}).get("cache_options") or "")
    ocr_source_changed = (effective_ocr_backend != previous_ocr_backend
                          or effective_paddle_model != previous_paddle_model
                          or effective_online_key != previous_online_key)
    effective_ocr_mode = normalize_ocr_mode(
        ocr_mode if ocr_mode is not None else saved_config.get("ocr_mode"))
    previous_ocr_mode = normalize_ocr_mode(saved_config.get("ocr_mode"))
    ocr_mode_changed = effective_ocr_mode != previous_ocr_mode
    effective_parallelism = _runtime_int_option(
        parallelism, saved_config.get("source_cleanup_concurrency"),
        "FOLIOTHREAD_LLM_CLEANUP_CONCURRENCY",
        _source_cleanup.DEFAULT_PARALLELISM, maximum=16)
    effective_ocr_workers = _runtime_int_option(
        ocr_workers, saved_config.get("ocr_workers"), "FOLIOTHREAD_OCR_WORKERS",
        OCR_WORKERS_DEFAULT, maximum=8)
    effective_ocr_queue_size = _runtime_int_option(
        ocr_queue_size, saved_config.get("ocr_queue_size"),
        "FOLIOTHREAD_OCR_QUEUE_SIZE", effective_ocr_workers * 2, maximum=32)
    effective_confidence_threshold = _runtime_float_option(
        confidence_threshold, saved_config.get("source_cleanup_confidence_threshold"),
        "FOLIOTHREAD_OCR_CONFIDENCE_THRESHOLD",
        OCR_CONFIDENCE_THRESHOLD_DEFAULT, maximum=100.0)
    effective_max_retries = _runtime_int_option(
        max_retries, saved_config.get("source_cleanup_max_retries"),
        "FOLIOTHREAD_LLM_CLEANUP_MAX_RETRIES",
        _source_cleanup.DEFAULT_MAX_RETRIES, minimum=0, maximum=6)
    effective_request_interval = _runtime_float_option(
        request_interval_seconds, saved_config.get("source_cleanup_request_interval"),
        "FOLIOTHREAD_LLM_CLEANUP_REQUEST_INTERVAL", 0.0, maximum=60.0)
    active_run = trace.start_run(
        job_id=job_id, filename=filename or resolved_filename,
        output_dir=job_dir(job_id)
    )
    profiler = PipelineProfiler(
        job_dir(job_id) / "performance.json", run_id=active_run.run_id
    )

    def report(label):
        if on_status:
            on_status(label)
        # A direct invocation has no worker thread, so mirror the same durable
        # progress surface that the normal worker provides.
        _runtime_status_callback(job_id, state, label)

    update_runtime_state(
        job_id, status="running", phase="source_cleanup",
        stage="document_processing", stage_id="document_processing",
        operation="source_cleanup", operation_id="source_cleanup",
        operation_label="重新扫描 PDF 原文", phase_label="正在扫描并纠错原文",
        cancel_requested=False, error=None, progress=True,
        event="开始重新扫描 PDF 原文", event_name="source_rescan_started",
        event_visibility="user", event_category="progress")
    set_llm_base_url(effective_cleanup_base_url)
    try:
        report("【阶段一】重新从 PDF 提取原文…")
        raw_paragraphs, extraction_warnings, extraction_report = \
            extract_document_paragraphs_with_report(
                resolved_filename, source, ocr_max_pages=None, on_progress=report,
                profiler=profiler, checkpoint_dir=job_dir(job_id),
                ocr_workers=effective_ocr_workers,
                ocr_queue_size=effective_ocr_queue_size,
                cancel_check=lambda: _runtime_cancel_requested(job_id),
                profile_cache_dir=ocr_profile_cache_dir(source),
                ocr_language=effective_ocr_language,
                ocr_dpi=effective_ocr_dpi,
                ocr_psm=effective_ocr_psm,
                ocr_mode=effective_ocr_mode,
                ocr_backend=effective_ocr_backend,
                ocr_paddle_model=effective_paddle_model,
                ocr_online_config=effective_online_config,
                source_language=(state.get("source_lang")
                                 or (resolved_ocr.label if resolved_ocr else "")))
        if not raw_paragraphs:
            detail = (extraction_failure_detail(extraction_warnings)
                      if extraction_warnings else _NO_TEXT_REASON)
            raise ValueError(f"未提取到有效文本：{detail}")

        report("【阶段一】使用最新纠错功能整理 OCR 错字与断行…")
        inherited_job_id = getattr(_RUNTIME_CTX, "job_id", None)
        inherited_reasoning = getattr(_LLM_CTX, "reasoning_effort", None)

        def cleanup_call(provider_name, api_key_value, model_name,
                         system_prompt, user_prompt, **kwargs):
            if inherited_job_id:
                _RUNTIME_CTX.job_id = inherited_job_id
            inherited_kwargs = {}
            if effective_cleanup_base_url:
                inherited_kwargs["base_url"] = effective_cleanup_base_url
            if inherited_reasoning:
                inherited_kwargs["reasoning_effort"] = inherited_reasoning
            return call_llm(
                provider_name, api_key_value, model_name, system_prompt,
                user_prompt, **inherited_kwargs, **kwargs)
        if extraction_report.get("semantic_source") in {"vlm", "mixed"}:
            cleaned, cleanup_meta, cleanup_warnings = raw_paragraphs, {
                "status":"skipped", "reason":"vlm_semantics_no_ocr_repair"}, []
        else:
            cleaned, cleanup_meta, cleanup_warnings = cleanup_source_paragraphs(
                raw_paragraphs, provider, api_key, model, call_llm_fn=cleanup_call,
                on_progress=report, max_batch_chars=max_batch_chars,
                max_batch_items=max_batch_items, parallelism=effective_parallelism,
                confidence_by_index=((extraction_report.get("ocr") or {}).get(
                    "paragraph_confidences") if isinstance(extraction_report, dict) else None),
                confidence_threshold=effective_confidence_threshold,
                checkpoint_dir=job_dir(job_id),
                cancel_check=lambda: _runtime_cancel_requested(job_id),
                max_retries=effective_max_retries,
                request_interval_seconds=effective_request_interval,
                profiler=profiler,
                ocr_language=effective_ocr_language or "",
                source_language=(resolved_ocr.label if resolved_ocr else ""),
                text_direction=(resolved_ocr.direction if resolved_ocr else ""),
                ocr_engine_code=(resolved_ocr.engine_code if resolved_ocr else ""),
                    ocr_backend=((extraction_report.get("ocr") or {}).get("actual_engine") or effective_ocr_backend))
        state["extraction_report"] = extraction_report
        cleanup_meta = dict(cleanup_meta or {})
        segments, segmentation_meta = _apply_source_segmentation(
            state, cleaned, mode=segmentation_mode)
        cleanup_meta.update({
            "forced_rescan": True,
            "previous_pair_count": previous_pairs,
            "previous_paragraph_count": previous_paragraphs,
            "extraction_warning_count": len(extraction_warnings),
            "parallelism": effective_parallelism,
            "confidence_threshold": effective_confidence_threshold,
            "ocr_workers": effective_ocr_workers,
            "ocr_queue_size": effective_ocr_queue_size,
            "ocr_language": effective_ocr_language or "",
            "ocr_language_engine_code": (
                resolved_ocr.engine_code if resolved_ocr else ""),
            "ocr_language_label": resolved_ocr.label if resolved_ocr else "",
            "ocr_text_direction": resolved_ocr.direction if resolved_ocr else "",
            "ocr_dpi": effective_ocr_dpi,
            "ocr_psm": effective_ocr_psm,
            "ocr_quality": effective_ocr_quality,
            "ocr_language_changed": ocr_language_changed,
            "previous_ocr_language": previous_ocr_language,
            "ocr_mode": effective_ocr_mode,
            "ocr_mode_changed": ocr_mode_changed,
            "previous_ocr_mode": previous_ocr_mode,
            "ocr_source": effective_ocr_backend,
            "ocr_source_changed": ocr_source_changed,
            "previous_ocr_source": previous_ocr_backend,
            "ocr_paddle_model": effective_paddle_model,
            "previous_ocr_paddle_model": previous_paddle_model,
            "ocr_online": effective_online_descriptor or {},
            "previous_ocr_online": previous_online_descriptor,
            "ocr_actual_engine": (extraction_report or {}).get("ocr", {}).get("actual_engine"),
            "ocr_actual_model": (extraction_report or {}).get("ocr", {}).get("model"),
            "ocr_trigger": (extraction_report or {}).get("ocr_trigger"),
            "pdf_text_layer_quality": (extraction_report or {}).get(
                "pdf_text_layer_quality"),
            "max_batch_chars": max_batch_chars,
            "max_batch_items": max_batch_items,
            "paragraph_count": len(cleaned),
            "segment_count": len(segments),
        })

        warnings = state.setdefault("warnings", [])
        for warning in extraction_warnings:
            if warning not in warnings:
                warnings.append(warning)
            _append_runtime_technical_log(job_id, f"document extraction: {warning}")
        for warning in cleanup_warnings:
            if warning not in warnings:
                warnings.append(warning)
            _append_runtime_technical_log(job_id, f"source cleanup: {warning}")

        _reset_source_dependent_state(state)
        state.update({
            "filename": resolved_filename,
            "source_paras": list(raw_paragraphs),
            "paras": list(segments),
            "source_cleanup": cleanup_meta,
            "source_cleanup_done": True,
            "segmentation": segmentation_meta,
            "extraction_report": extraction_report,
            "p1_done": True,
            "source_page_count": 0,
        })
        # The rescan rebuilt the positional contract, so the deterministic
        # non-translatable admission runs here too (after the reset above).
        _auto_exclude_non_translatable_segments(job_id, state)
        state["source_quality_gate"] = _audit_source_quality_for_state(state)
        cleanup_meta["segment_count"] = len(state.get("paras") or [])
        try:
            with fitz.open(stream=source, filetype="pdf") as source_pdf:
                state["source_page_count"] = source_pdf.page_count
        except Exception as exc:  # noqa: BLE001 - extraction already succeeded
            _append_runtime_technical_log(job_id, f"source page count: {exc}")
        pipeline_config = state.get("pipeline_config")
        if not isinstance(pipeline_config, dict):
            pipeline_config = {}
            state["pipeline_config"] = pipeline_config
        pipeline_config["enable_source_cleanup"] = True
        # The rescan is the moment the OCR language can legitimately change; the
        # new configuration is persisted so a later resume reuses it.
        pipeline_config["ocr_language"] = effective_ocr_language or ""
        pipeline_config["ocr_language_engine_code"] = (
            resolved_ocr.engine_code if resolved_ocr else "")
        pipeline_config["ocr_language_label"] = (
            resolved_ocr.label if resolved_ocr else "")
        pipeline_config["ocr_text_direction"] = (
            resolved_ocr.direction if resolved_ocr else "")
        pipeline_config["ocr_dpi"] = effective_ocr_dpi
        pipeline_config["ocr_psm"] = effective_ocr_psm
        pipeline_config["ocr_quality"] = effective_ocr_quality
        pipeline_config["ocr_backend"] = effective_ocr_backend
        pipeline_config["ocr_mode"] = effective_ocr_mode
        pipeline_config["ocr_paddle_model"] = effective_paddle_model
        if effective_online_descriptor:
            pipeline_config["ocr_online"] = dict(effective_online_descriptor)
        else:
            pipeline_config.pop("ocr_online", None)
        state["ocr_mode"] = effective_ocr_mode
        state["ocr_language"] = effective_ocr_language or ""
        state["ocr_config"] = {
            "ocr_language": effective_ocr_language or "",
            "ocr_language_engine_code": resolved_ocr.engine_code if resolved_ocr else "",
            "ocr_language_label": resolved_ocr.label if resolved_ocr else "",
            "ocr_text_direction": resolved_ocr.direction if resolved_ocr else "",
            "ocr_dpi": effective_ocr_dpi,
            "ocr_psm": effective_ocr_psm,
            "ocr_quality": effective_ocr_quality,
            "ocr_backend": effective_ocr_backend,
            "ocr_mode": effective_ocr_mode,
            "ocr_paddle_model": effective_paddle_model,
            "ocr_online": dict(effective_online_descriptor or {}),
        }
        save_source(job_id, source)
        state["performance"] = profiler.persist()
        trace.finish_run(job_id=job_id, status="checkpointed")
        save_job_state(job_id, state)
        update_runtime_state(
            job_id, status="idle_incomplete", phase="idle_incomplete",
            phase_label="原文纠错已完成，等待继续翻译", operation="source_cleanup",
            operation_id="source_cleanup", operation_label="原文纠错已完成",
            error=None, progress=True, heartbeat=True,
            event=(f"原文重新扫描完成：{len(raw_paragraphs)} 项提取为 "
                   f"{len(cleaned)} 段、{len(segments)} 个翻译单元，修复 "
                   f"{len(cleanup_meta.get('changed_segments') or [])} 项"),
            event_name="source_rescan_completed", event_visibility="user",
            event_category="progress",
            worker={"owner_pid": None, "worker_id": None,
                    "lease_expires_at": None})
        return state
    except BaseException as exc:
        trace.finish_run(job_id=job_id, status="failed", failure_reason=str(exc))
        _write_runtime_technical_log(job_id, exc)
        provider_status = provider_error_status(exc)
        user_error_message = (
            provider_error_message(exc, "原文重新扫描失败")
            if provider_status["status"] != "unknown"
            else str(exc)[:500] or "原文重新扫描失败")
        update_runtime_state(
            job_id, status="failed", phase="failed", phase_label="原文重新扫描失败",
            error={"type": type(exc).__name__, "message": user_error_message,
                   "stage": "document_processing", "operation": "source_cleanup",
                   "timestamp": _utc_now_iso(),
                   "technical_log": RUNTIME_TECHNICAL_LOG},
            event=f"原文重新扫描失败：{user_error_message[:180]}",
            event_name="source_rescan_failed", event_visibility="user",
            event_category="error", progress=True, heartbeat=True,
            worker={"owner_pid": None, "worker_id": None,
                    "lease_expires_at": None})
        raise
    finally:
        profiler.persist()
        set_llm_base_url(old_base_url)



def delivery_snapshot_asset_path(job_id, version, name):
    """On-disk path of a delivered asset, for native open / reveal actions."""
    return _snapshots.asset_path(job_dir(job_id), version, name)



def workflow_outcomes(state):
    digests = state.get("section_digests") or []
    degraded = [item.get("unit_id") for item in digests if item.get("status") == "deterministic_fallback"]
    synopsis_degraded = (state.get("document_synopsis") or {}).get("status") == "deterministic_fallback"
    annotation_results = state.get("annotation_batches") or {}
    failed = [int(offset) for offset, result in annotation_results.items()
              if result.get("status") in {"request_failed", "parse_failed", "stale"}]
    legacy_failures = bool(state.get("annotations_failed_batches") and not annotation_results)
    return {
        "understanding": {"status": "degraded" if degraded or synopsis_degraded else
            "success" if state.get("understanding_done") else "not_run",
            "degraded_units": degraded, "synopsis_degraded": synopsis_degraded,
            "total_units": len(digests)},
        "annotation": {"status": "partial" if failed or legacy_failures else
            "success" if state.get("annotations_done") else "not_run",
            "failed_offsets": failed, "failed_batches": len(failed) if annotation_results else
            int(state.get("annotations_failed_batches") or 0),
            "legacy_unlocated": legacy_failures},
    }


@_serialized_job_mutation
def prepare_workflow_retry(job_id, part):
    if _job_execution_alive(job_id):
        raise RuntimeError("任务仍在执行，请等待后台处理结束")
    state = load_job_state(job_id)
    if state is None:
        raise ValueError("找不到任务")
    state.setdefault("workflow_retry_history", []).append({
        "part": part, "previous_outcomes": workflow_outcomes(state),
        "requested_at": _finalization.now_iso()})
    if part == "understanding":
        state["understanding_retry_requested"] = True
    elif part == "annotation":
        if workflow_outcomes(state)["annotation"]["legacy_unlocated"]:
            state["annotations_done_offset"] = 0
            state["annotations"] = {}
            state["annotation_batches"] = {}
            state["annotations_failed_batches"] = 0
        state["annotations_done"] = False
    else:
        raise ValueError("未知的补跑步骤")
    save_job_state(job_id, state)
    return state



# ================= 交付状态与人工处理记录 =================
def _translation_review_finding(state, finding_id):
    """Resolve either the Core finding ID or the existing review-card ID."""
    from folith import delivery as _delivery
    selector = str(finding_id or "")
    candidates = [item for item in state.get("findings") or []
                  if isinstance(item, dict) and (
                      str(item.get("finding_id") or "") == selector
                      or _delivery.finding_id(item) == selector)]
    return next((item for item in reversed(candidates)
                 if isinstance(item.get("segment_id"), int)
                 and (_translation_evidence.current_review_event(
                     state, item["segment_id"]) or {}).get(
                         "review_event_id") == item.get("review_event_id")),
                candidates[-1] if candidates else None)


def _promote_human_reviewed_segment(state, segment_id, actor):
    """Promote one currently clean human-confirmed pair through existing TM truth."""
    from folith import delivery as _delivery
    pairs = state.get("pairs") or []
    if not isinstance(segment_id, int) or not 0 <= segment_id < len(pairs):
        return
    if any(_delivery._segment_index(item) == segment_id
           for item in _delivery.unresolved_findings(state)):
        return
    pair = pairs[segment_id]
    deterministic = check_translation_batch(
        [str(pair.get("source") or "")], [str(pair.get("target") or "")],
        state.get("glossary") or [], state.get("target_lang") or "简体中文")
    if any(item.get("severity") in {"blocking", "actionable"}
           for item in deterministic):
        return
    pair["reviewed"] = True
    pair["review_status"] = "reviewed_human"
    pair["accepted_target"] = pair.get("target") or ""
    pair["target_provenance"] = "human_accepted"
    pair["human_accepted"] = True
    pair["accepted_by_human"] = actor
    if state.get("use_tm", True) and _tm_eligible(pair.get("source"), pair.get("target")):
        tm = load_tm(tm_project_id(state))
        tm_put(tm, str(pair.get("source") or ""), str(pair.get("target") or ""),
               state_target_lang(state))
        save_tm(tm, tm_project_id(state))


@_serialized_job_mutation
def decide_translation_review_finding(
    job_id, finding_id, decision, actor, *, actor_type, note="",
):
    """Apply one fail-closed Translation Core HumanDecision to persisted state."""
    from folith import delivery as _delivery
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    finding = _translation_review_finding(state, finding_id)
    if finding is None or not finding.get("finding_id"):
        raise ValueError("finding has no Translation Core identity")
    state, updated, record = _translation_evidence.record_runtime_human_decision(
        state, finding["finding_id"], decision, actor, actor_type=actor_type,
        note=note)
    segment_id = updated.get("segment_id")
    if decision in {"accept_resolution", "dismiss"}:
        _promote_human_reviewed_segment(state, segment_id, actor)
    _recount_reviewed_segments(state)
    queue = _delivery.review_queue_findings(state)
    stats = state.setdefault("review_stats", {})
    for severity in ("blocking", "actionable", "informational"):
        stats[severity] = sum(item.get("severity") == severity for item in queue)
    state["has_blocking"] = bool(_delivery.unresolved_blocking(state))
    state["delivery_status"] = _delivery.compute_delivery_status(state)
    save_job_state(job_id, state)
    return state, updated, record


def _targeted_final_review_ids(state, limit=8):
    """Select a small, stable set of segments for the whole-book recheck.

    Batch review already covers ordinary semantic issues.  The final pass is
    reserved for risks that only become visible after all pairs are present:
    entity drift, locked-term conflicts, and low-authority knowledge conflicts.
    Resolved findings are deliberately excluded so a human decision is not
    reopened by a later resume.
    """
    pairs = state.get("pairs") or []
    if not pairs:
        return []
    blocked_source_indexes = {
        item.get("segment_index")
        for item in _audit_source_quality_for_state(state).get("flags") or []
        if isinstance(item, dict) and isinstance(item.get("segment_index"), int)
    }
    severity_rank = {"blocking": 0, "actionable": 1, "informational": 2}
    selected = {}
    for finding in state.get("findings") or []:
        if not isinstance(finding, dict) or finding.get("resolved"):
            continue
        finding_type = str(finding.get("type") or "").strip().lower()
        category = str(finding.get("category") or "").strip().lower()
        is_global_consistency_risk = (
            finding_type in {"knowledge_conflict", "entity_conflict"}
            or (finding_type == "glossary" and finding.get("conflict"))
            or (finding.get("conflict") and category == "terminology_consistency")
        )
        if not is_global_consistency_risk:
            continue
        raw_index = finding.get("segment_id", finding.get("segment_index"))
        try:
            segment_id = int(raw_index)
        except (TypeError, ValueError):
            continue
        if (isinstance(raw_index, bool) or not 0 <= segment_id < len(pairs)
                or segment_id in blocked_source_indexes):
            continue
        score = severity_rank.get(str(finding.get("severity") or ""), 3)
        previous = selected.get(segment_id)
        if previous is None or score < previous:
            selected[segment_id] = score
    bounded_limit = max(0, int(limit or 0))
    return [segment_id for segment_id, _score in sorted(
        selected.items(), key=lambda item: (item[1], item[0]))[:bounded_limit]]


@_serialized_job_mutation
def review_translation_segments(
    job_id, indexes, provider, api_key, model, target_lang, *, style_rules="",
    call_llm_fn=None, base_url=None, review_focus="", focus_findings=None,
):
    """Run the existing independent reviewer on current persisted segments.

    ``review_focus`` and ``focus_findings`` are optional context for the
    post-translation whole-book pass.  They do not change the normal manual
    refresh path and keep the global conflict visible to the reviewer.
    """
    from folith import delivery as _delivery
    from folith.terminology import (
        glossary_block as _glossary_block,
        select_glossary_for_segments as _select_glossary,
    )
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    pairs = state.get("pairs") or []
    segment_ids = sorted({int(index) for index in indexes
                          if str(index).lstrip("-").isdigit()
                          and 0 <= int(index) < len(pairs)})
    if not segment_ids:
        return state, {"reviewed_segment_ids": [], "failed_segment_ids": []}
    blocked_source_indexes = {
        item.get("segment_index")
        for item in _audit_source_quality_for_state(state).get("flags") or []
        if isinstance(item, dict) and isinstance(item.get("segment_index"), int)
    }
    blocked_requested = [index for index in segment_ids
                         if index in blocked_source_indexes]
    if blocked_requested:
        numbers = "、".join(str(index + 1) for index in blocked_requested)
        raise ValueError(f"第 {numbers} 段原文尚未核对，不能发送给审校模型")
    if base_url is None:
        pipeline_config = state.get("pipeline_config") or {}
        saved_reviewer = (state.get("reviewer_config")
                          or pipeline_config.get("reviewer") or {})
        saved_translator = (state.get("translator_config")
                            or pipeline_config.get("translator") or {})
        if str(saved_reviewer.get("provider") or "") == str(provider or ""):
            base_url = saved_reviewer.get("base_url") or ""
        elif str(saved_translator.get("provider") or "") == str(provider or ""):
            base_url = saved_translator.get("base_url") or ""
    review_call = call_llm_fn or _model_roles.make_role_call(
        call_llm, {"provider": provider, "api_key": api_key,
                   "model": model, "base_url": base_url})
    glossary = normalize_glossary(
        state.get("glossary") or (state.get("glossary_frozen") or {}).get(
            "entries") or [])
    reviewed, failed_ids = [], []
    tm = load_tm(tm_project_id(state)) if state.get("use_tm", True) else {}
    tm_changed = False
    for segment_id in segment_ids:
        pair = pairs[segment_id]
        source = str(pair.get("source") or "")
        target = str(pair.get("target") or "")
        section_profile = _batch_section_profile(
            state.get("document_profile"), segment_id, 1)
        selected, _ = _select_glossary(
            [source], glossary + _knowledge.provisional_hints(
                state.get("knowledge_candidates") or [],
                authoritative_entries=glossary),
            state.get("document_profile"), section_profile)
        glossary_text = _glossary_block(selected)
        deterministic = _globalize_batch_findings(check_translation_batch(
            [source], [target], glossary, target_lang,
            section_profile=section_profile), segment_id)
        review_context = _runtime_review_context(
            state, segment_id, 1, glossary_text, style_rules, target_lang)
        if review_focus:
            risk_items = [
                {
                    key: item.get(key)
                    for key in (
                        "segment_id", "type", "category", "severity", "reason",
                        "source", "preferred_target", "observed_target",
                        "observed_targets",
                    ) if item.get(key) is not None
                } for item in (focus_findings or [])
                if isinstance(item, dict)
                and item.get("segment_id", item.get("segment_index")) == segment_id
            ][:4]
            review_context["review_focus"] = str(review_focus)
            review_context["whole_book_consistency_risks"] = risk_items
        packet = _translation_evidence.build_runtime_review_packet(
            state, [pair], [segment_id], glossary,
            deterministic_checks=deterministic, review_context=review_context)
        evidence_index = _translation_evidence.TranslationEvidenceIndex(
            state.get("paras") or [item.get("source") or "" for item in pairs],
            pairs, glossary, state.get("document_profile"),
            state.get("document_synopsis"), state.get("section_digests"),
            state.get("findings"))
        findings, failed, trace = \
            _translation_evidence.review_translation_batch_with_evidence(
                [source], [target], glossary_text, style_rules, target_lang,
                provider, api_key, model, evidence_index,
                call_llm=review_call, segment_ids=[segment_id],
                translation_core_packet=packet)
        event_id = (
            f"translation-review-{job_id}-refresh-{segment_id}-"
            f"{len(state.get('review_evidence') or [])}")
        records = [_review_finding_record(item, event_id) for item in findings]
        state.setdefault("findings", []).extend(records)
        _translation_evidence.register_runtime_review_event(
            state, trace, records, event_id, [segment_id])
        focus_risk_present = bool(review_focus and any(
            isinstance(item, dict)
            and item.get("segment_id", item.get("segment_index")) == segment_id
            and not item.get("resolved")
            for item in (focus_findings or [])))
        if failed:
            failed_ids.append(segment_id)
            pair["reviewed"] = False
            pair["review_status"] = "review_failed"
        elif not focus_risk_present and not any(
                item.get("severity") in {"blocking", "actionable"}
                for item in [*deterministic, *records]):
            reviewed.append(segment_id)
            pair["reviewed"] = True
            pair["review_status"] = "reviewed_clean"
            pair["accepted_target"] = target
            pair["target_provenance"] = "reviewed"
            if state.get("use_tm", True) and _tm_eligible(source, target):
                if tm_put(tm, source, target, state_target_lang(state)):
                    tm_changed = True
        else:
            pair["reviewed"] = False
            pair["review_status"] = "reviewed_with_findings"
        _checkpoint.append_event(job_dir(job_id), {
            "phase": "semantic_review_done", "segment_ids": [segment_id],
            "refresh": True, "failed": failed,
            **({"review_focus": str(review_focus)} if review_focus else {}),
        })
    if tm_changed:
        save_tm(tm, tm_project_id(state))
    _recount_reviewed_segments(state)
    queue = _delivery.review_queue_findings(state)
    stats = state.setdefault("review_stats", {})
    for severity in ("blocking", "actionable", "informational"):
        stats[severity] = sum(item.get("severity") == severity for item in queue)
    stats["review_failed"] = sum(
        item.get("decision") == "failed" and _translation_evidence._current_translation_review(
            item) for item in state.get("review_evidence") or [] if isinstance(item, dict))
    state["has_blocking"] = bool(_delivery.unresolved_blocking(state))
    state["delivery_status"] = _delivery.compute_delivery_status(state)
    save_job_state(job_id, state)
    return state, {
        "reviewed_segment_ids": reviewed,
        "failed_segment_ids": failed_ids,
    }


def _run_targeted_final_review(
    job_id, state, reviewer_config, target_lang, style_rules, *,
    enable_review=True, on_status=None,
):
    """Run the bounded whole-book consistency pass once per translation.

    The existing segment reviewer is reused so evidence, review events and
    delivery gates remain on one path.  The usage ledger lives on the caller's
    state object; the review helper reloads state from disk, so it is merged
    back explicitly after the call to avoid losing the newly recorded usage.
    """
    record = state.get("targeted_final_review") or {}
    if record.get("status") in {"completed", "skipped"}:
        return state
    if not enable_review:
        state["targeted_final_review"] = {
            "status": "skipped", "reason": "review_disabled",
            "segment_ids": [], "reviewed_segment_ids": [],
            "failed_segment_ids": [],
        }
        save_job_state(job_id, state)
        return state
    segment_ids = _targeted_final_review_ids(state)
    if not segment_ids:
        state["targeted_final_review"] = {
            "status": "skipped", "reason": "no_cross_book_risk",
            "segment_ids": [], "reviewed_segment_ids": [],
            "failed_segment_ids": [],
        }
        save_job_state(job_id, state)
        return state
    if on_status:
        on_status(f"【阶段二终检】全书一致性定向复核（{len(segment_ids)} 段）...")
    focus_findings = []
    selected = set(segment_ids)
    for finding in state.get("findings") or []:
        if not isinstance(finding, dict) or finding.get("resolved"):
            continue
        raw_index = finding.get("segment_id", finding.get("segment_index"))
        try:
            segment_id = int(raw_index)
        except (TypeError, ValueError):
            continue
        if isinstance(raw_index, bool) or segment_id not in selected:
            continue
        item = {
            key: finding.get(key)
            for key in (
                "segment_id", "type", "category", "severity", "reason",
                "source", "preferred_target", "observed_target", "observed_targets",
            ) if finding.get(key) is not None
        }
        item["segment_id"] = segment_id
        focus_findings.append(item)
    state.setdefault("llm_usage", _usage.empty_usage())
    final_review_call = _model_roles.make_role_call(
        _usage.tracked_call(
            call_llm, state["llm_usage"], role="targeted_final_review"),
        reviewer_config)
    try:
        reviewed_state, result = review_translation_segments(
            job_id, segment_ids, reviewer_config["provider"],
            reviewer_config["api_key"], reviewer_config["model"], target_lang,
            style_rules=style_rules, call_llm_fn=final_review_call,
            base_url=reviewer_config.get("base_url"),
            review_focus="whole_book_consistency",
            focus_findings=focus_findings,
        )
        # review_translation_segments reloads and saves its own state.  Carry
        # over the in-memory ledger mutated by the tracked call wrapper.
        reviewed_state["llm_usage"] = state["llm_usage"]
        reviewed_ids = list(result.get("reviewed_segment_ids") or [])
        failed_ids = list(result.get("failed_segment_ids") or [])
        final_status = "failed" if failed_ids else "completed"
        reviewed_state["targeted_final_review"] = {
            "status": final_status,
            "segment_ids": list(segment_ids),
            "reviewed_segment_ids": reviewed_ids,
            "failed_segment_ids": failed_ids,
        }
        if failed_ids:
            warning = f"全书一致性定向复核有 {len(failed_ids)} 段未完成"
            warnings = reviewed_state.setdefault("warnings", [])
            if warning not in warnings:
                warnings.append(warning)
        save_job_state(job_id, reviewed_state)
        return reviewed_state
    except Exception as exc:
        state["targeted_final_review"] = {
            "status": "failed", "segment_ids": list(segment_ids),
            "reviewed_segment_ids": [], "failed_segment_ids": list(segment_ids),
            "error": str(exc)[:240],
        }
        warning = f"全书一致性定向复核失败：{str(exc)[:200]}"
        warnings = state.setdefault("warnings", [])
        if warning not in warnings:
            warnings.append(warning)
        save_job_state(job_id, state)
        return state


@_serialized_job_mutation
def mark_findings_resolved(job_id, finding_ids, action, note="", actor="user"):
    """Legacy action surface; Core findings route through HumanDecision."""
    from folith import delivery as _delivery
    state = load_job_state(job_id)
    if state is None:
        return None
    legacy_ids = []
    mapped = {"human_fixed": "accept_resolution", "preserved": "dismiss"}
    for selector in finding_ids or []:
        finding = _translation_review_finding(state, selector)
        if finding is not None and finding.get("finding_id") \
                and finding.get("input_fingerprint"):
            if action not in mapped:
                raise ValueError(f"Translation Core finding 不支持旧 action：{action}")
            state, updated, _record = \
                _translation_evidence.record_runtime_human_decision(
                    state, finding["finding_id"], mapped[action], actor,
                    actor_type="human", note=note)
            if mapped[action] in {"accept_resolution", "dismiss"}:
                _promote_human_reviewed_segment(
                    state, updated.get("segment_id"), actor)
        else:
            legacy_ids.append(selector)
    if legacy_ids:
        state, _marked = _delivery.mark_findings(
            state, legacy_ids, action, note, actor)
    _recount_reviewed_segments(state)
    state["has_blocking"] = bool(_delivery.unresolved_blocking(state))
    state["delivery_status"] = _delivery.compute_delivery_status(state)
    save_job_state(job_id, state)
    return state



def retranslate_segments(job_id, indexes, provider, api_key, model, target_lang,
                         style_rules="", glossary=None, on_status=None,
                         on_caption=None, actor="user", reviewer_provider=None,
                         reviewer_api_key=None, reviewer_model=None,
                         reviewer_base_url=None, translator_base_url=None):
    """定点重译（抽取自 scripts/fix_segments.py 的能力）。

    ``translator_base_url`` 与管线同名同义：自定义中转站必须显式传入，否则
    在非管线线程里调用会解析不到地址（此前会静默失败）。
    """
    from folith import delivery as _delivery
    old_base_url = getattr(_LLM_CTX, "base_url", None)
    if translator_base_url is not None:
        set_llm_base_url(translator_base_url)
    try:
        state, fixed = _delivery.retranslate_segments(
            job_id, indexes, provider, api_key, model, target_lang,
            style_rules, glossary, on_status, on_caption, actor,
            base_url=translator_base_url or "")
    finally:
        if translator_base_url is not None:
            set_llm_base_url(old_base_url)
    if fixed and state.get("translation_core_review_required"):
        review_provider = provider if reviewer_provider is None else reviewer_provider
        review_api_key = api_key if reviewer_api_key is None else reviewer_api_key
        review_model = model if reviewer_model is None else reviewer_model
        state, _ = review_translation_segments(
            job_id, fixed, review_provider, review_api_key, review_model,
            target_lang, style_rules=style_rules, base_url=reviewer_base_url)
    return state, fixed



@_serialized_job_mutation
def invalidate_academic_report(job_id, scope="all", section_id=None):
    """Invalidate academic artifacts only; translation stages remain intact."""
    from folith import academic_writer
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    academic_writer.invalidate_academic_state(state, scope, section_id)
    save_job_state(job_id, state)
    return state


def load_academic_artifact(job_id, name):
    """Read a canonical academic JSON artifact for UI/CLI inspection."""
    from folith import academic_writer
    ARTIFACT_FILES = academic_writer.ARTIFACT_FILES
    if name not in ARTIFACT_FILES:
        raise ValueError(f"未知学术 artifact：{name}")
    path = job_dir(job_id) / ARTIFACT_FILES[name]
    if not path.is_file():
        return None
    try:
        return academic_writer._read_artifact(path)
    except Exception:
        return None


@_serialized_job_mutation
def record_human_evidence(job_id, question_id, answer, interface="academic_workspace"):
    """Record a human author answer for an open evidence question.

    The answer is stored verbatim with provenance; the question is marked
    answered; the case-analysis/section staleness is propagated through the
    existing dependency-hash architecture (only affected sections rewrite).
    """
    from folith import academic_writer, human_evidence
    state = load_job_state(job_id)
    if state is None:
        raise ValueError(f"找不到任务 {job_id}")
    questions = academic_writer._read_artifact(
        job_dir(job_id) / academic_writer.ARTIFACT_FILES["human_evidence_questions"])
    if not questions:
        raise ValueError("当前没有待回答的人类证据问题；请先生成学术报告。")
    evidence = academic_writer._read_artifact(
        job_dir(job_id) / academic_writer.ARTIFACT_FILES["evidence"])
    entry, updated_questions = human_evidence.record_human_answer(
        questions, question_id, answer, evidence or {}, interface,
        existing=state.get("human_evidence") or [])
    entries = list(state.get("human_evidence") or [])
    entries = [x for x in entries if x.get("question_id") != question_id]
    entries.append(entry)
    state["human_evidence"] = entries
    academic_writer._write_artifact(
        job_dir(job_id) / academic_writer.ARTIFACT_FILES["human_evidence_questions"],
        updated_questions)
    record = state.get("academic_state", {}).get("artifacts", {}).get(
        "human_evidence_questions")
    if record:
        record["content_hash"] = updated_questions["content_hash"]
        record["updated_at"] = academic_writer._now()
    save_job_state(job_id, state)
    return entry


# ================= 阶段三：证据约束型学术写作 =================
def generate_mti_report(bilingual_pairs, termbase_dict, theory, provider, api_key,
                        model, state, job_id, on_status=None,
                        research_settings=None, literature_sources=None):
    """Compatibility wrapper for the evidence-grounded academic pipeline.

    ``bilingual_pairs`` and ``termbase_dict`` remain in the signature for old
    callers; the canonical inputs are now the saved translation state and the
    durable academic artifacts beside ``state.json``.
    """
    from folith import academic_writer
    report_md = academic_writer.run_academic_pipeline(
        state, job_id, theory, provider, api_key, model,
        artifact_dir=job_dir(job_id), call_llm=call_llm,
        save_state=lambda current: save_job_state(job_id, current),
        research_settings=research_settings, literature_sources=literature_sources,
        on_status=on_status,
    )
    save_compliance_record(job_id, load_job_state(job_id) or state)
    return report_md


# ================= 主流程：单文档完整流水线 =================
@_exclusive_job_execution(0)
def run_job_pipeline(job_id, filename, file_bytes, *, provider, api_key, model,
                     target_lang, auto_term, enable_report, translation_theory,
                     user_glossary=None, style_rules="", enable_review=True,
                     enable_annotate=True, use_tm=True,
                     strict_terminology_governance=False, mode=None,
                     research_settings=None, literature_sources=None,
                     delivery_config=None, batch_size=None, max_batch_chars=None,
                     enable_understanding=None, reviewer_provider=None,
                     reviewer_api_key=None, reviewer_model=None,
                     reviewer_base_url=None, translator_base_url=None,
                     auxiliary_provider=None, auxiliary_api_key=None,
                     auxiliary_model=None, auxiliary_base_url=None,
                     knowledge_feedback_interval=None,
                     enable_source_cleanup=None,
                     reasoning_effort=None, source_cleanup_concurrency=None,
                     ocr_workers=None, ocr_queue_size=None,
                     source_cleanup_confidence_threshold=None,
                     source_cleanup_max_retries=None,
                     source_cleanup_request_interval=None,
                     segmentation_mode=None,
                     translation_concurrency=None,
                     ocr_language=None, ocr_dpi=None, ocr_psm=None,
                     ocr_quality=None, ocr_mode=None, ocr_backend=None,
                     ocr_paddle_model=None, ocr_online=None,
                     on_status=None, on_caption=None):
    """执行单个文档的完整流程；每个里程碑实时落盘，刷新/重启后均可继续。

    strict_terminology_governance=True：翻译前建立文档画像，并要求自动候选
    术语完成审核/冻结。导入的锁定术语视为已固定，不会阻塞翻译。
    ``mode`` 仅保留给旧调用方；quality 映射到严格术语治理，quick 映射到关闭。
    ``auxiliary_*`` 可把画像、摘要、术语抽取与连续性观察路由到较低成本模型；
    不提供时沿用正文模型。``knowledge_feedback_interval`` 仅控制普通模式的
    批后连续性观察频率，严格审校模式仍逐批观察。
    """
    _ensure_output_dir()
    if translator_base_url is not None:
        set_llm_base_url(translator_base_url)
    set_llm_reasoning_effort(reasoning_effort)
    base = new_job_state(filename)
    state = load_job_state(job_id) or base
    state = {**base, **state}  # 兼容旧版本状态缺字段
    state = _state_migration.migrate_state(state)
    effective_segmentation_mode = _source_segmentation_mode(
        state, segmentation_mode)
    previous_performance = state.get("performance")
    saved_runtime_config = state.get("pipeline_config") or {}
    effective_ocr_language, effective_ocr_dpi, effective_ocr_psm, \
        effective_ocr_quality, resolved_ocr = _effective_ocr_config(
            ocr_language, ocr_dpi, ocr_psm, saved_runtime_config, ocr_quality)
    effective_ocr_backend, effective_paddle_model, effective_online_descriptor, \
        effective_online_config = _effective_ocr_source(
            ocr_backend, saved_runtime_config, ocr_paddle_model=ocr_paddle_model,
            ocr_online=ocr_online, prefer_saved=bool(saved_runtime_config))
    if (effective_ocr_backend == "online" and not effective_online_descriptor
            and not load_ocr_online_config()):
        try:
            visual_base = _resolve_online_ocr_base_url(
                provider, translator_base_url or getattr(_LLM_CTX, "base_url", None) or "")
            effective_online_config = _ocr_online.OnlineOCRConfig(
                provider=provider, model=model, api_key=api_key, base_url=visual_base,
                credential_source="translation_profile",
                reasoning_mode=("thinking_disabled" if provider == "OpenCode Go"
                                and model == "deepseek-v4.1-flash" else "provider_default"))
            effective_online_descriptor = effective_online_config.describe()
        except ValueError:
            pass  # A native document does not require a valid visual endpoint.
    effective_ocr_mode = normalize_ocr_mode(
        ocr_mode if ocr_mode is not None else saved_runtime_config.get("ocr_mode"))
    effective_source_language = str(
        state.get("source_lang")
        or (resolved_ocr.label if resolved_ocr else "") or "")
    effective_cleanup_concurrency = _runtime_int_option(
        source_cleanup_concurrency,
        saved_runtime_config.get("source_cleanup_concurrency"),
        "FOLIOTHREAD_LLM_CLEANUP_CONCURRENCY",
        _source_cleanup.DEFAULT_PARALLELISM, maximum=16)
    effective_ocr_workers = _runtime_int_option(
        ocr_workers, saved_runtime_config.get("ocr_workers"),
        "FOLIOTHREAD_OCR_WORKERS", OCR_WORKERS_DEFAULT, maximum=8)
    effective_ocr_queue_size = _runtime_int_option(
        ocr_queue_size, saved_runtime_config.get("ocr_queue_size"),
        "FOLIOTHREAD_OCR_QUEUE_SIZE", effective_ocr_workers * 2, maximum=32)
    effective_cleanup_confidence_threshold = _runtime_float_option(
        source_cleanup_confidence_threshold,
        saved_runtime_config.get("source_cleanup_confidence_threshold"),
        "FOLIOTHREAD_OCR_CONFIDENCE_THRESHOLD",
        OCR_CONFIDENCE_THRESHOLD_DEFAULT, maximum=100.0)
    effective_cleanup_max_retries = _runtime_int_option(
        source_cleanup_max_retries,
        saved_runtime_config.get("source_cleanup_max_retries"),
        "FOLIOTHREAD_LLM_CLEANUP_MAX_RETRIES",
        _source_cleanup.DEFAULT_MAX_RETRIES, minimum=0, maximum=6)
    effective_cleanup_request_interval = _runtime_float_option(
        source_cleanup_request_interval,
        saved_runtime_config.get("source_cleanup_request_interval"),
        "FOLIOTHREAD_LLM_CLEANUP_REQUEST_INTERVAL", 0.0, maximum=60.0)
    active_run = trace.start_run(
        job_id=job_id,
        filename=filename,
        file_size_bytes=len(file_bytes) if file_bytes else 0,
        target_lang=target_lang,
        output_dir=job_dir(job_id),
    )
    profiler = PipelineProfiler(
        job_dir(job_id) / "performance.json", run_id=active_run.run_id
    )

    def persist_performance(final_status=None):
        # Resuming a task after ingestion must not replace a useful prior
        # timing artifact with an empty run merely because later translation
        # stages do not touch the PDF profiler.
        if (isinstance(previous_performance, dict)
                and previous_performance.get("stages")
                and not profiler.snapshot().get("stages")):
            return previous_performance
        prof_res = profiler.persist()
        if final_status:
            trace.finish_run(job_id=job_id, status=final_status)
        return prof_res

    # An explicit UI-triggered source rescan is a checkpoint-only operation:
    # it must not fall through into profiling, terminology, or translation in
    # the same worker turn.  The request is persisted before the worker starts
    # so a browser refresh cannot lose the intent.
    if state.get("source_rescan_requested"):
        options = state.pop("source_rescan_options", {})
        state.pop("source_rescan_requested", None)
        save_job_state(job_id, state)
        return rescan_pdf_source(
            job_id, filename=filename, provider=provider, api_key=api_key,
            model=model, base_url=translator_base_url, on_status=on_status,
            max_batch_chars=(options.get("max_batch_chars")
                             if isinstance(options, dict) else None),
            max_batch_items=(options.get("max_batch_items")
                             if isinstance(options, dict) else None),
            parallelism=(options.get("parallelism", effective_cleanup_concurrency)
                         if isinstance(options, dict) else effective_cleanup_concurrency),
            ocr_workers=(options.get("ocr_workers") if isinstance(options, dict) else None),
            ocr_queue_size=(options.get("ocr_queue_size") if isinstance(options, dict) else None),
            confidence_threshold=(options.get("confidence_threshold")
                                  if isinstance(options, dict) else None),
            max_retries=(options.get("max_retries") if isinstance(options, dict) else None),
            request_interval_seconds=(options.get("request_interval_seconds")
                                      if isinstance(options, dict) else None),
            segmentation_mode=(options.get("segmentation_mode", segmentation_mode)
                               if isinstance(options, dict) else segmentation_mode),
            ocr_language=(options.get("ocr_language") if isinstance(options, dict) else None),
            ocr_dpi=(options.get("ocr_dpi") if isinstance(options, dict) else None),
            ocr_psm=(options.get("ocr_psm") if isinstance(options, dict) else None),
            ocr_quality=(options.get("ocr_quality") if isinstance(options, dict) else None),
            ocr_mode=(options.get("ocr_mode") if isinstance(options, dict) else None),
            ocr_backend=(options.get("ocr_backend") if isinstance(options, dict) else None),
            ocr_paddle_model=(options.get("ocr_paddle_model")
                              if isinstance(options, dict) else None),
            ocr_online=(options.get("ocr_online") if isinstance(options, dict) else None),
            _allow_active=True)

    # ---- 项目记忆注入（蓝图 §3.2 的跨任务复用）----
    # 必须在 state/pipeline_config 定型之前完成，否则 state["style_rules"] 与
    # 运行时实际使用的风格会不一致（存储值与实际值分叉）。
    # 只对尚未开始工作的任务注入：进行中的任务不能因为项目记忆变化而改术语。
    project_entries: list = []
    if not state.get("glossary") and not str(state.get("glossary_frozen") or ""):
        project_entries, project_style, project_meta = _project_memory_injection(state)
        if project_entries or project_style:
            known = {str(e.get("source") or "").casefold()
                     for e in normalize_glossary(list(user_glossary or []))}
            added = [e for e in project_entries
                     if str(e.get("source") or "").casefold() not in known]
            user_glossary = list(user_glossary or []) + added
            if project_style:
                style_rules = "；".join(
                    part for part in (project_style, str(style_rules or "").strip())
                    if part)
            state["project_memory"] = {
                **project_meta,
                "injected_entry_ids": [e["id"] for e in added],
                "injected_at": _project_now_iso(),
            }

    previous_target_lang = str(state.get("target_lang") or "")
    previous_style_rules = str(state.get("style_rules") or "")
    review_required = bool(state.get("translation_core_review_required")) \
        if state.get("p2_done") else bool(enable_review)
    saved_understanding = (state.get("pipeline_config") or {}).get(
        "enable_understanding")
    saved_source_cleanup = (state.get("pipeline_config") or {}).get(
        "enable_source_cleanup")
    is_pdf_for_config = str(filename or state.get("filename") or "").lower().endswith(".pdf")
    if mode is not None:
        strict_terminology_governance = mode == "quality"
    primary_base_url = (translator_base_url
                        or getattr(_LLM_CTX, "base_url", None) or "")
    translator_config = _model_roles.normalize_role_config(
        None, fallback_provider=provider, fallback_model=model,
        fallback_api_key=api_key, fallback_base_url=primary_base_url)
    resolved_reviewer_provider = reviewer_provider or provider
    reviewer_config = _model_roles.normalize_role_config(
        {
            "provider": reviewer_provider,
            "model": reviewer_model,
            "api_key": reviewer_api_key,
            "base_url": reviewer_base_url,
        },
        fallback_provider=provider, fallback_model=model,
        fallback_api_key=api_key,
        fallback_base_url=(primary_base_url
                           if resolved_reviewer_provider == provider else ""))
    saved_auxiliary = state.get("auxiliary_config") or {}
    resolved_auxiliary_provider = (auxiliary_provider
                                   or saved_auxiliary.get("provider") or provider)
    auxiliary_config = _model_roles.normalize_role_config(
        {
            "provider": auxiliary_provider or saved_auxiliary.get("provider"),
            "model": auxiliary_model or saved_auxiliary.get("model"),
            "api_key": auxiliary_api_key,
            "base_url": auxiliary_base_url or saved_auxiliary.get("base_url"),
        },
        fallback_provider=provider, fallback_model=model,
        fallback_api_key=api_key,
        fallback_base_url=(primary_base_url
                           if resolved_auxiliary_provider == provider else ""))
    understanding_runtime_config = _model_roles.public_role_config(auxiliary_config)
    has_fallback_digests = any(
        isinstance(item, dict) and item.get("status") == "deterministic_fallback"
        for item in state.get("section_digests") or [])
    retry_failed_understanding = bool(
        state.get("understanding_retry_requested") or (state.get("understanding_done")
        and has_fallback_digests
        and state.get("understanding_runtime_config") != understanding_runtime_config))
    pipeline_config = {
        "target_lang": target_lang,
        "auto_term": bool(auto_term),
        "enable_report": bool(enable_report),
        "translation_theory": translation_theory,
        "style_rules": style_rules,
        "enable_review": bool(enable_review),
        "enable_annotate": bool(enable_annotate),
        "use_tm": bool(use_tm),
        "strict_terminology_governance": bool(strict_terminology_governance),
        # 批次参数随任务保存：恢复任务必须沿用原配置，否则同一任务的批次划分
        # 会在中途变化。
        "batch_size": int(batch_size or BATCH_SIZE),
        "max_batch_chars": int(max_batch_chars or TRANSLATION_MAX_BATCH_CHARS),
        "enable_understanding": enable_understanding,
        "reasoning_effort": (str(reasoning_effort or "").strip().lower()
                              if str(reasoning_effort or "").strip().lower()
                              in _REASONING_EFFORT_VALUES else ""),
        "translator": _model_roles.public_role_config(translator_config),
        "reviewer": _model_roles.public_role_config(reviewer_config),
    }
    # Auxiliary work (profile, synopsis, glossary observation) can use a
    # cheaper model without changing the authoritative translator/reviewer.
    # Keep this outside the historical pipeline_config shape unless explicitly
    # configured, so old task snapshots remain byte-for-byte compatible.
    auxiliary_explicit = any(
        value is not None and str(value).strip()
        for value in (auxiliary_provider, auxiliary_model, auxiliary_base_url)
    )
    if auxiliary_explicit or saved_auxiliary:
        state["auxiliary_config"] = _model_roles.public_role_config(auxiliary_config)
    if knowledge_feedback_interval is not None:
        try:
            normalized_feedback_interval = max(1, int(knowledge_feedback_interval))
        except (TypeError, ValueError):
            normalized_feedback_interval = 1
        pipeline_config["knowledge_feedback_interval"] = normalized_feedback_interval
    if translation_concurrency is not None or "translation_concurrency" in saved_runtime_config:
        pipeline_config["translation_concurrency"] = _runtime_int_option(
            translation_concurrency,
            saved_runtime_config.get("translation_concurrency"),
            "FOLIOTHREAD_TRANSLATION_CONCURRENCY",
            DEFAULT_TRANSLATION_CONCURRENCY,
            minimum=1,
            maximum=16,
        )
    if segmentation_mode is not None or "segmentation_mode" in saved_runtime_config:
        pipeline_config["segmentation_mode"] = effective_segmentation_mode
    state["pipeline_config"] = pipeline_config
    state.setdefault("llm_usage", _usage.empty_usage())
    auxiliary_call = _model_roles.make_role_call(
        _usage.tracked_call(call_llm, state["llm_usage"], role="auxiliary"),
        auxiliary_config, job_id=job_id)
    state.update(
        target_lang=target_lang, auto_term_enabled=bool(auto_term),
        report_enabled=bool(enable_report), theory=translation_theory,
        style_rules=style_rules, enable_review=bool(enable_review),
        enable_annotate=bool(enable_annotate), use_tm=bool(use_tm),
        provider=provider, model=model,
        translation_core_review_required=review_required,
        translator_config=_model_roles.public_role_config(translator_config),
        reviewer_config=_model_roles.public_role_config(reviewer_config),
    )
    if state.get("p2_done") and review_required and (
            previous_target_lang != str(target_lang or "")
            or previous_style_rules != str(style_rules or "")):
        _invalidate_translation_reviews(
            state, list(range(len(state.get("pairs") or []))),
            "target language or style review context changed")
        _invalidate_final_delivery_state(state)
    if delivery_config is not None:
        state["delivery_config"] = normalize_delivery_config(
            delivery_config, enable_report=enable_report,
            enable_annotate=enable_annotate)
    elif state.get("delivery_config"):
        state["delivery_config"] = normalize_delivery_config(
            state["delivery_config"], enable_report=enable_report,
            enable_annotate=enable_annotate)
    if enable_understanding is None:
        if isinstance(saved_understanding, bool):
            # A resumed task keeps the decision it was created with.  This is
            # important for old Quick jobs whose legacy state has no new flag.
            enable_understanding = saved_understanding
        elif state.get("p1_done") or state.get("p2_done"):
            # Legacy in-progress jobs had no understanding pass outside strict
            # terminology governance; do not unexpectedly add API cost on resume.
            enable_understanding = bool(
                state.get("profile_done") or state.get("understanding_done")
                or state.get("quality_mode"))
        else:
            # New tasks: Quick skips it; Standard and Academic keep it enabled.
            enable_understanding = mode != "quick"
    retry_failed_understanding = bool(
        state.get("understanding_retry_requested") or (enable_understanding and state.get("understanding_done")
        and has_fallback_digests
        and state.get("understanding_runtime_config") != understanding_runtime_config))
    if enable_source_cleanup is None:
        if isinstance(saved_source_cleanup, bool):
            enable_source_cleanup = saved_source_cleanup
        else:
            enable_source_cleanup = str(filename or "").lower().endswith(".pdf")
    pipeline_config["enable_understanding"] = bool(enable_understanding)
    # Source cleanup is automatically enabled for PDF jobs.  Persist the key
    # only for PDF tasks (or an explicit opt-in) so legacy DOCX pipeline
    # configurations retain their stable shape.
    if is_pdf_for_config or enable_source_cleanup:
        pipeline_config["enable_source_cleanup"] = bool(enable_source_cleanup)
        pipeline_config.update({
            "source_cleanup_concurrency": effective_cleanup_concurrency,
            "ocr_workers": effective_ocr_workers,
            "ocr_queue_size": effective_ocr_queue_size,
            "source_cleanup_confidence_threshold": effective_cleanup_confidence_threshold,
            "source_cleanup_max_retries": effective_cleanup_max_retries,
            "source_cleanup_request_interval": effective_cleanup_request_interval,
            # The explicit OCR language contract travels with the task, so a
            # resume re-OCRs with the same language / direction / DPI / PSM.
            "ocr_language": effective_ocr_language or "",
            "ocr_language_engine_code": (
                resolved_ocr.engine_code if resolved_ocr else ""),
            "ocr_language_label": resolved_ocr.label if resolved_ocr else "",
            "ocr_text_direction": resolved_ocr.direction if resolved_ocr else "",
            "ocr_dpi": effective_ocr_dpi,
            "ocr_psm": effective_ocr_psm,
            "ocr_quality": effective_ocr_quality,
            "ocr_backend": effective_ocr_backend,
            "ocr_mode": effective_ocr_mode,
            "ocr_paddle_model": effective_paddle_model,
        })
        if effective_online_descriptor:
            pipeline_config["ocr_online"] = dict(effective_online_descriptor)
        else:
            pipeline_config.pop("ocr_online", None)
    else:
        pipeline_config.pop("enable_source_cleanup", None)
    state["pipeline_config"] = pipeline_config
    active_run.record_run_configuration({**pipeline_config, "auxiliary": auxiliary_config},
                                        resume=bool(saved_runtime_config))
    state["quality_mode"] = bool(strict_terminology_governance)
    state["ocr_language"] = effective_ocr_language or ""
    state["ocr_mode"] = effective_ocr_mode
    state["ocr_config"] = {
        "ocr_language": effective_ocr_language or "",
        "ocr_language_engine_code": resolved_ocr.engine_code if resolved_ocr else "",
        "ocr_language_label": resolved_ocr.label if resolved_ocr else "",
        "ocr_text_direction": resolved_ocr.direction if resolved_ocr else "",
        "ocr_dpi": effective_ocr_dpi,
        "ocr_psm": effective_ocr_psm,
        "ocr_quality": effective_ocr_quality,
        "ocr_backend": effective_ocr_backend,
        "ocr_mode": effective_ocr_mode,
        "ocr_paddle_model": effective_paddle_model,
        "ocr_online": dict(effective_online_descriptor or {}),
    }
    # Persist the selected mode before extraction starts.  If the worker is
    # cancelled while rasterizing/OCRing, the resumed call keeps the same CAT
    # contract even when the process environment or UI defaults changed.
    if (not state.get("p1_done")
            or (not state.get("source_cleanup_done")
                and not (state.get("segmentation") or {}).get("mode"))):
        state.setdefault("segmentation", {})["mode"] = effective_segmentation_mode
        state.setdefault("segmentation", {})["locked"] = True
    warnings = state.setdefault("warnings", [])

    if enable_report:
        from folith import academic_writer
        academic_writer.prepare_academic_inputs(
            state, translation_theory, research_settings, literature_sources)
        academic_writer.sync_versions(state)
    save_job_state(job_id, state)

    # 术语依赖失效：必须在“全部完成”早退之前执行，
    # 否则冻结术语表变更后的旧译文会继续以 reviewed/final/TM 状态存在。
    state, stale_segs = _apply_glossary_staleness(state, job_id)
    if stale_segs:
        save_job_state(job_id, state)

    # Legacy PDF tasks created before the source-cleanup checkpoint may already
    # have partial translation pairs.  Re-run the source stage once before
    # continuing, then invalidate every artifact derived from the old text so
    # translations cannot be silently paired with corrected source segments.
    is_pdf = str(filename or state.get("filename") or "").lower().endswith(".pdf")
    if (state.get("p1_done") and not state.get("source_cleanup_done")
            and is_pdf and enable_source_cleanup and not state.get("p2_done")
            and (state.get("extraction_report") or {}).get("semantic_source") not in {"vlm", "mixed"}):
        legacy_source = list(state.get("source_paras") or state.get("paras") or [])
        legacy_segments = list(state.get("paras") or [])
        if legacy_source:
            if on_status:
                on_status("【阶段一】恢复 LLM 原文纠错与断行整理…")
            cleaned, cleanup_meta, cleanup_warnings = cleanup_source_paragraphs(
                legacy_source, provider, api_key, model,
                on_progress=on_status,
                parallelism=effective_cleanup_concurrency,
                confidence_threshold=effective_cleanup_confidence_threshold,
                checkpoint_dir=job_dir(job_id),
                cancel_check=lambda: _runtime_cancel_requested(job_id),
                max_retries=effective_cleanup_max_retries,
                request_interval_seconds=effective_cleanup_request_interval,
                profiler=profiler,
                ocr_language=effective_ocr_language or "",
                source_language=(resolved_ocr.label if resolved_ocr else ""),
                text_direction=(resolved_ocr.direction if resolved_ocr else ""),
                ocr_engine_code=(resolved_ocr.engine_code if resolved_ocr else ""),
                ocr_backend=(((state.get("extraction_report") or {}).get("ocr") or {}).get("actual_engine") or effective_ocr_backend))
            segments, segmentation_meta = _apply_source_segmentation(
                state, cleaned, mode=effective_segmentation_mode)
            state["source_paras"] = legacy_source
            state["paras"] = segments
            state["segmentation"] = segmentation_meta
            state["source_cleanup"] = cleanup_meta
            state["source_cleanup_done"] = True
            for warning in cleanup_warnings:
                if warning not in warnings:
                    warnings.append(warning)
                _append_runtime_technical_log(job_id, f"source cleanup: {warning}")
            # A sentence split changes the positional source/pair contract even
            # when the OCR cleanup text itself is byte-for-byte unchanged.
            if cleaned != legacy_source or segments != legacy_segments:
                cleanup_meta["previous_pair_count"] = len(state.get("pairs") or [])
                _reset_source_dependent_state(state)
                state["paras"] = segments
                state["source_paras"] = legacy_source
                state["source_paragraphs"] = list(cleaned)
                state["segmentation"] = segmentation_meta
                state["source_cleanup"] = cleanup_meta
                state["source_cleanup_done"] = True
            # Deterministic non-translatable admission runs after any source
            # reset, so its exclusion ledger is never discarded with the old
            # translation contract.
            _auto_exclude_non_translatable_segments(job_id, state)
            state["source_quality_gate"] = _audit_source_quality_for_state(state)
            cleanup_meta.update({
                "paragraph_count": len(cleaned),
                "segment_count": len(state.get("paras") or []),
            })
            state["p1_done"] = True
            save_job_state(job_id, state)

    # Whole-book consistency review is a post-translation checkpoint.  Run it
    # before the completed-task fast path as well, so older jobs that already
    # have p2_done can receive the upgrade once without repeating it later.
    if state.get("p2_done"):
        state = _run_targeted_final_review(
            job_id, state, reviewer_config, target_lang, style_rules,
            enable_review=enable_review, on_status=on_status)

    # 全部完成 -> 直接返回
    if state["p1_done"] and state["p2_done"] and (not enable_report or state["p3_done"]) \
            and (not enable_annotate or state.get("annotations_done")) \
            and not retry_failed_understanding:
        state["stage"] = _state_migration.derive_stage(state)
        state["performance"] = persist_performance()
        save_job_state(job_id, state)
        return state

    # ---------------- 阶段一：排版清洗 ----------------
    if not state["p1_done"]:
        if file_bytes is None:
            file_bytes = load_source(job_id)
        if file_bytes is not None:
            try:
                save_source(job_id, file_bytes)
            except Exception:
                pass

        if state.get("source_paras"):
            raw_paragraphs = list(state.get("source_paras") or [])
            paragraphs = list(raw_paragraphs)
            extraction_warnings = []
            extraction_report = state.get("extraction_report") or {}
        else:
            if file_bytes is None:
                raise ValueError("缺少源文件，请重新上传后再继续")
            if on_status:
                on_status("【阶段一】排版解析与段落重建（确定性提取）...")

            # One extraction path for every format: the same call that produces the
            # paragraphs also produces the import-scope report, so a DOCX table can
            # never be dropped silently just because the pipeline took a shortcut.
            paragraphs, extraction_warnings, extraction_report = \
                extract_document_paragraphs_with_report(
                    filename, file_bytes, ocr_max_pages=None, on_progress=on_status,
                    profiler=profiler, checkpoint_dir=job_dir(job_id),
                    ocr_workers=effective_ocr_workers,
                    ocr_queue_size=effective_ocr_queue_size,
                    cancel_check=lambda: _runtime_cancel_requested(job_id),
                    profile_cache_dir=ocr_profile_cache_dir(file_bytes),
                    ocr_language=effective_ocr_language,
                    ocr_dpi=effective_ocr_dpi,
                    ocr_psm=effective_ocr_psm,
                    ocr_mode=effective_ocr_mode,
                    ocr_backend=effective_ocr_backend,
                    ocr_paddle_model=effective_paddle_model,
                    ocr_online_config=effective_online_config,
                    source_language=effective_source_language)

            raw_paragraphs = list(paragraphs)
            # Persist the extraction boundary before remote cleanup starts.  If the
            # process is cancelled while waiting for a model, OCR pages and cleanup
            # batches can resume from their own checkpoints without losing the raw
            # source audit trail.
            state["source_paras"] = raw_paragraphs
            state["extraction_report"] = extraction_report
            save_job_state(job_id, state)
        cleanup_meta = {
            "status": "skipped",
            "input_count": len(raw_paragraphs),
            "output_count": len(raw_paragraphs),
        }
        cleanup_warnings = []
        if extraction_report.get("semantic_source") in {"vlm", "mixed"}:
            cleanup_meta["reason"] = "vlm_semantics_no_ocr_repair"
        if (filename.lower().endswith(".pdf") and enable_source_cleanup and paragraphs
                and extraction_report.get("semantic_source") not in {"vlm", "mixed"}):
            if on_status:
                on_status("【阶段一】LLM 原文纠错与断行整理…")
            paragraphs, cleanup_meta, cleanup_warnings = cleanup_source_paragraphs(
                raw_paragraphs, provider, api_key, model,
                on_progress=on_status,
                parallelism=effective_cleanup_concurrency,
                confidence_by_index=((extraction_report.get("ocr") or {}).get(
                    "paragraph_confidences") if isinstance(extraction_report, dict) else None),
                confidence_threshold=effective_cleanup_confidence_threshold,
                checkpoint_dir=job_dir(job_id),
                cancel_check=lambda: _runtime_cancel_requested(job_id),
                max_retries=effective_cleanup_max_retries,
                request_interval_seconds=effective_cleanup_request_interval,
                profiler=profiler,
                ocr_language=effective_ocr_language or "",
                source_language=(resolved_ocr.label if resolved_ocr else ""),
                text_direction=(resolved_ocr.direction if resolved_ocr else ""),
                ocr_engine_code=(resolved_ocr.engine_code if resolved_ocr else ""),
                ocr_backend=((extraction_report.get("ocr") or {}).get("actual_engine") or effective_ocr_backend))
        elif profiler is not None:
            reason = ("vlm_semantics_no_ocr_repair" if extraction_report.get("semantic_source")
                      in {"vlm", "mixed"} else "source_cleanup_disabled")
            profiler.skipped("deterministic_cleanup", metadata={"reason": reason})
            profiler.skipped("llm_cleanup", metadata={"reason": reason})

        for warning in extraction_warnings:
            if warning not in warnings:
                warnings.append(warning)
            _append_runtime_technical_log(job_id, f"document extraction: {warning}")
        for warning in cleanup_warnings:
            if warning not in warnings:
                warnings.append(warning)
            _append_runtime_technical_log(job_id, f"source cleanup: {warning}")
        if not paragraphs:
            save_job_state(job_id, state)
            detail = (extraction_failure_detail(extraction_warnings)
                      if extraction_warnings else _NO_TEXT_REASON)
            raise ValueError(f"未提取到有效文本：{detail}")
        cleaned_paragraphs = list(paragraphs)
        segments, segmentation_meta = _apply_source_segmentation(
            state, cleaned_paragraphs, mode=effective_segmentation_mode)
        state["paras"] = segments
        # Segments the system can already prove are not translatable leave the
        # working set before the translation contract is built, so they never
        # become a translation request, a review obligation or a delivery
        # notice.
        _auto_exclude_non_translatable_segments(job_id, state)
        segments = state.get("paras") or []
        state["source_cleanup"] = cleanup_meta
        state["source_cleanup_done"] = True
        state["source_cleanup"].update({
            "paragraph_count": len(cleaned_paragraphs),
            "segment_count": len(segments),
        })
        if on_status:
            on_status(f"【阶段一】按句构建翻译单元（0/{len(segments)} 句）…")
        segment_stage = profiler.start_stage(
            "segment_build", concurrency=1, item_count=len(segments),
            metadata={
                "source": "document_ir" if extraction_report.get("semantic_source") in {"vlm", "mixed"} else "cleaned_paragraphs",
                "mode": segmentation_meta.get("mode"),
                "paragraph_count": len(cleaned_paragraphs),
            }
        ) if profiler is not None else None
        state["source_quality_gate"] = _audit_source_quality_for_state(state)
        if segment_stage is not None:
            for index in range(len(segments)):
                segment_stage.item(index, segment_stage.started_monotonic,
                                   metadata={"kind": "segment"})
            segment_stage.finish(status="completed")
        if on_status:
            on_status(f"【阶段一】按句构建翻译单元（{len(segments)}/{len(segments)} 句）…")
        if filename.lower().endswith(".pdf"):
            if file_bytes is not None:
                try:
                    with fitz.open(stream=file_bytes, filetype="pdf") as source_pdf:
                        state["source_page_count"] = source_pdf.page_count
                except Exception:
                    pass
            elif not state.get("source_page_count"):
                pages = ((extraction_report or {}).get("extracted") or {}).get("pages")
                if pages:
                    state["source_page_count"] = pages
                elif "total_pages" in ((extraction_report or {}).get("ocr") or {}):
                    state["source_page_count"] = extraction_report["ocr"]["total_pages"]
        state["p1_done"] = True
        if file_bytes is not None:
            try:
                save_source(job_id, file_bytes)  # 留存源文件，刷新后无需重新上传
            except Exception:
                pass
        for required_stage in ("pdf_classify", "layout_recovery", "rasterize", "ocr",
                               "deterministic_cleanup", "llm_cleanup", "segment_build"):
            if profiler.get(required_stage) is None:
                profiler.skipped(required_stage, metadata={"reason": "not_applicable"})
        state["performance"] = profiler.persist()
        save_job_state(job_id, state)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")

    # ---------------- 阶段 1.2：文档画像（长文理解；失败仅警告，不阻断） ----------------
    if enable_understanding and not state.get("profile_done"):
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        if on_status:
            on_status("【阶段1.2】文档画像（分布式采样 + 结构化校验）...")
        from folith.document_profile import profile_document
        profile, profile_warnings = profile_document(
            state["paras"], auxiliary_config["provider"],
            auxiliary_config["api_key"], auxiliary_config["model"], target_lang,
            call_llm=auxiliary_call,
            cancel_check=lambda: _runtime_cancel_requested(job_id))
        state["document_profile"] = profile
        state["profile_done"] = True
        for w in profile_warnings:
            if w not in warnings:
                warnings.append(w)
        if on_caption and profile:
            on_caption(f"✅ 文档画像完成：领域「{profile.get('domain') or '未知'}」"
                       f"· 文本类型「{profile.get('genre') or '未知'}」")
        elif on_caption:
            on_caption("⚠️ 文档画像失败，已跳过（可在 UI 中人工填写）。")
        save_job_state(job_id, state)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")

    # ---------------- 阶段 1.3：全文语义理解 ----------------
    if (enable_understanding or state.get("understanding_retry_requested")) and (
            not state.get("understanding_done") or retry_failed_understanding):
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        if on_status:
            on_status("【阶段1.3】全文语义理解（语义单元摘要 + 全文概要）...")
        if on_caption:
            on_caption("🧭 正在建立全文概要与当前单元摘要...")
        understanding_call = auxiliary_call
        units, digests, synopsis, understanding_warnings = \
            _context.build_document_understanding(
                state["paras"], state.get("document_profile"),
                auxiliary_config["provider"], auxiliary_config["api_key"],
                auxiliary_config["model"], target_lang, call_llm=understanding_call,
                checkpoint_dir=job_dir(job_id))
        state["semantic_units"] = units
        state["section_digests"] = digests
        state["document_synopsis"] = synopsis
        old_understanding_warnings = set(state.get("understanding_warnings") or [])
        warnings[:] = [warning for warning in warnings if warning not in old_understanding_warnings]
        state["understanding_warnings"] = list(understanding_warnings)
        state["understanding_done"] = True
        state["understanding_retry_requested"] = False
        state["understanding_status"] = "degraded" if understanding_warnings or any(
            item.get("status") == "deterministic_fallback" for item in digests
        ) or synopsis.get("status") == "deterministic_fallback" else "success"
        state["understanding_runtime_config"] = understanding_runtime_config
        for warning in understanding_warnings:
            if warning not in warnings:
                warnings.append(warning)
        _context.write_understanding_artifacts(
            job_dir(job_id), units, digests, synopsis)
        if on_caption:
            status = synopsis.get("status") or "unavailable"
            on_caption(f"{'⚠️ 全文理解降级' if state['understanding_status'] == 'degraded' else '✅ 全文理解完成'}："
                       f"{len(digests)} 个语义单元 · 概要状态 {status}")
        save_job_state(job_id, state)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")

    # ---------------- 阶段 1.5：智能抽取术语 ----------------
    if auto_term and not state["auto_terms"]:
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        if on_status:
            on_status("【阶段1.5】正在 AI 智能抽取全文核心术语...")
        if on_caption:
            on_caption("🤖 正在从全文分布式样本中提取专业术语...")
        from folith.terminology import extract_auto_terms_v2
        entries, extract_warnings = extract_auto_terms_v2(
            state["paras"], target_lang, auxiliary_config["provider"],
            auxiliary_config["api_key"], auxiliary_config["model"],
            document_profile=state.get("document_profile"),
            call_llm=auxiliary_call,
            cancel_check=lambda: _runtime_cancel_requested(job_id))
        state["auto_term_entries"] = entries
        state["auto_terms"] = {e["source"]: e["target"] for e in entries}
        if entries:
            if on_caption:
                on_caption(f"✅ 成功提取 {len(entries)} 个候选术语（全部出现位置已记录）")
        else:
            msg = "术语抽取失败（限流或返回格式异常），已跳过该步骤；可稍后点击“继续处理”重试。"
            if msg not in warnings and msg not in extract_warnings:
                warnings.append(msg)
        for w in extract_warnings:
            if w not in warnings:
                warnings.append(w)
        save_job_state(job_id, state)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")
    legacy_auto = [{"source": k, "target": v, "behavior": "translate",
                    "status": "provisional"}
                   for k, v in (state["auto_terms"] or {}).items()]
    auto_entries = normalize_glossary(state.get("auto_term_entries") or legacy_auto)
    user_entries = normalize_glossary(list(user_glossary or []))
    if not strict_terminology_governance:
        for e in auto_entries:
            if e["status"] == "candidate":
                e["status"] = "provisional"
    working = normalize_glossary(state.get("glossary") or [])
    if not working:
        # 项目记忆已在本函数开头注入到 user_glossary/style_rules，这里只做合并。
        working = normalize_glossary(user_entries + auto_entries)
    else:
        # 新上传/新抽取的术语若不在已保存审核表中，追加（不覆盖人工审核结果）
        known = {e["source"].casefold() for e in working}
        for e in user_entries + auto_entries:
            if e["source"].casefold() not in known:
                working.append(e)
    if not strict_terminology_governance:
        for e in working:
            if e["status"] == "candidate":
                e["status"] = "provisional"
    state["glossary"] = working
    glossary = working
    final_termbase = glossary_to_terms(glossary)

    # ---------------- 严格术语治理门禁：候选术语需人工审核/冻结后才能翻译 ----------------
    if strict_terminology_governance:
        pending = [e for e in working
                   if (e.get("status") or "").lower() == "candidate"]
        if not state.get("glossary_frozen") and not state.get("quality_bypass") \
                and pending:
            if on_status:
                on_status(f"⏸ 严格术语治理：{len(pending)} 条候选术语等待人工审核与冻结...")
            msg = (f"严格术语治理：术语表尚未冻结（{len(pending)} 条自动抽取的"
                   "候选术语待审核），翻译未开始。请在“术语准备与审核”面板"
                   "完成审核并冻结后继续；导入术语库中的锁定术语已视为固定，"
                   "无需再次审核。")
            if msg not in warnings:
                warnings.append(msg)
            state["stage"] = _state_migration.derive_stage(state)
            save_job_state(job_id, state)
            return state

    # ---------------- 阶段二：双语翻译（批次 + 确定性检查 + 独立审校 + 翻译记忆）----------------
    if not state["p2_done"]:
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        if on_status:
            on_status("【阶段二】双语翻译与术语严格注入（批次翻译 + 确定性检查 + 独立审校）...")
        _batch_cfg = state.get("pipeline_config") or {}
        translate_stage(state, job_id, glossary, provider, api_key, model, target_lang,
                        style_rules, enable_review, use_tm=use_tm,
                        document_profile=state.get("document_profile"),
                        translator_config=translator_config,
                        reviewer_config=reviewer_config,
                        auxiliary_config=auxiliary_config,
                        on_status=on_status, on_caption=on_caption,
                        batch_size=_batch_cfg.get("batch_size", batch_size),
                        max_batch_chars=_batch_cfg.get("max_batch_chars", max_batch_chars),
                        knowledge_feedback_interval=_batch_cfg.get(
                            "knowledge_feedback_interval", knowledge_feedback_interval),
                        translation_concurrency=_batch_cfg.get(
                            "translation_concurrency", translation_concurrency),
                        profiler=profiler,
                        review_overlap=_batch_cfg.get("review_overlap"),
                        max_review_in_flight=_batch_cfg.get("max_review_in_flight"),
                        review_overlap_strategy=_batch_cfg.get("review_overlap_strategy"))
        for _required_stage in ("translation_generation", "translation_deterministic_qa",
                                "translation_repair", "translation_semantic_review",
                                "knowledge_feedback", "state_commit", "tm_commit",
                                "runtime_telemetry"):
            if profiler.get(_required_stage) is None:
                profiler.skipped(_required_stage, metadata={"reason": "not_applicable"})
        state["performance"] = profiler.persist()
        state["p2_done"] = True
        from folith.epub import bind_source_references
        bind_source_references(state)
        from folith.document_ir import attach_pair_spans
        attach_pair_spans(state)
        from folith import delivery as _delivery
        state["delivery_status"] = _delivery.compute_delivery_status(state)
        save_job_state(job_id, state)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")

    if state.get("p2_done"):
        state = _run_targeted_final_review(
            job_id, state, reviewer_config, target_lang, style_rules,
            enable_review=enable_review, on_status=on_status)
        active_run.record_point("milestone", phase="translation_review_complete")

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")

    # ---------------- 阶段 2.5：三色自动标注 ----------------
    if enable_annotate and state["p2_done"] and not state.get("annotations_done"):
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        if on_status:
            on_status("【阶段 2.5】自动标注学习重点（红=生僻词 / 黄=专业名词 / 青绿=难点句）...")
        annotate_stage(state, job_id, glossary, provider, api_key, model, target_lang,
                       on_caption=on_caption)

    if _runtime_cancel_requested(job_id):
        raise RuntimeError("任务已请求取消")

    # ---------------- 阶段三：报告生成 ----------------
    if enable_report and not state["p3_done"]:
        if _runtime_cancel_requested(job_id):
            raise RuntimeError("任务已请求取消")
        if on_status:
            on_status(f"【阶段三】基于《{translation_theory}》生成报告...")
        report_md = generate_mti_report(state["pairs"], final_termbase, translation_theory,
                                        provider, api_key, model, state, job_id,
                                        on_status=on_status,
                                        research_settings=research_settings,
                                        literature_sources=literature_sources)
        if not report_md.strip():
            if state.get("report_status") == "blocked_final_case_policy":
                state["stage"] = _state_migration.derive_stage(state)
                save_job_state(job_id, state)
                return state
            raise RuntimeError("报告内容为空，请点击“继续处理”重试")
        state["p3_md"] = report_md
        state["theory"] = translation_theory
        state["p3_done"] = True
        save_job_state(job_id, state)

    state["stage"] = _state_migration.derive_stage(state)
    state["performance"] = persist_performance()
    save_job_state(job_id, state)
    active_run.record_point("milestone", phase="automatic_pipeline_complete")
    trace.finish_run(job_id=job_id, status="completed")
    return state
