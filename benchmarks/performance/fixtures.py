"""Deterministic fixtures for the Folith performance benchmark.

Nothing here is a mock of Folith: the documents are real DOCX files that the
real pipeline imports, segments and translates.  Only the *model* is replaced
(by the repository's own ``tests/offline_provider.py``), so a run is
reproducible and needs no network.

The document generator is seeded, so the same ``--segments`` value always
produces the same paragraphs and therefore the same amount of work.
"""
from __future__ import annotations

import random
from pathlib import Path
from typing import List

_SUBJECT = (
    "ecological succession", "the recovery of degraded peatlands",
    "riverine nutrient budgets", "the distribution of woodland birds",
    "long-term soil carbon dynamics", "coastal wetland accretion",
    "the response of alpine flora", "pollinator network structure",
    "sediment transport in braided rivers", "restoration of floodplain forests",
    "the measurement of canopy reflectance", "microbial turnover in permafrost",
    "the spread of invasive shrubs", "primary productivity in oligotrophic lakes",
    "fire regimes in boreal woodland", "the phenology of temperate hedgerows",
    "downstream sediment sorting", "the acoustic survey of bat colonies",
    "post-glacial vegetation succession", "the chemistry of headwater streams",
)

_VERB = (
    "The present chapter examines", "This section considers",
    "The following analysis describes", "The account below surveys",
    "The discussion that follows reviews", "The treatment given here evaluates",
)

_CLAUSE = (
    "the mechanisms that govern recovery", "the conditions under which stability is lost",
    "the relative weight of local and regional controls",
    "the consequences for management practice",
    "the limits of the available instrumental record",
    "the implications for long-term monitoring programmes",
    "the degree to which the two explanations can be distinguished",
    "the assumptions that the estimate depends upon",
    "the scale at which the pattern becomes apparent",
    "the sensitivity of the result to the choice of baseline",
)

_TAIL = (
    "Throughout, the emphasis falls on what the record can actually support rather "
    "than on the elegance of the model. Where the evidence is thin, the uncertainty "
    "is stated rather than smoothed away, because a reconstruction that hides its "
    "own error is of little use to the practitioner.",
    "It should be emphasised that the categories used here are practical rather "
    "than absolute. A single classification can conceal a gradient, and the reader "
    "is invited to weigh the boundaries as firmly as the evidence permits.",
    "Field practice has shaped the argument more than any single dataset. The "
    "difficult cases, encountered repeatedly over many seasons, have done more to "
    "refine the account than the idealised examples have done to extend it.",
    "A final caveat concerns the comparability of the records, which are drawn from "
    "a range of sources with markedly different resolution and coverage. Where a "
    "figure depends on the least complete series, it is flagged in the notes.",
)


def build_paragraph(index: int, rng: random.Random) -> str:
    subject = rng.choice(_SUBJECT)
    verb = rng.choice(_VERB)
    clause = rng.choice(_CLAUSE)
    body = (
        f"{verb} {subject}, with particular attention to {clause}. The record "
        f"assembled for the period is fragmentary, and the account given here is "
        f"built from a mixture of direct measurement, inference from associated "
        f"change, and comparison with better-documented analogues. Each of these "
        f"carries a different weight, and the argument is arranged so that the "
        f"strongest available evidence is brought to bear on the questions that "
        f"matter most, rather than spreading thinly across questions that the "
        f"data cannot settle. Against {subject}, the position taken here is that "
        f"{clause} remains open, and that {rng.choice(_CLAUSE)} is at least as "
        f"important a determinant as the process most often invoked."
    )
    return f"{index:05d}. {body} {rng.choice(_TAIL)}"


def build_docx(path: Path, paragraphs: int, *, seed: int = 20261022) -> Path:
    """Write a deterministic DOCX with ``paragraphs`` body paragraphs."""
    from docx import Document

    rng = random.Random(seed)
    document = Document()
    document.add_heading("Restoration Ecology: Field Observations", level=1)
    for index in range(paragraphs):
        document.add_paragraph(build_paragraph(index, rng))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def sample_segments(paragraphs: int) -> List[str]:
    """The source texts a job of this size will contain (used for expectations)."""
    rng = random.Random(20261022)
    return [build_paragraph(index, rng) for index in range(paragraphs)]
