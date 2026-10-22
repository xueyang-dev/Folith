#!/usr/bin/env python3
"""README 截图：用合成的演示数据打开 Qt 客户端，截取主要页面。

    venv/bin/python scripts/qt_showcase_screenshots.py [--out docs/assets] [--theme light]
    venv/bin/python scripts/qt_showcase_screenshots.py --lang en

演示文档与译文是为截图专门写的，不含任何真实用户文件；数据目录放在 tmp/ 下，
跑之前会清空。窗口会写真实的 QSettings（主题、语言），脚本在开头导出、结尾还原
``com.folith.QtDesktop``，不会改掉本机的偏好。

``--lang`` 决定界面语言与演示数据（语言代码取自 ``folith/qt/i18n.py``）：

* ``zh-CN``（默认）：基准语言，文件名与行为都和以前一样（``qt-01-tasks.png``）；
* ``en`` / ``ru`` / ``es``：界面文案走各自的语言包，演示数据（项目、术语表、
  原文与译文、文档画像、风险提示）也换成该语言，文件名多一个语言后缀
  （``qt-01-tasks.en.png``），免得译本 README 又指回中文截图。

演示内容全部是合成的：文档、译文、术语与项目名都是为截图写的，同一本书在
中文、英文、西班牙文、俄文里各有一个版本，彼此对得上。
"""
from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import sys
import time
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from folith.qt.selftest import open_job_and_wait  # noqa: E402

#: 界面语言 → 截图文件名后缀。中文是基准语言，文件名一个字都不改。
LANG_SUFFIX = {"zh-CN": "", "en": ".en", "ru": ".ru", "es": ".es"}

#: 演示项目的固定 ID：色点颜色由 ID 派生（``components.project_dot_token``），
#: 随机 ID 会让每次生成的截图颜色不同。书稿=蓝点（accent），合规=褐点（warning）。
DEMO_PROJECT_IDS = (
    "f0170000-0000-4000-8000-000000000005",
    "f0170000-0000-4000-8000-000000000002",
)

#: 窗口尺寸（逻辑像素）。中文界面 1440x900 就够（版面的最小宽度比它小）；英/西/俄
#: 的文案更长，版面的最小宽度分别是 1474 / 1546 / 1623，窗口太窄就会把按钮和
#: 卡片文字挤掉，所以三门外语统一用最宽的那个（1623x900），谁都不特殊。
WINDOW_SIZE = (1440, 900)
WINDOW_SIZE_LOCALIZED = (1623, 900)

#: 演示用的四种语言版本：同一本书的原文与译文。
EN_PARAGRAPHS = [
    "Long documents fail in translation for reasons that have little to do with any single sentence.",
    "A term that is rendered one way in chapter two and another way in chapter nine forces every reader to stop and wonder whether the author meant two different things.",
    "Folith treats a book as one piece of work rather than a stream of unrelated requests.",
    "Before the first paragraph is translated, the workspace reads the whole document, builds a profile of its domain and register, and proposes a glossary for the translator to confirm.",
    "Confirmed terms are locked for the project and injected into every request, so a locked term cannot drift between batches.",
    "Translation memory is kept per project: segments approved in one task are reused in the next without leaking into unrelated work.",
    "Each batch of paragraphs is sent with the neighbouring context, which lets the model resolve pronouns and abbreviations that span a page break.",
    "An independent reviewer reads the draft against the source and reports findings by category, severity and location.",
    "Blocking findings must be resolved by a person before the task can be delivered; informational notes never block.",
    "Edits made by hand are recorded as such, and an edited paragraph is reviewed again before it is trusted.",
    "Scanned pages are read by a vision model, while the local OCR engine only supplies the position of each line on the page.",
    "When a run is interrupted, the work saved so far stays on disk and the task continues from the last completed batch.",
    "The final delivery is frozen as an immutable snapshot with a manifest, so the files handed over can always be traced back to the reviewed text.",
    "The result is a translation that can be checked, explained and defended, not merely produced.",
]
ZH_PARAGRAPHS = [
    "长文档的翻译之所以失败，原因往往与某一个句子无关。",
    "同一个术语在第二章是一种译法、到第九章又换成另一种，读者每次都得停下来猜：作者是不是在说两件不同的事。",
    "Folith 把一本书当作一项完整的工作，而不是一串互不相干的请求。",
    "在翻译第一段之前，工作台会先通读全文，建立领域与语体画像，并提出一份术语表交给译者确认。",
    "确认后的术语在项目内锁定，并注入每一次请求，因此锁定术语不会在批次之间漂移。",
    "翻译记忆按项目保存：一个任务里通过审校的句段会在下一个任务中复用，而不会混入无关的工作。",
    "每一批段落都会附带前后文一起发送，模型因此能够处理跨页出现的代词和缩写。",
    "独立审校会对照原文通读初稿，并按类别、严重程度和位置报告问题。",
    "阻断性问题必须由人处理后任务才能交付；提示性说明永远不会阻断交付。",
    "人工修改会被如实记录，改过的段落在被采信之前会再次接受审校。",
    "扫描页由视觉模型阅读，本机 OCR 只负责给出每一行在页面上的位置。",
    "运行中断时，已经保存的工作留在磁盘上，任务会从最后完成的批次继续。",
    "最终交付会冻结为带清单的不可变快照，交出去的每个文件都能追溯到审校过的文本。",
    "这样得到的译文可以被检查、被解释、被辩护，而不只是被生成出来。",
]
ES_PARAGRAPHS = [
    "Los documentos largos fracasan en la traducción por razones que tienen poco que ver con una frase concreta.",
    "Un término traducido de una manera en el capítulo dos y de otra en el capítulo nueve obliga a cada lector a detenerse a preguntarse si el autor quería decir dos cosas distintas.",
    "Folith trata un libro como una sola obra y no como una sucesión de solicitudes sin relación entre sí.",
    "Antes de traducir el primer párrafo, el espacio de trabajo lee el documento completo, traza un perfil de su dominio y de su registro y propone un glosario para que el traductor lo confirme.",
    "Los términos confirmados quedan bloqueados para el proyecto y se inyectan en cada solicitud, de modo que un término bloqueado no puede desviarse entre lotes.",
    "La memoria de traducción se guarda por proyecto: los segmentos aprobados en una tarea se reutilizan en la siguiente sin mezclarse con trabajos ajenos.",
    "Cada lote de párrafos se envía junto con el contexto contiguo, lo que permite al modelo resolver pronombres y abreviaturas que atraviesan un salto de página.",
    "Un revisor independiente contrasta la traducción con el original y comunica sus hallazgos por categoría, gravedad y ubicación.",
    "Los hallazgos bloqueantes debe resolver una persona antes de poder entregar la tarea; las notas informativas nunca bloquean la entrega.",
    "Las correcciones hechas a mano quedan registradas como tales, y un párrafo corregido se revisa de nuevo antes de ser aceptado.",
    "Las páginas escaneadas las lee un modelo de visión, mientras que el motor de OCR local solo aporta la posición de cada línea en la página.",
    "Si una ejecución se interrumpe, el trabajo guardado hasta ese momento permanece en el disco y la tarea continúa desde el último lote completado.",
    "La entrega final se congela como una instantánea inmutable con su manifiesto, de modo que cada archivo entregado siempre puede rastrearse hasta el texto revisado.",
    "El resultado es una traducción que se puede comprobar, explicar y defender, no solo producir.",
]
RU_PARAGRAPHS = [
    "Длинные тексты не переводятся по причинам, которые почти не имеют отношения к отдельному предложению.",
    "Если термин во второй главе переведён одним образом, а в девятой — другим, каждому читателю приходится останавливаться и гадать, имел ли автор в виду две разные вещи.",
    "Folith воспринимает книгу как одну цельную работу, а не как череду несвязанных между собой запросов.",
    "Прежде чем перевести первый абзац, рабочее место читает документ целиком, составляет профиль его предметной области и стилистического регистра и предлагает глоссарий, который переводчик должен подтвердить.",
    "Подтверждённые термины фиксируются в проекте и подставляются в каждый запрос, поэтому зафиксированный термин не может «поплыть» от партии к партии.",
    "Память перевода ведётся по каждому проекту: сегменты, принятые в одной задаче, используются в следующей и не попадают в постороннюю работу.",
    "Каждая партия абзацев отправляется вместе с соседним контекстом — благодаря этому модель разбирает местоимения и сокращения, которые встречаются на стыке страниц.",
    "Независимый проверяющий сверяет черновик с оригиналом и сообщает о замечаниях с указанием категории, серьёзности и места.",
    "Блокирующие замечания должен устранить человек, иначе задачу нельзя сдать; замечания для сведения никогда не блокируют сдачу.",
    "Правки, внесённые вручную, помечаются как ручные, и исправленный абзац проверяют заново, прежде чем ему поверить.",
    "Страницы со сканами читает визуальная модель, а локальный модуль OCR сообщает только положение каждой строки на странице.",
    "Если запуск прерван, всё, что уже сохранено, остаётся на диске, и задача продолжается с последней завершённой партии.",
    "Итоговая сдача замораживается как неизменяемый снимок с манифестом, поэтому каждый переданный файл всегда можно проследить до проверенного текста.",
    "Получается перевод, который можно проверить, объяснить и отстоять, а не просто получить.",
]

#: 第二本书（金融合规）——每种语言各有一版原文与译文。名字里的语言代码是
#: 「原文语言 → 目标语言」。
SECOND_EN_ZH = [
    ("Supervisory authorities expect every model used in credit decisions to be documented, validated and monitored.",
     "监管机构要求，凡用于信贷决策的模型都必须有文档、经过验证并持续监测。"),
    ("Data lineage must be recorded from the original source to every derived report.",
     "数据血缘必须从原始来源一直记录到每一份派生报告。"),
    ("Exceptions are escalated to the risk committee within five business days.",
     "例外情况须在五个工作日内上报风险委员会。"),
]
SECOND_ES_EN = [
    ("Las autoridades de supervisión exigen que todos los modelos empleados en decisiones de crédito estén documentados, validados y monitorizados.",
     "Supervisory authorities expect every model used in credit decisions to be documented, validated and monitored."),
    ("La trazabilidad de los datos debe registrarse desde la fuente original hasta cada informe derivado.",
     "Data lineage must be recorded from the original source to every derived report."),
    ("Las excepciones se elevan al comité de riesgos en un plazo de cinco días hábiles.",
     "Exceptions are escalated to the risk committee within five business days."),
]
SECOND_EN_ES = [
    ("Supervisory authorities expect every model used in credit decisions to be documented, validated and monitored.",
     "Las autoridades de supervisión exigen que todos los modelos empleados en decisiones de crédito estén documentados, validados y monitorizados."),
    ("Data lineage must be recorded from the original source to every derived report.",
     "La trazabilidad de los datos debe registrarse desde la fuente original hasta cada informe derivado."),
    ("Exceptions are escalated to the risk committee within five business days.",
     "Las excepciones se elevan al comité de riesgos en un plazo de cinco días hábiles."),
]
SECOND_EN_RU = [
    ("Supervisory authorities expect every model used in credit decisions to be documented, validated and monitored.",
     "Надзорные органы ожидают, что каждая модель, используемая в кредитных решениях, задокументирована, проверена и находится под постоянным мониторингом."),
    ("Data lineage must be recorded from the original source to every derived report.",
     "Происхождение данных должно прослеживаться от первоисточника до каждого производного отчёта."),
    ("Exceptions are escalated to the risk committee within five business days.",
     "Исключения выносятся на комитет по рискам в течение пяти рабочих дней."),
]


@dataclass(frozen=True)
class Demo:
    """一套演示数据：项目、术语表、原文与译文，全是给截图写的合成内容。"""

    source_lang: str                       # 原文语言（空 = 自动识别）
    target_lang: str                       # 目标语言
    explicit_direction: bool                # 是否把语言方向写死在项目上
    projects: tuple[tuple[str, str], ...]  # (项目名, 项目描述)
    book_file: str                         # 第一本书的文件名
    source: tuple[str, ...]                # 第一本书的原文段落
    target: tuple[str, ...]                # 第一本书的译文段落
    terms: tuple[tuple[str, str, tuple[int, ...]], ...]
    second_file: str                       # 第二本书的文件名
    second: tuple[tuple[str, str], ...]
    prep_file: str                         # 译前准备页用的文档名
    domain: str                            # 术语领域
    scope: str                             # 术语作用域
    genre: str
    audience: str
    register: str
    finding: dict                          # 阻断性审校问题
    phase_completed: str
    phase_translating: str
    phase_ready: str


DEMOS: dict[str, Demo] = {
    "zh-CN": Demo(
        source_lang="",
        target_lang="简体中文",
        explicit_direction=False,
        projects=(
            ("技术专著本地化", "长篇技术写作，统一术语与翻译记忆"),
            ("金融合规白皮书", "监管与风控文件"),
        ),
        book_file="Why Long Documents Need a Workspace.docx",
        source=tuple(EN_PARAGRAPHS),
        target=tuple(ZH_PARAGRAPHS),
        terms=(
            ("translation memory", "翻译记忆", (5,)),
            ("glossary", "术语表", (3,)),
            ("independent reviewer", "独立审校", (7,)),
            ("blocking finding", "阻断性问题", (8,)),
            ("immutable snapshot", "不可变快照", (12,)),
            ("register", "语体", (3,)),
        ),
        second_file="Model Risk Governance Standard.docx",
        second=tuple(SECOND_EN_ZH),
        prep_file="Quarterly Product Whitepaper.docx",
        domain="本地化工程",
        scope="项目",
        genre="技术写作",
        audience="专业读者",
        register="正式书面语",
        finding={
            "type": "review", "category": "terminology", "severity": "blocking",
            "status": "open", "segment_id": 4, "segment_index": 4,
            "location_key": "term:lock", "requires_human_confirmation": True,
            "summary": "「漂移」在项目术语表中统一为「偏移」",
            "source_span": "drift", "target_span": "漂移",
            "explanation": "项目术语表把 drift 译为「偏移」，第 12 段也用了「偏移」；此处请保持一致，或同步修改术语表。",
            "recommendation": "采纳建议译文，或保留当前译文并记录理由。",
            "suggested_target": "确认后的术语在项目内锁定，并注入每一次请求，因此锁定术语不会在批次之间偏移。",
            "detector": "Folith Reviewer", "confidence": 0.9,
            "review_event_id": "rev-showcase-book", "reason": "术语一致性",
        },
        phase_completed="已完成",
        phase_translating="翻译中",
        phase_ready="待开始",
    ),
    "en": Demo(
        source_lang="Español",
        target_lang="English",
        explicit_direction=True,
        projects=(
            ("Document Localization",
             "Long-form technical writing, with one shared termbase and translation memory"),
            ("Financial Compliance",
             "Supervisory and risk-management documents"),
        ),
        book_file="Por qué los documentos largos necesitan un espacio de trabajo.docx",
        source=tuple(ES_PARAGRAPHS),
        target=tuple(EN_PARAGRAPHS),
        terms=(
            ("memoria de traducción", "translation memory", (5,)),
            ("glosario", "glossary", (3,)),
            ("revisor independiente", "independent reviewer", (7,)),
            ("hallazgo bloqueante", "blocking finding", (8,)),
            ("instantánea inmutable", "immutable snapshot", (12,)),
            ("registro", "register", (3,)),
        ),
        second_file="Norma de gobernanza del riesgo de modelos.docx",
        second=tuple(SECOND_ES_EN),
        prep_file="Informe trimestral de producto.docx",
        domain="Localization engineering",
        scope="Project",
        genre="Technical writing",
        audience="Professional readers",
        register="Formal written register",
        finding={
            "type": "review", "category": "terminology", "severity": "blocking",
            "status": "open", "segment_id": 4, "segment_index": 4,
            "location_key": "term:lock", "requires_human_confirmation": True,
            "summary": "The termbase locks «deriva» as «drift», but this segment says «desviarse»",
            "source_span": "desviarse", "target_span": "drift",
            "explanation": "The project termbase locks the source term «deriva» and renders it as «drift»; the source of this segment uses «desviarse» instead. Please keep them consistent or update the termbase.",
            "recommendation": "Accept the suggested translation, or keep the current one and record the reason.",
            "suggested_target": "Confirmed terms are locked for the project and injected into every request, so a locked term cannot drift between batches.",
            "detector": "Folith Reviewer", "confidence": 0.9,
            "review_event_id": "rev-showcase-book", "reason": "Terminology consistency",
        },
        phase_completed="Completed",
        phase_translating="Translating",
        phase_ready="Ready",
    ),
    "ru": Demo(
        source_lang="English",
        target_lang="Русский",
        explicit_direction=True,
        projects=(
            ("Локализация документов",
             "Объёмные технические тексты: единая терминология и память перевода"),
            ("Финансовый комплаенс",
             "Документы надзора и управления рисками"),
        ),
        book_file="Why Long Documents Need a Workspace.docx",
        source=tuple(EN_PARAGRAPHS),
        target=tuple(RU_PARAGRAPHS),
        terms=(
            ("translation memory", "память перевода", (5,)),
            ("glossary", "глоссарий", (3,)),
            ("independent reviewer", "независимый проверяющий", (7,)),
            ("blocking finding", "блокирующее замечание", (8,)),
            ("immutable snapshot", "неизменяемый снимок", (12,)),
            ("register", "регистр", (3,)),
        ),
        second_file="Model Risk Governance Standard.docx",
        second=tuple(SECOND_EN_RU),
        prep_file="Quarterly Product Whitepaper.docx",
        domain="Локализация",
        scope="Проект",
        genre="Техническая письменность",
        audience="Профессиональные читатели",
        register="Формальный письменный стиль",
        finding={
            "type": "review", "category": "terminology", "severity": "blocking",
            "status": "open", "segment_id": 4, "segment_index": 4,
            "location_key": "term:lock", "requires_human_confirmation": True,
            "summary": "В терминологии проекта «drift» передаётся как «смещение», а здесь стоит «поплыть»",
            "source_span": "drift", "target_span": "поплыть",
            "explanation": "Терминология проекта передаёт «drift» как «смещение», а в этом абзаце стоит «поплыть». Держитесь того же варианта или поправьте терминологию.",
            "recommendation": "Примите предложенный вариант или оставьте свой и запишите причину.",
            "suggested_target": "Подтверждённые термины фиксируются в проекте и подставляются в каждый запрос, поэтому зафиксированный термин не может «смещаться» от партии к партии.",
            "detector": "Folith Reviewer", "confidence": 0.9,
            "review_event_id": "rev-showcase-book", "reason": "Терминологическая согласованность",
        },
        phase_completed="Завершено",
        phase_translating="Перевод",
        phase_ready="Готово",
    ),
    "es": Demo(
        source_lang="English",
        target_lang="Español",
        explicit_direction=True,
        projects=(
            ("Traducción técnica",
             "Textos técnicos extensos, con terminología y memoria de traducción unificadas"),
            ("Cumplimiento financiero",
             "Documentos de supervisión y gestión de riesgos"),
        ),
        book_file="Why Long Documents Need a Workspace.docx",
        source=tuple(EN_PARAGRAPHS),
        target=tuple(ES_PARAGRAPHS),
        terms=(
            ("translation memory", "memoria de traducción", (5,)),
            ("glossary", "glosario", (3,)),
            ("independent reviewer", "revisor independiente", (7,)),
            ("blocking finding", "hallazgo bloqueante", (8,)),
            ("immutable snapshot", "instantánea inmutable", (12,)),
            ("register", "registro", (3,)),
        ),
        second_file="Model Risk Governance Standard.docx",
        second=tuple(SECOND_EN_ES),
        prep_file="Quarterly Product Whitepaper.docx",
        domain="Localización",
        scope="Proyecto",
        genre="Escritura técnica",
        audience="Lectores profesionales",
        register="Registro formal",
        finding={
            "type": "review", "category": "terminology", "severity": "blocking",
            "status": "open", "segment_id": 4, "segment_index": 4,
            "location_key": "term:lock", "requires_human_confirmation": True,
            "summary": "La base terminológica traduce «drift» como «deriva», pero aquí dice «desviarse»",
            "source_span": "drift", "target_span": "desviarse",
            "explanation": "La base terminológica del proyecto traduce «drift» como «deriva», mientras que en este párrafo aparece «desviarse». Mantén la coherencia o actualiza la base terminológica.",
            "recommendation": "Acepta la traducción sugerida o conserva la actual y deja constancia del motivo.",
            "suggested_target": "Los términos confirmados quedan bloqueados para el proyecto y se inyectan en cada solicitud, de modo que un término bloqueado no sufra derivas entre lotes.",
            "detector": "Folith Reviewer", "confidence": 0.9,
            "review_event_id": "rev-showcase-book", "reason": "Coherencia terminológica",
        },
        phase_completed="Completado",
        phase_translating="Traduciendo",
        phase_ready="Listo",
    ),
}


#: 记住核心原本的项目 ID 生成器，截图脚本用完就还回去。
_ORIGINAL_NEW_PROJECT_ID: list = []


def _remember_project_id_generator() -> None:
    from folith import project as project_module

    if not _ORIGINAL_NEW_PROJECT_ID:
        _ORIGINAL_NEW_PROJECT_ID.append(project_module.new_project_id)


def make_docx(path: Path, paragraphs: list[str], title: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=1)
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)
    return path.read_bytes()


def _direction(demo: Demo) -> dict:
    """项目上要不要写死语言方向。

    基准语言（zh-CN）不写：方向由任务推断，界面显示「自动 → 中」，和以前的
    截图一模一样。其它语言写死，免得界面上出现「自动」这种半截说明。
    """
    if not demo.explicit_direction:
        return {}
    return {"source_lang": demo.source_lang, "target_lang": demo.target_lang}


def _fixed_project_ids() -> None:
    """让演示项目用固定 ID（见 :data:`DEMO_PROJECT_IDS`）。

    项目色点的颜色由项目 ID 派生（``components.project_dot_token``），随机 ID
    会让每次生成的截图颜色都不一样。固定 ID 之后截图可复现。
    """
    from folith import project as project_module

    counter = {"n": 0}
    original = project_module.new_project_id

    def _next() -> str:
        index = counter["n"]
        counter["n"] += 1
        if index < len(DEMO_PROJECT_IDS):
            return DEMO_PROJECT_IDS[index]
        return original()

    project_module.new_project_id = _next


def _restore_project_ids() -> None:
    """把项目 ID 的生成器还回去（截图脚本不该改全局行为）。"""
    from folith import project as project_module

    project_module.new_project_id = _ORIGINAL_NEW_PROJECT_ID[0]


def _stamp_project(data_dir: Path, project_id: str, updated_at: str) -> None:
    """把项目记录里的更新时间改成指定值。

    核心每次保存都把 ``updated_at`` 写成"现在"（秒级），同一秒里建的两个项目
    只能靠 UUID 排序——任务库里的项目顺序、顶部"继续上次的工作"每次跑都可能
    换一张脸。截图要可复现，所以保存之后把这一项改回脚本指定的时间（记录结构
    不变，只是这个时间戳）。
    """
    import json

    from folith import project as project_module

    path = project_module.project_path(Path(data_dir), project_id)
    record = json.loads(path.read_text(encoding="utf-8"))
    record["updated_at"] = updated_at
    path.write_text(json.dumps(project_module.normalize_project(record),
                                ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _job_stamps() -> dict[str, str]:
    """演示任务的时间戳：书稿最新、合规其次、译前准备那份最旧。

    界面按"最近更新"排序（任务库顶部、项目分组），时间戳固定下来，截图里
    出现在哪儿才是数据决定的，而不是跑脚本的那一秒。
    """
    from datetime import datetime, timedelta, timezone

    base = datetime.now(timezone.utc).replace(microsecond=0)
    return {
        "showcase-book": base.isoformat(timespec="seconds"),
        "showcase-fin": (base - timedelta(seconds=1)).isoformat(timespec="seconds"),
        "showcase-prep": (base - timedelta(seconds=2)).isoformat(timespec="seconds"),
    }


def seed(core, data_dir: Path, demo: Demo) -> dict:
    from folith import translation_core

    fixtures = data_dir.parent / "fixtures"
    fixtures.mkdir(parents=True, exist_ok=True)
    (data_dir / ".onboarded").touch()
    stamps = _job_stamps()
    _remember_project_id_generator()
    _fixed_project_ids()

    book = core.create_project(demo.projects[0][0], description=demo.projects[0][1],
                               **_direction(demo))
    fin = core.create_project(demo.projects[1][0], description=demo.projects[1][1],
                              **_direction(demo))
    _restore_project_ids()
    glossary = [{
        "id": f"term-{i:02d}", "source": s, "target": t, "preferred": t,
        "status": "locked", "domain": demo.domain, "scope": demo.scope,
        "occurrences": list(occ), "note": "",
    } for i, (s, t, occ) in enumerate(demo.terms, 1)]

    pairs = []
    for i, (s, t) in enumerate(zip(demo.source, demo.target)):
        pairs.append({
            "segment_id": f"seg-{i:03d}", "source": s, "target": t,
            "initial_target": t, "reviewed": True,
            "review_status": "reviewed_clean",
            "target_provenance": "reviewed",
        })
    pairs[9]["target_provenance"] = "human"

    def job(job_id, filename, project, rows, *, review_upto, finding=None, status="completed"):
        source_bytes = make_docx(fixtures / filename, [r["source"] for r in rows],
                                 filename.rsplit(".", 1)[0])
        state = core.new_job_state(filename)
        state.update({
            "p1_done": True, "p2_done": True, "p3_done": False,
            "report_enabled": False, "target_lang": demo.target_lang,
            "project_id": project["project_id"],
            "paras": [r["source"] for r in rows], "source_paras": [r["source"] for r in rows],
            "pairs": rows, "glossary": glossary if project is book else [],
            "profile_done": True,
            "document_profile": {
                "display_name": filename.rsplit(".", 1)[0], "domain": demo.domain,
                "genre": demo.genre, "audience": demo.audience, "register": demo.register,
            },
            "translation_core_review_required": True,
            "review_evidence": [{
                "review_event_id": f"rev-{job_id}", "review_scope": "current_translation",
                "segment_ids": list(range(review_upto)),
                "created_at": "2026-10-17T10:00:00+08:00",
                "completion_receipt": {"status": "completed",
                                       "reviewed_segment_ids": list(range(review_upto))},
            }],
            "findings": [], "human_actions": [],
            "provider": "DeepSeek", "model": "deepseek-v4-flash",
            "stage": "TRANSLATED", "delivery_config": core.default_delivery_config(),
        })
        state["updated_at"] = stamps[job_id]
        if demo.explicit_direction:
            state["source_lang"] = demo.source_lang
        if finding is not None:
            index, raw = finding
            fp = translation_core.fingerprint({"segment": index, "target": rows[index]["target"]})
            normalized = translation_core.normalize_finding(raw, input_fingerprint=fp)
            normalized.update({"type": "review", "segment_index": index})
            state["findings"] = [normalized]
            state["has_blocking"] = raw.get("severity") == "blocking"
        core.save_job_state(job_id, state)
        core.save_source(job_id, source_bytes)
        core.update_runtime_state(job_id, status=status, phase=status,
                                  phase_label=demo.phase_completed if status == "completed"
                                  else demo.phase_translating,
                                  last_event=status)

    job("showcase-book", demo.book_file, book, pairs,
        review_upto=14, finding=(4, dict(demo.finding)))
    fin_rows = [{"segment_id": f"seg-{i:03d}", "source": s, "target": t, "initial_target": t,
                 "reviewed": True, "review_status": "reviewed_clean",
                 "target_provenance": "reviewed"} for i, (s, t) in enumerate(demo.second)]
    job("showcase-fin", demo.second_file, fin, fin_rows, review_upto=3)

    prep_name = demo.prep_file
    prep_bytes = make_docx(fixtures / prep_name, list(demo.source), prep_name.rsplit(".", 1)[0])
    prep = core.new_job_state(prep_name)
    prep.update({
        "target_lang": demo.target_lang,
        "project_id": fin["project_id"],
        "provider": "DeepSeek", "model": "deepseek-v4-flash",
        "delivery_config": core.default_delivery_config(),
    })
    if demo.explicit_direction:
        prep["source_lang"] = demo.source_lang
    prep["updated_at"] = stamps["showcase-prep"]
    core.save_job_state("showcase-prep", prep)
    core.save_source("showcase-prep", prep_bytes)
    core.update_runtime_state("showcase-prep", status="ready", phase="ready",
                              phase_label=demo.phase_ready, last_event="ready")

    book["glossary"] = glossary
    core.save_project(book)
    # 项目顺序（任务库顶部那排项目标签、侧栏「最近项目」）按记录的 updated_at
    # 倒序：合规项目最新、书稿项目次之，界面上的顺序就固定了。
    _stamp_project(data_dir, fin["project_id"], stamps["showcase-book"])
    _stamp_project(data_dir, book["project_id"], stamps["showcase-fin"])
    tm = core.load_tm(project_id=book["project_id"]) if hasattr(core, "load_tm") else {}
    for r in pairs[:10]:
        core.tm_put(tm, r["source"], r["target"], demo.target_lang)
    core.save_tm(tm, project_id=book["project_id"])
    return {"book": book["project_id"], "fin": fin["project_id"]}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=str(ROOT / "docs" / "assets"))
    parser.add_argument("--work", default=str(ROOT / "tmp" / "qt-showcase"))
    parser.add_argument("--theme", default="light", choices=("light", "dark"))
    parser.add_argument("--lang", default="zh-CN", choices=tuple(LANG_SUFFIX),
                        help="界面语言与演示数据语言（zh-CN 是基准，文件名不加后缀）")
    args = parser.parse_args()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    work = Path(args.work)
    shutil.rmtree(work, ignore_errors=True)
    data_dir = work / "data"
    data_dir.mkdir(parents=True)
    demo = DEMOS[args.lang]

    backup = work / "qsettings-backup.plist"
    had_prefs = subprocess.run(["defaults", "export", "com.folith.QtDesktop", str(backup)],
                               capture_output=True).returncode == 0

    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    os.environ.setdefault("QT_SCALE_FACTOR", "2")
    os.environ["FOLITH_QT_NONINTERACTIVE"] = "1"
    os.environ["FOLITH_EXIT_FORCE_PROMPT"] = "0"
    # 界面语言：环境变量优先于设置里的值（窗口在建视图之前读它），两边都写，
    # 保证这一轮截图从头到尾都是同一门语言。
    os.environ["FOLITH_UI_LANGUAGE"] = args.lang
    from folith.qt import bootstrap

    paths = bootstrap.prepare_environment(data_dir=data_dir, log_dir=work / "logs")
    from folith.qt import adapter
    from folith.qt import window as window_module
    from PySide6.QtCore import QSettings

    ids = seed(adapter.core, data_dir, demo)
    app = window_module.build_application([])
    app.setStyleSheet(app.styleSheet())
    settings = QSettings("Folith", "QtDesktop")
    settings.setValue("appearance/theme", args.theme)
    settings.setValue("appearance/language", args.lang)
    settings.sync()

    def pump(seconds: float) -> None:
        end = time.monotonic() + seconds
        while time.monotonic() < end:
            app.processEvents()
            time.sleep(0.02)

    try:
        config = adapter.EngineConfig(provider="DeepSeek", model="deepseek-v4-flash",
                                      api_key="demo", target_lang=demo.target_lang)
        window = window_module.MainWindow(config, paths, None)
        width, height = (WINDOW_SIZE if args.lang == "zh-CN" else WINDOW_SIZE_LOCALIZED)
        window.resize(width, height)
        window.show()
        pump(2.0)
        suffix = "" if args.theme == "light" else "-dark"
        lang_suffix = LANG_SUFFIX[args.lang]

        def shot(name: str) -> None:
            pump(1.2)
            image = window.grab()
            menu = window.menuBar()
            if menu is not None and menu.isVisible():
                ratio = image.devicePixelRatio()
                top = int(menu.height() * ratio)
                image = image.copy(0, top, image.width(), image.height() - top)
            image.save(str(out / f"{name}{lang_suffix}{suffix}.png"))
            print("saved", out / f"{name}{lang_suffix}{suffix}.png", image.width(), image.height())

        window.show_home()
        shot("qt-01-tasks")
        window.home.manage_projects_button.click()
        pump(1.0)
        window.projects.select(ids["book"])
        pump(1.5)
        window.projects._tab_changed("terms")
        shot("qt-02-projects")
        open_job_and_wait(window, "showcase-book")
        pump(1.5)
        window.workbench.show_page("translate")
        shot("qt-03-workbench")
        window.workbench.show_page("translate")
        window.workbench.show_page("deliver")
        shot("qt-04-delivery-gate")
        open_job_and_wait(window, "showcase-fin")
        pump(1.5)
        window.workbench.show_page("deliver")
        shot("qt-05-delivery-ready")
        open_job_and_wait(window, "showcase-prep")
        pump(1.5)
        window.workbench.show_page("translate")
        shot("qt-06-prepare")
        window.open_settings("models")
        dialog = window._settings_window
        dialog.resize(900, 420)
        pump(1.2)
        image = dialog.grab()
        image.save(str(out / f"qt-07-settings{lang_suffix}{suffix}.png"))
        print("saved", out / f"qt-07-settings{lang_suffix}{suffix}.png", image.width(), image.height())
        window.close()
        pump(0.5)
    finally:
        if had_prefs:
            subprocess.run(["defaults", "import", "com.folith.QtDesktop", str(backup)])
        else:
            subprocess.run(["defaults", "delete", "com.folith.QtDesktop"], capture_output=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
