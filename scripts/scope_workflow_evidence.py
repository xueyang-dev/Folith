"""Native Qt scope evidence, using an isolated DOCX and offline provider.

Run with QT_QPA_PLATFORM=cocoa on macOS or windows on Windows. No real provider,
user projects, settings or credentials are used. This is scripted acceptance,
not a human-operated study or a check of a packaged installer.
"""
from __future__ import annotations

import argparse
import json
import os
import platform
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='folith-scope-') as temporary:
        base = Path(temporary)
        os.environ.update(FOLITH_DESKTOP_DATA_DIR=str(base / 'data'),
                          FOLITH_DESKTOP_LOG_DIR=str(base / 'logs'),
                          FOLITH_OUTPUT_DIR=str(base / 'data'),
                          FOLITH_QT_SETTINGS_DIR=str(base / 'settings'),
                          FOLITH_QT_NONINTERACTIVE='1')
        from PySide6.QtCore import QTimer
        from PySide6.QtWidgets import QApplication
        from docx import Document
        import core
        from folith.desktop import paths
        from folith.qt import adapter, i18n
        from folith.qt.scope_dialog import ScopeDialog
        from folith.qt.selftest import open_job_and_wait
        from folith.qt.window import MainWindow

        app = QApplication.instance() or QApplication([])
        if app.platformName() in ('offscreen', 'minimal'):
            raise SystemExit('Native Qt platform required.')
        i18n.set_language('zh-CN')
        core.OUTPUT_DIR = base / 'data'
        source = base / 'scope.docx'
        doc = Document()
        for title, text in [('第一章', 'Alpha body.'), ('第二章', 'Beta body.')]:
            doc.add_heading(title, level=1)
            doc.add_paragraph(text)
        doc.save(source)
        config = adapter.EngineConfig(enable_report=False, enable_understanding=False,
                                      enable_source_cleanup=False)
        window = MainWindow(config, paths.resolve_paths(), None)
        window.resize(1200, 950)
        window.show()
        job = adapter.import_documents([source], config)[0].job_id
        open_job_and_wait(window, job)
        checked = []
        failures = []

        def dialog_action(restore=False):
            dialog = QApplication.activeModalWidget()
            if not isinstance(dialog, ScopeDialog):
                failures.append('scope dialog did not open')
                return
            try:
                if restore:
                    assert dialog._record_checks
                    for checkbox, _record in dialog._record_checks.values():
                        checkbox.setChecked(True)
                else:
                    key = next(k for k, (_, node) in dialog._node_checks.items()
                               if node['label'] == '第二章')
                    dialog._node_checks[key][0].setChecked(True)
                    assert '将排除所选 2 段' in dialog.summary_label.text()
                app.processEvents()
                name = '02-restore-plan.png' if restore else '01-select-scope.png'
                assert dialog.grab().save(str(output / name))
                (dialog.restore_button if restore else dialog.exclude_button).click()
                checked.append('restore' if restore else 'exclude')
            except Exception as exc:
                failures.append(str(exc))
                dialog.reject()

        try:
            before = adapter.source_size(job)
            for restore in (False, True, False):
                QTimer.singleShot(0, lambda restore=restore: dialog_action(restore))
                window.workbench.setup_panel.source_section.scope_button.click()
                assert not failures, failures
                state = core.load_job_state(job)
                assert not state['paras'] and not state['pairs']
                if restore:
                    assert not state['scope_exclusion_plan']['entries']
                else:
                    assert state['scope_exclusion_plan']['entries'][0]['label'] == '第二章'
                    assert adapter.source_size(job)['segments'] == before['segments'] - 2
            requests = []

            def offline_provider(*args, **kwargs):
                system, user = args[3:5]
                requests.append(user)
                if '学术翻译专家' in system:
                    units = re.findall(r'^\s*\d+\.\s+(.+?)\s*$',
                                       user.rsplit('待翻译段落', 1)[-1], re.MULTILINE)
                    return json.dumps([f'这是第 {i+1} 段中文译文。' for i in range(len(units))])
                return '[]'

            core.call_llm = offline_provider
            state = core.run_job_pipeline(
                job, source.name, source.read_bytes(), provider='DeepSeek', api_key='offline',
                model='deepseek-chat', target_lang='简体中文', auto_term=False,
                enable_report=False, translation_theory='目的论', enable_review=False,
                enable_annotate=False, use_tm=False, enable_understanding=False,
                enable_source_cleanup=False, delivery_config=core.default_delivery_config())
            assert state['p2_done'] and requests
            assert all('Beta body.' not in text and '第二章' not in text for text in requests)
            records = core.excluded_segment_records(state)
            assert len(records) == 2
            assert core.excluded_segments_manifest(state)['origins']['user'] == 2
            targets = [pair['target'] for pair in state['pairs']]
            adapter.restore_scope(job, [r['excluded_id'] for r in records])
            reopened = core.load_job_state(job)
            assert len(reopened['paras']) == len(reopened['pairs']) == 4
            assert [p['target'] for p in reopened['pairs'][:2]] == targets
            assert all(not p['target'] and not p.get('reviewed') for p in reopened['pairs'][2:])
            result = dict(platform=platform.platform(), qt_backend=app.platformName(),
                          synthetic_fixture=True, human_operated=False, real_provider_calls=0,
                          ui_actions=checked, pretranslation_plan=True, estimate_updates=True,
                          excluded_absent_from_model_requests=True, user_skip_manifest=True,
                          restore_preserves_existing_targets=True, reopened_pair_alignment=True)
            (output / 'evidence.json').write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
            print(json.dumps(result, ensure_ascii=False))
        finally:
            window.timer.stop()
            window._closing_accepted = True
            window.close()
            app.processEvents()


if __name__ == '__main__':
    main()
