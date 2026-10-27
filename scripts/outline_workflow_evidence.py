"""Native Qt document-outline panel evidence with a synthetic DOCX (#187).

Run with ``QT_QPA_PLATFORM=cocoa`` on macOS or ``windows`` on Windows.  The
fixture is a generated DOCX, the provider is offline, and the data/settings
directories are isolated, so no project, credential or user preference is
touched.  This is scripted acceptance, not a human-operated study and not a
check of a packaged installer.
"""
from __future__ import annotations

import argparse
import json
import os
import platform
import re
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

HEADINGS = [
    (1, "第一章 项目概况"),
    (2, "1.1 项目背景"),
    (3, "1.1.1 历史沿革"),
    (2, "1.2 目标读者"),
    (1, "第二章 实施计划"),
    (2, "2.1 阶段安排"),
    (3, "2.1.1 第一阶段"),
    (1, "第三章 附录"),
]
BODIES = [
    "本段介绍项目的基本情况。",
    "本段补充项目的范围与边界。",
    "本段说明项目为何启动。",
    "本段回顾此前的相关工作。",
    "本段记录历史沿革中的关键节点。",
    "本段再次确认沿革的结论。",
    "本段描述主要读者群体。",
    "本段说明读者的阅读场景。",
    "本段列出实施计划的总体安排。",
    "本段交代计划的资源前提。",
    "本段说明阶段的划分依据。",
    "本段给出阶段之间的依赖关系。",
    "本段描述第一阶段的交付内容。",
    "本段列出第一阶段的验收标准。",
    "本段汇总附录材料。",
    "本段说明附录的引用方式。",
]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="folith-outline-") as temporary:
        base = Path(temporary)
        os.environ.update(
            FOLITH_DESKTOP_DATA_DIR=str(base / "data"),
            FOLITH_DESKTOP_LOG_DIR=str(base / "logs"),
            FOLITH_OUTPUT_DIR=str(base / "data"),
            FOLITH_QT_SETTINGS_DIR=str(base / "settings"),
            FOLITH_QT_NONINTERACTIVE="1",
            FOLITH_UI_LANGUAGE="zh-CN",
        )
        os.environ.pop("FOLITH_EXIT_FORCE_PROMPT", None)
        from docx import Document
        from PySide6.QtWidgets import QApplication

        import core
        from folith import outline_anchor
        from folith.desktop import paths
        from folith.qt import adapter, i18n, onboarding, platform_ui, preferences
        from folith.qt.outline_panel import NODE_ROLE
        from folith.qt.selftest import open_job_and_wait
        from folith.qt.window import MainWindow

        app = QApplication.instance() or QApplication([])
        if app.platformName() in ("offscreen", "minimal"):
            raise SystemExit("Native Qt platform required.")
        i18n.set_language("zh-CN")
        core.OUTPUT_DIR = base / "data"
        onboarding.mark_completed(preferences.settings(),
                                  adapter.app_identity()["app_version"])

        source = base / "outline-evidence.docx"
        document = Document()
        for level, title in HEADINGS:
            document.add_heading(title, level=level)
            for text in BODIES[:2]:
                document.add_paragraph(text)
        document.save(source)

        config = adapter.EngineConfig(
            enable_report=False, enable_understanding=False,
            enable_source_cleanup=False)
        window = MainWindow(config, paths.resolve_paths(), None)
        window.resize(1360, 900)
        window.show()
        job = adapter.import_documents([source], config)[0].job_id

        state = core.load_job_state(job)
        assert not state["paras"] and not state["pairs"]

        def offline_provider(*call_args, **_kwargs):
            system, user = call_args[3:5]
            if "学术翻译专家" in system:
                units = re.findall(r"^\s*\d+\.\s+(.+?)\s*$",
                                   user.rsplit("待翻译段落", 1)[-1], re.MULTILINE)
                return json.dumps([f"这是第 {index + 1} 段中文译文。"
                                   for index in range(len(units))])
            return "[]"

        core.call_llm = offline_provider
        core.run_job_pipeline(
            job, source.name, source.read_bytes(), provider="DeepSeek",
            api_key="offline", model="deepseek-chat", target_lang="简体中文",
            auto_term=False, enable_report=False, translation_theory="目的论",
            enable_review=False, enable_annotate=False, use_tm=False,
            enable_understanding=False, enable_source_cleanup=False,
            delivery_config=core.default_delivery_config())
        state = core.load_job_state(job)
        assert state["p2_done"] and state["pairs"]

        open_job_and_wait(window, job)
        bench = window.workbench
        panel = bench.outline_panel

        def pump(predicate, timeout=20.0):
            deadline = time.monotonic() + timeout
            while time.monotonic() < deadline:
                app.processEvents()
                if predicate():
                    return
                time.sleep(0.02)
            raise AssertionError("native outline operation timed out")

        def shot(name):
            app.processEvents()
            assert window.grab().save(str(output / name))
            return name

        records = {"platform": platform.platform(),
                   "qt_backend": app.platformName(),
                   "synthetic_fixture": True,
                   "human_operated": False,
                   "real_provider_calls": 0}

        records["closed_by_default"] = not bench.outline_open()
        records["toolbar_visible"] = bench.outline_button.isVisible()
        from PySide6.QtGui import QKeySequence

        shortcut = QKeySequence(platform_ui.shortcut("outline"))
        hint = platform_ui.shortcut_text("outline")
        records["shortcut_binding"] = shortcut.toString()
        records["shortcut_hint"] = hint
        records["tooltip_shows_hint"] = hint in bench.outline_button.toolTip()
        records["action_shortcut_matches"] = (
            bench.outline_shortcut.shortcut().toString() == shortcut.toString())
        records["sidecar_written_by_open_worker"] = (
            outline_anchor.outline_sidecar_path(job).is_file())
        records["open_parse"] = {
            "parsed": bool(bench._outline_payload.get("parsed")),
            "invalidated": str(bench._outline_payload.get("invalidated") or ""),
        }

        bench.outline_button.click()
        pump(lambda: bench.outline_open()
             and panel.isVisible() and bench.left_stack.width() > 200)
        nodes = [node for node in
                 (bench._outline_payload.get("outline") or {}).get("nodes") or []
                 if isinstance(node, dict)]
        records["nodes"] = len(nodes)
        records["top_level_items"] = panel.tree.topLevelItemCount()
        records["tree_built_lazily"] = (
            panel.tree.topLevelItemCount() < len(nodes))
        records["max_level"] = max(int(node.get("level") or 0) for node in nodes)
        shot("01-outline-open.png")

        by_title = {str(node.get("title")): node for node in nodes}
        entries = (bench._outline_mapping or {}).get("nodes") or {}
        target = by_title["1.1 项目背景"]
        target_entry = entries.get(str(target.get("id"))) or {}
        assert target_entry.get("state") in ("resolved", "partial")
        panel.activate_node(str(target.get("id")))
        pump(lambda: bench._current_index == target_entry.get("first_segment_index"))
        records["node_click_navigates"] = True
        records["clicked_segment"] = bench._current_index
        shot("02-node-navigation.png")

        implementation = by_title["2.1 阶段安排"]
        implementation_entry = entries.get(str(implementation.get("id"))) or {}
        bench.select_segment(int(implementation_entry["first_segment_index"]))
        pump(lambda: panel.tree.currentItem() is not None)

        def owning_node(index):
            best = None
            for node_id, entry in entries.items():
                if entry.get("state") not in ("resolved", "partial"):
                    continue
                first = entry.get("first_segment_index")
                last = entry.get("last_segment_index", first)
                if isinstance(first, int) and first <= index <= int(last):
                    if best is None or first > best[0]:
                        best = (first, str(node_id))
            return best[1] if best else ""

        current_item = panel.tree.currentItem()
        highlighted = (str(current_item.data(0, NODE_ROLE) or "")
                       if current_item is not None else "")
        records["grid_selection_highlights_owner"] = (
            highlighted == owning_node(bench._current_index))
        records["highlighted_node"] = highlighted
        shot("03-grid-highlight.png")

        bench.search_edit.setText("第一章")
        pump(lambda: bool(bench.search_edit.text()))
        filtered_target = entries.get(str(by_title["2.1 阶段安排"].get("id"))) or {}
        assert bench.select_segment(int(filtered_target["first_segment_index"])) is False
        panel.activate_node(str(by_title["2.1 阶段安排"].get("id")))
        pump(lambda: not bench.search_edit.text())
        notice = str(bench._task_notice[1] or "")
        records["filtered_target_clears_filter"] = "已清除筛选" in notice
        records["filter_outcome_notice"] = notice
        records["filtered_click_segment"] = bench._current_index
        shot("04-filter-cleared.png")

        bench.outline_button.click()
        pump(lambda: not bench.outline_open())
        records["closed_state_saved"] = (
            bench.settings.value(bench.OUTLINE_OPEN_KEY, False, type=bool) is False)
        bench.outline_button.click()
        pump(lambda: bench.outline_open())
        records["reopened_state_saved"] = (
            bench.settings.value(bench.OUTLINE_OPEN_KEY, False, type=bool) is True)
        before = core.job_state_path(job).read_text(encoding="utf-8")
        records["preference_not_task_data"] = "outline_open" not in before
        preferences.settings().sync()

        window.timer.stop()
        window._closing_accepted = True
        window.close()
        app.processEvents()

        second = MainWindow(config, paths.resolve_paths(), None)
        second.resize(1360, 900)
        second.show()
        open_job_and_wait(second, job)
        second_bench = second.workbench
        records["preference_survives_reopen"] = second_bench.outline_open()
        records["second_open_parsed"] = bool(
            second_bench._outline_payload.get("parsed"))
        records["second_open_invalidated"] = str(
            second_bench._outline_payload.get("invalidated") or "")
        pump(lambda: second_bench.outline_open()
             and second_bench.outline_panel.isVisible()
             and second_bench.left_stack.width() > 200)
        app.processEvents()
        assert second.grab().save(str(output / "05-reopen-remembered.png"))
        second.timer.stop()
        second._closing_accepted = True
        second.close()
        app.processEvents()

        (output / "evidence.json").write_text(
            json.dumps(records, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8")
        print(json.dumps(records, ensure_ascii=False))


if __name__ == "__main__":
    main()
