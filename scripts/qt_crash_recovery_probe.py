#!/usr/bin/env python3
"""Qt recovery: hard-kill after durable progress, reopen, resume and export.

Roadmap §6.C: "强制结束后既有检查点可恢复" and §6.A.4: the task must show a
credible recovery entry point.  This is done at the *process* level — the client
is killed with SIGKILL/TerminateProcess while a real task is in flight,
then a fresh window is opened on the same data directory.

Two phases in one script:

    --phase run     child: build the window, import, start the task, keep pumping
    --phase verify  parent: kill the child mid-flight, reopen, assert and resume

Usage:
    QT_QPA_PLATFORM=offscreen venv/bin/python scripts/qt_crash_recovery_probe.py \
        --data-dir /tmp/qt-crash/data --out /tmp/qt-crash/report.json

For a native long-document run, add ``--window --exercise-lifecycle
--export-working --fixture <document>``. The driver verifies live workbench and
card progress, pauses through in-flight commits, continues while minimized,
and saves an untouched ``after-kill`` copy before recovery. Committed pair
fields must survive, committed batches must not be regenerated, and every
non-empty target must occur in the working DOCX in source order. Source content
and logs stay in the local evidence directory. QSettings are isolated too.

Only an explicit ``--connection-data-dir <existing Folith data directory>``
uses its saved real API connections. It reads that configuration without
copying credentials into the evidence directory; the default uses a local stub.
"""
from __future__ import annotations

import argparse
import dataclasses
import json
import os
import shutil
import subprocess
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from folith.qt.selftest import open_job_and_wait  # noqa: E402

ACTIVE = {"resume_requested", "queued", "starting", "running",
          "waiting_external", "cancelling", "pausing", "paused"}


class UiHeartbeat:
    """Measure event-loop gaps; this is evidence, not an input-latency claim."""

    def __init__(self, app):
        from PySide6.QtCore import QTimer
        self.gaps = []
        self.last = time.monotonic()
        self.timer = QTimer(app)
        self.timer.setInterval(50)
        self.timer.timeout.connect(self._tick)
        self.timer.start()

    def _tick(self):
        now = time.monotonic()
        self.gaps.append(now - self.last)
        self.last = now

    def snapshot(self, *, start=0):
        gaps = sorted(self.gaps[start:] + [time.monotonic() - self.last])
        return {"timer_interval_ms": 50, "ticks": len(gaps) - 1,
                "max_gap_ms": round(1000 * gaps[-1], 2),
                "p95_gap_ms": round(1000 * gaps[int(.95 * (len(gaps) - 1))], 2)}


class Stub:
    def __init__(self, out_dir: Path, delay: float = 0.6) -> None:
        self.out_dir = out_dir
        self.delay = delay
        self.proc: subprocess.Popen | None = None
        self.url = ""
        self._handle = None

    def start(self) -> str:
        ready = self.out_dir / "stub-ready.json"
        ready.unlink(missing_ok=True)
        self._handle = (self.out_dir / "stub-server.log").open("ab")
        self.proc = subprocess.Popen(
            [sys.executable, str(ROOT / "scripts" / "d0_fake_llm_server.py"),
             "--port", "0", "--log", str(self.out_dir / "stub-llm.jsonl"),
             "--ready-file", str(ready), "--delay", str(self.delay)],
            stdout=self._handle, stderr=subprocess.STDOUT)
        deadline = time.monotonic() + 60
        while time.monotonic() < deadline:
            if ready.exists():
                self.url = json.loads(ready.read_text(encoding="utf-8"))["url"]
                return self.url
            if self.proc.poll() is not None:
                raise RuntimeError("stub server exited")
            time.sleep(0.1)
        raise RuntimeError("stub server did not become ready")

    def stop(self) -> None:
        if self.proc is not None and self.proc.poll() is None:
            self.proc.terminate()
            try:
                self.proc.wait(timeout=10)
            except subprocess.TimeoutExpired:  # pragma: no cover
                self.proc.kill()
        if self._handle is not None:
            self._handle.close()


def _resume_finished(adapter, job_id):
    # Poll only the small runtime record while the worker is active. Replaying
    # and copying a full book every 50 ms makes the driver itself stall Qt.
    runtime = adapter.core.load_runtime_state(job_id) or {}
    if runtime.get("status") in adapter.core.RUNTIME_ACTIVE_STATUSES:
        return False
    return bool((adapter.load_state(job_id) or {}).get("p2_done"))


def _quiet_window(app, window, heartbeat, seconds=3.0):
    """Measure real window work without a full-state driver predicate."""
    first_gap = len(heartbeat.gaps)
    started, cpu = time.monotonic(), time.process_time()
    before = len(window.workbench._state.get("pairs") or [])
    while time.monotonic() - started < seconds:
        app.processEvents()
        time.sleep(0.01)
    wall, used = time.monotonic() - started, time.process_time() - cpu
    gaps = heartbeat.gaps[first_gap:] + [time.monotonic() - heartbeat.last]
    return {"wall_s": round(wall, 4), "process_cpu_s": round(used, 4),
            "process_cpu_percent_one_core": round(100 * used / wall, 2),
            "max_gap_ms": round(1000 * max(gaps, default=0), 2),
            "visible_pairs_before": before,
            "visible_pairs_after": len(window.workbench._state.get("pairs") or []),
            "scope": "driver only pumps Qt; CPU includes the real core worker"}


def _prepare_env(data_dir: Path, log_dir: Path, *, native: bool = False) -> None:
    if native:
        os.environ.pop("QT_QPA_PLATFORM", None)
    else:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    os.environ["FOLITH_QT_NONINTERACTIVE"] = "1"
    os.environ["FOLITH_EXIT_FORCE_PROMPT"] = "0"
    os.environ["FOLITH_DESKTOP_DATA_DIR"] = str(data_dir)
    os.environ["FOLITH_DESKTOP_LOG_DIR"] = str(log_dir)
    os.environ["FOLITH_OUTPUT_DIR"] = str(data_dir)
    # The stored task defaults belong to whoever owns this machine; a leftover
    # task/enable_report would put the task into the report stage and change what
    # the recovery run proves.  The child inherits this directory, so both
    # windows in a recovery run share one private preference file.
    os.environ["FOLITH_QT_SETTINGS_DIR"] = str(log_dir / "settings")


def _engine_config(adapter, connection_data_dir=None, stub_url=""):
    """Read an explicitly selected saved connection without saving credentials."""
    if connection_data_dir:
        previous = adapter.core.OUTPUT_DIR
        try:
            adapter.core.OUTPUT_DIR = Path(connection_data_dir).expanduser().resolve()
            config = adapter.EngineConfig.load()
        finally:
            adapter.core.OUTPUT_DIR = previous
        if config.connection_problem():
            raise ValueError("The saved model connection is incomplete")
    else:
        config = adapter.EngineConfig(
            provider="自定义中转站", model="stub-model", api_key="stub-key",
            base_url=stub_url, target_lang="简体中文")
    return dataclasses.replace(config, enable_review=True, enable_report=False,
                               enable_understanding=False, enable_source_cleanup=False)


def _private_window(w, adapter, config, log_dir):
    """The probe must not read or overwrite the operator's task preferences.

    ``FOLITH_QT_SETTINGS_DIR`` (set in ``_prepare_env``) already redirects every
    preference reader, so the window is built normally here.
    """
    return w.MainWindow(config, adapter.desktop_paths.resolve_paths(), None)


def _pump(app, seconds, predicate):
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        app.processEvents()
        if predicate():
            return True
        time.sleep(0.05)
    return bool(predicate())


def _durable_state(data_dir, job_id):
    """Read a stable durability prefix without changing any task files."""
    from folith.storage import state_journal as journal
    state = json.loads((data_dir / job_id / "state.json").read_text(encoding="utf-8"))
    start = journal.cursor(state)
    records = journal.read_mutations(job_id, after_seq=start, output_dir=data_dir)
    if records:
        sequences = [record["seq"] for record in records]
        if sequences != list(range(start + 1, sequences[-1] + 1)):
            raise AssertionError("journal sequence is not a complete ordered suffix")
    report = journal.replay(job_id, state, output_dir=data_dir)
    if report["gap"] or journal.cursor(state) < journal.max_seq(job_id, output_dir=data_dir):
        raise AssertionError(f"journal replay did not reach its durable tail: {report}")
    return state, report


def _assert_committed_prefix_preserved(expected, current):
    """All committed fields survive; resume may add derived source references."""
    assert len(current) >= len(expected), "resume lost committed pairs"
    added_fields = set()
    for old, new in zip(expected, current):
        for key, value in old.items():
            assert key in new and new[key] == value, f"resume changed committed field {key}"
        added_fields.update(new.keys() - old.keys())
    return sorted(added_fields)


def _paused_state_after_inflight_commits(app, adapter, job_id):
    """Already received batches may still commit after the pause request."""
    state = adapter.load_state(job_id)
    quiet_since = time.monotonic()
    deadline = quiet_since + 60
    while time.monotonic() < deadline:
        app.processEvents()
        latest = adapter.load_state(job_id)
        if latest.get("pairs") != state.get("pairs"):
            state = latest
            quiet_since = time.monotonic()
        if time.monotonic() - quiet_since >= 1.5:
            return latest
        time.sleep(0.1)
    raise AssertionError("committed pairs did not settle while paused")


def run_child(args) -> int:
    """Import a document, start the task, and keep the window alive."""
    data_dir = Path(args.data_dir).expanduser().resolve()
    log_dir = Path(args.log_dir).expanduser().resolve()
    _prepare_env(data_dir, log_dir, native=args.window)
    from folith.qt import adapter, selftest, window as w

    app = w.build_application([])
    heartbeat = UiHeartbeat(app)
    config = _engine_config(adapter, args.connection_data_dir,
                            os.environ.get("FOLITH_QT_CRASH_STUB_URL", ""))
    window = _private_window(w, adapter, config, log_dir)
    window.show()
    fixture = Path(args.fixture).expanduser()
    window.import_paths([str(fixture)], autostart=False)
    deadline = time.monotonic() + 120
    while time.monotonic() < deadline:
        app.processEvents()
        jobs = adapter.list_jobs()
        if jobs and adapter.core.load_source(jobs[0].job_id):
            break
        time.sleep(0.05)
    else:
        raise SystemExit("child: import did not complete")
    job_id = adapter.list_jobs()[0].job_id
    open_job_and_wait(window, job_id)
    window.start_or_resume()
    (Path(args.log_dir) / "child-job.txt").write_text(job_id, encoding="utf-8")
    # Keep the process alive while the parent watches for durable progress.
    lifecycle_done = not args.exercise_lifecycle
    next_check = 0.0
    # A real provider may spend several minutes finishing in-flight requests.
    # This is a probe coordination budget, not a product timeout change.
    control_budget = min(args.timeout, 600.0) if args.connection_data_dir else 60.0
    while True:
        app.processEvents()
        if not lifecycle_done and time.monotonic() >= next_check:
            next_check = time.monotonic() + 0.2
            current = window.workbench._state
            if len(current.get("pairs") or []) >= args.min_committed_pairs:
                worker = adapter.core._RUNTIME_WORKERS.get(job_id)
                assert worker is not None and worker.is_alive()
                window.pause_job()
                assert _pump(app, control_budget, lambda: (
                    adapter.core.load_runtime_state(job_id) or {}).get("status") == "paused")
                paused = _paused_state_after_inflight_commits(app, adapter, job_id)
                pairs = paused.get("pairs") or []
                # The normal poll, not an unconditional reopen, must catch up.
                assert _pump(app, 15, lambda: window.workbench._state.get("pairs") == pairs)
                card = next(row for row in adapter.list_jobs() if row.job_id == job_id)
                translated = sum(bool(str(pair.get("target") or "").strip()) for pair in pairs)
                total = max(len(paused.get("paras") or []), len(pairs))
                assert card.progress_percent == round(100 * translated / total)
                assert not paused.get("p2_done"), "task completed before the lifecycle checkpoint"
                window.grab().save(str(log_dir / "paused-window.png"))
                until = time.monotonic() + 1.0
                _pump(app, 1.1, lambda: time.monotonic() >= until)
                assert adapter.load_state(job_id).get("pairs") == pairs
                window.showMinimized()
                assert _pump(app, 5, window.isMinimized)
                window.continue_job()
                assert adapter.core._RUNTIME_WORKERS.get(job_id) is worker
                assert _pump(app, control_budget, lambda: len(
                    window.workbench._state.get("pairs") or []) > len(pairs))
                lifecycle = {"paused_pairs": len(pairs), "stable_while_paused": True,
                             "workbench_caught_up": True, "card_percent": card.progress_percent,
                             "minimized": window.isMinimized(),
                             "progressed_while_minimized": True,
                             "same_worker_after_continue": True,
                             "ui_heartbeat": heartbeat.snapshot(),
                             "qt_platform": app.platformName()}
                assert lifecycle["minimized"]
                (log_dir / "lifecycle.json").write_text(json.dumps(lifecycle), encoding="utf-8")
                lifecycle_done = True
        time.sleep(0.02)
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Qt crash recovery probe")
    parser.add_argument("--data-dir", required=True)
    parser.add_argument("--log-dir", default=None)
    parser.add_argument("--out", default=None)
    parser.add_argument("--fixture", default=None)
    parser.add_argument("--phase", default="verify", choices=("run", "verify"))
    parser.add_argument("--paragraphs", type=int, default=60)
    parser.add_argument("--timeout", type=float, default=600.0)
    parser.add_argument("--window", action="store_true", help="use the native Qt platform")
    parser.add_argument("--stub-delay", type=float, default=0.6)
    parser.add_argument("--connection-data-dir", default=None,
                        help="explicitly use real saved API connections from this Folith data directory")
    parser.add_argument("--min-committed-pairs", type=int, default=1)
    parser.add_argument("--exercise-lifecycle", action="store_true",
                        help="pause, verify live progress, continue while minimized, then kill")
    parser.add_argument("--export-working", action="store_true",
                        help="export working assets instead of freezing a final delivery")
    args = parser.parse_args(argv)

    data_dir = Path(args.data_dir).expanduser().resolve()
    log_dir = Path(args.log_dir or (data_dir.parent / "logs")).expanduser().resolve()
    log_dir.mkdir(parents=True, exist_ok=True)
    if args.phase == "run":
        return run_child(args)
    if data_dir.exists() and any(data_dir.iterdir()):
        parser.error("use an empty data directory; existing tasks will not be changed")
    _prepare_env(data_dir, log_dir, native=args.window)

    out_dir = Path(args.out).parent if args.out else data_dir.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    if (out_dir / "after-kill").exists() or any(
            (log_dir / name).exists() for name in ("child-job.txt", "lifecycle.json")):
        parser.error("use a fresh evidence directory; previous results will not be overwritten")
    report: dict = {"started_at": time.time(), "steps": []}
    heartbeat = None

    def step(name: str, **extra) -> None:
        report["steps"].append(
            {"at": round(time.time() - report["started_at"], 2),
             "step": name, **extra})

    stub = None
    from folith.qt import adapter
    if args.connection_data_dir:
        config = _engine_config(adapter, args.connection_data_dir)
        url = ""
        report["connection"] = {"kind": "saved", "provider": config.provider,
                                "model": config.model}
    else:
        stub = Stub(out_dir, delay=args.stub_delay)
        url = stub.start()
        config = _engine_config(adapter, stub_url=url)
        report["stub_url"] = url
    fixture = Path(args.fixture).expanduser() if args.fixture else out_dir / "crash.docx"
    if not fixture.exists():
        from folith.qt import selftest

        selftest.make_docx(fixture, [f"第 {n} 段，术语 MT。" for n in range(args.paragraphs)])

    env = dict(os.environ)
    env["FOLITH_QT_CRASH_STUB_URL"] = url
    child_command = [sys.executable, str(Path(__file__).resolve()),
         "--phase", "run", "--data-dir", str(data_dir), "--log-dir", str(log_dir),
         "--fixture", str(fixture), "--paragraphs", str(args.paragraphs),
         "--min-committed-pairs", str(args.min_committed_pairs)]
    if args.window:
        child_command.append("--window")
    if args.exercise_lifecycle:
        child_command.append("--exercise-lifecycle")
    if args.connection_data_dir:
        child_command.extend(["--connection-data-dir",
                              str(Path(args.connection_data_dir).expanduser().resolve())])
    child = subprocess.Popen(child_command,
        env=env, stdout=(log_dir / "child.log").open("ab"),
        stderr=subprocess.STDOUT)
    report["child_pid"] = child.pid
    try:
        job_file = log_dir / "child-job.txt"
        deadline = time.monotonic() + 180
        while time.monotonic() < deadline and not job_file.exists():
            if child.poll() is not None:
                raise AssertionError(
                    "child exited before starting a task: "
                    + (log_dir / "child.log").read_text(errors="replace")[-500:])
            time.sleep(0.2)
        job_id = job_file.read_text(encoding="utf-8").strip()
        report["job_id"] = job_id

        # Wait until the task is running, then let it make progress before the
        # kill, so the crash lands *inside* the pipeline rather than at startup.
        # Progress during translation is reported in `stage_progress`, not in
        # `completed_units`, so "active" is the honest signal here.
        runtime_path = data_dir / job_id / "runtime_state.json"

        def read_runtime() -> dict:
            try:
                return json.loads(runtime_path.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                return {}

        runtime: dict = {}
        deadline = time.monotonic() + 180
        while time.monotonic() < deadline:
            runtime = read_runtime()
            if runtime.get("status") in ACTIVE:
                break
            if child.poll() is not None:
                raise AssertionError("child exited while the task was starting")
            time.sleep(0.2)
        else:
            raise AssertionError(
                f"the task never became active: {runtime.get('status')!r}")
        # An active worker alone is insufficient: force-kill recovery must
        # preserve actual committed results, not just restart an empty task.
        deadline = time.monotonic() + args.timeout
        while time.monotonic() < deadline:
            if child.poll() is not None:
                raise AssertionError("child exited before durable progress: " +
                                     (log_dir / "child.log").read_text(errors="replace")[-1000:])
            try:
                observed, _ = _durable_state(data_dir, job_id)
            except (OSError, ValueError, AssertionError):
                observed = {}
            lifecycle_ready = not args.exercise_lifecycle or (log_dir / "lifecycle.json").is_file()
            if len(observed.get("pairs") or []) >= args.min_committed_pairs and lifecycle_ready:
                break
            if observed.get("p2_done"):
                raise AssertionError("task completed before the requested kill checkpoint")
            time.sleep(0.2)
        else:
            raise AssertionError("no durable translation checkpoint before timeout")
        if args.exercise_lifecycle:
            report["lifecycle"] = json.loads((log_dir / "lifecycle.json").read_text())
        runtime = read_runtime()
        report["before_kill"] = {
            "status": runtime.get("status"),
            "completed_units": runtime.get("completed_units"),
            "total_units": runtime.get("total_units"),
            "stage_progress": runtime.get("stage_progress"),
            "phase": runtime.get("phase"),
            "still_active": runtime.get("status") in ACTIVE,
            "observed_committed_pairs": len(observed["pairs"]),
        }
        step("mid-flight", **report["before_kill"])
        if runtime.get("status") not in ACTIVE:
            raise AssertionError(
                "the task finished before it could be killed; widen the window "
                f"(status={runtime.get('status')!r})")

        # Hard kill: no cleanup, no chance to flush anything.
        child.kill()  # SIGKILL on POSIX, TerminateProcess on Windows.
        child.wait(timeout=30)
        report["child_exit"] = child.returncode
        step("killed", exit_code=child.returncode)
        shutil.copytree(data_dir, out_dir / "after-kill")
        expected, replay = _durable_state(data_dir, job_id)
        expected_pairs = expected.get("pairs") or []
        assert len(expected_pairs) >= args.min_committed_pairs
        report["durable_after_kill"] = {"pairs": len(expected_pairs), "replay": replay,
                                        "segments": len(expected.get("paras") or [])}
        events_path = data_dir / job_id / "events.jsonl"
        events_offset = events_path.stat().st_size if events_path.exists() else 0

        # Reopen on the same data directory with a fresh window.
        _prepare_env(data_dir, log_dir, native=args.window)
        from folith.qt import adapter, window as w

        app = w.build_application([])
        heartbeat = UiHeartbeat(app)
        window = _private_window(w, adapter, config, log_dir)
        window.show()
        app.processEvents()

        def pump(seconds: float, predicate=None, interval: float = 0.05) -> bool:
            deadline = time.monotonic() + seconds
            while time.monotonic() < deadline:
                app.processEvents()
                if predicate is not None and predicate():
                    return True
                time.sleep(interval)
            app.processEvents()
            return bool(predicate()) if predicate is not None else True

        window.refresh_home()
        state = adapter.load_state(job_id)
        assert state.get("pairs") == expected_pairs, "reopen changed committed pairs"
        report["committed_pairs_equal_after_reopen"] = True
        report["qt_platform"] = app.platformName()
        runtime = adapter.runtime_view(job_id, state)
        report["after_reopen"] = {
            "status": runtime.get("status"),
            "status_label": runtime.get("status_label"),
            "p1_done": bool(state.get("p1_done")),
            "p2_done": bool(state.get("p2_done")),
            "available_actions": runtime.get("available_actions"),
        }
        step("reopened", **report["after_reopen"])
        # The job must be reported as interrupted (or resumable), never complete.
        assert not (state.get("p1_done") and state.get("p2_done")), \
            "a killed task was reported as complete"
        assert runtime.get("status") in ("interrupted", "stalled",
                                         "idle_incomplete", "failed"), \
            f"unexpected status after a crash: {runtime.get('status')}"

        opened_at = time.monotonic()
        open_job_and_wait(window, job_id)
        report["open_job_call_s"] = round(time.monotonic() - opened_at, 4)
        app.processEvents()
        report["recovery_entry"] = {
            "retry_visible": window.workbench.retry_button.isVisible(),
            "start_text": window.workbench.start_button.text(),
            "status_text": window.workbench.status_label.text(),
        }
        step("recovery-entry", **report["recovery_entry"])

        # Resume through the same entry point a user would use.
        resume_gap_start = len(heartbeat.gaps)
        window.start_or_resume()
        report["quiet_running_window"] = _quiet_window(app, window, heartbeat)
        report["resumed"] = pump(args.timeout, lambda: _resume_finished(adapter, job_id))
        # Capture before the driver's full-state comparisons and direct export.
        report["ui_heartbeat_during_resume"] = heartbeat.snapshot(start=resume_gap_start)
        state = adapter.load_state(job_id)
        report["added_pair_fields_after_resume"] = _assert_committed_prefix_preserved(
            expected_pairs, state.get("pairs") or [])
        report["committed_fields_preserved_after_resume"] = True
        assert len(state.get("pairs") or []) == len(state.get("paras") or []), \
            "completed task does not contain every source segment"
        with events_path.open("rb") as stream:
            stream.seek(events_offset)
            new_events = [json.loads(line) for line in stream if line.strip()]
        generation_offsets = [event["offset"] for event in new_events
                              if event.get("phase") == "generation_started"]
        assert all(offset >= len(expected_pairs) for offset in generation_offsets), \
            "resume regenerated an already committed batch"
        report["committed_batches_not_regenerated"] = True
        report["after_resume"] = {
            "p1_done": bool(state.get("p1_done")),
            "p2_done": bool(state.get("p2_done")),
            "pairs": len(state.get("pairs") or []),
            "reviewed": sum(1 for pair in state.get("pairs") or []
                            if pair.get("reviewed")),
        }
        step("resumed", **report["after_resume"])

        # And the recovered task can be delivered.
        if args.export_working:
            written = adapter.export_current_assets(job_id, out_dir / "exported")
            assert written and all(path.is_file() and path.stat().st_size for path in written)
            from docx import Document
            docx_name = adapter.delivery_file_name(state.get("filename"), "translation.docx")
            docx = next(path for path in written if path.name == docx_name)
            text = "\n".join(paragraph.text for paragraph in Document(docx).paragraphs)
            position = 0
            targets = [str(pair.get("target") or "") for pair in state.get("pairs") or []]
            for target in targets:
                if not target:
                    continue
                found = text.find(target, position)
                assert found >= 0, "exported DOCX omitted or reordered a committed target"
                position = found + len(target)
            report["exported_targets_in_source_order"] = sum(bool(target) for target in targets)
            report["delivery"] = {"kind": "working export", "files": [p.name for p in written]}
        else:
            window.deliver()
            pump(180, lambda: bool(adapter.delivery_snapshots(job_id)))
            report["delivery"] = {
                "kind": "final snapshot",
                "snapshots": [item.get("snapshot_version")
                              for item in adapter.delivery_snapshots(job_id)],
            }
        step("delivered", **report["delivery"])
        report["ok"] = bool(report["resumed"]
                            and report["after_resume"]["p2_done"]
                            and (report["delivery"].get("files")
                                 or report["delivery"].get("snapshots")))
    except BaseException as exc:  # noqa: BLE001 - the report is the evidence
        import traceback

        report["ok"] = False
        report["error"] = f"{type(exc).__name__}: {exc}"
        report["traceback"] = traceback.format_exc()
    finally:
        if heartbeat is not None:
            report["ui_heartbeat"] = heartbeat.snapshot()
            heartbeat.timer.stop()
        if stub is not None:
            stub.stop()
        if child.poll() is None:
            child.kill()
        report.setdefault("ok", False)
        report["finished_at"] = time.time()
        report["duration_s"] = round(report["finished_at"] - report["started_at"], 1)

    text = json.dumps(report, ensure_ascii=False, indent=2, default=str)
    for secret in {config.api_key, config.reviewer_api_key,
                   *(str(c.get("api_key") or "") for c in config.connections)}:
        if secret:
            text = text.replace(secret, "[REDACTED]")
    if args.out:
        Path(args.out).write_text(text + "\n", encoding="utf-8")
    print(text, flush=True)
    return 0 if report.get("ok") else 1


if __name__ == "__main__":
    sys.exit(main())
