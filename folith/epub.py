"""Reflowable translation EPUBs built from source order and current targets.

Media never goes through the model. Source paragraph references survive CAT
structure edits; image positions are paragraph boundaries, not target offsets.
"""
from __future__ import annotations

import collections
import io
import re
import zipfile
from copy import deepcopy
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from docx.oxml.ns import qn

from . import pdf_ingestion, segmentation, translation_target
from .textual import text_direction, language_tag

SOURCE_REFS = "epub_source_paragraphs"


class EpubExportError(RuntimeError):
    pass


def _text(value: Any) -> str:
    # XML 1.0 characters, including supplementary-plane Unicode.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f\ud800-\udfff\ufffe\uffff]", "", str(value or ""))


def _normal(value: Any) -> str:
    return " ".join(_text(value).split())


def _bind_from_paragraph_ranges(pairs, origins, paragraph_ranges, cleaned) -> bool:
    """Attach raw-paragraph references from the persisted segment→paragraph map.

    ``segmentation.paragraph_ranges`` is the same positional contract the CAT
    list itself uses, so an exclusion that maintains it (the segment leaves the
    map with the pair) can still bind images.  The map is only trusted when it
    still describes the current text: joining a paragraph's segments has to
    reproduce that cleaned paragraph.  A split or merge breaks that, so those
    keep failing closed instead of binding an image to the wrong paragraph.
    """
    if not isinstance(paragraph_ranges, list) or not paragraph_ranges:
        return False
    owner: dict = {}
    by_paragraph: dict = {}
    for row in paragraph_ranges:
        if not isinstance(row, dict):
            return False
        start, end = row.get("start"), row.get("end")
        paragraph_index = row.get("paragraph_index")
        if isinstance(start, bool) or isinstance(end, bool) \
                or not isinstance(start, int) or not isinstance(end, int) \
                or not isinstance(paragraph_index, int):
            return False
        if not 0 <= paragraph_index < len(origins) or start > end:
            return False
        indexes = []
        for index in range(max(0, start), min(end, len(pairs))):
            if index in owner:
                return False
            owner[index] = paragraph_index
            indexes.append(index)
        by_paragraph[paragraph_index] = indexes
    if set(owner) != set(range(len(pairs))):
        return False
    for paragraph_index, indexes in by_paragraph.items():
        if not indexes:
            continue
        joined = _normal(" ".join(str(pairs[index].get("source") or "")
                                  for index in indexes))
        if paragraph_index >= len(cleaned) \
                or joined != _normal(cleaned[paragraph_index]):
            return False
    changed = False
    for index, paragraph_index in owner.items():
        if SOURCE_REFS not in pairs[index]:
            pairs[index][SOURCE_REFS] = list(dict.fromkeys(origins[paragraph_index]))
            changed = True
    return changed


#: Inputs for which :func:`bind_source_references` already found that the pairs
#: cannot be bound (exact keys, most recent last).
_UNBINDABLE: "collections.OrderedDict[tuple, None]" = collections.OrderedDict()


def bind_source_references(state: dict) -> bool:
    """Attach raw paragraph ordinals before an edit changes CAT ordinals.

    Only a verified, unedited initial segmentation can be backfilled. Existing
    references are copied by edits and never recomputed from current ordinals.
    No source parsing or filesystem writes are needed here.

    An exclusion is not "unedited" — the pair is gone — but its positional map
    is maintained, so it is bound from ``segmentation.paragraph_ranges``
    instead of being refused: otherwise excluding one structural fragment would
    silently cost the task its image/text association for every other paragraph.
    """
    if not isinstance(state, dict):
        return False
    pairs = state.get("pairs") or []
    if any(not isinstance(pair, dict) for pair in pairs):
        return False
    if not pairs or all(SOURCE_REFS in pair for pair in pairs):
        return False
    raw = state.get("source_paras") or []
    cleaned = state.get("source_paragraphs") or []
    if not raw:
        return False
    if not cleaned:
        cleaned = raw
    cleanup = state.get("source_cleanup") or {}
    origins = cleanup.get("output_input_indices")
    if origins is None:
        if [_normal(p) for p in raw] != [_normal(p) for p in cleaned]:
            return False
        origins = [[i] for i in range(len(raw))]
    if len(origins) != len(cleaned) or any(
        not refs or any(type(ref) is not int or not 0 <= ref < len(raw) for ref in refs)
        for refs in origins
    ):
        return False
    meta = state.get("segmentation") or {}
    if state.get("segment_structure_history"):
        # A recorded structure edit (split/merge/source rewrite) means the map
        # below is stale by construction: never backfill from it.
        return False
    if state.get("excluded_segments"):
        return _bind_from_paragraph_ranges(
            pairs, origins, meta.get("paragraph_ranges") or [], cleaned)
    # Re-segmenting the whole document costs ~0.3 s on a book and runs before
    # every saved edit; when the pairs cannot be bound (they do not match the
    # initial segmentation) the answer is the same until one of the inputs
    # changes, so a refusal is remembered for exactly these inputs.
    key = (tuple(cleaned), tuple(tuple(refs) for refs in origins),
           meta.get("mode") or "paragraph", state.get("source_lang"),
           tuple(pair.get("source") for pair in pairs))
    if key in _UNBINDABLE:
        return False
    units, rebuilt = segmentation.segment_paragraphs(
        cleaned, mode=meta.get("mode") or "paragraph",
        source_lang=state.get("source_lang"))
    # Partial translation batches can also be bound while they still occupy
    # the initial prefix. Never accept matching length alone.
    if len(pairs) > len(units) or any(
        _normal(pair.get("source")) != _normal(units[i])
        for i, pair in enumerate(pairs)
    ):
        _UNBINDABLE[key] = None
        while len(_UNBINDABLE) > 8:
            _UNBINDABLE.popitem(last=False)
        return False
    changed = False
    for item in rebuilt["paragraph_ranges"]:
        refs = list(dict.fromkeys(origins[item["paragraph_index"]]))
        for index in range(item["start"], min(item["end"], len(pairs))):
            if SOURCE_REFS not in pairs[index]:
                pairs[index][SOURCE_REFS] = list(refs)
                changed = True
    return changed


def _docx_source(source: bytes):
    document = Document(io.BytesIO(source))
    nodes, paragraphs, media, warnings = [], [], {}, []
    for paragraph in document.paragraphs:
        parts = [_text(part).strip() for part in re.split(r"\n+", paragraph.text)]
        # Mirrors the documented extraction boundary. Empty image paragraphs
        # still participate in the content stream below.
        kept = [part for part in parts if len(part) > 1 and not pdf_ingestion.ORNAMENT_RE.fullmatch(part)]
        for part in kept:
            ref = len(paragraphs)
            paragraphs.append(part)
            style = paragraph.style.style_id if paragraph.style is not None else ""
            role = "heading" if style.startswith("Heading") or style == "Title" else "body"
            nodes.append({"kind": "text", "ref": ref, "role": role})
        # Inline/floating pictures are emitted after their complete paragraph.
        # Each occurrence is kept even when the package reuses one media part.
        picture_nodes = list(paragraph._p.iter(qn("a:blip")))
        picture_nodes += list(paragraph._p.iter("{urn:schemas-microsoft-com:vml}imagedata"))
        for picture in picture_nodes:
            rid = picture.get(qn("r:embed")) or picture.get(qn("r:id"))
            if not rid:
                raise EpubExportError("EPUB 无法读取 DOCX 的外链图片，请先将图片嵌入源文件。")
            part = document.part.related_parts.get(rid)
            if part is None:
                raise EpubExportError("EPUB 图片资源缺失，请检查源 DOCX。")
            data = part.blob
            if data.startswith(b"\x89PNG\r\n\x1a\n"):
                ext, mime = "png", "image/png"
            elif data.startswith(b"\xff\xd8"):
                ext, mime = "jpg", "image/jpeg"
            else:
                try:
                    pixmap = fitz.Pixmap(data)
                    if pixmap.colorspace and pixmap.colorspace.n > 3:
                        pixmap = fitz.Pixmap(fitz.csRGB, pixmap)
                    data = pixmap.tobytes("png")
                    ext, mime = "png", "image/png"
                except Exception as exc:
                    raise EpubExportError("EPUB 暂不支持该 DOCX 图片格式，请将其转换为 PNG/JPEG 后重新导入。") from exc
            name = f"images/image-{len(media) + 1:04d}.{ext}"
            media[name] = (data, mime)
            nodes.append({"kind": "image", "name": name})
        if any(True for _ in paragraph._p.iter(qn("wp:anchor"))):
            warnings.append("DOCX 浮动图片已放在其锚定段落之后。")
    if document.tables:
        warnings.append("DOCX 表格及表格内图片未纳入此版 EPUB；导入范围说明仍适用。")
    if any(True for _ in document.element.body.iter(qn("w:txbxContent"))):
        warnings.append("DOCX 文本框内容未纳入此版 EPUB。")
    cover = None
    if nodes and nodes[0]["kind"] == "image":
        cover = nodes.pop(0)["name"]
    return paragraphs, nodes, media, warnings, cover


def _visual_pdf_source(state: dict, source: bytes):
    """Use the persisted semantic source, rather than re-reading a text layer.

    Whole scan pages are not illustrations. Only figures with real source
    geometry can be cropped; page-only figures are explicitly reported omitted.
    Paragraph references still bind against the original semantic paragraphs,
    so cleanup and CAT structure edits retain their existing export contract.
    """
    from .document_ir import normalized_bbox

    document = (state.get("extraction_report") or {}).get("document_ir") or {}
    blocks = document.get("blocks") or []
    by_id = {block["block_id"]: block for block in blocks}
    ids = document.get("translation_block_ids") or []
    if not ids or len(set(ids)) != len(ids) or any(bid not in by_id for bid in ids):
        raise EpubExportError("EPUB 缺少已识别正文的文档结构，无法导出扫描件。")
    paragraphs = [_text(by_id[bid]["text"]) for bid in ids]
    refs = {bid: index for index, bid in enumerate(ids)}
    pages = {page["page_id"]: page for page in document.get("pages") or []}
    visual_pages = {span["page_id"] for block in blocks
                    if (block.get("provenance") or {}).get("provider") == "vlm"
                    for span in (block.get("source_anchor") or {}).get("source_spans") or []}

    def source_order(block):
        spans = (block.get("source_anchor") or {}).get("source_spans") or []
        page_index = min((pages[s["page_id"]]["page_index"] for s in spans
                          if s.get("page_id") in pages), default=len(pages))
        # Native drawing order need not be reading order (a later paragraph
        # can be drawn before an earlier image). Keep the existing native EPUB
        # geometry order; semantic blocks use the provider's reading order.
        if (block.get("provenance") or {}).get("provider") == "native_pdf":
            boxes = [s["bbox"] for s in spans if s.get("bbox") and s.get("page_id") in pages
                     and pages[s["page_id"]]["page_index"] == page_index]
            if boxes:
                first = min(boxes, key=lambda box: (box[1], box[0]))
                return page_index, first[1], first[0]
        return page_index, block["reading_order"], 0

    nodes, media = [], {}
    warnings = ["扫描件 EPUB 使用已识别正文重建；不保留原页排版，也不把整张扫描页作为插图。"]
    with fitz.open(stream=source, filetype="pdf") as pdf:
        for block in sorted(blocks, key=source_order):
            bid = block["block_id"]
            if bid in refs:
                role = "heading" if block["type"] in {"title", "heading"} else "body"
                nodes.append({"kind": "text", "ref": refs[bid], "role": role})
            elif block["type"] == "figure":
                spans = (block.get("source_anchor") or {}).get("source_spans") or []
                if ((block.get("provenance") or {}).get("provider") == "native_pdf"
                        and any(s.get("page_id") in visual_pages for s in spans)):
                    # Mixed PDFs retain raw native layout evidence for scanned
                    # pages too. The full-page scan is not a semantic figure.
                    continue
                crops = []
                complete = bool(spans)
                for span in spans:
                    page = pages.get(span.get("page_id"))
                    box = normalized_bbox(span.get("bbox")) if span.get(
                        "grounding_quality") in {"exact", "strong", "approximate"} else None
                    if page and box and 0 <= page["page_index"] < len(pdf):
                        crop = (page["page_index"], tuple(box))
                        if crop not in crops:
                            crops.append(crop)
                    else:
                        complete = False
                if not complete:
                    numbers = sorted({pages[s["page_id"]]["page_index"] + 1 for s in spans
                                      if s.get("page_id") in pages})
                    location = "PDF 第 " + "、".join(map(str, numbers)) + " 页的" if numbers else ""
                    warnings.append(f"{location}插图缺少完整区域坐标，未纳入 EPUB。")
                    continue
                for page_index, box in crops:
                    page = pdf[page_index]
                    rotation = page.rotation
                    page.set_rotation(0)
                    try:
                        rect = fitz.Rect(box[0] * page.rect.width, box[1] * page.rect.height,
                                         box[2] * page.rect.width, box[3] * page.rect.height)
                        data = page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=rect,
                                              alpha=False).tobytes("png")
                    finally:
                        page.set_rotation(rotation)
                    name = f"images/image-{len(media) + 1:04d}.png"
                    media[name] = (data, "image/png")
                    nodes.append({"kind": "image", "name": name})
    return paragraphs, nodes, media, warnings, None


def _pdf_source(source: bytes):
    blocks = pdf_ingestion.classify_blocks(pdf_ingestion.extract_layout_blocks(source))
    records = [record for record in pdf_ingestion.reconstruct_paragraph_records(blocks)
               if record.get("text") and not pdf_ingestion.ORNAMENT_RE.fullmatch(record["text"])]
    paragraphs = [_text(record["text"]) for record in records]
    by_id = {block["block_id"]: block for block in blocks}
    ordered, media, warnings = [], {}, []
    cover = None
    for ref, record in enumerate(records):
        first = by_id[record["block_ids"][0]]
        ordered.append((first["page_number"], first["y0"], first["x0"], 0,
                        {"kind": "text", "ref": ref, "role": record["role"]}))
    with fitz.open(stream=source, filetype="pdf") as document:
        vector_pages = []
        for page_number, page in enumerate(document, 1):
            if page.get_drawings():
                vector_pages.append(page_number)
            page_blocks = [block for block in blocks if block["page_number"] == page_number]
            images = [block for block in page_blocks if block["role"] == "image"]
            if images and not records:
                raise EpubExportError("EPUB 原文件只有图像、没有可关联的文本；扫描页不属于普通文本 PDF 导出范围。")
            # A leading image-only page is the original book cover. Render
            # the whole page, including margins and vector decorations, once.
            if page_number == 1 and images and not any(1 in record["page_numbers"] for record in records):
                cover = "images/cover.png"
                media[cover] = (page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes("png"), "image/png")
                if page_number in vector_pages:
                    vector_pages.remove(page_number)
                continue
            if images and not any(page_number in record["page_numbers"] for record in records):
                warnings.append(f"PDF 第 {page_number} 页没有文本层，该页图片按原样保留，图片内文字未翻译。")
            for block in images:
                name = f"images/image-{len(media) + 1:04d}.png"
                rectangle = fitz.Rect(block["x0"], block["y0"], block["x1"], block["y1"])
                # Rendering the image's rectangle preserves PDF masks, rotation
                # and clipping instead of exporting an untransformed image blob.
                data = page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=rectangle, alpha=False).tobytes("png")
                media[name] = (data, "image/png")
                ordered.append((page_number, block["y0"], block["x0"], 1,
                                {"kind": "image", "name": name, "page": page_number}))
            for block in page_blocks:
                if block["role"] == "caption":
                    ordered.append((page_number, block["y0"], block["x0"], 2,
                                    {"kind": "caption", "text": _text(block["text"]), "original": True}))
        if vector_pages:
            warnings.append("PDF 矢量绘图未自动提取（页码：" + "、".join(map(str, vector_pages)) + "），请核对原文。")
    return paragraphs, [item[-1] for item in sorted(ordered, key=lambda item: item[:4])], media, warnings, cover


def _toc_links(paragraphs: list[str], nodes: list[dict]) -> dict[int, int]:
    """Link source contents entries only to an unambiguous later heading."""
    def key(value):
        value = re.sub(r"(?:\s+[.·…]*\s*|[.·…]{2,})\d+\s*$", "", value)
        return _normal(value).casefold()

    text_nodes = [node for node in nodes if node["kind"] == "text"]
    links = {}
    for start, node in enumerate(text_nodes):
        if key(paragraphs[node["ref"]]) not in {"contents", "table of contents", "目录", "目錄"}:
            continue
        end = next((i for i in range(start + 1, len(text_nodes))
                    if text_nodes[i]["role"] == "heading"), len(text_nodes))
        headings = {}
        for offset in range(end, len(text_nodes)):
            candidate = text_nodes[offset]
            if candidate["role"] != "heading":
                continue
            headings.setdefault(key(paragraphs[candidate["ref"]]), []).append(candidate["ref"])
            # PDF extraction can split one wrapped title into adjacent heading
            # paragraphs. Match their complete text to the contents entry.
            parts = [paragraphs[candidate["ref"]]]
            for following in text_nodes[offset + 1:]:
                if following["role"] != "heading":
                    break
                parts.append(paragraphs[following["ref"]])
                headings.setdefault(key(" ".join(parts)), []).append(candidate["ref"])
        for entry in text_nodes[start + 1:end]:
            matches = headings.get(key(paragraphs[entry["ref"]]), [])
            if len(matches) == 1:
                links[entry["ref"]] = matches[0]
    return links


def _join_targets(values: list[str]) -> str:
    result = ""
    for value in values:
        value = _text(value).strip()
        if not value:
            continue
        if result and not (re.search(r"[\u3400-\u9fff。！？；：，、]$", result) and re.match(r"[\u3400-\u9fff]", value)):
            result += " "
        result += value
    return result


def _markup(value: Any) -> str:
    return escape(_text(value)).replace("\n", "<br/>")


_language = language_tag

def _xhtml(title: str, body: str, language: str, direction: str | None = None) -> bytes:
    return (f'<?xml version="1.0" encoding="utf-8"?>\n'
            f'<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" '
            f'xml:lang="{language}" lang="{language}" '
            f'dir="{direction or text_direction(re.sub(r"<[^>]+>", " ", body), language)}"><head><title>{_markup(title)}</title>'
            f'<link rel="stylesheet" type="text/css" href="style.css"/></head><body>{body}</body></html>').encode("utf-8")


def build_translation_epub(state: dict, source: bytes, *, job_id: str = "") -> tuple[bytes, dict]:
    """Create a self-contained EPUB 3 and a human-readable scope report."""
    if not source:
        raise EpubExportError("EPUB 需要原始文件才能保留图片，请重新上传源文件。")
    validation = translation_target.validate_translation_pairs(state.get("pairs") or [], allow_empty=True)
    if validation["blocking"]:
        raise EpubExportError("EPUB 译文未通过正文格式检查。")
    filename = state.get("filename") or ""
    try:
        if filename.lower().endswith(".docx"):
            paragraphs, nodes, media, warnings, cover = _docx_source(source)
        elif filename.lower().endswith(".pdf"):
            if (state.get("extraction_report") or {}).get("semantic_source") in {"vlm", "mixed"}:
                paragraphs, nodes, media, warnings, cover = _visual_pdf_source(state, source)
            else:
                paragraphs, nodes, media, warnings, cover = _pdf_source(source)
        elif filename.lower().endswith((".epub", ".txt", ".md", ".markdown")):
            from .text_import import extract_text_source
            paragraphs, omitted, _facts = extract_text_source(filename, source)
            paragraphs = [_text(text) for text in paragraphs]
            nodes = [{"kind": "text", "ref": index, "role": "body"}
                     for index in range(len(paragraphs))]
            media, cover = {}, None
            warnings = ["此来源导出为重建的纯文字 EPUB；不保留源书排版、图片或音视频。"]
            warnings.extend(f'{item["label"]}（{item["count"]}）未纳入译文。' for item in omitted)
        else:
            raise EpubExportError("EPUB 导出不支持此来源格式。")
    except EpubExportError:
        raise
    except Exception as exc:
        raise EpubExportError(f"EPUB 无法读取源文件：{exc}") from exc
    working = deepcopy(state)
    raw = working.get("source_paras")
    if raw is None or not raw:
        # An untouched legacy task can be matched to the source; no guesses
        # based on mutable paragraph indexes or fuzzy string search.
        working["source_paras"] = list(paragraphs)
        raw = paragraphs
    if [_normal(p) for p in raw] != [_normal(p) for p in paragraphs]:
        raise EpubExportError("EPUB 原文段落与任务来源不一致，无法可靠定位图片；请重新导入源文件建立图文关联。")
    bind_source_references(working)
    pairs = working.get("pairs") or []
    if not pairs or any(SOURCE_REFS not in pair for pair in pairs):
        raise EpubExportError("EPUB 缺少图文来源关联；此任务的段落结构已变更或来源映射不完整，请重新导入建立关联。")
    groups = []
    for pair in pairs:
        refs = tuple(pair[SOURCE_REFS])
        if not refs or any(type(ref) is not int or not 0 <= ref < len(paragraphs) for ref in refs):
            raise EpubExportError("EPUB 段落来源关联无效，请重新导入源文件。")
        if groups and groups[-1]["refs"] == refs:
            groups[-1]["targets"].append(pair.get("target") or "")
        else:
            groups.append({"refs": refs, "targets": [pair.get("target") or ""]})
    groups_by_ref = {}
    for index, group in enumerate(groups):
        for ref in group["refs"]:
            groups_by_ref.setdefault(ref, []).append(index)
    title = str(state.get("book_title") or Path(filename).stem or "译文")
    language = _language(str(state.get("target_lang") or "简体中文"))
    direction = text_direction(" ".join(str(pair.get("target") or "") for pair in pairs),
                               state.get("target_lang"))
    source_links = _toc_links(paragraphs, nodes)
    group_links = {}
    for index, group in enumerate(groups):
        if len(group["refs"]) != 1:
            continue
        destination_ref = source_links.get(group["refs"][0])
        destinations = groups_by_ref.get(destination_ref, [])
        if destinations and _join_targets(groups[destinations[0]]["targets"]):
            group_links[index] = destinations[0]
    unit_locations = {}
    emitted = set()
    chapters = []
    current_title, current_parts, current_size = "正文", [], 0
    caption_count = sum(node["kind"] == "caption" for node in nodes)
    for node in nodes:
        if node["kind"] == "text":
            for index in groups_by_ref.get(node["ref"], []):
                if index in emitted:
                    continue
                emitted.add(index)
                target = _join_targets(groups[index]["targets"])
                if not target:
                    current_parts.append('<p class="missing" aria-label="尚无译文"></p>')
                    continue
                heading = node["role"] == "heading"
                if current_parts and (heading or current_size > 100_000):
                    chapters.append((current_title, "".join(current_parts)))
                    current_parts, current_size = [], 0
                    current_title = target if heading else f"正文 {len(chapters) + 1}"
                elif heading:
                    current_title = target
                tag = "h1" if heading else "p"
                unit_locations[index] = f"chapter-{len(chapters) + 1:04d}.xhtml#unit-{index}"
                markup = _markup(target)
                if index in group_links:
                    markup = f'<a href="epub-unit:{group_links[index]}">{markup}</a>'
                current_parts.append(f'<{tag} id="unit-{index}">{markup}</{tag}>')
                current_size += len(target)
        elif node["kind"] == "image":
            alt = f"原文插图 {list(media).index(node['name']) + 1}"
            current_parts.append(f'<figure><img src="{node["name"]}" alt="{alt}"/></figure>')
        else:
            caption_direction = text_direction(node["text"], state.get("source_lang"))
            current_parts.append(f'<p class="caption" lang="und" dir="{caption_direction}">{_markup(node["text"])}</p>')
    if current_parts:
        chapters.append((current_title, "".join(current_parts)))
    if len(emitted) != len(groups):
        raise EpubExportError("EPUB 有译文无法定位到原文，已停止导出以避免遗漏。")
    if not chapters:
        raise EpubExportError("EPUB 尚无可导出的译文或图片。")
    chapters = [(label, re.sub(r'href="epub-unit:(\d+)"', lambda match: f'href="{unit_locations[int(match[1])]}"', body))
                for label, body in chapters]
    warnings = list(dict.fromkeys(warnings))
    if caption_count:
        warnings.append(f"{caption_count} 条 PDF 图注保留原文；图片内部文字保持原样。")
    elif media:
        warnings.append("图片内部文字保持原样。")
    missing = sum(bool(str(pair.get("source") or "").strip()) and not str(pair.get("target") or "").strip() for pair in pairs)
    if missing:
        warnings.append(f"{missing} 个翻译单元尚无译文，正文中留空。")
    report = {"format": "EPUB 3", "layout": "reflowable", "source_filename": filename,
              "image_count": len(media), "has_cover": bool(cover), "original_caption_count": caption_count,
              "chapter_count": len(chapters), "linked_toc_count": len(group_links), "missing_translation_count": missing,
              "placement": "图片位于对应完整段落之后；跨图片合并的译文整体位于图片之前。",
              "warnings": warnings}
    identifier = "urn:folith:" + (job_id or "translation")
    files, entries = {}, []
    if cover:
        files["cover.xhtml"] = _xhtml(title, f'<section epub:type="cover" class="cover"><img src="{cover}" alt="原书封面"/></section>', language, direction)
        entries.append(("cover.xhtml", "封面"))
    for index, (chapter_title, body) in enumerate(chapters, 1):
        name = f"chapter-{index:04d}.xhtml"
        files[name] = _xhtml(chapter_title, body, language, direction)
        entries.append((name, chapter_title))
    toc = '<nav epub:type="toc" id="toc"><h1>目录</h1><ol>' + "".join(
        f'<li><a href="{name}">{_markup(label)}</a></li>' for name, label in entries) + '</ol></nav>'
    if cover:
        toc += '<nav epub:type="landmarks" hidden="hidden"><ol><li><a epub:type="cover" href="cover.xhtml">封面</a></li></ol></nav>'
    files["nav.xhtml"] = _xhtml("目录", toc, language, direction)
    files["style.css"] = b'body { margin: 0 5%; line-height: 1.6; } p { margin: 0.7em 0; } h1,h2 { line-height: 1.3; } figure { margin: 1em 0; text-align: center; break-inside: avoid; } img { max-width: 100%; height: auto; } .caption { text-align: center; font-size: 0.9em; } .cover { text-align: center; margin: 0; padding: 0; } .cover img { max-height: 95vh; object-fit: contain; }'
    manifest = ['<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>',
                '<item id="css" href="style.css" media-type="text/css"/>',
                '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>']
    spine = []
    for index, (name, _) in enumerate(entries):
        manifest.append(f'<item id="content-{index}" href="{name}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="content-{index}"/>')
    for index, (name, (_, mime)) in enumerate(media.items()):
        properties = ' properties="cover-image"' if name == cover else ''
        manifest.append(f'<item id="image-{index}" href="{name}" media-type="{mime}"{properties}/>')
    modified = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    files["content.opf"] = (f'<?xml version="1.0" encoding="utf-8"?>'
        f'<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book-id">'
        f'<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="book-id">{escape(_text(identifier))}</dc:identifier>'
        f'<dc:title>{escape(_text(title))}</dc:title><dc:language>{language}</dc:language>'
        f'<meta property="dcterms:modified">{modified}</meta><meta property="rendition:layout">reflowable</meta></metadata>'
        f'<manifest>{"".join(manifest)}</manifest><spine toc="ncx" page-progression-direction="{direction}">{"".join(spine)}</spine></package>').encode("utf-8")
    files["toc.ncx"] = (f'<?xml version="1.0" encoding="utf-8"?>'
        f'<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1"><head>'
        f'<meta name="dtb:uid" content="{escape(identifier, quote=True)}"/></head>'
        f'<docTitle><text>{_markup(title)}</text></docTitle><navMap>' + "".join(
        f'<navPoint id="nav-{index}" playOrder="{index + 1}"><navLabel><text>{_markup(label)}</text></navLabel>'
        f'<content src="{name}"/></navPoint>' for index, (name, label) in enumerate(entries)) + '</navMap></ncx>').encode("utf-8")
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", b"application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/container.xml", b'<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles><rootfile full-path="EPUB/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        for name, data in files.items():
            archive.writestr("EPUB/" + name, data)
        for name, (data, _) in media.items():
            archive.writestr("EPUB/" + name, data)
    return output.getvalue(), report
