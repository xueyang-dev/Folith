"""Split out of core.py.

Moved verbatim out of ``core.py``; ``core`` re-exports every public and private
name defined here.  Anything still owned by ``core`` is read late-bound through
``_c`` so that tests (and callers) that patch ``core.X`` keep working.
"""

import fitz
import io
import pandas as pd
import re
from docx import Document
from folith.textual import text_direction
from html import escape as html_escape
from folith.engine._link import core as _c


def _docx_part_texts(document, needle):
    """读取 word/<needle> 部件里的全部文本（脚注/尾注/批注不在正文 XML 里）。"""
    try:
        package = getattr(document.part, "package", None)
        if package is None:
            return []
        blob = None
        for part in package.parts:
            if needle in str(getattr(part, "partname", "")):
                blob = part.blob
                break
        if not blob:
            return []
        from lxml import etree
        root = etree.fromstring(bytes(blob))
    except Exception:  # noqa: BLE001 - report is best-effort, never blocks import
        return []
    return [str(node.text or "").strip() for node in root.iter(f"{_c._W_NS}t")
            if str(node.text or "").strip()]


def _docx_unsupported_parts(document):
    """Inspect a DOCX for content a paragraph-only extractor cannot translate.

    The point is not to support these structures yet: it is to make sure the
    user is never told "文档已导入" while tables, running heads or footnotes were
    silently dropped.  Counting is deliberately structural (no text heuristics).
    """
    parts = []

    def add(kind, label, count, detail, sample=""):
        if count <= 0:
            return
        parts.append({"kind": kind, "label": label, "count": int(count),
                      "detail": detail, "sample": str(sample or "")[:160]})

    try:
        body = document.element.body
        tables = list(getattr(document, "tables", []) or [])
        table_cells = sum(len(row.cells) for table in tables for row in table.rows)
        table_sample = next(
            (str(cell.text).strip() for table in tables for row in table.rows
             for cell in row.cells if str(cell.text).strip()), "")
        add("table", "表格内容", len(tables),
            f"{len(tables)} 个表格、约 {table_cells} 个单元格；表格文字不进入翻译",
            table_sample)

        headers = footers = 0
        header_sample = footer_sample = ""
        for section in getattr(document, "sections", []) or []:
            for paragraph in getattr(section.header, "paragraphs", []) or []:
                text = str(paragraph.text or "").strip()
                if text:
                    headers += 1
                    header_sample = header_sample or text
            for paragraph in getattr(section.footer, "paragraphs", []) or []:
                text = str(paragraph.text or "").strip()
                if text:
                    footers += 1
                    footer_sample = footer_sample or text
        add("header", "页眉", headers,
            f"页眉有 {headers} 段文字（重复页眉是扫描件最常见的问题来源）；"
            "页眉不进入翻译单元，译文不保留页眉",
            header_sample)
        add("footer", "页脚", footers,
            f"页脚有 {footers} 段文字（通常含页码）；"
            "页脚不进入翻译单元，译文不保留页脚",
            footer_sample)

        boxes = body.findall(f".//{_c._W_NS}txbxContent")
        box_sample = next(
            (" ".join(str(node.text or "") for node in box.iter(f"{_c._W_NS}t")).strip()
             for box in boxes
             if " ".join(str(node.text or "") for node in box.iter(f"{_c._W_NS}t")).strip()),
            "")
        add("textbox", "文本框", len(boxes),
            f"{len(boxes)} 个文本框；文本框文字不在正文段落里", box_sample)

        images = body.findall(f".//{_c._A_NS}blip")
        add("image", "图片", len(images),
            f"{len(images)} 张内嵌图片按原样保留在源文件里，但图片内的文字不参与翻译")

        equations = body.findall(f".//{_c._M_NS}oMath")
        add("equation", "公式", len(equations),
            f"{len(equations)} 处公式不参与翻译")

        footnotes = _docx_part_texts(document, "footnotes.xml")
        add("footnote", "脚注", len(footnotes),
            f"{len(footnotes)} 条脚注不参与翻译", footnotes[0] if footnotes else "")
        endnotes = _docx_part_texts(document, "endnotes.xml")
        add("endnote", "尾注", len(endnotes),
            f"{len(endnotes)} 条尾注不参与翻译", endnotes[0] if endnotes else "")
        comments = _docx_part_texts(document, "comments.xml")
        add("comment", "批注", len(comments),
            f"{len(comments)} 条批注不参与翻译", comments[0] if comments else "")
    except Exception:  # noqa: BLE001
        return parts
    return parts


def findings_report_md(state):
    """把审查结果渲染成 Markdown 报告（下载/展示用）。"""
    from folith.translation_evidence import translation_review_readiness

    stats = state.get("review_stats") or {}
    readiness = translation_review_readiness(state)
    skipped = int(readiness.get("skipped_count") or 0)
    reason_labels = {
        "already_reviewed_tm": "复用已审校翻译记忆",
        "clean_deterministic": "无确定性风险信号",
        "unknown": "未记录跳过原因",
    }
    lines = [
        "# 翻译审查报告", "",
        "## 概览",
        f"- 已审校段落：{stats.get('reviewed_segments', 0)}",
        f"- 审校批次：{stats.get('batches_reviewed', 0)}",
        f"- 审校失败批次：{stats.get('review_failed', 0)}",
        f"- 翻译记忆复用：{state.get('tm_used_count', 0)} 段",
        f"- 未执行本轮独立语义审校：{skipped} 段",
        f"- 审校覆盖状态：{readiness.get('coverage_status')}",
        f"- 审校覆盖完整：{'是' if readiness.get('coverage_complete') else '否'}",
        f"- blocking：{stats.get('blocking', 0)}",
        f"- actionable：{stats.get('actionable', 0)}",
        f"- informational：{stats.get('informational', 0)}",
        "", "## 审校覆盖",
    ]
    if skipped:
        lines.append(f"{skipped} 段未执行本轮独立语义审校（按策略跳过）：")
        for reason, count in sorted(
                (readiness.get("skipped_reason_counts") or {}).items()):
            lines.append(f"- {reason_labels.get(reason, reason)}：{count} 段")
        lines.append("- 跳过不等于审校通过，也不产生独立审校 Finding。")
    elif readiness.get("coverage_status") == "not_required":
        lines.append("当前任务未启用独立审校。")
    else:
        lines.append(
            f"独立审校覆盖 {readiness.get('reviewed_count', 0)} / "
            f"{len(readiness.get('expected_segment_ids') or [])} 段。")
    lines.extend(["", "## 待处理问题"])
    findings = state.get("findings") or []
    if not findings:
        lines.append("无。")
    else:
        for f in findings:
            line = f"- 第 {f.get('segment_index', -1) + 1} 段 [{f.get('severity')}] {f.get('reason')}"
            if f.get("suggested_target"):
                line += f"（建议译文：{f['suggested_target']}）"
            lines.append(line)
    return "\n".join(lines)


def _apply_doc_fonts(doc):
    """默认字体与可编辑学术文档的基础版式。"""
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
    style_sizes = {
        "Normal": 12, "Heading 1": 15, "Heading 2": 13.5,
        "Heading 3": 12.5, "Heading 4": 12, "Title": 18,
    }
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3",
                       "Heading 4", "Title", "Intense Quote", "List Bullet",
                       "List Number"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        style.font.name = _c.EN_FONT
        if style_name in style_sizes:
            style.font.size = Pt(style_sizes[style_name])
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), _c.EN_FONT)
        rfonts.set(qn("w:hAnsi"), _c.EN_FONT)
        rfonts.set(qn("w:eastAsia"), _c.CN_FONT)
        if style_name == "Normal":
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.first_line_indent = Cm(0.74)
        elif style_name not in {"Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"}:
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.first_line_indent = None


def _apply_run_fonts(run, rtl=False):
    """单个 run 的字体（表格单元格里的 run 不受 Normal 样式继承影响时兜底）。"""
    import unicodedata
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run.font.name = _c.EN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), _c.EN_FONT)
    rfonts.set(qn("w:hAnsi"), _c.EN_FONT)
    rfonts.set(qn("w:eastAsia"), _c.CN_FONT)
    rfonts.set(qn("w:cs"), "Arial")
    if rtl or any(unicodedata.bidirectional(c) in {"R", "AL"} for c in str(run.text or "")):
        rtl_el = rpr.find(qn("w:rtl"))
        if rtl_el is None:
            rtl_el = OxmlElement("w:rtl")
            rpr.append(rtl_el)
        rtl_el.set(qn("w:val"), "1")


def _apply_paragraph_direction(paragraph, language=None):
    """Set the direction on the individual paragraph, including table cells."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rtl = text_direction(paragraph.text, language) == "rtl"
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1" if rtl else "0")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    if rtl:
        # Logical leading alignment stays on the right in RTL renderers;
        # literal "right" is mirrored by paragraph BiDi in LibreOffice.
        ppr.find(qn("w:jc")).set(qn("w:val"), "start")
    for run in paragraph.runs:
        _apply_run_fonts(run, rtl=rtl)


def dict_to_excel(term_dict):
    df = pd.DataFrame(list(term_dict.items()), columns=["Source", "Target"])
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False)
    output.seek(0)
    return output


def paragraphs_to_word(paragraphs, source_lang=None):
    doc = Document()
    _apply_doc_fonts(doc)
    doc.add_heading('阶段一：清洗后原文提取', 0)
    for p in paragraphs:
        _apply_paragraph_direction(doc.add_paragraph(p), source_lang)
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


def translations_to_word(pairs, target_lang=None):
    doc = Document()
    _apply_doc_fonts(doc)
    doc.add_heading("译文", 0)
    for pair in pairs or []:
        _apply_paragraph_direction(
            doc.add_paragraph(str(pair.get("target") or "")), target_lang)
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


def build_translation_docx(state, source):
    """Rebuild translated DOCX in source order, retaining supported pictures."""
    from folith.epub import (
        EpubExportError, ImageOnlySourceError, _docx_source, _pdf_source, _join_targets,
        _translation_groups,
    )
    from docx.shared import Emu
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    filename = str(state.get("filename") or "").lower()
    pairs = state.get("pairs") or []
    report = {"image_count": 0, "warnings": []}
    if not source or not filename.endswith((".pdf", ".docx")):
        report["warnings"] = ["缺少原始文件，无法恢复图片。" if not source else
                              "此来源的译文 DOCX 为纯文字重建，不保留源文件图片。"]
        return translations_to_word(pairs, state.get("target_lang")), report
    try:
        paragraphs, nodes, media, warnings, cover = (
            _docx_source(source) if filename.endswith(".docx") else _pdf_source(source))
        report.update(image_count=len(media), warnings=[
            warning.replace("EPUB", "译文 DOCX") for warning in warnings])
        if not media:
            return translations_to_word(pairs, state.get("target_lang")), report
        groups, groups_by_ref = _translation_groups(state, paragraphs)
    except ImageOnlySourceError as exc:
        # 扫描件没有可关联的文本：退回纯文字重建（加图片之前的行为），不阻断交付。
        report.update(image_count=0, warnings=[
            str(exc).replace("EPUB", "译文 DOCX") + " 已改为纯文字重建，不保留图片。"])
        return translations_to_word(pairs, state.get("target_lang")), report
    except EpubExportError as exc:
        message = str(exc)
        if "来源不一致" not in message:
            raise RuntimeError(message.replace("EPUB", "译文 DOCX")) from exc
        # After re-reading failed pages the source paragraph mapping can no
        # longer be matched to the original file.  Keep the translation
        # deliverable: rebuild a text-only DOCX and report the missing images
        # instead of blocking the whole delivery.  A missing association is a
        # different, structural defect and still fails explicitly.
        report["warnings"].append(
            message.replace("EPUB", "译文 DOCX")
            + " 已改为纯文字重建，不保留图片。")
        return translations_to_word(pairs, state.get("target_lang")), report
    except Exception as exc:
        raise RuntimeError(f"译文 DOCX 无法读取源文件：{exc}") from exc

    doc = Document()
    _apply_doc_fonts(doc)
    doc.add_heading("译文", 0)
    section = doc.sections[0]
    max_width = section.page_width - section.left_margin - section.right_margin
    max_height = section.page_height - section.top_margin - section.bottom_margin
    # Leave space for the picture paragraph's line box and spacing.
    max_height -= Emu(360000)
    emitted = set()
    if cover:
        nodes = [{"kind": "image", "name": cover}] + nodes
    for node in nodes:
        if node["kind"] == "text":
            for index in groups_by_ref.get(node["ref"], []):
                if index in emitted:
                    continue
                emitted.add(index)
                target = _join_targets(groups[index]["targets"])
                paragraph = (doc.add_heading(target, level=1)
                             if target and node["role"] == "heading"
                             else doc.add_paragraph(target))
                _apply_paragraph_direction(paragraph, state.get("target_lang"))
        elif node["kind"] == "image":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.first_line_indent = 0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture = paragraph.add_run().add_picture(io.BytesIO(media[node["name"]][0]))
            scale = min(1, max_width / picture.width, max_height / picture.height)
            picture.width = Emu(round(picture.width * scale))
            picture.height = Emu(round(picture.height * scale))
        elif node["kind"] == "caption":
            paragraph = doc.add_paragraph(node["text"], style="Caption")
            _apply_paragraph_direction(paragraph, state.get("source_lang"))
    if len(emitted) != len(groups):
        raise RuntimeError("译文 DOCX 有译文无法定位到原文，已停止导出以避免遗漏。")
    report["placement"] = "图片位于对应完整段落之后；跨图片合并的译文整体位于图片之前。"
    report["warnings"] = list(dict.fromkeys(
        [warning.replace("EPUB", "译文 DOCX") for warning in warnings]
        + ["图片内部文字保持原样，PDF 图注保留原文；文档重新排版，不保留原页码和坐标。"]
    ))
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out, report


def translations_to_pdf(pairs, target_lang=None):
    """Render a paginated, Unicode-aware translation PDF with PyMuPDF Story."""
    paragraphs = "".join(
        '<table class="entry"><tr>'
        f'<td dir="{text_direction(pair.get("target"), target_lang)}" '
        'style="text-align: start">'
        f'{html_escape(str(pair.get("target") or ""))}</td></tr></table>'
        for pair in pairs or [])
    html = f'<article><h1>译文</h1>{paragraphs}</article>'
    css = (
        "@page { size: A4; } "
        "body { font-family: sans-serif; font-size: 11pt; line-height: 1.65; "
        "color: #1f2937; } h1 { font-size: 20pt; margin: 0 0 20pt; "
        "color: #111827; } table.entry { width: 100%; border-collapse: collapse; "
        "margin: 0 0 10pt; break-inside: avoid; page-break-inside: avoid; } "
        "td { padding: 0; }"
    )
    story = fitz.Story(html=html, user_css=css, em=11)
    page_rect = fitz.paper_rect("a4")
    content_rect = fitz.Rect(54, 54, page_rect.width - 54, page_rect.height - 54)

    def rectfn(_rect_num, _filled):
        return page_rect, content_rect, None

    pdf = story.write_with_links(rectfn)
    for number, page in enumerate(pdf, start=1):
        page.insert_text(
            (page_rect.width / 2 - 4, page_rect.height - 25), str(number),
            fontsize=9, fontname="helv", color=(0.45, 0.48, 0.52))
    data = pdf.tobytes(garbage=4, deflate=True)
    pdf.close()
    return data


def glossary_to_excel(entries, fallback=None):
    normalized = _c.normalize_glossary(entries or [])
    if normalized:
        rows = [{
            "Source": entry.get("source") or "",
            "Target": entry.get("preferred") or entry.get("target") or "",
            "Status": entry.get("status") or "",
            "Domain": entry.get("domain") or "",
            "Note": entry.get("note") or "",
        } for entry in normalized]
        frame = pd.DataFrame(rows)
    else:
        frame = pd.DataFrame(list((fallback or {}).items()),
                             columns=["Source", "Target"])
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False)
    output.seek(0)
    return output


def _is_latin_word_char(ch):
    if not ch:
        return False
    if ch.isascii() and (ch.isalnum() or ch == "_"):
        return True
    return "\u00C0" <= ch <= "\u024F"


def _expand_span_end(text, start, end, last_is_latin):
    if last_is_latin:
        while end < len(text) and _is_latin_word_char(text[end]):
            end += 1
    return start, end


def _find_span(text, needle):
    """宽容定位子串：统一引号/破折号/省略号、折叠空白后查找，返回 (start, end) 或 None。
    对于西文字词，确保不从词中截断，且若命中词首则完整包含整个单词（如 Egypt 命中 Egyptian MiGs 时标全 Egyptian）。
    """
    if not needle or not text:
        return None

    first_is_latin = _is_latin_word_char(needle[0])
    last_is_latin = _is_latin_word_char(needle[-1])

    # 1. 直接子串查找（首字符为西文字母时，确保位于词首边界，避免在 smart 中误中 art）
    if first_is_latin:
        curr = 0
        while True:
            pos = text.find(needle, curr)
            if pos < 0:
                break
            if pos == 0 or not _is_latin_word_char(text[pos - 1]):
                return _expand_span_end(text, pos, pos + len(needle), last_is_latin)
            curr = pos + 1
    else:
        pos = text.find(needle)
        if pos >= 0:
            return _expand_span_end(text, pos, pos + len(needle), last_is_latin)

    # 2. 宽容归一化查找（引号、破折号、省略号统一及空白折叠）
    mapping = []
    norm_chars = []
    for orig_idx, ch in enumerate(text):
        ch2 = ch.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
        ch2 = ch2.replace("–", "-").replace("—", "-").replace("…", "...")
        if ch2.isspace():
            if norm_chars and norm_chars[-1] != " ":
                norm_chars.append(" ")
                mapping.append(None)
            continue
        norm_chars.append(ch2)
        mapping.append(orig_idx)
    norm_text = "".join(norm_chars)
    needle2 = needle.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    needle2 = needle2.replace("–", "-").replace("—", "-").replace("…", "...")
    needle2 = re.sub(r"\s+", " ", needle2).strip()

    curr = 0
    pos = -1
    while True:
        pos = norm_text.find(needle2, curr)
        if pos < 0:
            break
        if first_is_latin and pos > 0 and _is_latin_word_char(norm_text[pos - 1]):
            curr = pos + 1
            continue
        break

    if pos < 0:
        return None

    start = None
    for i in range(pos, len(mapping)):
        if mapping[i] is not None:
            start = mapping[i]
            break
    end = None
    for i in range(min(pos + len(needle2), len(mapping)) - 1, -1, -1):
        if mapping[i] is not None:
            end = mapping[i] + 1
            break
    if start is None or end is None:
        return None
    return _expand_span_end(text, start, end, last_is_latin)


def _colored_cell(cell, text, spans, colors=None):
    """把一个单元格按 spans（(start,end,type) 已排序不重叠）拆成带色 run。"""
    from docx.shared import RGBColor
    palette = colors or _c.ANNOTATION_COLORS
    cursor = 0
    first = True
    for start, end, atype in spans:
        if start > cursor:
            run = cell.paragraphs[0].add_run(text[cursor:start])
            _apply_run_fonts(run)
        run = cell.paragraphs[0].add_run(text[start:end])
        _apply_run_fonts(run)
        run.font.color.rgb = RGBColor.from_string(
            str(palette.get(atype, _c.ANNOTATION_COLORS[atype])).lstrip("#"))
        if atype in ("rare", "domain"):
            run.bold = True
        elif atype == "hard":
            run.underline = True
        cursor = end
        first = False
    if cursor < len(text):
        run = cell.paragraphs[0].add_run(text[cursor:])
        _apply_run_fonts(run)
    if first and not text:
        cell.paragraphs[0].add_run("")


def pairs_to_word(pairs, annotations=None, colors=None, *, source_lang=None, target_lang=None):
    """双语对照表 -> Word 表格。

    annotations: {seg: [{"type": "rare|domain|hard", "src_span": [s,e]|None,
                         "tgt_span": [s,e]|None, "note": str}]}
    colors: {"rare"|"domain"|"hard": "RRGGBB"}，可自定义三类标注颜色。
    """
    palette = dict(_c.ANNOTATION_COLORS)
    if colors:
        palette.update({k: str(v).lstrip("#") for k, v in colors.items()})
    doc = Document()
    _apply_doc_fonts(doc)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = "原文"
    table.rows[0].cells[1].text = "译文"
    annot = _c._clean_annotations(annotations, pairs) if annotations else {}
    for i, pair in enumerate(pairs):
        row = table.add_row().cells
        seg_annot = annot.get(i) or []
        src_spans = _c._compose_spans(
            [(it["src_span"][0], it["src_span"][1], it["type"])
             for it in seg_annot if it.get("src_span")], len(pair['source']))
        tgt_spans = _c._compose_spans(
            [(it["tgt_span"][0], it["tgt_span"][1], it["type"])
             for it in seg_annot if it.get("tgt_span")], len(pair['target']))
        _colored_cell(row[0], pair['source'], src_spans, palette)
        _colored_cell(row[1], pair['target'], tgt_spans, palette)
        for paragraph in row[0].paragraphs:
            _apply_paragraph_direction(paragraph, source_lang)
        for paragraph in row[1].paragraphs:
            _apply_paragraph_direction(paragraph, target_lang)
    # 图例（放表格后，避免挤占首行）
    p_legend = doc.add_paragraph()
    legend_parts = [
        f"{_c.ANNOTATION_LABELS[k]}（#{palette.get(k, _c.ANNOTATION_COLORS[k])}）"
        for k in ("rare", "domain", "hard")]
    run = p_legend.add_run("图例：" + "；".join(legend_parts) + "。")
    _apply_run_fonts(run)
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


def _add_formatted_runs(paragraph, text):
    parts = text.split('**')
    for i, part in enumerate(parts):
        run = paragraph.add_run(part)
        _apply_run_fonts(run)
        if i % 2 != 0:
            run.bold = True


def _markdown_table_row(line):
    value = str(line or "").strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|"):
        value = value[:-1]
    return [cell.strip() for cell in value.split("|")]


def _add_markdown_table(doc, rows):
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=columns)
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for col_index in range(columns):
            cell = cells[col_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.35
            paragraph.paragraph_format.space_after = 0
            paragraph.paragraph_format.first_line_indent = None
            run = paragraph.add_run(values[col_index] if col_index < len(values) else "")
            _apply_run_fonts(run)
            if row_index == 0:
                run.bold = True
    return table


def markdown_to_word(md_text, theory):
    doc = Document()
    _apply_doc_fonts(doc)
    md_text = re.sub(r'```markdown|```', '', md_text).strip()
    md_text = re.sub(r'<!--.*?-->', '', md_text, flags=re.DOTALL)
    quote_labels = {
        "SYNTHETIC_SOURCE": "真实源文",
        "SIMULATED": "模拟初译",
        "OPTIMIZED": "优化译文",
    }
    quote_labels.update({"SOURCE": "原文", "INITIAL": "初译", "TARGET": "终译"})
    md_text = re.sub(
        r'(?m)^>\s*\[(SYNTHETIC_SOURCE|SIMULATED|OPTIMIZED|SOURCE|INITIAL|TARGET)\s+'
        r'(?:SC-\d{4,}|seg-[A-Za-z0-9_-]+)\]:\s*',
        lambda match: f"> {quote_labels[match.group(1)]}：", md_text)
    lines = md_text.splitlines()
    while lines and not lines[0].strip():
        lines.pop(0)
    if lines and re.match(r"^#\s+翻译实践报告", lines[0], re.IGNORECASE):
        lines.pop(0)
    deduped = []
    for line in lines:
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line.strip())
        if heading:
            title = re.sub(r"\s+#+\s*$", "", heading.group(2)).strip()
            previous = next((value for value in reversed(deduped) if value.strip()), "")
            previous_heading = re.match(r"^#{1,6}\s+(.+?)\s*$", previous.strip())
            if previous_heading:
                previous_title = re.sub(r"\s+#+\s*$", "", previous_heading.group(1)).strip()
                if re.sub(r"^\d+(?:\.\d+)*[.、．]?\s*", "", title).casefold() == \
                        re.sub(r"^\d+(?:\.\d+)*[.、．]?\s*", "", previous_title).casefold():
                    continue
        deduped.append(line)
    title = doc.add_heading(f'翻译实践报告：基于{theory}', 0)
    title.alignment = 1
    index = 0
    while index < len(deduped):
        line = deduped[index].strip()
        if ("|" in line and index + 1 < len(deduped)
                and _c._is_markdown_table_separator(deduped[index + 1])):
            rows = [_markdown_table_row(line)]
            index += 2
            while index < len(deduped) and "|" in deduped[index] and deduped[index].strip():
                rows.append(_markdown_table_row(deduped[index]))
                index += 1
            _add_markdown_table(doc, rows)
            continue
        line = line.strip()
        if not line:
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading:
            heading_level = max(1, min(4, len(heading.group(1)) - 1))
            doc.add_heading(heading.group(2).strip(), level=heading_level)
        elif line.startswith(('- ', '* ')):
            p = doc.add_paragraph(style='List Bullet')
            _add_formatted_runs(p, line[2:])
        elif line.startswith('> '):
            p = doc.add_paragraph(style='Intense Quote')
            _add_formatted_runs(p, line[2:])
        else:
            p = doc.add_paragraph()
            _add_formatted_runs(p, line)
        index += 1
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io
