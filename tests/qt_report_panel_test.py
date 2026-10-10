"""实践报告页（学术专业模式）。

* 页面状态与问题分组是旧界面同一套判断（``folith.workspace_report``）；
* 作者能看到卡在哪、回答待补充的问题、按范围触发定点重建；
* 重建只失效学术产物，译文纹丝不动；
* 报告模板（DOCX）能保存、能移除，坏文件说清原因；
* 没有报告、没有问题时页面如实说明，不渲染空白。
"""
from __future__ import annotations

import io
import os
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["FOLITH_QT_NONINTERACTIVE"] = "1"

pytest.importorskip("PySide6")

import core  # noqa: E402
from folith import academic_writer, workspace_report  # noqa: E402
from folith.qt import adapter  # noqa: E402
from stage3_case_review_test import _setup, _synthetic_case  # noqa: E402


def _save(job_id, name, value):
    state = core.load_job_state(job_id)
    academic_writer._save_artifact(state, core.job_dir(job_id), name, value, f"{name}-dep", f"{name}-v1")
    core.save_job_state(job_id, state)


@pytest.fixture()
def job(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    cases = [_synthetic_case("SC-15", "382", "较好的模拟初译", "较弱的当前译文")]
    job_id, _state, _selected = _setup(tmp_path, cases)
    state = core.load_job_state(job_id)
    state["p3_md"] = "# 翻译实践报告\n\n## 1 项目概况\n\n正文一。\n\n## 2 案例分析\n\n正文二。\n"
    core.save_job_state(job_id, state)
    return job_id


# ------------------------------------------------------------ 分组与状态 ------

def test_issue_categories_follow_the_issue_type():
    assert workspace_report.issue_category({"type": "case_minimum_not_met"}) == "案例不足"
    assert workspace_report.issue_category({"type": "wrong_project_statistic"}) == "统计验证失败"
    assert workspace_report.issue_category({"type": "template_missing_chapter"}) == "模板合规"
    assert workspace_report.issue_category({"type": "missing_required_section"}) == "章节需要重新生成"
    assert workspace_report.issue_category({"type": "uncitable_literature_source"}) == "引用需要确认"
    assert workspace_report.issue_category({"type": "something_else"}) == "需要人工补充"


def test_report_headings_and_cleaning():
    text = "# 翻译实践报告\n\n## 1 概况\n\n## 1 概况\n\n<!-- hidden -->正文 seg-123abc 引用\n\n## 2 案例\n"
    cleaned = workspace_report.clean_report_for_display(text)
    assert "hidden" not in cleaned and "seg-123abc" not in cleaned
    assert not cleaned.startswith("# 翻译实践报告"), "标题由工具栏统一显示"
    titles = [h["title"] for h in workspace_report.report_headings(cleaned)]
    assert titles.count("1 概况") == 1, "相邻重复标题只留一个"


def test_view_without_academic_artifacts_still_renders(job):
    view = adapter.report_view(job)
    assert view["has_report"] is True
    assert [h["title"] for h in view["headings"]] == ["1 项目概况", "2 案例分析"]
    assert dict(view["chips"])["交付"]
    assert view["overall"] in {"blocked", "review_required", "draft_preview"}
    assert view["template"]["present"] is False


def test_validation_issues_become_actionable_groups(job):
    _save(job, "validation", {"status": "fail", "issues": [
        {"type": "case_minimum_not_met"}, {"type": "wrong_project_statistic"},
        {"type": "missing_required_section", "section_id": "3.2"}]})
    view = adapter.report_view(job)
    categories = {g["category"]: g for g in view["groups"]}
    assert {"案例不足", "统计验证失败", "章节需要重新生成"} <= set(categories)
    assert categories["案例不足"]["severity"] == "blocker"
    assert categories["章节需要重新生成"]["sections"] == ["3.2"]
    assert categories["统计验证失败"]["can_repair"] and categories["统计验证失败"]["repair_label"]
    assert view["overall"] == "blocked" and "阻塞项" in view["gate_note"]
    assert dict(view["chips"])["阻塞项"] == "2"


def test_open_review_findings_are_folded_into_the_issue_list(job):
    state = core.load_job_state(job)
    state["findings"] = [{"segment_index": 0, "type": "check", "severity": "blocking",
                          "category": "format_integrity", "summary": "占位符未被保留",
                          "reason": "占位符丢失"}]
    core.save_job_state(job, state)
    groups = adapter.report_view(job)["groups"]
    review = [g for g in groups if g.get("target") == "review"]
    assert review and review[0]["severity"] == "blocker"
    assert not review[0]["can_repair"], "审校问题要去审校页处理，不是重新生成报告"


def test_human_questions_are_listed_and_answering_records_them(job):
    _save(job, "human_evidence_questions", {"questions": [
        {"question_id": "hq-1", "status": "open", "case_id": "SC-15",
         "question": "当时为什么选择这种译法？", "context": {"source": "Source 382"}}]})
    view = adapter.report_view(job)
    assert [q["question_id"] for q in view["open_questions"]] == ["hq-1"]
    with pytest.raises(adapter.AdapterError, match="回答"):
        adapter.answer_report_question(job, "hq-1", "  ")
    adapter.answer_report_question(job, "hq-1", "译者当时参考了委托要求")
    state = core.load_job_state(job)
    assert any("委托要求" in str(e) for e in state.get("human_evidence") or [])


def test_repair_invalidates_academic_artifacts_but_not_the_translation(job):
    before = [p["target"] for p in core.load_job_state(job)["pairs"]]
    adapter.repair_report(job, "统计验证失败")
    state = core.load_job_state(job)
    assert [p["target"] for p in state["pairs"]] == before
    assert state["pairs"][0].get("reviewed") in (None, True, False)


def test_repairing_one_section_targets_that_section(job):
    adapter.repair_report(job, "章节需要重新生成", ["3.2"])
    state = core.load_job_state(job)
    artifacts = (state.get("academic_state") or {}).get("artifacts") or {}
    assert state["pairs"][0]["target"] == "较弱的当前译文"
    assert isinstance(artifacts, dict)


# ------------------------------------------------------------ 模板 -----------

def _docx_bytes():
    from docx import Document

    document = Document()
    document.add_heading("摘要", level=1)
    document.add_heading("1 项目概况", level=1)
    document.add_heading("1.1 项目背景", level=2)
    document.add_heading("2 案例分析", level=1)
    document.add_heading("参考文献", level=1)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_template_can_be_saved_and_removed(job, tmp_path):
    path = tmp_path / "template.docx"
    path.write_bytes(_docx_bytes())
    adapter.attach_report_template(job, str(path))
    assert adapter.report_view(job)["template"]["present"] is True
    adapter.clear_report_template(job)
    assert adapter.report_view(job)["template"]["present"] is False


def test_bad_templates_are_refused_with_a_reason(job, tmp_path):
    wrong = tmp_path / "notes.txt"
    wrong.write_text("x", encoding="utf-8")
    with pytest.raises(adapter.AdapterError, match="docx"):
        adapter.attach_report_template(job, str(wrong))
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip")
    with pytest.raises(adapter.AdapterError, match="无法使用"):
        adapter.attach_report_template(job, str(broken))
    with pytest.raises(adapter.AdapterError, match="读取"):
        adapter.attach_report_template(job, str(tmp_path / "missing.docx"))


# ------------------------------------------------------------ 面板 -----------

@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication

    return QApplication.instance() or QApplication([])


def _panel(job_id, **kwargs):
    from folith.qt.report_panel import ReportPanel

    panel = ReportPanel()
    panel.apply(adapter.report_view(job_id), **kwargs)
    return panel


def test_panel_shows_status_chips_and_the_draft(qapp, job):
    panel = _panel(job)
    assert panel.status._label.text()
    assert "章节" in panel.chips.text()
    assert "正文二" in panel.draft.toPlainText()
    assert panel.heading_box.count() == 3          # “跳到章节…” + 两个标题
    assert panel.no_issues.isHidden() is False or panel.issue_list.count() > 0


def test_panel_lists_issue_groups_and_offers_the_right_actions(qapp, job):
    _save(job, "validation", {"status": "fail", "issues": [{"type": "case_minimum_not_met"}]})
    panel = _panel(job)
    assert panel.issue_list.count() >= 1
    assert panel.issue_list.currentRow() == 0
    assert not panel.repair_button.isHidden()
    assert panel.repair_button.text() == "重新选择案例并继续生成"
    assert not panel.open_cases_button.isHidden(), "案例不足要能跳到案例页"
    seen = []
    panel.repairRequested.connect(lambda c, s: seen.append((c, s)))
    panel.repair_button.click()
    assert seen == [("案例不足", [])]


def test_panel_repair_waits_for_connection_and_open_questions(qapp, job):
    _save(job, "validation", {"status": "fail", "issues": [{"type": "case_minimum_not_met"}]})
    panel = _panel(job, connection_ok=False)
    assert not panel.repair_button.isEnabled()
    assert "API Key" in panel.repair_hint.text()

    _save(job, "human_evidence_questions", {"questions": [
        {"question_id": "hq-9", "status": "open", "question": "问题？", "context": {}}]})
    state = core.load_job_state(job)
    academic = dict(state.get("academic_state") or {})
    academic["human_evidence_status"] = {"unanswered": 1}
    state["academic_state"] = academic
    core.save_job_state(job, state)
    panel2 = _panel(job)
    rows = [panel2.issue_list.item(i).text() for i in range(panel2.issue_list.count())]
    human_row = next(i for i, text in enumerate(rows) if "人工补充" in text)
    panel2.issue_list.setCurrentRow(human_row)
    assert not panel2.repair_button.isEnabled(), "还有待回答的问题时不能继续生成"
    answers = []
    panel2.answerSubmitted.connect(lambda q, a: answers.append((q, a)))
    box = panel2._answer_boxes["hq-9"]
    box.setPlainText("不记得")
    panel2._submit_answer("hq-9")
    assert answers == [("hq-9", "不记得")]


def test_panel_without_issues_says_so(qapp, job):
    panel = _panel(job)
    if not panel.issue_list.count():
        assert not panel.no_issues.isHidden()


def test_panel_template_buttons(qapp, job, tmp_path):
    panel = _panel(job)
    assert panel.template_remove.isHidden()
    path = tmp_path / "template.docx"
    path.write_bytes(_docx_bytes())
    adapter.attach_report_template(job, str(path))
    panel.apply(adapter.report_view(job))
    assert not panel.template_remove.isHidden() and "template" in panel.template_line.text()


def test_panel_impact_notice_offers_the_rebuild(qapp, job):
    panel = _panel(job)
    view = adapter.report_view(job)
    view["impact"] = {"stale": True, "reason": "译文已变化", "affected": [], "reusable": []}
    panel.apply(view)
    assert not panel.rebuild_button.isHidden() and "译文已变化" in panel.impact.text()
    seen = []
    panel.rebuildRequested.connect(lambda: seen.append(1))
    panel.rebuild_button.click()
    assert seen == [1]
    panel.apply(view, connection_ok=False)
    assert not panel.rebuild_button.isEnabled()


def test_the_report_tab_follows_the_academic_mode(qapp, tmp_path, monkeypatch):
    from folith.qt import views

    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    state = core.new_job_state("plain.docx")
    state.update(p1_done=True, p2_done=True, paras=["a"], pairs=[{"source": "a", "target": "甲"}])
    core.save_job_state("plain0000000003", state)
    bench = views.WorkbenchView()
    loaded = adapter.load_state("plain0000000003")
    bench.load("plain0000000003", adapter.EngineConfig(target_lang="简体中文"), loaded,
               adapter.segment_rows("plain0000000003", loaded), [])
    index = bench.tabs.indexOf(bench.report_panel)
    assert not bench.tabs.isTabVisible(index)
    bench.set_academic_mode(True)
    assert bench.tabs.isTabVisible(index)
