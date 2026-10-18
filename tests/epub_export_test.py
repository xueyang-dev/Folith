"""Source-to-delivery EPUB regressions with actual PDF/DOCX image fixtures."""
import io
import json
from copy import deepcopy
from pathlib import PurePosixPath
from xml.etree import ElementTree as ET
from zipfile import ZIP_STORED, ZipFile

import fitz
import pytest
from docx import Document
from docx.shared import Inches

import core
from folith.epub import SOURCE_REFS, EpubExportError, bind_source_references, build_translation_epub

XHTML = {"h": "http://www.w3.org/1999/xhtml"}
OPF = {"o": "http://www.idpf.org/2007/opf"}


def illustration():
    with fitz.open() as doc:
        page = doc.new_page(width=240, height=120)
        page.draw_rect(fitz.Rect(10, 10, 230, 110), color=(0.1, 0.35, 0.65), fill=(0.85, 0.93, 1))
        page.insert_text((30, 65), "Original illustration", fontsize=17)
        return page.get_pixmap(alpha=False).tobytes("png")


def docx_source():
    doc = Document()
    doc.add_heading("Chapter 1: An Illustrated Book", level=1)
    doc.add_paragraph("The first paragraph describes the picture. It has two sentences.")
    doc.add_picture(io.BytesIO(illustration()), width=Inches(3))
    doc.add_paragraph("Figure 1. A blue illustration.", style="Caption")
    doc.add_paragraph("The second paragraph follows the picture.")
    doc.add_heading("Chapter 2: Reading", level=1)
    doc.add_paragraph("The font size can change while the picture stays in order.")
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def pdf_source():
    with fitz.open() as doc:
        page = doc.new_page(width=500, height=650)
        page.insert_text((50, 70), "Chapter 1: An Illustrated Book", fontsize=19)
        page.insert_text((50, 115), "The first paragraph describes the picture.", fontsize=11)
        # Deliberately add later text before the picture to exercise geometric
        # ordering rather than assuming PDF drawing order is reading order.
        page.insert_text((50, 380), "The second paragraph follows the picture.", fontsize=11)
        page.insert_image(fitz.Rect(70, 150, 430, 330), stream=illustration())
        page.insert_text((80, 350), "Figure 1. A blue illustration.", fontsize=9)
        return doc.tobytes()


def translated_state(filename, source):
    raw, _, report = core.extract_document_paragraphs_with_report(filename, source)
    state = core.new_job_state(filename)
    state["source_paras"] = raw
    units, _ = core._apply_source_segmentation(state, raw, mode="sentence")
    state.update(paras=units, extraction_report=report, target_lang="简体中文", p1_done=True, p2_done=True,
                 pairs=[{"source": unit, "target": f"译文{i + 1}：{unit}", "reviewed": True,
                         "target_provenance": "human_edit"} for i, unit in enumerate(units)])
    return state


def content(data):
    with ZipFile(io.BytesIO(data)) as archive:
        pages = [archive.read(name).decode() for name in sorted(archive.namelist())
                 if "/chapter-" in name]
    return "".join(pages)


def validate_package(data):
    with ZipFile(io.BytesIO(data)) as archive:
        info = archive.infolist()[0]
        assert info.filename == "mimetype"
        assert info.compress_type == ZIP_STORED
        assert archive.read("mimetype") == b"application/epub+zip"
        container = ET.fromstring(archive.read("META-INF/container.xml"))
        rootfile = container.find(".//{urn:oasis:names:tc:opendocument:xmlns:container}rootfile")
        assert rootfile.get("full-path") == "EPUB/content.opf"
        package = ET.fromstring(archive.read(rootfile.get("full-path")))
        manifest = {item.get("id"): item for item in package.findall("o:manifest/o:item", OPF)}
        assert sum(item.get("properties") == "nav" for item in manifest.values()) == 1
        for item in manifest.values():
            assert "EPUB/" + item.get("href") in archive.namelist()
        for item in package.findall("o:spine/o:itemref", OPF):
            assert item.get("idref") in manifest
        for name in archive.namelist():
            if name.endswith((".xhtml", ".ncx", ".opf", ".xml")):
                root = ET.fromstring(archive.read(name))
                for tag, attr in (("h:img", "src"), ("h:a", "href")):
                    for node in root.findall(".//" + tag, XHTML):
                        path = str(PurePosixPath(name).parent / node.get(attr).split("#")[0])
                        assert path in archive.namelist()
                        if "#" in node.get(attr):
                            fragment = node.get(attr).split("#", 1)[1]
                            destination = ET.fromstring(archive.read(path))
                            assert any(element.get("id") == fragment for element in destination.iter())


@pytest.mark.parametrize("filename,maker", [("illustrated.docx", docx_source), ("illustrated.pdf", pdf_source)])
def test_actual_source_to_epub_preserves_image_position(filename, maker):
    source = maker()
    state = translated_state(filename, source)
    before = deepcopy(state)
    data, report = build_translation_epub(state, source, job_id="illustrated")
    validate_package(data)
    assert state == before  # building a preview cannot invalidate a snapshot
    assert report["image_count"] == 1
    book = content(data)
    first = next(pair["target"] for pair in state["pairs"] if "first paragraph" in pair["source"])
    second = next(pair["target"] for pair in state["pairs"] if "second paragraph" in pair["source"])
    assert book.index(first) < book.index("<img ") < book.index(second)
    assert book.count("<img ") == 1
    assert book.index("<img ") < book.index("Figure 1.") < book.index(second)
    if filename.endswith(".pdf"):
        assert report["original_caption_count"] == 1
    else:
        assert report["chapter_count"] == 2
        # Both translated sentences in the same source paragraph form one p.
        with ZipFile(io.BytesIO(data)) as archive:
            root = ET.fromstring(archive.read("EPUB/chapter-0001.xhtml"))
            assert any("first paragraph" in "".join(p.itertext()) and "two sentences" in "".join(p.itertext())
                       for p in root.findall(".//h:p", XHTML))


def test_structure_edits_keep_sources_and_never_duplicate_images(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    source = docx_source()
    state = translated_state("edited.docx", source)
    core.save_source("edited", source)
    core.save_job_state("edited", state)
    state = core.mutate_translation_segments("edited", 1, "split", source_a="The first paragraph",
                                            source_b="describes the picture.", target_a="图前第一部分", target_b="图前第二部分")
    assert state["pairs"][1][SOURCE_REFS] == state["pairs"][2][SOURCE_REFS]
    state = core.mutate_translation_segments("edited", 1, "merge")
    data, _ = build_translation_epub(state, source)
    assert content(data).count("<img ") == 1
    assert content(data).index("图前第一部分") < content(data).index("<img ")
    refs = state["pairs"][1][SOURCE_REFS]
    state = core.mutate_translation_segments("edited", 1, "exclude", exclude_reason="测试排除")
    data, _ = build_translation_epub(state, source)
    assert content(data).count("<img ") == 1
    assert "图前第一部分" not in content(data)
    state = core.mutate_translation_segments("edited", 1, "include", excluded_id=state["excluded_segments"][0]["excluded_id"])
    assert state["pairs"][1][SOURCE_REFS] == refs
    data, _ = build_translation_epub(state, source)
    assert "图前第一部分" in content(data)
    state = core.undo_translation_segment_structure("edited")
    data, _ = build_translation_epub(state, source)
    assert "图前第一部分" not in content(data)
    assert content(data).count("<img ") == 1
    state = core.undo_translation_segment_structure("edited")
    data, _ = build_translation_epub(state, source)
    assert "图前第一部分" in content(data)


def test_cleanup_merge_and_manual_merge_across_picture_keep_target_once(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    source = docx_source()
    state = translated_state("merged.docx", source)
    # Two source paragraphs become one cleanup paragraph without losing either
    # origin. The picture remains after that whole translated paragraph.
    raw = state["source_paras"]
    cleaned = [raw[0], raw[1] + " " + raw[2], *raw[3:]]
    units, _ = core._apply_source_segmentation(state, cleaned, mode="sentence")
    state["source_cleanup"] = {"output_input_indices": [[0], [1, 2], *[[i] for i in range(3, len(raw))]]}
    state["paras"] = units
    state["pairs"] = [{"source": unit, "target": f"合并译文{i}"} for i, unit in enumerate(units)]
    assert bind_source_references(state)
    assert state["pairs"][1][SOURCE_REFS] == [1, 2]
    core.save_source("merged", source)
    core.save_job_state("merged", state)
    state = core.mutate_translation_segments("merged", 3, "merge")
    assert len(state["pairs"][3][SOURCE_REFS]) > 1
    data, _ = build_translation_epub(state, source)
    assert content(data).count("合并译文3") == 1
    assert content(data).count("<img ") == 1


def test_unmapped_legacy_edit_and_wrong_source_fail_explicitly():
    source = docx_source()
    state = translated_state("legacy.docx", source)
    state["segment_structure_history"] = [{"operation": "merge"}]
    with pytest.raises(EpubExportError, match="来源关联"):
        build_translation_epub(state, source)
    state.pop("segment_structure_history")
    state["source_paras"][0] = "A different source"
    with pytest.raises(EpubExportError, match="来源不一致"):
        build_translation_epub(state, source)


def test_escaping_missing_targets_and_unchanged_media():
    source = docx_source()
    state = translated_state("escape.docx", source)
    state["pairs"][1]["target"] = '<literal & text> 😀'
    state["pairs"][2]["target"] = ""
    data, report = build_translation_epub(state, source)
    validate_package(data)
    assert "&lt;literal &amp; text&gt; 😀" in content(data)
    assert report["missing_translation_count"] == 1
    with ZipFile(io.BytesIO(data)) as archive:
        assert archive.read("EPUB/images/image-0001.png") == illustration()


def test_delivery_option_bundle_and_frozen_snapshot(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    source = docx_source()
    state = translated_state("delivery.docx", source)
    config = {key: False for key in core.default_delivery_config()}
    config["deliver_epub"] = True
    state["delivery_config"] = config
    core.save_source("epubdelivery", source)
    core.save_job_state("epubdelivery", state)
    bundle = core.build_delivery_assets("epubdelivery", state)
    assert set(bundle) == {"translation.epub", "epub_export_report.json", "delivery_manifest.json"}
    assert json.loads(bundle["epub_export_report.json"])["image_count"] == 1
    manifest = json.loads(bundle["delivery_manifest.json"])
    assert "translation.epub" in str(manifest)
    snapshot = core.create_delivery_snapshot("epubdelivery", state)
    assets = core.delivery_snapshot_assets("epubdelivery", snapshot["snapshot_version"])
    validate_package(assets["translation.epub"])
    assert core.default_delivery_config()["deliver_epub"] is False


def test_scanned_page_does_not_masquerade_as_an_inline_picture():
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_image(page.rect, stream=illustration())
        source = doc.tobytes()
    with pytest.raises(EpubExportError, match="扫描页"):
        build_translation_epub({"filename": "scan.pdf", "pairs": [{"source": "OCR", "target": "识别译文"}]}, source)


def visual_pdf_state():
    from folith.document_ir import PageEnvelope, SemanticBlock
    from folith.document_understanding import VisualDocumentAssembler
    from folith.document_ir import PageUnderstanding
    from folith.spatial_index import SpatialTextIndex
    import time

    with fitz.open() as pdf:
        page = pdf.new_page(width=500, height=650)
        page.insert_image(page.rect, stream=illustration())
        source = pdf.tobytes()
    assembler = VisualDocumentAssembler("scan.pdf")
    blocks = [SemanticBlock("title", "A scanned chapter"),
              SemanticBlock("paragraph", "Text before the picture. Another sentence."),
              SemanticBlock("figure", ""), SemanticBlock("caption", "A caption."),
              SemanticBlock("paragraph", "Text after the picture."),
              SemanticBlock("figure", "")]
    anchors = [[{"page_id": "p1", "bbox": None, "grounding_method": "page_only",
                 "grounding_quality": "page_only"}] for _ in blocks]
    anchors[2] = [{"page_id": "p1", "bbox": [0.2, 0.3, 0.8, 0.5],
                   "grounding_method": "region_alignment", "grounding_quality": "strong"}]
    assembler.add_page(PageEnvelope("scan.pdf", "p1", 0, 500, 650),
                       (PageUnderstanding(blocks, "vlm"), anchors,
                        SpatialTextIndex("p1"), time.perf_counter()))
    raw = assembler.paragraphs
    state = core.new_job_state("scan.pdf")
    state["source_paras"] = raw
    state["extraction_report"] = {"semantic_source": "vlm", "document_ir": assembler.document}
    units, _ = core._apply_source_segmentation(state, raw, mode="sentence")
    state.update(paras=units, target_lang="简体中文", p1_done=True, p2_done=True,
                 pairs=[{"source": text, "target": f"识别译文{i}"} for i, text in enumerate(units)])
    return source, state


@pytest.mark.parametrize("semantic_source", ["vlm", "mixed"])
def test_semantic_pdf_epub_uses_recognized_text_and_only_grounded_figures(semantic_source):
    source, state = visual_pdf_state()
    state["extraction_report"]["semantic_source"] = semantic_source
    before = deepcopy(state)
    data, report = build_translation_epub(state, source)
    validate_package(data)
    assert state == before
    book = content(data)
    body = ET.fromstring(book).find("h:body", XHTML)
    body_text = "".join(body.itertext())
    for pair in state["pairs"]:
        assert body_text.count(pair["target"]) == 1
    assert report["image_count"] == 1
    assert book.index("识别译文1") < book.index("<img ") < book.index("识别译文3")
    assert "<h1" in book
    assert any("缺少完整区域坐标" in warning for warning in report["warnings"])
    with ZipFile(io.BytesIO(data)) as archive:
        pix = fitz.Pixmap(archive.read("EPUB/images/image-0001.png"))
        assert (pix.width, pix.height) == (600, 260)  # region, not the whole scan


def test_semantic_epub_keeps_cat_merge_and_exclusion_references(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    source, state = visual_pdf_state()
    core.save_source("semantic", source)
    # First bind the source identity before structure edits, as the editor does.
    assert bind_source_references(state)
    core.save_job_state("semantic", state)
    state = core.mutate_translation_segments("semantic", 1, "merge")
    data, _ = build_translation_epub(state, source)
    assert content(data).count("识别译文1") == 1
    assert content(data).count("识别译文2") == 1
    state = core.mutate_translation_segments("semantic", 1, "exclude", exclude_reason="test")
    data, _ = build_translation_epub(state, source)
    assert "识别译文1" not in content(data)
    assert content(data).count("<img ") == 1


def test_semantic_epub_rejects_mismatched_canonical_source():
    source, state = visual_pdf_state()
    state["source_paras"][0] = "Different recognized source"
    with pytest.raises(EpubExportError, match="来源不一致"):
        build_translation_epub(state, source)


def test_mixed_pdf_uses_page_order_and_does_not_embed_full_scan():
    from folith.document_ir import PageUnderstanding, SemanticBlock
    from folith.document_understanding import understand_mixed_pdf

    class Provider:
        def understand_page(self, page):
            return PageUnderstanding([SemanticBlock("paragraph", "Recognized middle page.")], "vlm")

    with fitz.open(stream=pdf_source(), filetype="pdf") as pdf:
        page = pdf.new_page(width=500, height=650)
        page.insert_image(page.rect, stream=illustration())
        page = pdf.new_page(width=500, height=650)
        page.insert_text((50, 70), "The final native paragraph.")
        source = pdf.tobytes()
    raw, document = understand_mixed_pdf("mixed.pdf", source, [1], Provider())
    state = core.new_job_state("mixed.pdf")
    state.update(source_paras=raw, target_lang="简体中文",
                 extraction_report={"semantic_source": "mixed", "document_ir": document})
    units, _ = core._apply_source_segmentation(state, raw, mode="paragraph")
    state.update(paras=units, pairs=[{"source": text, "target": f"译文{i}"}
                                     for i, text in enumerate(units)])
    data, report = build_translation_epub(state, source)
    validate_package(data)
    book = content(data)
    # The native illustration survives; the full-page scan does not.
    assert report["image_count"] == 1
    bodies = []
    with ZipFile(io.BytesIO(data)) as archive:
        for name in sorted(archive.namelist()):
            if "/chapter-" in name:
                bodies.extend("".join(node.itertext()) for node in
                              ET.fromstring(archive.read(name)).find("h:body", XHTML)
                              if node.tag in {f'{{{XHTML["h"]}}}p', f'{{{XHTML["h"]}}}h1'})
    assert bodies == [pair["target"] for pair in state["pairs"]]
    first = next(pair["target"] for pair in state["pairs"] if "first paragraph" in pair["source"])
    second = next(pair["target"] for pair in state["pairs"] if "second paragraph" in pair["source"])
    assert book.index(first) < book.index("<img ") < book.index(second)


def test_image_only_plate_in_text_pdf_is_preserved():
    with fitz.open(stream=pdf_source(), filetype="pdf") as doc:
        page = doc.new_page(width=500, height=650)
        page.insert_image(fitz.Rect(50, 100, 450, 300), stream=illustration())
        source = doc.tobytes()
    state = translated_state("plate.pdf", source)
    data, report = build_translation_epub(state, source)
    assert report["image_count"] == 2
    assert content(data).count("<img ") == 2
    assert any("第 2 页" in warning and "原样保留" in warning for warning in report["warnings"])


@pytest.mark.parametrize("filename", ["cover.pdf", "cover.docx"])
def test_original_cover_is_first_without_generated_reading_notes(filename):
    if filename.endswith(".pdf"):
        with fitz.open(stream=pdf_source(), filetype="pdf") as doc:
            page = doc.new_page(pno=0, width=500, height=650)
            page.insert_image(page.rect, stream=illustration())
            page.draw_rect(fitz.Rect(20, 20, 480, 630), color=(1, 0, 0))
            expected = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes("png")
            source = doc.tobytes()
    else:
        doc = Document(io.BytesIO(docx_source()))
        picture = doc.add_picture(io.BytesIO(illustration()), width=Inches(3))
        paragraph = picture._inline.getparent().getparent().getparent()
        doc.element.body.insert(0, paragraph)
        stream = io.BytesIO()
        doc.save(stream)
        source, expected = stream.getvalue(), illustration()
    data, report = build_translation_epub(translated_state(filename, source), source)
    validate_package(data)
    assert report["has_cover"] is True
    assert report["image_count"] == 2
    with ZipFile(io.BytesIO(data)) as archive:
        package = ET.fromstring(archive.read("EPUB/content.opf"))
        manifest = {item.get("id"): item for item in package.findall("o:manifest/o:item", OPF)}
        first = package.find("o:spine/o:itemref", OPF).get("idref")
        assert manifest[first].get("href") == "cover.xhtml"
        cover = next(item for item in manifest.values() if item.get("properties") == "cover-image")
        assert archive.read("EPUB/" + cover.get("href")) == expected
        assert cover.get("href") not in content(data)
        assert "EPUB/title.xhtml" not in archive.namelist()
        assert all("阅读说明" not in archive.read(name).decode() for name in archive.namelist() if name.endswith(".xhtml"))
    assert content(data).count("<img ") == 1


def test_contents_entries_link_to_source_matched_chapter_anchors():
    doc = Document()
    doc.add_heading("CONTENTS", level=1)
    doc.add_paragraph("Vertigo ... 9")
    doc.add_paragraph("Back From the Moon")
    doc.add_paragraph("Unmatched title")
    doc.add_heading("VERTIGO", level=1)
    doc.add_paragraph("A night of extraordinary clarity.")
    doc.add_heading("BACK FROM", level=1)
    doc.add_heading("THE MOON", level=1)
    doc.add_paragraph("The final days of our training.")
    stream = io.BytesIO()
    doc.save(stream)
    source = stream.getvalue()
    state = translated_state("contents.docx", source)
    units, _ = core._apply_source_segmentation(state, state["source_paras"], mode="paragraph")
    state["paras"] = units
    state["pairs"] = [{"source": unit, "target": f"译文：{unit}"} for unit in units]
    # Directory and chapter translations may differ; correspondence uses source.
    for pair in state["pairs"]:
        if pair["source"] == "Vertigo ... 9":
            pair["target"] = "眩晕"
        elif pair["source"] == "VERTIGO":
            pair["target"] = "第一章：眩晕"
    data, report = build_translation_epub(state, source)
    validate_package(data)
    assert report["linked_toc_count"] == 2
    with ZipFile(io.BytesIO(data)) as archive:
        root = ET.fromstring(archive.read("EPUB/chapter-0001.xhtml"))
        links = root.findall(".//h:a", XHTML)
        assert len(links) == 2
        link = next(link for link in links if link.text == "眩晕")
        path, fragment = link.get("href").split("#")
        target = ET.fromstring(archive.read("EPUB/" + path))
        assert next(node for node in target.iter() if node.get("id") == fragment).text == "第一章：眩晕"
        assert "Unmatched title" in "".join(root.itertext())
        package = ET.fromstring(archive.read("EPUB/content.opf"))
        assert not any(item.get("properties") == "cover-image" for item in package.findall("o:manifest/o:item", OPF))
