"""DOCX write-back pairs translations with paragraphs by identity, not position.

Segments are sentence-level; a paragraph with three sentences is three pairs.
Pairing the i-th unit with the i-th pair shifts every translation after the
first multi-sentence paragraph.  These tests assert the text of *each output
paragraph*, through the real segmentation and CAT edit functions.
"""
from __future__ import annotations

import io

import pytest
from docx import Document

import core
from folith.docx_roundtrip import (
    InvalidAnchorError,
    build_roundtrip_docx,
    verify_roundtrip_docx,
)

PARAS = [
    ("Chapter One", "Heading 1"),
    ("Alpha went home. Beta stayed behind. Gamma left early.", None),
    ("Delta was the last one.", None),
    ("Appendix", "Heading 1"),
    ("Epsilon closes the book.", None),
]


def _source() -> bytes:
    doc = Document()
    for text, style in PARAS:
        if style:
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _state(source: bytes, name="identity.docx") -> dict:
    raw, _, report = core.extract_document_paragraphs_with_report(name, source)
    state = core.new_job_state(name)
    state["source_paras"] = raw
    units, _ = core._apply_source_segmentation(state, raw, mode="sentence")
    state.update(
        paras=units, extraction_report=report, target_lang="简体中文",
        p1_done=True, p2_done=True,
        pairs=[{"source": unit, "target": f"译[{unit}]", "reviewed": True,
                "target_provenance": "human_edit"} for unit in units])
    return state


def _texts(data) -> list[str]:
    raw = data.getvalue() if hasattr(data, "getvalue") else data
    return [p.text for p in Document(io.BytesIO(raw)).paragraphs]


def _expected(*, drop=()) -> list[str]:
    """Each paragraph: its sentences' translations in order, or the source."""
    out = []
    for index, (text, _style) in enumerate(PARAS):
        if index in drop:
            out.append(text)
            continue
        sentences = [s for s in ("Alpha went home.", "Beta stayed behind.",
                                 "Gamma left early.") if s in text] or [text]
        out.append(" ".join(f"译[{s}]" for s in sentences))
    return out


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)


def test_multi_sentence_paragraph_keeps_every_translation_in_place():
    source = _source()
    state = _state(source)
    assert len(state["pairs"]) == 7        # the case that used to shift
    data, report = build_roundtrip_docx(state, source)
    assert _texts(data) == _expected()
    assert report["untranslated_units"] == []
    assert verify_roundtrip_docx(source, data.getvalue(), report=report)["passed"]


def test_range_excluded_paragraph_keeps_source_and_is_reported(env):
    source = _source()
    state = _state(source)
    core.save_source("rng", source)
    core.save_job_state("rng", state)
    # Appendix (pair 5) leaves the translation scope.
    state = core.mutate_translation_segments("rng", 5, "exclude", exclude_reason="范围排除")
    data, report = build_roundtrip_docx(state, source)
    assert _texts(data) == _expected(drop=(3,))
    assert [(u["instance_id"], u["reason"]) for u in report["untranslated_units"]] == [
        ("u:doc:p3", "excluded")]


def test_manually_excluded_first_paragraph(env):
    source = _source()
    state = _state(source)
    core.save_source("man", source)
    core.save_job_state("man", state)
    state = core.mutate_translation_segments("man", 0, "exclude", exclude_reason="手动排除")
    data, report = build_roundtrip_docx(state, source)
    assert _texts(data) == _expected(drop=(0,))
    assert [u["instance_id"] for u in report["untranslated_units"]] == ["u:doc:p0"]


def test_split_inside_a_paragraph_joins_both_halves(env):
    source = _source()
    state = _state(source)
    core.save_source("spl", source)
    core.save_job_state("spl", state)
    state = core.mutate_translation_segments(
        "spl", 6, "split", source_a="Epsilon closes", source_b="the book.",
        target_a="结尾甲", target_b="结尾乙")
    data, report = build_roundtrip_docx(state, source)
    expected = _expected()
    expected[4] = "结尾甲结尾乙"
    assert _texts(data) == expected
    assert report["untranslated_units"] == []


def test_merge_inside_a_paragraph_is_one_translation(env):
    source = _source()
    state = _state(source)
    core.save_source("mrg", source)
    core.save_job_state("mrg", state)
    state = core.mutate_translation_segments("mrg", 1, "merge")   # Alpha + Beta
    data, _report = build_roundtrip_docx(state, source)
    texts = _texts(data)
    assert len(texts) == 5 and texts[2] == _expected()[2] and texts[4] == _expected()[4]
    assert "译[Gamma left early.]" in texts[1] and texts[1].startswith("译[Alpha")


def test_merge_across_paragraphs_fails_explicitly(env):
    source = _source()
    state = _state(source)
    core.save_source("xm", source)
    core.save_job_state("xm", state)
    state = core.mutate_translation_segments("xm", 3, "merge")    # Gamma + Delta
    with pytest.raises(InvalidAnchorError, match="跨段落合并"):
        build_roundtrip_docx(state, source)


def test_pairs_that_do_not_belong_to_the_source_fail_explicitly():
    source = _source()
    state = _state(source)
    state["source_paras"][0] = "A different source"
    with pytest.raises(InvalidAnchorError, match="原格式回写不会按位置猜测"):
        build_roundtrip_docx(state, source)


def test_delivery_writes_each_paragraph_correctly(env):
    source = _source()
    state = _state(source)
    state["delivery_config"] = dict(core.default_delivery_config(),
                                    deliver_original_docx=True)
    core.save_source("dlv", source)
    core.save_job_state("dlv", state)
    bundle = core.build_delivery_assets("dlv", state)
    assert _texts(bundle["original_format.docx"]) == _expected()
