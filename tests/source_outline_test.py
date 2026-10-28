"""原文结构解析回归（#27）：模型无关、扫描页如实、与 core 提取逐字对齐。

这一批守护的是「准备」页只读原文结构的两条底线：

- **不猜**：本地读不出的（扫描页、加密、损坏文件）必须如实返回空/未识别，
  绝不把扫描页当成已识别内容，也绝不为看结构去调模型；
- **可对齐**：结构解析的段落与 :func:`core.extract_document_paragraphs_with_report`
  在文字层路径上必须逐字相同——按结构消费这份结果的界面（#28 的翻译范围、
  #61 的原位置定位）靠这份一致性直接复用锚点，不需要第二套解析。
"""
import io

import fitz
import pytest
from docx import Document

import core
from folith.source_outline import (empty_outline, outline_fingerprint,
                                   parse_source_outline)


def _docx(*blocks) -> bytes:
    document = Document()
    for kind, text in blocks:
        if kind == "h":
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf(pages: int, *, toc=None, textless=()) -> bytes:
    """每页有可读正文（页与页内容不同，避免被当页眉页脚归一）。"""
    document = fitz.open()
    for index in range(pages):
        page = document.new_page()
        if index + 1 in textless:
            continue
        page.insert_text((72, 72), f"第{index + 1}章 章节标题文字", fontsize=16)
        for line in range(4):
            page.insert_text(
                (72, 130 + line * 20),
                f"这是第{index + 1}页的第{line + 1}句话，页页不同，用来当正文。",
                fontsize=11)
    if toc:
        document.set_toc(toc)
    return document.tobytes()


# ---------------------------------------------------------------- TXT / Markdown

def test_txt_paragraph_groups_are_identified_structure():
    data = "\n\n".join(f"段落 {index} 的内容。" for index in range(1, 46)).encode()
    outline = parse_source_outline("book.txt", data)
    assert outline["total_paragraphs"] == 45
    assert len(outline["nodes"]) == 3
    first = outline["nodes"][0]
    assert first["kind"] == "group"
    assert first["para_start"] == 0 and first["para_count"] == 20
    assert first["paragraphs"][0] == "段落 1 的内容。"
    # 分组必须恰好覆盖全部段落，一段不漏。
    assert sum(node["para_count"] for node in outline["nodes"]) == 45


def test_markdown_headings_bound_chapters():
    markdown = (b"# Intro\n\nFirst paragraph.\n\n"
                b"## Chapter\n\nBody text here.\n\n"
                b"# Fin\n\nLast words.")
    outline = parse_source_outline("doc.md", markdown)
    assert [(node["title"], node["level"]) for node in outline["nodes"]] == [
        ("Intro", 1), ("Chapter", 2), ("Fin", 1)]
    # 标题本身就是段落序列的一员（与 core 提取一致），落在自己章节的开头。
    assert outline["nodes"][1]["para_start"] == 2
    assert outline["paragraphs"][2] == "Chapter"


# ---------------------------------------------------------------- EPUB

def _epub() -> bytes:
    import zipfile

    chapter = ("<?xml version='1.0' encoding='utf-8'?>"
               "<html xmlns='http://www.w3.org/1999/xhtml'><body>"
               "<h1>{title}</h1><p>{body}</p>"
               "</body></html>")
    container = ("<?xml version='1.0'?>"
                 "<container version='1.0' xmlns='urn:oasis:names:tc:opendocument:"
                 "xmlns:container'><rootfiles><rootfile full-path='OEBPS/content.opf'"
                 " media-type='application/oebps-package+xml'/></rootfiles></container>")
    opf = ("<?xml version='1.0'?>"
           "<package xmlns='http://www.idpf.org/2007/opf' version='2.0' "
           "unique-identifier='id'><metadata/>"
           "<manifest><item id='c1' href='c1.xhtml' "
           "media-type='application/xhtml+xml'/>"
           "<item id='c2' href='c2.xhtml' "
           "media-type='application/xhtml+xml'/></manifest>"
           "<spine><itemref idref='c1'/><itemref idref='c2'/></spine></package>")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as book:
        book.writestr("META-INF/container.xml", container)
        book.writestr("OEBPS/content.opf", opf)
        book.writestr("OEBPS/c1.xhtml",
                      chapter.format(title="第一章", body="第一章的正文内容，足够长。"))
        book.writestr("OEBPS/c2.xhtml",
                      chapter.format(title="第二章", body="第二章的正文内容，足够长。"))
    return buffer.getvalue()


def test_epub_spine_chapters_bound_nodes_and_match_core_extraction():
    data = _epub()
    outline = parse_source_outline("book.epub", data)
    assert [node["title"] for node in outline["nodes"]] == ["第一章", "第二章"]
    assert outline["chapters"] == 2
    # 对齐契约：EPUB 段落序列与 core 提取（extract_text_source 同一批函数）逐字相同。
    extracted, _warnings, _report = core.extract_document_paragraphs_with_report(
        "book.epub", data, ocr_mode="off", ocr_max_pages=0)
    assert outline["paragraphs"] == [str(item) for item in extracted]


# ---------------------------------------------------------------- DOCX

def test_docx_heading_styles_bound_chapters_and_match_core_extraction():
    data = _docx(
        ("p", "开篇没有标题的一句正文，长度足够。"),
        ("h", "第一章"),
        ("p", "第一章的正文内容，长度足够被保留下来。"),
        ("h", "第二章"),
        ("p", "第二章的正文内容，长度也足够被保留。"),
    )
    outline = parse_source_outline("sample.docx", data)
    assert [(node["title"], node["para_start"], node["para_count"])
            for node in outline["nodes"]] == [
        ("未分章内容", 0, 1),
        ("第一章", 1, 2),
        ("第二章", 3, 2),
    ]
    # 对齐契约：DOCX 段落序列与 core 提取逐字相同。
    extracted, _warnings, _report = core.extract_document_paragraphs_with_report(
        "sample.docx", data, ocr_mode="off", ocr_max_pages=0)
    assert outline["paragraphs"] == list(extracted)


def test_docx_non_heading_styles_do_not_split():
    data = _docx(("h", "唯一一章"), ("p", "第一段。"), ("p", "第二段。"))
    outline = parse_source_outline("sample.docx", data)
    assert len(outline["nodes"]) == 1
    assert outline["nodes"][0]["title"] == "唯一一章"


def test_docx_reuses_styles_but_honors_each_paragraph_outline_level(monkeypatch):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    document = Document()
    for index in range(60):
        paragraph = document.add_paragraph(f"Paragraph {index}")
        if index in (10, 30):
            level = OxmlElement("w:outlineLvl")
            level.set(qn("w:val"), "0" if index == 10 else "1")
            paragraph._p.get_or_add_pPr().append(level)
    data = io.BytesIO()
    document.save(data)
    original = Paragraph.style
    calls = []
    def get_style(paragraph):
        calls.append(1)
        return original.fget(paragraph)
    monkeypatch.setattr(Paragraph, "style", property(get_style, original.fset))
    outline = parse_source_outline("book.docx", data.getvalue())
    assert [(node["title"], node["level"]) for node in outline["nodes"]] == [
        ("未分章内容", 1), ("Paragraph 10", 1), ("Paragraph 30", 2)]
    assert len(calls) == 1


# ---------------------------------------------------------------- PDF

def test_pdf_without_toc_uses_page_nodes():
    data = _pdf(3)
    outline = parse_source_outline("book.pdf", data)
    assert outline["pages"] == 3 and not outline["scanned"]
    assert [node["page_start"] for node in outline["nodes"]] == [1, 2, 3]
    assert all(node["para_count"] >= 1 for node in outline["nodes"])
    # 对齐契约：段落序列与 core 文字层路径逐字相同。
    extracted, _warnings, _report = core.extract_document_paragraphs_with_report(
        "book.pdf", data, ocr_mode="off", ocr_max_pages=0)
    assert outline["paragraphs"] == [str(item) for item in extracted]


def test_pdf_toc_chapters_cover_pages():
    data = _pdf(4, toc=[[1, "上篇", 1], [1, "下篇", 3]])
    outline = parse_source_outline("book.pdf", data)
    assert [node["title"] for node in outline["nodes"]] == ["上篇", "下篇"]
    assert (outline["nodes"][0]["page_start"], outline["nodes"][0]["page_end"]) == (1, 2)
    assert (outline["nodes"][1]["page_start"], outline["nodes"][1]["page_end"]) == (3, 4)
    # 章节覆盖全部段落：按页分给章节的段落总数 = 总段数（本文无游离段落）。
    assert sum(node["para_count"] for node in outline["nodes"]) \
        == outline["total_paragraphs"]


def test_pdf_textless_pages_reported_not_faked():
    data = _pdf(3, textless={2})
    outline = parse_source_outline("book.pdf", data)
    assert outline["textless_pages"] == [2]
    assert not outline["scanned"]
    page_two = next(node for node in outline["nodes"] if node["page_start"] == 2)
    assert page_two["para_count"] == 0
    assert page_two["scanned"] is True  # 单页无文字层 → 节点如实标记"未识别"


def test_fully_scanned_pdf_states_it_clearly():
    data = _pdf(2, textless={1, 2})
    outline = parse_source_outline("book.pdf", data)
    assert outline["scanned"] is True
    assert outline["total_paragraphs"] == 0
    assert all(node["scanned"] for node in outline["nodes"])


def test_encrypted_pdf_returns_no_content_without_error():
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "secret body text that is long enough")
    buffer = io.BytesIO()
    document.save(buffer, encryption=fitz.PDF_ENCRYPT_AES_128,
                  owner_pw="x", user_pw="y")
    outline = parse_source_outline("locked.pdf", buffer.getvalue())
    assert outline.get("encrypted") is True
    assert outline["total_paragraphs"] == 0
    assert outline["nodes"] == []


def test_broken_file_returns_empty_outline_without_error():
    outline = parse_source_outline("book.pdf", b"this is not a pdf at all")
    assert outline["nodes"] == [] and outline["paragraphs"] == []
    assert outline["fingerprint"] == outline_fingerprint(empty_outline("book.pdf"))


# ---------------------------------------------------------------- 指纹与缓存口径

def test_fingerprint_stable_for_same_source_changes_with_structure():
    data = _docx(("h", "第一章"), ("p", "第一段正文。"))
    first = parse_source_outline("sample.docx", data)
    second = parse_source_outline("sample.docx", data)
    assert first["fingerprint"] == second["fingerprint"]
    changed = _docx(("h", "第一章"), ("p", "第一段正文。"), ("p", "多出来的一段。"))
    assert parse_source_outline("sample.docx", changed)["fingerprint"] \
        != first["fingerprint"]


def test_unknown_suffix_returns_empty_outline():
    outline = parse_source_outline("photo.jpg", b"\xff\xd8\xff")
    assert outline["nodes"] == [] and outline["paragraphs"] == []
    assert outline["format"] == "JPG"


@pytest.mark.parametrize("filename", ["book.pdf", "sample.docx", "book.epub",
                                      "doc.md", "notes.txt"])
def test_every_supported_format_returns_finished_outline(filename):
    """统一收尾（total/fingerprint/id 顺序落定）对每种格式都成立，缺省安全。"""
    data = {"book.pdf": _pdf(1), "sample.docx": _docx(("h", "章"), ("p", "正文段落。")),
            "book.epub": _epub(),
            "doc.md": b"# T\n\nBody.", "notes.txt": "一段正文。".encode()}[filename]
    outline = parse_source_outline(filename, data)
    assert outline["fingerprint"]
    assert outline["total_paragraphs"] == len(outline["paragraphs"])
    assert [node["id"] for node in outline["nodes"]] == \
        [f"n{index:04d}" for index in range(len(outline["nodes"]))]



# ---------------------------------------------------------------- #185: 原文层级结构与溯源

def test_pdf_hierarchical_bookmarks_three_levels_and_page_end_bounds():
    """PDF 书签多层解析：至少三层、正确 parent_id、page_end 截断至同级或更高级书签。"""
    toc = [
        [1, "Part 1", 1],
        [2, "Chapter 1.1", 1],
        [3, "Section 1.1.1", 1],
        [3, "Section 1.1.2", 2],
        [2, "Chapter 1.2", 3],
        [1, "Part 2", 5],
    ]
    data = _pdf(6, toc=toc)
    outline = parse_source_outline("book.pdf", data)
    nodes = outline["nodes"]
    assert len(nodes) == 6

    # 1. 至少三个独立层级
    levels = {node["level"] for node in nodes}
    assert {1, 2, 3}.issubset(levels)

    # 2. 节点标题与层级
    expected_tuples = [
        ("Part 1", 1),
        ("Chapter 1.1", 2),
        ("Section 1.1.1", 3),
        ("Section 1.1.2", 3),
        ("Chapter 1.2", 2),
        ("Part 2", 1),
    ]
    assert [(n["title"], n["level"]) for n in nodes] == expected_tuples

    # 3. 父子指针 parent_id
    assert nodes[0]["parent_id"] == ""
    assert nodes[1]["parent_id"] == nodes[0]["id"]
    assert nodes[2]["parent_id"] == nodes[1]["id"]
    assert nodes[3]["parent_id"] == nodes[1]["id"]  # 同级兄弟共用父节点
    assert nodes[4]["parent_id"] == nodes[0]["id"]
    assert nodes[5]["parent_id"] == ""

    # 4. page_end 截止在同级或更高级节点的下一页前
    # Part 1 (lvl 1) -> stops before Part 2 (page 5) -> page 4
    assert (nodes[0]["page_start"], nodes[0]["page_end"]) == (1, 4)
    # Chapter 1.1 (lvl 2) -> stops before Chapter 1.2 (page 3) -> page 2
    assert (nodes[1]["page_start"], nodes[1]["page_end"]) == (1, 2)
    # Section 1.1.1 (lvl 3) -> stops before Section 1.1.2 (page 2) -> page 1
    assert (nodes[2]["page_start"], nodes[2]["page_end"]) == (1, 1)
    # Section 1.1.2 (lvl 3) -> stops before Chapter 1.2 (lvl 2 <= 3, page 3) -> page 2
    assert (nodes[3]["page_start"], nodes[3]["page_end"]) == (2, 2)
    # Chapter 1.2 (lvl 2) -> stops before Part 2 (lvl 1 <= 2, page 5) -> page 4
    assert (nodes[4]["page_start"], nodes[4]["page_end"]) == (3, 4)
    # Part 2 (lvl 1) -> goes to end of doc (page 6)
    assert (nodes[5]["page_start"], nodes[5]["page_end"]) == (5, 6)

    # 5. Provenance
    assert all(n["provenance"] == "native_toc" for n in nodes)


def test_pdf_front_matter_retains_page_group_before_toc():
    """PDF 首页之前的前置页保留为独立节点，provenance 为 page_group。"""
    toc = [[1, "第一章", 2]]
    data = _pdf(3, toc=toc)
    outline = parse_source_outline("book.pdf", data)
    assert len(outline["nodes"]) == 2
    front, ch1 = outline["nodes"]
    assert front["title"] == "前置页"
    assert front["level"] == 1
    assert front["parent_id"] == ""
    assert front["provenance"] == "page_group"
    assert (front["page_start"], front["page_end"]) == (1, 1)

    assert ch1["title"] == "第一章"
    assert ch1["level"] == 1
    assert ch1["parent_id"] == ""
    assert ch1["provenance"] == "native_toc"
    assert (ch1["page_start"], ch1["page_end"]) == (2, 3)


def _epub3_fixture(*, with_nav: bool = True, with_ncx: bool = False,
                   headings_in_html: bool = True) -> bytes:
    import zipfile

    container = ("<?xml version='1.0'?>"
                 "<container version='1.0' xmlns='urn:oasis:names:tc:opendocument:"
                 "xmlns:container'><rootfiles><rootfile full-path='OEBPS/content.opf'"
                 " media-type='application/oebps-package+xml'/></rootfiles></container>")

    manifest_items = [
        "<item id='c1' href='c1.xhtml' media-type='application/xhtml+xml'/>",
        "<item id='c2' href='c2.xhtml' media-type='application/xhtml+xml'/>",
    ]
    if with_nav:
        manifest_items.append("<item id='nav' href='nav.xhtml' media-type='application/xhtml+xml' properties='nav'/>")
    if with_ncx:
        manifest_items.append("<item id='ncx' href='toc.ncx' media-type='application/x-dtbncx+xml'/>")

    spine_attr = " toc='ncx'" if with_ncx else ""
    opf = (f"<?xml version='1.0'?>"
           f"<package xmlns='http://www.idpf.org/2007/opf' version='3.0' unique-identifier='id'>"
           f"<metadata xmlns:dc='http://purl.org/dc/elements/1.1/'><dc:title>Test</dc:title></metadata>"
           f"<manifest>{''.join(manifest_items)}</manifest>"
           f"<spine{spine_attr}><itemref idref='c1'/><itemref idref='c2'/></spine></package>")

    if headings_in_html:
        c1_content = ("<?xml version='1.0' encoding='utf-8'?>"
                      "<html xmlns='http://www.w3.org/1999/xhtml'><body>"
                      "<h1 id='p1'>Part 1 Heading</h1><p>Part 1 text body paragraph.</p>"
                      "<h2 id='s1_1'>Chapter 1.1 Heading</h2><p>Chapter 1.1 text body paragraph.</p>"
                      "<h3 id='s1_1_1'>Section 1.1.1 Heading</h3><p>Deep nested paragraph.</p>"
                      "</body></html>")
        c2_content = ("<?xml version='1.0' encoding='utf-8'?>"
                      "<html xmlns='http://www.w3.org/1999/xhtml'><body>"
                      "<h1 id='p2'>Part 2 Heading</h1><p>Part 2 text body paragraph.</p>"
                      "</body></html>")
    else:
        c1_content = ("<?xml version='1.0' encoding='utf-8'?>"
                      "<html xmlns='http://www.w3.org/1999/xhtml'><body>"
                      "<p id='p1'>First paragraph without heading tag.</p>"
                      "<p id='s1_1'>Second paragraph in first chapter.</p>"
                      "</body></html>")
        c2_content = ("<?xml version='1.0' encoding='utf-8'?>"
                      "<html xmlns='http://www.w3.org/1999/xhtml'><body>"
                      "<p id='p2'>Body paragraph in second chapter.</p>"
                      "</body></html>")

    nav_content = ("<?xml version='1.0' encoding='utf-8'?>"
                   "<html xmlns='http://www.w3.org/1999/xhtml' xmlns:epub='http://www.idpf.org/2007/ops'>"
                   "<body><nav epub:type='toc' id='toc'><h1>TOC</h1><ol>"
                   "<li><a href='c1.xhtml#p1'>Part 1 Verbatim Label</a><ol>"
                   "<li><a href='c1.xhtml#s1_1'>Chapter 1.1 Verbatim Label</a><ol>"
                   "<li><a href='c1.xhtml#s1_1_1'>Section 1.1.1 Verbatim Label</a></li>"
                   "</ol></li></ol></li>"
                   "<li><a href='c2.xhtml#p2'>Part 2 Verbatim Label</a></li>"
                   "</ol></nav></body></html>")

    ncx_content = ("<?xml version='1.0' encoding='utf-8'?>"
                   "<ncx xmlns='http://www.daisy.org/z3986/2005/ncx/' version='2005-1'>"
                   "<head><meta name='dtb:uid' content='12345'/></head>"
                   "<docTitle><text>TOC</text></docTitle><navMap>"
                   "<navPoint id='np1' playOrder='1'>"
                   "<navLabel><text>Part 1 Verbatim Label</text></navLabel>"
                   "<content src='c1.xhtml#p1'/>"
                   "<navPoint id='np1_1' playOrder='2'>"
                   "<navLabel><text>Chapter 1.1 Verbatim Label</text></navLabel>"
                   "<content src='c1.xhtml#s1_1'/>"
                   "<navPoint id='np1_1_1' playOrder='3'>"
                   "<navLabel><text>Section 1.1.1 Verbatim Label</text></navLabel>"
                   "<content src='c1.xhtml#s1_1_1'/>"
                   "</navPoint></navPoint></navPoint>"
                   "<navPoint id='np2' playOrder='4'>"
                   "<navLabel><text>Part 2 Verbatim Label</text></navLabel>"
                   "<content src='c2.xhtml#p2'/>"
                   "</navPoint></navMap></ncx>")

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as book:
        book.writestr("META-INF/container.xml", container)
        book.writestr("OEBPS/content.opf", opf)
        book.writestr("OEBPS/c1.xhtml", c1_content)
        book.writestr("OEBPS/c2.xhtml", c2_content)
        if with_nav:
            book.writestr("OEBPS/nav.xhtml", nav_content)
        if with_ncx:
            book.writestr("OEBPS/toc.ncx", ncx_content)
    return buffer.getvalue()


def test_epub3_nav_xhtml_nested_ol_hierarchy():
    """EPUB3 nav.xhtml 嵌套 <ol> 产生准确层级与 verbatim 标题。"""
    data = _epub3_fixture(with_nav=True, with_ncx=False)
    outline = parse_source_outline("book.epub", data)
    nodes = outline["nodes"]
    assert len(nodes) == 4
    assert [(n["title"], n["level"], n["provenance"]) for n in nodes] == [
        ("Part 1 Verbatim Label", 1, "native_toc"),
        ("Chapter 1.1 Verbatim Label", 2, "native_toc"),
        ("Section 1.1.1 Verbatim Label", 3, "native_toc"),
        ("Part 2 Verbatim Label", 1, "native_toc"),
    ]
    assert nodes[0]["parent_id"] == ""
    assert nodes[1]["parent_id"] == nodes[0]["id"]
    assert nodes[2]["parent_id"] == nodes[1]["id"]
    assert nodes[3]["parent_id"] == ""


def test_epub2_ncx_produces_identical_tree_to_epub3():
    """仅有 toc.ncx 的 EPUB2 产生与 EPUB3 完全相同的目录树。"""
    epub3_data = _epub3_fixture(with_nav=True, with_ncx=False)
    epub2_data = _epub3_fixture(with_nav=False, with_ncx=True)

    o3 = parse_source_outline("book3.epub", epub3_data)
    o2 = parse_source_outline("book2.epub", epub2_data)

    tree3 = [(n["title"], n["level"], n["parent_id"], n["provenance"]) for n in o3["nodes"]]
    tree2 = [(n["title"], n["level"], n["parent_id"], n["provenance"]) for n in o2["nodes"]]
    assert tree3 == tree2


def test_epub_heading_hierarchy_fallback_without_toc():
    """没有 TOC 的 EPUB 回落到 XHTML 的 h1–h6 标题层级而非 spine 边界。"""
    data = _epub3_fixture(with_nav=False, with_ncx=False, headings_in_html=True)
    outline = parse_source_outline("book.epub", data)
    nodes = outline["nodes"]
    assert len(nodes) == 4
    assert [(n["title"], n["level"], n["provenance"]) for n in nodes] == [
        ("Part 1 Heading", 1, "native_heading_style"),
        ("Chapter 1.1 Heading", 2, "native_heading_style"),
        ("Section 1.1.1 Heading", 3, "native_heading_style"),
        ("Part 2 Heading", 1, "native_heading_style"),
    ]
    assert nodes[0]["parent_id"] == ""
    assert nodes[1]["parent_id"] == nodes[0]["id"]
    assert nodes[2]["parent_id"] == nodes[1]["id"]
    assert nodes[3]["parent_id"] == ""


def test_epub_unlisted_spine_item_keeps_spine_boundary_node():
    """TOC 未引用的阅读顺序章节保留为 spine_boundary 节点。"""
    import zipfile

    container = ("<?xml version='1.0'?>"
                 "<container version='1.0' xmlns='urn:oasis:names:tc:opendocument:"
                 "xmlns:container'><rootfiles><rootfile full-path='OEBPS/content.opf'"
                 " media-type='application/oebps-package+xml'/></rootfiles></container>")
    opf = ("<?xml version='1.0'?>"
           "<package xmlns='http://www.idpf.org/2007/opf' version='3.0' unique-identifier='id'>"
           "<metadata xmlns:dc='http://purl.org/dc/elements/1.1/'><dc:title>Test</dc:title></metadata>"
           "<manifest>"
           "<item id='cover' href='cover.xhtml' media-type='application/xhtml+xml'/>"
           "<item id='c1' href='c1.xhtml' media-type='application/xhtml+xml'/>"
           "<item id='nav' href='nav.xhtml' media-type='application/xhtml+xml' properties='nav'/>"
           "</manifest>"
           "<spine><itemref idref='cover'/><itemref idref='c1'/></spine></package>")

    cover_xhtml = ("<?xml version='1.0' encoding='utf-8'?>"
                   "<html xmlns='http://www.w3.org/1999/xhtml'><body>"
                   "<h1>封面</h1><p>封面正文说明文字。</p>"
                   "</body></html>")
    c1_xhtml = ("<?xml version='1.0' encoding='utf-8'?>"
                "<html xmlns='http://www.w3.org/1999/xhtml'><body>"
                "<h1>第一章</h1><p>第一章正文内容。</p>"
                "</body></html>")
    nav_xhtml = ("<?xml version='1.0' encoding='utf-8'?>"
                 "<html xmlns='http://www.w3.org/1999/xhtml' xmlns:epub='http://www.idpf.org/2007/ops'>"
                 "<body><nav epub:type='toc' id='toc'><ol>"
                 "<li><a href='c1.xhtml'>第一章目录条目</a></li>"
                 "</ol></nav></body></html>")

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as book:
        book.writestr("META-INF/container.xml", container)
        book.writestr("OEBPS/content.opf", opf)
        book.writestr("OEBPS/cover.xhtml", cover_xhtml)
        book.writestr("OEBPS/c1.xhtml", c1_xhtml)
        book.writestr("OEBPS/nav.xhtml", nav_xhtml)

    outline = parse_source_outline("book.epub", buffer.getvalue())
    nodes = outline["nodes"]
    assert len(nodes) == 2
    assert nodes[0]["title"] == "封面"
    assert nodes[0]["provenance"] == "spine_boundary"
    assert nodes[1]["title"] == "第一章目录条目"
    assert nodes[1]["provenance"] == "native_toc"


def test_docx_custom_style_outline_lvl_and_numbering():
    """DOCX 自定义样式 w:outlineLvl 以及多级编号 w:numPr/w:ilvl 结构提取。"""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    # 1. 样式带有 w:outlineLvl val="0" (Level 1)
    custom_style = doc.styles.add_style("CustomHeadingL1", WD_STYLE_TYPE.PARAGRAPH)
    s_ppr = custom_style.element.get_or_add_pPr()
    s_lvl = OxmlElement("w:outlineLvl")
    s_lvl.set(qn("w:val"), "0")
    s_ppr.append(s_lvl)

    doc.add_paragraph("Custom Style Title", style="CustomHeadingL1")
    doc.add_paragraph("Body under custom style title.")

    # 2. 直接格式带有 w:outlineLvl val="1" (Level 2)
    p2 = doc.add_paragraph("Direct Outline Level 2")
    p2_ppr = p2._p.get_or_add_pPr()
    p2_lvl = OxmlElement("w:outlineLvl")
    p2_lvl.set(qn("w:val"), "1")
    p2_ppr.append(p2_lvl)
    doc.add_paragraph("Body under direct outline level.")

    # 3. 多级编号 w:numPr / w:ilvl
    p3 = doc.add_paragraph("Numbered 1.0 Heading")
    p3_ppr = p3._p.get_or_add_pPr()
    num_pr3 = OxmlElement("w:numPr")
    ilvl3 = OxmlElement("w:ilvl")
    ilvl3.set(qn("w:val"), "0")  # Level 1
    num_pr3.append(ilvl3)
    p3_ppr.append(num_pr3)

    p4 = doc.add_paragraph("Numbered 1.1 Subheading")
    p4_ppr = p4._p.get_or_add_pPr()
    num_pr4 = OxmlElement("w:numPr")
    ilvl4 = OxmlElement("w:ilvl")
    ilvl4.set(qn("w:val"), "1")  # Level 2
    num_pr4.append(ilvl4)
    p4_ppr.append(num_pr4)

    buffer = io.BytesIO()
    doc.save(buffer)
    data = buffer.getvalue()

    outline = parse_source_outline("test.docx", data)
    nodes = outline["nodes"]
    assert len(nodes) == 4
    assert [(n["title"], n["level"], n["provenance"]) for n in nodes] == [
        ("Custom Style Title", 1, "native_outline_level"),
        ("Direct Outline Level 2", 2, "native_outline_level"),
        ("Numbered 1.0 Heading", 1, "native_outline_level"),
        ("Numbered 1.1 Subheading", 2, "native_outline_level"),
    ]
    assert nodes[0]["parent_id"] == ""
    assert nodes[1]["parent_id"] == nodes[0]["id"]
    assert nodes[2]["parent_id"] == ""
    assert nodes[3]["parent_id"] == nodes[2]["id"]


def test_outline_parser_version_constant_and_exported():
    """OUTLINE_PARSER_VERSION 常量导出并在每次解析结果中携带。"""
    from folith.source_outline import OUTLINE_PARSER_VERSION
    assert isinstance(OUTLINE_PARSER_VERSION, int)
    assert OUTLINE_PARSER_VERSION >= 1

    outline = parse_source_outline("doc.txt", b"Hello world\n\nSecond paragraph")
    assert outline["parser_version"] == OUTLINE_PARSER_VERSION


def test_no_model_network_or_ocr_called_during_outline_parse(monkeypatch):
    """断言解析源大纲时绝不调用大模型、网络或 OCR。"""
    import socket
    import urllib.request

    def _fail(*args, **kwargs):
        pytest.fail("Network / OCR / Model call was made during outline parsing!")

    monkeypatch.setattr(socket.socket, "connect", _fail)
    monkeypatch.setattr(urllib.request, "urlopen", _fail)

    # 尝试解析所有格式
    parse_source_outline("book.pdf", _pdf(2))
    parse_source_outline("sample.docx", _docx(("h", "章"), ("p", "正文")))
    parse_source_outline("book.epub", _epub())
    parse_source_outline("doc.md", b"# Header\n\nContent")
    parse_source_outline("notes.txt", b"First para\n\nSecond para")


def test_every_node_carries_valid_provenance():
    """断言每种格式产出的每个节点都具备合法的 provenance 值。"""
    valid_provenances = {
        "native_toc",
        "native_heading_style",
        "native_outline_level",
        "spine_boundary",
        "page_group",
    }
    cases = [
        ("book.pdf", _pdf(2, toc=[[1, "C1", 1]])),
        ("plain.pdf", _pdf(2)),
        ("sample.docx", _docx(("h", "C1"), ("p", "P1"))),
        ("book.epub", _epub3_fixture(with_nav=True)),
        ("headings.epub", _epub3_fixture(with_nav=False, headings_in_html=True)),
        ("doc.md", b"# Heading\n\nBody"),
        ("notes.txt", b"P1\n\nP2"),
    ]
    for filename, data in cases:
        outline = parse_source_outline(filename, data)
        assert outline["nodes"], f"{filename} produced no nodes"
        for node in outline["nodes"]:
            prov = node.get("provenance")
            assert prov in valid_provenances, f"Invalid provenance {prov} in {filename} node {node['title']}"
