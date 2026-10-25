"""流式源的「原位置」面板与工作台入口（#61）。

* 面板：解析在后台预热，定位展示章节 / 标题路径 / 行号 + 相邻段落并高亮；
  未加载完先如实说「正在读取」，空任务、无匹配都不崩溃；
* 适配器：``source_context_available`` / ``segment_source_context`` 只认带源
  文件的流式任务；
* 工作台：DOCX 任务的「查看原位置」可用并在左栏展开上下文面板；没有可用
  源文件的任务按钮禁用并解释边界，而不是悄悄消失。
"""
from __future__ import annotations

import io
import os
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

pytest.importorskip("PySide6")

import core  # noqa: E402
from folith.qt import adapter  # noqa: E402

SENTENCE = "The committee agreed that every revision must leave a visible trace."


def _docx_bytes(first: str = SENTENCE) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Findings", level=1)
    document.add_paragraph(first)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture()
def docx_job(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "srcpanel00000001"
    state = core.new_job_state("note.docx")
    state.update(p1_done=True, p2_done=True, paras=[SENTENCE],
                 pairs=[{"source": SENTENCE, "target": "译文"}])
    core.save_job_state(job_id, state)
    core.save_source(job_id, _docx_bytes())
    return job_id


@pytest.fixture()
def legacy_job(tmp_path, monkeypatch):
    """没有留存源文件的旧任务：两种定位都不适用。"""
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "srcpanel00000002"
    state = core.new_job_state("legacy.docx")
    state.update(p1_done=True, pairs=[{"source": SENTENCE, "target": "译文"}])
    core.save_job_state(job_id, state)
    return job_id


def test_adapter_availability_and_context(docx_job, legacy_job):
    assert adapter.source_context_available(docx_job) is True
    assert adapter.source_context_available(legacy_job) is False
    assert adapter.source_context_available("") is False
    context = adapter.segment_source_context(docx_job, 0)
    assert context["found"] and context["title"] == "Findings"
    missing = adapter.segment_source_context(docx_job, 7)
    assert missing["found"] is False and "没有原文" in missing["hint"]
    assert adapter.source_document_blocks(docx_job), "预热解析能拿到顺序块"


@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication

    return QApplication.instance() or QApplication([])


def test_panel_locates_and_highlights(qapp, docx_job):
    from folith.qt.source_panel import SourcePanel

    panel = SourcePanel()
    panel.load(docx_job)
    assert not panel.locate_button.isEnabled(), "没有选中段落时不能定位"
    panel.set_segment(0)
    assert panel.locate_button.isEnabled()
    panel.locate(0)  # offscreen：同步读取，行为可预期
    assert panel.location.text() == "Findings"
    assert "已定位" in panel.hint.text()
    plain = panel.view.toPlainText()
    assert SENTENCE in plain, "匹配块要真的展示出来"
    assert "highlight" not in plain, "高亮是富文本标记，不进纯文本"
    assert panel.view.toHtml().count("#ffe08a") >= 1, "匹配块有高亮底色"


def test_panel_without_a_task_stays_empty(qapp):
    from folith.qt.source_panel import SourcePanel

    panel = SourcePanel()
    panel.load("")
    assert not panel.locate_button.isEnabled()
    panel.set_segment(3)
    panel.locate(3)
    assert panel.view.toPlainText() == ""


def test_panel_reports_an_unmatched_segment(qapp, docx_job):
    from folith.qt.source_panel import SourcePanel

    panel = SourcePanel()
    panel.load(docx_job)
    panel.locate(0)  # 先建立 job 上下文
    state = adapter.load_state(docx_job)
    state["pairs"].append({"source": "zzzz qqqq nothing like this", "target": "?"})
    core.save_job_state(docx_job, state)
    panel.locate(1)
    assert "没能" in panel.hint.text() and panel.view.toPlainText() == ""


@pytest.fixture()
def client(tmp_path, monkeypatch, qapp):
    data_dir = tmp_path / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setenv("FOLITH_DESKTOP_DATA_DIR", str(data_dir))
    monkeypatch.setenv("FOLITH_DESKTOP_LOG_DIR", str(tmp_path / "logs"))
    monkeypatch.setenv("FOLITH_OUTPUT_DIR", str(data_dir))
    monkeypatch.setattr(core, "OUTPUT_DIR", data_dir, raising=False)

    from folith.desktop import paths as desktop_paths
    from folith.qt import window as window_module
    from folith.qt import i18n as i18n_module

    i18n_module.set_language("zh-CN")
    paths = desktop_paths.resolve_paths()
    config = adapter.EngineConfig(
        provider="自定义中转站", model="stub-model", api_key="stub-key",
        base_url="", target_lang="简体中文", enable_review=True,
        enable_report=False, enable_understanding=False,
        enable_source_cleanup=False)
    main = window_module.MainWindow(config, paths, None)
    main.show()
    qapp.processEvents()
    yield {"window": main, "adapter": adapter, "core": core, "app": qapp}
    sys.path.insert(0, str(ROOT / "tests"))
    from qt_workers import drain_job_workers
    drain_job_workers(core)


def _seed(client, filename: str, source: bytes | None, job_id="srcpanel-ui-job"):
    adapter = client["adapter"]
    core = client["core"]
    state = core.new_job_state(filename)
    state["pairs"] = [{"source": SENTENCE, "target": "译文"}]
    state["paras"] = [SENTENCE]
    core.save_job_state(job_id, state)
    if source is not None:
        core.save_source(job_id, source)
    return job_id


def test_workbench_opens_the_context_panel_for_docx(client):
    from folith.qt.selftest import open_job_and_wait

    window = client["window"]
    job_id = _seed(client, "note.docx", _docx_bytes())
    open_job_and_wait(window, job_id)
    workbench = window.workbench
    assert workbench._stream_ok and not workbench._pdf_ok
    workbench.select_segment(0)
    client["app"].processEvents()
    assert not workbench.view_source_button.isHidden()
    assert workbench.view_source_button.isEnabled()
    workbench.show_source_location(0)
    client["app"].processEvents()
    assert workbench.source_open()
    assert workbench.left_stack.currentWidget() is workbench.source_panel
    assert workbench.source_panel.location.text() == "Findings"


def test_workbench_explains_when_no_source_survived(client):
    from folith.qt.selftest import open_job_and_wait

    window = client["window"]
    job_id = _seed(client, "legacy.docx", None)
    open_job_and_wait(window, job_id)
    workbench = window.workbench
    assert not workbench._source_available
    workbench.select_segment(0)
    client["app"].processEvents()
    assert not workbench.view_source_button.isHidden(), "解释而非隐藏"
    assert not workbench.view_source_button.isEnabled()
    assert "目前支持" in workbench.view_source_button.toolTip()


if __name__ == "__main__":
    sys.exit(pytest.main([__file__]))
