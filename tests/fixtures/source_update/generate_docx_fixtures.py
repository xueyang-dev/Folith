"""Generate synthetic DOCX fixtures for source update tests."""
from pathlib import Path
from docx import Document


def generate_docx_pair(out_dir: Path):
    out_dir.mkdir(parents=True, exist_ok=True)
    v1_path = out_dir / "sample_v1.docx"
    v2_path = out_dir / "sample_v2.docx"

    # Document v1
    doc1 = Document()
    doc1.add_heading("Title: Project Overview.", level=1)
    doc1.add_paragraph("Introduction to Folith workbench.")
    doc1.add_paragraph("Legacy section to be removed.")
    doc1.add_paragraph("Core architecture and data integrity principles.")
    doc1.add_paragraph("Important safety and privacy notice.")
    doc1.add_paragraph("Future development roadmap and conclusion.")
    doc1.save(str(v1_path))

    # Document v2
    doc2 = Document()
    doc2.add_heading("Title: Project Overview.", level=1)
    doc2.add_paragraph("Important safety and privacy notice.")  # Moved
    doc2.add_paragraph("Introduction to Folith professional workbench.")  # Changed
    doc2.add_paragraph("Core architecture and data integrity principles.")  # Unchanged
    doc2.add_paragraph("New section on professional CAT migration.")  # Added
    doc2.add_paragraph("Future development roadmap and conclusion.")  # Unchanged
    doc2.save(str(v2_path))

    return v1_path, v2_path


if __name__ == "__main__":
    generate_docx_pair(Path(__file__).parent)
    print("Generated DOCX fixtures successfully.")
