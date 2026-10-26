"""Cross-script translation transport and RTL delivery regressions."""
import io
import json
import unicodedata
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document

import core
from folith.segmentation import segment_paragraphs
from folith.textual import text_direction

ARABIC = "OpenAI API: مهلة الطلب 30 ثانية (انظر القسم 2)."
#: Every selectable translation language except Arabic itself (21 - 1).
#: Keeping this aligned with ``adapter.TARGET_LANGUAGES`` makes the "all
#: language combinations" claim checkable instead of a sample.
SELECTABLE_PARTNERS = [
    ("English", "The report is ready."),
    ("简体中文", "报告已经准备好了。"),
    ("繁体中文", "報告已經準備好了。"),
    ("日本語", "報告書の準備ができました。"),
    ("한국어", "보고서가 준비되었습니다."),
    ("Deutsch", "Der Bericht ist fertig."),
    ("Français", "Le rapport est prêt."),
    ("Español", "El informe está listo."),
    ("Italiano", "Il rapporto è pronto."),
    ("Português", "O relatório está pronto."),
    ("Русский", "Отчёт готов."),
    ("Nederlands", "Het rapport is klaar."),
    ("Polski", "Raport jest gotowy."),
    ("Українська", "Звіт готовий."),
    ("Tiếng Việt", "Báo cáo đã sẵn sàng."),
    ("Türkçe", "Rapor hazır."),
    ("فارسی", "گزارش آماده است."),
    ("עברית", "הדוח מוכן."),
    ("اردو", "رپورٹ تیار ہے۔"),
    ("ئۇيغۇرچە", "دوكلات تەييار بولدى."),
]
#: Valid language names outside the selectable catalog. The engine accepts any
#: declared target, so these still need a protocol round trip.
ADDITIONAL_LANGUAGE_NAMES = [
    ("हिन्दी", "रिपोर्ट तैयार है।"),
    ("বাংলা", "প্রতিবেদন প্রস্তুত।"),
    ("ภาษาไทย", "รายงานพร้อมแล้ว"),
    ("Kiswahili", "Ripoti iko tayari."),
    ("Ελληνικά", "Η έκθεση είναι έτοιμη."),
]
W = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


@pytest.mark.parametrize("language,text", SELECTABLE_PARTNERS)
@pytest.mark.parametrize("reverse", [False, True])
def test_every_selectable_language_pairs_with_arabic(language, text, reverse):
    source, target, target_language = (ARABIC, text, language) if reverse else (text, ARABIC, "العربية")

    def response(provider, key, model, system, user, **kwargs):
        assert target_language in system
        assert source in user
        return json.dumps([target], ensure_ascii=False)

    assert core.translate_batch([source], [], [], "", "", target_language,
                                "OpenAI", "unused", "gpt-4o-mini",
                                call_llm_fn=response) == [target]


@pytest.mark.parametrize("language,text", ADDITIONAL_LANGUAGE_NAMES)
@pytest.mark.parametrize("reverse", [False, True])
def test_additional_language_names_reach_translation_protocol(language, text, reverse):
    source, target, target_language = (ARABIC, text, language) if reverse else (text, ARABIC, "العربية")

    def response(provider, key, model, system, user, **kwargs):
        assert target_language in system
        assert source in user
        return json.dumps([target], ensure_ascii=False)

    assert core.translate_batch([source], [], [], "", "", target_language,
                                "OpenAI", "unused", "gpt-4o-mini",
                                call_llm_fn=response) == [target]


@pytest.mark.parametrize("text,language,expected", [
    (ARABIC, "العربية", "rtl"), ("OpenAI API", "ar-SA", "rtl"),
    ("", "现代标准阿拉伯语", "rtl"), (ARABIC, None, "rtl"),
    ("AI: הדוח מוכן", "he", "rtl"), ("API: گزارش آماده است", "fa", "rtl"),
    ("مرحبا", "ar-Latn", "ltr"), ("", "az-Arab", "rtl"),
    ("漢字 ١٢٣٤", None, "ltr"), ("日本語", "日本語", "ltr"),
    ("1234", None, "ltr"), ("", "ar", "rtl"),
])
def test_direction_respects_language_script_and_mixed_prefixes(text, language, expected):
    assert text_direction(text, language) == expected


def test_docx_bilingual_columns_and_annotation_runs_keep_independent_directions():
    pairs = [{"source": "The API request timeout is 30 seconds.", "target": ARABIC}]
    output = core.pairs_to_word(pairs, source_lang="en", target_lang="ar",
                               annotations={0: [{"type": "domain", "src_span": [4, 7],
                                                 "tgt_span": [0, 10]}]})
    doc = Document(output)
    source, target = doc.tables[0].rows[1].cells
    assert source.text == pairs[0]["source"] and target.text == ARABIC
    assert source.paragraphs[0].alignment == 0
    assert target.paragraphs[0]._p.find("w:pPr/w:jc", W).get(f'{{{W["w"]}}}val') == "start"
    assert len(target.paragraphs[0].runs) > 1
    with ZipFile(output) as archive:
        xml = ET.fromstring(archive.read("word/document.xml"))
    assert xml.findall('.//w:bidi[@w:val="1"]', W)


def test_delivery_bundle_passes_language_to_single_bilingual_and_pdf_exports(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    # A Latin-only preserved paragraph still belongs to an Arabic document.
    doc = Document()
    doc.add_paragraph("API")
    source = io.BytesIO()
    doc.save(source)
    state = core.new_job_state("source.docx")
    state.update(p1_done=True, p2_done=True, target_lang="ar", source_lang="en",
                 source_paras=["API"], paras=["API"],
                 pairs=[{"source": "API", "target": "API", "reviewed": True}],
                 delivery_config=core.normalize_delivery_config({
                     "deliver_plain_docx": True, "deliver_bilingual_docx": True,
                     "deliver_pdf": True, "deliver_epub": True}))
    core.save_source("rtl-bundle", source.getvalue())
    bundle = core._delivery_asset_bundle("rtl-bundle", state, "", "", "")
    for name in ["translation.docx", "bilingual.docx"]:
        with ZipFile(io.BytesIO(bundle[name])) as archive:
            xml = ET.fromstring(archive.read("word/document.xml"))
        assert xml.findall('.//w:bidi[@w:val="1"]', W)
        assert xml.findall('.//w:jc[@w:val="start"]', W)
    assert bundle["translation.pdf"].startswith(b"%PDF")
    with ZipFile(io.BytesIO(bundle["translation.epub"])) as archive:
        chapters = [ET.fromstring(archive.read(n)) for n in archive.namelist()
                    if n.endswith(".xhtml")]
        package = ET.fromstring(archive.read("EPUB/content.opf"))
    assert all(c.get("dir") == "rtl" for c in chapters)
    assert package.find("{http://www.idpf.org/2007/opf}spine").get("page-progression-direction") == "rtl"


def test_arabic_question_marks_split_without_losing_text():
    first = "متى يبدأ الاجتماع؟"
    second = "من المسؤول عن تقديم التقرير؟"
    units, _ = segment_paragraphs([first + " " + second], mode="sentence", source_lang="ar")
    assert units == [first, second]
    assert core._count_sentences(first + " " + second) == 2


def test_pdf_arabic_with_latin_prefix_reaches_right_margin():
    data = core.translations_to_pdf(
        [{"target": "API: مهلة الطلب 30 ثانية."}], target_lang="ar")
    with core.fitz.open(stream=data, filetype="pdf") as doc:
        spans = [span for block in doc[0].get_text("dict")["blocks"]
                 for line in block.get("lines", []) for span in line["spans"]
                 if any(unicodedata.bidirectional(c) == "AL" for c in span["text"])]
        assert spans
        assert max(span["bbox"][2] for span in spans) > doc[0].rect.width * 0.85


def test_docx_runs_receive_rtl_flag_for_complex_script():
    pairs = [{"source": "API request", "target": ARABIC}]
    output = core.pairs_to_word(pairs, source_lang="en", target_lang="ar")
    with ZipFile(output) as archive:
        xml = ET.fromstring(archive.read("word/document.xml"))
    rtl_runs = xml.findall(".//w:r/w:rPr/w:rtl", W)
    assert len(rtl_runs) >= 1
    assert any(node.get(f'{{{W["w"]}}}val') == "1" for node in rtl_runs)


def test_cell_direction_and_workbench_delegate_respect_column_language():
    from PySide6.QtCore import Qt, QRect
    from PySide6.QtGui import QStandardItemModel, QStandardItem
    from PySide6.QtWidgets import QApplication, QStyleOptionViewItem
    from folith.qt.segment_editor import WorkbenchRowDelegate, cell_direction

    _app = QApplication.instance() or QApplication([])
    # Direct cell_direction checks
    assert cell_direction("API", "ar") == Qt.LayoutDirection.RightToLeft
    assert cell_direction("API", "en") == Qt.LayoutDirection.LeftToRight
    assert cell_direction(ARABIC, None) == Qt.LayoutDirection.RightToLeft

    # WorkbenchRowDelegate checks
    model = QStandardItemModel(1, 3)
    model.setItem(0, 0, QStandardItem("1"))
    model.setItem(0, 1, QStandardItem("API"))
    model.setItem(0, 2, QStandardItem("API"))

    delegate = WorkbenchRowDelegate()
    delegate.source_language = "en"
    delegate.target_language = "ar"

    opt_src = QStyleOptionViewItem()
    opt_src.rect = QRect(0, 0, 200, 30)
    delegate.initStyleOption(opt_src, model.index(0, 1))
    assert opt_src.direction == Qt.LayoutDirection.LeftToRight

    opt_tgt = QStyleOptionViewItem()
    opt_tgt.rect = QRect(0, 0, 200, 30)
    delegate.initStyleOption(opt_tgt, model.index(0, 2))
    assert opt_tgt.direction == Qt.LayoutDirection.RightToLeft


def test_pdf_long_arabic_paragraph_paginates_without_truncation():
    # Extraction returns shaped presentation forms, so compare per-segment
    # numeric markers rather than character counts.
    sentence = "هذا نص عربي طويل لاختبار تعدد الصفحات والاتجاه من اليمين إلى اليسار."
    markers = [f"{index:04d}" for index in range(90)]
    text = " ".join(f"{sentence} {marker}." for marker in markers)
    data = core.translations_to_pdf([{"target": text}], target_lang="ar")
    with core.fitz.open(stream=data, filetype="pdf") as doc:
        assert len(doc) >= 2
        extracted = "".join(page.get_text() for page in doc)
    compact = "".join(extracted.split())
    assert [marker for marker in markers if marker not in compact] == []


def test_pdf_many_arabic_paragraphs_keep_the_last_segment():
    sentence = "هذا نص عربي طويل لاختبار تعدد الصفحات والاتجاه من اليمين إلى اليسار. "
    pairs = [{"target": f"{sentence} رقم {index}."} for index in range(60)]
    data = core.translations_to_pdf(pairs, target_lang="ar")
    with core.fitz.open(stream=data, filetype="pdf") as doc:
        assert len(doc) >= 3
        extracted = "".join(page.get_text() for page in doc)
        for page in doc:
            arabic = [span for block in page.get_text("dict")["blocks"]
                      for line in block.get("lines", []) for span in line["spans"]
                      if any(0x0600 <= ord(char) <= 0x06FF
                             or 0xFE70 <= ord(char) <= 0xFEFF
                             for char in span["text"])]
            if arabic:
                assert max(span["bbox"][2] for span in arabic) > page.rect.width * 0.85
    assert "59" in "".join(extracted.split())


def test_academic_report_arabic_excerpt_direction():
    from folith.report_template import _insert_content_after
    from docx import Document

    doc = Document()
    anchor = doc.add_paragraph("Heading")
    content = "> [TARGET seg-1]: " + ARABIC
    inserted = _insert_content_after(anchor, content)
    bidi = inserted._p.find("w:pPr/w:bidi", W)
    assert bidi is not None
    assert bidi.get(f'{{{W["w"]}}}val') == "1"


def test_epub_arabic_localized_toc_and_cover():
    from folith.epub import build_translation_epub

    state = {
        "filename": "test.docx",
        "target_lang": "ar",
        "source_lang": "en",
        "source_paras": ["Hello"],
        "pairs": [{"source": "Hello", "target": ARABIC, "source_refs": [0]}],
    }
    doc = Document()
    doc.add_paragraph("Hello")
    buf = io.BytesIO()
    doc.save(buf)
    epub_bytes, _ = build_translation_epub(state, buf.getvalue())
    with ZipFile(io.BytesIO(epub_bytes)) as archive:
        nav = archive.read("EPUB/nav.xhtml").decode("utf-8")
        assert "جدول المحتويات" in nav


def test_font_stack_includes_cross_platform_arabic_coverage():
    from folith.qt import theme
    assert any("Arabic" in f or "Geeza" in f or "Segoe UI" in f for f in theme.FONT_STACK)


def test_an_installed_font_in_the_stack_covers_arabic_glyphs():
    from PySide6.QtGui import QFont, QFontDatabase, QRawFont
    from PySide6.QtWidgets import QApplication

    from folith.qt import theme

    _app = QApplication.instance() or QApplication([])
    sample = "العربيةقَرَأَ؟،١٢٣٤"
    installed = set(QFontDatabase.families())
    covering = []
    for family in theme.FONT_STACK:
        if family not in installed:
            continue
        font = QFont(family)
        font.setPixelSize(13)
        raw = QRawFont.fromFont(font)
        if raw.isValid() and all(raw.supportsCharacter(ord(char)) for char in sample):
            covering.append(family)
    if not covering:
        pytest.skip("no installed font in FONT_STACK covers the Arabic sample")
    assert covering



