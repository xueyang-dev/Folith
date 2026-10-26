"""Tests for DocumentOutline persistence and segment anchoring (#186).

Acceptance criteria verification:
- Deterministic outline persistence in sidecar (source_outline.json).
- Deterministic node -> paragraph -> segment resolution.
- Honest partial results for mixed PDFs / unrecognized paragraphs.
- Segment UID durability; position is never used as a fallback after structural changes.
- Split / merge refusal: reports 'structure_edited' instead of jumping to off-by-N segment.
- Paragraph exclusion shift followed by resolved nodes.
- Reopening task reuses persisted outline without re-parsing.
- Source change or bumped parser version invalidates sidecar.
- _reset_source_dependent_state deletes sidecar on rescan and re-import.
- No new state key carries outline nodes or their paragraph texts.
- Documentation boundary: navigation identity only, never source-update or round-trip anchor.
"""
from __future__ import annotations

import hashlib
import io
import json
import time
from unittest.mock import MagicMock, patch

import fitz
import pytest
from docx import Document

import core
from folith.outline_anchor import (
    OUTLINE_PARSER_VERSION,
    SOURCE_OUTLINE_FILE,
    SOURCE_OUTLINE_VERSION,
    clear_source_outline,
    get_or_create_source_outline,
    load_source_outline,
    match_paragraph_positions,
    node_source_range,
    resolve_outline_mapping,
    save_source_outline,
    source_outline_path,
    structure_revision,
)
from folith.source_outline import empty_outline, outline_fingerprint, parse_source_outline


# ---------------------------------------------------------------------------
# Fixtures and helpers
# ---------------------------------------------------------------------------

def _docx(*blocks) -> bytes:
    doc = Document()
    for kind, text in blocks:
        if kind == "h":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf(pages: int, *, toc=None) -> bytes:
    doc = fitz.open()
    for index in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"Chapter {index + 1} Heading", fontsize=16)
        for line in range(2):
            page.insert_text(
                (72, 120 + line * 20),
                f"Body text for page {index + 1} line {line + 1}.",
                fontsize=11)
    if toc:
        doc.set_toc(toc)
    return doc.tobytes()


def _epub_fixture() -> bytes:
    from tests.source_outline_test import _epub3_fixture
    return _epub3_fixture(with_nav=True, with_ncx=False)


def _setup_task(tmp_path, job_id, filename, file_bytes, *, raw_paras, cleaned_paras=None, pairs=None):
    job_d = tmp_path / job_id
    job_d.mkdir(parents=True, exist_ok=True)
    (job_d / "source.bin").write_bytes(file_bytes)

    if cleaned_paras is None:
        cleaned_paras = list(raw_paras)

    if pairs is None:
        pairs = [{"source": p, "target": f"Trans: {p}"} for p in cleaned_paras]

    state = {
        "job_id": job_id,
        "filename": filename,
        "p1_done": True,
        "source_paras": list(raw_paras),
        "source_paragraphs": list(cleaned_paras),
        "paras": [p.get("source") for p in pairs],
        "pairs": pairs,
        "source_cleanup": {
            "output_input_indices": [[i] for i in range(len(raw_paras))],
        },
        "segmentation": {
            "mode": "paragraph",
            "paragraph_ranges": [
                {"paragraph_index": i, "start": i, "end": i + 1}
                for i in range(len(cleaned_paras))
            ],
        },
        "segment_structure_history": [],
        "excluded_segments": [],
    }
    core.assign_missing_segment_uids(state, job_id=job_id)
    (job_d / "state.json").write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    return state


# ---------------------------------------------------------------------------
# Acceptance criteria 1: Resolution across PDF, DOCX, and EPUB
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("fmt,builder", [
    ("pdf", lambda: ("sample.pdf", _pdf(3, toc=[[1, "Chapter 1", 1], [1, "Chapter 2", 2], [1, "Chapter 3", 3]]))),
    ("docx", lambda: ("sample.docx", _docx(("h", "Chapter 1"), ("p", "Text 1"), ("h", "Chapter 2"), ("p", "Text 2")))),
    ("epub", lambda: ("sample.epub", _epub_fixture())),
])
def test_outline_resolution_across_pdf_docx_epub(tmp_path, monkeypatch, fmt, builder):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    filename, file_bytes = builder()
    job_id = f"job-{fmt}-test"

    outline = parse_source_outline(filename, file_bytes)
    assert outline.get("nodes"), f"{fmt} should produce outline nodes"

    raw_paras = list(outline.get("paragraphs") or [])
    state = _setup_task(tmp_path, job_id, filename, file_bytes, raw_paras=raw_paras)

    sidecar = get_or_create_source_outline(job_id, state=state, source_bytes=file_bytes, filename=filename)
    mapping = sidecar["mapping"]

    # Every outline node resolves to a first segment or is reported as partial/unresolved with a reason
    # No node is silently dropped
    node_ids = {node["id"] for node in outline["nodes"]}
    mapped_ids = set(mapping["nodes"].keys())
    assert node_ids == mapped_ids, "Every outline node must exist in mapping['nodes']"

    for node_id, entry in mapping["nodes"].items():
        assert entry["state"] in {"resolved", "partial", "unresolved"}
        if entry["state"] in {"resolved", "partial"}:
            assert isinstance(entry["first_segment_index"], int)
            assert 0 <= entry["first_segment_index"] < len(state["pairs"])
            assert entry["segment_uids"]
        else:
            assert entry["reason"] in {"no_paragraph_match", "cleanup_lineage_missing", "structure_edited", "node_excluded"}


# ---------------------------------------------------------------------------
# Acceptance criteria 2: Resolution performance (no re-extraction, no model calls)
# ---------------------------------------------------------------------------

def test_resolution_performs_no_reextraction_and_no_model_calls(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "job-perf-2000"

    # Create 2000 segments across 20 chapters
    raw_paras = [f"Paragraph text {i} with unique content." for i in range(2000)]
    outline_nodes = []
    for ch in range(20):
        start = ch * 100
        outline_nodes.append({
            "id": f"n{ch:04d}",
            "parent_id": "",
            "kind": "chapter",
            "title": f"Chapter {ch + 1}",
            "level": 1,
            "para_start": start,
            "para_count": 100,
            "paragraphs": raw_paras[start:start + 100],
        })
    outline = {
        "version": OUTLINE_PARSER_VERSION,
        "total_paragraphs": 2000,
        "nodes": outline_nodes,
        "paragraphs": raw_paras,
    }
    outline["fingerprint"] = outline_fingerprint(outline)

    state = {
        "job_id": job_id,
        "filename": "perf.txt",
        "p1_done": True,
        "source_paras": raw_paras,
        "source_paragraphs": raw_paras,
        "paras": raw_paras,
        "pairs": [{"source": p, "target": f"T: {p}"} for p in raw_paras],
        "source_cleanup": {
            "output_input_indices": [[i] for i in range(2000)],
        },
        "segmentation": {
            "paragraph_ranges": [{"paragraph_index": i, "start": i, "end": i + 1} for i in range(2000)],
        },
        "segment_structure_history": [],
        "excluded_segments": [],
    }
    core.assign_missing_segment_uids(state, job_id=job_id)

    # Instrument: ensure parse_source_outline and any LLM calls are 0
    with patch("folith.source_outline.parse_source_outline") as mock_parse, \
         patch("core._LLM_CTX") as mock_llm:
        t0 = time.perf_counter()
        mapping = resolve_outline_mapping(outline, state)
        elapsed = time.perf_counter() - t0

        assert mock_parse.call_count == 0, "Resolution must not perform source re-extraction"
        assert mock_llm.call_count == 0, "Resolution must not perform model calls"

    # Must be deterministic and sub-second
    assert elapsed < 0.5, f"2000-segment resolution took {elapsed:.4f}s, expected < 0.5s"
    assert len(mapping["nodes"]) == 20
    assert mapping["nodes"]["n0000"]["first_segment_index"] == 0
    assert mapping["nodes"]["n0019"]["first_segment_index"] == 1900


# ---------------------------------------------------------------------------
# Acceptance criteria 3: Split or merge refuses stale positional path
# ---------------------------------------------------------------------------

def test_split_merge_refuses_positional_path_reports_structure_edited(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "job-split-refusal"
    raw_paras = ["Alpha 1", "Alpha 2", "Beta 1", "Beta 2"]
    file_bytes = b"Alpha 1\n\nAlpha 2\n\nBeta 1\n\nBeta 2\n"

    state = _setup_task(tmp_path, job_id, "test.txt", file_bytes, raw_paras=raw_paras)
    outline = {
        "nodes": [
            {"id": "n0000", "title": "Chap 1", "para_start": 0, "para_count": 2, "paragraphs": ["Alpha 1", "Alpha 2"]},
            {"id": "n0001", "title": "Chap 2", "para_start": 2, "para_count": 2, "paragraphs": ["Beta 1", "Beta 2"]},
        ],
        "paragraphs": raw_paras,
    }
    outline["fingerprint"] = outline_fingerprint(outline)

    # Initial mapping with unchanged structure
    sidecar = {
        "version": SOURCE_OUTLINE_VERSION,
        "parser_version": OUTLINE_PARSER_VERSION,
        "source_sha256": hashlib.sha256(file_bytes).hexdigest(),
        "fingerprint": outline["fingerprint"],
        "outline": outline,
        "mapping": resolve_outline_mapping(outline, state),
    }
    save_source_outline(job_id, sidecar)
    assert sidecar["mapping"]["nodes"]["n0000"]["first_segment_index"] == 0
    assert sidecar["mapping"]["nodes"]["n0001"]["first_segment_index"] == 2

    # Perform a split on segment 0 (inside node n0000)
    # This mints new UIDs and records a history entry
    core.mutate_translation_segments(
        job_id,
        0,
        "split",
        source_a="Alpha",
        source_b="1",
        actor="test_user",
    )
    edited_state = core.load_job_state(job_id)
    assert len(edited_state["segment_structure_history"]) == 1
    assert structure_revision(edited_state) == 1

    # Reload outline: structure_revision changed!
    reloaded = get_or_create_source_outline(job_id, state=edited_state, source_bytes=file_bytes)
    mapping = reloaded["mapping"]

    # Node n0000 was split: original uid seg-job-split-refusal-0000 was destroyed!
    # Resolver MUST refuse the stale positional path and report structure_edited
    n0 = mapping["nodes"]["n0000"]
    assert n0["state"] in {"partial", "unresolved"}
    assert n0["reason"] == "structure_edited"

    # Node n0001 (Chap 2): UIDs were intact (Beta 1, Beta 2), so it resolves to its new shifted position (index 3)
    n1 = mapping["nodes"]["n0001"]
    assert n1["state"] == "resolved"
    assert n1["first_segment_index"] == 3, "Node n0001 must resolve by UID to shifted index 3, not stale index 2"


# ---------------------------------------------------------------------------
# Acceptance criteria 4: Exclusion shift followed by resolved nodes
# ---------------------------------------------------------------------------

def test_exclusion_shift_followed_by_resolved_nodes(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "job-exclude-shift"
    raw_paras = ["Intro 1", "Body 1", "Body 2"]
    file_bytes = b"# Intro\n\nIntro 1\n\n# Body\n\nBody 1\n\nBody 2\n"

    state = _setup_task(tmp_path, job_id, "test.md", file_bytes, raw_paras=raw_paras)
    outline = parse_source_outline("test.md", file_bytes)

    sidecar = get_or_create_source_outline(job_id, state=state, source_bytes=file_bytes)
    assert sidecar["mapping"]["nodes"]["n0000"]["first_segment_index"] == 0
    assert sidecar["mapping"]["nodes"]["n0001"]["first_segment_index"] == 1

    # Exclude segment 0 (Intro 1)
    core.mutate_translation_segments(
        job_id,
        0,
        "exclude",
        exclude_reason="Not for translation",
        actor="test_user",
    )
    excluded_state = core.load_job_state(job_id)
    assert len(excluded_state["excluded_segments"]) == 1

    reloaded = get_or_create_source_outline(job_id, state=excluded_state, source_bytes=file_bytes)
    mapping = reloaded["mapping"]

    # Node 0 was excluded: reports node_excluded
    assert mapping["nodes"]["n0000"]["state"] == "unresolved"
    assert mapping["nodes"]["n0000"]["reason"] == "node_excluded"
    assert mapping["nodes"]["n0000"]["first_segment_index"] is None

    # Node 1: points at shifted segment 0 (which is now Body 1)
    assert mapping["nodes"]["n0001"]["state"] == "resolved"
    assert mapping["nodes"]["n0001"]["first_segment_index"] == 0
    first_uid = mapping["nodes"]["n0001"]["segment_uids"][0]
    assert excluded_state["pairs"][0]["segment_uid"] == first_uid


# ---------------------------------------------------------------------------
# Acceptance criteria 5 & 10: Re-importing identical file & residual documentation
# ---------------------------------------------------------------------------

def test_reimport_identical_file_points_to_same_text(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    file_bytes = b"Chapter 1\n\nContent paragraph.\n"
    raw_paras = ["Chapter 1", "Content paragraph."]

    # First task
    job_1 = "job-import-1"
    state_1 = _setup_task(tmp_path, job_1, "test.txt", file_bytes, raw_paras=raw_paras)
    sidecar_1 = get_or_create_source_outline(job_1, state=state_1, source_bytes=file_bytes)

    # Second task (re-imported identical file under another job ID)
    job_2 = "job-import-2"
    state_2 = _setup_task(tmp_path, job_2, "test.txt", file_bytes, raw_paras=raw_paras)
    sidecar_2 = get_or_create_source_outline(job_2, state=state_2, source_bytes=file_bytes)

    # Both resolve to the same segment index and same text
    n1_entry = sidecar_1["mapping"]["nodes"]["n0000"]
    n2_entry = sidecar_2["mapping"]["nodes"]["n0000"]
    assert n1_entry["first_segment_index"] == n2_entry["first_segment_index"] == 0

    seg1_text = state_1["pairs"][n1_entry["first_segment_index"]]["source"]
    seg2_text = state_2["pairs"][n2_entry["first_segment_index"]]["source"]
    assert seg1_text == seg2_text == "Chapter 1"

    # Residual difference from positional/job-scoped identity:
    # Segment UIDs are job-scoped instance identifiers (seg-<job_id>-<index>),
    # so they differ between jobs, while the navigation text & index match exactly.
    assert n1_entry["segment_uids"][0] != n2_entry["segment_uids"][0]
    assert n1_entry["segment_uids"][0] == f"seg-{job_1}-0000"
    assert n2_entry["segment_uids"][0] == f"seg-{job_2}-0000"


# ---------------------------------------------------------------------------
# Acceptance criteria 6: Reopening task reuses persisted outline without reparsing
# ---------------------------------------------------------------------------

def test_reopen_task_reuses_persisted_outline_without_reparsing(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "job-reopen-cache"
    file_bytes = b"# Chapter\n\nBody text.\n"
    raw_paras = ["Chapter", "Body text."]

    state = _setup_task(tmp_path, job_id, "test.md", file_bytes, raw_paras=raw_paras)

    with patch("folith.outline_anchor.parse_source_outline", wraps=parse_source_outline) as spy_parse:
        # First call: parses and persists sidecar
        first = get_or_create_source_outline(job_id, state=state, source_bytes=file_bytes)
        assert spy_parse.call_count == 1
        assert (core.job_dir(job_id) / SOURCE_OUTLINE_FILE).exists()

        # Reopen: should reuse persisted outline without invoking parser
        second = get_or_create_source_outline(job_id, state=state, source_bytes=file_bytes)
        assert spy_parse.call_count == 1, "Parser must NOT be invoked again for unchanged source on reopen"
        assert first["fingerprint"] == second["fingerprint"]


# ---------------------------------------------------------------------------
# Acceptance criteria 7: Changed source or bumped parser version invalidates sidecar
# ---------------------------------------------------------------------------

def test_source_change_invalidates_persisted_outline(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "job-source-invalidate"
    file_bytes = b"# Old Title\n\nOld content.\n"
    state = _setup_task(tmp_path, job_id, "test.md", file_bytes, raw_paras=["Old Title", "Old content."])

    # Initial sidecar
    get_or_create_source_outline(job_id, state=state, source_bytes=file_bytes)
    old_sidecar = load_source_outline(job_id)
    assert old_sidecar["source_sha256"] == hashlib.sha256(file_bytes).hexdigest()

    # Source bytes change
    new_bytes = b"# New Title\n\nNew content.\n"
    (core.job_dir(job_id) / "source.bin").write_bytes(new_bytes)

    # Re-call: must re-parse and overwrite sidecar
    new_sidecar = get_or_create_source_outline(job_id, state=state, source_bytes=new_bytes)
    assert new_sidecar["source_sha256"] == hashlib.sha256(new_bytes).hexdigest()
    assert new_sidecar["outline"]["nodes"][0]["title"] == "New Title"


def test_bumped_parser_version_invalidates_persisted_outline(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "job-version-invalidate"
    file_bytes = b"# Chap\n\nContent.\n"
    state = _setup_task(tmp_path, job_id, "test.md", file_bytes, raw_paras=["Chap", "Content."])

    get_or_create_source_outline(job_id, state=state, source_bytes=file_bytes)
    sidecar = load_source_outline(job_id)
    # Simulate an older parser version on disk
    sidecar["parser_version"] = 0
    save_source_outline(job_id, sidecar)

    # Next load detects parser_version mismatch and re-parses
    with patch("folith.outline_anchor.parse_source_outline", wraps=parse_source_outline) as spy_parse:
        reloaded = get_or_create_source_outline(job_id, state=state, source_bytes=file_bytes)
        assert spy_parse.call_count == 1
        assert reloaded["parser_version"] == OUTLINE_PARSER_VERSION


# ---------------------------------------------------------------------------
# Acceptance criteria 8: Atomic write, schema shape, no new outline key in state
# ---------------------------------------------------------------------------

def test_atomic_write_schema_shape_and_no_outline_nodes_in_state(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "job-shape-check"
    file_bytes = b"# Chap 1\n\nPara 1\n"
    state = _setup_task(tmp_path, job_id, "test.md", file_bytes, raw_paras=["Chap 1", "Para 1"])

    sidecar = get_or_create_source_outline(job_id, state=state, source_bytes=file_bytes)
    path = source_outline_path(job_id)
    assert path.is_file()

    # Verify sidecar shape
    assert set(sidecar.keys()) == {
        "version", "parser_version", "source_sha256", "fingerprint", "outline", "mapping"
    }
    assert sidecar["version"] == SOURCE_OUTLINE_VERSION
    assert sidecar["parser_version"] == OUTLINE_PARSER_VERSION

    mapping = sidecar["mapping"]
    assert set(mapping.keys()) == {"fingerprint", "structure_revision", "nodes"}
    assert mapping["structure_revision"] == 0
    for node_id, entry in mapping["nodes"].items():
        assert "state" in entry
        assert "reason" in entry
        assert "segment_uids" in entry
        assert "first_segment_index" in entry

    # Verify state.json does NOT contain outline nodes or paragraph texts
    persisted_state = json.loads((core.job_dir(job_id) / "state.json").read_text(encoding="utf-8"))
    assert "source_outline" not in persisted_state
    assert "outline" not in persisted_state
    for key, val in persisted_state.items():
        if isinstance(val, dict):
            assert "nodes" not in val, f"State key '{key}' must not carry outline nodes"


# ---------------------------------------------------------------------------
# Acceptance criteria 9: _reset_source_dependent_state deletes sidecar
# ---------------------------------------------------------------------------

def test_reset_source_dependent_state_deletes_sidecar_on_rescan_and_reimport(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "job-reset-deletion"
    file_bytes = b"# Chap\n\nText\n"
    state = _setup_task(tmp_path, job_id, "test.md", file_bytes, raw_paras=["Chap", "Text"])

    get_or_create_source_outline(job_id, state=state, source_bytes=file_bytes)
    assert source_outline_path(job_id).is_file()

    # Calling _reset_source_dependent_state must delete the sidecar
    core._reset_source_dependent_state(state, job_id=job_id)
    assert not source_outline_path(job_id).exists(), "Sidecar must be deleted on source reset"


# ---------------------------------------------------------------------------
# Acceptance criteria 11: Legacy task without UIDs uses positional fallback
# ---------------------------------------------------------------------------

def test_legacy_task_without_uids_positional_fallback(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "job-legacy-nouids"
    raw_paras = ["Header", "Paragraph"]

    # State with no segment_uid fields on pairs
    state = {
        "job_id": job_id,
        "filename": "legacy.txt",
        "p1_done": True,
        "source_paras": raw_paras,
        "source_paragraphs": raw_paras,
        "paras": raw_paras,
        "pairs": [{"source": "Header"}, {"source": "Paragraph"}],
        "source_cleanup": {"output_input_indices": [[0], [1]]},
        "segmentation": {"paragraph_ranges": [{"paragraph_index": 0, "start": 0, "end": 1},
                                              {"paragraph_index": 1, "start": 1, "end": 2}]},
        "segment_structure_history": [],
        "excluded_segments": [],
    }
    outline = {
        "nodes": [
            {"id": "n0000", "title": "Header", "para_start": 0, "para_count": 1, "paragraphs": ["Header"]},
        ],
        "paragraphs": raw_paras,
    }

    mapping = resolve_outline_mapping(outline, state)
    entry = mapping["nodes"]["n0000"]
    assert entry["state"] == "resolved"
    assert entry["reason"] == "positional_fallback"
    assert entry["first_segment_index"] == 0

    # If structure changed without UIDs, positional fallback is strictly refused
    state["segment_structure_history"] = [{"operation": "manual_edit"}]
    mapping_after_edit = resolve_outline_mapping(outline, state)
    entry_after = mapping_after_edit["nodes"]["n0000"]
    assert entry_after["state"] == "unresolved"
    assert entry_after["reason"] == "structure_edited"
    assert entry_after["first_segment_index"] is None


# ---------------------------------------------------------------------------
# Acceptance criteria 1: Partial resolution honesty
# ---------------------------------------------------------------------------

def test_partial_resolution_honest_span_and_unplaced_count(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "job-partial-honesty"

    # Raw paragraphs in state: only 2 of the 3 paragraphs from the node exist
    raw_paras = ["Title", "Para 1 of node", "Para 3 of node", "Other"]
    state = _setup_task(tmp_path, job_id, "test.txt", b"dummy", raw_paras=raw_paras)

    # Outline node has 3 paragraphs; Para 2 was unplaced (e.g. mixed PDF text layer vs VLM)
    node = {
        "id": "n0000",
        "title": "Node with unplaced paragraph",
        "para_start": 1,
        "para_count": 3,
        "paragraphs": ["Para 1 of node", "Unrecognized Para 2", "Para 3 of node"],
    }
    outline = {
        "nodes": [node],
        "paragraphs": ["Title", "Para 1 of node", "Unrecognized Para 2", "Para 3 of node"],
    }

    mapping = resolve_outline_mapping(outline, state)
    entry = mapping["nodes"]["n0000"]

    assert entry["state"] == "partial", "Partial match must NOT be rounded up to resolved"
    assert entry["matched_paragraphs"] == 2
    assert entry["unplaced_paragraphs"] == 1
    assert entry["first_segment_index"] == 1
    assert entry["matched_span"] == [1, 3]  # spans segment 1 to 2
    assert len(entry["segment_uids"]) == 2


# ---------------------------------------------------------------------------
# Acceptance criteria 12: Architecture boundary documentation
# ---------------------------------------------------------------------------

def test_architecture_boundary_documented_as_navigation_only():
    import folith.outline_anchor as anchor
    doc = anchor.__doc__ or ""
    assert "Navigation identity only" in doc
    assert "docs/architecture-boundaries.md" in doc
    assert "Source Update" in doc or "source-update" in doc
