"""翻译范围回归（issue #28）：按页/章/节的译前批量排除与恢复。

范围选择必须是**既有排除账本的批量入口**，不另设"排除真相"，这一批断言守护
的就是这条契约：

- 范围树是从文档结构（页血缘 + 标题启发式）算出的只读投影，映射到当前工作
  列表的连续索引区间；排除之后树跟着收缩，恢复之后重新算出归属；
- 批量排除与逐段排除同语义：移出 pairs/paras、进 ``excluded_segments`` 留档
  （含原文锚点与范围标签），因此不进翻译、不计入待完成、交付时如实登记；
- 批量恢复按原始位置精确还原（锚点插回），译前与译后两条路径都要成立；
- 交付清单把「用户跳过 / 系统判定 / 解析失败」分开陈述——跳过永远不计为
  已翻译，解析失败不是排除。
"""
import json

import core
import pytest


def _sources():
    return ["第一章 缘起", "Alpha body.", "第二章 争论", "Beta body."]


def _state(*, translated=False):
    sources = _sources()
    state = core.new_job_state("translation-scope.pdf")
    state.update(
        p1_done=True,
        paras=list(sources),
        review_stats={"reviewed_segments": 0},
        delivery_status="draft",
    )
    if translated:
        state.update(
            p2_done=True,
            pairs=[{"source": source, "target": f"译文{index}",
                    "initial_target": f"译文{index}", "reviewed": True,
                    "from_tm": False, "glossary_entry_ids": []}
                   for index, source in enumerate(sources)],
        )
    return state


def _persist(tmp_path, job_id, state):
    old_output = core.OUTPUT_DIR
    core.OUTPUT_DIR = tmp_path
    core.save_source(job_id, b"translation scope source")
    core.save_job_state(job_id, state)
    return old_output


def _chapter_nodes(tree):
    return next((group for group in tree["groups"] if group["kind"] == "chapter"), None)


# ---------------------------------------------------------------- 范围树

def test_scope_tree_maps_chapters_to_contiguous_ranges(tmp_path):
    job_id = "scope-tree-chapters"
    old_output = _persist(tmp_path, job_id, _state())
    try:
        tree = core.translation_scope_tree(core.load_job_state(job_id))
        assert tree["available"] is True
        assert tree["total"] == {"segments": 4,
                                 "chars": sum(len(item) for item in _sources())}
        chapter = _chapter_nodes(tree)
        assert [(node["start"], node["end"]) for node in chapter["nodes"]] == \
            [(0, 2), (2, 4)]
        assert chapter["nodes"][0]["label"] == "第一章 缘起"
        assert chapter["nodes"][1]["label"] == "第二章 争论"
        assert chapter["nodes"][1]["key"] == "chapter:0001"
        # 没有页血缘（segment_source_spans 缺失）时不提供按页分组。
        assert not any(group["kind"] == "page" for group in tree["groups"])
    finally:
        core.OUTPUT_DIR = old_output


def test_scope_tree_pages_come_from_segment_source_spans(tmp_path):
    job_id = "scope-tree-pages"
    state = _state()
    state["segment_source_spans"] = [
        [{"page_id": "p1"}], [{"page_id": "p1"}],
        [{"page_id": "p2"}], [{"page_id": "p2"}],
    ]
    old_output = _persist(tmp_path, job_id, state)
    try:
        tree = core.translation_scope_tree(core.load_job_state(job_id))
        pages = next(group for group in tree["groups"] if group["kind"] == "page")
        assert [(node["start"], node["end"]) for node in pages["nodes"]] == \
            [(0, 2), (2, 4)]
        assert [node["label"] for node in pages["nodes"]] == ["第 1 页", "第 2 页"]
    finally:
        core.OUTPUT_DIR = old_output


def test_scope_tree_shrinks_after_exclusion_and_reflects_excluded(tmp_path):
    job_id = "scope-tree-shrink"
    old_output = _persist(tmp_path, job_id, _state())
    try:
        core.exclude_translation_scope(job_id, ["chapter:0001"])
        tree = core.translation_scope_tree(core.load_job_state(job_id))
        assert tree["total"]["segments"] == 2
        chapter = _chapter_nodes(tree)
        assert [(node["start"], node["end"]) for node in chapter["nodes"]] == [(0, 2)]
        assert tree["excluded"]["count"] == 2
    finally:
        core.OUTPUT_DIR = old_output


def test_scope_tree_empty_state_is_unavailable():
    state = core.new_job_state("scope-empty.pdf")
    tree = core.translation_scope_tree(state)
    assert tree["available"] is False
    assert tree["groups"] == []


# ---------------------------------------------------------------- 批量排除

def test_scope_exclude_pre_translation_moves_segments_into_ledger(tmp_path):
    job_id = "scope-exclude-paras"
    old_output = _persist(tmp_path, job_id, _state())
    try:
        result = core.exclude_translation_scope(job_id, ["chapter:0001"])
        assert result["excluded"] == 2
        assert result["remaining"] == 2
        state = core.load_job_state(job_id)
        assert state["paras"] == ["第一章 缘起", "Alpha body."]
        # 与逐段排除同一条账本：原文留档、原因带范围标签、锚点齐全。
        records = core.excluded_segment_records(state)
        assert [record["source"] for record in records] == \
            ["第二章 争论", "Beta body."]
        assert {record["reason"] for record in records} == \
            {"范围排除：第二章 争论"}
        assert all(record["restorable"] for record in records)
        # 排除不计为已翻译：没有 pairs、没有 reviewed 标记。
        assert state.get("pairs") == []
        assert not any(record.get("target") for record in records)
        # 历史留档但不接入撤销（批量恢复走锚点，不需要内容快照）。
        assert state["segment_structure_history"][-1]["operation"] == "scope_exclude"
        assert "restore_pairs" not in state["segment_structure_history"][-1]
        assert not core.can_undo_translation_segment_structure(state)["available"]
    finally:
        core.OUTPUT_DIR = old_output


def test_scope_exclude_translated_state_invalidates_reviews(tmp_path):
    job_id = "scope-exclude-pairs"
    old_output = _persist(tmp_path, job_id, _state(translated=True))
    try:
        result = core.exclude_translation_scope(job_id, ["chapter:0001"])
        state = core.load_job_state(job_id)
        assert [pair["source"] for pair in state["pairs"]] == \
            ["第一章 缘起", "Alpha body."]
        # 排除时的译文完整留档，恢复是精确还原。
        records = core.excluded_segment_records(state)
        assert [record["target"] for record in records] == ["译文2", "译文3"]
        # 交付范围变化的既有契约：真值变更留痕（action=scope_exclude）。
        assert state["translation_truth"]["last_change"]["action"] == "scope_exclude"
        assert result["excluded"] == 2
    finally:
        core.OUTPUT_DIR = old_output


def test_scope_exclude_rejects_unknown_keys_and_running_tasks(tmp_path):
    job_id = "scope-exclude-guard"
    old_output = _persist(tmp_path, job_id, _state())
    try:
        with pytest.raises(ValueError, match="请选择要排除的范围"):
            core.exclude_translation_scope(job_id, [])
        with pytest.raises(ValueError, match="范围与当前段落不匹配"):
            core.exclude_translation_scope(job_id, ["chapter:0009"])
        # 已排除一段后再用旧树键恢复，键已不存在，必须拒绝而不是错位排除。
        core.exclude_translation_scope(job_id, ["chapter:0001"])
        with pytest.raises(ValueError, match="范围与当前段落不匹配"):
            core.exclude_translation_scope(job_id, ["chapter:0001"])
    finally:
        core.OUTPUT_DIR = old_output


# ---------------------------------------------------------------- 批量恢复

@pytest.mark.parametrize("translated", [False, True])
def test_scope_restore_puts_segments_back_in_original_order(
        tmp_path, translated):
    job_id = "scope-restore-order"
    old_output = _persist(tmp_path, job_id, _state(translated=translated))
    try:
        # 范围键跟着工作列表重算：排除一章之后剩下的章拿的是**新的**键
        #（对话框每次打开都重新读树，这一点由 adapter.translation_scope 保证）。
        state = core.load_job_state(job_id)
        first = _chapter_nodes(core.translation_scope_tree(state))["nodes"][0]["key"]
        core.exclude_translation_scope(job_id, [first])
        state = core.load_job_state(job_id)
        second = _chapter_nodes(core.translation_scope_tree(state))["nodes"][-1]["key"]
        core.exclude_translation_scope(job_id, [second])
        state = core.load_job_state(job_id)
        assert state["paras"] == []
        excluded_ids = [record["excluded_id"]
                        for record in core.excluded_segment_records(state)]
        # 打乱顺序恢复（后排除的先恢复），位置仍由锚点决定。
        result = core.restore_translation_scope(job_id, list(reversed(excluded_ids)))
        assert result["restored"] == 4
        state = core.load_job_state(job_id)
        assert state["paras"] == _sources()
        if translated:
            assert [pair["source"] for pair in state["pairs"]] == _sources()
            assert [pair["target"] for pair in state["pairs"]] == \
                ["译文0", "译文1", "译文2", "译文3"]
            # 恢复的内容原样回来，但旧的"已审校"不跟着点亮。
            assert all(not pair["reviewed"] for pair in state["pairs"])
        assert core.excluded_segment_records(state) == []
        assert state["translation_truth"]["last_change"]["action"] == "scope_include"
    finally:
        core.OUTPUT_DIR = old_output


def test_scope_restore_rejects_unknown_ids(tmp_path):
    job_id = "scope-restore-guard"
    old_output = _persist(tmp_path, job_id, _state())
    try:
        with pytest.raises(ValueError, match="请选择要恢复的已排除段落"):
            core.restore_translation_scope(job_id, [])
        with pytest.raises(ValueError, match="找不到要恢复的已排除段落"):
            core.restore_translation_scope(job_id, ["x999999"])
    finally:
        core.OUTPUT_DIR = old_output


# ---------------------------------------------------------------- 交付清单

def test_manifest_separates_user_skips_system_and_parse_failures(tmp_path):
    job_id = "scope-manifest-origins"
    old_output = _persist(tmp_path, job_id, _state())
    try:
        core.exclude_translation_scope(job_id, ["chapter:0001"])
        state = core.load_job_state(job_id)
        # 系统判定的结构性排除走同一条账本，excluded_by 带 system: 前缀。
        state["excluded_segments"] = list(state["excluded_segments"]) + [{
            "excluded_id": "x900000", "index": 0,
            "source": "Running head", "target": "",
            "reason": "系统判定为无需翻译的结构性片段（page_header）",
            "excluded_at": "2026-10-25T00:00:00+00:00",
            "excluded_by": "system:source_quality",
        }]
        state["extraction_report"] = {
            "status": "partial",
            "unsupported": [{"kind": "figure", "count": 2}],
        }
        state["warnings"] = ["OCR 部分页面失败：p.12"]
        core.save_job_state(job_id, state)

        manifest = core.excluded_segments_manifest(state)
        assert manifest["count"] == 3
        assert manifest["origins"] == {"user": 2, "system": 1}
        assert {record["origin"] for record in manifest["records"]} == \
            {"user", "system"}
        # 解析失败单列，与跳过互不归属。
        assert manifest["parse_failures"]["extraction_status"] == "partial"
        assert manifest["parse_failures"]["unsupported_count"] == 2
        assert manifest["parse_failures"]["warnings"] == ["OCR 部分页面失败：p.12"]
        assert "跳过永远不计为已翻译" in manifest["rule"]
        assert "不计为已翻译" in manifest["skip_semantics"]
        assert "解析失败" in manifest["skip_semantics"]

        text = core._excluded_segments_md(manifest)
        assert "用户跳过 2" in text and "系统判定 1" in text
        assert "| 来源 |" in text and "用户跳过" in text and "系统判定" in text
        assert "## 解析失败（不是排除）" in text
        assert "figure：2 处" in text
    finally:
        core.OUTPUT_DIR = old_output


def test_manifest_json_stays_serializable_after_scope_exclude(tmp_path):
    job_id = "scope-manifest-bundle"
    old_output = _persist(tmp_path, job_id, _state())
    try:
        core.exclude_translation_scope(job_id, ["chapter:0000"])
        state = core.load_job_state(job_id)
        manifest = core.excluded_segments_manifest(state)
        payload = json.loads(json.dumps(manifest, ensure_ascii=False))
        assert payload["count"] == 2
        assert payload["records"][0]["reason"] == "范围排除：第一章 缘起"
    finally:
        core.OUTPUT_DIR = old_output


# ---------------------------------------------------------------- 译前闭环（#28）

def _docx_outline_state(job_id):
    """新导入、还没开始的两章 DOCX 任务：没有 paras/pairs，只有源文件。

    大纲形状与 ``folith.source_outline.parse_source_outline`` 的 DOCX 分支
    一致：段落列表 + 每章一个节点（para_start / para_count / 指纹）。
    """
    state = core.new_job_state("two-chapters.docx")
    state["filename"] = "two-chapters.docx"
    return state


def _two_chapter_outline():
    paragraphs = [
        "第一章 缘起", "Alpha body one.", "Alpha body two.",
        "第二章 争论", "Beta body one.", "Beta body two.",
    ]
    outline = {
        "format": "DOCX",
        "nodes": [
            {"id": "n0000", "kind": "chapter", "title": "第一章 缘起",
             "level": 1, "para_start": 0, "para_count": 3,
             "paragraphs": paragraphs[:3], "chars": 30,
             "page_start": None, "page_end": None},
            {"id": "n0001", "kind": "chapter", "title": "第二章 争论",
             "level": 1, "para_start": 3, "para_count": 3,
             "paragraphs": paragraphs[3:], "chars": 30,
             "page_start": None, "page_end": None},
        ],
        "paragraphs": paragraphs,
        "total_paragraphs": len(paragraphs),
        "total_chars": sum(len(item) for item in paragraphs),
        "pages": 0, "scanned": False, "partial": False,
    }
    from folith.source_outline import outline_fingerprint
    outline["fingerprint"] = outline_fingerprint(outline)
    return outline


def test_pre_start_scope_tree_comes_from_source_outline(tmp_path):
    """开始前：范围树从原文大纲读出两章，不再返回 available=false。"""
    job_id = "scope-prestart-tree"
    state = _docx_outline_state(job_id)
    state["source_outline"] = _two_chapter_outline()
    old_output = _persist(tmp_path, job_id, state)
    try:
        tree = core.translation_scope_tree(core.load_job_state(job_id))
        assert tree["available"] is True
        assert tree["pre_start"] is True
        chapter = _chapter_nodes(tree)
        assert [(node["key"], node["start"], node["end"]) for node in chapter["nodes"]] == [
            ("outline:n0000", 0, 3), ("outline:n0001", 3, 6)]
        assert tree["total"] == {"segments": 6,
                                 "chars": sum(len(p) for p in [
                                     "第一章 缘起", "Alpha body one.",
                                     "Alpha body two.", "第二章 争论",
                                     "Beta body one.", "Beta body two."])}
    finally:
        core.OUTPUT_DIR = old_output


def test_renamed_outline_keeps_pre_start_exclusion_anchors(tmp_path):
    from folith import outline_corrections

    job_id = "scope-corrected-name"
    state = _docx_outline_state(job_id)
    outline = state["source_outline"] = _two_chapter_outline()
    old_output = _persist(tmp_path, job_id, state)
    try:
        outline_corrections.save(job_id, outline, "n0001", title="新标题")
        core.exclude_translation_scope(job_id, ["outline:n0001"])
        state = core.load_job_state(job_id)
        entry = state["scope_exclusion_plan"]["entries"][0]
        assert entry["label"] == "新标题"
        assert entry["anchor_first"] == "第二章 争论"
        state["paras"] = list(outline["paragraphs"])
        state["source_paras"] = list(state["paras"])
        state["source_paragraphs"] = list(state["paras"])
        state["source_cleanup"] = {"output_input_indices": [[i] for i in range(6)]}
        state["segmentation"] = {"paragraph_ranges": [
            {"paragraph_index": i, "start": i, "end": i + 1} for i in range(6)]}
        result = core.apply_scope_exclusion_plan(job_id, state)
        assert not result["unapplied"]
        assert state["paras"] == outline["paragraphs"][:3]
    finally:
        core.OUTPUT_DIR = old_output


def test_pre_start_exclusion_refuses_unreadable_corrections(tmp_path, monkeypatch):
    from folith import outline_corrections
    from folith.qt import adapter
    job_id = "scope-corrupt-corrections"
    state = _docx_outline_state(job_id)
    state["source_outline"] = _two_chapter_outline()
    old_output = _persist(tmp_path, job_id, state)
    try:
        path = core.job_dir(job_id) / outline_corrections.FILENAME
        damaged = b'{"entries": [broken'
        path.write_bytes(damaged)
        monkeypatch.setattr(core, "ensure_source_outline", lambda *_: state)
        assert adapter.translation_scope(job_id, state)["available"] is False
        with pytest.raises(ValueError, match="无法读取结构修正文件"):
            core.exclude_translation_scope(job_id, ["outline:n0001"])
        assert not core.load_job_state(job_id).get("scope_exclusion_plan")
        assert path.read_bytes() == damaged
    finally:
        core.OUTPUT_DIR = old_output


def test_pre_start_exclusion_is_a_plan_and_applies_when_units_are_built(tmp_path):
    """操作闭环：译前排除记成计划 → 任务开始建好翻译单元 → 落进排除账本。

    审校探针的场景：新导入两章 DOCX，准备页选范围排除第二章，之后才开始翻译。
    落账后第二章不进入翻译（不会为它发起任何付费模型调用），账本、交付与
    恢复入口都按既有契约成立。
    """
    job_id = "scope-prestart-plan"
    state = _docx_outline_state(job_id)
    state["source_outline"] = _two_chapter_outline()
    old_output = _persist(tmp_path, job_id, state)
    try:
        # 1) 译前排除第二章：还不触碰排除账本，只记计划。
        result = core.exclude_translation_scope(job_id, ["outline:n0001"])
        assert result["planned"] == 1 and result["excluded"] == 0
        assert result["remaining"] == 3
        state = core.load_job_state(job_id)
        assert state.get("excluded_segments") in (None, []), \
            "译前排除不落账：账本只描述已移出的翻译单元"
        plan = state["scope_exclusion_plan"]
        assert [entry["label"] for entry in plan["entries"]] == ["第二章 争论"]
        assert plan["entries"][0]["para_start"] == 3
        assert plan["entries"][0]["anchor_first"] == "第二章 争论"
        tree = core.translation_scope_tree(state)
        assert tree["excluded"]["count"] == 3, "总量按原文段落，记录按计划范围"

        # 2) 任务开始：阶段一建好翻译单元（paras + paragraph_ranges + 清洗映射）。
        state = core.load_job_state(job_id)
        state["paras"] = ["第一章 缘起", "Alpha body one.", "Alpha body two.",
                          "第二章 争论", "Beta body one.", "Beta body two."]
        state["source_paras"] = list(state["paras"])
        state["source_paragraphs"] = list(state["paras"])
        state["source_cleanup"] = {
            "output_input_indices": [[i] for i in range(6)]}
        state["segmentation"] = {"paragraph_ranges": [
            {"paragraph_index": i, "start": i, "end": i + 1}
            for i in range(6)]}
        core.save_job_state(job_id, state)
        outcome = core.apply_scope_exclusion_plan(job_id, state)

        # 3) 落账：第二章移出工作列表，第一章保持不动。
        assert [item["label"] for item in outcome["applied"]] == ["第二章 争论"]
        assert outcome["unapplied"] == []
        # apply 不落盘（调用方随阶段收尾统一保存），这里保存后从盘上读回验证。
        core.save_job_state(job_id, state)
        state = core.load_job_state(job_id)
        assert state["paras"] == ["第一章 缘起", "Alpha body one.",
                                  "Alpha body two."]
        records = core.excluded_segment_records(state)
        assert [record["source"] for record in records] == \
            ["第二章 争论", "Beta body one.", "Beta body two."]
        assert {record["reason"] for record in records} == \
            {"范围排除（译前）：第二章 争论"}
        assert all(record["restorable"] for record in records)
        assert state.get("scope_exclusion_plan") is None
        assert state.get("source_outline") is None, "落账后范围入口回到工作列表口径"
        applied_audit = state["scope_exclusion_applied"]
        assert applied_audit[0]["excluded_count"] == 3

        # 4) 已落账的排除可以照常恢复（与译后范围排除同一入口）。
        core.restore_translation_scope(
            job_id, [record["excluded_id"] for record in records])
        state = core.load_job_state(job_id)
        assert len(state["paras"]) == 6
        assert state.get("excluded_segments") in (None, [])
    finally:
        core.OUTPUT_DIR = old_output


def test_pre_start_plan_survives_until_units_exist_and_restores_before_start(
        tmp_path):
    """计划撤下走同一个恢复入口；任务开始前不产生任何排除记录。"""
    job_id = "scope-prestart-unplan"
    state = _docx_outline_state(job_id)
    state["source_outline"] = _two_chapter_outline()
    old_output = _persist(tmp_path, job_id, state)
    try:
        core.exclude_translation_scope(job_id, ["outline:n0000"])
        state = core.load_job_state(job_id)
        excluded_id = state["scope_exclusion_plan"]["entries"][0]["excluded_id"]
        result = core.restore_translation_scope(job_id, [excluded_id])
        assert result["restored"] == 1 and result["remaining"] == 6
        state = core.load_job_state(job_id)
        assert state["scope_exclusion_plan"]["entries"] == []
        assert state.get("excluded_segments") in (None, [])
    finally:
        core.OUTPUT_DIR = old_output


def test_pre_start_plan_unmappable_structure_is_reported_not_guessed(tmp_path):
    """导入结构与大纲对不上（锚点验证失败）时如实报未应用，绝不猜位置。"""
    job_id = "scope-prestart-mismatch"
    state = _docx_outline_state(job_id)
    state["source_outline"] = _two_chapter_outline()
    old_output = _persist(tmp_path, job_id, state)
    try:
        core.exclude_translation_scope(job_id, ["outline:n0001"])
        state = core.load_job_state(job_id)
        # 任务开始，但提取出的结构与大纲不一致：锚点对不上。
        state["paras"] = ["完全不同的开头", "另一段正文"]
        state["source_paras"] = list(state["paras"])
        state["source_paragraphs"] = list(state["paras"])
        core.save_job_state(job_id, state)
        outcome = core.apply_scope_exclusion_plan(job_id, state)
        assert outcome["applied"] == []
        assert [item["label"] for item in outcome["unapplied"]] == ["第二章 争论"]
        core.save_job_state(job_id, state)
        state = core.load_job_state(job_id)
        assert state["paras"] == ["完全不同的开头", "另一段正文"], "不猜位置"
        assert state.get("excluded_segments") in (None, [])
        assert any("译前范围排除未能应用" in w for w in state["warnings"])
        assert state["scope_exclusion_applied"][-1]["unapplied"] is True
    finally:
        core.OUTPUT_DIR = old_output


def test_pre_start_plan_never_applies_after_pairs_exist(tmp_path):
    """位置契约已建立后计划不得再落账：事后位移会打乱已持久化的引用。"""
    job_id = "scope-prestart-late"
    state = _docx_outline_state(job_id)
    state["source_outline"] = _two_chapter_outline()
    old_output = _persist(tmp_path, job_id, state)
    try:
        core.exclude_translation_scope(job_id, ["outline:n0001"])
        state = core.load_job_state(job_id)
        state["paras"] = ["第一章 缘起", "Alpha body one.", "Alpha body two.",
                          "第二章 争论", "Beta body one.", "Beta body two."]
        state["pairs"] = [{"source": text, "target": f"译文{i}"}
                          for i, text in enumerate(state["paras"])]
        core.save_job_state(job_id, state)
        outcome = core.apply_scope_exclusion_plan(job_id, state)
        assert outcome["applied"] == []
        assert len(outcome["unapplied"]) == 1
        state = core.load_job_state(job_id)
        assert len(state["pairs"]) == 6, "pairs 一个都不能少"
        assert state.get("scope_exclusion_plan") is not None, "计划保留待查"
    finally:
        core.OUTPUT_DIR = old_output


def test_ensure_source_outline_parses_and_persists_once(tmp_path):
    """首次打开准备页时解析大纲并落盘；再开一次、或任务已开始时不再重解析。"""
    job_id = "scope-ensure-outline"
    state = _docx_outline_state(job_id)
    state["filename"] = "plan.md"  # 内容按 Markdown 给（本地文字层路径）
    old_output = _persist(tmp_path, job_id, state)
    try:
        core.save_source(job_id, (
            "# 第一章 缘起\n\nAlpha body one.\n\nAlpha body two.\n\n"
            "# 第二章 争论\n\nBeta body one.\n\nBeta body two.\n"
        ).encode("utf-8"))
        first = core.ensure_source_outline(job_id)
        assert first["source_outline"]["total_paragraphs"] == 6
        saved_at = (core.job_dir(job_id) / "state.json").stat().st_mtime_ns
        again = core.ensure_source_outline(job_id)
        assert again["source_outline"]["total_paragraphs"] == 6
        assert (core.job_dir(job_id) / "state.json").stat().st_mtime_ns == saved_at, \
            "已有大纲时不得再次写盘"

        started_state = core.load_job_state(job_id)
        started_state.pop("source_outline", None)
        started_state.update(p1_done=True, paras=["已开始"])
        core.save_job_state(job_id, started_state)
        started = core.ensure_source_outline(job_id, first)  # intentionally stale snapshot
        assert "source_outline" not in started
        assert started["paras"] == ["已开始"]
    finally:
        core.OUTPUT_DIR = old_output


# ------------------------------------------------- 补充探针回归（验收二轮）

def test_multi_entry_plan_applies_all_entries_without_index_drift(tmp_path):
    """[P1] 多条目一次落账：先汇总全部索引再统一移除，逐条删除不再错位。

    审校探针的场景：真实两章 DOCX 全选排除后，账本四段全部排除，但第二章
    两段仍到达翻译入口。修复后工作列表必须清空，账本来源与原文一一对应。
    """
    job_id = "scope-multi-entry"
    sources = ["第一章 缘起", "Alpha body.", "第二章 争论", "Beta body."]
    outline = _two_chapter_outline()
    outline["paragraphs"] = list(sources)
    outline["total_paragraphs"] = len(sources)
    outline["total_chars"] = sum(len(item) for item in sources)
    outline["nodes"][0].update(paragraphs=sources[:2], para_count=2,
                               chars=sum(len(item) for item in sources[:2]))
    outline["nodes"][1].update(para_start=2, paragraphs=sources[2:],
                               para_count=2,
                               chars=sum(len(item) for item in sources[2:]))
    state = _docx_outline_state(job_id)
    state["source_outline"] = outline
    old_output = _persist(tmp_path, job_id, state)
    try:
        # 先在译前把两章都记成计划（此时还没有翻译单元，p1_done 为假）。
        core.exclude_translation_scope(
            job_id, ["outline:n0000", "outline:n0001"])
        state = core.load_job_state(job_id)
        assert len(state["scope_exclusion_plan"]["entries"]) == 2
        # 任务开始：建好翻译单元（与大纲逐段对应）。
        state = core.load_job_state(job_id)
        state["paras"] = list(sources)
        state["source_paras"] = list(sources)
        state["source_paragraphs"] = list(sources)
        state["p1_done"] = True
        state["segmentation"] = {"mode": "sentence", "locked": True,
                                 "paragraph_ranges": [
            {"paragraph_index": i, "start": i, "end": i + 1}
            for i in range(4)]}
        core.save_job_state(job_id, state)
        state = core.load_job_state(job_id)
        outcome = core.apply_scope_exclusion_plan(job_id, state)
        core.save_job_state(job_id, state)

        assert len(outcome["applied"]) == 2 and outcome["unapplied"] == []
        state = core.load_job_state(job_id)
        # 工作列表彻底清空：两章的段都到不了翻译入口。
        assert state["paras"] == []
        records = core.excluded_segment_records(state)
        assert [record["source"] for record in records] == list(sources)
        reasons = {record["source"]: record["reason"] for record in records}
        assert reasons["第一章 缘起"] == "范围排除（译前）：第一章 缘起"
        assert reasons["Beta body."] == "范围排除（译前）：第二章 争论"
        assert state.get("scope_exclusion_plan") is None
        # 审计里两条目各 2 段，没有把四段算到一个条目头上。
        counts = sorted(item["excluded_count"] for item in state["scope_exclusion_applied"])
        assert counts == [2, 2]
    finally:
        core.OUTPUT_DIR = old_output


def test_plan_applies_through_many_to_one_cleanup_mapping(tmp_path):
    """[P1] 清洗映射（output_input_indices）在计划应用时已可读。

    探针场景：``[[0, 1], [2]]`` 的清洗映射把前两段原始段落并成一个翻译单元；
    排除第三页（原始第 3 段）必须经映射落到清洗后单元 1，而不是被报成
    "结构不一致"后仍进入翻译。
    """
    job_id = "scope-plan-cleanup-map"
    raw = ["Alpha one.", "Alpha two.", "Beta body."]
    cleaned = ["Alpha one.\nAlpha two.", "Beta body."]
    state = core.new_job_state("three-pages.docx")
    state["filename"] = "three-pages.docx"
    old_output = _persist(tmp_path, job_id, state)
    try:
        state = core.load_job_state(job_id)
        state["source_outline"] = {
            "format": "DOCX", "total_paragraphs": 3,
            "total_chars": sum(len(item) for item in raw),
            "paragraphs": list(raw),
            "nodes": [{"id": "n0000", "kind": "chapter", "title": "Beta",
                       "level": 1, "para_start": 2, "para_count": 1,
                       "paragraphs": raw[2:], "chars": len(raw[2]),
                       "page_start": None, "page_end": None}],
            "pages": 0, "scanned": False, "partial": False, "fingerprint": "x",
        }
        core.save_job_state(job_id, state)
        core.exclude_translation_scope(job_id, ["outline:n0000"])
        # 任务开始：清洗把 0、1 并成单元 0，原始第 2 段是单元 1。
        state = core.load_job_state(job_id)
        state["paras"] = list(cleaned)
        state["source_paras"] = list(raw)
        state["source_paragraphs"] = list(cleaned)
        state["source_cleanup"] = {"output_input_indices": [[0, 1], [2]]}
        state["segmentation"] = {"mode": "sentence", "locked": True,
                                 "paragraph_ranges": [
            {"paragraph_index": 0, "start": 0, "end": 1},
            {"paragraph_index": 1, "start": 1, "end": 2}]}
        state["p1_done"] = True
        core.save_job_state(job_id, state)
        state = core.load_job_state(job_id)
        outcome = core.apply_scope_exclusion_plan(job_id, state)
        core.save_job_state(job_id, state)

        assert outcome["unapplied"] == [], outcome
        assert [item["label"] for item in outcome["applied"]] == ["Beta"]
        state = core.load_job_state(job_id)
        assert state["paras"] == ["Alpha one.\nAlpha two."]
        records = core.excluded_segment_records(state)
        assert [record["source"] for record in records] == ["Beta body."]
    finally:
        core.OUTPUT_DIR = old_output


def test_started_task_with_all_excluded_keeps_ledger_visible_in_scope(tmp_path):
    """[P2] 已开始的任务全部排除后：范围树回到账本口径，恢复记录不丢。"""
    job_id = "scope-all-excluded-ledger"
    old_output = _persist(tmp_path, job_id, _state())
    try:
        first = _chapter_nodes(core.translation_scope_tree(
            core.load_job_state(job_id)))["nodes"][0]["key"]
        core.exclude_translation_scope(job_id, [first])
        # 范围键跟着工作列表重算：排除一章后剩下的章拿新的键。
        second = _chapter_nodes(core.translation_scope_tree(
            core.load_job_state(job_id)))["nodes"][0]["key"]
        core.exclude_translation_scope(job_id, [second])
        state = core.load_job_state(job_id)
        assert state["paras"] == [] and state["p1_done"]
        tree = core.translation_scope_tree(state)
        assert tree["available"] is False
        assert tree["pre_start"] is False, "已开始的任务不得落进译前分支"
        assert tree["excluded"]["count"] == 4, "账本记录必须能在对话框里恢复"
        assert all(record["source"] for record in tree["excluded"]["records"])
    finally:
        core.OUTPUT_DIR = old_output


def test_pre_start_totals_reflect_the_exclusion_plan(tmp_path):
    """[P2] 译前排除后的总量与预估都要扣除计划范围。

    预估按**正常使用顺序**验证：排除前先看一次预估（填充缓存），排除第二章
    之后**不清缓存**再读一次——缓存命中也必须走计划扣减（审校探针的场景：
    缓存命中绕过扣减，预估停留在四段）。
    """
    job_id = "scope-plan-totals"
    state = _docx_outline_state(job_id)
    state["filename"] = "plan.md"  # 预估走本地文字层提取（txt/md 直接读）
    state["source_outline"] = _two_chapter_outline()
    old_output = _persist(tmp_path, job_id, state)
    try:
        from folith.qt import adapter as qt_adapter
        qt_adapter._source_size_cache.clear()
        core.save_source(job_id, "第一章 缘起\n\nAlpha body one.\n\nAlpha body two.\n\n"
                         "第二章 争论\n\nBeta body one.\n\nBeta body two.\n".encode("utf-8"))
        before = qt_adapter.source_size(job_id)
        assert before and before["segments"] == 6, before

        core.exclude_translation_scope(job_id, ["outline:n0001"])
        state = core.load_job_state(job_id)
        tree = core.translation_scope_tree(state)
        # 总量扣掉第二章：剩第一章的 3 段。
        assert tree["total"]["segments"] == 3
        assert tree["total"]["chars"] == sum(len(item) for item in
                                             state["source_outline"]["paragraphs"][:3])
        totals = core.scope_plan_totals(state)
        assert totals["segments"] == 3
        assert totals["chars"] == sum(len(item) for item in
                                      state["source_outline"]["paragraphs"][3:])

        # 缓存命中路径：预估同步扣减（排除前后的差异即计划范围）。
        after = qt_adapter.source_size(job_id)
        assert after["segments"] == 3, after
        assert after["chars"] == before["chars"] - totals["chars"], after
    finally:
        core.OUTPUT_DIR = old_output


def _cleanup_mapping_state(job_id, tmp_path, raw, cleaned, mapping):
    """建好翻译单元的任务：清洗映射为 ``mapping``（清洗后段 → 原始段列表）。"""
    state = core.new_job_state("three-pages.pdf")
    state["filename"] = "three-pages.pdf"
    old_output = _persist(tmp_path, job_id, state)
    state = core.load_job_state(job_id)
    state["source_outline"] = {
        "format": "PDF", "total_paragraphs": len(raw),
        "total_chars": sum(len(item) for item in raw),
        "paragraphs": list(raw),
        "nodes": [{"id": "n0000", "kind": "chapter", "title": "第一节",
                   "level": 1, "para_start": 0, "para_count": 2,
                   "paragraphs": raw[:2], "chars": sum(len(x) for x in raw[:2]),
                   "page_start": None, "page_end": None}],
        "pages": 0, "scanned": False, "partial": False, "fingerprint": "x",
    }
    core.save_job_state(job_id, state)
    return old_output, state


def test_partial_cleanup_merge_never_deletes_kept_content(tmp_path):
    """[P1] 清洗合并跨越排除边界：部分覆盖整条拒用，保留内容一个不删。

    审校探针的复现：三页 PDF，清洗映射 ``[[0, 1], [2]]`` 把第 1、2 页合并成
    一个翻译单元；只排除第二页时，删掉合并单元会把用户保留的第一页一起
    移走。修复后整条计划按未应用上报（带警告），翻译列表原样保留。
    """
    job_id = "scope-partial-merge"
    raw = ["第一页的内容。", "第二页的内容。", "第三页的内容。"]
    cleaned = ["第一页的内容。\n第二页的内容。", "第三页的内容。"]
    old_output, _state = _cleanup_mapping_state(job_id, tmp_path, raw, cleaned,
                                                [[0, 1], [2]])
    try:
        # 大纲里"第二节"覆盖原始第 2 段（第二页）。
        state = core.load_job_state(job_id)
        state["source_outline"]["nodes"].append(
            {"id": "n0001", "kind": "chapter", "title": "第二节",
             "level": 1, "para_start": 1, "para_count": 1,
             "paragraphs": raw[1:2], "chars": len(raw[1]),
             "page_start": None, "page_end": None})
        core.save_job_state(job_id, state)
        core.exclude_translation_scope(job_id, ["outline:n0001"])
        # 任务开始：清洗合并了第 1、2 页。
        state = core.load_job_state(job_id)
        state["paras"] = list(cleaned)
        state["source_paras"] = list(raw)
        state["source_paragraphs"] = list(cleaned)
        state["source_cleanup"] = {"output_input_indices": [[0, 1], [2]]}
        state["segmentation"] = {"mode": "sentence", "locked": True,
                                 "paragraph_ranges": [
            {"paragraph_index": 0, "start": 0, "end": 1},
            {"paragraph_index": 1, "start": 1, "end": 2}]}
        state["p1_done"] = True
        core.save_job_state(job_id, state)
        state = core.load_job_state(job_id)
        outcome = core.apply_scope_exclusion_plan(job_id, state)
        core.save_job_state(job_id, state)

        assert outcome["applied"] == []
        assert [item["label"] for item in outcome["unapplied"]] == ["第二节"]
        state = core.load_job_state(job_id)
        # 保留的内容一个不动：合并单元（含第一页）原样留在翻译列表。
        assert state["paras"] == cleaned
        assert state.get("excluded_segments") in (None, [])
        assert any("合并" in w for w in state["warnings"]), state["warnings"]
        assert state["scope_exclusion_applied"][-1]["partial_coverage"]

        # 正例：把第一、二页一起排除（并集完整覆盖合并单元）则合法应用。
        # （apply 之后 source_outline 已被清出；直接注入计划，apply 只对锚点。）
        state = core.load_job_state(job_id)
        state["excluded_segments"] = []
        state["scope_exclusion_applied"] = []
        state["scope_exclusion_plan"] = {
            "entries": [{"excluded_id": "x900001", "key": "outline:n0000",
                         "kind": "chapter", "label": "第一二节",
                         "para_start": 0, "para_end": 2,
                         "anchor_first": raw[0], "anchor_last": raw[1],
                         "reason": "范围排除（译前）：第一二节",
                         "planned_at": "2026-10-25T00:00:00+00:00",
                         "planned_by": "user"}],
            "fingerprint": "x"}
        core.save_job_state(job_id, state)
        state = core.load_job_state(job_id)
        outcome = core.apply_scope_exclusion_plan(job_id, state)
        core.save_job_state(job_id, state)
        assert outcome["unapplied"] == []
        state = core.load_job_state(job_id)
        assert state["paras"] == ["第三页的内容。"], \
            "并集完整覆盖合并单元时才允许删除"
        # 账本记录的是被移出的**翻译单元**（合并后的文本），一条。
        records = core.excluded_segment_records(state)
        assert len(records) == 1
        assert records[0]["source"] == cleaned[0]
        assert records[0]["reason"] == "范围排除（译前）：第一二节"
    finally:
        core.OUTPUT_DIR = old_output


def test_imported_docx_plan_runs_through_the_real_pipeline(tmp_path, monkeypatch):
    """真实 DOCX → 本地范围 → 管线钩子 → 翻译请求 → 交付账本 → 恢复。"""
    import io
    import re
    from docx import Document
    from folith.qt import adapter

    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job_id = "scope-real-pipeline"
    document = Document()
    for title, body in [("第一章 缘起", "Alpha body one."),
                        ("第二章 争论", "Beta body one.")]:
        document.add_heading(title, level=1)
        document.add_paragraph(body)
    stream = io.BytesIO()
    document.save(stream)
    data = stream.getvalue()
    core.save_source(job_id, data)
    core.save_job_state(job_id, core.new_job_state("scope.docx"))
    tree = adapter.translation_scope(job_id)
    second = tree["groups"][0]["nodes"][1]
    adapter.exclude_scope(job_id, [second["key"]])
    requests = []

    def offline_provider(*args, **kwargs):
        system, user = args[3:5]
        requests.append(user)
        if "学术翻译专家" in system:
            sources = re.findall(r"^\s*\d+\.\s+(.+?)\s*$",
                                 user.rsplit("待翻译段落", 1)[-1], re.MULTILINE)
            return json.dumps([f"这是第 {i+1} 段中文译文。" for i in range(len(sources))])
        return "[]"

    monkeypatch.setattr(core, "call_llm", offline_provider)
    state = core.run_job_pipeline(
        job_id, "scope.docx", data, provider="DeepSeek", api_key="offline",
        model="deepseek-chat", target_lang="简体中文", auto_term=False,
        enable_report=False, translation_theory="目的论", enable_review=False,
        enable_annotate=False, use_tm=False, enable_understanding=False,
        enable_source_cleanup=False, delivery_config=core.default_delivery_config())
    assert state["p2_done"]
    assert requests, "管线确实调用过 provider 替身"
    assert all("Beta body" not in user and "第二章 争论" not in user for user in requests)
    assert [pair["source"] for pair in state["pairs"]] == ["第一章 缘起", "Alpha body one."]
    records = core.excluded_segment_records(state)
    assert [r["source"] for r in records] == ["第二章 争论", "Beta body one."]
    manifest = core.excluded_segments_manifest(state)
    assert manifest["origins"]["user"] == 2
    result = adapter.restore_scope(job_id, [r["excluded_id"] for r in records])
    assert result["restored"] == 2
    reopened = core.load_job_state(job_id)
    assert reopened["paras"] == ["第一章 缘起", "Alpha body one.", "第二章 争论", "Beta body one."]

    assert len(reopened["pairs"]) == len(reopened["paras"])
    assert [p["source"] for p in reopened["pairs"]] == reopened["paras"]
    assert [p["target"] for p in reopened["pairs"][:2]] == [p["target"] for p in state["pairs"]]
    assert all(not p.get("target") and not p.get("reviewed") for p in reopened["pairs"][2:])


def test_restore_pretranslation_exclusion_does_not_shift_later_translations(tmp_path):
    state = _state(translated=True)
    before_pairs = json.loads(json.dumps(state["pairs"]))
    # Simulate a source-only exclusion from before translation, in the middle.
    state["paras"].pop(1)
    state["pairs"].pop(1)
    state["excluded_segments"] = [{"excluded_id": "x000001", "index": 1,
        "source": "Alpha body.", "target": "", "pair": None,
        "before_source": "第一章 缘起", "after_source": "第二章 争论",
        "excluded_by": "user"}]
    old_output = _persist(tmp_path, "scope-restore-middle", state)
    try:
        restored = core.restore_translation_scope("scope-restore-middle", ["x000001"])["state"]
        assert len(restored["pairs"]) == len(restored["paras"]) == 4
        assert [p["source"] for p in restored["pairs"]] == restored["paras"]
        assert restored["pairs"][1]["target"] == ""
        assert restored["pairs"][2]["target"] == before_pairs[2]["target"]
        assert restored["pairs"][3]["target"] == before_pairs[3]["target"]
    finally:
        core.OUTPUT_DIR = old_output


def test_unapplied_plan_stops_real_pipeline_until_explicitly_withdrawn(tmp_path, monkeypatch):
    from folith.qt import adapter
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path)
    job = "scope-unmapped-stop"
    state = core.new_job_state("scope.md")
    core.save_source(job, b"# Chapter One\n\nAlpha body.\n\n# Chapter Two\n\nBeta body.\n")
    core.save_job_state(job, state)
    tree = adapter.translation_scope(job)
    adapter.exclude_scope(job, [tree["groups"][0]["nodes"][-1]["key"]])
    state = core.load_job_state(job)
    state["scope_exclusion_plan"]["entries"][0]["anchor_first"] = "mismatched source"
    core.save_job_state(job, state)
    calls = []
    def offline(*args, **kwargs):
        calls.append(args)
        raise AssertionError("Unapplied exclusion must not reach a provider")
    monkeypatch.setattr(core, "call_llm", offline)
    config = dict(provider="DeepSeek", api_key="offline", model="deepseek-chat",
        target_lang="简体中文", auto_term=False, enable_report=False,
        translation_theory="目的论", enable_review=False, enable_annotate=False,
        use_tm=False, enable_understanding=True, enable_source_cleanup=False)
    with pytest.raises(RuntimeError, match="译前范围排除未能应用"):
        core.run_job_pipeline(job, "scope.md", core.load_source(job), **config)
    assert calls == []
    tree = adapter.translation_scope(job)
    assert tree["excluded"]["unapplied_plan"] is True
    pending = [r for r in tree["excluded"]["records"] if r.get("planned")]
    assert len(pending) == 1
    # Merely retrying is not consent to translate excluded content.
    with pytest.raises(RuntimeError, match="译前范围排除未能应用"):
        core.run_job_pipeline(job, "scope.md", None, **config)
    assert calls == []
    adapter.restore_scope(job, [pending[0]["excluded_id"]])
    assert not core.load_job_state(job)["scope_exclusion_plan"]["entries"]
    assert not adapter.translation_scope(job)["excluded"]["unapplied_plan"]


def test_chapter_scope_node_id_in_exclusion_records(tmp_path):
    """[Issue #190] 排除记录携带结构化 node_id 与 scope_key，支持章节粒度衍生统计与报告。"""
    from folith.engine.segments import _excluded_segment_record

    # 1. 直接构造单条记录：未指定时缺省，指定时记录
    state = {"exclusion_seq": 1}
    rec_plain = _excluded_segment_record(state, 0, reason="test", actor="user")
    assert "node_id" not in rec_plain
    assert "scope_key" not in rec_plain

    rec_scoped = _excluded_segment_record(state, 1, reason="test", actor="user",
                                          node_id="n0002", scope_key="outline:n0002")
    assert rec_scoped.get("node_id") == "n0002"
    assert rec_scoped.get("scope_key") == "outline:n0002"

    # 2. 译前大纲计划排除：落账后携带 node_id 与 scope_key
    job_id = "scope-node-id-plan"
    raw = ["Chapter 1", "Body 1", "Chapter 2", "Body 2"]
    st = core.new_job_state("doc.docx")
    st["filename"] = "doc.docx"
    old_output = _persist(tmp_path, job_id, st)
    try:
        st = core.load_job_state(job_id)
        st["source_outline"] = {
            "format": "DOCX", "total_paragraphs": 4,
            "total_chars": sum(len(x) for x in raw),
            "paragraphs": list(raw),
            "nodes": [
                {"id": "n0001", "kind": "chapter", "title": "Chapter 1",
                 "level": 1, "para_start": 0, "para_count": 2,
                 "paragraphs": raw[:2], "chars": sum(len(x) for x in raw[:2])},
                {"id": "n0002", "kind": "chapter", "title": "Chapter 2",
                 "level": 1, "para_start": 2, "para_count": 2,
                 "paragraphs": raw[2:], "chars": sum(len(x) for x in raw[2:])},
            ],
            "pages": 0, "scanned": False, "partial": False, "fingerprint": "fp1",
        }
        core.save_job_state(job_id, st)
        # 排除 Chapter 2
        core.exclude_translation_scope(job_id, ["outline:n0002"])
        plan_st = core.load_job_state(job_id)
        entries = plan_st["scope_exclusion_plan"]["entries"]
        assert len(entries) == 1
        assert entries[0].get("node_id") == "n0002"
        assert entries[0].get("key") == "outline:n0002"

        # 模拟建立翻译单元并应用计划
        plan_st["paras"] = list(raw)
        plan_st["source_paras"] = list(raw)
        plan_st["source_paragraphs"] = list(raw)
        plan_st["segmentation"] = {
            "mode": "sentence", "locked": True,
            "paragraph_ranges": [{"paragraph_index": i, "start": i, "end": i + 1}
                                 for i in range(len(raw))],
        }
        plan_st["p1_done"] = True
        core.save_job_state(job_id, plan_st)

        st_applied = core.load_job_state(job_id)
        outcome = core.apply_scope_exclusion_plan(job_id, st_applied)
        assert len(outcome["applied"]) == 1
        core.save_job_state(job_id, st_applied)

        final_st = core.load_job_state(job_id)
        records = core.excluded_segment_records(final_st)
        assert len(records) == 2
        for r in records:
            assert r.get("node_id") == "n0002"
            assert r.get("scope_key") == "outline:n0002"

        # 3. 任务运行后按范围排除（post-start）：同样携带结构化标识
        core.exclude_translation_scope(job_id, ["chapter:0000"])
        post_st = core.load_job_state(job_id)
        post_records = core.excluded_segment_records(post_st)
        assert len(post_records) == 4
        new_records = [r for r in post_records if r.get("scope_key") == "chapter:0000"]
        assert len(new_records) == 2

        # 4. 恢复测试：恢复不破坏兼容性
        core.restore_translation_scope(job_id, [r["excluded_id"] for r in new_records])
        restored_st = core.load_job_state(job_id)
        assert len(restored_st["paras"]) == 2
    finally:
        core.OUTPUT_DIR = old_output
