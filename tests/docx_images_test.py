"""Actual source-to-DOCX ordering and delivery regressions."""
import io
import json
from copy import deepcopy
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

import core
from folith.engine.document_export import build_translation_docx
from folith.epub import bind_source_references
from epub_export_test import docx_source, pdf_source, translated_state, illustration


def flow(payload):
    document = Document(io.BytesIO(payload.getvalue()))
    return document, ["[image]" if paragraph._p.findall('.//' + qn('a:blip'))
                      else paragraph.text for paragraph in document.paragraphs]


@pytest.mark.parametrize('filename,maker', [('book.docx', docx_source), ('book.pdf', pdf_source)])
def test_source_images_between_translated_paragraphs(filename, maker):
    source = maker()
    state = translated_state(filename, source)
    before = deepcopy(state)
    payload, report = build_translation_docx(state, source)
    document, parts = flow(payload)
    first = next(i for i, text in enumerate(parts) if 'first paragraph' in text)
    second = next(i for i, text in enumerate(parts) if 'second paragraph' in text)
    assert first < parts.index('[image]') < second
    assert report['image_count'] == len(document.inline_shapes) == 1
    assert state == before
    if filename.endswith('.docx'):
        with ZipFile(payload) as archive:
            assert illustration() in [archive.read(name) for name in archive.namelist()
                                      if name.startswith('word/media/')]
    else:
        assert parts.index('[image]') < parts.index('Figure 1. A blue illustration.') < second


def test_split_merge_exclude_restore_do_not_duplicate_or_remove_picture(tmp_path, monkeypatch):
    monkeypatch.setattr(core, 'OUTPUT_DIR', tmp_path)
    source = docx_source()
    state = translated_state('book.docx', source)
    core.save_source('pictures', source)
    core.save_job_state('pictures', state)
    state = core.mutate_translation_segments('pictures', 1, 'split',
        source_a='The first paragraph', source_b='describes the picture.',
        target_a='图前第一部分', target_b='图前第二部分')
    for operation in ('merge', 'exclude', 'include'):
        kwargs = {'exclude_reason': '测试'} if operation == 'exclude' else {}
        if operation == 'include':
            kwargs = {'excluded_id': state['excluded_segments'][0]['excluded_id']}
        state = core.mutate_translation_segments('pictures', 1, operation, **kwargs)
        _, parts = flow(build_translation_docx(state, source)[0])
        assert parts.count('[image]') == 1
        assert ('图前第一部分' in '\n'.join(parts)) == (operation != 'exclude')


def test_cross_picture_merge_target_is_emitted_once():
    source = docx_source()
    state = translated_state('book.docx', source)
    bind_source_references(state)
    # Merge the first body paragraph and following caption across the picture.
    first = next(i for i, pair in enumerate(state['pairs']) if 'first paragraph' in pair['source'])
    last = next(i for i, pair in enumerate(state['pairs']) if 'Figure 1.' in pair['source'])
    refs = list(dict.fromkeys(ref for pair in state['pairs'][first:last+1]
                              for ref in pair['epub_source_paragraphs']))
    state['pairs'][first:last+1] = [dict(source='merged', target='完整合并译文', epub_source_paragraphs=refs)]
    _, parts = flow(build_translation_docx(state, source)[0])
    assert parts.count('完整合并译文') == 1
    assert parts.index('完整合并译文') < parts.index('[image]')


def test_unreliable_source_association_fails_without_losing_images():
    source = docx_source()
    state = translated_state('book.docx', source)
    state['segment_structure_history'] = [{'operation': 'merge'}]
    with pytest.raises(RuntimeError, match='译文 DOCX.*来源关联'):
        build_translation_docx(state, source)
    state.pop('segment_structure_history')
    state['source_paras'][0] = 'different source'
    with pytest.raises(RuntimeError, match='来源不一致'):
        build_translation_docx(state, source)


def test_cover_repeated_and_tall_images_fit_page():
    doc = Document()
    doc.add_picture(io.BytesIO(illustration()), width=Inches(20), height=Inches(40))
    doc.add_paragraph('A paragraph between the two pictures.')
    doc.add_picture(io.BytesIO(illustration()), width=Inches(3))
    source_io = io.BytesIO()
    doc.save(source_io)
    source = source_io.getvalue()
    state = translated_state('cover.docx', source)
    payload, report = build_translation_docx(state, source)
    document, parts = flow(payload)
    assert parts[1] == '[image]'
    assert parts.count('[image]') == report['image_count'] == 2
    section = document.sections[0]
    for shape in document.inline_shapes:
        assert shape.width <= section.page_width - section.left_margin - section.right_margin
        assert shape.height <= section.page_height - section.top_margin - section.bottom_margin


def test_delivery_bundle_and_snapshot_contain_pictures(tmp_path, monkeypatch):
    monkeypatch.setattr(core, 'OUTPUT_DIR', tmp_path)
    source = docx_source()
    state = translated_state('book.docx', source)
    state['delivery_config'] = core.normalize_delivery_config({'deliver_plain_docx': True})
    core.save_source('pictures', source)
    core.save_job_state('pictures', state)
    bundle = core._delivery_asset_bundle('pictures', state, '简体中文', '', '')
    assert len(Document(io.BytesIO(bundle['translation.docx'])).inline_shapes) == 1
    assert json.loads(bundle['docx_export_report.json'])['image_count'] == 1
    snapshot = core.create_delivery_snapshot('pictures', state)
    assets = core.delivery_snapshot_assets('pictures', snapshot['snapshot_version'])
    assert len(Document(io.BytesIO(assets['translation.docx'])).inline_shapes) == 1


def test_text_source_and_no_picture_docx_keep_plain_translation():
    doc = Document()
    doc.add_paragraph('This document has no pictures.')
    source_io = io.BytesIO()
    doc.save(source_io)
    for state, source in [(translated_state('plain.docx', source_io.getvalue()), source_io.getvalue()),
                          ({'filename': 'plain.txt', 'pairs': [{'target': '纯文字译文'}]}, b'source')]:
        payload, report = build_translation_docx(state, source)
        document, parts = flow(payload)
        assert not document.inline_shapes
        assert report['image_count'] == 0
        assert parts[1:] == [pair['target'] for pair in state['pairs']]


def test_missing_translation_stays_blank_and_missing_source_is_reported():
    source = docx_source()
    state = translated_state('book.docx', source)
    state['pairs'][1]['target'] = ''
    payload, report = build_translation_docx(state, source)
    document, parts = flow(payload)
    assert len(document.inline_shapes) == 1
    assert '译文2' not in '\n'.join(parts)
    payload, report = build_translation_docx(state, None)
    assert not flow(payload)[0].inline_shapes
    assert any('缺少原始文件' in warning for warning in report['warnings'])


def test_tall_bitmap_preserves_aspect_ratio_and_fits_page():
    import fitz
    with fitz.open() as image_doc:
        page = image_doc.new_page(width=200, height=2000)
        page.draw_rect(page.rect, fill=(0.2, 0.4, 0.6))
        image = page.get_pixmap().tobytes('png')
    doc = Document()
    doc.add_paragraph('Text before the very tall picture.')
    doc.add_picture(io.BytesIO(image))
    doc.add_paragraph('Text after the very tall picture.')
    buffer = io.BytesIO()
    doc.save(buffer)
    source = buffer.getvalue()
    state = translated_state('tall.docx', source)
    payload, _ = build_translation_docx(state, source)
    document, parts = flow(payload)
    picture = document.inline_shapes[0]
    section = document.sections[0]
    assert picture.height < section.page_height - section.top_margin - section.bottom_margin
    assert picture.width / picture.height == pytest.approx(0.1, abs=0.00001)
    assert parts.index('[image]') == 2


def test_image_only_source_falls_back_to_plain_docx(monkeypatch):
    from folith import epub

    def no_text(_source):
        raise epub.ImageOnlySourceError('EPUB 原文件只有图像、没有可关联的文本')

    monkeypatch.setattr(epub, '_pdf_source', no_text)
    state = {'filename': 'scan.pdf', 'target_lang': '简体中文',
             'pairs': [{'source': 'Scanned line.', 'target': '扫描行。'}]}
    payload, report = build_translation_docx(state, b'%PDF-1.4 scanned')
    assert '扫描行。' in [p.text for p in Document(io.BytesIO(payload if isinstance(payload, bytes) else payload.getvalue())).paragraphs]
    assert report['image_count'] == 0 and '纯文字重建' in report['warnings'][0]
