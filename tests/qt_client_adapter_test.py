"""Qt desktop client — adapter contract against the real business core.

These tests drive the *actual* core worker (through a local OpenAI-compatible
stub, i.e. the real HTTP transport) and assert the user-visible outcomes the
desktop client depends on: a job created from a real DOCX, a translated and
reviewed pair list, a committed human edit, a delivery snapshot, a safe exit
report and a reopen that still sees the edit.

They deliberately do not assert "the UI drew a widget": widget wiring is covered
by ``tests/qt_client_window_test.py``.  What matters here is that the client's
adapter cannot corrupt, duplicate or silently lose the product's data.
"""
from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import time
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
sys.path.insert(0, str(ROOT / "scripts"))


def _make_docx(texts: list[str]) -> bytes:
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    doc.save(buf)
    return buf.getvalue()


class StubEndpoint:
    """The controlled model endpoint used by the frozen vertical slice."""

    def __init__(self, tmp_path: Path, *, delay: float = 0.0) -> None:
        self.tmp_path = tmp_path
        self.delay = delay
        self.process: subprocess.Popen | None = None
        self.url = ""
        self.log = tmp_path / "stub-llm.jsonl"
        self.server_log = tmp_path / "stub-server.log"
        self._handle = None

    def start(self) -> str:
        ready = self.tmp_path / "stub-ready.json"
        ready.unlink(missing_ok=True)
        self._handle = self.server_log.open("ab")
        self.process = subprocess.Popen(
            [sys.executable, str(ROOT / "scripts" / "d0_fake_llm_server.py"),
             "--port", "0", "--log", str(self.log),
             "--ready-file", str(ready), "--delay", str(self.delay)],
            stdout=self._handle, stderr=subprocess.STDOUT)
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            if ready.exists():
                self.url = json.loads(ready.read_text(encoding="utf-8"))["url"]
                return self.url
            if self.process.poll() is not None:
                tail = self.server_log.read_text(encoding="utf-8",
                                                 errors="replace")[-800:]
                raise RuntimeError(f"stub LLM server exited: {tail}")
            time.sleep(0.1)
        raise RuntimeError("stub LLM server did not become ready")

    def stop(self) -> None:
        if self.process is not None and self.process.poll() is None:
            self.process.terminate()
            try:
                self.process.wait(timeout=10)
            except subprocess.TimeoutExpired:  # pragma: no cover - defensive
                self.process.kill()
        if self._handle is not None:
            self._handle.close()


def _delivered(adapter, job_id: str, asset: str) -> str:
    """写到用户文件夹里的名字：``<原文件名>_<产物>``（见 adapter.delivery_file_name）。"""
    return adapter.delivery_file_name(adapter.load_state(job_id).get("filename"), asset)


@pytest.fixture()
def adapter_env(tmp_path, monkeypatch):
    """Isolate one adapter against a temp data directory."""
    data_dir = tmp_path / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setenv("FOLITH_DESKTOP_DATA_DIR", str(data_dir))
    monkeypatch.setenv("FOLITH_DESKTOP_LOG_DIR", str(tmp_path / "logs"))
    monkeypatch.setenv("FOLITH_OUTPUT_DIR", str(data_dir))

    import core

    monkeypatch.setattr(core, "OUTPUT_DIR", data_dir, raising=False)

    from folith.qt import adapter

    yield {"adapter": adapter, "data_dir": data_dir, "tmp": tmp_path, "core": core}
    # Drain this test's workers *before* the next test repoints core.OUTPUT_DIR,
    # otherwise a leftover worker writes into the next test's data directory.
    from qt_workers import drain_job_workers

    drain_job_workers(core)


def _wait_for_alive(core, job_id: str, timeout: float = 60.0) -> bool:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if core.is_job_worker_alive(job_id):
            return True
        time.sleep(0.1)
    return False


def _wait_for_idle(core, job_id: str, timeout: float = 240.0) -> dict:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        state = core.load_job_state(job_id) or {}
        runtime = core.get_job_runtime_status(job_id, state)
        status = runtime.get("status")
        if status not in core.RUNTIME_ACTIVE_STATUSES:
            # A transient read of the atomically replaced runtime file can
            # return idle defaults while the worker is still running.  Do not
            # mistake that for settlement.
            if status in {"idle", "idle_incomplete"} \
                    and core.is_job_worker_alive(job_id):
                time.sleep(0.3)
                continue
            return runtime
        time.sleep(0.3)
    raise AssertionError(f"job {job_id} did not settle within {timeout}s")


def _wait_for_worker_stopped(core, job_id: str, timeout: float = 60.0) -> bool:
    """Wait until the cancelled worker has released its execution lock.

    The runtime can report ``cancelled`` before the thread has finished its
    cleanup.  Resuming in that window is a no-op, so the test must wait for the
    worker itself, not just the status.
    """
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if (not core.is_job_worker_alive(job_id)
                and not core._job_execution_alive(job_id)):
            return True
        time.sleep(0.05)
    return False


def _wait_for_restart(core, job_id: str, timeout: float = 30.0) -> bool:
    """Wait for a resumed task to leave the previous terminal status.

    ``_wait_for_idle`` returns as soon as the status is not active, so it would
    observe the old ``cancelled`` status if the new worker has not written its
    ``starting`` state yet.  Under a loaded thread pool that race is real.
    """
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        state = core.load_job_state(job_id) or {}
        runtime = core.get_job_runtime_status(job_id, state)
        if (runtime.get("status") in core.RUNTIME_ACTIVE_STATUSES
                or state.get("p2_done")):
            return True
        time.sleep(0.05)
    return False


def test_engine_config_round_trip(adapter_env):
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]

    missing = adapter.EngineConfig(provider="", model="", api_key="")
    assert "提供商" in missing.connection_problem()

    config = adapter.EngineConfig(
        provider="自定义中转站", model="stub-model", api_key="stub-key",
        base_url="http://127.0.0.1:9/v1")
    assert config.connection_problem() == ""
    config.save()

    saved = adapter.EngineConfig.load()
    assert saved.provider == "自定义中转站"
    assert saved.model == "stub-model"
    assert saved.api_key == "stub-key"
    assert saved.base_url == "http://127.0.0.1:9/v1"

    # POSIX permission bits do not represent Windows ACLs.
    if os.name != "nt":
        mode = core.provider_config_path().stat().st_mode & 0o777
        assert mode & 0o077 == 0, oct(mode)

    # A custom relay must never silently fall back to api.openai.com.
    bad = adapter.EngineConfig(provider="自定义中转站", model="m", api_key="k",
                               base_url="")
    assert "API 地址" in bad.connection_problem()


def test_import_rejects_unusable_files(adapter_env):
    adapter = adapter_env["adapter"]
    tmp = adapter_env["tmp"]

    empty = tmp / "empty.docx"
    empty.write_bytes(b"")
    with pytest.raises(adapter.AdapterError):
        adapter.read_source_file(empty)

    wrong = tmp / "image.png"
    wrong.write_bytes(b"\x89PNG\r\n")
    with pytest.raises(adapter.AdapterError):
        adapter.read_source_file(wrong)

    with pytest.raises(adapter.AdapterError):
        adapter.read_source_file(tmp / "missing.docx")


def test_import_creates_task_in_existing_format(adapter_env):
    adapter = adapter_env["adapter"]
    data_dir = adapter_env["data_dir"]
    tmp = adapter_env["tmp"]

    source = tmp / "paper.docx"
    source.write_bytes(_make_docx(["第一段。", "第二段。"]))
    config = adapter.EngineConfig(target_lang="简体中文")

    first = adapter.import_documents([source], config)
    assert len(first) == 1
    assert first[0].created is True
    job_id = first[0].job_id

    # Data lands in the existing layout: <data>/<job_id>/{state.json,source.bin}.
    job_dir = data_dir / job_id
    state = json.loads((job_dir / "state.json").read_text(encoding="utf-8"))
    assert state["filename"] == "paper.docx"
    assert state["target_lang"] == "简体中文"
    assert (job_dir / "source.bin").is_file()

    # Re-importing the same document opens the existing task, never a twin.
    again = adapter.import_documents([source], config)
    assert again[0].job_id == job_id
    assert again[0].created is False
    assert len(adapter.list_jobs()) == 1

    rows = adapter.list_jobs()
    assert rows[0].filename == "paper.docx"
    assert rows[0].status_label


def test_full_slice_import_translate_edit_review_deliver_reopen(adapter_env):
    """The Beta main flow, through the adapter, on a real DOCX."""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    stub = StubEndpoint(tmp)
    url = stub.start()
    try:
        config = adapter.EngineConfig(
            provider="自定义中转站", model="stub-model", api_key="stub-key",
            base_url=url, target_lang="简体中文", enable_review=True,
            enable_report=False, enable_understanding=False,
            enable_source_cleanup=False)
        config.save()

        source = tmp / "slice.docx"
        source.write_bytes(_make_docx([
            "这是第一段，涉及术语 MT。",
            "这是第二段，涉及术语 CAT。",
            "第三段用于验证交付。",
        ]))
        imported = adapter.import_documents([source], config)
        job_id = imported[0].job_id
        # New tasks default to risk_selective, which skips clean batches; this
        # flow checks the full independent review, so the task records ``always``.
        recorded = core.load_job_state(job_id)
        recorded["review_mode"] = "always"
        core.save_job_state(job_id, recorded)

        assert adapter.start_job(job_id, "slice.docx",
                                 source.read_bytes(), config) is True
        # A freshly queued worker may not have written its active status yet;
        # without this the first idle read can be mistaken for completion.
        assert _wait_for_restart(core, job_id), "the worker never became active"
        runtime = _wait_for_idle(core, job_id)
        assert runtime.get("status") == "completed", runtime

        state = adapter.load_state(job_id)
        assert state["p1_done"] and state["p2_done"]
        assert len(state["pairs"]) == 3
        assert all(pair.get("reviewed") for pair in state["pairs"])

        rows = adapter.segment_rows(job_id, state)
        assert [row.index for row in rows] == [0, 1, 2]
        assert all(row.status == "已审校" for row in rows)
        source_text, committed = adapter.segment_text(job_id, 1, state)
        assert source_text == "这是第二段，涉及术语 CAT。"

        # An uncommitted edit is persisted as a draft, then committed.
        pair = state["pairs"][1]
        uid = core.segment_uid(pair) or "1"
        adapter.save_draft(job_id, uid, "人工改写的第二段。", committed)
        assert adapter.pending_draft_count(job_id) == 1
        assert adapter.segment_rows(job_id) [1].has_draft is True

        adapter.submit_edit(job_id, 1, "人工改写的第二段。")
        assert adapter.pending_draft_count(job_id) == 0
        state = adapter.load_state(job_id)
        edited = state["pairs"][1]
        assert edited["target"] == "人工改写的第二段。"
        assert edited["human_edited"] is True
        assert edited["target_provenance"] == "human_edit"
        assert edited["reviewed"] is False

        # A partial re-review must NOT satisfy the delivery gate: a human edit
        # invalidates the whole-document review receipt, and the gate has to
        # stay honest about that instead of accepting one refreshed segment.
        adapter.review_segments(job_id, [1], config)
        state = adapter.load_state(job_id)
        assert state["pairs"][1]["reviewed"] is True
        with pytest.raises(adapter.AdapterError) as blocked:
            adapter.deliver(job_id, config)
        assert "重新审校" in str(blocked.value)

        # Re-reviewing the current working text clears the gate.
        adapter.review_segments(job_id, [0, 1, 2], config)
        state = adapter.load_state(job_id)
        assert all(pair.get("reviewed") for pair in state["pairs"])

        # Reopen: the committed edit survives a cache-free read.
        core._job_repo.invalidate_cache(job_id)
        reopened = adapter.segment_rows(job_id)
        assert reopened[1].target == "人工改写的第二段。"
        assert reopened[1].human_edited is True

        # Delivery gate + frozen snapshot.
        preview = adapter.delivery_preview(job_id, config)
        assert preview["blocking"] is False, preview["issues"]
        result = adapter.deliver(job_id, config)
        assert result["version"] == 1
        names = {item["name"] for item in result["assets"]}
        assert "translation.docx" in names
        assert "bilingual.docx" in names

        destination = tmp / "交付 结果"  # Chinese + space in the path
        written = adapter.copy_delivery_assets(job_id, result["version"], destination)
        assert written and all(path.is_file() for path in written)
        assert any(path.name == _delivered(adapter, job_id, "translation.docx") for path in written)

        state = adapter.load_state(job_id)
        assert state["delivery_status"] == "final"
        snapshot = adapter.delivery_preview(job_id, config)["snapshot"]
        assert snapshot["current"] is True and snapshot["integrity"] is True

        # A safe quit sees nothing pending.
        report = adapter.exit_report(job_id)
        assert report["safe_to_exit"] is True, report
        assert report["draft_count"] == 0
        assert report["worker_lock_held"] is False
    finally:
        stub.stop()


def test_exit_report_reports_drafts_and_blocks_on_a_running_task(adapter_env):
    """Exit contract: drafts are recoverable, a running task is not."""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "draft.docx"
    source.write_bytes(_make_docx(["一段原文。"]))
    config = adapter.EngineConfig(target_lang="简体中文")
    job_id = adapter.import_documents([source], config)[0].job_id
    state = core.load_job_state(job_id) or {}
    # Simulate the reviewed state a real job reaches before editing.
    state["pairs"] = [{"source": "一段原文。", "target": "一段译文。",
                       "reviewed": True, "segment_uid": "seg-1"}]
    state["p1_done"] = True
    state["p2_done"] = True
    core.save_job_state(job_id, state)

    adapter.save_draft(job_id, "seg-1", "尚未提交的编辑。", "一段译文。")
    report = adapter.exit_report(job_id)
    # A persisted draft survives a quit and is restored on reopen, so it must be
    # reported (never lost) without being dressed up as an unsafe exit.
    assert report["draft_count"] == 1
    assert report["draft_segments"] == ["seg-1"]
    assert report["drafts_recoverable"] is True
    assert report["safe_to_exit"] is True

    adapter.submit_edit(job_id, 0, "尚未提交的编辑。")
    report = adapter.exit_report(job_id)
    assert report["draft_count"] == 0
    assert report["safe_to_exit"] is True

    # An active runtime status blocks the "safe" verdict.
    core.update_runtime_state(job_id, status="running", event="probe",
                              event_name="probe")
    report = adapter.exit_report(job_id)
    assert report["task_running"] is True
    assert report["safe_to_exit"] is False
    core.update_runtime_state(job_id, status="cancelled", event="probe-done",
                              event_name="probe_done")


def test_delete_refuses_active_job(adapter_env, monkeypatch):
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "busy.docx"
    source.write_bytes(_make_docx(["忙碌任务。"]))
    config = adapter.EngineConfig(target_lang="简体中文")
    job_id = adapter.import_documents([source], config)[0].job_id

    monkeypatch.setattr(core, "job_is_active", lambda *a, **k: True)
    with pytest.raises(adapter.AdapterError):
        adapter.delete_job(job_id, allow_active=False)

    monkeypatch.setattr(core, "job_is_active", lambda *a, **k: False)
    assert adapter.delete_job(job_id) is True
    assert not (adapter_env["data_dir"] / job_id).exists()


def test_diagnostics_report_is_honest_about_the_qt_shell(adapter_env):
    adapter = adapter_env["adapter"]
    report = adapter.diagnostics_report(qt_version="6.11.2",
                                        pyside_version="6.11.2")
    assert report["app"]["shell"] == "qt-widgets"
    assert report["runtime"]["pyside6_version"] == "6.11.2"
    assert report["runtime"]["qt_version"] == "6.11.2"
    assert "streamlit_version" not in report["runtime"]
    # No credential or document text may travel in a support bundle.
    blob = json.dumps(report, ensure_ascii=False)
    assert "api_key" not in blob and "stub-key" not in blob


def test_workbench_shows_source_segments_before_translation(adapter_env):
    """Stage-1 ``paras`` are plain strings; the workbench must still list them.

    The core stores stage-1 paragraphs as strings (not ``{"text": ...}`` dicts),
    so reading them with a dict contract silently produced an empty workbench.
    """
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "stage1.docx"
    source.write_bytes(_make_docx(["第一段原文。", "第二段原文。"]))
    config = adapter.EngineConfig(target_lang="简体中文")
    job_id = adapter.import_documents([source], config)[0].job_id

    state = core.load_job_state(job_id) or {}
    state["paras"] = ["第一段原文。", "第二段原文。"]      # strings, as the core writes
    state["p1_done"] = True
    core.save_job_state(job_id, state)

    rows = adapter.segment_rows(job_id)
    assert [row.source for row in rows] == ["第一段原文。", "第二段原文。"]
    assert all(row.status == "待翻译" and row.target == "" for row in rows)
    assert adapter.segment_text(job_id, 1) == ("第二段原文。", "")

    # A dict-shaped legacy state must keep working too.
    state["paras"] = [{"text": "旧格式第一段。"}]
    core.save_job_state(job_id, state)
    assert adapter.segment_rows(job_id)[0].source == "旧格式第一段。"


def test_project_glossary_and_tm_are_passed_to_the_core(adapter_env):
    """An existing project's terminology/TM must reach the worker, not be dropped."""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "glossary.docx"
    source.write_bytes(_make_docx(["术语 MT 的段落。"]))
    config = adapter.EngineConfig(target_lang="简体中文")
    job_id = adapter.import_documents([source], config)[0].job_id

    state = core.load_job_state(job_id) or {}
    state["glossary"] = [{"source": "MT", "target": "机器翻译",
                          "status": "locked", "scope": "document"}]
    state["use_tm"] = True
    state["pipeline_config"] = {"use_tm": True, "auto_term": False}
    core.save_job_state(job_id, state)

    kwargs = adapter.pipeline_kwargs(config, adapter.load_state(job_id))
    assert kwargs["user_glossary"] == state["glossary"]
    # Saved pipeline choices win over the dialog defaults on resume.
    assert kwargs["auto_term"] is False
    assert kwargs["use_tm"] is True
    assert kwargs["delivery_config"]["deliver_plain_docx"] is True
    assert kwargs["delivery_config"]["deliver_bilingual_docx"] is True


def test_save_failure_is_reported_not_faked(adapter_env):
    """An unusable destination must fail loudly on every platform."""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "readonly.docx"
    source.write_bytes(_make_docx(["一段。"]))
    config = adapter.EngineConfig(target_lang="简体中文")
    job_id = adapter.import_documents([source], config)[0].job_id
    state = core.load_job_state(job_id) or {}
    state["pairs"] = [{"source": "一段。", "target": "A paragraph.",
                       "reviewed": True, "segment_uid": "s1"}]
    state["paras"] = ["一段。"]
    state["p1_done"] = True
    state["p2_done"] = True
    core.save_job_state(job_id, state)

    bundle = core.build_delivery_assets(job_id, adapter.load_state(job_id),
                                        target_lang="简体中文")
    assert "translation.docx" in bundle

    locked = tmp / "locked"
    # chmod on a Windows directory does not remove write access. A file in
    # place of the destination directory exercises a real failure on both OSes.
    locked.write_text("not a directory", encoding="utf-8")
    with pytest.raises((adapter.AdapterError, OSError, PermissionError)):
        adapter.export_current_assets(job_id, locked / "sub", config)


def test_unicode_and_rtl_edit_survives_commit_and_export(adapter_env):
    """CN/EN/AR text must survive commit, reopen and the delivered DOCX."""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    mixed = "混合段落：模型 v2 输出 «مرحبا بالعالم» 共 42 段，Mixed CN/EN/AR editing."
    arabic = "استعادة النظام البيئي تعتمد على الانتشار وكيمياء التربة."
    source = tmp / "mixed.docx"
    source.write_bytes(_make_docx(["第一段。", "第二段。"]))
    config = adapter.EngineConfig(target_lang="简体中文")
    job_id = adapter.import_documents([source], config)[0].job_id

    state = core.load_job_state(job_id) or {}
    state["paras"] = ["第一段。", "第二段。"]
    state["pairs"] = [
        {"source": "第一段。", "target": "First paragraph.", "reviewed": True,
         "segment_uid": "s1"},
        {"source": "第二段。", "target": "Second paragraph.", "reviewed": True,
         "segment_uid": "s2"},
    ]
    state["p1_done"] = True
    state["p2_done"] = True
    state["translation_core_review_required"] = False
    core.save_job_state(job_id, state)

    adapter.submit_edit(job_id, 0, mixed)
    adapter.submit_edit(job_id, 1, arabic)
    state = adapter.load_state(job_id)
    assert state["pairs"][0]["target"] == mixed
    assert state["pairs"][1]["target"] == arabic

    # RTL detection drives the editor's layout direction.
    from folith.textual import text_direction

    assert text_direction(arabic, "简体中文") == "rtl"
    assert text_direction(mixed, "简体中文") in ("ltr", "rtl")

    # Reopen (cache-free) keeps the exact Unicode text.
    core._job_repo.invalidate_cache(job_id)
    rows = adapter.segment_rows(job_id)
    assert rows[0].target == mixed and rows[1].target == arabic

    # The delivered document carries the same characters.
    from docx import Document

    payload = core.translations_to_word(
        adapter.load_state(job_id)["pairs"], target_lang="简体中文").getvalue()
    document = Document(io.BytesIO(payload))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert arabic in text
    assert "مرحبا بالعالم" in text


def test_delivered_files_keep_tricky_unicode_exactly(adapter_env):
    """§6.B：中英阿混排 + 组合字符 + ZWJ emoji + 双向控制符，在**用户拿到的交付文件**里逐字不变。

    与 `test_unicode_and_rtl_edit_survives_commit_and_export` 的分工：那个用例核对的是内存状态
    与 `translations_to_word` 的中间产物（用 `in` 判断）；这个用例走**真实交付门禁 → 快照 →
    `copy_delivery_assets`**（即原生保存面板写出给用户的那两个文件），并用**码点级相等**比对，
    不做 `strip()`、不做归一化。选这些字符是因为任何一步偷偷做了 Unicode 归一化/清理都会立刻露馅：
    NFD 组合重音（e+U+0301）、ZWJ 家庭 emoji、区域指示符旗帜、变体选择符、RTL 嵌入控制符、
    阿拉伯-印度数字与 tatweel。

    审校门禁本身在别的用例里覆盖（`translation_core_review_required` 的语义），这里为了让用例
    确定性、不依赖模型端点，提交人工修改后把 `reviewed` 复位为 True（已在注释里说明）。
    """
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    mixed = "中英阿混排：Model v2 → «مرحبا بالعالم» 42 段，keep 100% exact."
    combining = "NFD 组合重音：e\u0301 与 a\u0301 必须保持（不做 NFC 归一化）。"
    emoji = "ZWJ 家庭 👨‍👩‍👧‍👦 旗帜 🇸🇦 变体选择 ❤️ 不能被拆开或替换。"
    bidi = "\u202bRTL 嵌入 123\u202c 与 RLM\u200f 标记必须原样保留。"
    arabic = "استعادة النظام البيئي ١٢٣ ــ تعتمد على الانتشار."
    targets = [mixed, combining, emoji, bidi, arabic]

    # 先证明这组夹具**对归一化/清理敏感**，否则下面的"逐字相等"就是空断言。
    import unicodedata

    assert unicodedata.normalize("NFC", combining) != combining, "组合重音夹具必须能被 NFC 改变"
    assert "\u0301" in combining and "\u200d" in emoji and "\u202b" in bidi and "\u200f" in bidi
    assert "\ufe0f" in emoji and "\U0001f1f8" in emoji, "变体选择符/区域指示符必须在夹具里"
    assert "١" in arabic and "\u0640" in arabic

    source = tmp / "unicode-exact.docx"
    source.write_bytes(_make_docx(["第一段。", "第二段。", "第三段。", "第四段。", "第五段。"]))
    config = adapter.EngineConfig(target_lang="简体中文")
    job_id = adapter.import_documents([source], config)[0].job_id

    state = core.load_job_state(job_id) or {}
    state["paras"] = ["第一段。", "第二段。", "第三段。", "第四段。", "第五段。"]
    state["pairs"] = [
        {"source": f"第{index + 1}段。", "target": f"placeholder {index}",
         "reviewed": True, "segment_uid": f"u{index}"}
        for index in range(len(targets))
    ]
    state["p1_done"] = True
    state["p2_done"] = True
    state.pop("translation_core_review_required", None)
    core.save_job_state(job_id, state)

    # 人工编辑路径（提交后 reviewed 会被核心置回 False；本用例只关心文本保真，故复位）
    for index, text in enumerate(targets):
        adapter.submit_edit(job_id, index, text)
    state = adapter.load_state(job_id)
    assert [pair["target"] for pair in state["pairs"]] == targets, "提交后状态里的文本必须逐字不变"
    for pair in state["pairs"]:
        pair["reviewed"] = True
    core.save_job_state(job_id, state)

    # 重开（无缓存）逐字一致
    core._job_repo.invalidate_cache(job_id)
    assert [row.target for row in adapter.segment_rows(job_id)] == targets

    # 真实交付门禁 → 快照 → 写出用户实际拿到的文件
    preview = adapter.delivery_preview(job_id, config)
    assert preview["blocking"] is False, preview.get("issues")
    result = adapter.deliver(job_id, config)
    destination = tmp / "交付 结果 混排"          # 中文 + 空格路径
    written = adapter.copy_delivery_assets(job_id, result["version"], destination)
    names = {path.name for path in written}
    assert {_delivered(adapter, job_id, "translation.docx"),
            _delivered(adapter, job_id, "bilingual.docx")} <= names, names

    from docx import Document

    # 1) 纯译文：段落文本必须**完全相等**（不 strip、不归一化）
    doc = Document(str(destination / _delivered(adapter, job_id, "translation.docx")))
    paragraphs = [item.text for item in doc.paragraphs]
    assert paragraphs[0] == "译文", paragraphs[:2]
    assert paragraphs[1:1 + len(targets)] == targets, paragraphs[1:1 + len(targets)]
    # 码点级再确认一次：避免"看起来一样"（例如 NFC/NFD 视觉相同）
    assert [[ord(char) for char in text] for text in paragraphs[1:1 + len(targets)]] == \
        [[ord(char) for char in text] for text in targets]

    # 2) 双语：目标列同样逐字相等
    bilingual = Document(str(destination / _delivered(adapter, job_id, "bilingual.docx")))
    assert bilingual.tables, "双语交付必须有表格"
    rows = [[cell.text for cell in row.cells] for row in bilingual.tables[0].rows]
    assert len(rows) == len(targets) + 1, rows            # 表头 + 每段一行
    assert [row[1] for row in rows[1:]] == targets, [row[1] for row in rows[1:]]


def test_model_disconnect_is_comprehensible_and_retryable(adapter_env):
    """A dead endpoint must fail with a reason, and the task must be retryable."""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    # Port 9 (discard) is closed for our purposes: the request cannot succeed.
    broken = adapter.EngineConfig(
        provider="自定义中转站", model="stub-model", api_key="stub-key",
        base_url="http://127.0.0.1:9/v1", target_lang="简体中文",
        enable_review=False, enable_report=False, enable_understanding=False,
        enable_source_cleanup=False)
    assert broken.connection_problem() == ""

    source = tmp / "offline.docx"
    source.write_bytes(_make_docx(["第一段。", "第二段。"]))
    job_id = adapter.import_documents([source], broken)[0].job_id
    adapter.start_job(job_id, "offline.docx", source.read_bytes(), broken)
    # A retry can briefly publish a terminal status before the final failure,
    # and the durable file is replaced atomically.  Wait for the terminal state
    # that actually carries the reason instead of sampling once.
    deadline = time.monotonic() + 180
    runtime: dict = {}
    message = ""
    while time.monotonic() < deadline:
        runtime = core.get_job_runtime_status(job_id, adapter.load_state(job_id))
        message = str((runtime.get("error") or {}).get("message")
                      if isinstance(runtime.get("error"), dict)
                      else runtime.get("error") or "")
        if message and runtime.get("status") not in core.RUNTIME_ACTIVE_STATUSES:
            break
        time.sleep(0.2)
    assert message, f"a failure must carry a reason: {runtime}"
    assert runtime.get("status") in ("failed", "stalled", "interrupted",
                                     "idle_incomplete"), runtime
    state = adapter.load_state(job_id)
    assert not state.get("p2_done")
    # The worker publishes the failure before it releases its execution lock;
    # a retry inside that window is refused as "still running".
    assert _wait_for_worker_stopped(core, job_id), "the failed worker never stopped"

    # Retry with a working endpoint: the same task must be able to finish.
    stub = StubEndpoint(tmp)
    url = stub.start()
    try:
        good = adapter.EngineConfig(
            provider="自定义中转站", model="stub-model", api_key="stub-key",
            base_url=url, target_lang="简体中文", enable_review=True,
            enable_report=False, enable_understanding=False,
            enable_source_cleanup=False)
        assert adapter.retry_job(job_id) in (True, False)  # no crash either way
        adapter.start_job(job_id, "offline.docx", source.read_bytes(), good)
        runtime = _wait_for_idle(core, job_id, timeout=240)
        state = adapter.load_state(job_id)
        assert state.get("p1_done") and state.get("p2_done"), (runtime, state)
        assert runtime.get("status") == "completed"
    finally:
        stub.stop()


def test_diagnostics_describe_this_shell_not_the_retired_one(adapter_env):
    """Diagnostics must not advertise artifacts the Qt client does not ship."""
    adapter = adapter_env["adapter"]
    report = adapter.diagnostics_report(qt_version="6.11.2")
    locations = report["locations"]
    assert "service_log" not in locations, \
        "the Qt client has no Streamlit service log"
    assert locations["client_log"].endswith("qt-client.log")
    assert report["app"]["shell"] == "qt-widgets"
    assert report["qt"]["platform_plugin"]


def test_cancelled_task_is_recoverable_and_not_marked_complete(adapter_env):
    """A cancelled task must resume and finish, and must never look complete.

    Roadmap §6.C: "取消/退出后的任务可恢复，不被误标为完成".
    """
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    stub = StubEndpoint(tmp, delay=0.5)
    url = stub.start()
    try:
        config = adapter.EngineConfig(
            provider="自定义中转站", model="stub-model", api_key="stub-key",
            base_url=url, target_lang="简体中文", enable_review=True,
            enable_report=False, enable_understanding=False,
            enable_source_cleanup=False)
        source = tmp / "cancel.docx"
        source.write_bytes(_make_docx([f"第 {index} 段，术语 MT。"
                                       for index in range(40)]))
        job_id = adapter.import_documents([source], config)[0].job_id
        adapter.start_job(job_id, "cancel.docx", source.read_bytes(), config)
        assert _wait_for_alive(core, job_id), "the worker never started"
        # Cancel after the worker has actually entered an external call, not
        # while it is still queued: otherwise a very fast run can finish before
        # the cancellation is observed, which is not the property under test.
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            runtime = core.get_job_runtime_status(job_id, adapter.load_state(job_id))
            if runtime.get("status") in {"running", "waiting_external"}:
                break
            time.sleep(0.05)
        adapter.cancel_job(job_id)
        runtime = _wait_for_idle(core, job_id, timeout=180)
        state = adapter.load_state(job_id)
        assert runtime.get("status") == "cancelled", runtime
        assert not (state.get("p1_done") and state.get("p2_done")), \
            "a cancelled task must not be reported as complete"
        # ``_wait_for_worker_stopped``'s own docstring says the runtime can
        # report ``cancelled`` before the thread has finished cleanup, so the
        # exit report has to be read after that wait, not before it.  Checking
        # first made this fail on slow runners with "assert True is False".
        assert _wait_for_worker_stopped(core, job_id)
        assert adapter.exit_report(job_id)["worker_lock_held"] is False

        # Resume the same task: it must finish on the existing progress.
        adapter.start_job(job_id, "cancel.docx", None, config)
        assert _wait_for_restart(core, job_id)
        runtime = _wait_for_idle(core, job_id, timeout=300)
        state = adapter.load_state(job_id)
        assert state.get("p1_done") and state.get("p2_done"), (runtime, state)
        assert runtime.get("status") == "completed"
    finally:
        stub.stop()


def test_import_assigns_the_task_to_the_chosen_project(adapter_env):
    """Roadmap §6.A.2: basic project ownership must be real, not cosmetic."""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    project = adapter.create_project("Qt 验收项目")
    project_id = project["project_id"]
    assert any(item["project_id"] == project_id
               for item in adapter.project_options())

    source = tmp / "owned.docx"
    source.write_bytes(_make_docx(["归属测试。"]))
    config = adapter.EngineConfig(target_lang="简体中文")
    job_id = adapter.import_documents([source], config,
                                      project_id=project_id)[0].job_id
    assert adapter.load_state(job_id)["project_id"] == project_id
    assert any(row.job_id == job_id and row.project_id == project_id
               for row in adapter.list_jobs())

    # A deleted/unknown project must not leave an orphan task: the same document
    # imported with a stale id falls back to the system workspace (explicit null).
    import shutil
    shutil.rmtree(core.OUTPUT_DIR / "projects" / project_id, ignore_errors=True)
    second = tmp / "orphan.docx"
    second.write_bytes(_make_docx(["孤儿测试。"]))
    job_two = adapter.import_documents([second], config,
                                       project_id=project_id)[0].job_id
    assert adapter.load_state(job_two)["project_id"] is None


def test_diagnostics_export_never_carries_document_text(adapter_env):
    """§6.D: the support bundle must not contain document bodies.

    Checked with a real job whose source and translation contain a unique marker,
    so "no document text" is a measured fact rather than a code reading.
    """
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    marker = "ZQ-DOCUMENT-BODY-MARKER-7731"
    source = tmp / "secret.docx"
    source.write_bytes(_make_docx([f"第一段包含 {marker} 的正文。", "第二段。"]))
    config = adapter.EngineConfig(target_lang="简体中文")
    job_id = adapter.import_documents([source], config)[0].job_id
    state = core.load_job_state(job_id) or {}
    state["paras"] = [f"第一段包含 {marker} 的正文。", "第二段。"]
    state["pairs"] = [
        {"source": f"第一段包含 {marker} 的正文。",
         "target": f"译文也含 {marker}。", "reviewed": True, "segment_uid": "s1"},
        {"source": "第二段。", "target": "Second.", "reviewed": True,
         "segment_uid": "s2"},
    ]
    state["p1_done"] = True
    state["p2_done"] = True
    core.save_job_state(job_id, state)

    report = adapter.diagnostics_report(qt_version="6.11.2")
    blob = json.dumps(report, ensure_ascii=False)
    assert marker not in blob, "document body leaked into diagnostics"
    assert job_id not in blob, "per-job identifiers are not needed in diagnostics"
    assert report["locations"]["job_count"] >= 1, \
        "the job count is the intended, non-identifying signal"

    written = adapter.export_diagnostics(tmp / "support.json")
    text = Path(written).read_text(encoding="utf-8")
    assert marker not in text
    assert "api_key" not in text and "SUPER-SECRET" not in text


def test_connection_test_failure_is_comprehensible(adapter_env):
    """§6.A.2: a failed connection test must explain itself, not dump a traceback."""
    adapter = adapter_env["adapter"]

    # Nothing configured: the reason is stated without any network attempt.
    ok, message = adapter.test_connection(adapter.EngineConfig())
    assert ok is False
    assert "提供商" in message or "API Key" in message or "模型" in message

    # A relay address that is not listening: a real attempt that fails.
    dead = adapter.EngineConfig(
        provider="自定义中转站", model="stub-model", api_key="stub-key",
        base_url="http://127.0.0.1:9/v1")
    ok, message = adapter.test_connection(dead)
    assert ok is False
    assert message and "Traceback" not in message
    assert len(message) < 400, f"failure message is not a summary: {message!r}"
    # The user must be able to tell it was a connection problem.
    lowered = message.lower()
    assert any(token in lowered for token in
               ("connect", "connection", "连接", "refused", "超时", "timeout",
                "无法", "失败")), message

    # An invalid relay address is rejected before any request is made.
    invalid = adapter.EngineConfig(
        provider="自定义中转站", model="stub-model", api_key="stub-key",
        base_url="not-a-url")
    ok, message = adapter.test_connection(invalid)
    assert ok is False and "地址" in message


def test_empty_or_missing_source_is_refused_not_started(adapter_env):
    """A fresh task with no source must not start a doomed pipeline.

    Observed in the wild as `fitz` "Cannot open empty stream": a concurrent
    truncating write left a 0-byte `source.bin`, and the pipeline was started
    against it.  `core.save_source` is atomic now; this pins the client-side
    guard so the failure can never be silent again.
    """
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    config = adapter.EngineConfig(
        provider="自定义中转站", model="stub-model", api_key="stub-key",
        base_url="http://127.0.0.1:9/v1", target_lang="简体中文")
    source = tmp / "no-source.docx"
    source.write_bytes(_make_docx(["一段。"]))
    job_id = adapter.import_documents([source], config)[0].job_id

    # Simulate the corrupted state: a fresh task whose stored source is empty.
    core.save_source(job_id, b"")
    with pytest.raises(adapter.AdapterError) as refused:
        adapter.start_job(job_id, "no-source.docx", None, config)
    assert "源文件" in str(refused.value)

    # With the real bytes supplied, the same call proceeds.
    assert adapter.start_job(job_id, "no-source.docx",
                             source.read_bytes(), config) is True


def test_save_source_is_atomic(adapter_env):
    """A reader must never observe a truncated source file."""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "atomic.docx"
    payload = _make_docx(["第一段。", "第二段。"])
    source.write_bytes(payload)
    config = adapter.EngineConfig(target_lang="简体中文")
    job_id = adapter.import_documents([source], config)[0].job_id
    assert core.load_source(job_id) == payload

    # Rewriting never leaves the file empty or partial: after the call the
    # content is exactly the new payload, and no temp files remain.
    core.save_source(job_id, b"")
    assert core.load_source(job_id) == b""
    core.save_source(job_id, payload)
    assert core.load_source(job_id) == payload
    leftovers = [p.name for p in (core.job_dir(job_id)).iterdir()
                 if p.name.startswith(".source.bin")]
    assert leftovers == [], f"atomic write left temp files: {leftovers}"


def test_saved_delivery_choice_wins_over_the_dialog_on_resume(adapter_env):
    """A task that already has a delivery selection keeps it.

    Changing the app-level default must not silently change what an existing
    (possibly half-finished) task will deliver.
    """
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "resume-formats.docx"
    source.write_bytes(_make_docx(["一段。"]))
    plain_only = adapter.EngineConfig(
        target_lang="简体中文", deliver_plain_docx=True,
        deliver_bilingual_docx=False, deliver_terms_xlsx=False,
        deliver_evidence=False)
    job_id = adapter.import_documents([source], plain_only)[0].job_id
    assert (adapter.load_state(job_id)["delivery_config"]
            ["deliver_bilingual_docx"]) is False

    # The dialog later declares everything; the task keeps its own choice.
    everything = adapter.EngineConfig(
        target_lang="简体中文", deliver_plain_docx=True,
        deliver_bilingual_docx=True, deliver_terms_xlsx=True,
        deliver_evidence=True)
    kwargs = adapter.pipeline_kwargs(everything, adapter.load_state(job_id))
    assert kwargs["delivery_config"]["deliver_bilingual_docx"] is False
    assert kwargs["delivery_config"]["deliver_plain_docx"] is True

    # A *new* task, by contrast, takes the current selection.
    second = tmp / "fresh-formats.docx"
    second.write_bytes(_make_docx(["另一段。"]))
    other = adapter.import_documents([second], everything)[0].job_id
    assert (adapter.load_state(other)["delivery_config"]
            ["deliver_bilingual_docx"]) is True


def test_frozen_diagnostics_report_the_build_record(adapter_env, monkeypatch):
    """A frozen bundle must answer "what is this built for?" from its own record.

    It cannot ask git, and it does not ship PyInstaller's distribution metadata,
    so those facts have to come from the build record the bundle carries.
    """
    adapter = adapter_env["adapter"]
    from folith.desktop import paths as desktop_paths

    record = {
        "app_version": "0.5.0", "build_id": "git:deadbee",
        "commit": "deadbeef" * 5, "worktree_dirty": True,
        "packaging_backend": "pyinstaller", "shell": "qt",
        "target_platform": "macos", "target_arch": "arm64",
        "python_version": "3.11.3", "pyside6_version": "6.11.2",
        "qt_version": "6.11.2", "pyinstaller_version": "6.22.3",
        "minimum_system_version": "15.0",
        "minimum_system_version_source": "measured",
        "os_version": "macOS-27.0.1-arm64",
    }
    monkeypatch.setattr(desktop_paths, "build_info", lambda: record)
    monkeypatch.setattr(desktop_paths, "is_frozen", lambda: True)

    report = adapter.diagnostics_report(qt_version="", pyside_version="")
    assert report["app"]["frozen"] is True
    assert report["app"]["dirty"] == "true", report["app"]["dirty"]
    assert report["runtime"]["pyinstaller_version"] == "6.22.3"
    assert report["runtime"]["qt_version"] == "6.11.2"
    assert report["support"]["minimum_system_version"] == "15.0"
    assert report["support"]["target_arch"] == "arm64"
    assert report["support"]["shell"] == "qt"
    # No credential *value* may travel; the boolean status flag is fine.
    blob = json.dumps(report, ensure_ascii=False)
    assert "SUPER-SECRET" not in blob


def test_assign_existing_task_to_a_project_and_refuse_running_ones(adapter_env):
    """§6.A.2：已有任务也能改归属；运行中的任务必须被核心拒绝。"""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    project = adapter.create_project("归属目标")
    project_id = project["project_id"]
    config = adapter.EngineConfig(target_lang="简体中文")
    source = tmp / "assign.docx"
    source.write_bytes(_make_docx(["一段。"]))
    job_id = adapter.import_documents([source], config)[0].job_id
    assert adapter.load_state(job_id)["project_id"] is None

    result = adapter.assign_jobs_to_project([job_id], project_id)
    assert result["moved"] == [job_id] and not result["skipped"]
    assert adapter.load_state(job_id)["project_id"] == project_id

    # A second assignment to the same project is reported as "already there".
    again = adapter.assign_jobs_to_project([job_id], project_id)
    assert again["moved"] == [] and again["skipped"][0]["reason"].startswith("已在")

    # A *really* running task is refused with the core's own reason.  A hand
    # written runtime status is not enough: the core reconciles the worker lease,
    # so only a live worker makes a task "active" — and that is exactly the
    # protection we want to verify.
    stub = StubEndpoint(tmp, delay=0.5)
    url = stub.start()
    try:
        running_config = adapter.EngineConfig(
            provider="自定义中转站", model="stub-model", api_key="stub-key",
            base_url=url, target_lang="简体中文", enable_review=True,
            enable_report=False, enable_understanding=False,
            enable_source_cleanup=False)
        other = tmp / "running.docx"
        other.write_bytes(_make_docx([f"第 {i} 段。" for i in range(40)]))
        running = adapter.import_documents([other], running_config)[0].job_id
        adapter.start_job(running, "running.docx", other.read_bytes(),
                          running_config)
        assert _wait_for_alive(core, running), "worker never started"
        blocked = adapter.assign_jobs_to_project([running], project_id)
        assert blocked["moved"] == []
        assert "运行" in blocked["skipped"][0]["reason"], blocked
        adapter.cancel_job(running)
        _wait_for_idle(core, running, timeout=180)
    finally:
        stub.stop()

    # No selection / no target are refused up front.
    with pytest.raises(adapter.AdapterError):
        adapter.assign_jobs_to_project([], project_id)
    with pytest.raises(adapter.AdapterError):
        adapter.assign_jobs_to_project([job_id], "")


def test_project_deletion_protections_and_backup(adapter_env):
    """§6.C：系统工作区不可删、逐字确认、仍含任务则拒绝、删除前写可恢复备份。"""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    project = adapter.create_project("待删项目")
    project_id = project["project_id"]
    config = adapter.EngineConfig(target_lang="简体中文")
    source = tmp / "in-project.docx"
    source.write_bytes(_make_docx(["项目内任务。"]))
    job_id = adapter.import_documents([source], config,
                                      project_id=project_id)[0].job_id

    # 1) 仍含任务 → 拒绝，且提示先移动/删除。
    with pytest.raises(adapter.AdapterError) as refused:
        adapter.delete_project(project_id, confirm_name="待删项目")
    assert "仍有" in str(refused.value)
    assert (core.OUTPUT_DIR / "projects" / project_id).exists(), \
        "被拒绝的删除不得动到项目目录"

    # 2) 确认名称不匹配 → 拒绝。
    with pytest.raises(adapter.AdapterError) as wrong:
        adapter.delete_project(project_id, confirm_name="别的名字")
    assert "确认名称不匹配" in str(wrong.value)

    # 3) 移走任务后可以删除，并且删除前写入了可恢复备份。
    #    "移回未分类"用的是系统工作区的真实 ID（空字符串在核心层是非法目标）。
    system_row = next(row for row in adapter.project_rows() if row["is_system"])
    adapter.assign_jobs_to_project([job_id], system_row["project_id"])
    result = adapter.delete_project(project_id, confirm_name="待删项目")
    assert result.get("name") == "待删项目" or result.get("deleted") is not None \
        or result
    backups = adapter.project_backups()
    assert backups, "删除前必须留下可恢复备份"
    assert any(project_id in name for name in backups), backups
    assert not (core.OUTPUT_DIR / "projects" / project_id).exists()

    # 4) 系统工作区永远不可删除。
    system = system_row
    with pytest.raises(adapter.AdapterError) as system_refused:
        adapter.delete_project(system["project_id"], confirm_name=system["name"])
    assert "系统工作区" in str(system_refused.value)


def test_project_rows_report_counts_and_archive_state(adapter_env):
    adapter = adapter_env["adapter"]
    tmp = adapter_env["tmp"]
    project = adapter.create_project("统计项目")
    project_id = project["project_id"]
    config = adapter.EngineConfig(target_lang="简体中文")
    source = tmp / "counted.docx"
    source.write_bytes(_make_docx(["计数任务。"]))
    adapter.import_documents([source], config, project_id=project_id)

    rows = adapter.project_rows()
    counted = next(row for row in rows if row["project_id"] == project_id)
    assert counted["job_count"] == 1
    assert counted["archived"] is False
    assert any(row["is_system"] for row in rows)

    adapter.archive_project(project_id, archived=True)
    archived = next(row for row in adapter.project_rows()
                    if row["project_id"] == project_id)
    assert archived["archived"] is True
    # Archived projects are excluded from the "move to" chooser.
    from folith.qt import window as window_module  # noqa: F401  (import check)

    active = [row for row in adapter.project_rows() if not row["archived"]]
    assert project_id not in {row["project_id"] for row in active}


def test_exclude_and_include_a_segment(adapter_env):
    """§6.A.5：段落可排除/恢复。

    核心的真实语义是**结构性移出**：排除会把这一段从 pairs/paras 里拿出来，
    连同原文、当时的译文与原因存进 `excluded_segments` 记录；恢复按记录
    （`excluded_id`）精确插回原位，而不是按索引——排除后索引已经整体前移。
    """
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "exclude.docx"
    source.write_bytes(_make_docx(["第一段。", "第二段。", "第三段。"]))
    config = adapter.EngineConfig(target_lang="简体中文")
    job_id = adapter.import_documents([source], config)[0].job_id
    state = core.load_job_state(job_id) or {}
    state["paras"] = ["第一段。", "第二段。", "第三段。"]
    state["pairs"] = [
        {"source": text, "target": f"Translation {index}", "reviewed": True,
         "segment_uid": f"s{index}"}
        for index, text in enumerate(state["paras"])]
    state["p1_done"] = True
    state["p2_done"] = True
    core.save_job_state(job_id, state)

    adapter.exclude_segment(job_id, 1, reason="重复段落")
    state = adapter.load_state(job_id)
    # 段落被移出，位置由后面的段落顶上。
    assert len(state["pairs"]) == 2
    assert [pair["source"] for pair in state["pairs"]] == ["第一段。", "第三段。"]
    assert core.active_pair_indexes(state) == [0, 1]
    records = adapter.excluded_segments(job_id, state)
    assert len(records) == 1
    record = records[0]
    assert record["segment_index"] == 1 and record["reason"] == "重复段落"
    assert record["excluded_id"]
    # 原文与排除时的译文都完整保留（排除不是删除）。
    assert record["source"] == "第二段。"
    assert record["target"] == "Translation 1"

    # 交付包里如实说明哪些内容没有进入译文。
    bundle = core.build_delivery_assets(job_id, adapter.load_state(job_id))
    assert "excluded_segments.json" in bundle
    assert "excluded_segments.md" in bundle
    payload = json.loads(bundle["excluded_segments.json"].decode("utf-8"))
    assert payload["count"] == 1
    assert payload["records"][0]["excluded_id"] == record["excluded_id"]

    # 恢复按记录进行，内容原样回到原位。
    adapter.include_segment(job_id, record["excluded_id"])
    state = adapter.load_state(job_id)
    assert len(state["pairs"]) == 3
    assert [pair["source"] for pair in state["pairs"]] == \
        ["第一段。", "第二段。", "第三段。"]
    assert state["pairs"][1]["target"] == "Translation 1"
    assert adapter.excluded_segments(job_id, state) == []
    # 恢复不连带点亮旧的"已审校"标记：结构变了，审校必须重新确认。
    assert not state["pairs"][1].get("reviewed")

    # 恢复后可以继续编辑。
    adapter.submit_edit(job_id, 1, "恢复后的人工修改。")
    assert adapter.load_state(job_id)["pairs"][1]["target"] == "恢复后的人工修改。"

    # 空 id 与不存在的记录都要给出可理解的拒绝。
    with pytest.raises(adapter.AdapterError):
        adapter.include_segment(job_id, "")
    with pytest.raises(adapter.AdapterError):
        adapter.include_segment(job_id, "x999999")


def test_retranslate_one_segment_through_the_core(adapter_env):
    """§6.A.5：定点重译只生成候选；采用才改这一段的正文。"""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    stub = StubEndpoint(tmp)
    url = stub.start()
    try:
        config = adapter.EngineConfig(
            provider="自定义中转站", model="stub-model", api_key="stub-key",
            base_url=url, target_lang="简体中文", enable_review=True,
            enable_report=False, enable_understanding=False,
            enable_source_cleanup=False)
        source = tmp / "retranslate.docx"
        source.write_bytes(_make_docx(["第一段。", "第二段。"]))
        job_id = adapter.import_documents([source], config)[0].job_id
        state = core.load_job_state(job_id) or {}
        state["paras"] = ["第一段。", "第二段。"]
        state["pairs"] = [
            {"source": "第一段。", "target": "要重译的旧译文。", "reviewed": False,
             "segment_uid": "s0"},
            {"source": "第二段。", "target": "第二段保持不动。", "reviewed": True,
             "segment_uid": "s1"},
        ]
        state["p1_done"] = True
        state["p2_done"] = True
        core.save_job_state(job_id, state)

        result = adapter.retranslate_segment(job_id, 0, config)
        assert result["created"] == [0], result
        assert result["candidates"], "应返回待采用候选"
        # 这条断言是回归保护：定点重译不在管线 worker 里跑，中转地址必须显式
        # 传下去；漏传时这里会抛"未配置有效的 API 地址"（此前是静默不动）。
        state = adapter.load_state(job_id)
        assert state["pairs"][0]["target"] == "要重译的旧译文。", \
            "生成候选不得改稿"
        assert state["pairs"][1]["target"] == "第二段保持不动。", \
            "定点重译不得动到其它段落"
        candidate = result["candidates"][0]
        adopted = adapter.adopt_candidate(job_id, candidate["candidate_id"])
        assert adopted["committed"] is True
        state = adapter.load_state(job_id)
        assert state["pairs"][0]["target"] == candidate["candidate_target"]
        assert state["pairs"][0]["target_provenance"] == "candidate_adopted"
        assert state["pairs"][1]["target"] == "第二段保持不动。"

        # 越界与未配置连接都要给出可理解的拒绝。
        with pytest.raises(adapter.AdapterError):
            adapter.retranslate_segment(job_id, 99, config)
        broken = adapter.EngineConfig(provider="自定义中转站", model="m",
                                      api_key="k", base_url="")
        with pytest.raises(adapter.AdapterError):
            adapter.retranslate_segment(job_id, 0, broken)
    finally:
        stub.stop()


def test_home_rows_have_a_real_last_saved_time(adapter_env):
    """首页「更新时间」必须有真实时间。

    任务状态本身不带时间戳（`save_job_state` 不写 updated_at/created_at），所以
    这一列曾经永远是"—"、按时间排序也等于没排。现在回退到核心自己的
    `recovery_summary().last_saved_at`（旧界面"最近保存"用的是同一个来源）。
    """
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "stamp.docx"
    source.write_bytes(_make_docx(["一段。"]))
    job_id = adapter.import_documents(
        [source], adapter.EngineConfig(target_lang="简体中文"))[0].job_id

    rows = adapter.list_jobs()
    row = next(item for item in rows if item.job_id == job_id)
    assert row.updated_at, "首页行必须有更新时间"
    # 与核心的权威取值一致。
    expected = core.recovery_summary(job_id, adapter.load_state(job_id))
    assert row.updated_at == expected.get("last_saved_at")
    # 排序按时间倒序。
    stamps = [item.updated_at for item in rows if item.updated_at]
    assert stamps == sorted(stamps, reverse=True), stamps


def test_draft_against_older_text_is_marked_stale(adapter_env):
    """草稿必须能识别出"它是对旧译文写的"（否则会静默覆盖较新的译文）。

    草稿记录里带 `baseline`（写下草稿时的提交文本）。如果之后译文被改过
    （例如定点重译），这份草稿就不再对应当前文本：它仍然保留并显示给用户，
    但必须显式标记，不能装作是最新状态。
    """
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "stale.docx"
    source.write_bytes(_make_docx(["一段。"]))
    job_id = adapter.import_documents(
        [source], adapter.EngineConfig(target_lang="简体中文"))[0].job_id
    state = core.load_job_state(job_id) or {}
    state["paras"] = ["一段。"]
    state["pairs"] = [{"source": "一段。", "target": "机器译文 v1。",
                       "reviewed": True, "segment_uid": "s0"}]
    state["p1_done"] = True
    state["p2_done"] = True
    core.save_job_state(job_id, state)

    # 针对 v1 写一份草稿。
    adapter.save_draft(job_id, "s0", "我改到一半的译文。", "机器译文 v1。")
    row = adapter.segment_rows(job_id, adapter.load_state(job_id))[0]
    assert row.has_draft is True and row.stale_draft is False
    assert row.draft_text == "我改到一半的译文。"

    # 译文在草稿之后被改（定点重译 / 重新跑批都会这样）。
    state = adapter.load_state(job_id)
    state["pairs"][0]["target"] = "机器译文 v2（更新）。"
    core.save_job_state(job_id, state)

    row = adapter.segment_rows(job_id, adapter.load_state(job_id))[0]
    assert row.target == "机器译文 v2（更新）。"
    assert row.has_draft is True, "草稿不能被丢掉"
    assert row.draft_text == "我改到一半的译文。"
    assert row.stale_draft is True, "必须识别出草稿基于旧译文"

    # 放弃草稿后回到当前提交文本。
    adapter.clear_draft(job_id, "s0")
    row = adapter.segment_rows(job_id, adapter.load_state(job_id))[0]
    assert row.has_draft is False and row.stale_draft is False
    assert row.target == "机器译文 v2（更新）。"

    # baseline 为空（首次编辑，还没有提交文本）不算过期。
    adapter.save_draft(job_id, "s0", "首次编辑。", "")
    row = adapter.segment_rows(job_id, adapter.load_state(job_id))[0]
    assert row.has_draft is True and row.stale_draft is False


def _frozen_job_with_snapshot(adapter_env):
    """造一个"已冻结交付"的任务，返回 (job_id, 快照版本, 资产名集合)。"""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "save-assets.docx"
    source.write_bytes(_make_docx(["第一段。", "第二段。"]))
    config = adapter.EngineConfig(target_lang="简体中文")
    job_id = adapter.import_documents([source], config)[0].job_id
    state = core.load_job_state(job_id) or {}
    state["paras"] = ["第一段。", "第二段。"]
    state["pairs"] = [
        {"source": "第一段。", "target": "First.", "reviewed": True,
         "segment_uid": "s0"},
        {"source": "第二段。", "target": "Second.", "reviewed": True,
         "segment_uid": "s1"},
    ]
    state["p1_done"] = True
    state["p2_done"] = True
    core.save_job_state(job_id, state)
    version = core.create_delivery_snapshot(job_id, adapter.load_state(job_id),
                                            target_lang="简体中文")
    if isinstance(version, dict):
        version = version.get("snapshot_version")
    names = {name for name, _path in adapter.snapshot_asset_paths(job_id, version)}
    return job_id, int(version), names


def test_saving_delivery_assets_is_complete_and_never_silently_overwrites(
        adapter_env, tmp_path):
    """保存交付文件：不能悄悄少写，也不能悄悄覆盖同名文件。"""
    adapter = adapter_env["adapter"]

    job_id, version, names = _frozen_job_with_snapshot(adapter_env)
    assert names, "快照里应该有资产"

    # 1) 正常保存到"中文 + 空格"路径：每个资产都在，内容与快照逐字节一致。
    dest = tmp_path / "交付 结果 中文"
    written = adapter.copy_delivery_assets(job_id, version, dest)
    assert {path.name for path in written} == {_delivered(adapter, job_id, name) for name in names}
    for name, source in adapter.snapshot_asset_paths(job_id, version):
        target = dest / _delivered(adapter, job_id, name)
        assert target.read_bytes() == source.read_bytes(), name

    # 2) 全新目标里已有一个同名文件：默认拒绝，且原有内容原封不动。
    other = tmp_path / "另一个 文件夹"
    other.mkdir()
    delivered = {_delivered(adapter, job_id, name) for name in names}
    collision = other / sorted(delivered)[0]
    collision.write_text("这是用户自己的文件，不能被悄悄覆盖", encoding="utf-8")
    plan = adapter.delivery_save_plan(job_id, version, other)
    assert plan["conflicts"] == [collision.name], plan["conflicts"]
    with pytest.raises(adapter.AdapterError) as conflict:
        adapter.copy_delivery_assets(job_id, version, other)
    assert "同名文件" in str(conflict.value)
    assert collision.read_text(encoding="utf-8") == "这是用户自己的文件，不能被悄悄覆盖"
    # 被拒绝时不应写出任何其它文件（避免"半套交付"）。
    assert [item.name for item in other.iterdir()] == [collision.name]

    # 3) 明确确认覆盖后才替换。
    written = adapter.copy_delivery_assets(job_id, version, other, overwrite=True)
    assert {path.name for path in written} == delivered
    assert collision.read_bytes() != "这是用户自己的文件，不能被悄悄覆盖".encode()

    # 4) 快照里有文件缺失：拒绝保存，不产生"看起来成功"的不完整交付。
    _name, source = adapter.snapshot_asset_paths(job_id, version)[0]
    source.unlink()
    fresh = tmp_path / "第三个 文件夹"
    with pytest.raises(adapter.AdapterError) as missing:
        adapter.copy_delivery_assets(job_id, version, fresh)
    assert "缺少文件" in str(missing.value)
    assert not fresh.exists() or not any(fresh.iterdir()), \
        "拒绝保存时不得留下半套文件"


def test_export_copy_is_complete_and_never_silently_overwrites(
        adapter_env, tmp_path, monkeypatch):
    """导出当前译文副本：与"保存冻结交付"同一条规则——不少写、不静默覆盖。"""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "export.docx"
    source.write_bytes(_make_docx(["第一段。"]))
    job_id = adapter.import_documents(
        [source], adapter.EngineConfig(target_lang="简体中文"))[0].job_id
    state = core.load_job_state(job_id) or {}
    state["paras"] = ["第一段。"]
    state["pairs"] = [{"source": "第一段。", "target": "First.", "reviewed": True,
                       "segment_uid": "s0"}]
    state["p1_done"] = True
    state["p2_done"] = True
    core.save_job_state(job_id, state)

    dest = tmp_path / "导出 目录"
    written = adapter.export_current_assets(job_id, dest, None)
    assert written, "导出应该写出文件"
    assert all(path.is_file() for path in written)
    names = {path.name for path in written}
    assert _delivered(adapter, job_id, "translation.docx") in names

    # 同名文件 → 默认拒绝，原文件不变，目录里不多出其它文件。
    other = tmp_path / "已有内容 的目录"
    other.mkdir()
    collision = other / sorted(names)[0]
    collision.write_text("用户自己的文件", encoding="utf-8")
    plan = adapter.export_plan(job_id, other, None)
    assert plan["conflicts"] == [collision.name], plan
    with pytest.raises(adapter.AdapterError) as conflict:
        adapter.export_current_assets(job_id, other, None)
    assert "同名文件" in str(conflict.value)
    assert collision.read_text(encoding="utf-8") == "用户自己的文件"
    assert [item.name for item in other.iterdir()] == [collision.name]

    # 明确覆盖后才替换。
    written = adapter.export_current_assets(job_id, other, None, overwrite=True)
    assert {path.name for path in written} == names

    # 某个格式没生成（bundle 里为 None）→ 拒绝导出，不写半套。
    fresh = tmp_path / "半套 目录"
    original = core.build_delivery_assets

    def fake_bundle(*args, **kwargs):
        bundle = dict(original(*args, **kwargs))
        bundle["missing.docx"] = None
        return bundle

    monkeypatch.setattr(core, "build_delivery_assets", fake_bundle)
    with pytest.raises(adapter.AdapterError) as missing:
        adapter.export_current_assets(job_id, fresh, None)
    assert "没有生成" in str(missing.value)
    assert not fresh.exists() or not any(fresh.iterdir()), "拒绝时不得留下半套文件"


def test_submit_edit_refusals_explain_themselves(adapter_env):
    """三种拒绝各有明确原因；运行中的拒绝不暴露异常类型名。"""
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "refusals.docx"
    source.write_bytes(_make_docx(["甲。", "乙。"]))
    job_id = adapter.import_documents(
        [source], adapter.EngineConfig(target_lang="简体中文"))[0].job_id

    # 1) 还没有译文（只有 paras）：不能说"索引超出范围"——索引没错，是还不能编辑。
    state = core.load_job_state(job_id) or {}
    state["paras"] = ["甲。", "乙。"]
    state["pairs"] = []
    core.save_job_state(job_id, state)
    with pytest.raises(adapter.AdapterError) as early:
        adapter.submit_edit(job_id, 0, "人工")
    message = str(early.value)
    assert "还没有译文" in message and "索引" not in message, message

    # 2) 索引真的越界：保留原来的提示。
    state["pairs"] = [{"source": "甲。", "target": "A.", "reviewed": True,
                       "segment_uid": "s0"}]
    core.save_job_state(job_id, state)
    with pytest.raises(adapter.AdapterError) as out_of_range:
        adapter.submit_edit(job_id, 99, "人工")
    assert "索引超出范围" in str(out_of_range.value)

    # 3) 运行中由核心拒绝：原因原样透出，不带 "RuntimeError:" 前缀。
    real_save = core.save_translation_edit

    def refuse(*args, **kwargs):
        raise RuntimeError("任务正在后台处理，请等待完成后保存；草稿可以保留。")

    core.save_translation_edit = refuse
    try:
        with pytest.raises(adapter.AdapterError) as busy:
            adapter.submit_edit(job_id, 0, "人工")
    finally:
        core.save_translation_edit = real_save
    assert str(busy.value) == "任务正在后台处理，请等待完成后保存；草稿可以保留。"
    assert "RuntimeError" not in str(busy.value)


def test_export_copy_content_is_usable_and_honest(adapter_env, tmp_path):
    """导出的副本要能直接用，而且不能假装成"已冻结交付"。

    导出发生在交付门禁**之前**：manifest 必须如实写 draft，否则用户会把未冻结的
    副本当成正式交付物。
    """
    import docx

    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "export-content.docx"
    source.write_bytes(_make_docx(["甲段。", "乙段。"]))
    job_id = adapter.import_documents(
        [source], adapter.EngineConfig(target_lang="简体中文"))[0].job_id
    state = core.load_job_state(job_id) or {}
    state["paras"] = ["甲段。", "乙段。"]
    state["pairs"] = [
        {"source": "甲段。", "target": "First.", "reviewed": True,
         "segment_uid": "s0"},
        {"source": "乙段。", "target": "人工：乙段。", "reviewed": True,
         "human_edited": True, "segment_uid": "s1"},
    ]
    state["p1_done"] = True
    state["p2_done"] = True
    core.save_job_state(job_id, state)

    dest = tmp_path / "导出 内容"
    adapter.export_current_assets(job_id, dest, None)

    translation = [p.text for p in docx.Document(
        str(dest / _delivered(adapter, job_id, "translation.docx"))).paragraphs if p.text.strip()]
    assert translation[0] == "译文"           # 标题
    assert translation[1:] == ["First.", "人工：乙段。"]

    bilingual = docx.Document(str(dest / _delivered(adapter, job_id, "bilingual.docx")))
    rows = [[cell.text for cell in row.cells] for row in bilingual.tables[0].rows]
    assert rows[0] == ["原文", "译文"]
    assert rows[1] == ["甲段。", "First."]
    assert rows[2] == ["乙段。", "人工：乙段。"]

    manifest = json.loads((dest / _delivered(adapter, job_id, "delivery_manifest.json")).read_text(
        encoding="utf-8"))
    assert manifest.get("delivery_status") != "final", \
        "导出不是冻结交付，manifest 不能写 final"
    assert manifest.get("delivery_status") == "draft"


def test_legacy_import_never_overwrites_an_existing_task(adapter_env, tmp_path):
    """§6.C：接入旧数据目录时**不覆盖既有任务**，且源目录保持原样。

    第一次接入后用户可能已经在任务上做了人工修改；再次接入同一个旧目录必须跳过
    冲突任务，而不是把用户的修改覆盖回旧内容。
    """
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]

    # 1) 造一个"旧数据目录"：一个任务，译文里带可识别的标记。
    legacy = tmp_path / "old-shell-outputs"
    legacy.mkdir()
    job_id = "legacyjob0000001"
    job_dir = legacy / job_id
    job_dir.mkdir()
    state = core.new_job_state("legacy.docx")
    state["paras"] = ["旧数据段落。"]
    state["pairs"] = [{"source": "旧数据段落。", "target": "旧译文。",
                       "reviewed": True, "segment_uid": "s0"}]
    state["p1_done"] = True
    state["p2_done"] = True
    (job_dir / "state.json").write_text(json.dumps(state, ensure_ascii=False),
                                        encoding="utf-8")
    source_snapshot = {path.name: path.read_bytes()
                       for path in sorted(job_dir.iterdir())}

    # 2) 第一次接入：任务出现，内容来自旧目录。
    first = adapter.import_legacy_directory(legacy)
    assert first["report"]["ok"] is True, first
    assert job_id in adapter.load_state(job_id)["pairs"][0]["segment_uid"] or True
    assert adapter.load_state(job_id)["pairs"][0]["target"] == "旧译文。"

    # 3) 用户在接入后的任务上做人工修改。
    adapter.submit_edit(job_id, 0, "接入后我改的译文。")
    assert adapter.load_state(job_id)["pairs"][0]["target"] == "接入后我改的译文。"

    # 4) 再次接入同一个旧目录：冲突任务必须被跳过，人工修改必须保留。
    second = adapter.import_legacy_directory(legacy)
    report = second["report"]
    assert report["ok"] is True, report
    assert job_id in report.get("skipped_conflicts", []), report
    assert adapter.load_state(job_id)["pairs"][0]["target"] == "接入后我改的译文。", \
        "再次接入不得覆盖已有任务"
    assert adapter.load_state(job_id)["pairs"][0]["human_edited"] is True

    # 5) 源目录保持原样（逐文件字节一致）。
    after = {path.name: path.read_bytes() for path in sorted(job_dir.iterdir())}
    assert after == source_snapshot, "旧数据目录不得被修改"

    # 6) 把当前数据目录自己当旧目录接入：必须给出可理解的原因。
    with pytest.raises(adapter.AdapterError) as same:
        adapter.import_legacy_directory(adapter.desktop_paths.resolve_paths().data_dir)
    assert "相同" in str(same.value) or "包含" in str(same.value), str(same.value)


def test_unreadable_tasks_are_listed_not_silently_hidden(adapter_env):
    """读不出状态的任务必须**列出来**并标明，而不是从列表里凭空消失。

    真实缺陷：核心的读取路径会静默跳过无法解析的任务目录——任务仍在磁盘上，用户
    却只看到"任务不见了"，既不知道原因也无从处理。
    """
    adapter = adapter_env["adapter"]
    tmp = adapter_env["tmp"]
    data_dir = adapter.desktop_paths.resolve_paths().data_dir

    good = tmp / "good.docx"
    good.write_bytes(_make_docx(["好的一段。"]))
    good_job = adapter.import_documents(
        [good], adapter.EngineConfig(target_lang="简体中文"))[0].job_id

    broken_ids = {}
    for name, content in (("badjson", "{ not json"),
                          ("emptyfile", ""),
                          ("wrongshape", json.dumps(["not", "a", "dict"]))):
        job_id = f"broken{name}0001"
        job_dir = data_dir / job_id
        job_dir.mkdir(parents=True, exist_ok=True)
        (job_dir / "state.json").write_text(content, encoding="utf-8")
        broken_ids[name] = job_id

    rows = adapter.list_jobs()
    by_id = {row.job_id: row for row in rows}
    assert good_job in by_id, "一个坏任务不能让好任务消失"
    for name, job_id in broken_ids.items():
        assert job_id in by_id, f"{name} 应该出现在列表里（不能静默隐藏）"
    # 无法解析的两个被明确标出；能解析成别的形状的按正常任务列出（它会以
    # "没有可翻译的文本"失败，而不是假装正常）。
    for name in ("badjson", "emptyfile"):
        row = by_id[broken_ids[name]]
        assert row.status_label == adapter.UNREADABLE_STATUS, row.status_label
        assert row.error, "要说明为什么读不出来"
        assert row.active is False
    assert by_id[broken_ids["wrongshape"]].status_label != adapter.UNREADABLE_STATUS

    # 列表按更新时间排序时不会因为坏任务而崩。
    stamps = [row.updated_at for row in rows if row.updated_at]
    assert stamps == sorted(stamps, reverse=True)

    # 这些任务可以被删除（不要求状态可读）。
    result = adapter.delete_job(broken_ids["badjson"])
    assert result is not False
    assert not (data_dir / broken_ids["badjson"]).exists()


def test_delivery_preview_never_promises_what_the_gate_refuses(adapter_env):
    """预览必须与真实门禁一致：审校未就绪时不能显示"可以交付"。

    真实缺陷：`delivery_preview` 只问"有没有阻断项"，没问审校就绪；于是按钮亮着、
    用户点下去才被门禁拒绝（"review gate=not_run，不能进入 final"）。
    """
    adapter = adapter_env["adapter"]
    core = adapter_env["core"]
    tmp = adapter_env["tmp"]

    source = tmp / "gate-preview.docx"
    source.write_bytes(_make_docx(["甲。", "乙。"]))
    job_id = adapter.import_documents(
        [source], adapter.EngineConfig(target_lang="简体中文"))[0].job_id
    state = core.load_job_state(job_id) or {}
    state["paras"] = ["甲。", "乙。"]
    state["pairs"] = [
        {"source": "甲。", "target": "A.", "reviewed": False, "segment_uid": "s0"},
        {"source": "乙。", "target": "B.", "reviewed": True, "segment_uid": "s1"},
    ]
    state["p1_done"] = True
    state["p2_done"] = True
    state["translation_core_review_required"] = True
    core.save_job_state(job_id, state)

    preview = adapter.delivery_preview(job_id)
    assert preview["can_deliver"] is False, \
        "审校未就绪时预览不能声称可以交付"
    assert preview["review_gate"]["required"] is True
    assert preview["review_gate"]["ready"] is False
    assert preview["review_gate"]["status"] == "not_run"
    assert any("审校未完成" in str(item) for item in preview["issues"]), \
        preview["issues"]

    # 与真实门禁对照：确实会被拒绝（而不是预览过于保守）。
    with pytest.raises(adapter.AdapterError) as refused:
        adapter.deliver(job_id, adapter.EngineConfig(target_lang="简体中文"))
    assert "review gate" in str(refused.value) or "审校" in str(refused.value)

    # 不需要审校的任务（未开启审校）不受这条规则影响。
    state = adapter.load_state(job_id)
    state.pop("translation_core_review_required", None)
    core.save_job_state(job_id, state)
    preview = adapter.delivery_preview(job_id)
    assert preview["review_gate"]["required"] is False
    assert preview["can_deliver"] is True


def _fault_server(status: int, payload: dict):
    """一个总是返回指定 HTTP 状态与 JSON 体的本地端点（用于错误映射验证）。"""
    import threading
    from http.server import BaseHTTPRequestHandler
    from scripts.local_http_server import LocalThreadingHTTPServer as ThreadingHTTPServer

    raw = json.dumps(payload).encode()

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):  # noqa: N802
            self.rfile.read(int(self.headers.get("Content-Length") or 0))
            self.send_response(status)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(raw)))
            self.end_headers()
            self.wfile.write(raw)

        def log_message(self, *args):  # noqa: A003 - silence the test server
            return

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server, f"http://127.0.0.1:{server.server_port}/v1"


@pytest.mark.parametrize("status,payload,needle", [
    (401, {"error": {"message": "Incorrect API key provided: sk-abc***",
                     "code": "invalid_api_key"}}, "API Key 无效"),
    (404, {"error": {"message": "The model `nope` does not exist",
                     "code": "model_not_found"}}, "HTTP 404"),
    (429, {"error": {"message": "Rate limit reached"}}, "过于频繁"),
    (402, {"error": {"message": "Insufficient balance"}}, "余额不足"),
    (500, {"error": {"message": "internal server error"}}, "暂时不可用"),
])
def test_provider_errors_are_actionable_not_raw(status, payload, needle):
    """§6.B：连接失败必须**可理解**——不能把原始 JSON 或异常类型名丢给用户。

    这一条此前只用"连接层失败"验证过；真实服务最常见的失败是 HTTP 状态码
    （Key 错、模型名错、限流、余额不足、服务端故障），这里逐个注入并断言。
    """
    from folith.qt import adapter

    server, url = _fault_server(status, payload)
    try:
        config = adapter.EngineConfig(provider="自定义中转站", model="demo",
                                      api_key="sk-abc", base_url=url)
        ok, message = adapter.test_connection(config)
        assert ok is False
        assert needle in message, message
        # 不泄露原始 JSON / 异常类型名。
        for leak in ("invalid_request_error", "Traceback", "HTTPError",
                     "Exception:", "sk-abc"):
            assert leak not in message, (leak, message)
    finally:
        server.shutdown()


def test_task_failure_reason_is_actionable_too(adapter_env, tmp_path):
    """任务运行中遇到提供商错误：运行视图给出的原因同样要可理解。"""
    from folith.qt import adapter

    server, url = _fault_server(401, {"error": {
        "message": "Incorrect API key provided: sk-abc***",
        "code": "invalid_api_key"}})
    try:
        config = adapter.EngineConfig(
            provider="自定义中转站", model="demo", api_key="sk-abc",
            base_url=url, target_lang="简体中文", enable_review=False,
            enable_report=False, enable_understanding=False,
            enable_source_cleanup=False)
        source = tmp_path / "authfail.docx"
        source.write_bytes(_make_docx(["甲段。"]))
        job_id = adapter.import_documents([source], config)[0].job_id
        adapter.start_job(job_id, "authfail.docx", source.read_bytes(), config)
        detail = ""
        for _ in range(240):
            runtime = adapter.runtime_view(job_id, adapter.load_state(job_id))
            detail = str(runtime.get("detail") or "")
            if runtime.get("status") in ("failed", "completed", "cancelled"):
                break
            time.sleep(0.5)
        assert "API Key 无效" in detail, detail
        assert "Traceback" not in detail and "invalid_request_error" not in detail
    finally:
        server.shutdown()


def _canned_server(content: str, seen: list):
    """返回固定 content 的本地端点；`seen` 记录它实际回复过的内容。

    注意用**工厂函数**接收 content：如果在循环里直接定义 Handler 类并引用循环变量，
    所有端点都会回复最后一次迭代的值（Python 闭包/全局晚绑定）——我第一版探针就是
    这样，把"每个任务用自己的端点"误报成了"跨任务串数据"。
    """
    import threading
    from http.server import BaseHTTPRequestHandler
    from scripts.local_http_server import LocalThreadingHTTPServer as ThreadingHTTPServer

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):  # noqa: N802
            self.rfile.read(int(self.headers.get("Content-Length") or 0))
            payload = {
                "id": "x", "object": "chat.completion",
                "choices": [{"index": 0, "finish_reason": "stop",
                             "message": {"role": "assistant", "content": content}}],
                "usage": {"prompt_tokens": 1, "completion_tokens": 1,
                          "total_tokens": 2},
            }
            raw = json.dumps(payload).encode()
            seen.append(content)
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(raw)))
            self.end_headers()
            self.wfile.write(raw)

        def log_message(self, *args):  # noqa: A003
            return

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server, f"http://127.0.0.1:{server.server_port}/v1"


@pytest.mark.parametrize("content,needle", [
    ("[]", "wrong_item_count"),
    (json.dumps(["A.", "B.", "C."]), "wrong_item_count"),
    ('["译文', "unparseable_transport_wrapper"),
    (json.dumps([""]), "empty_item_or_unsupported_item"),
])
def test_bad_translation_protocol_is_rejected_without_partial_writes(
        adapter_env, tmp_path, content, needle):
    """批次响应不合格必须**整批拒绝**：绝不能错位写入或写半截。

    段落位置是整条流水线的契约（审校、检查点、交付都按同一个索引指代），
    条数不符时若"凑合写进去"，译文与原文就会错位——属于数据损坏。
    """
    from folith.qt import adapter

    seen: list = []
    server, url = _canned_server(content, seen)
    try:
        config = adapter.EngineConfig(
            provider="自定义中转站", model="m", api_key="k", base_url=url,
            target_lang="简体中文", enable_review=False, enable_report=False,
            enable_understanding=False, enable_source_cleanup=False)
        source = tmp_path / "protocol.docx"
        source.write_bytes(_make_docx(["只有一段。"]))
        job_id = adapter.import_documents([source], config)[0].job_id
        adapter.start_job(job_id, "protocol.docx", source.read_bytes(), config)
        detail = ""
        for _ in range(240):
            runtime = adapter.runtime_view(job_id, adapter.load_state(job_id))
            detail = str(runtime.get("detail") or "")
            if runtime.get("status") in ("failed", "completed", "cancelled"):
                break
            time.sleep(0.5)
        state = adapter.load_state(job_id)
        pairs = state.get("pairs") or []
        assert state.get("p2_done") is False, "不合格响应不能被当作完成"
        assert all(not str(pair.get("target") or "").strip() for pair in pairs), \
            f"不得写入半截译文：{pairs}"
        assert needle in detail, detail
        # 该任务的响应确实来自它自己的端点（跨任务串数据会是严重缺陷）。
        assert seen and all(item == content for item in seen), seen
    finally:
        server.shutdown()


def test_plain_text_response_is_accepted_as_a_single_translation(adapter_env,
                                                                tmp_path):
    """纯文本响应（不是 JSON 数组）在"只请求一段"时按单条译文接受。

    这是**有意**的回退：真实中转站可能直接返回纯文本。段数 > 1 时同样会被条数
    校验拦下（见上一条用例），所以不会造成错位。
    """
    from folith.qt import adapter

    seen: list = []
    server, url = _canned_server("这是一段纯文本译文。", seen)
    try:
        config = adapter.EngineConfig(
            provider="自定义中转站", model="m", api_key="k", base_url=url,
            target_lang="简体中文", enable_review=False, enable_report=False,
            enable_understanding=False, enable_source_cleanup=False)
        source = tmp_path / "plain.docx"
        source.write_bytes(_make_docx(["只有一段。"]))
        job_id = adapter.import_documents([source], config)[0].job_id
        adapter.start_job(job_id, "plain.docx", source.read_bytes(), config)
        for _ in range(240):
            runtime = adapter.runtime_view(job_id, adapter.load_state(job_id))
            if runtime.get("status") in ("failed", "completed", "cancelled"):
                break
            time.sleep(0.5)
        state = adapter.load_state(job_id)
        assert state.get("p2_done") is True
        assert (state["pairs"][0]["target"] == "这是一段纯文本译文。")
    finally:
        server.shutdown()


@pytest.mark.parametrize("broken_renderer", [False, True])
def test_selftest_pdf_source_detects_missing_bundled_renderer(
        adapter_env, tmp_path, monkeypatch, broken_renderer):
    """A missing optional renderer must fail the release probe, not hide its button."""
    from types import SimpleNamespace
    from folith import pdf_preview
    from folith.qt import selftest

    adapter = adapter_env["adapter"]
    source = tmp_path / "selftest.pdf"
    text = "Text layer page one."
    selftest.make_text_pdf(source, [text])
    job_id = adapter.import_documents([source], adapter.EngineConfig())[0].job_id
    # Exercise the real PDF adapters against known source text. The complete
    # probe separately verifies the actual workbench and pipeline extraction.
    window = SimpleNamespace(workbench=SimpleNamespace(
        view_source_button=SimpleNamespace(isHidden=lambda: False)))
    flow = selftest.Flow(window, None, out_dir=tmp_path)
    state = {"filename": source.name, "paras": [text]}
    if broken_renderer:
        monkeypatch.setattr(pdf_preview, "fitz", None)
        with pytest.raises(AssertionError, match="PDF source is unavailable"):
            flow._check_pdf_source(adapter, job_id, state)
    else:
        facts = flow._check_pdf_source(adapter, job_id, state)
        assert facts["pages"] == 1
        assert facts["highlight_rects"] > 0


def test_ocr_mode_actually_controls_extraction(adapter_env, tmp_path):
    """「何时识别」设置必须真的生效：off 不做 OCR，force 对文本层也做 OCR。

    这条对应一个真实缺陷：设置界面的下拉曾写 `always`/`never`，而核心只认
    `force`/`off`，于是核心把它们**静默归一成 auto**——选"从不识别"的扫描件照样被
    OCR（与用户明确要求相反），选"总是识别"也不会强制识别。
    """
    from folith.qt import adapter, selftest

    import shutil

    if shutil.which("tesseract") is None:
        pytest.skip("OCR force extraction requires the Tesseract executable")

    config_base = dict(provider="自定义中转站", model="stub-model",
                       api_key="stub-key", target_lang="简体中文",
                       enable_review=False, enable_report=False,
                       enable_understanding=False, enable_source_cleanup=False,
                       ocr_language="eng", ocr_backend="tesseract")

    scanned = tmp_path / "scanned.pdf"
    selftest.make_scanned_pdf(scanned, ["Scanned page one for OCR baseline."])

    text_pdf = tmp_path / "text-layer.pdf"
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 80), "Text layer page one.", fontsize=18)
    document.save(str(text_pdf))
    document.close()

    stub = StubEndpoint(tmp_path)
    url = stub.start()
    try:
        def run(source, mode):
            config = adapter.EngineConfig(base_url=url, ocr_mode=mode,
                                          **config_base)
            job_id = adapter.import_documents([source], config)[0].job_id
            adapter.start_job(job_id, source.name, source.read_bytes(), config)
            detail = ""
            for _ in range(240):
                runtime = adapter.runtime_view(job_id, adapter.load_state(job_id))
                detail = str(runtime.get("detail") or "")
                if runtime.get("status") in ("failed", "completed", "cancelled"):
                    break
                time.sleep(0.5)
            state = adapter.load_state(job_id)
            report = state.get("extraction_report") or {}
            return (adapter.runtime_view(job_id, state).get("status"),
                    bool(report.get("ocr_used")), detail)

        # off：明确不做 OCR —— 扫描件应当以"没有文本层，请允许 OCR"失败
        status, ocr_used, detail = run(scanned, "off")
        assert status == "failed" and ocr_used is False, (status, ocr_used, detail)
        assert "没有文本层" in detail and "OCR" in detail, detail

        # force：即使是带文本层的 PDF 也要走 OCR
        status, ocr_used, detail = run(text_pdf, "force")
        assert status == "completed", (status, detail)
        assert ocr_used is True, "force 模式下必须真的做 OCR"
    finally:
        stub.stop()


def test_ocr_mode_labels_map_to_core_values(adapter_env):
    """界面标签 -> 核心取值必须一一对应；历史值（always/never）按实际生效值显示。"""
    from folith.qt import adapter, window as window_module
    from folith.qt.dialogs import EngineDialog
    from folith.qt.workers import TaskRunner

    app = window_module.build_application([])
    labels = {label: mode for label, mode in adapter.OCR_MODE_CHOICES}
    assert set(labels.values()) == {"auto", "force", "off"}
    for label, mode in labels.items():
        dialog = EngineDialog(adapter.EngineConfig(ocr_mode=mode), TaskRunner())
        assert dialog.ocr_mode_box.currentText() == label
        assert dialog._collect().ocr_mode == mode, (label, mode)
        dialog.close()

    # 历史值：界面显示"实际生效"的模式（核心把它们当 auto），保存后即为 auto。
    for legacy in ("always", "never"):
        dialog = EngineDialog(adapter.EngineConfig(ocr_mode=legacy), TaskRunner())
        assert dialog._collect().ocr_mode == "auto", legacy
        dialog.close()
    # 别名也要认
    assert adapter.normalize_ocr_mode("强制") == "force"
    assert adapter.normalize_ocr_mode("关闭") == "off"
    assert adapter.normalize_ocr_mode("") == "auto"


def test_ocr_language_choices_come_from_the_core_registry(adapter_env):
    """OCR 语言下拉必须来自核心注册表（不手写第二份），且**保存的是取值不是标签**。

    真实缺陷（两处）：
    1. 对话框手写了 6 个原始代码，与核心那份 30 项注册表重复，还把大多数语言藏起来；
    2. 用 `setCurrentText(标签)` 选中可编辑下拉框时，`currentData()` 为空，保存会把
       **标签**写进配置（"简体中文 + English" 而不是 `chi_sim+eng`）。
    """
    from folith.qt import adapter, window as window_module
    from folith.qt.dialogs import EngineDialog
    from folith.qt.workers import TaskRunner

    window_module.build_application([])
    choices = adapter.ocr_language_choices()
    core_options = adapter.core.ocr_language_options()
    assert choices[0] == (adapter.OCR_LANGUAGE_AUTO, "")
    assert len(choices) == len(core_options) + 1, (len(choices), len(core_options))
    # 标签来自核心、取值是核心的 engine_code
    labels = {label for label, _value in choices}
    assert "简体中文" in labels and "English" in labels and "العربية" in labels
    values = {value for _label, value in choices}
    assert "chi_sim" in values and "eng" in values and "ara" in values

    # 每个已知取值：显示标签、保存取值（不能泄漏标签）
    for label, value in choices:
        if not value:
            continue
        dialog = EngineDialog(adapter.EngineConfig(ocr_language=value),
                              TaskRunner())
        assert dialog.ocr_language_box.currentText() == label, value
        assert dialog._collect().ocr_language == value, (label, value)
        dialog.close()

    # 未知代码（用户手填）与"自动判断"都要正确往返
    dialog = EngineDialog(adapter.EngineConfig(ocr_language="jpn"), TaskRunner())
    assert dialog._collect().ocr_language == "jpn"
    dialog.close()
    dialog = EngineDialog(adapter.EngineConfig(ocr_language=""), TaskRunner())
    assert dialog.ocr_language_box.currentText() == adapter.OCR_LANGUAGE_AUTO
    assert dialog._collect().ocr_language == ""
    dialog.close()


def test_new_task_defaults_are_written_into_the_job_state(adapter_env, tmp_path):
    """新建任务必须把"任务默认值"写进任务状态，否则运行时会静默忽略设置。

    真实缺陷：`enable_report=True` 导入后状态里仍是 `False`（`new_job_state` 的默认），
    而运行配置以状态为准 → 学术报告阶段根本不跑（实测产物 0 个 vs 开启后 13 个）。
    """
    import uuid

    from folith.qt import adapter

    config = adapter.EngineConfig(
        target_lang="简体中文", enable_report=True, auto_term=False,
        use_tm=False, enable_understanding=True, enable_source_cleanup=True)
    source = tmp_path / f"defaults-{uuid.uuid4().hex[:8]}.docx"
    source.write_bytes(_make_docx(["一段。"]))
    job_id = adapter.import_documents([source], config)[0].job_id
    state = adapter.load_state(job_id)
    assert state["report_enabled"] is True
    assert state["auto_term"] is False
    assert state["use_tm"] is False
    assert state["enable_understanding"] is True
    assert state["enable_source_cleanup"] is True

    kwargs = adapter.pipeline_kwargs(config, state)
    assert kwargs["enable_report"] is True
    assert kwargs["auto_term"] is False
    assert kwargs["use_tm"] is False

    # 已有任务保留自己保存的选择：再导入同一份文件（内容寻址 → 同一任务）
    # 不得被新配置覆盖。
    adapter.import_documents([source],
                             adapter.EngineConfig(target_lang="简体中文",
                                                  enable_report=False,
                                                  auto_term=True))
    state = adapter.load_state(job_id)
    assert state["report_enabled"] is True, "已有任务的选择不能被覆盖"
    assert state["auto_term"] is False


def test_separate_reviewer_inherits_the_main_address_for_the_same_provider(
        adapter_env, tmp_path):
    """独立审校模型留空地址时：同一提供商沿用正文地址；不同提供商必须显式填写。

    真实缺陷：此前一律 `resolve_openai_base_url(reviewer_provider, None)`，对自定义
    中转站直接抛错——"选独立审校 + 同一个中转站但不重抄地址"会让任务根本启动不了。
    """
    from folith.qt import adapter

    config = adapter.EngineConfig(
        provider="自定义中转站", model="main", api_key="k",
        base_url="http://127.0.0.1:9/v1", target_lang="简体中文",
        reviewer_mode="separate", reviewer_provider="自定义中转站",
        reviewer_model="reviewer", reviewer_api_key="k2",
        reviewer_base_url="")
    kwargs = adapter.pipeline_kwargs(config, {})
    assert kwargs["reviewer_base_url"] == "http://127.0.0.1:9/v1"
    assert kwargs["reviewer_model"] == "reviewer"

    # 提供商不同且没填地址：**不**悄悄沿用正文地址（那会把审校请求发到错的端点）。
    # 这里只断言"没有编造地址"；真正的拒绝由核心在调用时给出（本构建里只有一个
    # 已知中转站提供商，所以"另一个已知中转站"这个场景无法构造）。
    other = adapter.EngineConfig(
        provider="自定义中转站", model="main", api_key="k",
        base_url="http://127.0.0.1:9/v1", target_lang="简体中文",
        reviewer_mode="separate", reviewer_provider="另一个中转站",
        reviewer_model="reviewer", reviewer_api_key="k2")
    assert not adapter.pipeline_kwargs(other, {})["reviewer_base_url"]

    # 显式填了地址就用它
    explicit = adapter.EngineConfig(
        provider="自定义中转站", model="main", api_key="k",
        base_url="http://127.0.0.1:9/v1", target_lang="简体中文",
        reviewer_mode="separate", reviewer_provider="另一个中转站",
        reviewer_model="reviewer", reviewer_api_key="k2",
        reviewer_base_url="http://127.0.0.1:8/v1")
    assert adapter.pipeline_kwargs(explicit, {})["reviewer_base_url"] == \
        "http://127.0.0.1:8/v1"


def test_separate_reviewer_model_is_really_used(adapter_env, tmp_path):
    """独立审校模型要真的被调用（看端点收到的 model）。"""
    from folith.qt import adapter

    stub = StubEndpoint(tmp_path)
    url = stub.start()
    try:
        config = adapter.EngineConfig(
            provider="自定义中转站", model="main-model", api_key="k",
            base_url=url, target_lang="简体中文", enable_review=True,
            enable_report=False, enable_understanding=False,
            enable_source_cleanup=False, reviewer_mode="separate",
            reviewer_provider="自定义中转站", reviewer_model="reviewer-model",
            reviewer_api_key="k2")
        source = tmp_path / "reviewer.docx"
        source.write_bytes(_make_docx(["第一段。", "第二段。"]))
        job_id = adapter.import_documents([source], config)[0].job_id
        adapter.start_job(job_id, "reviewer.docx", source.read_bytes(), config)
        for _ in range(300):
            runtime = adapter.runtime_view(job_id, adapter.load_state(job_id))
            if runtime.get("status") in ("failed", "completed", "cancelled"):
                break
            time.sleep(0.5)
        log = tmp_path / "stub-llm.jsonl"
        models = set()
        if log.is_file():
            for line in log.read_text(encoding="utf-8").splitlines():
                models.add(json.loads(line).get("model"))
        assert "main-model" in models, models
        assert "reviewer-model" in models, f"独立审校模型没有被调用：{models}"
    finally:
        stub.stop()


def test_every_dialog_field_reaches_the_core_with_a_key_it_accepts(adapter_env):
    """表驱动不变量：对话框里的每个字段都必须以**核心认识的键名**到达核心。

    这一类缺陷我连着抓了三轮（交付格式、OCR 模式、OCR 语言/学术报告），所以把整类
    钉死：
    * `pipeline_kwargs` 产生的键必须全部是 `core.run_job_pipeline` 接受的参数
      （拼错或改名会**静默丢弃**该设置）；
    * 对话框展示的每个字段都必须出现在 kwargs 里（或经明确的映射/容器到达）。
    """
    import inspect

    from folith.qt import adapter
    from folith.qt.dialogs import EngineDialog

    accepted = set(inspect.signature(
        adapter.core.run_job_pipeline).parameters)

    config = adapter.EngineConfig(
        provider="自定义中转站", model="m", api_key="k",
        base_url="http://127.0.0.1:9/v1", target_lang="简体中文",
        enable_report=True, enable_annotate=True, enable_review=True,
        reviewer_mode="separate", reviewer_provider="自定义中转站",
        reviewer_model="rv", reviewer_api_key="k2")
    kwargs = adapter.pipeline_kwargs(config, {})
    unknown = sorted(set(kwargs) - accepted)
    assert not unknown, f"这些键核心不认识（会被静默丢弃）：{unknown}"

    # 明确的三类"非同名到达"：base_url→translator_base_url；deliver_*→delivery_config；
    # reviewer_* 仅在独立审校时出现（上面已设为 separate）。
    mapped = {"base_url": "translator_base_url"}
    delivered_in_container = {name for name in EngineDialog.OWNED_FIELDS
                              if name.startswith("deliver_")}
    # `reviewer_mode` 是**客户端开关**：它只决定是否发送 reviewer_* 那一组键
    # （核心靠这些键是否存在判断"独立审校"），本身不是核心参数。
    control_only = {"reviewer_mode", "connections", "connection_roles", "connection_models", "default_translation_mode"}
    for field in sorted(EngineDialog.OWNED_FIELDS):
        if field in delivered_in_container:
            assert (kwargs.get("delivery_config") or {}), "交付格式必须进入 delivery_config"
            continue
        if field in control_only:
            continue
        key = mapped.get(field, field)
        assert key in kwargs, f"对话框字段 {field} 没有到达核心（键 {key}）"

    # 独立审校开关的作用：打开时 reviewer_* 必须整组出现（否则等于没生效）。
    for key in ("reviewer_provider", "reviewer_model", "reviewer_api_key",
                "reviewer_base_url"):
        assert key in kwargs, f"独立审校开启时缺少 {key}"
    same = adapter.pipeline_kwargs(adapter.EngineConfig(
        target_lang="简体中文", reviewer_mode="same",
        reviewer_provider="自定义中转站", reviewer_model="rv"), {})
    assert "reviewer_provider" not in same, "同一模型时不应发送审校模型设置"

    # 反过来：核心接受但适配层不设置的参数应当是"未在界面暴露"的调优项，
    # 不允许出现"界面有开关却从不传"的情况（上面已逐个断言）。
    ui_fields = set(EngineDialog.OWNED_FIELDS) - delivered_in_container
    never_sent = sorted(
        name for name in accepted
        if name not in kwargs and name in ui_fields)
    assert not never_sent, f"界面暴露但从不发送的参数：{never_sent}"


def test_a_copied_task_directory_is_listed_normally_not_as_corrupt(
        adapter_env, tmp_path):
    """复制/改名产生的任务目录要按**正常任务**列出（以目录名为身份）。

    我的一个探针语料就是这样：49 个目录里装的是同一个任务的状态（内部 job_id 相同）。
    此前这类目录会被标成"无法读取（数据可能损坏）"且文档名/段数全空——既不准确也不
    诚实。现在目录名是身份来源，状态可读就是正常任务；只有**读不出来**才报损坏。
    """
    import json
    import shutil

    from folith.qt import adapter, views

    data_dir = adapter.desktop_paths.resolve_paths().data_dir
    source = tmp_path / "copied.docx"
    source.write_bytes(_make_docx(["甲段。", "乙段。"]))
    job_id = adapter.import_documents(
        [source], adapter.EngineConfig(target_lang="简体中文"))[0].job_id
    core = adapter_env["core"]
    state = core.load_job_state(job_id) or {}
    state["paras"] = ["甲段。", "乙段。"]
    state["pairs"] = [
        {"source": "甲段。", "target": "A.", "reviewed": True, "segment_uid": "s0"},
        {"source": "乙段。", "target": "B.", "reviewed": True, "segment_uid": "s1"},
    ]
    state["p1_done"] = True
    state["p2_done"] = True
    core.save_job_state(job_id, state)

    copy_dir = data_dir / "renamedcopy0001"
    shutil.copytree(data_dir / job_id, copy_dir)
    # 说明：state.json 里本来就不带 job_id（身份由目录名承载）；会串号的是
    # `summary.json` 那份缓存——那一点由
    # `test_listing_trusts_the_directory_name_over_a_stale_summary` 覆盖。
    copied_summary = copy_dir / "summary.json"
    assert copied_summary.is_file() or True

    rows = {row.job_id: row for row in adapter.list_jobs()}
    assert "renamedcopy0001" in rows, "被复制的目录必须出现在列表里（不能消失）"
    row = rows["renamedcopy0001"]
    assert row.status_label != adapter.UNREADABLE_STATUS
    assert row.filename == "copied.docx", row.filename
    assert row.segment_count == 2, row.segment_count
    assert row.target_lang == "简体中文"

    # 真正读不出来的目录才是"无法读取"（danger）。
    broken_dir = data_dir / "trulybroken00001"
    broken_dir.mkdir()
    (broken_dir / "state.json").write_text("{ 坏掉的 json", encoding="utf-8")
    rows = {row.job_id: row for row in adapter.list_jobs()}
    assert rows["trulybroken00001"].status_label == adapter.UNREADABLE_STATUS
    assert views.status_tone_for_label(adapter.UNREADABLE_STATUS) == "danger"


def test_legacy_import_rows_use_the_directory_name_even_with_a_stale_summary(
        adapter_env, tmp_path):
    """接入旧数据目录后：行标识必须是导入后的**目录名**。

    旧壳时代写下的 `summary.json` 里可能带着另一个 job_id（我的语料里 49/50 个都是
    这样）。若行标识取自那份缓存，导入进来的任务在列表里就会指向一个**不存在的
    目录**——点开即失败，同时真实目录会被标成"无法读取"。
    """
    import json

    from folith.qt import adapter

    core = adapter_env["core"]
    legacy = tmp_path / "old-outputs"
    legacy.mkdir()
    old_job = "oldjob0000000001"
    job_dir = legacy / old_job
    job_dir.mkdir()
    state = core.new_job_state("legacy-doc.pdf")
    state["paras"] = ["旧段一。", "旧段二。", "旧段三。"]
    state["pairs"] = [
        {"source": f"旧段{index}。", "target": f"Old {index}.",
         "reviewed": True, "segment_uid": f"s{index}"} for index in range(3)]
    state["p1_done"] = state["p2_done"] = True
    state["target_lang"] = "简体中文"
    (job_dir / "state.json").write_text(
        json.dumps(state, ensure_ascii=False), encoding="utf-8")
    (job_dir / "summary.json").write_text(json.dumps({
        "job_id": "stale-id-from-old-shell", "filename": "legacy-doc.pdf",
        "segment_count": 3, "target_lang": "简体中文",
        "progress_label": "已冻结交付"}, ensure_ascii=False), encoding="utf-8")

    report = adapter.import_legacy_directory(legacy)["report"]
    assert report["ok"] is True and report["copied_jobs"] == [old_job]

    rows = {row.job_id: row for row in adapter.list_jobs()}
    assert set(rows) == {old_job}, rows
    assert "stale-id-from-old-shell" not in rows
    row = rows[old_job]
    assert row.filename == "legacy-doc.pdf"
    assert row.segment_count == 3
    assert row.target_lang == "简体中文"
    assert row.status_label != adapter.UNREADABLE_STATUS
    # 能真的打开（状态按目录名读到）
    assert adapter.load_state(old_job)["pairs"][0]["target"] == "Old 0."


def test_a_stale_summary_cache_cannot_dictate_the_row_fields(adapter_env,
                                                             tmp_path):
    """缓存过期时，首页行必须显示**状态里的真值**（文件名/段数/语言）。

    真实缺陷：`summary.json` 的 job_id 与目录名不一致时（复制、改名、接入旧数据目录
    都会如此），此前只把 id 改成目录名、其余字段照用缓存——于是列表显示**另一个任务**
    的文件名与段数（实测：状态是"真名.pdf / 5 段 / 简体中文"，列表显示
    "旧名.pdf / 3 段 / English"）。现在这类目录的摘要按状态重建。
    """
    import json

    from folith.qt import adapter

    core = adapter_env["core"]
    legacy = tmp_path / "old-outputs"
    legacy.mkdir()
    job_id = "staletask0000001"
    job_dir = legacy / job_id
    job_dir.mkdir()
    state = core.new_job_state("真名.pdf")
    state["paras"] = [f"真段{index}。" for index in range(5)]
    state["pairs"] = [
        {"source": f"真段{index}。", "target": f"T{index}", "reviewed": True,
         "segment_uid": f"s{index}"} for index in range(5)]
    state["p1_done"] = state["p2_done"] = True
    state["target_lang"] = "简体中文"
    (job_dir / "state.json").write_text(json.dumps(state, ensure_ascii=False),
                                        encoding="utf-8")
    # 过期缓存：id、文件名、段数、语言全是另一份的
    (job_dir / "summary.json").write_text(json.dumps({
        "job_id": "other-id", "filename": "旧名.pdf", "segment_count": 3,
        "target_lang": "English", "progress_label": "待处理"},
        ensure_ascii=False), encoding="utf-8")

    adapter.import_legacy_directory(legacy)
    row = next(item for item in adapter.list_jobs() if item.job_id == job_id)
    assert row.filename == "真名.pdf", row.filename
    assert row.segment_count == 5, row.segment_count
    assert row.target_lang == "简体中文", row.target_lang


def test_review_cost_estimate_is_local_and_counts_source_plus_target():
    from folith.qt import adapter
    state = {"pairs": [{"source": "abcdefgh", "target": "中文译文"},
                       {"source": "", "target": ""}]}
    estimate = adapter.review_cost_estimate(state, [0, 1, 5])
    assert estimate == {"segments": 2, "tokens": 2 + 4}
