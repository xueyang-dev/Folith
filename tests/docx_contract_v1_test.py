"""DOCX Contract v1: 原格式保真往返（Surgical round-trip）验收测试套件。

遵循 docs/docx-contract-v1.md 验收标准与 6 项验收核查：
1. 支持范围内全部内容完成提取、编辑/翻译和写回，保留项不计入翻译完成数。
2. 未修改 ZIP 条目逐项比较解压内容（100% 字节一致）；修改 part 检查结构保真。
3. 丢 tag、失效 anchor、未决修订等故障输入必须明确失败。
4. 未翻译/保留报告可用于交付核查。
5. 文件可正常打开并保持结构完整。
6. 现有审校、人工批准、失效传播与交付快照路径成立。
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import zipfile
from pathlib import Path

import pytest
import core
from docx import Document
from folith.docx_roundtrip import (
    CONTRACT_VERSION,
    DocxUnit,
    InvalidAnchorError,
    TagQAError,
    TrackChangesPendingError,
    build_roundtrip_docx,
    extract_roundtrip_units,
    inspect_docx_package,
    validate_target_tags,
    verify_roundtrip_docx,
    write_roundtrip_docx,
)
from folith.engine.document_export import build_translation_docx
from tests.fixtures.docx_contract.generate_fixtures import generate_all_fixtures

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures" / "docx_contract"


@pytest.fixture(scope="module", autouse=True)
def ensure_fixtures():
    """Ensure all 9 Contract v1 fixtures exist."""
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    generate_all_fixtures(FIXTURES_DIR)


def _load_fixture(name: str) -> bytes:
    path = FIXTURES_DIR / name
    assert path.exists(), f"Fixture {name} not found in {FIXTURES_DIR}"
    return path.read_bytes()


# ---------------------------------------------------------------------------
# 1. 未决修订（Track Changes）明确拒绝（契约 line 47）
# ---------------------------------------------------------------------------

def test_contract_v1_track_changes_rejection():
    """未决修订必须在翻译前明确拒绝；禁止静默接受、拒绝或遗漏。"""
    data = _load_fixture("docx-track-changes.docx")

    # 1. 底层检查必须抛出 TrackChangesPendingError
    with pytest.raises(TrackChangesPendingError, match="未决修订"):
        inspect_docx_package(data)

    with pytest.raises(TrackChangesPendingError, match="未决修订"):
        extract_roundtrip_units(data)

    # 2. 核心层 extract_document_paragraphs_with_report 必须如实置 failed 并说明原因
    paragraphs, warnings, report = core.extract_document_paragraphs_with_report(
        "docx-track-changes.docx", data
    )
    assert report.get("status") == "failed"
    assert "未决修订" in report.get("error", "")


# ---------------------------------------------------------------------------
# 2. 正文与标题层级保真写回 (docx-simple.docx)
# ---------------------------------------------------------------------------

def test_contract_v1_simple_body_and_headings():
    """普通正文与标题层级提取、翻译及原格式写回，保留样式和段落顺序。"""
    data = _load_fixture("docx-simple.docx")

    manifest = inspect_docx_package(data)
    assert manifest["counts"]["body_paragraphs"] == 4
    units, meta = extract_roundtrip_units(data)
    assert len(units) == 4

    # 模拟翻译
    translations = [
        "Folith 文档往返保真验证",
        "一、文档保真度导论",
        "这是带有标准格式的第一段正文。",
        "这是测试段落顺序与身份的第二段正文。",
    ]
    for unit, target in zip(units, translations):
        unit.target = target

    out_bytes, report = write_roundtrip_docx(data, units, target_lang="zh-CN")
    res = verify_roundtrip_docx(data, out_bytes, units, report)
    assert res["passed"] is True
    assert res["unmodified_entries_identical"] is True

    # 检查输出文档内容
    out_doc = Document(io.BytesIO(out_bytes))
    out_texts = [p.text for p in out_doc.paragraphs if p.text.strip()]
    assert out_texts == translations
    # 标题样式保留
    assert out_doc.paragraphs[0].style.name == "Title"
    assert out_doc.paragraphs[1].style.name == "Heading 1"


# ---------------------------------------------------------------------------
# 3. 表格与合并单元格保真写回 (docx-table.docx)
# ---------------------------------------------------------------------------

def test_contract_v1_tables_and_merged_cells():
    """表格文字翻译，表格结构、行列及单元格合并（gridSpan/vMerge）完好保留。"""
    data = _load_fixture("docx-table.docx")

    manifest = inspect_docx_package(data)
    assert manifest["counts"]["tables"] == 2
    assert manifest["counts"]["merged_cells"] >= 1

    units, meta = extract_roundtrip_units(data)
    table_units = [u for u in units if u.source_anchor["kind"] == "table_cell"]
    assert len(table_units) >= 8

    # 翻译表格单元
    for u in units:
        if "Quarter" in u.source:
            u.target = "季度"
        elif "Revenue" in u.source:
            u.target = "营收（美元）"
        elif "Global Division Overview" in u.source:
            u.target = "全球部门总览"
        elif "Region" in u.source:
            u.target = "地区"
        elif "Staff Count" in u.source:
            u.target = "员工数"
        elif "Operating Status" in u.source:
            u.target = "运营状态"
        elif "Active" in u.source:
            u.target = "正常运营"
        else:
            u.target = u.source

    out_bytes, report = write_roundtrip_docx(data, units, target_lang="zh-CN")
    res = verify_roundtrip_docx(data, out_bytes, units, report)
    assert res["passed"] is True
    assert res["structures_preserved"] is True

    out_doc = Document(io.BytesIO(out_bytes))
    assert len(out_doc.tables) == 2
    t1 = out_doc.tables[0]
    assert t1.cell(0, 0).text == "季度"
    assert t1.cell(0, 1).text == "营收（美元）"

    t2 = out_doc.tables[1]
    # 合并单元格内容
    assert t2.cell(0, 0).text == "全球部门总览"
    assert t2.cell(1, 0).text == "地区"


# ---------------------------------------------------------------------------
# 4. Inline 格式与目标词序变化 (docx-inline-formatting.docx)
# ---------------------------------------------------------------------------

def test_contract_v1_inline_formatting_and_word_order():
    """基本 inline 格式（粗体/斜体/下划线）在目标语言词序调整后语义归属正确。"""
    data = _load_fixture("docx-inline-formatting.docx")
    units, meta = extract_roundtrip_units(data)

    # 寻找包含粗体/斜体的段落
    p1 = next(u for u in units if "bold text" in u.source)
    assert "<b>bold text</b>" in p1.source
    assert "<i>italic styling</i>" in p1.source
    p1.target = "此段落测试<b>粗体文本</b>以及<i>斜体样式</i>。"

    # 词序重排段落：
    # 源文: We offer <b>twenty-four hour support</b> for all global enterprise clients.
    # 译文: 我们为所有全球企业客户提供<b>全天候二十四小时支持</b>。
    p_reorder = next(u for u in units if "twenty-four hour support" in u.source)
    p_reorder.target = "我们为所有全球企业客户提供<b>全天候二十四小时支持</b>。"

    out_bytes, report = write_roundtrip_docx(data, units, target_lang="zh-CN")
    res = verify_roundtrip_docx(data, out_bytes, units, report)
    assert res["passed"] is True

    # 检查输出 run 级格式
    out_doc = Document(io.BytesIO(out_bytes))
    # 验证 p1
    found_p1 = next(p for p in out_doc.paragraphs if "粗体文本" in p.text)
    bold_runs = [r for r in found_p1.runs if r.bold]
    assert len(bold_runs) >= 1
    assert "粗体文本" in "".join(r.text for r in bold_runs)

    italic_runs = [r for r in found_p1.runs if r.italic]
    assert len(italic_runs) >= 1
    assert "斜体样式" in "".join(r.text for r in italic_runs)

    # 验证词序重排段落
    found_p4 = next(p for p in out_doc.paragraphs if "全天候二十四小时支持" in p.text)
    reorder_bold = [r for r in found_p4.runs if r.bold]
    assert len(reorder_bold) >= 1
    assert "全天候二十四小时支持" in "".join(r.text for r in reorder_bold)


# ---------------------------------------------------------------------------
# 5. 超链接显示文字与链接目标保真 (docx-hyperlinks.docx)
# ---------------------------------------------------------------------------

def test_contract_v1_hyperlinks_target_preservation():
    """超链接显示文字翻译，外链 relationship 与内部 anchor 100% 保持不变。"""
    data = _load_fixture("docx-hyperlinks.docx")
    units, meta = extract_roundtrip_units(data)

    u_link = next(u for u in units if "<link id=" in u.source)
    assert "Example Documentation" in u_link.source
    link_id = u_link.tags[0]["id"]
    u_link.target = f'访问官方文档：<link id="{link_id}">示例技术文档</link>了解细节。'

    u_anchor = next(u for u in units if "<link anchor=" in u.source)
    u_anchor.target = '参见<link anchor="sec_compliance">内部合规指引</link>获取操作规范。'

    out_bytes, report = write_roundtrip_docx(data, units, target_lang="zh-CN")
    res = verify_roundtrip_docx(data, out_bytes, units, report)
    assert res["passed"] is True

    # 检查 relationships 字节完全相同
    with zipfile.ZipFile(io.BytesIO(data)) as sz, zipfile.ZipFile(io.BytesIO(out_bytes)) as oz:
        assert sz.read("word/_rels/document.xml.rels") == oz.read("word/_rels/document.xml.rels")

    out_doc = Document(io.BytesIO(out_bytes))
    p_text = " ".join(p.text for p in out_doc.paragraphs)
    assert "示例技术文档" in p_text
    assert "内部合规指引" in p_text


# ---------------------------------------------------------------------------
# 6. 脚注文本与引用关系保真 (docx-footnotes.docx)
# ---------------------------------------------------------------------------

def test_contract_v1_footnotes_preservation_and_translation():
    """普通脚注文本翻译，保留引用 ID、引用标记及分隔结构。"""
    data = _load_fixture("docx-footnotes.docx")
    units, meta = extract_roundtrip_units(data)

    fn_units = [u for u in units if u.source_anchor["kind"] == "footnote"]
    assert len(fn_units) == 2

    for u in fn_units:
        if "Smith" in u.source:
            u.target = "基线方法论参见 Smith 与 Brown (2025)。"
        elif "Empirical validation" in u.source:
            u.target = "实验室条件下的实证验证。"

    out_bytes, report = write_roundtrip_docx(data, units, target_lang="zh-CN")
    res = verify_roundtrip_docx(data, out_bytes, units, report)
    assert res["passed"] is True

    with zipfile.ZipFile(io.BytesIO(out_bytes)) as oz:
        fn_xml = oz.read("word/footnotes.xml").decode("utf-8")
        assert "基线方法论参见 Smith 与 Brown" in fn_xml
        assert "实验室条件下的实证验证" in fn_xml
        assert 'w:type="separator"' in fn_xml
        assert 'w:type="continuationSeparator"' in fn_xml


# ---------------------------------------------------------------------------
# 7. 页眉页脚共享去重与页码字段保护 (docx-headers-footers.docx)
# ---------------------------------------------------------------------------

def test_contract_v1_headers_footers_and_fields():
    """页眉页脚翻译普通文字，保留节引用，页码字段受保护。"""
    data = _load_fixture("docx-headers-footers.docx")
    units, meta = extract_roundtrip_units(data)

    hdr_unit = next(u for u in units if u.source_anchor["kind"] == "header")
    hdr_unit.target = "Folith 企业级智能翻译工作台"

    out_bytes, report = write_roundtrip_docx(data, units, target_lang="zh-CN")
    res = verify_roundtrip_docx(data, out_bytes, units, report)
    assert res["passed"] is True

    out_doc = Document(io.BytesIO(out_bytes))
    header_text = out_doc.sections[0].header.paragraphs[0].text
    assert "Folith 企业级智能翻译工作台" in header_text

    # 页脚字段指令仍保留
    with zipfile.ZipFile(io.BytesIO(out_bytes)) as oz:
        footer_xml = oz.read("word/footer1.xml").decode("utf-8")
        assert 'w:instr="PAGE"' in footer_xml
        assert 'w:instr="NUMPAGES"' in footer_xml


# ---------------------------------------------------------------------------
# 8. 不支持/保留元素安全保留与报告呈现 (docx-unsupported-features.docx)
# ---------------------------------------------------------------------------

def test_contract_v1_unsupported_features_preserved_and_reported():
    """文本框、批注、公式、图片安全保留在原位置；报告如实声明，不计入翻译。"""
    data = _load_fixture("docx-unsupported-features.docx")
    manifest = inspect_docx_package(data)

    assert manifest["counts"]["textboxes"] >= 1
    assert manifest["counts"]["comments"] >= 1
    assert manifest["counts"]["formulas"] >= 1
    assert manifest["counts"]["drawings"] >= 1

    preserved_kinds = {item["kind"] for item in manifest["preserved"]}
    assert "textbox" in preserved_kinds
    assert "comment" in preserved_kinds
    assert "formula" in preserved_kinds
    assert "image" in preserved_kinds

    units, meta = extract_roundtrip_units(data)
    for u in units:
        u.target = "契约 v1 保留功能与文本验证。"

    out_bytes, report = write_roundtrip_docx(data, units, target_lang="zh-CN")
    res = verify_roundtrip_docx(data, out_bytes, units, report)
    assert res["passed"] is True

    # 验证未修改条目（如批注 comments.xml、图片 media/）字节完全相同
    with zipfile.ZipFile(io.BytesIO(data)) as sz, zipfile.ZipFile(io.BytesIO(out_bytes)) as oz:
        assert sz.read("word/comments.xml") == oz.read("word/comments.xml")
        media_names = [n for n in sz.namelist() if n.startswith("word/media/")]
        for m in media_names:
            assert sz.read(m) == oz.read(m)


# ---------------------------------------------------------------------------
# 9. 复合结构保真 (docx-mixed-structures.docx)
# ---------------------------------------------------------------------------

def test_contract_v1_mixed_structures():
    """表格中超链接、脚注附近格式切换、页眉字段等组合结构无损坏。"""
    data = _load_fixture("docx-mixed-structures.docx")
    units, meta = extract_roundtrip_units(data)

    for u in units:
        if "<link" in u.source:
            lid = next((t["id"] for t in u.tags if t.get("type") == "hyperlink"), "rId1")
            u.target = f'仔细阅读<link id="{lid}">技术规格书</link>。' 
        elif "Approved" in u.source:
            u.target = "状态：<b>已批准</b>"
        elif "neural parsing" in u.source:
            u.target = '先进的<i>神经解析技术</i><ref type="footnote" id="1"/>产出核验结果。'
        elif "Neural parsing model" in u.source:
            u.target = "神经解析模型版本 4.2。"
        else:
            u.target = u.source

    out_bytes, report = write_roundtrip_docx(data, units, target_lang="zh-CN")
    res = verify_roundtrip_docx(data, out_bytes, units, report)
    assert res["passed"] is True

    out_doc = Document(io.BytesIO(out_bytes))
    t = out_doc.tables[0]
    assert "技术规格书" in t.cell(0, 0).text
    assert "已批准" in t.cell(0, 1).text


# ---------------------------------------------------------------------------
# 10. Tag QA 故障注入测试
# ---------------------------------------------------------------------------

def test_contract_v1_tag_qa_fault_injection():
    """丢 tag、未知 tag、超链接 ID 篡改等故障输入必须明确失败。"""
    data = _load_fixture("docx-inline-formatting.docx")
    units, meta = extract_roundtrip_units(data)
    u = units[1]

    # 1. 标签未闭合
    with pytest.raises(TagQAError, match="未闭合或嵌套错误"):
        validate_target_tags(u, "这是<b>损坏的标签")

    # 2. 标签嵌套混乱
    with pytest.raises(TagQAError, match="未闭合或嵌套错误"):
        validate_target_tags(u, "这是<b><i>混乱嵌套</b></i>")

    # 3. 未知标签
    with pytest.raises(TagQAError, match="未知或不合法标签"):
        validate_target_tags(u, "这是<danger>未知标签</danger>")

    # 4. 超链接丢失测试
    data_link = _load_fixture("docx-hyperlinks.docx")
    units_link, _ = extract_roundtrip_units(data_link)
    u_link = next(u for u in units_link if "<link" in u.source)

    with pytest.raises(TagQAError, match="缺失该超链接标签"):
        validate_target_tags(u_link, "超链接被粗暴丢弃的译文。")

    # 5. 超链接 ID 篡改测试
    with pytest.raises(TagQAError, match="超链接 ID 不匹配"):
        validate_target_tags(u_link, '<link id="rId999">篡改 ID</link>')

    # 6. 失效 anchor 测试
    u_bad_anchor = copy.deepcopy(u)
    u_bad_anchor.source_anchor["index"] = 99999
    with pytest.raises(InvalidAnchorError, match="超出范围"):
        write_roundtrip_docx(data, [u_bad_anchor])


# ---------------------------------------------------------------------------
# 11. 人工编辑 / 重新翻译与再次交付结构保真
# ---------------------------------------------------------------------------

def test_contract_v1_human_edit_and_redelivery():
    """人工在审校工作台修改译文后再次交付，原格式结构与未修改条目依然完好受保护。"""
    data = _load_fixture("docx-simple.docx")
    units, meta = extract_roundtrip_units(data)

    # 第一次翻译交付
    for u in units:
        u.target = "第一版译文：" + u.source_plain
    out1_bytes, rep1 = write_roundtrip_docx(data, units)
    assert verify_roundtrip_docx(data, out1_bytes, units, rep1)["passed"] is True

    # 人工在工作台对第 2 段进行编辑修改
    units[2].target = "人工审校修订后的正式译文。"
    out2_bytes, rep2 = write_roundtrip_docx(data, units)
    assert verify_roundtrip_docx(data, out2_bytes, units, rep2)["passed"] is True

    # 验证修订生效且结构依然完好
    doc2 = Document(io.BytesIO(out2_bytes))
    assert "人工审校修订后的正式译文。" in doc2.paragraphs[2].text
    assert doc2.paragraphs[0].style.name == "Title"


# ---------------------------------------------------------------------------
# 12. 接入现有交付门禁与不可变交付清单 (core.freeze_delivery)
# ---------------------------------------------------------------------------

def test_contract_v1_integration_with_freeze_delivery(tmp_path, monkeypatch):
    """复用既有翻译、审校与交付流程；freeze_delivery 产出合规原格式保真成果包。"""
    monkeypatch.setattr(core, "OUTPUT_DIR", tmp_path / "output")
    Path(core.OUTPUT_DIR).mkdir(parents=True, exist_ok=True)

    data = _load_fixture("docx-simple.docx")
    job_id = "job-docx-contract-test"
    core.save_source(job_id, data)

    state = core.new_job_state("docx-simple.docx")
    state["source_paras"] = ["段落1", "段落2", "段落3", "段落4"]
    state["paras"] = ["段落1", "段落2", "段落3", "段落4"]
    state["pairs"] = [
        {"source": "S1", "target": "交付译文1", "reviewed": True, "target_provenance": "human_edit"},
        {"source": "S2", "target": "交付译文2", "reviewed": True, "target_provenance": "model"},
        {"source": "S3", "target": "交付译文3", "reviewed": True, "target_provenance": "human_edit"},
        {"source": "S4", "target": "交付译文4", "reviewed": True, "target_provenance": "human_edit"},
    ]
    state["p1_done"] = True
    state["p2_done"] = True
    state["target_lang"] = "zh-CN"
    # 标记启用 DOCX 原格式保真写回
    state["docx_roundtrip"] = True
    core.save_job_state(job_id, state)

    state["delivery_config"] = core.normalize_delivery_config({
        "deliver_plain_docx": True,
    })
    core.save_job_state(job_id, state)

    # 1. 构建交付资产包
    bundle = core.build_delivery_assets(job_id, state)
    assert "translation.docx" in bundle
    assert "docx_export_report.json" in bundle
    assert "delivery_manifest.json" in bundle

    manifest = json.loads(bundle["delivery_manifest.json"].decode("utf-8"))
    assert "translation.docx" in manifest["generated_assets"]

    # 验证交付物 translation.docx 是真正的原格式保真输出
    trans_docx_bytes = bundle["translation.docx"]
    verify_roundtrip_docx(data, trans_docx_bytes)
    doc = Document(io.BytesIO(trans_docx_bytes))
    assert doc.paragraphs[0].style.name == "Title"
    assert "交付译文1" in doc.paragraphs[0].text

    # 2. 创建不可变交付快照
    snapshot = core.create_delivery_snapshot(job_id, state)
    assert snapshot.get("snapshot_version")
    assets = core.delivery_snapshot_assets(job_id, snapshot["snapshot_version"])
    assert "translation.docx" in assets
    assert "delivery_manifest.json" in assets
