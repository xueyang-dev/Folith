"""DOCX 页眉/页脚“不进入翻译单元”的契约测试（B7-a）。

契约（现状即契约，零行为改动）：

- DOCX/EPUB 的页眉、页脚**不进入翻译单元**——因此不会出现在源文段落里、
  不会进 QA、不进审校覆盖。
- 代价是译文 DOCX/EPUB **不还原**页眉页脚；这一点必须由导入/交付报告的
  ``unsupported`` 条目明确告知用户。
"""

import io

import core
from docx import Document
from docx.enum.section import WD_SECTION


def _docx_with_header_and_footer():
    document = Document()
    document.add_paragraph("Body paragraph one.")
    document.add_paragraph("Body paragraph two.")
    section = document.sections[0]
    section.header.paragraphs[0].text = "Running head text"
    section.footer.paragraphs[0].text = "Footer page marker"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_headers_and_footers_never_become_translation_units():
    data = _docx_with_header_and_footer()

    paragraphs, _warnings, report = core.extract_document_paragraphs_with_report(
        "headers-footers.docx", data)

    # 1) 页眉/页脚不进入源文段落。
    assert paragraphs == ["Body paragraph one.", "Body paragraph two."]
    joined = "\n".join(paragraphs)
    assert "Running head text" not in joined
    assert "Footer page marker" not in joined

    # 2) 但必须在 unsupported 里被点名，且明说译文不保留。
    unsupported = report.get("extracted", {}).get("unsupported") \
        or report.get("unsupported") or []
    kinds = {str(item.get("kind") or "") for item in unsupported}
    assert "header" in kinds
    assert "footer" in kinds
    details = " ".join(str(item.get("detail") or "") for item in unsupported)
    assert "页眉" in details and "页脚" in details
    for item in unsupported:
        if str(item.get("kind") or "") in {"header", "footer"}:
            assert "不保留" in str(item.get("detail") or "")


def test_headers_and_footers_do_not_reach_the_source_paragraph_state():
    data = _docx_with_header_and_footer()

    paragraphs, _warnings, report = core.extract_document_paragraphs_with_report(
        "headers-footers.docx", data)
    state = core.new_job_state("headers-footers.docx")
    state["source_paras"] = list(paragraphs)
    state["paras"] = list(paragraphs)

    assert all("Running head text" not in str(item or "")
               for item in state["source_paras"])
    assert all("Footer page marker" not in str(item or "")
               for item in state["paras"])
    # 页眉页脚不产生翻译单元 ⇒ 不影响审校覆盖度。
    assert report.get("extracted", {}).get("paragraph_count", len(paragraphs)) \
        == len(paragraphs)
