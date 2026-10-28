"""准备页「原文结构」区块回归（#27 的界面层）。

钉住的是四件事：

- **无需 API 也能看**：结构来自本地解析（后台线程 + 缓存），断网/未配置模型时
  照样可读；每个节点都能点开看那部分原文；
- **扫描页不装懂**：无文字层的页如实显示"未识别"，预览解释视觉模型路径；
- **只读**：浏览结构不写任务状态（任务文件逐字节不变），列表没有勾选入口；
- **缓存复用**：同一份源文件第二次请求直接命中缓存，不再排队解析。

界面层测试用 offscreen Qt 平台（与 qt_task_setup_test 同一套约定）。
"""
from __future__ import annotations

import io
import os
import sys
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["FOLITH_QT_NONINTERACTIVE"] = "1"

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from PySide6.QtCore import QPoint, Qt  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

import core  # noqa: E402
from folith.application import commands as cmd  # noqa: E402
from folith.application.settings import SettingsSources  # noqa: E402
from folith.qt import adapter  # noqa: E402
from folith.qt.source_outline_panel import (  # noqa: E402
    SourceOutlineSection, node_item_text, node_preview_text, outline_summary)
from folith.qt import source_outline_panel as panel_module  # noqa: E402

pytestmark = pytest.mark.filterwarnings("ignore::DeprecationWarning")


@pytest.fixture(scope="module")
def qt_app():
    yield QApplication.instance() or QApplication([])


@pytest.fixture()
def data_env(tmp_path, monkeypatch):
    data_dir = tmp_path / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setenv("FOLITH_OUTPUT_DIR", str(data_dir))
    monkeypatch.setattr(core, "OUTPUT_DIR", data_dir, raising=False)
    return data_dir


def _make_docx() -> bytes:
    from docx import Document

    buffer = io.BytesIO()
    document = Document()
    document.add_heading("第一章", level=1)
    document.add_paragraph("第一章的正文内容，长度足够被保留下来。")
    document.add_heading("第二章", level=1)
    document.add_paragraph("第二章的正文内容，长度也足够被保留。")
    document.save(buffer)
    return buffer.getvalue()


def _new_task(data_env) -> str:
    imported = cmd.import_document(
        "sample.docx", _make_docx(), SettingsSources(explicit={}))
    return imported.job_id


def _wait(qt_app):
    panel_module.loader().wait(5000)
    qt_app.processEvents()


def _patch_loader(monkeypatch):
    """测试任务各自独立输出目录，按需给一个**本测试内稳定**的 loader：
    既有同进程缓存不串任务，又不至于让 Qt 把无父对象的后台加载器提前回收。"""
    holder = []

    def factory():
        if not holder:
            holder.append(panel_module.PreviewLoader(
                render=panel_module._render_outline))
        return holder[0]

    monkeypatch.setattr(panel_module, "loader", factory)
    return holder


def test_unreadable_corrections_reported_in_preparation(qt_app, data_env):
    from folith import outline_corrections, source_outline
    job = _new_task(data_env)
    path = core.job_dir(job) / outline_corrections.FILENAME
    path.write_bytes(b'{"entries": [broken')
    section = SourceOutlineSection()
    section._job_id = job
    section._apply_outline(source_outline.parse_source_outline("sample.docx", _make_docx()))
    assert section.node_list.count() > 0
    assert "无法读取结构修正文件" in section.status_label.text()
    assert path.read_bytes() == b'{"entries": [broken'


# ---------------------------------------------------------------- 纯函数

def test_scanned_node_honestly_says_not_recognized():
    node = {"id": "n0000", "title": "第 2 页", "kind": "page", "para_count": 0,
            "scanned": True, "page_start": 2, "page_end": 2, "paragraphs": []}
    assert "未识别" in node_item_text(node)
    preview = node_preview_text(node)
    assert "未识别" in preview and "视觉模型" in preview
    assert node.get("paragraphs") == []  # 绝不假装有原文


def test_text_node_text_and_preview():
    node = {"id": "n0001", "title": "第一章", "kind": "chapter", "para_count": 3,
            "scanned": False, "page_start": 1, "page_end": 2,
            "paragraphs": ["第一段原文。", "第二段原文。"]}
    text = node_item_text(node)
    assert "第一章" in text and "第 1–2 页" in text and "3 段" in text
    assert "第一段原文。" in node_preview_text(node)


def test_outline_summary_counts_and_scanned_case():
    outline = {"nodes": [{"kind": "chapter"}, {"kind": "chapter"}, {"kind": "page"}],
               "total_paragraphs": 42, "scanned": False}
    summary = outline_summary(outline)
    assert "2 章" in summary and "1 页" in summary and "42 段" in summary
    assert "未调用模型" in summary
    scanned = {"nodes": [{"kind": "page"}], "total_paragraphs": 0, "scanned": True,
               "pages": 9}
    assert "视觉模型" in outline_summary(scanned)


def test_preparation_rename_persists_and_is_used_by_scope(data_env, qt_app, monkeypatch):
    from PySide6.QtWidgets import QInputDialog
    job_id = _new_task(data_env)
    _patch_loader(monkeypatch)
    section = SourceOutlineSection()
    view = adapter.task_setup_view(adapter.EngineConfig(), job_id)
    section.load(job_id, "sample.docx", view)
    _wait(qt_app)
    monkeypatch.setattr(QInputDialog, "getText", lambda *args, **kwargs: ("手动名称", True))
    section._edit_outline("rename")
    assert section.outline()["nodes"][0]["title"] == "手动名称"
    reopened = SourceOutlineSection()
    reopened.load(job_id, "sample.docx", view)
    _wait(qt_app)
    assert reopened.outline()["nodes"][0]["title"] == "手动名称"
    scope = adapter.translation_scope(job_id)
    assert scope["groups"][0]["nodes"][0]["label"] == "手动名称"
    section.deleteLater()
    reopened.deleteLater()


# ---------------------------------------------------------------- 区块行为

def test_section_lists_nodes_readonly_and_shows_preview(
        data_env, qt_app, monkeypatch):
    job_id = _new_task(data_env)
    source = adapter.source_path(job_id)
    _patch_loader(monkeypatch)
    section = SourceOutlineSection()
    view = adapter.task_setup_view(adapter.EngineConfig(), job_id)
    state_file = Path(core.job_dir(job_id)) / "state.json"
    before = state_file.read_bytes()
    section.load(job_id, "sample.docx", view)
    _wait(qt_app)

    assert section.node_list.count() == 2  # 第一章 + 第二章
    # 只读：没有勾选入口（不是"默认全选"，是根本没有复选框）。
    assert all(not (section.node_list.item(index).flags()
                    & Qt.ItemFlag.ItemIsUserCheckable)
               for index in range(section.node_list.count()))
    assert "本地读取" in section.section.summary_label.text()

    # 点开一个节点能看到原文（本地解析，与任务状态无关）。
    section.node_list.setCurrentRow(1)
    assert "第二章" in section.preview.toPlainText()

    # 只读浏览不写任务状态：任务文件逐字节不变（没有模型调用，也没有状态变更）。
    after = state_file.read_bytes()
    assert after == before


def test_second_request_hits_cache_without_reparsing(
        data_env, qt_app, monkeypatch):
    job_id = _new_task(data_env)
    holder = _patch_loader(monkeypatch)
    view = adapter.task_setup_view(adapter.EngineConfig(), job_id)

    first = SourceOutlineSection()
    first.load(job_id, "sample.docx", view)
    _wait(qt_app)
    assert first.node_list.count() == 2
    loader = holder[0]
    parsed = [(key, value) for key, value in loader._cache.items()
              if key[0] == job_id]
    assert parsed  # 解析结果按源文件身份进了缓存

    # 同一份源文件再开一次（窗口刷新会重建区块）：同步命中缓存，列表立即就绪。
    second = SourceOutlineSection()
    second.load(job_id, "sample.docx", view)
    qt_app.processEvents()
    assert second.node_list.count() == 2
    assert loader._cache[parsed[0][0]] is parsed[0][1]  # 缓存条目未被重算替换


def test_section_without_source_file_says_so_and_does_not_hang(
        data_env, qt_app, monkeypatch, tmp_path):
    job_id = _new_task(data_env)
    _patch_loader(monkeypatch)
    source = adapter.source_path(job_id)
    source.unlink()
    section = SourceOutlineSection()
    section.load(job_id, "sample.docx",
                 adapter.task_setup_view(adapter.EngineConfig(), job_id))
    qt_app.processEvents()
    assert section.node_list.count() == 0
    assert "没有可解析的结构" in section.section.summary_label.text()
    assert "源文件" in section.status_label.text()


def test_section_is_readonly_after_start(data_env, qt_app, monkeypatch):
    job_id = _new_task(data_env)
    state = core.load_job_state(job_id)
    state["p1_done"] = True
    core.save_job_state(job_id, state)
    _patch_loader(monkeypatch)
    section = SourceOutlineSection()
    section.load(job_id, "sample.docx", adapter.task_setup_view(
        adapter.EngineConfig(), job_id))
    qt_app.processEvents()
    assert not section.isEnabled()
    assert section.node_list.count() == 0  # 已开始不再请求结构
    assert "工作台" in section.status_label.text()


# ---------------------------------------------------------------- 面板集成

def _setup_panel(data_env, qt_app):
    from folith.qt.setup_panel import TaskSetupPanel

    job_id = _new_task(data_env)
    panel = TaskSetupPanel()
    panel.set_choices(adapter.setup_choices())
    view = adapter.task_setup_view(adapter.EngineConfig(), job_id)
    panel.load(view)
    panel.resize(1100, 900)
    panel.show()
    qt_app.processEvents()
    return panel, job_id, view


def test_setup_panel_hosts_the_structure_section(data_env, qt_app, monkeypatch):
    import folith.qt.setup_panel as setup_panel_module
    monkeypatch.setattr(setup_panel_module.SourceOutlineSection, "load",
                        lambda self, *args, **kwargs: None)
    panel, job_id, view = _setup_panel(data_env, qt_app)
    # 区块在语言选择之后、翻译模式之前（先看内容，再做决定）：按面板内纵坐标比。
    source_y = panel.source_section.mapTo(panel, QPoint(0, 0)).y()
    mode_y = panel.mode_selector.mapTo(panel, QPoint(0, 0)).y()
    assert 0 < source_y < mode_y
    assert panel.source_section.section.title_label.text() == "原文结构"


# ---------------------------------------------------------------- 范围入口（#28）

def test_scope_entry_emits_intent_and_wires_through_panel(data_env, qt_app):
    """准备页的「翻译范围…」按钮发出意图信号；面板把它转发给 window。

    #28 的验收要求"开始前排除"有操作入口：入口就放在结构区块里。真正的
    对话框与状态变更在 window/引擎，这里只验证意图链路是通的。
    """
    section = SourceOutlineSection()
    emitted = []
    section.scopeRequested.connect(lambda: emitted.append("section"))
    section.scope_button.click()
    assert emitted == ["section"]

    from folith.qt.setup_panel import TaskSetupPanel
    panel = TaskSetupPanel()
    panel.set_choices(adapter.setup_choices())
    forwarded = []
    panel.scopeRequested.connect(lambda: forwarded.append("panel"))
    panel.source_section.scopeRequested.emit()
    assert forwarded == ["panel"]


def test_scope_entry_disabled_without_source_file(data_env, qt_app, monkeypatch):
    job_id = _new_task(data_env)
    _patch_loader(monkeypatch)
    source = adapter.source_path(job_id)
    source.unlink()
    section = SourceOutlineSection()
    section.load(job_id, "sample.docx",
                 adapter.task_setup_view(adapter.EngineConfig(), job_id))
    qt_app.processEvents()
    assert not section.scope_button.isEnabled(), "源文件不在盘上时没有可依据的结构"


def test_pre_start_scope_via_adapter_returns_outline_tree(data_env, qt_app):
    """适配层：任务开始前 adapter.translation_scope 返回大纲树（available=True）。

    审校探针的场景是"任务 paras=0、范围接口 available=false"——这条断言
    钉住修复后的接口契约。
    """
    job_id = _new_task(data_env)
    state = core.load_job_state(job_id)
    assert not state.get("paras") and not state.get("pairs")
    core.ensure_source_outline(job_id, state)
    scope = adapter.translation_scope(job_id)
    assert scope["available"] is True
    assert scope["pre_start"] is True
    chapter = next(group for group in scope["groups"] if group["kind"] == "chapter")
    assert [node["label"] for node in chapter["nodes"]] == ["第一章", "第二章"]
    # 排除走范围键（outline:…），结果记成计划，账本保持为空。
    result = adapter.exclude_scope(job_id, [chapter["nodes"][1]["key"]])
    assert result["planned"] == 1 and result["excluded"] == 0
    assert core.load_job_state(job_id).get("excluded_segments") in (None, [])
